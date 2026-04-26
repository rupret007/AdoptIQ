"""Round 16 / Phase 1 -- cross-format file-level consistency tests.

The Round-15 design wired both the Excel ``Summary`` sheet and the Word
"Executive Summary" table through ``report_export_styling.build_summary_rows``
so the two formats agree by construction.  But the only existing parity
tests work at the ``canonical_metrics`` layer
(`tests/test_cross_report_parity.py`) -- nothing actually generates a
``.docx`` and ``.xlsx`` from the same fixture and compares the strings the
human reader will see.

These tests close that gap.  They build small deterministic in-memory
frames, drive the same fixture through both writers, and assert that the
displayed KPI strings match exactly.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest

import report_export_styling as styling

try:
    import docx  # python-docx
except Exception:  # pragma: no cover - dep missing only in a degraded env
    docx = None  # type: ignore[assignment]

try:
    import report_word_styling as word_styling
except Exception:  # pragma: no cover
    word_styling = None  # type: ignore[assignment]


_REPO_ROOT = Path(__file__).resolve().parent.parent


def _read(path: str) -> str:
    return (_REPO_ROOT / path).read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Deterministic, hand-authored fixture frames.
#
# The frame shapes below match the *minimum* canonical_metrics needs to
# return non-trivial counts: ``count_customers``, ``count_total_barriers``,
# ``count_critical_barriers``, ``count_open_barriers``, ``count_total_tac``,
# ``count_p1``, ``count_open_tac``, ``count_escalated``, ``count_bems``.
# ---------------------------------------------------------------------------


def _build_fixture_frames() -> dict[str, Any]:
    ab_df = pd.DataFrame(
        {
            "ACCOUNT_NAME": ["Acme Corp", "Beta Inc", "Gamma Ltd", "Acme Corp"],
            "STATUS": ["Open", "Open", "Closed", "Open"],
            "PRIORITY": ["Critical", "High", "Medium", "Critical"],
        }
    )
    csone_df = pd.DataFrame(
        {
            "Customer Name": ["Acme Corp", "Beta Inc", "Beta Inc", "Gamma Ltd", "Delta Co"],
            "Severity": ["1", "2", "1", "3", "1"],
            "Status": ["Open", "Open", "Closed", "Open", "Open"],
            "TAC Case Number": ["T1", "T2", "T3", "T4", "T5"],
            "Sub Case Type": [
                "Break/Fix",
                "Configuration",
                "Provisioning",
                "Configuration",
                "Break/Fix",
            ],
            "Escalated": ["Y", "N", "N", "N", "Y"],
        }
    )
    cs_pulse = pd.DataFrame(
        {
            "ACCOUNT_NAME": ["Acme Corp", "Beta Inc", "Gamma Ltd", "Epsilon Co"],
            "HEALTH_SCORE": [3.0, 5.0, 8.0, 9.0],
        }
    )
    ext_bugs = pd.DataFrame({"id": ["b1", "b2"], "title": ["bug1", "bug2"]})
    ext_incidents = pd.DataFrame({"id": ["i1", "i2", "i3"]})

    return {
        "sheets": {
            "AB_Detail_All": ab_df,
            "CSOne_Detail_All": csone_df,
            "External_Bugs": ext_bugs,
            "External_Incidents": ext_incidents,
        },
        "csconsole_data": {"customer_pulse": cs_pulse},
    }


def _kpi_row_dict(rows: list[tuple[str, str]]) -> dict[str, str]:
    """Convert ``[(label, value), ...]`` to ``{label: value}`` for cross-
    format string-equality assertions."""

    return {str(label): str(value) for (label, value) in rows}


# ---------------------------------------------------------------------------
# Phase 1.1 -- in-memory parity: build_summary_rows == build_executive_summary_rows.
# ---------------------------------------------------------------------------


def test_phase_1_1_excel_and_word_summary_rows_are_identical_in_memory():
    """Round 16 / Phase 1.1 -- the Excel ``build_summary_rows`` and the
    Word ``build_executive_summary_rows`` must return identical
    ``(label, value)`` tuples for the same fixture, in the same order.

    Both helpers delegate to canonical_metrics under the hood, so any
    drift between them indicates a regression in the Word/Excel parity
    layer (R15 / Phase 3.7).
    """

    if word_styling is None:
        pytest.skip("report_word_styling not importable in this environment")

    fixture = _build_fixture_frames()

    excel_rows = styling.build_summary_rows(
        fixture["sheets"],
        fixture["csconsole_data"],
        manager="Test Manager",
        tech="Contact Center",
        days=90,
        generated_at_utc_iso_z="2026-04-26T00:00:00Z",
    )
    word_rows = word_styling.build_executive_summary_rows(
        fixture["sheets"],
        fixture["csconsole_data"],
        manager="Test Manager",
        tech="Contact Center",
        days=90,
        generated_at_utc_iso_z="2026-04-26T00:00:00Z",
    )

    assert excel_rows, "Excel summary rows must be non-empty for this fixture"
    assert word_rows, "Word summary rows must be non-empty for this fixture"
    excel_tuples = [tuple(r) for r in excel_rows]
    word_tuples = [tuple(r) for r in word_rows]
    assert excel_tuples == word_tuples, (
        "Round 16 / Phase 1.1 cross-format drift: Excel build_summary_rows "
        "and Word build_executive_summary_rows returned different rows for "
        "the same fixture.\n"
        f"Excel: {excel_tuples}\nWord:  {word_tuples}"
    )


def test_phase_1_1_summary_rows_include_real_canonical_values():
    """Round 16 / Phase 1.1 corollary -- the parity test above is only
    meaningful if at least one KPI surfaces a non-placeholder value.
    Pin that the customer count is *not* ``--`` for the canonical
    fixture so a future regression that silently breaks the
    canonical-metrics call chain (the R15-015 class of bug) is caught
    here in addition to in the Round-15 correctness suite.
    """

    fixture = _build_fixture_frames()
    rows = styling.build_summary_rows(
        fixture["sheets"], fixture["csconsole_data"], manager="m", tech="t", days=90
    )
    label_to_value = _kpi_row_dict(rows)
    assert "Customers in portfolio" in label_to_value
    assert label_to_value["Customers in portfolio"] != "--", (
        "Round 16 / Phase 1.1 canary: customer count rendered as `--` for a "
        "fixture that clearly carries customers; canonical_metrics call "
        "chain may have regressed (cf. R15-015)."
    )


# ---------------------------------------------------------------------------
# Phase 1.2 -- file-level parity: write to xlsx + docx and read them back.
# ---------------------------------------------------------------------------


def _xlsx_summary_dict(xlsx_path: Path) -> dict[str, str]:
    """Read a workbook, find the ``Summary`` sheet, and return a
    label->value dict from the (label, value) pairs."""

    wb = openpyxl.load_workbook(xlsx_path, data_only=False, read_only=True)
    try:
        assert "Summary" in wb.sheetnames, (
            f"Expected `Summary` sheet, got {wb.sheetnames}"
        )
        sheet = wb["Summary"]
        out: dict[str, str] = {}
        for row in sheet.iter_rows(values_only=True):
            if not row:
                continue
            if len(row) < 2:
                continue
            label, value = row[0], row[1]
            if label is None or str(label).strip() == "":
                continue
            label_s = str(label).strip()
            value_s = "" if value is None else str(value).strip()
            # Skip the section heading / decorative cells.
            if not value_s:
                continue
            out[label_s] = value_s
        return out
    finally:
        wb.close()


def _docx_summary_dict(docx_path: Path) -> dict[str, str]:
    """Read a docx, find the executive-summary 2-col table (whose
    header row is exactly ``Metric|Value``), and return a label->value
    dict."""

    assert docx is not None, "python-docx must be installed for this test"
    document = docx.Document(str(docx_path))
    for table in document.tables:
        if not table.rows:
            continue
        header_cells = [c.text.strip() for c in table.rows[0].cells]
        if len(header_cells) >= 2 and header_cells[:2] == ["Metric", "Value"]:
            out: dict[str, str] = {}
            for row in table.rows[1:]:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) < 2:
                    continue
                if not cells[0]:
                    continue
                out[cells[0]] = cells[1]
            return out
    raise AssertionError(
        "Round 16 / Phase 1.2: no `Metric/Value` executive-summary table "
        "found in docx."
    )


def test_phase_1_2_summary_sheet_and_docx_table_render_same_strings(tmp_path):
    """Round 16 / Phase 1.2 -- write the Round-15 Summary sheet to xlsx
    and the executive-summary table to a docx using the SAME fixture,
    then read both files back and assert the displayed strings agree
    label-for-label and value-for-value.

    This is the file-level cross-format consistency test.  It catches
    drift that the in-memory parity test (1.1) cannot -- e.g. a
    formatter that secretly truncates strings, or an xlsx writer
    that coerces strings to numbers and loses leading zeros.
    """

    if word_styling is None:
        pytest.skip("report_word_styling not importable in this environment")
    if docx is None:
        pytest.skip("python-docx not installed in this environment")

    fixture = _build_fixture_frames()

    # Write xlsx with the Summary sheet.
    xlsx_path = tmp_path / "round16_cross_format.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="xlsxwriter") as writer:
        # Empty placeholder sheet so the workbook is valid.
        pd.DataFrame({"placeholder": [None]}).to_excel(
            writer, sheet_name="Placeholder", index=False
        )
        wrote = styling.write_summary_sheet(
            writer,
            fixture["sheets"],
            fixture["csconsole_data"],
            manager="Test Manager",
            tech="Contact Center",
            days=90,
            generated_at_utc_iso_z="2026-04-26T00:00:00Z",
        )
    assert wrote is True, "write_summary_sheet should report success on xlsxwriter"

    # Write docx with the executive-summary table.
    docx_path = tmp_path / "round16_cross_format.docx"
    document = docx.Document()
    rows = word_styling.build_executive_summary_rows(
        fixture["sheets"],
        fixture["csconsole_data"],
        manager="Test Manager",
        tech="Contact Center",
        days=90,
        generated_at_utc_iso_z="2026-04-26T00:00:00Z",
    )
    word_styling.add_executive_summary_table(document, rows)
    document.save(str(docx_path))

    xlsx_dict = _xlsx_summary_dict(xlsx_path)
    docx_dict = _docx_summary_dict(docx_path)

    # Both must report at least the canonical KPIs we built fixtures for.
    expected_labels = {
        "Manager scope",
        "Technology scope",
        "Window (days)",
        "Customers in portfolio",
        "Adoption barriers (total)",
        "TAC cases (total)",
        "TAC cases (P1)",
        "TAC cases (open)",
        "External incidents (rows)",
    }
    missing_xlsx = expected_labels - xlsx_dict.keys()
    missing_docx = expected_labels - docx_dict.keys()
    assert not missing_xlsx, f"xlsx Summary missing labels: {missing_xlsx}"
    assert not missing_docx, f"docx executive summary missing labels: {missing_docx}"

    # The intersection of labels must be non-empty AND every shared label
    # must carry the same displayed value in both formats.
    shared = set(xlsx_dict.keys()) & set(docx_dict.keys())
    assert shared, "no shared labels between xlsx Summary and docx exec table"
    diffs = {
        label: (xlsx_dict[label], docx_dict[label])
        for label in shared
        if xlsx_dict[label] != docx_dict[label]
    }
    assert not diffs, (
        "Round 16 / Phase 1.2 cross-format drift: the Excel Summary tab and "
        "the Word executive-summary table render different strings for the "
        f"same fixture.\nDrifted labels (xlsx, docx): {diffs}"
    )


# ---------------------------------------------------------------------------
# Phase 1.3 -- column curation parity: the curated sheet schema for
# ``AB_Detail_All`` returned to the Word renderer (via column metadata)
# must match what the Excel writer ships.
# ---------------------------------------------------------------------------


def test_phase_1_3_curated_columns_module_consistent():
    """Round 16 / Phase 1.3 -- ``report_export_schema.CURATED_COLUMNS``
    is the single source of truth for which columns ship in each
    workbook tab.  Pin two structural invariants:

    1. Every curated column list is non-empty.
    2. No curated list contains a name that ``is_internal_column``
       would reject (Round 15 / Phase 1 catches drift here in a
       different test, but Round 16 re-asserts it as part of the
       cross-format invariant set: the Word tables shouldn't render
       columns the Excel side has explicitly denied).
    """

    import report_export_schema as schema

    assert schema.CURATED_COLUMNS, "CURATED_COLUMNS must not be empty"
    for sheet_name, cols in schema.CURATED_COLUMNS.items():
        assert cols, f"curated column list for {sheet_name!r} is empty"
        for col in cols:
            assert not schema.is_internal_column(col), (
                f"curated column {col!r} for {sheet_name!r} would be rejected "
                "by is_internal_column()"
            )


# ---------------------------------------------------------------------------
# Phase 1.x -- marker presence.
# ---------------------------------------------------------------------------


def test_phase_1_marker_in_test_module():
    """Self-pin: this module documents Phase 1.x markers."""
    contents = _read("tests/test_round16_cross_format_consistency.py")
    assert "Round 16 / Phase 1.1" in contents
    assert "Round 16 / Phase 1.2" in contents
    assert "Round 16 / Phase 1.3" in contents
