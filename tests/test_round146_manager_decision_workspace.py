"""Round 146 Manager Decision Workspace domain contracts."""

from __future__ import annotations
from source_shape_utils import assert_in_source

from pathlib import Path

import pytest
from openpyxl import Workbook

import manager_decision_workspace as mdw


MANAGERS = ["Manager One", "Manager Two"]
TECHNOLOGIES = ["All", "All Contact Center"]


def _selection(**overrides):
    values = {
        "report_type": "leader",
        "manager": "Manager One",
        "technology": "All Contact Center",
        "days": 90,
        "scope_type": "team",
        "scope_value": "",
        "subscription_id": "",
        "allowed_managers": MANAGERS,
        "allowed_technologies": TECHNOLOGIES,
    }
    values.update(overrides)
    return mdw.validate_workspace_selection(**values)


def _write_snapshot(
    path: Path,
    *,
    fingerprint: str,
    open_aps: int,
    source_state: str = "available",
    ap_rows: list[list[object]] | None = None,
    formula: bool = False,
    scope_type: str = "team",
    scope_value: str = "",
    source_states: dict[str, str] | None = None,
    insight_rows: list[list[object]] | None = None,
    evidence_rows: list[list[object]] | None = None,
) -> Path:
    declared_source_states = {
        "Subscriptions": "available",
        "Action_Plans": source_state,
        "Adoption_Barriers": "available",
        "Customer_Pulse": "available",
        "TAC_Cases": "available",
        "Success_Priorities": "available",
    }
    declared_source_states.update(source_states or {})
    workbook = Workbook()
    info = workbook.active
    info.title = "Report_Info"
    info.append(["Item", "Value", "Detail"])
    for row in (
        ["Report_Type", "leader", ""],
        ["Manager", "Manager One", ""],
        ["Technology", "All Contact Center", ""],
        ["Scope_Type", scope_type, ""],
        ["Scope_Value", scope_value, ""],
        ["Days", 90, ""],
        ["Data_As_Of_UTC", "2026-08-03T12:00:00Z", ""],
        ["Fact_Contract_SHA256", fingerprint, ""],
    ):
        info.append(row)
    for source, state in declared_source_states.items():
        info.append([f"Source_State:{source}", state, ""])
    lineage = workbook.create_sheet("Metric_Lineage")
    lineage.append([
        "Metric_Key", "Display_Label", "Metric_Value", "Unit",
        "Source_State", "Source_Sheet", "Caveat",
    ])
    lineage.append([
        "kpi.action_plans_open", "Open Action Plans", open_aps,
        "records", source_state, "Action_Plans", "",
    ])
    lineage.append([
        "kpi.customers", "Customers", 3,
        "records", "available", "Subscriptions", "",
    ])
    for row in insight_rows or []:
        lineage.append(row)
    if formula:
        lineage.append(["kpi.test_formula", "Formula", "=1+1", "records", "available", "Report_Info", ""])
    chart = workbook.create_sheet("Chart_Data")
    chart.append([
        "Chart_ID", "Metric_Key", "Display_Label", "Series", "Category",
        "Period_Start", "Value", "Unit", "Source_State",
    ])
    chart.append([
        "action_plan_status_aging", "chart.action_plan_status.open",
        "Action Plan status and aging — Open", "Status", "Open", "",
        open_aps, "records", source_state,
    ])
    actions = workbook.create_sheet("Action_Plans")
    actions.append([
        "ID", "BU_NAME", "SUBJECT_C", "STATUS_C", "OWNER_NAME_C",
        "DUE_DATE_C", "PRIORITY_C", "NEXT_ACTION_C", "AdoptIQ_Is_Overdue",
    ])
    for row in ap_rows or []:
        actions.append(row)
    accounts = workbook.create_sheet("Account_Summary")
    risk_state = (
        "available"
        if all(
            declared_source_states[source] in {"available", "zero"}
            for source in (
                "Subscriptions",
                "Action_Plans",
                "Adoption_Barriers",
                "Customer_Pulse",
                "TAC_Cases",
            )
        )
        else "partial"
    )
    accounts.append([
        "Account", "Risk_Band", "Risk_Score_0_100", "Open_AP", "Overdue_AP",
        "Critical_High_Barriers", "TAC_Cases", "Risk_Band_Source_State",
        "Risk_Score_0_100_Source_State", "Open_AP_Source_State",
        "Overdue_AP_Source_State", "Critical_High_Barriers_Source_State",
        "TAC_Cases_Source_State",
    ])
    accounts.append([
        "Acme", "HIGH", 78, open_aps, 1, 3, 2, risk_state, risk_state,
        declared_source_states["Action_Plans"],
        declared_source_states["Action_Plans"],
        declared_source_states["Adoption_Barriers"],
        declared_source_states["TAC_Cases"],
    ])
    if evidence_rows is not None:
        evidence = workbook.create_sheet("Evidence_Links")
        evidence.append([
            "Evidence_Key", "Evidence_Type", "Display_Label", "Source_State",
            "Metric_Value", "Unit", "Source_Sheet", "Source_Row_Number",
        ])
        for row in evidence_rows:
            evidence.append(row)
    workbook.save(path)
    workbook.close()
    return path


def test_workspace_selection_supports_leader_drilldown_and_rejects_missing_value():
    selection = _selection(scope_type="member", scope_value="member@example.invalid")
    assert selection["scope_type"] == "member"
    assert selection["scope_value"] == "member@example.invalid"
    with pytest.raises(ValueError, match="individual member"):
        _selection(scope_type="member", scope_value="")


@pytest.mark.parametrize(
    ("report_type", "scope_type"),
    [("compact", "member"), ("renewal", "team"), ("subscription", "customer")],
)
def test_workspace_selection_rejects_unsupported_report_scope(report_type, scope_type):
    with pytest.raises(ValueError, match="does not support"):
        _selection(report_type=report_type, scope_type=scope_type, scope_value="Acme")


def test_scope_preview_is_decision_oriented_and_fixture_honest():
    preview = mdw.build_scope_preview(
        _selection(),
        members=[{"email": "one@example.invalid"}, {"email": "two@example.invalid"}],
        customers=[{"value": "Acme"}, {"value": "Beta"}],
        source_mode="local_fixture_guarded",
        source_state="partial",
        warning="Customer Pulse is stale.",
    )
    assert preview["member_count"] == 2
    assert preview["customer_count"] == 2
    assert preview["scope_label"] == "Manager One team"
    assert preview["live_validation_performed"] is False
    assert any("no live validation" in item.lower() for item in preview["limitations"])
    assert "Action Plans" in preview["expected_sources"]


def test_load_workbook_snapshot_projects_metrics_sources_action_plans_and_risk(tmp_path):
    path = _write_snapshot(
        tmp_path / "source.xlsx",
        fingerprint="abc123",
        open_aps=1,
        ap_rows=[
            ["AP-1", "Acme", "Close adoption gap", "On Track", "Owner A", "2026-08-10", "High", "Confirm training", False],
            ["", "Beta", "", "Overdue", "", "2026-07-01", "Medium", "Escalate", True],
        ],
    )
    snapshot = mdw.load_workbook_snapshot(path)
    assert snapshot["fact_fingerprint"] == "abc123"
    assert snapshot["formula_cells"] == 0
    assert snapshot["source_states"] == {
        "Action_Plans": "available",
        "Adoption_Barriers": "available",
        "Customer_Pulse": "available",
        "Subscriptions": "available",
        "Success_Priorities": "available",
        "TAC_Cases": "available",
    }
    assert snapshot["metrics"][0]["metric_key"] == "kpi.action_plans_open"
    assert snapshot["canonical_snapshot"] is True
    assert snapshot["charts"][0]["chart_key"] == "action_plan_status_aging"
    assert snapshot["charts"][0]["series"][0]["value"] == 1
    assert snapshot["action_plan_count"] == 2
    assert snapshot["missing_action_plan_id_count"] == 1
    assert snapshot["action_plans"][0]["is_overdue"] is True
    assert snapshot["action_plans"][0]["title"] == "Title unavailable"
    assert snapshot["accounts"][0]["risk_score_0_100"] == 78
    assert snapshot["accounts"][0]["critical_high_barriers"] == 3
    assert set(snapshot["accounts"][0]["field_states"].values()) == {"available"}


def test_canonical_snapshot_projects_only_allowlisted_frozen_insights(tmp_path):
    claim = "Support  themes:\nConfiguration appeared in 3 exact TAC case records."
    path = _write_snapshot(
        tmp_path / "insights.xlsx",
        fingerprint="a" * 64,
        open_aps=1,
        insight_rows=[
            [
                "insight.support_themes",
                "Support themes (TAC)",
                claim,
                "frozen claim",
                "available",
                "TAC_Cases",
                "",
            ],
            [
                "insight.unreviewed_new_claim",
                "Unreviewed",
                "This must not cross the web boundary.",
                "frozen claim",
                "available",
                "TAC_Cases",
                "",
            ],
        ],
        evidence_rows=[
            [
                "insight.support_themes",
                "insight",
                "Support themes (TAC)",
                "available",
                claim,
                "frozen claim",
                "TAC_Cases",
                2,
            ],
        ],
    )

    snapshot = mdw.load_workbook_snapshot(path)

    assert snapshot["decision_insight_integrity"] == "verified_shape"
    assert snapshot["decision_insights"] == [
        {
            "insight_key": "insight.support_themes",
            "evidence_key": "insight.support_themes",
            "evidence_count": 1,
            "label": "Support themes (TAC)",
            "claim": claim,
            "caveat": "",
            "source_state": "available",
            "source_sheets": ["TAC_Cases"],
            "provenance": "canonical Metric_Lineage",
        }
    ]


def test_mixed_available_and_zero_insight_evidence_remains_complete(tmp_path):
    claim = "Momentum: TAC activity increased while Customer Pulse returned zero records."
    snapshot = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "mixed-complete-insight.xlsx",
            fingerprint="9" * 64,
            open_aps=1,
            source_states={"Customer_Pulse": "zero"},
            insight_rows=[
                [
                    "insight.window_momentum",
                    "Momentum within this window",
                    claim,
                    "frozen claim",
                    "available",
                    "TAC_Cases; Customer_Pulse",
                    "",
                ],
            ],
            evidence_rows=[
                [
                    "insight.window_momentum",
                    "insight",
                    "Momentum within this window",
                    "available",
                    claim,
                    "frozen claim",
                    "TAC_Cases",
                    2,
                ],
                [
                    "insight.window_momentum",
                    "insight",
                    "Momentum within this window",
                    "zero",
                    claim,
                    "frozen claim",
                    "Customer_Pulse",
                    None,
                ],
            ],
        )
    )

    assert snapshot["decision_insight_integrity"] == "verified_shape"
    assert snapshot["decision_insights"][0]["claim"] == claim
    assert snapshot["decision_insights"][0]["source_state"] == "available"
    assert snapshot["decision_insights"][0]["evidence_count"] == 1


def test_single_source_support_insight_preserves_stale_state(tmp_path):
    claim = "Support themes: source is stale, so this exact claim remains internal."
    snapshot = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "stale-support-insight.xlsx",
            fingerprint="7" * 64,
            open_aps=1,
            source_states={"TAC_Cases": "stale"},
            insight_rows=[
                ["insight.support_themes", "Support themes", claim, "frozen claim", "stale", "TAC_Cases", ""],
            ],
            evidence_rows=[
                ["insight.support_themes", "insight", "Support themes", "stale", claim, "frozen claim", "TAC_Cases", 2],
            ],
        )
    )

    assert snapshot["decision_insight_integrity"] == "verified_shape"
    assert snapshot["decision_insights"][0]["source_state"] == "stale"


def test_real_canonical_source_data_round_trips_frozen_insight_and_evidence(tmp_path):
    from tests.test_round157_reporting_ask_ai import _facts_with_tac

    delivery, facts = _facts_with_tac([
        {
            "SR Number": "TAC-1",
            "BU_NAME": "Acme",
            "Severity": "P1",
            "Case Status": "Open",
            "Date/Time Opened": "2026-07-30",
            "Tech.": "Webex Calling",
        },
        {
            "SR Number": "TAC-2",
            "BU_NAME": "Acme",
            "Severity": "P3",
            "Case Status": "Open",
            "Date/Time Opened": "2026-07-28",
            "Tech.": "Webex Calling",
        },
    ])
    sheets = delivery.build_source_data_sheets(facts)
    path = tmp_path / "canonical-insights.xlsx"
    delivery.write_source_data_workbook(path, sheets)

    snapshot = mdw.load_workbook_snapshot(path)
    expected = facts["decision_insights"]["support_themes"]
    projected = {
        item["insight_key"]: item for item in snapshot["decision_insights"]
    }

    assert snapshot["decision_insight_integrity"] == "verified_shape"
    assert projected["insight.support_themes"]["claim"] == expected["paragraph_text"]
    assert projected["insight.support_themes"]["evidence_key"] == "insight.support_themes"
    assert projected["insight.support_themes"]["evidence_count"] > 0

    view = mdw.snapshot_from_status(
        {
            "analysis_id": "canonical-insights",
            "status": "completed",
            "report_type": "leader",
            "scope_type": "team",
            "excel_hash": snapshot["workbook_sha256"],
        },
        snapshot,
    )
    assert view["decision_insights"][0]["claim"] == expected["paragraph_text"]
    assert view["customer_share_readiness"]["customer_shareable"] is False
    assert view["customer_share_readiness"]["live_validation_performed"] is False


def test_real_predictive_tac_block_keeps_unavailable_claim_and_other_insights(tmp_path):
    import decision_report_delivery as delivery
    from tests.test_decision_insight_contract import _predictive_team_data
    from tests.test_round160_predictive_engine import AS_OF

    team_data = _predictive_team_data()
    tac = team_data["Alex Rivera"]["tac_cases"].iloc[0:0].copy()
    tac.attrs["source_unavailable"] = True
    tac.attrs["source_unavailable_detail"] = "source unavailable in local replay"
    team_data["Alex Rivera"]["tac_cases"] = tac
    facts = delivery.build_report_facts(
        team_data,
        report_type="Comprehensive",
        scope_type="team",
        scope_value="Alex Rivera's Team",
        manager_name="Alex Rivera",
        days=90,
        as_of=AS_OF,
        data_as_of_utc=AS_OF.isoformat(),
        data_as_of_state="partial",
    )
    sheets = delivery.build_source_data_sheets(facts)
    path = tmp_path / "predictive-tac-block.xlsx"
    delivery.write_source_data_workbook(path, sheets)

    snapshot = mdw.load_workbook_snapshot(path)
    projected = {
        item["insight_key"]: item for item in snapshot["decision_insights"]
    }

    assert snapshot["decision_insight_integrity"] == "verified_shape"
    assert projected["insight.predictive_outlook_30d"]["source_state"] == "unavailable"
    assert "Forecast unavailable" in projected["insight.predictive_outlook_30d"]["claim"]
    assert "insight.window_momentum" in projected


def test_duplicate_canonical_insight_key_fails_closed(tmp_path):
    snapshot = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "duplicate-insight.xlsx",
            fingerprint="b" * 64,
            open_aps=1,
            insight_rows=[
                ["insight.window_momentum", "Momentum", "Claim one.", "frozen claim", "available", "TAC_Cases", ""],
                ["insight.window_momentum", "Momentum", "Claim two.", "frozen claim", "available", "TAC_Cases", ""],
            ],
            evidence_rows=[],
        )
    )

    assert snapshot["decision_insights"] == []
    assert snapshot["decision_insight_integrity"] == "duplicate_keys"


def test_overlong_frozen_insight_is_withheld_instead_of_truncated(tmp_path):
    claim = "X" * 4_001
    snapshot = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "overlong-insight.xlsx",
            fingerprint="8" * 64,
            open_aps=1,
            insight_rows=[
                ["insight.support_themes", "Support themes", claim, "frozen claim", "available", "TAC_Cases", ""],
            ],
            evidence_rows=[
                ["insight.support_themes", "insight", "Support themes", "available", claim, "frozen claim", "TAC_Cases", 2],
            ],
        )
    )

    assert snapshot["decision_insights"] == []
    assert snapshot["decision_insight_integrity"] == "evidence_mismatch"


def test_decision_view_withholds_insight_when_lineage_and_evidence_disagree(tmp_path):
    workbook = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "unverified-insight.xlsx",
            fingerprint="c" * 64,
            open_aps=1,
            scope_type="customer",
            scope_value="Acme",
            insight_rows=[
                [
                    "insight.support_operating_health",
                    "Support operating health",
                    "Support operating health: median closure time was 2 days.",
                    "frozen claim",
                    "available",
                    "TAC_Cases",
                    "",
                ],
            ],
            evidence_rows=[
                [
                    "insight.support_operating_health",
                    "insight",
                    "Support operating health",
                    "available",
                    2,
                    "days",
                    "TAC_Cases",
                    2,
                ],
            ],
        )
    )
    assert workbook["decision_insights"] == []
    assert workbook["decision_insight_integrity"] == "evidence_mismatch"

    view = mdw.snapshot_from_status(
        {
            "analysis_id": "customer-insight",
            "status": "completed",
            "report_type": "leader",
            "scope_type": "customer",
            "scope_value": "Acme",
            "excel_hash": workbook["workbook_sha256"],
        },
        workbook,
    )

    assert view["decision_insights"] == []
    assert view["decision_insight_integrity"] == "evidence_mismatch"
    assert any("withheld" in warning for warning in view["source_warnings"])
    assert view["customer_share_readiness"]["customer_shareable"] is False
    assert view["customer_share_readiness"]["production_accuracy_claimed"] is False
    assert view["customer_share_readiness"]["release_ready"] is False


def test_customer_share_readiness_keeps_offline_preview_fail_closed():
    snapshot = {
        "status": "completed",
        "scope_type": "customer",
        "canonical_snapshot": True,
        "fact_fingerprint": "c" * 64,
        "workbook_sha256": "d" * 64,
        "persisted_workbook_hash_verified": True,
        "evidence_integrity_verified": True,
        "evidence_available": True,
        "formula_cells": 0,
        "data_as_of_state": "available",
        "source_states": {"TAC_Cases": "available"},
        "source_warnings": [],
        "decision_insight_integrity": "verified_shape",
        "decision_insights": [{
            "insight_key": "insight.support_themes",
            "evidence_key": "insight.support_themes",
            "evidence_count": 2,
            "claim": "Support themes: 2 exact TAC case records.",
            "source_state": "available",
            "source_sheets": ["TAC_Cases"],
        }],
    }

    readiness = mdw.assess_customer_share_readiness(snapshot)

    assert readiness["state"] == "internal_preview"
    assert readiness["customer_shareable"] is False
    assert readiness["live_validation_performed"] is False
    assert readiness["production_accuracy_claimed"] is False
    assert readiness["release_ready"] is False
    assert any("receipt" in reason for reason in readiness["reasons"])


def test_untrusted_truthy_receipt_cannot_enable_customer_sharing():
    snapshot = {
        "status": "completed",
        "scope_type": "subscription",
        "canonical_snapshot": True,
        "fact_fingerprint": "e" * 64,
        "workbook_sha256": "f" * 64,
        "persisted_workbook_hash_verified": True,
        "evidence_integrity_verified": True,
        "evidence_available": True,
        "formula_cells": 0,
        "data_as_of_state": "zero",
        "source_states": {"TAC_Cases": "zero"},
        "source_warnings": [],
        "decision_insight_integrity": "verified_shape",
        "decision_insights": [{
            "insight_key": "insight.predictive_outlook_30d",
            "evidence_key": "insight.predictive_outlook_30d",
            "evidence_count": 0,
            "claim": "Predictive outlook: no elevated customer signal was asserted.",
            "source_state": "zero",
            "source_sheets": ["TAC_Cases"],
        }],
        "customer_share_validation_receipt": {
            "mode_executed": "live",
            "fact_fingerprint": "e" * 64,
            "workbook_sha256": "f" * 64,
            "live_validation_performed": True,
            "production_accuracy_claimed": True,
            "release_ready": True,
            "manual_source_reconciliation_complete": True,
            "owner_customer_share_approved": True,
        },
    }

    readiness = mdw.assess_customer_share_readiness(snapshot)
    assert readiness["customer_shareable"] is False
    assert readiness["live_validation_performed"] is False
    assert readiness["production_accuracy_claimed"] is False
    assert readiness["release_ready"] is False
    assert any("not enabled" in reason for reason in readiness["reasons"])


@pytest.mark.parametrize(
    ("degraded_source", "degraded_state", "withheld_fields", "preserved_fields"),
    [
        (
            "Action_Plans",
            "partial",
            {"risk_band", "risk_score_0_100", "open_action_plans", "overdue_action_plans"},
            {"critical_high_barriers": 3, "tac_cases": 2},
        ),
        (
            "TAC_Cases",
            "failed",
            {"risk_band", "risk_score_0_100", "tac_cases"},
            {"open_action_plans": 2, "overdue_action_plans": 1, "critical_high_barriers": 3},
        ),
        (
            "Adoption_Barriers",
            "stale",
            {"risk_band", "risk_score_0_100", "critical_high_barriers"},
            {"open_action_plans": 2, "overdue_action_plans": 1, "tac_cases": 2},
        ),
        (
            "Customer_Pulse",
            "partial",
            {"risk_band", "risk_score_0_100"},
            {
                "open_action_plans": 2,
                "overdue_action_plans": 1,
                "critical_high_barriers": 3,
                "tac_cases": 2,
            },
        ),
    ],
)
def test_canonical_account_projection_withholds_only_degraded_derived_fields(
    tmp_path,
    degraded_source,
    degraded_state,
    withheld_fields,
    preserved_fields,
):
    snapshot = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / f"{degraded_source}.xlsx",
            fingerprint=f"degraded-{degraded_source}",
            open_aps=2,
            source_states={degraded_source: degraded_state},
        )
    )
    account = snapshot["accounts"][0]
    for field in withheld_fields:
        assert account[field] in {None, ""}
        assert account["field_states"][field] not in {"available", "zero"}
    for field, expected in preserved_fields.items():
        assert account[field] == expected
        assert account["field_states"][field] == "available"


def test_load_workbook_snapshot_discloses_formula_cells(tmp_path):
    snapshot = mdw.load_workbook_snapshot(
        _write_snapshot(tmp_path / "formula.xlsx", fingerprint="formula", open_aps=0, formula=True)
    )
    assert snapshot["formula_cells"] == 1
    assert snapshot["warnings"]


def test_snapshot_from_status_builds_decision_cards_and_scope_binding(tmp_path):
    workbook = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "view.xlsx",
            fingerprint="view",
            open_aps=2,
            scope_value="Manager One team",
        )
    )
    snapshot = mdw.snapshot_from_status(
        {
            "analysis_id": "Leader_Manager_90d",
            "status": "completed",
            "scope_type": "team",
            "scope_value": "",
            "word_report": "/safe/report.docx",
            "excel_report": str(tmp_path / "view.xlsx"),
            "partial_data_warnings": [{"dataset": "pulse", "kind": "stale"}],
        },
        workbook,
    )
    binding = mdw.ask_ai_binding(snapshot)
    assert snapshot["word_available"] is True
    assert snapshot["excel_available"] is True
    assert snapshot["decision_metrics"]
    assert binding["analysis_id"] == "Leader_Manager_90d"
    assert binding["fact_fingerprint"] == "view"
    assert binding["scope_type"] == "team"
    assert binding["scope_value"] == ""
    assert snapshot["scope_label"] == "Manager One team"


def test_legacy_compact_snapshot_projects_kpis_charts_and_public_warnings(tmp_path):
    path = tmp_path / "legacy-compact.xlsx"
    workbook = Workbook()
    info = workbook.active
    info.title = "Report_Info"
    info.append(["Item", "Value"])
    for row in (
        ["Export type", "Standard (Compact)"],
        ["Manager", "Manager One"],
        ["Technology", "All Contact Center"],
        ["Days", 90],
        ["Report_Generated_At_UTC", "2026-08-03T12:00:00Z"],
        [
            "Partial_Data_Warning",
            "adoption_barriers (tech_filter_scope_excluded): internal filter detail",
        ],
    ):
        info.append(row)
    dashboard = workbook.create_sheet("Executive_Dashboard")
    dashboard.append(["Metric", "Value", "Status"])
    dashboard.append(["Total Customers Analyzed", 2, "Portfolio"])
    risk = workbook.create_sheet("Risk_Summary")
    risk.append(["Customer", "Overall_Risk_Score", "Risk_Band", "Support_Cases"])
    risk.append(["Acme", 7.5, "HIGH", 1])
    risk.append(["Beta", 2.0, "LOW", 0])
    actions = workbook.create_sheet("Action_Plans")
    actions.append([
        "ID", "Customer Name", "Subject", "Status", "Priority",
        "Due Date", "Next Action", "Next Action Owner",
    ])
    actions.append([
        "AP-1", "Acme", "Recover adoption", "Open", "P1",
        "2026-08-01", "Confirm owner", "Owner A",
    ])
    actions.append([
        "AP-2", "Beta", "Close workshop", "Completed - Successful", "P3",
        "2026-07-15", "Archive", "Owner B",
    ])
    barriers = workbook.create_sheet("All_Adoption_Barriers")
    barriers.append(["ID", "BU_NAME"])
    barriers.append(["AB-1", "Acme"])
    cases = workbook.create_sheet("All_Support_Cases")
    cases.append(["Case Number", "Customer"])
    cases.append(["TAC-1", "Acme"])
    workbook.save(path)
    workbook.close()

    workbook_snapshot = mdw.load_workbook_snapshot(path)
    snapshot = mdw.snapshot_from_status(
        {
            "analysis_id": "legacy-compact",
            "status": "completed",
            "report_type": "compact",
            "manager": "Manager One",
            "technology": "All Contact Center",
            "days": 90,
        },
        workbook_snapshot,
    )
    metric_values = {
        item["metric_key"]: item["value"] for item in snapshot["decision_metrics"]
    }
    assert snapshot["canonical_snapshot"] is False
    assert metric_values["kpi.customers"] == 2
    assert metric_values["kpi.action_plans_open"] == 1
    assert metric_values["kpi.action_plans_overdue"] == 1
    assert metric_values["kpi.tac_cases"] == 1
    assert {item["chart_key"] for item in snapshot["charts"]} == {
        "activity_mix", "action_plan_status_aging", "risk_distribution",
    }
    assert snapshot["accounts"][0]["risk_score_0_100"] == 75
    warning_copy = " ".join(snapshot["source_warnings"]).casefold()
    assert "adoption barriers" in warning_copy
    assert "tech_filter_scope_excluded" not in warning_copy
    assert "internal filter detail" not in warning_copy
    assert any("compatibility projection" in item for item in snapshot["source_warnings"])


def test_rehydrated_leader_member_recovers_email_but_composite_customer_fails_closed(tmp_path):
    member_workbook = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "member.xlsx",
            fingerprint="member-fingerprint",
            open_aps=1,
            scope_type="member",
            scope_value="Alice Able (alice@example.invalid)",
        )
    )
    member = mdw.snapshot_from_status(
        {
            "analysis_id": "old-member",
            "status": "completed",
            "report_type": "leader",
            "manager": "Manager One",
            "_rehydrated_from_audit": True,
        },
        member_workbook,
    )
    assert member["scope_value"] == "alice@example.invalid"
    assert mdw.ask_ai_binding(member)["binding_available"] is True

    customer_workbook = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "customer.xlsx",
            fingerprint="customer-fingerprint",
            open_aps=1,
            scope_type="customer",
            scope_value="Acme Corporation (Alice Able)",
        )
    )
    customer = mdw.snapshot_from_status(
        {
            "analysis_id": "old-customer",
            "status": "completed",
            "report_type": "leader",
            "manager": "Manager One",
            "_rehydrated_from_audit": True,
        },
        customer_workbook,
    )
    binding = mdw.ask_ai_binding(customer)
    assert binding["binding_available"] is False
    assert binding["scope_value"] == ""


def test_compare_snapshots_classifies_business_and_source_changes(tmp_path):
    before = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "before.xlsx",
            fingerprint="before",
            open_aps=2,
            source_state="available",
            ap_rows=[
                ["AP-1", "Acme", "Plan 1", "On Track", "Owner A", "2026-08-10", "High", "Next", False],
                ["AP-2", "Beta", "Plan 2", "Completed", "Owner B", "2026-07-01", "Low", "Done", False],
            ],
        )
    )
    after = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "after.xlsx",
            fingerprint="after",
            open_aps=1,
            source_state="partial",
            ap_rows=[
                ["AP-1", "Acme", "Plan 1", "Completed", "Owner C", "2026-08-12", "High", "Done", False],
                ["AP-2", "Beta", "Plan 2", "On Track", "Owner B", "2026-07-01", "Low", "Restart", True],
                ["AP-3", "Gamma", "Plan 3", "New", "Owner D", "2026-09-01", "Medium", "Start", False],
            ],
        )
    )
    comparison = mdw.compare_snapshots(before, after)
    change_names = {item["change"] for item in comparison["source_driven_action_plan_changes"]}
    assert {"completed", "reopened", "became_overdue", "owner_changed", "due_date_changed", "new"} <= change_names
    assert comparison["metric_changes"] == []
    assert comparison["action_plan_changes"] == []
    assert comparison["source_driven_metric_changes"][0]["delta"] == -1
    assert comparison["source_driven_metric_changes"][0]["classification"] == "source_availability"
    assert comparison["business_change_count"] == 0
    assert comparison["source_driven_change_count"] == (
        len(comparison["source_driven_metric_changes"])
        + len(comparison["source_driven_action_plan_changes"])
    )
    assert comparison["source_state_changes"] == [
        {"source": "Action_Plans", "before": "available", "after": "partial"}
    ]
    assert any("Source coverage changed" in caveat for caveat in comparison["caveats"])


def test_compare_snapshots_keeps_like_for_like_changes_as_business_movement(tmp_path):
    before = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "business-before.xlsx",
            fingerprint="business-before",
            open_aps=1,
            ap_rows=[
                ["AP-1", "Acme", "Plan 1", "On Track", "Owner A", "2026-08-10", "High", "Next", False],
            ],
        )
    )
    after = mdw.load_workbook_snapshot(
        _write_snapshot(
            tmp_path / "business-after.xlsx",
            fingerprint="business-after",
            open_aps=0,
            ap_rows=[
                ["AP-1", "Acme", "Plan 1", "Completed", "Owner A", "2026-08-10", "High", "Done", False],
            ],
        )
    )

    comparison = mdw.compare_snapshots(before, after)

    assert comparison["metric_changes"][0]["delta"] == -1
    assert {item["change"] for item in comparison["action_plan_changes"]} >= {"completed", "status_changed"}
    assert comparison["business_change_count"] == (
        len(comparison["metric_changes"]) + len(comparison["action_plan_changes"])
    )
    assert comparison["source_driven_metric_changes"] == []
    assert comparison["source_driven_action_plan_changes"] == []
    assert comparison["source_driven_change_count"] == 0


def test_compare_snapshots_does_not_report_outage_rows_as_absent_business_changes():
    before = {
        "metrics": [{
            "metric_key": "kpi.action_plans_open",
            "label": "Open Action Plans",
            "value": 1,
            "source_state": "available",
            "source_sheet": "Action_Plans",
        }],
        "source_states": {"Action_Plans": "available"},
        "action_plans": [{"record_id": "AP-1", "customer": "Acme", "title": "Plan 1"}],
    }
    after = {
        "metrics": [{
            "metric_key": "kpi.action_plans_open",
            "label": "Open Action Plans",
            "value": 0,
            "source_state": "unavailable",
            "source_sheet": "Action_Plans",
        }],
        "source_states": {"Action_Plans": "unavailable"},
        "action_plans": [],
    }

    comparison = mdw.compare_snapshots(before, after)

    assert comparison["business_change_count"] == 0
    assert comparison["metric_changes"] == []
    assert comparison["action_plan_changes"] == []
    assert comparison["source_driven_metric_changes"][0]["delta"] == -1
    assert comparison["source_driven_action_plan_changes"] == [{
        "record_id": "AP-1",
        "change": "absent",
        "customer": "Acme",
        "title": "Plan 1",
        "classification": "source_availability",
        "before_source_state": "available",
        "after_source_state": "unavailable",
    }]


def test_compare_snapshots_fails_closed_when_scope_differs():
    before = {
        "manager": "One", "technology": "All", "scope_type": "team",
        "scope_value": "", "days": 90, "metrics": [{"metric_key": "kpi.customers", "value": 1}],
    }
    after = {
        "manager": "Two", "technology": "All", "scope_type": "team",
        "scope_value": "", "days": 90, "metrics": [{"metric_key": "kpi.customers", "value": 999}],
    }
    comparison = mdw.compare_snapshots(before, after)
    assert comparison["same_scope"] is False
    assert comparison["comparison_state"] == "incompatible_scope"
    assert comparison["incompatible_fields"] == ["manager"]
    assert comparison["metric_changes"] == []
    assert comparison["action_plan_changes"] == []
    assert comparison["business_change_count"] == 0
    assert any("do not have the same" in caveat for caveat in comparison["caveats"])


def test_history_filter_supports_manager_report_scope_customer_technology_and_status():
    records = [
        {
            "request_id": "one",
            "report_type": "leader",
            "manager": "Manager One",
            "technology": "All",
            "customer_name": "Acme",
            "status": "completed",
            "days": 90,
            "created_at": "2026-08-03T12:00:00Z",
            "word_path": "/safe/one.docx",
            "excel_path": "/safe/one.xlsx",
        },
        {
            "request_id": "two",
            "report_type": "compact",
            "manager": "Manager Two",
            "technology": "All Contact Center",
            "customer_name": "Beta",
            "status": "error",
            "days": 30,
            "created_at": "2026-08-02T12:00:00Z",
        },
    ]
    filtered = mdw.filter_history(
        records,
        manager="Manager One",
        report_type="leader",
        scope_type="customer",
        customer="acm",
        technology="All",
        status="completed",
    )
    assert [item["analysis_id"] for item in filtered] == ["one"]
    assert filtered[0]["word_available"] is True
    assert filtered[0]["excel_available"] is True
    assert "word_path" not in filtered[0]


def test_source_shape_keeps_round146_contract_markers():
    source = Path(mdw.__file__).read_text(encoding="utf-8")
    assert_in_source(source, "Round 146", label='source')
    assert_in_source(source, "WORKSPACE_SCHEMA", label='source')
    assert_in_source(source, "compare_snapshots", label='source')
    assert_in_source(source, "load_workbook_snapshot", label='source')
