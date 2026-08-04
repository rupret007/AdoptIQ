"""Manager-feedback acceptance gates, centered on the Leader report."""

from __future__ import annotations

import base64
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import decision_report_delivery as delivery
from scripts import generate_offline_acceptance_artifacts as harness


AS_OF = "2026-08-03T12:00:00Z"
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _fake_chart_renderer(_chart_id: str, _rows: pd.DataFrame, target: Path) -> bool:
    target.write_bytes(_TINY_PNG)
    return True


def _visible_text(document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(parts)


def _leader_facts(scope: str, *, warnings: list[dict] | None = None) -> dict:
    payload = harness.load_sanitized_fixture()
    team_data, labels = harness.build_scope_fixture(payload, scope, delivery)
    return delivery.build_report_facts(
        team_data,
        report_type=labels["report_type"],
        scope_type=labels["scope_type"],
        scope_value=labels["scope_value"],
        manager_name=str(payload["manager_name"]),
        days=int(payload["days"]),
        as_of=AS_OF,
        external_incidents=payload["external_incidents"],
        external_bugs=payload["external_bugs"],
        partial_data_warnings=warnings or [],
    )


def test_leader_report_directly_closes_manager_feedback(monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep decisions in Word and complete records in Source Data."""

    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    facts = _leader_facts(
        "team",
        warnings=[
            {
                "dataset": "csone_tac",
                "kind": "tac_unmatched_after_subscription_join",
                "effect": "1 TAC row had no SUBSCRIPTION_ID / ACCOUNT_ID_C attribution",
            }
        ],
    )
    document = delivery.build_concise_word_document(facts)
    sheets = delivery.build_source_data_sheets(facts)
    contract = delivery.validate_cross_artifact_contract(facts, sheets, document)
    text = _visible_text(document)
    lower_text = text.casefold()

    assert contract["ok"], contract["errors"]
    assert contract["word"]["word_count"] <= 1500
    assert len(document.inline_shapes) == 4
    assert {
        "Activity Mix by Source",
        "Action Plan Status and Aging",
        "Customer Risk Distribution",
        "Activity Trend",
    }.issubset({paragraph.text for paragraph in document.paragraphs})

    # The screenshot's repeated raw activities/cases must not return to Word.
    assert "no subject" not in lower_text
    assert "all action plans" not in lower_text
    assert "all tac cases" not in lower_text
    assert "complete activity, case, and source records" in lower_text

    # Action Plans remain decision-ready in Word and complete in Source Data.
    for label in ("Record ID", "Account", "Owner", "Status", "Due", "Priority", "Next action"):
        assert label.casefold() in lower_text
    assert len(facts["top_action_plans"]) == 5
    assert "AP-005" not in text
    assert len(sheets["Action_Plans"]) == facts["action_plan_lifecycle"]["total"]
    assert "AP-005" in set(sheets["Action_Plans"]["Record_ID"].astype(str))
    assert {
        "Record_ID",
        "AdoptIQ_Title",
        "AdoptIQ_Status_Bucket",
        "AdoptIQ_Due_Date",
        "PRIORITY_C",
    }.issubset(sheets["Action_Plans"].columns)
    assert len(sheets["TAC_Cases"]) == facts["kpis"]["tac_cases"]

    # A manager sees an honest, plain-language limitation—not implementation tokens.
    assert "coverage has disclosed limitations" in lower_text
    assert "coverage is complete" not in lower_text
    assert "partial assignment" in lower_text
    assert "unassigned portfolio" in lower_text
    for implementation_token in (
        "csone_tac",
        "tac_unmatched_after_subscription_join",
        "subscription_id",
        "account_id_c",
    ):
        assert implementation_token not in lower_text


@pytest.mark.parametrize("scope", ["team", "member", "customer"])
def test_leader_team_member_and_customer_deep_dives_remain_concise_and_complete(
    scope: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    facts = _leader_facts(scope)
    document = delivery.build_concise_word_document(facts)
    sheets = delivery.build_source_data_sheets(facts)
    contract = delivery.validate_cross_artifact_contract(facts, sheets, document)

    assert facts["report_type"] == "Leader"
    assert facts["scope_type"] == scope
    assert contract["ok"], contract["errors"]
    assert contract["word"]["word_count"] <= 1500
    assert len(document.inline_shapes) == 4
    assert "1 are" not in _visible_text(document)
    assert set(delivery.SOURCE_DATA_SHEET_NAMES) == set(sheets)
    assert len(sheets["Action_Plans"]) == facts["action_plan_lifecycle"]["total"]


def test_source_data_dates_render_and_risk_lineage_uses_score_units(tmp_path: Path) -> None:
    facts = _leader_facts("team")
    sheets = delivery.build_source_data_sheets(facts)
    target = tmp_path / "leader-source-data.xlsx"
    delivery.write_source_data_workbook(target, sheets)

    lineage = sheets["Metric_Lineage"]
    risk_rows = lineage.loc[
        lineage["Metric_Key"].astype(str).str.match(r"summary\.account\..*\.risk_score$")
    ]
    assert not risk_rows.empty
    assert set(risk_rows["Unit"].astype(str)) == {"score (0–100)"}

    workbook = load_workbook(target, read_only=False, data_only=False)
    try:
        chart_sheet = workbook["Chart_Data"]
        headers = [cell.value for cell in chart_sheet[1]]
        column_number = headers.index("Period_Start") + 1
        column_letter = chart_sheet.cell(1, column_number).column_letter
        populated = [
            chart_sheet.cell(row_number, column_number)
            for row_number in range(2, chart_sheet.max_row + 1)
            if chart_sheet.cell(row_number, column_number).value is not None
        ]
        assert populated
        assert all(cell.number_format == "yyyy-mm-dd" for cell in populated)
        assert float(chart_sheet.column_dimensions[column_letter].width or 0) >= 14
    finally:
        workbook.close()
