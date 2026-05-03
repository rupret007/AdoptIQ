"""Round 64 / Phase 2 (B2) -- Comprehensive XLSX dedicated Action_Plans sheet.

The Build-36 Comprehensive XLSX consistently reported
``Action plans (open): 0`` in its ``Summary`` sheet because:

1. The R62 / B helper ``canonical_metrics.count_open_action_plans``
   read ``AB_Detail_All.Action Plan Title``, but the comprehensive
   flow's ``AB_Detail_All`` sheet is pure adoption barriers (the AB
   join never populates the Action Plan Title cell), so all 70 rows
   were empty.
2. The comprehensive XLSX never wrote a dedicated ``Action_Plans``
   sheet (mirroring the Compact / Leader naming convention) so the
   helper had no real action-plan source to derive from. The Leader
   XLSX for the same portfolio had 365 AP rows (~120 currently open).

R64 / Phase 2 (B2) closes the structural gap:

* Add a dedicated ``Action_Plans`` sheet to the comprehensive XLSX
  (sourced from ``filtered_action_plans``, scoped to the manager +
  technology by ``_filter_csconsole_data_by_technology``).
* Extend ``count_open_action_plans(ab_df, ap_df=None)`` with an
  optional ``ap_df`` parameter that counts non-closed-status rows
  when present (case-insensitive match against the canonical closed
  set: ``Completed - Successful``, ``Completed - Unsuccessful``,
  ``Closed - Cancelled``, ``Closed``).
* Fall back to the R62 / B AB_Detail_All path when ``ap_df`` is
  ``None`` / empty so older workbooks still resolve.
* Wire the new sheet through ``report_export_styling.build_summary_rows``
  so the Summary cell now reflects real AP data.

These tests pin the contract end-to-end without requiring a live
Snowflake / Flask environment.

Round 64 / Phase 2.  Made-with: Cursor.
"""

from __future__ import annotations

import pandas as pd

import canonical_metrics as cm
import report_export_styling


# ---------------------------------------------------------------------------
# Extended count_open_action_plans contracts
# ---------------------------------------------------------------------------


def test_count_open_action_plans_with_ap_df_and_status_column_counts_non_closed():
    """Given a real Action_Plans frame, count rows whose STATUS_C is NOT
    in the canonical closed set."""
    ap_df = pd.DataFrame({
        "AP_NAME_C": [f"AP-{i}" for i in range(8)],
        "STATUS_C": [
            "New Request",                  # open
            "On Track",                     # open
            "On Hold",                      # open
            "Off Trajectory",               # open
            "Completed - Successful",       # closed
            "Completed - Unsuccessful",     # closed
            "Closed - Cancelled",           # closed
            "Closed",                       # closed
        ],
    })

    assert cm.count_open_action_plans(None, ap_df=ap_df) == 4


def test_count_open_action_plans_with_ap_df_is_case_insensitive():
    """Status comparison must be case-insensitive (CSConsole exports
    sometimes upper-case the value)."""
    ap_df = pd.DataFrame({"STATUS_C": ["new request", "COMPLETED - SUCCESSFUL", "Closed", "on track"]})
    # 4 rows, 2 closed (Completed-Successful, Closed) -> 2 open
    assert cm.count_open_action_plans(None, ap_df=ap_df) == 2


def test_count_open_action_plans_with_ap_df_handles_double_space_variants():
    """Whitespace-collapse: 'completed -  successful' (double space) -> closed."""
    ap_df = pd.DataFrame({"STATUS_C": ["completed -  successful", "On Track"]})
    assert cm.count_open_action_plans(None, ap_df=ap_df) == 1


def test_count_open_action_plans_with_ap_df_treats_blank_status_as_open():
    """An action plan with no status is in-flight, not closed (safer half).

    This pins the R64/B2 design choice: never silently mark a status-
    less row as closed -- doing so would under-count open APs and the
    Summary cell would understate work-in-progress.
    """
    ap_df = pd.DataFrame({"STATUS_C": ["", None, "  ", "Closed"]})
    # 3 blanks (open) + 1 Closed (closed) -> 3 open
    assert cm.count_open_action_plans(None, ap_df=ap_df) == 3


def test_count_open_action_plans_with_ap_df_no_status_column_returns_zero():
    """Round 71 / Phase 4 (#20): when the AP frame has no recognizable
    status column, return 0 + structured warning -- pre-R71 we returned
    ``len(ap_df)`` (silently marking every row as "open"), which
    inflated the Summary KPI any time the upstream Snowflake projection
    was missing STATUS_C.  Returning 0 + a warning is the correct
    fail-safe: the Summary cell now reads honestly while the operator
    has greppable provenance for the schema drift."""
    ap_df = pd.DataFrame({"AP_NAME_C": ["AP-1", "AP-2", "AP-3"]})
    assert cm.count_open_action_plans(None, ap_df=ap_df) == 0


def test_count_open_action_plans_tolerates_alternate_status_column_names():
    """STATUS_C / AP_STATUS_C / Status / STATUS / status / case_status_norm
    / status_norm should all resolve."""
    for col_name in ("STATUS_C", "AP_STATUS_C", "Status", "STATUS", "status",
                     "case_status_norm", "status_norm"):
        ap_df = pd.DataFrame({col_name: ["New Request", "Completed - Successful", "On Track"]})
        assert cm.count_open_action_plans(None, ap_df=ap_df) == 2, (
            f"helper failed to recognize status column {col_name!r}"
        )


def test_count_open_action_plans_falls_back_to_ab_when_ap_df_is_none():
    """R62/B back-compat: when ap_df=None, count from AB_Detail_All's
    Action Plan Title column (the renewal scenario depends on this
    fallback)."""
    ab = pd.DataFrame({
        "ID": ["AB-1", "AB-2", "AB-3"],
        "Action Plan Title": ["Plan A", "", "Plan B"],
    })
    assert cm.count_open_action_plans(ab, ap_df=None) == 2
    # And without the kwarg at all (positional-only first arg back-compat).
    assert cm.count_open_action_plans(ab) == 2


def test_count_open_action_plans_falls_back_to_ab_when_ap_df_is_empty():
    """An empty Action_Plans frame must NOT short-circuit to 0 -- the
    R62/B AB fallback should still try."""
    ab = pd.DataFrame({"Action Plan Title": ["Real Plan", "Another Plan"]})
    assert cm.count_open_action_plans(ab, ap_df=pd.DataFrame()) == 2


def test_count_open_action_plans_returns_zero_when_both_sources_empty():
    """All inputs empty -> 0, no exceptions."""
    assert cm.count_open_action_plans(pd.DataFrame(), ap_df=pd.DataFrame()) == 0
    assert cm.count_open_action_plans(None, ap_df=None) == 0


# ---------------------------------------------------------------------------
# build_summary_rows wiring contracts
# ---------------------------------------------------------------------------


def _ap_with_n_open(n_open: int, n_closed: int = 0) -> pd.DataFrame:
    """Build an Action_Plans-shaped frame with N open + M closed rows."""
    statuses = ["New Request"] * n_open + ["Completed - Successful"] * n_closed
    return pd.DataFrame({
        "AP_NAME_C": [f"AP-{i}" for i in range(n_open + n_closed)],
        "STATUS_C": statuses,
    })


def test_build_summary_rows_uses_action_plans_sheet_when_present():
    """Comprehensive scenario: when the Action_Plans sheet has 120 open
    rows, the Summary cell must read 120 -- not the 0 the AB-only path
    would have produced (since AB_Detail_All has empty Action Plan
    Title cells for every row in this scenario)."""
    ab = pd.DataFrame({
        "ID": [f"AB-{i}" for i in range(70)],
        "Action Plan Title": [""] * 70,  # the production data shape
    })
    ap = _ap_with_n_open(120, n_closed=245)  # 365 total rows, 120 open

    rows = report_export_styling.build_summary_rows(
        sheets={"AB_Detail_All": ab, "Action_Plans": ap},
        manager="Brian Frazier",
        tech="All Contact Center",
        days=90,
    )
    by_label = dict(rows)
    assert by_label.get("Action plans (open)") == "120", (
        "Round 64 / Phase 2 (B2): Action_Plans sheet must take precedence "
        f"over the empty AB_Detail_All path; got {by_label.get('Action plans (open)')!r}"
    )


def test_build_summary_rows_falls_back_to_ab_when_no_action_plans_sheet():
    """R62/B back-compat: when the workbook has no Action_Plans sheet
    (renewal scenario, older comprehensive workbooks), the Summary
    cell still derives from AB_Detail_All's Action Plan Title column.
    Reuses the existing R62/B path."""
    ab = pd.DataFrame({
        "ID": ["AB-1", "AB-2", "AB-3"],
        "Action Plan Title": ["Plan A", "", "Plan B"],
    })
    rows = report_export_styling.build_summary_rows(
        sheets={"AB_Detail_All": ab},
        manager="Test", tech="Test", days=90,
    )
    by_label = dict(rows)
    assert by_label.get("Action plans (open)") == "2", (
        "Round 64 / Phase 2 (B2) MUST preserve the R62/B AB-fallback for "
        "workbooks that pre-date the dedicated Action_Plans sheet"
    )


def test_build_summary_rows_falls_back_to_csconsole_action_plans_sheet():
    """Belt-and-suspenders: if the workbook has a CSConsole_Action_Plans
    sheet but no Action_Plans sheet (mid-rollout shape), the Summary
    helper still resolves the AP data via the alternate sheet name."""
    ab = pd.DataFrame({"Action Plan Title": [""] * 5})
    ap_via_csconsole = _ap_with_n_open(7, n_closed=2)
    rows = report_export_styling.build_summary_rows(
        sheets={"AB_Detail_All": ab, "CSConsole_Action_Plans": ap_via_csconsole},
        manager="Test", tech="Test", days=90,
    )
    by_label = dict(rows)
    assert by_label.get("Action plans (open)") == "7"


def test_build_summary_rows_falls_back_to_csconsole_data_dict():
    """And one more level of back-compat: callers that thread the
    Action_Plans data only via ``csconsole_data['action_plans']`` (the
    historical Comprehensive shape) still get a valid count."""
    ab = pd.DataFrame({"Action Plan Title": [""] * 3})
    ap = _ap_with_n_open(5, n_closed=1)
    rows = report_export_styling.build_summary_rows(
        sheets={"AB_Detail_All": ab},
        csconsole_data={"action_plans": ap},
        manager="Test", tech="Test", days=90,
    )
    by_label = dict(rows)
    assert by_label.get("Action plans (open)") == "5"


def test_build_summary_rows_action_plans_zero_when_ap_sheet_all_closed():
    """Edge case: Action_Plans sheet present but every row is closed ->
    0 open. The Summary cell must surface "0" (deterministic), not "--".
    """
    ab = pd.DataFrame({"Action Plan Title": [""] * 5})
    ap = _ap_with_n_open(0, n_closed=10)
    rows = report_export_styling.build_summary_rows(
        sheets={"AB_Detail_All": ab, "Action_Plans": ap},
        manager="Test", tech="Test", days=90,
    )
    by_label = dict(rows)
    assert by_label.get("Action plans (open)") == "0"


def test_build_summary_rows_position_invariant_holds_with_action_plans_sheet():
    """The R62/B golden-fixture position (Action plans (open) sits
    between Adoption barriers (open) and TAC cases (total)) MUST
    survive the R64/B2 wiring change."""
    ab = pd.DataFrame({"ID": ["AB-1"], "Action Plan Title": [""]})
    ap = _ap_with_n_open(3)
    rows = report_export_styling.build_summary_rows(
        sheets={"AB_Detail_All": ab, "Action_Plans": ap},
        manager="Test", tech="Test", days=90,
    )
    labels = [label for label, _ in rows]
    ap_idx = labels.index("Action plans (open)")
    open_barriers_idx = labels.index("Adoption barriers (open)")
    tac_total_idx = labels.index("TAC cases (total)")
    assert open_barriers_idx < ap_idx < tac_total_idx, (
        f"Round 62 / B golden-fixture position drift -- 'Action plans (open)' "
        f"must sit between 'Adoption barriers (open)' ({open_barriers_idx}) and "
        f"'TAC cases (total)' ({tac_total_idx}); actually at {ap_idx}.  "
        f"Full sequence: {labels}"
    )


# ---------------------------------------------------------------------------
# End-to-end: comprehensive XLSX writer puts Action_Plans into the workbook
# ---------------------------------------------------------------------------


def test_write_excel_workbook_includes_action_plans_sheet_when_supplied(tmp_path):
    """Pin that ``adoptiq_backend.write_excel_workbook`` honors the new
    ``Action_Plans`` key in the sheets dict.  The comprehensive flow
    in ``app_simple.run_comprehensive_analysis`` adds this key when
    ``filtered_action_plans`` is non-empty."""
    from openpyxl import load_workbook  # noqa: PLC0415

    import adoptiq_backend  # noqa: PLC0415

    sheets = {
        "AB_Detail_All": pd.DataFrame({
            "ID": ["AB-1", "AB-2"],
            "Customer Name": ["Acme", "Beta"],
        }),
        "CSOne_Detail_All": pd.DataFrame({
            "Case Number": ["CS-1"],
            "customer_name": ["Acme"],
        }),
        "Action_Plans": _ap_with_n_open(5, n_closed=3),
    }
    base = str(tmp_path / "round64_action_plans_smoke")
    out_path = adoptiq_backend.write_excel_workbook(base, sheets)

    wb = load_workbook(out_path, data_only=True)
    assert "Action_Plans" in wb.sheetnames, (
        "Round 64 / Phase 2 (B2): write_excel_workbook must surface the "
        f"Action_Plans sheet when supplied; sheetnames={wb.sheetnames}"
    )
    # And the Summary sheet should reflect the open count from this sheet.
    summary_ws = wb["Summary"]
    summary_rows = list(summary_ws.iter_rows(values_only=True))
    summary_dict = {row[0]: row[1] for row in summary_rows if row and len(row) >= 2}
    assert summary_dict.get("Action plans (open)") == "5", (
        f"Round 64 / Phase 2 (B2): Summary 'Action plans (open)' must "
        f"reflect the Action_Plans sheet's open count; got {summary_dict.get('Action plans (open)')!r}"
    )
