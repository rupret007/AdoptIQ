"""Focused acceptance tests for the legacy-to-canonical report adapter."""

from __future__ import annotations

import base64
from pathlib import Path
from typing import Any

import pandas as pd
import pytest
from docx import Document

import canonical_metrics as cm
import decision_report_delivery as delivery
import canonical_report_adapter as adapter
from canonical_report_adapter import (
    CanonicalReportAdapterError,
    canonicalize_legacy_artifacts,
)
from risk_scoring import compute_customer_risk_profile


AS_OF = "2026-08-03T21:00:00Z"
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _fake_chart_renderer(_chart_id: str, _rows: pd.DataFrame, target: Path) -> bool:
    target.write_bytes(_TINY_PNG)
    return True


def _document_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(parts)


def _source_frames(marker: str) -> dict[str, pd.DataFrame]:
    return {
        "action": pd.DataFrame(
            [
                {
                    "ID": f"AP-{marker}",
                    "CSSM": "Jordan CSSM",
                    "Customer Name": "Acme Corporation",
                    "Account ID": "ACC-001",
                    "Subject": "Resolve onboarding gap",
                    "Description": f"RAW ACTION DETAIL {marker}",
                    "Status": "Open",
                    "Priority": "P1",
                    "Created Date": "2026-07-01",
                    "Due Date": "2026-08-01",
                    "Next Action": "Confirm the revised enablement date",
                    "Next Action Owner": "Alex Rivera",
                }
            ]
        ),
        "barrier": pd.DataFrame(
            [
                {
                    "ID": f"AB-{marker}",
                    "CSSM": "Jordan CSSM",
                    "Customer Name": "Acme Corporation",
                    "Account ID": "ACC-001",
                    "Subject": "Calling migration dependency",
                    "Description": f"RAW BARRIER EVIDENCE {marker}",
                    "Adoption Barrier Status": "Open",
                    "Status": "Open",
                    "Severity": "P1",
                    "Open Date": "2026-07-05",
                }
            ]
        ),
        "pulse": pd.DataFrame(
            [
                {
                    "ID": f"CP-{marker}",
                    "CSSM": "Jordan CSSM",
                    "Account ID": "ACC-001",
                    "Customer Name": "Acme Corporation",
                    "Customer Pulse": "Red",
                    "Pulse Rating": "Poor",
                    "Pulse Score": 1,
                    "Pulse Date": "2026-07-10",
                    "Comments": f"RAW PULSE EVIDENCE {marker}",
                }
            ]
        ),
        "tac": pd.DataFrame(
            [
                {
                    "SR Number": f"CASE-{marker}",
                    "CSSM": "Jordan CSSM",
                    "Customer Name": "Acme Corporation",
                    "Account ID": "ACC-001",
                    "Title": "Media path instability",
                    "Severity": "P1",
                    "Case Status": "Open",
                    "Date/Time Opened": "2026-07-15",
                    "Problem Description": f"RAW TAC EVIDENCE {marker}",
                }
            ]
        ),
        "priority": pd.DataFrame(
            [
                {
                    "ID": f"SP-{marker}",
                    "CSSM": "Jordan CSSM",
                    "Account ID": "ACC-001",
                    "Customer Name": "Acme Corporation",
                    "Success Priority Title": "Complete migration",
                    "Status": "Active",
                }
            ]
        ),
        "incident": pd.DataFrame(
            [
                {
                    "id": f"INC-{marker}",
                    "title": "Sanitized service incident",
                    "status": "resolved",
                    "description": f"RAW INCIDENT EVIDENCE {marker}",
                }
            ]
        ),
        "bug": pd.DataFrame(
            [
                {
                    "bug_id": f"BUG-{marker}",
                    "title": "Sanitized external defect",
                    "source": "help.webex.com",
                }
            ]
        ),
    }


def _fixture_subscription(marker: str) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "SUBSCRIPTION_ID": f"SUB-{marker}",
                "ACCOUNT_ID_C": "ACC-001",
                "BU_NAME": "Acme Corporation",
                "CSSM": "Jordan CSSM",
                "TECHNOLOGY_C": "Webex Calling",
                "SUB_TECHNOLOGY_C": "Webex Calling",
                "STATUS_C": "Active",
            }
        ]
    )


def _canonical_fixture_risk(
    frames: dict[str, pd.DataFrame],
    subscriptions: pd.DataFrame,
) -> tuple[float, str]:
    """Derive the legacy fixture claim from the canonical scoped inputs."""

    incident_state = cm.source_data_state(frames["incident"])
    scoped_incidents = delivery._customer_incidents(  # noqa: SLF001
        frames["incident"],
        "Acme Corporation",
        identity={"account_ids": ("acc-001",)},
    )
    profile = compute_customer_risk_profile(
        "Acme Corporation",
        customer_ab=frames["barrier"],
        customer_csone=frames["tac"],
        customer_pulse=frames["pulse"],
        customer_action_plans=frames["action"],
        customer_subs=subscriptions,
        ext_incidents=scoped_incidents,
        incident_source_state=str(incident_state["state"]),
        incident_source_detail=str(incident_state["detail"]),
        recent_window_days=90,
        as_of=AS_OF,
    )
    return float(profile["risk_score_0_100"]), str(profile["risk_band"])


def _write_legacy_pair(
    tmp_path: Path,
    *,
    family: str,
    marker: str,
    risk_score: float | None = None,
    risk_band: str | None = None,
) -> tuple[Path, Path, dict[str, str]]:
    word_path = tmp_path / f"AdoptIQ_Report_{marker}.docx"
    legacy_document = Document()
    legacy_document.add_paragraph(f"LEGACY RAW APPENDIX {marker}")
    legacy_document.save(word_path)

    frames = _source_frames(marker)
    workbook_path = tmp_path / f"AdoptIQ_Legacy_{marker}.xlsx"
    if family == "compact":
        names = {
            "subscriptions": "Subscriptions",
            "family_facts": "Risk_Summary",
            "action_plans": "Action_Plans",
            "adoption_barriers": "All_Adoption_Barriers",
            "customer_pulse": "Customer_Pulse",
            "tac_cases": "All_Support_Cases",
            "success_priorities": "Success_Priorities",
        }
        subscriptions = _fixture_subscription(marker)
        derived_score, derived_band = _canonical_fixture_risk(
            frames,
            subscriptions,
        )
        resolved_score = derived_score if risk_score is None else risk_score
        resolved_band = derived_band if risk_band is None else risk_band
        compact_display_score = resolved_score / 10.0
        family_summary = pd.DataFrame(
            [
                {
                    "Customer": "Acme Corporation",
                    "CSSM": "Jordan CSSM",
                    "Overall_Risk_Score": compact_display_score,
                    "Risk_Band": resolved_band,
                }
            ]
        )
    elif family == "renewal":
        names = {
            "subscriptions": "Subscriptions",
            "family_facts": "Renewal_Summary",
            "action_plans": "Customer_Action_Plans",
            "adoption_barriers": "Customer_Adoption_Barriers",
            "customer_pulse": "Customer_Customer_Pulse",
            "tac_cases": "Customer_Support_Cases",
            "success_priorities": "Customer_Success_Priorities",
        }
        subscriptions = _fixture_subscription(marker)
        derived_score, derived_band = _canonical_fixture_risk(
            frames,
            subscriptions,
        )
        resolved_score = derived_score if risk_score is None else risk_score
        resolved_band = derived_band if risk_band is None else risk_band
        family_summary = pd.DataFrame(
            [
                {
                    "Customer": "Acme Corporation",
                    "CSSM": "Jordan CSSM",
                    "Overall_Risk_Score": resolved_score,
                    "Risk_Band": resolved_band,
                }
            ]
        )
    else:
        names = {
            "subscriptions": "Summary",
            "action_plans": "Action_Plans",
            "adoption_barriers": "Adoption_Barriers",
            "customer_pulse": "Customer_Pulse",
            "tac_cases": "TAC_Cases",
            "success_priorities": "Success_Priorities",
        }
        subscriptions = pd.DataFrame(
            {
                "Metric": [
                    "Customer Name",
                    "Subscription ID",
                    "Technology",
                    "Status",
                    "Renewal Date",
                    "Annual Recurring Revenue (ARR)",
                    "Contract Term",
                    "Adoption Risk",
                    "Recommendation",
                ],
                "Value": [
                    "Acme Corporation",
                    "SUB-001",
                    "Webex Calling",
                    "Active",
                    "2026-12-31",
                    125000,
                    "36 months",
                    "Medium",
                    "Complete the adoption workshop",
                ],
            }
        )

    report_info = pd.DataFrame(
        [
            {"Item": "Export type", "Value": family},
            {"Item": "Selected Account ID", "Value": "ACC-001"},
            {"Item": "Generated at (UTC)", "Value": AS_OF},
        ]
    )
    action_export = frames["action"]
    barrier_export = frames["barrier"]
    pulse_export = frames["pulse"]
    if family == "subscription":
        # Pre-R146 Subscription artifacts restamped raw headers over values
        # that had already been placed in curated-column order.
        action_export = pd.DataFrame(
            [
                {
                    "ID": f"AP-{marker}",
                    "CSSM": "Jordan CSSM",
                    "ACCOUNT_ID_C": "Acme Corporation",
                    "BU_NAME": "ACC-001",
                    "SUBJECT_C": "Resolve onboarding gap",
                    "STATUS_C": "Open",
                    "DUE_DATE_C": "P1",
                    "CREATED_DATE_C": "2026-07-01",
                    "PRIORITY_C": "2026-08-01",
                    "NEXT_ACTION_C": "Confirm the revised enablement date",
                    "NEXT_ACTION_OWNER_C": "Alex Rivera",
                }
            ]
        )
        barrier_export = pd.DataFrame(
            [
                {
                    "ID": f"AB-{marker}",
                    "CSSM": "Jordan CSSM",
                    "ACCOUNT_ID_C": "Acme Corporation",
                    "BU_NAME": "ACC-001",
                    "SUBJECT_C": "Calling migration dependency",
                    "SEVERITY_C": f"RAW BARRIER EVIDENCE {marker}",
                    "STATUS_C": "Open",
                    "OPEN_DATE_C": "P1",
                    "DESCRIPTION_C": "2026-07-05",
                }
            ]
        )
        pulse_export = pd.DataFrame(
            [
                {
                    "ID": f"CP-{marker}",
                    "CSSM": "Jordan CSSM",
                    "ACCOUNT__C": "Acme Corporation",
                    "BU_NAME": "ACC-001",
                    "CUSTOMER_PULSE__C": "Red",
                    "PULSE_RATING__C": "Poor",
                    "SCORE__C": 1,
                    "PULSE_DATE_C": "2026-07-10",
                }
            ]
        )
    with pd.ExcelWriter(workbook_path, engine="xlsxwriter") as writer:
        report_info.to_excel(writer, sheet_name="Report_Info", index=False)
        subscriptions.to_excel(writer, sheet_name=names["subscriptions"], index=False)
        if family in {"compact", "renewal"}:
            family_summary.to_excel(
                writer,
                sheet_name=names["family_facts"],
                index=False,
            )
        action_export.to_excel(writer, sheet_name=names["action_plans"], index=False)
        barrier_export.to_excel(writer, sheet_name=names["adoption_barriers"], index=False)
        pulse_export.to_excel(writer, sheet_name=names["customer_pulse"], index=False)
        frames["tac"].to_excel(writer, sheet_name=names["tac_cases"], index=False)
        frames["priority"].to_excel(writer, sheet_name=names["success_priorities"], index=False)
        frames["incident"].to_excel(writer, sheet_name="External Incidents", index=False)
        frames["bug"].to_excel(writer, sheet_name="External Defects", index=False)
    return word_path, workbook_path, names


@pytest.mark.parametrize(
    ("family", "report_type", "scope_type", "scope_value", "marker"),
    [
        ("compact", "Compact", "team", "Entire team", "COMPACT"),
        ("renewal", "Renewal Portfolio", "team", "Entire team", "RENEWAL-PORTFOLIO"),
        ("renewal", "Renewal", "customer", "Acme Corporation", "RENEWAL-CUSTOMER"),
        ("subscription", "Subscription", "subscription", "SUB-001", "SUBSCRIPTION"),
    ],
)
def test_legacy_report_families_become_one_validated_canonical_contract(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    family: str,
    report_type: str,
    scope_type: str,
    scope_value: str,
    marker: str,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, names = _write_legacy_pair(
        tmp_path,
        family=family,
        marker=marker,
    )

    result = canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type=report_type,
        manager_name="Local Fixture Manager",
        technology="Webex Calling",
        scope_type=scope_type,
        scope_value=scope_value,
        days=90,
        as_of=AS_OF,
    )

    source_path = Path(result["source_data_path"])
    assert result["contract"]["ok"] is True
    assert Path(result["word_path"]) == word_path.resolve()
    assert source_path.is_file()
    assert not workbook_path.exists()
    assert result["legacy_workbook_retired"] is True
    assert result["contract"]["legacy_workbook_retired"] is True
    assert result["contract"]["footer"] == {
        "ok": True,
        "sections": 1,
        "stamped_sections": 1,
    }
    assert result["contract"]["mapped_sheets"]["subscriptions"] == names["subscriptions"]
    assert result["facts"].get("technology", "Webex Calling") == "Webex Calling"
    assert result["facts"]["kpis"]["team_members"] == 1
    if family == "subscription":
        assert {
            "action_plans:customer-account-position",
            "action_plans:priority-due-date-position",
            "adoption_barriers:customer-account-position",
            "adoption_barriers:description-severity-open-date-position",
            "customer_pulse:customer-account-position",
        }.issubset(result["facts"]["legacy_adapter"]["header_repairs"])

    with pd.ExcelFile(source_path) as workbook:
        assert {
            "Report_Info",
            "Metric_Lineage",
            "Chart_Data",
            "Subscriptions",
            "Action_Plans",
            "Adoption_Barriers",
            "Customer_Pulse",
            "TAC_Cases",
            "Success_Priorities",
        }.issubset(workbook.sheet_names)
        action_plans = pd.read_excel(workbook, sheet_name="Action_Plans")
        barriers = pd.read_excel(workbook, sheet_name="Adoption_Barriers")
        tac_cases = pd.read_excel(workbook, sheet_name="TAC_Cases")
        subscriptions = pd.read_excel(workbook, sheet_name="Subscriptions")
        lineage = pd.read_excel(workbook, sheet_name="Metric_Lineage")
        evidence = pd.read_excel(workbook, sheet_name="Evidence_Links")

    assert action_plans["Record_ID"].tolist() == [f"AP-{marker}"]
    assert action_plans["CSSM"].tolist() == ["Jordan CSSM"]
    assert action_plans.loc[0, "Customer Name"] == "Acme Corporation"
    assert action_plans.loc[0, "Account ID"] == "ACC-001"
    assert action_plans.loc[0, "Priority"] == "P1"
    assert str(action_plans.loc[0, "Due Date"]).startswith("2026-08-01")
    assert barriers["Record_ID"].tolist() == [f"AB-{marker}"]
    assert tac_cases["Record_ID"].tolist() == [f"CASE-{marker}"]
    assert f"RAW BARRIER EVIDENCE {marker}" in set(barriers["Description"].astype(str))
    family_facts = subscriptions.loc[
        subscriptions["Legacy_Record_Type"].fillna("")
        == "Family-specific reported fact"
    ]
    assert not family_facts.empty
    assert set(family_facts["Legacy_Source_Sheet"]) == {
        names.get("family_facts", names["subscriptions"])
    }
    family_keys = set(family_facts["Metric_Key"].dropna().astype(str))
    assert family_keys
    assert family_keys.issubset(set(lineage["Metric_Key"].dropna().astype(str)))
    assert family_keys.issubset(set(evidence["Evidence_Key"].dropna().astype(str)))
    if family in {"compact", "renewal"}:
        core_subscriptions = subscriptions.loc[
            subscriptions["Legacy_Record_Type"].fillna("") == ""
        ]
        assert core_subscriptions["Record_ID"].tolist() == [f"SUB-{marker}"]
        assert core_subscriptions["Account ID"].tolist() == ["ACC-001"]
        assert core_subscriptions["TECHNOLOGY_C"].tolist() == ["Webex Calling"]
    if family == "compact":
        assert "Overall_Risk_Score" not in subscriptions.columns
        assert "Risk_Score_0_10" not in subscriptions.columns
        assert result["facts"]["legacy_adapter"][
            "superseded_family_risk_fields"
        ] == {
            "Risk_Summary": ["Overall_Risk_Score"],
        }

    document = Document(word_path)
    report_specific_tables = [
        table
        for table in document.tables
        if [cell.text for cell in table.rows[0].cells]
        == ["Decision fact", "Reported value"]
    ]
    if family in {"renewal", "subscription"}:
        assert len(report_specific_tables) == 1
        report_specific_rows = [
            [cell.text for cell in row.cells]
            for row in report_specific_tables[0].rows[1:]
        ]
        assert 1 <= len(report_specific_rows) <= delivery.REPORT_SPECIFIC_FACT_LIMIT
        if family == "subscription":
            visible_values = {row[1] for row in report_specific_rows}
            assert {
                "SUB-001",
                "Active",
                "2026-12-31",
                "125000",
                "36 months",
                "Medium",
                "Complete the adoption workshop",
            }.issubset(visible_values)
    else:
        assert report_specific_tables == []

    word_text = _document_text(word_path)
    assert f"LEGACY RAW APPENDIX {marker}" not in word_text
    assert f"RAW BARRIER EVIDENCE {marker}" not in word_text
    assert "Complete selected-scope records" in word_text
    expected_family_heading = f"{family.title()} Decision Facts"
    assert (expected_family_heading in word_text) is (
        family in {"renewal", "subscription"}
    )
    if scope_type in {"customer", "subscription"}:
        assert "team member" not in word_text.casefold()
    else:
        assert "1 team member" in word_text.casefold()

    assert all(
        any(
            paragraph.text.startswith("AdoptIQ v")
            for paragraph in section.footer.paragraphs
        )
        for section in document.sections
    )
    action_table = next(
        table
        for table in document.tables
        if [cell.text for cell in table.rows[0].cells][:3]
        == ["Record ID", "Account", "Owner"]
    )
    assert action_table.rows[1].cells[2].text == "Alex Rivera"


def test_renewal_claim_uses_scoped_no_smear_risk_and_reconciles_exactly(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="renewal",
        marker="NO-INCIDENT-SMEAR",
    )

    result = canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type="Renewal",
        manager_name="Local Fixture Manager",
        technology="Webex Calling",
        scope_type="customer",
        scope_value="Acme Corporation",
        days=90,
        as_of=AS_OF,
    )

    profile = result["facts"]["risk_profiles"]["Acme Corporation"]
    score_claim = next(
        claim
        for claim in result["facts"]["legacy_adapter"][
            "risk_claims_reconciled"
        ]
        if claim["sheet"] == "Renewal_Summary" and claim["kind"] == "score"
    )
    assert len(result["facts"]["external_incidents"]) == 1
    assert profile["components"]["incidents"]["details"]["count"] == 0
    assert score_claim["value"] == profile["risk_score_0_100"] == 50.9
    assert profile["risk_band"] == "MEDIUM"


def test_fetch_warning_marks_source_partial_and_withholds_decision_metrics(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="renewal",
        marker="WARNING-STATE",
    )

    result = canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type="Renewal",
        manager_name="Local Fixture Manager",
        technology="Webex Calling",
        scope_type="customer",
        scope_value="Acme Corporation",
        days=90,
        as_of=AS_OF,
        partial_data_warnings=[
            {
                "dataset": "csconsole_action_plans",
                "kind": "fetch_failed",
                "effect": "A later source page failed after one retained record.",
            }
        ],
    )

    facts = result["facts"]
    action_coverage = facts["source_coverage"].set_index("Source_Sheet").loc[
        "Action_Plans"
    ]
    assert action_coverage["Source_State"] == "partial"
    assert action_coverage["Record_Count"] == 1
    assert facts["action_plan_lifecycle"]["source_state"] == "partial"

    action_chart = facts["chart_data"].loc[
        facts["chart_data"]["Chart_ID"] == "action_plan_status_aging"
    ]
    assert set(action_chart["Source_State"]) == {"partial"}
    assert action_chart["Value"].isna().all()

    source_path = Path(result["source_data_path"])
    with pd.ExcelFile(source_path) as workbook:
        report_info = pd.read_excel(workbook, sheet_name="Report_Info")
        lineage = pd.read_excel(workbook, sheet_name="Metric_Lineage")
        evidence = pd.read_excel(workbook, sheet_name="Evidence_Links")

    report_state = report_info.loc[
        report_info["Item"] == "Source_State:Action_Plans", "Value"
    ]
    assert report_state.tolist() == ["partial"]
    action_lineage = lineage.loc[
        lineage["Metric_Key"] == "kpi.action_plans_total"
    ].iloc[0]
    assert action_lineage["Source_State"] == "partial"
    assert pd.isna(action_lineage["Metric_Value"])
    action_evidence = evidence.loc[
        evidence["Evidence_Key"] == "kpi.action_plans_total"
    ]
    assert set(action_evidence["Source_State"]) == {"partial"}
    assert action_evidence["Metric_Value"].isna().all()
    assert action_evidence["Contribution_Value"].isna().all()

    document = Document(word_path)
    assert len(document.inline_shapes) == 0
    word_text = _document_text(word_path)
    assert "Unavailable (Partial)" in word_text
    assert "Chart withheld" in word_text


def test_renewal_bundle_failure_keeps_empty_sources_failed_not_zero(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, names = _write_legacy_pair(
        tmp_path,
        family="renewal",
        marker="FAILED-BUNDLE",
        risk_score=25.9,
        risk_band="LOW",
    )
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        for key in ("action_plans", "customer_pulse", "success_priorities"):
            pd.DataFrame(columns=["Record_ID"]).to_excel(
                writer,
                sheet_name=names[key],
                index=False,
            )

    result = canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type="Renewal",
        manager_name="Local Fixture Manager",
        technology="Webex Calling",
        scope_type="customer",
        scope_value="Acme Corporation",
        days=90,
        as_of=AS_OF,
        partial_data_warnings=[
            {
                "dataset": "csconsole_bundle",
                "kind": "fetch_failed",
                "effect": "Renewal CSConsole sources were unavailable, not zero.",
            }
        ],
    )

    coverage = result["facts"]["source_coverage"].set_index("Source_Sheet")
    assert coverage.loc["Action_Plans", "Source_State"] == "failed"
    assert pd.isna(coverage.loc["Action_Plans", "Record_Count"])
    assert coverage.loc["Customer_Pulse", "Source_State"] == "failed"
    assert coverage.loc["Success_Priorities", "Source_State"] == "failed"
    assert coverage.loc["Adoption_Barriers", "Source_State"] == "partial"
    assert result["facts"]["chart_data"]["Value"].isna().all()

    with pd.ExcelFile(result["source_data_path"]) as workbook:
        report_info = pd.read_excel(workbook, sheet_name="Report_Info")
    states = report_info.loc[
        report_info["Item"].astype(str).str.startswith("Source_State:"),
        ["Item", "Value"],
    ].set_index("Item")["Value"]
    assert states["Source_State:Action_Plans"] == "failed"
    assert states["Source_State:Customer_Pulse"] == "failed"
    assert states["Source_State:Success_Priorities"] == "failed"
    assert states["Source_State:Adoption_Barriers"] == "partial"


def test_compact_legacy_0_10_risk_drift_is_explicitly_superseded() -> None:
    legacy = pd.DataFrame(
        [
            {
                "Customer": "Beta Industries",
                "Overall_Risk_Score": 2.4,
                "Risk_Score_0_10": 2.4,
                "Risk_Level": "LOW",
                "Risk_Band": "LOW",
            },
            {
                "Customer": "Acme Corporation",
                "Overall_Risk_Score": 2.7,
                "Risk_Score_0_10": 2.7,
                "Risk_Level": "LOW",
                "Risk_Band": "LOW",
            },
            {
                "Customer": "Gamma Public Sector",
                "Overall_Risk_Score": 6,
                "Risk_Score_0_10": 6,
                "Risk_Level": "HIGH",
                "Risk_Band": "HIGH",
            },
        ]
    )
    canonical = {
        "account_summary_all": [
            ["Beta Industries", "LOW", 22.3, 0, 0, 0, 0],
            ["Acme Corporation", "LOW", 25.6, 0, 0, 0, 0],
            ["Gamma Public Sector", "HIGH", 58.2, 0, 0, 1, 4],
        ]
    }

    assert adapter._risk_score(
        2.4,
        "Overall_Risk_Score",
        sheet_name="Risk_Summary",
    ) == 24.0
    assert adapter._risk_values_conflict(24.0, 22.3, kind="score") is True

    reconciled, dropped = adapter._supersede_compact_display_risk_scores(
        legacy,
        sheet_name="Risk_Summary",
    )
    assert dropped == ("Overall_Risk_Score", "Risk_Score_0_10")
    assert set(reconciled.columns) == {
        "Customer",
        "Risk_Level",
        "Risk_Band",
    }
    claims = adapter._validate_family_risk_claims(
        {"Risk_Summary": reconciled},
        canonical,
        info={},
        scope_type="team",
        scope_value="Entire team",
    )
    assert {(claim["customer"], claim["kind"]) for claim in claims} == {
        ("Beta Industries", "band"),
        ("Acme Corporation", "band"),
        ("Gamma Public Sector", "band"),
    }

    facts, _lineage, _counts = adapter._family_fact_projection(
        {"Risk_Summary": reconciled},
        family="compact",
        info={},
        scope_type="team",
        scope_value="Entire team",
    )
    assert not set(facts["Legacy_Source_Field"]).intersection(
        {"Overall_Risk_Score", "Risk_Score_0_10"}
    )


def test_compact_ambiguous_customer_risk_is_quarantined_with_partial_warning(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="compact",
        marker="AMBIGUOUS-RISK",
    )
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        pd.DataFrame(
            [
                {
                    "Customer": "ACME CORPORATION",
                    "Overall_Risk_Score": 0.3,
                    "Risk_Score_0_10": 0.3,
                    "Risk_Level": "HEALTHY",
                    "Risk_Band": "HEALTHY",
                },
                {
                    "Customer": "Acme Corporation",
                    "Overall_Risk_Score": 5.09,
                    "Risk_Score_0_10": 5.09,
                    "Risk_Level": "MEDIUM",
                    "Risk_Band": "MEDIUM",
                },
            ]
        ).to_excel(writer, sheet_name="Risk_Summary", index=False)

    result = canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type="Compact",
        manager_name="Local Fixture Manager",
        technology="All",
        scope_type="team",
        scope_value="Entire team",
        days=90,
        as_of=AS_OF,
    )

    assert result["contract"]["ok"] is True
    assert result["facts"]["risk_summary"]["source_state"] == "partial"
    assert any(
        warning.get("dataset") == "Subscriptions"
        and warning.get("kind") == "ambiguous_customer_name"
        and "family-risk cells were quarantined" in warning.get("effect", "")
        for warning in result["facts"]["partial_data_warnings"]
    )
    adapter_contract = result["facts"]["legacy_adapter"]
    assert adapter_contract["ambiguous_family_customer_labels"] == {
        "acmecorporation": ["ACME CORPORATION", "Acme Corporation"]
    }
    assert adapter_contract["ambiguous_family_risk_quarantine"] == {
        "rows": 2,
        "fields": ["Risk_Band", "Risk_Level"],
    }

    subscriptions = pd.read_excel(
        result["source_data_path"],
        sheet_name="Subscriptions",
    )
    # Risk_Summary remains family evidence only. It can no longer masquerade
    # as core subscription records, so the true subscription row survives and
    # the ambiguous summary risk fields are quarantined from the reported-fact
    # projection below.
    core_rows = subscriptions.loc[
        subscriptions["Legacy_Record_Type"].fillna("") == ""
    ]
    assert core_rows["Record_ID"].dropna().tolist() == ["SUB-AMBIGUOUS-RISK"]
    assert core_rows["Account ID"].dropna().tolist() == ["ACC-001"]
    reported = subscriptions.loc[
        subscriptions["Legacy_Record_Type"].fillna("")
        == "Family-specific reported fact"
    ]
    ambiguous_reported = reported.loc[
        reported["Legacy_Source_Sheet"] == "Risk_Summary"
    ]
    assert not set(ambiguous_reported["Legacy_Source_Field"]).intersection(
        {"Overall_Risk_Score", "Risk_Score_0_10", "Risk_Level", "Risk_Band"}
    )


def test_compact_duplicate_exact_label_contradiction_is_not_ambiguity() -> None:
    family_frame = pd.DataFrame(
        [
            {"Customer": "Acme Corporation", "Risk_Band": "HEALTHY"},
            {"Customer": "Acme Corporation", "Risk_Band": "LOW"},
        ]
    )
    assert adapter._ambiguous_family_customer_labels(
        {"Risk_Summary": family_frame}
    ) == {}

    with pytest.raises(
        CanonicalReportAdapterError,
        match="unresolved family-specific risk contradiction",
    ):
        adapter._validate_family_risk_claims(
            {"Risk_Summary": family_frame},
            {
                "account_summary_all": [
                    ["Acme Corporation", "LOW", 25.6, 0, 0, 0, 0]
                ]
            },
            info={},
            scope_type="team",
            scope_value="Entire team",
        )


def test_compact_explicit_0_100_risk_contradiction_still_fails_closed() -> None:
    legacy = pd.DataFrame(
        [
            {
                "Customer": "Beta Industries",
                "Overall_Risk_Score": 2.4,
                "Risk_Score_0_10": 2.4,
                "Risk_Score_0_100": 82,
                "Risk_Band": "LOW",
            }
        ]
    )
    reconciled, _dropped = adapter._supersede_compact_display_risk_scores(
        legacy,
        sheet_name="Risk_Summary",
    )

    with pytest.raises(
        CanonicalReportAdapterError,
        match=r"Risk_Score_0_100=82.*canonical Account_Summary.*22\.3",
    ):
        adapter._validate_family_risk_claims(
            {"Risk_Summary": reconciled},
            {
                "account_summary_all": [
                    ["Beta Industries", "LOW", 22.3, 0, 0, 0, 0]
                ]
            },
            info={},
            scope_type="team",
            scope_value="Entire team",
        )


@pytest.mark.parametrize(
    "legacy",
    [
        {
            "Customer": "Beta Industries",
            "Overall_Risk_Score": 2.4,
            "Risk_Score_0_10": 8.2,
            "Risk_Band": "LOW",
        },
        {
            "Customer": "Beta Industries",
            "Overall_Risk_Score": 8.2,
            "Risk_Score_0_10": 8.2,
            "Risk_Band": "LOW",
        },
    ],
)
def test_compact_supersession_refuses_internal_display_contradictions(
    legacy: dict[str, Any],
) -> None:
    with pytest.raises(
        CanonicalReportAdapterError,
        match="unresolved Compact display-risk contradiction",
    ):
        adapter._supersede_compact_display_risk_scores(
            pd.DataFrame([legacy]),
            sheet_name="Risk_Summary",
        )


def test_exact_family_risk_contradiction_fails_closed_and_keeps_sources(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="renewal",
        marker="RISK-CONTRADICTION",
        risk_score=82,
        risk_band="Critical",
    )
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        pd.DataFrame(
            [
                {
                    "Account": "Acme Corporation",
                    "Risk_Score_0_100": 51.4,
                    "Risk_Band": "MEDIUM",
                }
            ]
        ).to_excel(writer, sheet_name="Account_Summary", index=False)

    original_word = word_path.read_bytes()
    source_path = delivery.source_data_path_for_word(word_path)
    with pytest.raises(
        CanonicalReportAdapterError,
        match="unresolved family-specific risk contradiction",
    ) as error:
        canonicalize_legacy_artifacts(
            word_path,
            workbook_path,
            report_type="Renewal",
            manager_name="Local Fixture Manager",
            technology="Webex Calling",
            scope_type="customer",
            scope_value="Acme Corporation",
            days=90,
            as_of=AS_OF,
        )

    message = str(error.value)
    assert "Renewal_Summary" in message
    assert "Overall_Risk_Score=82" in message
    assert "Account_Summary" in message
    assert "51.4" in message
    assert "Critical" in message
    assert "MEDIUM" in message
    assert word_path.read_bytes() == original_word
    assert workbook_path.is_file()
    assert not source_path.exists()
    assert not list(tmp_path.glob(".*.canonical-*"))


def test_all_substantive_family_sheets_are_lineage_mapped_before_retirement(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="renewal",
        marker="FAMILY-FACTS",
    )
    fixture_score, fixture_band = _canonical_fixture_risk(
        _source_frames("FAMILY-FACTS"),
        _fixture_subscription("FAMILY-FACTS"),
    )
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        pd.DataFrame(
            [
                {
                    "Account": "Acme Corporation",
                    "Risk_Score_0_100": fixture_score,
                    "Risk_Band": fixture_band,
                }
            ]
        ).to_excel(writer, sheet_name="Account_Summary", index=False)
        pd.DataFrame(
            [
                {
                    "Customer": "Acme Corporation",
                    "Recommendation": "Complete the migration readiness review",
                    "Owner": "Jordan CSSM",
                }
            ]
        ).to_excel(writer, sheet_name="Recommendations", index=False)
        pd.DataFrame(
            [
                {
                    "Metric": "Renewal ARR at Risk",
                    "Value": 125000,
                    "Customer": "Acme Corporation",
                },
                {
                    "Metric": "Renewal Date",
                    "Value": "2026-12-31",
                    "Customer": "Acme Corporation",
                },
                {
                    "Metric": "Contract Status",
                    "Value": "At Risk",
                    "Customer": "Acme Corporation",
                },
                {
                    "Metric": "Remaining Term",
                    "Value": "5 months",
                    "Customer": "Acme Corporation",
                },
                {
                    "Metric": "Subscription ID",
                    "Value": "SUB-RENEW-001",
                    "Customer": "Acme Corporation",
                },
            ]
        ).to_excel(writer, sheet_name="Key_Metrics", index=False)

    result = canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type="Renewal",
        manager_name="Local Fixture Manager",
        technology="Webex Calling",
        scope_type="customer",
        scope_value="Acme Corporation",
        days=90,
        as_of=AS_OF,
    )

    assert result["legacy_workbook_retired"] is True
    assert not workbook_path.exists()
    assert result["contract"]["quarantined_sheets"] == []
    dispositions = result["contract"]["sheet_disposition"]
    assert dispositions["Renewal_Summary"] == "mapped_family_facts"
    assert dispositions["Account_Summary"] == "mapped_family_facts"
    assert dispositions["Recommendations"] == "mapped_family_facts"
    assert dispositions["Key_Metrics"] == "mapped_family_facts"

    source_path = Path(result["source_data_path"])
    subscriptions = pd.read_excel(source_path, sheet_name="Subscriptions")
    lineage = pd.read_excel(source_path, sheet_name="Metric_Lineage")
    evidence = pd.read_excel(source_path, sheet_name="Evidence_Links")
    reported = subscriptions.loc[
        subscriptions["Legacy_Record_Type"].fillna("")
        == "Family-specific reported fact"
    ]
    assert {
        "Renewal_Summary",
        "Account_Summary",
        "Recommendations",
        "Key_Metrics",
    }.issubset(set(reported["Legacy_Source_Sheet"]))
    assert "Complete the migration readiness review" in set(
        reported["Legacy_Fact_Value"].astype(str)
    )
    assert "125000" in set(reported["Legacy_Fact_Value"].astype(str))
    family_keys = set(reported["Metric_Key"].dropna().astype(str))
    family_lineage = lineage.loc[lineage["Metric_Key"].isin(family_keys)]
    assert set(family_lineage["Metric_Key"]) == family_keys
    assert set(family_lineage["Source_Sheet"]) == {"Subscriptions"}
    assert family_keys.issubset(set(evidence["Evidence_Key"].dropna().astype(str)))

    document = Document(word_path)
    report_specific_table = next(
        table
        for table in document.tables
        if [cell.text for cell in table.rows[0].cells]
        == ["Decision fact", "Reported value"]
    )
    report_specific_rows = [
        [cell.text for cell in row.cells]
        for row in report_specific_table.rows[1:]
    ]
    visible_values = {row[1] for row in report_specific_rows}
    assert {
        "SUB-RENEW-001",
        "125000",
        "2026-12-31",
        "At Risk",
        "5 months",
        "MEDIUM",
        "Complete the migration readiness review",
    }.issubset(visible_values)
    assert len(report_specific_rows) <= delivery.REPORT_SPECIFIC_FACT_LIMIT
    semantic = delivery.validate_word_semantics(result["facts"], document)
    assert semantic["ok"] is True
    assert semantic["report_specific_fact_count"] == len(report_specific_rows)
    assert semantic["report_specific_heading_validated"] is True

    # The semantic gate must reject a renderer that changes a visible value
    # while the exact evidence keys remain in the paired source contract.
    report_specific_table.rows[1].cells[1].text = "999999"
    tampered = delivery.validate_word_semantics(result["facts"], document)
    assert tampered["ok"] is False
    assert any(
        "Word visible cells differ from canonical facts: Decision fact"
        in error
        for error in tampered["errors"]
    )


def test_unmapped_substantive_sheet_is_quarantined_and_source_is_retained(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="renewal",
        marker="QUARANTINE",
    )
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        pd.DataFrame(
            [{"Material analyst fact": "Do not discard this source-owned note"}]
        ).to_excel(writer, sheet_name="Analyst_Notes", index=False)

    result = canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type="Renewal",
        manager_name="Local Fixture Manager",
        technology="Webex Calling",
        scope_type="customer",
        scope_value="Acme Corporation",
        days=90,
        as_of=AS_OF,
    )

    assert result["legacy_workbook_retired"] is False
    assert workbook_path.is_file()
    assert Path(result["source_data_path"]).is_file()
    assert result["contract"]["quarantined_sheets"] == ["Analyst_Notes"]
    assert (
        result["contract"]["sheet_disposition"]["Analyst_Notes"]
        == "quarantined_unmapped_substantive"
    )
    assert any(
        warning.get("dataset") == "Analyst_Notes"
        and warning.get("kind") == "legacy_sheet_quarantined"
        for warning in result["facts"]["partial_data_warnings"]
    )


def _write_placeholder_pair(tmp_path: Path) -> tuple[Path, Path]:
    word_path = tmp_path / "AdoptIQ_Report_Compact_Placeholders.docx"
    document = Document()
    document.add_paragraph("ORIGINAL WORD MUST SURVIVE FAILED VALIDATION")
    document.save(word_path)
    workbook_path = tmp_path / "legacy-placeholders.xlsx"
    with pd.ExcelWriter(workbook_path, engine="xlsxwriter") as writer:
        pd.DataFrame(
            [{"Item": "Export type", "Value": "compact"}]
        ).to_excel(writer, sheet_name="Report_Info", index=False)
        pd.DataFrame(
            [{"Customer": "Acme Corporation", "Risk_Band": "Unknown"}]
        ).to_excel(writer, sheet_name="Risk_Summary", index=False)
        pd.DataFrame(
            [
                {
                    "Status": "EMPTY",
                    "Dataset": "Action_Plans",
                    "Message": "No action plan records were returned for the selected scope.",
                    "Generated_At": AS_OF,
                }
            ]
        ).to_excel(writer, sheet_name="Action_Plans", index=False)
        pd.DataFrame(
            [
                {
                    "Status": "EMPTY",
                    "Dataset": "All_Adoption_Barriers",
                    "Message": "No barrier records were returned; PLACEHOLDER-AB-ROW.",
                    "Generated_At": AS_OF,
                }
            ]
        ).to_excel(writer, sheet_name="All_Adoption_Barriers", index=False)
        pd.DataFrame(
            [
                {
                    "Source State": "Unavailable",
                    "Record Count": 0,
                    "Coverage Note": "Support source unavailable; PLACEHOLDER-TAC-ROW.",
                }
            ]
        ).to_excel(writer, sheet_name="All_Support_Cases", index=False)
        pd.DataFrame(columns=["ID", "Customer Name"]).to_excel(
            writer,
            sheet_name="Customer_Pulse",
            index=False,
        )
        pd.DataFrame(
            [{"Message": "No data available for this analysis"}]
        ).to_excel(writer, sheet_name="Success_Priorities", index=False)
    return word_path, workbook_path


def test_placeholder_rows_are_removed_without_turning_unavailable_into_zero(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path = _write_placeholder_pair(tmp_path)

    result = canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type="Compact",
        manager_name="Local Fixture Manager",
        technology="All",
        scope_type="team",
        scope_value="Entire team",
        days=90,
        as_of=AS_OF,
    )

    frames = result["facts"]["frames"]
    assert cm.source_data_state(frames["action_plans"])["state"] == "zero"
    assert cm.source_data_state(frames["adoption_barriers"])["state"] == "zero"
    assert cm.source_data_state(frames["customer_pulse"])["state"] == "zero"
    assert cm.source_data_state(frames["tac_cases"])["state"] == "unavailable"

    source_path = Path(result["source_data_path"])
    barriers = pd.read_excel(source_path, sheet_name="Adoption_Barriers")
    tac_cases = pd.read_excel(source_path, sheet_name="TAC_Cases")
    report_info = pd.read_excel(source_path, sheet_name="Report_Info")
    try:
        info = dict(zip(report_info["Item"], report_info["Value"], strict=True))
    except TypeError:  # pragma: no cover - Python 3.9 compatibility in local fixtures
        info = dict(zip(report_info["Item"], report_info["Value"]))
    assert barriers.empty
    assert tac_cases.empty
    assert info["Source_State:Adoption_Barriers"] == "zero"
    assert info["Source_State:TAC_Cases"] == "unavailable"
    public_text = "\n".join(
        (
            _document_text(word_path),
            barriers.to_csv(index=False),
            tac_cases.to_csv(index=False),
        )
    )
    assert "PLACEHOLDER-AB-ROW" not in public_text
    assert "PLACEHOLDER-TAC-ROW" not in public_text


def test_legacy_workbook_already_at_source_data_path_is_replaced_not_deleted(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Compact/Renewal commonly reuse the final Source Data filename."""

    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="compact",
        marker="SAME-PATH",
    )
    source_path = delivery.source_data_path_for_word(word_path)
    workbook_path.replace(source_path)

    result = canonicalize_legacy_artifacts(
        word_path,
        source_path,
        report_type="Compact",
        manager_name="Local Fixture Manager",
        technology="All",
        scope_type="team",
        scope_value="Entire team",
        days=90,
        as_of=AS_OF,
    )

    assert Path(result["source_data_path"]) == source_path.resolve()
    assert source_path.is_file()
    assert result["legacy_workbook_retired"] is False
    assert result["contract"]["legacy_workbook_retired"] is False
    with pd.ExcelFile(source_path) as workbook:
        assert tuple(workbook.sheet_names) == delivery.SOURCE_DATA_SHEET_NAMES


def test_validation_failure_leaves_the_existing_word_target_untouched(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="compact",
        marker="FAIL-CLOSED",
    )
    original_word = word_path.read_bytes()
    source_path = delivery.source_data_path_for_word(word_path)

    monkeypatch.setattr(
        delivery,
        "validate_written_source_workbook",
        lambda *_args, **_kwargs: {"ok": False, "errors": ["forced validation failure"]},
    )
    with pytest.raises(CanonicalReportAdapterError, match="forced validation failure"):
        canonicalize_legacy_artifacts(
            word_path,
            workbook_path,
            report_type="Compact",
            manager_name="Local Fixture Manager",
            technology="All",
            scope_type="team",
            scope_value="Entire team",
            days=90,
            as_of=AS_OF,
        )

    assert word_path.read_bytes() == original_word
    assert workbook_path.is_file()
    assert not source_path.exists()
    assert not list(tmp_path.glob(".*.canonical-*"))


@pytest.mark.parametrize("failing_target", ["source_data", "word"])
def test_pair_install_failure_restores_both_original_targets(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    failing_target: str,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="compact",
        marker=f"PAIR-ROLLBACK-{failing_target}",
    )
    source_path = delivery.source_data_path_for_word(word_path)
    source_path.write_bytes(b"ORIGINAL PAIRED SOURCE TARGET")
    original_word = word_path.read_bytes()
    original_source = source_path.read_bytes()
    real_replace = adapter.os.replace
    blocked_target = source_path if failing_target == "source_data" else word_path

    def fail_target_install(source: Any, target: Any) -> None:
        source_candidate = Path(source)
        target_candidate = Path(target)
        if (
            target_candidate == blocked_target
            and ".canonical-" in source_candidate.name
        ):
            raise PermissionError("simulated locked final target")
        real_replace(source, target)

    monkeypatch.setattr(adapter.os, "replace", fail_target_install)

    with pytest.raises(
        CanonicalReportAdapterError,
        match=r"pair installation failed: PermissionError; rollback completed",
    ):
        canonicalize_legacy_artifacts(
            word_path,
            workbook_path,
            report_type="Compact",
            manager_name="Local Fixture Manager",
            technology="All",
            scope_type="team",
            scope_value="Entire team",
            days=90,
            as_of=AS_OF,
        )

    assert word_path.read_bytes() == original_word
    assert source_path.read_bytes() == original_source
    assert workbook_path.is_file()
    assert not list(tmp_path.glob(".*.canonical-*"))
    assert not list(tmp_path.glob(".*.rollback-*"))


def test_round149_footer_applies_via_r74_when_in_memory_stamp_fails(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Round 149 / Build 111: post-save R74 enforcer when apply_word_footer returns False."""

    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    monkeypatch.setattr(
        "_r68_build_label.apply_word_footer",
        lambda _doc: False,
    )
    word_path, workbook_path, _names = _write_legacy_pair(
        tmp_path,
        family="compact",
        marker="round149-footer-fallback",
    )

    result = canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type="Compact",
        manager_name="Local Fixture Manager",
        technology="All",
        scope_type="team",
        scope_value="Entire team",
        days=90,
        as_of=AS_OF,
    )

    assert result["contract"]["ok"] is True
    assert result["contract"]["footer"]["ok"] is True
    assert result["contract"]["footer"]["stamped_sections"] >= 1
