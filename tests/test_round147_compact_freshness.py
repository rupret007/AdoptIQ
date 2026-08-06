"""Round 147 Compact freshness never substitutes report-generation time."""

from __future__ import annotations
from source_shape_utils import assert_in_source

import json
from dataclasses import replace
from pathlib import Path

from docx import Document

import app_simple as app_mod
import ask_ai_grounded as grounded
import canonical_report_adapter as adapter
import decision_report_delivery as delivery
import manager_decision_workspace as workspace
from tests.test_round142_decision_report_delivery import _team_fixture
from tests.test_canonical_report_adapter import (
    _fake_chart_renderer,
    _write_legacy_pair,
)
from tests.test_round147_report_ai_usefulness import _request, _risk_group


SOURCE_CLOCK = "2026-08-03T21:00:00Z"
ATTEMPT_CLOCK = "2026-08-03T20:59:30Z"
EVALUATION_CLOCK = "2026-08-03T20:59:00Z"


def _freshness_facts(freshness: dict) -> dict:
    warning = freshness.get("warning")
    return delivery.build_report_facts(
        _team_fixture(),
        report_type="Compact",
        scope_type="team",
        scope_value="Dana Manager team",
        manager_name="Dana Manager",
        days=90,
        as_of=freshness["evaluation_as_of_utc"],
        data_as_of_utc=freshness["data_as_of_utc"],
        data_as_of_state=freshness["data_as_of_state"],
        data_as_of_detail=freshness["data_as_of_detail"],
        retrieval_attempted_at_utc=freshness["retrieval_attempted_at"],
        partial_data_warnings=[warning] if warning else [],
    )


def _document_text(document: Document) -> str:
    return "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )


def test_healthy_compact_prefetch_keeps_source_clock_available() -> None:
    freshness = app_mod._r147_compact_prefetch_freshness(
        {
            "data_retrieved_at": SOURCE_CLOCK,
            "attempted_at": ATTEMPT_CLOCK,
        },
        outcome="success",
        evaluation_clock=EVALUATION_CLOCK,
    )

    assert freshness == {
        "data_as_of_utc": SOURCE_CLOCK,
        "data_as_of_state": "available",
        "data_as_of_detail": (
            "Source retrieval completed at the recorded UTC timestamp."
        ),
        "data_retrieved_at": SOURCE_CLOCK,
        "retrieval_attempted_at": ATTEMPT_CLOCK,
        "evaluation_as_of_utc": SOURCE_CLOCK,
        "outcome": "success",
        "warning": None,
    }
    facts = _freshness_facts(freshness)
    assert facts["report_type"] == "Compact"
    assert facts["as_of_utc"] == "2026-08-03T21:00:00+00:00"
    assert facts["data_as_of_state"] == "available"
    assert facts["evaluation_as_of_utc"] == "2026-08-03T21:00:00+00:00"


def test_compact_timeout_preserves_known_source_clock_but_marks_partial() -> None:
    freshness = app_mod._r147_compact_prefetch_freshness(
        {
            "data_retrieved_at": SOURCE_CLOCK,
            "attempted_at": ATTEMPT_CLOCK,
        },
        outcome="timeout",
        evaluation_clock=EVALUATION_CLOCK,
    )

    assert freshness["data_as_of_utc"] == SOURCE_CLOCK
    assert freshness["data_as_of_state"] == "partial"
    assert freshness["evaluation_as_of_utc"] == SOURCE_CLOCK
    assert freshness["warning"]["kind"] == "freshness_partial"
    assert "timed" not in freshness["data_as_of_utc"].casefold()


def test_compact_runtime_without_source_clock_keeps_data_as_of_blank() -> None:
    freshness = app_mod._r147_compact_prefetch_freshness(
        {
            "data_retrieved_at": None,
            "attempted_at": ATTEMPT_CLOCK,
        },
        outcome="runtime",
        evaluation_clock=EVALUATION_CLOCK,
    )

    assert freshness["data_as_of_utc"] == ""
    assert freshness["data_retrieved_at"] == ""
    assert freshness["retrieval_attempted_at"] == ATTEMPT_CLOCK
    assert freshness["evaluation_as_of_utc"] == ATTEMPT_CLOCK
    assert freshness["data_as_of_state"] == "unavailable"
    assert freshness["warning"]["freshness"] == "unavailable"


def test_compact_no_retrieval_clocks_never_promotes_evaluation_clock() -> None:
    freshness = app_mod._r147_compact_prefetch_freshness(
        {},
        outcome="runtime",
        evaluation_clock=EVALUATION_CLOCK,
    )

    assert freshness["data_as_of_utc"] == ""
    assert freshness["retrieval_attempted_at"] == ""
    assert freshness["evaluation_as_of_utc"] == EVALUATION_CLOCK
    assert freshness["data_as_of_state"] == "unavailable"


def test_unavailable_freshness_reaches_word_workbook_workspace_and_ai_binding(
    tmp_path: Path,
) -> None:
    freshness = app_mod._r147_compact_prefetch_freshness(
        {"attempted_at": ATTEMPT_CLOCK},
        outcome="runtime",
        evaluation_clock=EVALUATION_CLOCK,
    )
    facts = _freshness_facts(freshness)
    sheets = delivery.build_source_data_sheets(facts)
    info = sheets["Report_Info"].set_index("Item")

    assert info.loc["Data_As_Of_UTC", "Value"] == ""
    assert info.loc["Data_As_Of_State", "Value"] == "unavailable"
    assert (
        info.loc["Retrieval_Attempted_At_UTC", "Value"]
        == "2026-08-03T20:59:30+00:00"
    )
    assert (
        info.loc["Evaluation_As_Of_UTC", "Value"]
        == "2026-08-03T20:59:30+00:00"
    )

    word_text = _document_text(delivery.build_concise_word_document(facts))
    assert "Data as of unavailable" in word_text
    assert "retrieval attempted 2026-08-03 20:59 UTC" in word_text
    assert "Data as of 2026-08-03 20:59 UTC" not in word_text

    source_path = tmp_path / "AdoptIQ_Source_Data_Compact_Freshness.xlsx"
    delivery.write_source_data_workbook(source_path, sheets)
    loaded = workspace.load_workbook_snapshot(source_path)
    snapshot = workspace.snapshot_from_status(
        {
            "analysis_id": "compact-freshness-runtime",
            "status": "completed",
            "report_type": "compact",
            "manager": "Dana Manager",
            "technology": "All",
            "days": 90,
            "data_as_of_utc": "",
            "data_as_of_state": "unavailable",
            "retrieval_attempted_at_utc": ATTEMPT_CLOCK,
        },
        loaded,
    )
    binding = workspace.ask_ai_binding(snapshot)
    assert snapshot["data_as_of_utc"] == ""
    assert snapshot["data_as_of_state"] == "unavailable"
    assert snapshot["retrieval_attempted_at_utc"].startswith("2026-08-03T20:59:30")
    assert snapshot["source_states"]["Data_Freshness"] == "unavailable"
    assert binding["data_as_of_utc"] == ""
    assert binding["data_as_of_state"] == "unavailable"
    assert binding["source_states"]["Data_Freshness"] == "unavailable"


def test_compact_adapter_keeps_unavailable_data_as_of_blank(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    word_path, workbook_path, _ = _write_legacy_pair(
        tmp_path,
        family="compact",
        marker="FRESHNESS-RUNTIME",
    )
    freshness = app_mod._r147_compact_prefetch_freshness(
        {"attempted_at": ATTEMPT_CLOCK},
        outcome="runtime",
        evaluation_clock=EVALUATION_CLOCK,
    )

    result = adapter.canonicalize_legacy_artifacts(
        word_path,
        workbook_path,
        report_type="Compact",
        manager_name="Dana Manager",
        technology="All",
        scope_type="team",
        scope_value="Dana Manager team",
        days=90,
        as_of=freshness["evaluation_as_of_utc"],
        data_as_of_utc=freshness["data_as_of_utc"],
        data_as_of_state=freshness["data_as_of_state"],
        data_as_of_detail=freshness["data_as_of_detail"],
        retrieval_attempted_at_utc=freshness["retrieval_attempted_at"],
        partial_data_warnings=[freshness["warning"]],
    )

    assert result["contract"]["ok"] is True
    assert result["facts"]["as_of_utc"] == ""
    assert result["facts"]["data_as_of_state"] == "unavailable"
    assert result["facts"]["retrieval_attempted_at_utc"].startswith(
        "2026-08-03T20:59:30"
    )
    assert "Data as of unavailable" in _document_text(Document(result["word_path"]))


def test_compact_branch_order_finalizes_after_all_prefetch_outcomes() -> None:
    source = Path(app_mod.__file__).read_text(encoding="utf-8")
    start = source.index("def run_compact_analysis(")
    holder = source.index("_compact_prefetch_meta = {", start)
    timeout = source.index("except FutureTimeoutError:", holder)
    runtime = source.index("except Exception as e:", timeout)
    finalize = source.index("finalize freshness once", runtime)
    report_context = source.index("_r23_ctx = {", finalize)

    assert start < holder < timeout < runtime < finalize < report_context
    canonical_region_start = source.index(
        "_r147_compact_evaluation_as_of =", report_context
    )
    canonical_region_end = source.index(
        "_r147_compact = _r147_canonicalize_legacy_delivery(",
        canonical_region_start,
    )
    assert "_now_utc_iso_z()" not in source[
        canonical_region_start:canonical_region_end
    ]


def test_progress_discloses_unavailable_freshness_and_attempt_clock(client) -> None:
    analysis_id = "round147-compact-freshness-progress"
    with app_mod.analysis_status_lock:
        previous = dict(app_mod.analysis_status)
        app_mod.analysis_status[analysis_id] = {
            "status": "completed",
            "progress": 100,
            "message": "Compact report generated with partial data.",
            "current_step": "Complete",
            "manager": "Dana Manager",
            "technology": "All",
            "days": 90,
            "start_time": "2026-08-03T20:58:00Z",
            "completion_time": "2026-08-03T21:05:00Z",
            "data_retrieved_at": "",
            "data_as_of_utc": "",
            "data_as_of_state": "unavailable",
            "retrieval_attempted_at_utc": ATTEMPT_CLOCK,
            "evaluation_as_of_utc": "1999-01-01T00:00:00Z",
            "report_type": "compact",
        }
    try:
        response = client.get(f"/progress/{analysis_id}")
    finally:
        with app_mod.analysis_status_lock:
            app_mod.analysis_status.clear()
            app_mod.analysis_status.update(previous)

    assert response.status_code == 200
    body = response.get_data(as_text=True)
    assert_in_source(body, 'id="provenance-data-as-of">unavailable</span>', label='body')
    assert_in_source(body, 'id="provenance-data-state">(unavailable)</span>', label='body')
    assert_in_source(body, "2026-08-03T20:59:30", label='body')
    assert "1999-01-01T00:00:00" not in body
    assert_in_source(body, "not a source data-as-of claim", label='body')


def test_report_bound_ask_ai_names_unavailable_freshness_without_fake_as_of() -> None:
    request = _request(
        "What is Acme's risk score?",
        [_risk_group("Acme", 54.2)],
    )
    bundle = json.loads(request.report_fact_bundle)
    bundle["data_as_of_utc"] = ""
    bundle["data_as_of_state"] = "unavailable"
    bundle["retrieval_attempted_at_utc"] = ATTEMPT_CLOCK
    bundle["source_states"]["Data_Freshness"] = "unavailable"
    bundle["exact_evidence"]["data_as_of_utc"] = ""
    request = replace(
        request,
        data_as_of_utc="",
        report_fact_bundle=json.dumps(bundle, sort_keys=True),
    )

    result = grounded._r146_report_bound_snapshot_answer(  # noqa: SLF001
        request,
        {"scope_type": "team", "report_analysis_id": request.report_analysis_id},
    )

    assert result["ok"] is True
    assert result["response_state"] == "partial"
    assert result["confidence"]["level"] != "High"
    assert "source data-as-of unavailable" in result["answer"]
    assert f"retrieval was attempted at {ATTEMPT_CLOCK}" in result["answer"]
    assert "Data_Freshness: unavailable" in result["partial_data_warnings"]
    assert EVALUATION_CLOCK not in result["answer"]
