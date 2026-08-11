"""Round 147 exact claim-to-source-row evidence contract."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import decision_report_delivery as delivery
import manager_decision_workspace as workspace
from report_iteration_loop import evaluate_report_quality
from tests.test_round142_decision_report_delivery import (
    AS_OF,
    _facts,
    _fake_chart_renderer,
    _team_fixture,
)


def test_every_visible_lineage_key_resolves_and_count_claims_reconcile() -> None:
    facts = _facts()
    sheets = delivery.build_source_data_sheets(facts)

    result = delivery.validate_evidence_links(sheets)
    assert result["ok"], result["errors"]
    assert result["lineage_keys"] == len(facts["metric_lineage"])
    assert result["resolved_rows"] > result["lineage_keys"]
    assert set(facts["metric_lineage"]["Metric_Key"]).issubset(
        set(sheets["Evidence_Links"]["Evidence_Key"])
    )
    overdue = sheets["Evidence_Links"].loc[
        sheets["Evidence_Links"]["Evidence_Key"] == "kpi.action_plans_overdue"
    ]
    assert set(overdue["Record_ID"]) == {"AP-001", "AP-003"}
    assert set(overdue["Source_Row_Number"]) == {2, 4}


def test_written_evidence_lookup_returns_exact_bounded_rows(tmp_path: Path) -> None:
    facts = _facts()
    source_path = tmp_path / "AdoptIQ_Source_Data_Round147.xlsx"
    delivery.write_source_data_workbook(
        source_path,
        delivery.build_source_data_sheets(facts),
    )
    snapshot = workspace.load_workbook_snapshot(source_path)

    evidence = workspace.load_workbook_evidence(
        source_path,
        "chart.activity_trend.action_plans.2026-07-06",
        expected_fingerprint=snapshot["fact_fingerprint"],
        limit=3,
    )

    assert snapshot["evidence_available"] is True
    assert evidence["total_records"] == 7
    assert len(evidence["records"]) == 3
    assert evidence["truncated"] is True
    assert all(record["source_sheet"] == "Action_Plans" for record in evidence["records"])
    assert all(record["record_id"].startswith("AP-") for record in evidence["records"])
    assert any("first 3 of 7" in item for item in evidence["limitations"])


def test_report_bound_selector_returns_only_verified_exact_rows(tmp_path: Path) -> None:
    facts = _facts()
    source_path = tmp_path / "AdoptIQ_Source_Data_Ask_AI.xlsx"
    delivery.write_source_data_workbook(
        source_path,
        delivery.build_source_data_sheets(facts),
    )
    snapshot = workspace.load_workbook_snapshot(source_path)

    selected = workspace.select_report_bound_evidence(
        source_path,
        snapshot,
        "Show the actual overdue Action Plan records and their owners.",
        preferred_evidence_key="kpi.action_plans_overdue",
        max_keys=3,
        max_records=5,
    )

    assert selected["schema"] == "report-bound-evidence/v1"
    assert selected["evidence_contract"] == "canonical-evidence-links/v1"
    assert selected["fact_fingerprint"] == snapshot["fact_fingerprint"]
    assert selected["groups"][0]["evidence_key"] == "kpi.action_plans_overdue"
    assert selected["groups"][0]["metric_value"] == 2
    assert selected["groups"][0]["unit"] == "records"
    assert {record["record_id"] for record in selected["groups"][0]["records"]} == {
        "AP-001", "AP-003",
    }
    assert all(
        record["source_sheet"] == "Action_Plans"
        for record in selected["groups"][0]["records"]
    )


def test_report_bound_selector_rejects_unbound_evidence_key(tmp_path: Path) -> None:
    facts = _facts()
    source_path = tmp_path / "AdoptIQ_Source_Data_Ask_AI_Missing.xlsx"
    delivery.write_source_data_workbook(
        source_path,
        delivery.build_source_data_sheets(facts),
    )
    snapshot = workspace.load_workbook_snapshot(source_path)

    with pytest.raises(workspace.EvidenceNotFoundError):
        workspace.select_report_bound_evidence(
            source_path,
            snapshot,
            "Show evidence",
            preferred_evidence_key="recommendation.not-in-this-report",
        )


def test_evidence_lookup_fails_closed_after_source_row_tamper(tmp_path: Path) -> None:
    facts = _facts()
    source_path = tmp_path / "AdoptIQ_Source_Data_Tampered.xlsx"
    delivery.write_source_data_workbook(
        source_path,
        delivery.build_source_data_sheets(facts),
    )
    workbook = load_workbook(source_path)
    action_sheet = workbook["Action_Plans"]
    title_column = next(
        cell.column for cell in action_sheet[1] if cell.value == "AdoptIQ_Title"
    )
    action_sheet.cell(2, title_column).value = "Substituted after generation"
    workbook.save(source_path)
    workbook.close()

    with pytest.raises(
        workspace.EvidenceIntegrityError,
        match="digest failed|integrity failed",
    ):
        workspace.load_workbook_evidence(source_path, "kpi.action_plans_overdue")


def test_evidence_lookup_fails_closed_after_link_metric_tamper(tmp_path: Path) -> None:
    facts = _facts()
    source_path = tmp_path / "AdoptIQ_Source_Data_Link_Tampered.xlsx"
    delivery.write_source_data_workbook(
        source_path,
        delivery.build_source_data_sheets(facts),
    )
    workbook = load_workbook(source_path)
    links = workbook["Evidence_Links"]
    headers = {cell.value: cell.column for cell in links[1]}
    target_row = next(
        row
        for row in range(2, links.max_row + 1)
        if links.cell(row, headers["Evidence_Key"]).value == "kpi.action_plans_overdue"
    )
    links.cell(target_row, headers["Metric_Value"]).value = 999
    workbook.save(source_path)
    workbook.close()

    with pytest.raises(workspace.EvidenceIntegrityError, match="Evidence_Links"):
        workspace.load_workbook_evidence(
            source_path,
            "kpi.action_plans_overdue",
            expected_fingerprint=delivery.fact_contract_fingerprint(facts),
        )


def test_unavailable_source_is_explicit_derivation_not_fake_zero() -> None:
    fixture = _team_fixture()
    failed = pd.DataFrame()
    failed.attrs["fetch_error"] = "controlled fixture outage"
    fixture["Alex Rivera"]["adoption_barriers"] = failed
    facts = delivery.build_report_facts(
        fixture,
        report_type="Leader",
        scope_type="team",
        scope_value="Dana Manager team",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
    )
    sheets = delivery.build_source_data_sheets(facts)
    links = sheets["Evidence_Links"].loc[
        sheets["Evidence_Links"]["Evidence_Key"] == "kpi.adoption_barriers"
    ]

    assert len(links) == 1
    assert links.iloc[0]["Evidence_Role"] == "unavailable_state"
    assert links.iloc[0]["Source_State"] == "failed"
    assert pd.isna(links.iloc[0]["Source_Row_Number"])
    assert pd.isna(links.iloc[0]["Contribution_Value"])


def test_complete_summary_sheets_publish_values_with_per_field_states() -> None:
    sheets = delivery.build_source_data_sheets(_facts())
    account = sheets["Account_Summary"].loc[
        sheets["Account_Summary"]["Account"] == "Acme Corporation"
    ].iloc[0]
    member = sheets["Member_Summary"].loc[
        sheets["Member_Summary"]["Team_Member"] == "Alex Rivera"
    ].iloc[0]

    assert account["Risk_Score_0_100"] is not None
    assert account["Open_AP"] == 7
    assert account["Critical_High_Barriers"] == 1
    assert account["TAC_Cases"] == 1
    assert {
        account["Risk_Band_Source_State"],
        account["Risk_Score_0_100_Source_State"],
        account["Open_AP_Source_State"],
        account["Overdue_AP_Source_State"],
        account["Critical_High_Barriers_Source_State"],
        account["TAC_Cases_Source_State"],
    } == {"available"}
    assert member["Customers"] == 2
    assert member["Open_AP"] == 10
    assert member["Barriers"] == 1
    assert member["TAC_Cases"] == 1
    assert {
        member["Customers_Source_State"],
        member["Open_AP_Source_State"],
        member["Overdue_AP_Source_State"],
        member["Barriers_Source_State"],
        member["TAC_Cases_Source_State"],
    } == {"available"}


@pytest.mark.parametrize(
    ("source_key", "source_state", "account_fields", "member_fields"),
    [
        (
            "action_plans",
            "partial",
            ("Risk_Band", "Risk_Score_0_100", "Open_AP", "Overdue_AP"),
            ("Customers", "Open_AP", "Overdue_AP"),
        ),
        (
            "adoption_barriers",
            "partial",
            ("Risk_Band", "Risk_Score_0_100", "Critical_High_Barriers"),
            ("Customers", "Barriers"),
        ),
        (
            "tac_cases",
            "failed",
            ("Risk_Band", "Risk_Score_0_100", "TAC_Cases"),
            ("Customers", "TAC_Cases"),
        ),
    ],
)
def test_computed_summary_sheets_fail_closed_per_degraded_contributor(
    source_key,
    source_state,
    account_fields,
    member_fields,
) -> None:
    fixture = _team_fixture()
    source = fixture["Alex Rivera"][source_key].copy()
    if source_state == "failed":
        source = source.iloc[0:0].copy()
        source.attrs["fetch_error"] = "controlled source failure"
    else:
        source.attrs["fetch_error"] = "controlled partial source"
        source.attrs["fetch_error_partial"] = True
    fixture["Alex Rivera"][source_key] = source
    facts = delivery.build_report_facts(
        fixture,
        report_type="Leader",
        scope_type="team",
        scope_value="Dana Manager team",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
    )
    sheets = delivery.build_source_data_sheets(facts)
    account = sheets["Account_Summary"].loc[
        sheets["Account_Summary"]["Account"] == "Acme Corporation"
    ].iloc[0]
    member = sheets["Member_Summary"].loc[
        sheets["Member_Summary"]["Team_Member"] == "Alex Rivera"
    ].iloc[0]

    assert all(pd.isna(account[field]) for field in account_fields)
    assert all(pd.isna(member[field]) for field in member_fields)
    assert account["Risk_Score_0_100_Source_State"] == "partial"
    assert member["Customers_Source_State"] == "partial"

    unaffected_account_fields = {
        "Open_AP": "Open_AP_Source_State",
        "Overdue_AP": "Overdue_AP_Source_State",
        "Critical_High_Barriers": "Critical_High_Barriers_Source_State",
        "TAC_Cases": "TAC_Cases_Source_State",
    }
    for field, state_field in unaffected_account_fields.items():
        if field not in account_fields:
            assert pd.notna(account[field])
            assert account[state_field] in {"available", "zero"}


def test_populated_chart_render_failure_prevents_report_publication(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    facts = _facts()
    monkeypatch.setattr(
        delivery,
        "_render_chart_image",
        lambda _chart_id, _rows, _target: False,
    )

    with pytest.raises(ValueError, match="Expected populated chart"):
        delivery.build_concise_word_document(facts)


def test_visible_action_owner_tamper_fails_semantic_word_contract() -> None:
    facts = _facts()
    sheets = delivery.build_source_data_sheets(facts)
    document = delivery.build_concise_word_document(facts)
    action_table = next(
        table
        for table in document.tables
        if [cell.text for cell in table.rows[0].cells]
        == ["Record ID", "Account", "Owner", "Status", "Due", "Age (days)", "Priority"]
    )
    action_table.rows[1].cells[2].text = "Invented Manager Owner"

    semantic = delivery.validate_word_semantics(facts, document)
    contract = delivery.validate_cross_artifact_contract(facts, sheets, document)

    assert semantic["ok"] is False
    assert any(
        "visible cells differ" in error and "Owner" in error
        for error in semantic["errors"]
    )
    assert contract["ok"] is False
    assert contract["word_semantics"]["ok"] is False


def test_leader_technology_scope_is_visible_and_semantically_enforced(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    facts = _facts()
    facts["technology"] = "All Contact Center"
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    document = delivery.build_concise_word_document(facts)

    assert document.paragraphs[1].text == (
        "Dana Manager • Team: Dana Manager team • "
        "Technology: All Contact Center • 90-day window"
    )
    assert delivery.validate_word_semantics(facts, document)["ok"] is True

    document.paragraphs[1].text = document.paragraphs[1].text.replace(
        "All Contact Center",
        "All",
    )
    semantic = delivery.validate_word_semantics(facts, document)

    assert semantic["ok"] is False
    assert semantic["scope_subtitle_validated"] is False
    assert any("scope subtitle differs" in error for error in semantic["errors"])


def test_partial_risk_decisions_carry_claim_level_state_and_are_validated(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    facts = _facts()
    facts["risk_summary"]["source_state"] = "partial"
    facts["risk_summary"]["source_state_detail"] = "Customer Pulse is partial"
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    document = delivery.build_concise_word_document(facts)
    header = (
        "Account",
        "Risk",
        "Score",
        "Evidence state",
        # Round 153 / Tier 1: a "Top risk drivers" column now sits between
        # the evidence state and the next action.
        "Top risk drivers",
        "Evidence-backed next action",
    )
    table = next(
        table
        for table in document.tables
        if tuple(cell.text for cell in table.rows[0].cells) == header
    )

    assert all("(Partial)" in row.cells[1].text for row in table.rows[1:])
    assert all("(Partial)" in row.cells[2].text for row in table.rows[1:])
    assert all(row.cells[3].text == "Partial" for row in table.rows[1:])
    assert delivery.validate_word_semantics(facts, document)["ok"] is True

    table.rows[1].cells[3].text = "Available"
    semantic = delivery.validate_word_semantics(facts, document)

    assert semantic["ok"] is False
    assert any(
        "visible cells differ" in error and "Evidence state" in error
        for error in semantic["errors"]
    )


@pytest.mark.parametrize(
    ("header", "value_column"),
    [
        (("Metric", "Value", "Lineage key"), 1),
        (("Source", "State", "Distinct records"), 2),
        (("Lifecycle", "Distinct plans"), 1),
        (("Team member", "Customers", "Open AP", "Overdue AP", "Barriers", "TAC"), 1),
    ],
)
def test_visible_summary_numeric_tamper_fails_semantic_word_contract(
    header: tuple[str, ...],
    value_column: int,
) -> None:
    facts = _facts()
    sheets = delivery.build_source_data_sheets(facts)
    document = delivery.build_concise_word_document(facts)
    table = next(
        table
        for table in document.tables
        if tuple(cell.text for cell in table.rows[0].cells) == header
    )
    table.rows[1].cells[value_column].text = "999999"

    contract = delivery.validate_cross_artifact_contract(facts, sheets, document)

    assert contract["ok"] is False
    assert contract["word_semantics"]["ok"] is False
    assert any("visible cells differ" in error for error in contract["word_semantics"]["errors"])


def test_visible_account_summary_numeric_tamper_fails_semantic_word_contract() -> None:
    facts = delivery.build_report_facts(
        _team_fixture(),
        report_type="Leader",
        scope_type="customer",
        scope_value="Acme Corporation",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
    )
    sheets = delivery.build_source_data_sheets(facts)
    document = delivery.build_concise_word_document(facts)
    header = (
        "Account",
        "Risk band",
        "Risk score",
        "Open AP",
        "Overdue AP",
        "Critical/high barriers",
        "TAC",
    )
    table = next(
        table
        for table in document.tables
        if tuple(cell.text for cell in table.rows[0].cells) == header
    )
    table.rows[1].cells[2].text = "999999"

    contract = delivery.validate_cross_artifact_contract(facts, sheets, document)

    assert contract["ok"] is False
    assert contract["word_semantics"]["ok"] is False
    assert any(
        "visible cells differ" in error and "Risk score" in error
        for error in contract["word_semantics"]["errors"]
    )


def test_concise_word_layout_uses_compact_readable_body_spacing() -> None:
    document = delivery.build_concise_word_document(_facts())
    normal = document.styles["Normal"]

    assert normal.font.size is not None
    assert normal.font.size.pt == pytest.approx(10.0)
    assert normal.paragraph_format.space_after is not None
    assert normal.paragraph_format.space_after.pt == pytest.approx(2.0)


def test_unassigned_portfolio_records_do_not_masquerade_as_team_members() -> None:
    fixture = _team_fixture()
    fixture["__Unassigned_Portfolio__"] = {
        "action_plans": pd.DataFrame(
            [
                {
                    "ID": "AP-UNASSIGNED",
                    "BU_NAME": "Unassigned Customer",
                    "STATUS_C": "Open",
                    "OWNER_NAME": "",
                }
            ]
        ),
    }
    facts = delivery.build_report_facts(
        fixture,
        report_type="Leader",
        scope_type="team",
        scope_value="Dana Manager team",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
    )
    sheets = delivery.build_source_data_sheets(facts)
    document = delivery.build_concise_word_document(facts)

    assert facts["unassigned_portfolio_summary"][0] == "Unassigned / Portfolio"
    assert all(row[0] != "Unassigned / Portfolio" for row in facts["member_summary_all"])
    assert "Unassigned / Portfolio" not in set(sheets["Member_Summary"]["Team_Member"])
    assert "Unassigned / Portfolio" in set(sheets["Action_Plans"]["CSSM"])
    assert any(
        "excluded from this ranking" in paragraph.text
        for paragraph in document.paragraphs
    )


def test_member_summary_has_adjacent_exact_lineage_and_passes_strict_audit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fixture = _team_fixture()
    fixture["Alex Rivera"]["action_plans"] = fixture["Alex Rivera"][
        "action_plans"
    ].assign(
        SUBJECT_C="Resolve adoption dependency",
        NEXT_ACTION_C="Confirm accountable owner",
    )
    fixture["__legacy_unassigned__"] = {
        "action_plans": pd.DataFrame(
            [
                {
                    "ID": "AP-UNASSIGNED-AUDIT",
                    "BU_NAME": "Unassigned Customer",
                    "STATUS_C": "Open",
                    "OWNER_NAME": "",
                }
            ]
        ),
    }
    facts = delivery.build_report_facts(
        fixture,
        report_type="Compact",
        scope_type="team",
        scope_value="Local Fixture Manager team",
        manager_name="Local Fixture Manager",
        days=90,
        as_of=AS_OF,
    )
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    document = delivery.build_concise_word_document(facts)
    member_table = next(
        table
        for table in document.tables
        if [cell.text for cell in table.rows[0].cells]
        == ["Team member", "Customers", "Open AP", "Overdue AP", "Barriers", "TAC"]
    )

    next_element = member_table._tbl.getnext()  # noqa: SLF001
    adjacent_text = "".join(next_element.itertext()).strip()
    assert adjacent_text.startswith("[Source: Source Data File → Metric_Lineage /")
    assert "summary.member.* → Member_Summary" in adjacent_text

    semantic = delivery.validate_word_semantics(facts, document)
    assert semantic["ok"], semantic["errors"]

    word_path = tmp_path / "AdoptIQ_Report_Compact_Member_Audit.docx"
    source_path = delivery.source_data_path_for_word(word_path)
    document.save(word_path)
    delivery.write_source_data_workbook(
        source_path,
        delivery.build_source_data_sheets(facts),
    )
    payload, gate = evaluate_report_quality(
        docx_path=word_path,
        xlsx_path=source_path,
        scenario_key="compact",
        strict=True,
    )

    assert payload["unbacked_metric_claim_count"] == 0, payload[
        "unbacked_metric_claims"
    ]
    assert gate.passed, gate.details

    member_table.rows[1].cells[1].text = "999"
    tampered = delivery.validate_word_semantics(facts, document)
    assert tampered["ok"] is False
    assert any("visible cells differ" in error for error in tampered["errors"])


def test_punctuation_colliding_member_names_get_distinct_evidence_identities() -> None:
    team_data = {
        "Alex Rivera": {
            "subscriptions": pd.DataFrame(
                [
                    {
                        "SUBSCRIPTION_ID": "SUB-ONE",
                        "ACCOUNT_ID_C": "ACC-ONE",
                        "BU_NAME": "Customer One",
                    }
                ]
            )
        },
        "Alex-Rivera": {
            "subscriptions": pd.DataFrame(
                [
                    {
                        "SUBSCRIPTION_ID": "SUB-TWO",
                        "ACCOUNT_ID_C": "ACC-TWO",
                        "BU_NAME": "Customer Two",
                    }
                ]
            )
        },
    }
    facts = delivery.build_report_facts(
        team_data,
        report_type="Leader",
        scope_type="team",
        scope_value="Collision team",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
    )
    sheets = delivery.build_source_data_sheets(facts)
    document = delivery.build_concise_word_document(facts)
    contract = delivery.validate_cross_artifact_contract(facts, sheets, document)

    member_keys = sheets["Member_Summary"]["Metric_Key"].tolist()
    member_lineage = sheets["Metric_Lineage"].loc[
        sheets["Metric_Lineage"]["Metric_Key"].str.startswith("summary.member.")
    ]
    assert len(member_keys) == 2
    assert len(set(member_keys)) == 2
    assert member_lineage["Metric_Key"].is_unique
    assert contract["ok"], contract["errors"]
