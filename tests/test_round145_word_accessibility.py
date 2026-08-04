"""Round 145 final Word accessibility normalization contracts."""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from report_word_styling import add_heading_on_new_page, apply_document_accessibility


def test_new_page_heading_uses_paragraph_format_instead_of_break_paragraph() -> None:
    doc = Document()
    doc.add_paragraph("Full preceding page content")

    heading = add_heading_on_new_page(doc, "Next section", level=1)

    assert heading.paragraph_format.page_break_before is True
    assert len(doc.paragraphs) == 2
    assert all('w:type="page"' not in paragraph._p.xml for paragraph in doc.paragraphs)


def test_accessibility_pass_repairs_heading_skips_and_table_header_semantics() -> None:
    doc = Document()
    doc.add_paragraph("Report title", style="Title")
    doc.add_paragraph("Primary section", style="Heading 1")
    skipped = doc.add_paragraph("Skipped subsection", style="Heading 3")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "ID"
    table.rows[0].cells[1].text = "Status"

    changed = apply_document_accessibility(doc)

    assert skipped.style.name == "Heading 2"
    assert changed == {"heading_levels": 1, "table_headers": 1}
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    assert header_properties.find(qn("w:tblHeader")) is not None
    for row in table.rows:
        assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None


def test_accessibility_pass_is_idempotent() -> None:
    doc = Document()
    doc.add_paragraph("Section", style="Heading 1")
    doc.add_table(rows=1, cols=1)

    first = apply_document_accessibility(doc)
    second = apply_document_accessibility(doc)

    assert first["table_headers"] == 1
    assert second == {"heading_levels": 0, "table_headers": 0}
