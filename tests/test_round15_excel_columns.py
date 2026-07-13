"""Round 15 / Phase 1 -- Excel column-curation tests.

Locks the customer-facing column SSoT (``report_export_schema``) and the
wiring of the four primary Excel writers in ``adoptiq_backend.py`` and
``app_simple.py`` so a future change can't silently re-leak Salesforce
plumbing into the customer workbook.

Two layers of coverage:

  * marker tests -- every Round 15 / Phase 1.x source comment is present
    in the file we expect, so reverting the wiring is a noisy diff.
  * behavioral tests -- the SSoT actually drops the SF/ETL/internal
    columns the gold fixture revealed.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from report_export_schema import (
    CURATED_COLUMNS,
    INTERNAL_COLUMN_DENYLIST,
    INTERNAL_COLUMN_PREFIXES,
    SHEET_HEADER_RENAMES,
    apply_export_schema,
    filter_columns,
    is_internal_column,
)

REPO_ROOT = Path(__file__).resolve().parent.parent


# ---------------------------------------------------------------------------
# SSoT correctness
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "name",
    [
        "IS_DELETED",
        "ISDELETED",
        "MAY_EDIT",
        "MAYEDIT",
        "IS_LOCKED",
        "ISLOCKED",
        "RECORD_TYPE_ID",
        "SYSTEM_MODSTAMP",
        "SYSTEMMODSTAMP",
        "LAST_VIEWED_DATE",
        "LAST_REFERENCED_DATE",
        "CONNECTION_RECEIVED_ID",
        "CONNECTION_SENT_ID",
        "ETL_ID",
        "DELETE_FLAG",
        "RECURSIVE_FLAG",
        "EDWSF_BATCH_ID",
        "EDWSF_CREATE_DTM",
        "_FIVETRAN_SYNCED",
        "_FIVETRAN_DELETED",
        "_stale_storage",
        "_from_storage",
        "_window_meta",
        "publication_id",
        "Title_Lower",
        "Desc_Lower",
        "col_2",
        "col_32",
    ],
)
def test_phase_1_ssot_denies_known_internal_columns(name):
    """Round 15 / Phase 1: SSoT denylist covers every leak from the gold xlsx."""
    assert is_internal_column(name), f"{name!r} should be on the SSoT denylist"


@pytest.mark.parametrize(
    "name",
    [
        "ID",
        "NAME",
        "SUBJECT_C",
        "SEVERITY_C",
        "AB_STATUS_C",
        "customer_name",
        "open_age_days",
        "BU_NAME",
        "ACCOUNT__C",
        "CUSTOMER_PULSE__C",
    ],
)
def test_phase_1_ssot_keeps_business_columns(name):
    """Round 15 / Phase 1: SSoT must not over-deny customer-facing columns."""
    assert not is_internal_column(name), f"{name!r} must survive the denylist"


def test_phase_1_ssot_underscore_prefix_blocks_unknown_internal_markers():
    """Defense-in-depth: any new ``_`` prefix gets dropped automatically."""
    assert is_internal_column("_brand_new_internal_marker")
    assert is_internal_column("_audit_only")


def test_phase_1_ssot_handles_non_string_inputs_safely():
    """``is_internal_column`` must never raise on non-string headers."""
    assert is_internal_column(None) is False
    assert is_internal_column(42) is False
    assert is_internal_column(float("nan")) is False
    assert is_internal_column("") is False


def test_phase_1_curated_ab_detail_all_has_no_denylist_overlap():
    """Curated allowlists must not include anything the denylist drops."""
    for sheet, cols in CURATED_COLUMNS.items():
        for c in cols:
            assert not is_internal_column(c), (
                f"{sheet!r} curated column {c!r} is on the denylist; "
                "fix the SSoT before shipping."
            )


def test_phase_1_curated_columns_are_unique_per_sheet():
    """Curated allowlists must be unique within each sheet."""
    for sheet, cols in CURATED_COLUMNS.items():
        assert len(cols) == len(set(cols)), f"{sheet!r} has duplicate curated columns"


# ---------------------------------------------------------------------------
# Filter / apply behaviour
# ---------------------------------------------------------------------------


def test_phase_1_filter_columns_drops_denied_and_preserves_order():
    cols = ["ID", "IS_DELETED", "_stale_storage", "SUBJECT_C", "SEVERITY_C"]
    out = filter_columns(cols)
    assert out == ["ID", "SUBJECT_C", "SEVERITY_C"]


def test_phase_1_filter_columns_projects_curated_for_known_sheets():
    cols = ["IS_DELETED", "SEVERITY_C", "ID", "ZZZ_UNKNOWN", "SUBJECT_C"]
    out = filter_columns(cols, sheet_name="AB_Detail_All")
    # Curated order from the SSoT, only including columns we passed in
    assert out == ["ID", "SUBJECT_C", "SEVERITY_C"]
    assert "ZZZ_UNKNOWN" not in out
    assert "IS_DELETED" not in out


def test_phase_1_apply_export_schema_renames_csone_headers():
    df = pd.DataFrame(
        [
            {
                "Customer Name: Customer Name": "Acme",
                "Product: Product Name": "ProdX",
                "SR Number": "TAC1",
                "Title": "boom",
                "col_2": "x",
                "Title_Lower": "boom",
            }
        ]
    )
    out = apply_export_schema(df, sheet_name="CSOne_Detail_All")
    assert "Customer" in out.columns
    assert "Product" in out.columns
    assert "Customer Name: Customer Name" not in out.columns
    assert "col_2" not in out.columns
    assert "Title_Lower" not in out.columns
    assert "SR Number" in out.columns


def test_phase_1_apply_export_schema_no_curated_passes_through_minus_denied():
    """Sheet without a curated entry: denylist still applies."""
    df = pd.DataFrame(
        [
            {
                "ID": 1,
                "IS_DELETED": False,
                "BUSINESS_COL": "keep",
                "_internal": "drop",
            }
        ]
    )
    out = apply_export_schema(df, sheet_name="UNRELATED_SHEET_NAME")
    assert list(out.columns) == ["ID", "BUSINESS_COL"]


def test_phase_1_apply_export_schema_handles_none_and_non_dataframe():
    assert apply_export_schema(None) is None
    converted = apply_export_schema([{"ID": 1, "IS_DELETED": True}], sheet_name=None)
    assert list(converted.columns) == ["ID"]


# ---------------------------------------------------------------------------
# Marker presence -- writers actually call the SSoT
# ---------------------------------------------------------------------------


def _read(rel: str) -> str:
    return (REPO_ROOT / rel).read_text(encoding="utf-8")


def test_phase_1_marker_adoptiq_backend_imports_ssot():
    src = _read("adoptiq_backend.py")
    assert "from report_export_schema import apply_export_schema as _r15_apply_export_schema" in src
    assert "Round 15 / Phase 1.1" in src
    assert src.count("_r15_apply_export_schema(df_copy, sheet_name=") >= 2


def test_phase_1_marker_app_simple_imports_ssot_and_curates():
    src = _read("app_simple.py")
    assert "from report_export_schema import apply_export_schema as _r15_apply_export_schema" in src
    assert "Round 15 / Phase 1.3" in src
    assert "Round 15 / Phase 1.4" in src
    assert "Round 15 / Phase 1.5" in src


def test_phase_1_marker_ssot_module_present_and_self_describing():
    src = _read("report_export_schema.py")
    assert "Round 15 / Phase 1" in src
    assert "INTERNAL_COLUMN_DENYLIST" in src
    assert "CURATED_COLUMNS" in src


# ---------------------------------------------------------------------------
# Regression -- gold fixture, post-curation
# ---------------------------------------------------------------------------


def test_phase_1_gold_ab_detail_all_drops_to_curated_set_after_filter():
    """Apply the schema to the gold fixture's headers and assert the bad
    columns are gone and only curated entries survive."""
    import openpyxl

    wb = openpyxl.load_workbook(REPO_ROOT / "tests/fixtures/round15/gold_data.xlsx", read_only=True)
    try:
        headers = [c.value for c in next(wb["AB_Detail_All"].iter_rows(min_row=1, max_row=1))]
    finally:
        wb.close()

    out = filter_columns(headers, sheet_name="AB_Detail_All")
    for leak in (
        "IS_DELETED",
        "MAY_EDIT",
        "IS_LOCKED",
        "RECORD_TYPE_ID",
        "SYSTEM_MODSTAMP",
        "LAST_VIEWED_DATE",
        "LAST_REFERENCED_DATE",
        "CONNECTION_RECEIVED_ID",
        "CONNECTION_SENT_ID",
        "_FIVETRAN_SYNCED",
        "_FIVETRAN_DELETED",
    ):
        assert leak not in out, f"{leak!r} survived the curation filter"
    assert len(out) <= len(CURATED_COLUMNS["AB_Detail_All"])
    # And we kept the most important business columns
    assert "ID" in out
    assert "SUBJECT_C" in out
    assert "SEVERITY_C" in out
    assert "AB_STATUS_C" in out


def test_phase_1_gold_csconsole_pulse_drops_etl_and_sf_audit():
    import openpyxl

    wb = openpyxl.load_workbook(REPO_ROOT / "tests/fixtures/round15/gold_data.xlsx", read_only=True)
    try:
        headers = [c.value for c in next(wb["CSConsole_Customer_Pulse"].iter_rows(min_row=1, max_row=1))]
    finally:
        wb.close()

    out = filter_columns(headers, sheet_name="CSConsole_Customer_Pulse")
    for leak in (
        "ETL_ID",
        "DELETE_FLAG",
        "RECURSIVE_FLAG",
        "ISDELETED",
        "MAYEDIT",
        "ISLOCKED",
        "CONNECTIONRECEIVEDID",
        "CONNECTIONSENTID",
        "CREATEDBYID",
        "LASTMODIFIEDBYID",
        "SYSTEMMODSTAMP",
        "EDWSF_BATCH_ID",
        "EDWSF_CREATE_DTM",
        "EDWSF_CREATE_USER",
        "EDWSF_UPDATE_DTM",
        "EDWSF_UPDATE_USER",
        "EDWSF_SOURCE_DELETED_FLAG",
    ):
        assert leak not in out, f"{leak!r} survived CSConsole curation"
    assert "BU_NAME" in out
    assert "CUSTOMER_PULSE__C" in out


def test_phase_1_gold_external_incidents_drops_internal_markers():
    import openpyxl

    wb = openpyxl.load_workbook(REPO_ROOT / "tests/fixtures/round15/gold_data.xlsx", read_only=True)
    try:
        headers = [c.value for c in next(wb["External_Incidents"].iter_rows(min_row=1, max_row=1))]
    finally:
        wb.close()

    out = filter_columns(headers, sheet_name="External_Incidents")
    for leak in ("_stale_storage", "_from_storage", "_window_meta", "publication_id"):
        assert leak not in out
    assert "title" in out
    assert "status" in out


def test_phase_1_gold_csone_renames_relationship_labels():
    """The CSOne header rename collapses SOQL relationship labels."""
    import openpyxl

    wb = openpyxl.load_workbook(REPO_ROOT / "tests/fixtures/round15/gold_data.xlsx", read_only=True)
    try:
        ws = wb["CSOne_Detail_All"]
        rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
        headers = list(rows[0])
        first_row = list(rows[1]) if len(rows) > 1 else [None] * len(headers)
    finally:
        wb.close()

    df = pd.DataFrame([dict(zip(headers, first_row))])
    out = apply_export_schema(df, sheet_name="CSOne_Detail_All")
    assert "Customer" in out.columns or "customer_name" in out.columns
    assert "Customer Name: Customer Name" not in out.columns
    assert "col_2" not in out.columns
    assert "col_32" not in out.columns
    assert "Title_Lower" not in out.columns
    assert "Desc_Lower" not in out.columns


# ---------------------------------------------------------------------------
# Coverage of public API surface
# ---------------------------------------------------------------------------


def test_phase_1_public_api_surface_is_minimal():
    """SSoT exports just what the writers need, no more.

    Round 45 / Phase 5 added ``friendly_header`` for callers that need
    to map a single raw Snowflake column name to its director-friendly
    label (e.g. cell-comment generators, chart legends).  The helper
    is documented in ``report_export_schema._FRIENDLY_HEADER_LABELS``
    and is the public counterpart to the cross-sheet rename that
    ``apply_export_schema`` applies as its final step.
    """
    import report_export_schema as schema

    expected = {
        "INTERNAL_COLUMN_DENYLIST",
        "INTERNAL_COLUMN_PREFIXES",
        "CURATED_COLUMNS",
        "SHEET_HEADER_RENAMES",
        "is_internal_column",
        "filter_columns",
        "apply_export_schema",
        # Round 45 / Phase 5: cross-sheet friendly-label helper.
        "friendly_header",
    }
    assert set(schema.__all__) == expected


def test_phase_1_internal_column_prefixes_includes_underscore_and_edwsf():
    """The two prefix denials we rely on are wired."""
    assert "_" in INTERNAL_COLUMN_PREFIXES
    assert "EDWSF_" in INTERNAL_COLUMN_PREFIXES


def test_phase_1_sheet_header_renames_csone_present():
    """CSOne rename map must cover the relationship-label leaks."""
    csone = SHEET_HEADER_RENAMES.get("CSOne_Detail_All", {})
    assert csone.get("Customer Name: Customer Name") == "Customer"
    assert csone.get("Product: Product Name") == "Product"
    assert csone.get("Case Owner: Full Name") == "Case Owner"


def test_phase_1_denylist_is_frozen_set():
    """``INTERNAL_COLUMN_DENYLIST`` is immutable to prevent monkey-patch drift."""
    assert isinstance(INTERNAL_COLUMN_DENYLIST, frozenset)
