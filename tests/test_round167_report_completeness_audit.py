"""Round 167 report completeness and source drill-through gates."""

from __future__ import annotations

from pathlib import Path
import re

import pandas as pd
import pytest

import app_simple
import decision_report_delivery as delivery
from report_completeness_audit import (
    audit_source_data_frames,
    audit_word_placeholders,
    audit_written_source_hyperlinks,
)
from tests.test_round142_decision_report_delivery import _facts, _fake_chart_renderer


def test_canonical_fixture_has_only_explained_unknowns_and_complete_links() -> None:
    audit = audit_source_data_frames(delivery.build_source_data_sheets(_facts()))

    assert audit["ok"], audit["errors"]
    assert audit["undefined_cells"] == []
    assert audit["unexplained_action_plan_status_rows"] == 0
    assert audit["tac_non_record_rows"] == 0
    assert audit["source_record_link_rows"] > 0
    assert audit["source_record_link_errors"] == 0


def test_audit_rejects_renderer_placeholder_unexplained_unknown_and_footer() -> None:
    sheets = delivery.build_source_data_sheets(_facts())
    sheets["Customer_Pulse"] = sheets["Customer_Pulse"].copy()
    sheets["Customer_Pulse"].loc[0, "PULSE_RATING__C"] = "undefined"
    sheets["Action_Plans"] = sheets["Action_Plans"].copy()
    unknown_index = sheets["Action_Plans"].index[sheets["Action_Plans"]["AdoptIQ_Status_Bucket"].eq("Unknown")][0]
    sheets["Action_Plans"].loc[unknown_index, "AdoptIQ_Data_Quality"] = "OK"
    footer = pd.DataFrame(
        [
            {
                "Record_ID": "",
                "Record_ID_Data_Quality": "Missing stable source ID",
                "Customer": "Export summary",
            }
        ]
    )
    sheets["TAC_Cases"] = pd.concat([sheets["TAC_Cases"], footer], ignore_index=True, sort=False)

    audit = audit_source_data_frames(sheets)

    assert audit["ok"] is False
    assert audit["undefined_cells"]
    assert audit["unexplained_action_plan_status_rows"] == 1
    assert audit["tac_non_record_rows"] == 1


def test_word_audit_rejects_bare_unknown_even_outside_identity_cells() -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("Status: Unknown")

    audit = audit_word_placeholders(document)

    assert audit["ok"] is False
    assert audit["generic_unknown_locations"] == ["paragraph:1"]


def test_written_workbook_and_word_prove_quality_contract(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    facts = _facts()
    sheets = delivery.build_source_data_sheets(facts)
    source_path = tmp_path / "source.xlsx"
    delivery.write_source_data_workbook(source_path, sheets)
    workbook_audit = audit_written_source_hyperlinks(source_path)

    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    document = delivery.build_concise_word_document(facts)
    word_audit = audit_word_placeholders(document)

    assert workbook_audit["ok"], workbook_audit["errors"]
    assert workbook_audit["url_cells"] == workbook_audit["clickable_cells"]
    assert workbook_audit["url_cells"] > 0
    assert word_audit["ok"], word_audit["errors"]


def test_manager_word_uses_explicit_reasons_instead_of_bare_unknown(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    facts = _facts()
    assert (
        delivery._action_plan_status_display(  # noqa: SLF001 - presentation contract
            {
                "AdoptIQ_Status_Bucket": "Unknown",
                "STATUS_C": "Awaiting business validation",
            }
        )
        == "Unmapped: Awaiting business validation"
    )
    assert (
        delivery._action_plan_status_display(  # noqa: SLF001
            {"AdoptIQ_Status_Bucket": "Unknown", "STATUS_C": ""}
        )
        == "Status not provided"
    )
    unresolved_labels = facts["chart_data"].loc[
        facts["chart_data"]["Metric_Key"].astype(str).str.endswith(".unknown"),
        "Display_Label",
    ]
    assert any("Status unresolved" in value for value in unresolved_labels)
    assert any("Created date unavailable" in value for value in unresolved_labels)

    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    document = delivery.build_concise_word_document(facts)
    visible_text = [paragraph.text for paragraph in document.paragraphs]
    visible_text.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert not re.search(r"\bunknown\b", "\n".join(visible_text), re.IGNORECASE)


def test_awaiting_business_validation_is_active_work_not_unknown() -> None:
    from canonical_metrics import build_action_plan_lifecycle

    lifecycle = build_action_plan_lifecycle(
        pd.DataFrame(
            [
                {
                    "ID": "AP-VALIDATE-1",
                    "SUBJECT_C": "Confirm business outcome",
                    "STATUS_C": "Awaiting business validation",
                    "CREATED_DATE_C": "2026-07-15T12:00:00Z",
                }
            ]
        ),
        as_of="2026-08-03T12:00:00Z",
    )

    assert lifecycle["bucket_counts"]["Open"] == 1
    assert lifecycle["bucket_counts"]["Unknown"] == 0
    assert lifecycle["records"].iloc[0]["AdoptIQ_Data_Quality"] == "OK"


def test_partial_warning_detail_uses_effect_and_never_serializes_null() -> None:
    warning = {
        "dataset": "data_freshness",
        "kind": "freshness_partial",
        "effect": "One source has an older observation timestamp.",
    }

    assert app_simple._r167_partial_warning_detail(warning) == (  # noqa: SLF001
        "One source has an older observation timestamp."
    )
    fallback = app_simple._r167_partial_warning_detail(  # noqa: SLF001
        {"dataset": "customer_pulse", "kind": "coverage_partial", "error": None}
    )
    assert fallback == "customer pulse coverage is incomplete (coverage partial)."
    assert fallback.casefold() not in {"none", "null", "nan", "unknown"}


def test_public_chart_categories_explain_unresolved_action_plan_values() -> None:
    sheets = delivery.build_source_data_sheets(_facts())
    chart_data = sheets["Chart_Data"]
    evidence_links = sheets["Evidence_Links"]

    unresolved_chart = chart_data.loc[
        chart_data["Metric_Key"].astype(str).str.endswith(".unknown")
    ]
    unresolved_evidence = evidence_links.loc[
        evidence_links["Evidence_Key"].astype(str).str.endswith(".unknown")
    ]

    assert set(unresolved_chart["Category"]) == {
        "Status unresolved",
        "Created date unavailable",
    }
    assert set(unresolved_evidence["Chart_Category"]) == {
        "Status unresolved",
        "Created date unavailable",
    }
    assert not chart_data.astype("string").fillna("").eq("Unknown").any().any()


def test_manager_decision_text_is_bounded_without_changing_source_rows() -> None:
    source_value = "Detailed source narrative " * 100

    rendered = delivery._bounded_manager_text(  # noqa: SLF001
        source_value,
        limit=140,
        fallback="Not provided",
    )

    assert len(rendered) <= 140
    assert rendered.endswith("[complete in Source Data]")
    assert len(source_value) > len(rendered)


def test_filtered_is_an_allowed_honest_source_state() -> None:
    sheets = delivery.build_source_data_sheets(_facts())
    report_info = sheets["Report_Info"].copy()
    source_row = report_info["Item"].astype(str).str.startswith("Source_State:")
    assert source_row.any()
    report_info.loc[source_row.idxmax(), "Value"] = "filtered"
    sheets["Report_Info"] = report_info

    audit = audit_source_data_frames(sheets)

    assert not any("invalid source state" in item for item in audit["errors"])


def test_completeness_blocks_unassigned_account_rows_and_roster_drift() -> None:
    sheets = delivery.build_source_data_sheets(_facts())
    action_plans = sheets["Action_Plans"].copy()
    subscriptions = sheets["Subscriptions"]
    subscription_account_column = next(
        column
        for column in subscriptions.columns
        if str(column).replace("_", " ").casefold() in {"account id c", "account id"}
    )
    attributed_account_ids = set(
        subscriptions.loc[
            subscriptions["Attributed_Team_Members"].fillna("").astype(str).str.strip().ne(""),
            subscription_account_column,
        ]
        .dropna()
        .astype(str)
    )
    target_index = action_plans.index[0]
    action_plans["ACCOUNT_ID_C"] = ""
    action_plans.loc[target_index, "ACCOUNT_ID_C"] = sorted(attributed_account_ids)[0]
    action_plans.loc[target_index, "Attributed_Team_Members"] = "Unassigned / Portfolio"
    sheets["Action_Plans"] = action_plans
    member_summary = sheets["Member_Summary"].copy()
    extra = {column: "" for column in member_summary.columns}
    extra["Team_Member"] = "Out Of Scope Member"
    sheets["Member_Summary"] = pd.concat(
        [member_summary, pd.DataFrame([extra])],
        ignore_index=True,
    )

    audit = audit_source_data_frames(sheets)

    assert audit["ok"] is False
    assert audit["attribution_mismatch_rows"] >= 1
    assert audit["member_summary_mismatch"] is True


def test_completeness_blocks_tac_rows_outside_window_or_without_opened_date() -> None:
    sheets = delivery.build_source_data_sheets(_facts())
    tac = sheets["TAC_Cases"].copy()
    opened_column = next(
        column
        for column in tac.columns
        if str(column).replace("_", " ").casefold()
        in {"date/time opened", "created date", "date opened", "open date c"}
    )
    report_info = dict(
        zip(
            sheets["Report_Info"]["Item"].astype(str),
            sheets["Report_Info"]["Value"],
        )
    )
    as_of = pd.to_datetime(report_info["Evaluation_As_Of_UTC"], utc=True)
    days = int(report_info["Days"])
    outside = tac.iloc[[0]].copy()
    outside["Record_ID"] = "TAC-OUTSIDE-WINDOW"
    outside[opened_column] = as_of - pd.Timedelta(days + 1, unit="D")
    missing = tac.iloc[[0]].copy()
    missing["Record_ID"] = "TAC-MISSING-DATE"
    missing[opened_column] = pd.NaT
    sheets["TAC_Cases"] = pd.concat(
        [tac, outside, missing],
        ignore_index=True,
        sort=False,
    )

    audit = audit_source_data_frames(sheets)

    assert audit["ok"] is False
    assert audit["tac_outside_window_rows"] == 1
    assert audit["tac_missing_open_date_rows"] == 1


def test_completeness_blocks_out_of_scope_customer_commercial_fact() -> None:
    sheets = delivery.build_source_data_sheets(_facts())
    report_info = sheets["Report_Info"].copy()
    report_info.loc[report_info["Item"].eq("Scope_Type"), "Value"] = "customer"
    report_info.loc[report_info["Item"].eq("Scope_Value"), "Value"] = "Acme Corporation"
    sheets["Report_Info"] = report_info
    subscriptions = sheets["Subscriptions"].copy()
    extra = {column: "" for column in subscriptions.columns}
    extra.update(
        {
            "Record_ID": "legacy:commercial:outside",
            "Legacy_Record_Type": "Family-specific reported fact",
            "Account ID": "ACC-OUTSIDE",
            "Customer Name": "Different Customer",
            "Legacy_Source_Sheet": "Renewal_Commercial_Facts",
            "Legacy_Source_Field": "ARR Amount",
            "Legacy_Fact_Value": 999999,
        }
    )
    sheets["Subscriptions"] = pd.concat(
        [subscriptions, pd.DataFrame([extra])],
        ignore_index=True,
    )

    audit = audit_source_data_frames(sheets)

    assert audit["ok"] is False
    assert audit["customer_scope_mismatch_rows"] >= 1
    assert audit["customer_scope_mismatch_by_sheet"]["Subscriptions"] >= 1
