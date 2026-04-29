"""Round 48 / F-RP-PARTIAL-BANNER-MISSING regression tests.

Pin that the renewal Word renderer surfaces the same Partial Data
Warning banner the compact / EI Word renderers already emit (R46
F-COMP-DQ-BANNER pattern at app_simple.py L5510 and
executive_intelligence_formatter.py L1610).  Pre-Round 48 the
renewal Word silently consumed a "successful" report built on
partial data while the Excel ``Report_Info`` sheet correctly
showed the warnings -- the two artifacts disagreed.
"""

from __future__ import annotations

import inspect
from pathlib import Path

import pandas as pd
import pytest

import app_simple


_APP_SIMPLE_PATH = Path(app_simple.__file__).resolve()


@pytest.fixture(scope="module")
def app_simple_source() -> str:
    return _APP_SIMPLE_PATH.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Function signature -- the kwarg must exist
# ---------------------------------------------------------------------------


def test_round48_create_simple_renewal_report_accepts_partial_data_warnings_kwarg():
    """``_create_simple_renewal_report`` must accept the kwarg
    so the harvest result from ``run_customer_renewal_analysis``
    can be threaded in.
    """

    sig = inspect.signature(app_simple._create_simple_renewal_report)
    assert "partial_data_warnings" in sig.parameters, (
        "_create_simple_renewal_report does not accept "
        "partial_data_warnings kwarg; R48 wiring is unplugged"
    )
    param = sig.parameters["partial_data_warnings"]
    assert param.default is None, (
        f"partial_data_warnings kwarg default should be None; got {param.default!r}"
    )


def test_round48_renewal_word_fix_anchor_present(app_simple_source: str):
    assert "F-RP-PARTIAL-BANNER-MISSING" in app_simple_source, (
        "Round 48 fix anchor F-RP-PARTIAL-BANNER-MISSING missing "
        "from app_simple.py"
    )


def test_round48_renewal_caller_forwards_kwarg(app_simple_source: str):
    """The single ``_create_simple_renewal_report(...)`` call site
    in ``run_customer_renewal_analysis`` must pass
    ``partial_data_warnings=_r48_word_pdw or None`` so the renderer
    actually receives the harvested list.
    """

    assert "partial_data_warnings=_r48_word_pdw or None" in app_simple_source, (
        "run_customer_renewal_analysis does not forward the harvested "
        "_r48_word_pdw to _create_simple_renewal_report; banner will "
        "not render"
    )


# ---------------------------------------------------------------------------
# Banner rendered into the Word document body
# ---------------------------------------------------------------------------


def _build_sample_pdw():
    return [
        {
            "dataset": "customer_pulse",
            "error": "schema_drift:customer_pulse: missing slot(s) rating",
            "kind": "schema_drift",
        },
        {
            "dataset": "adoption_barriers",
            "error": "schema_drift:adoption_barriers: missing slot(s) customer",
            "kind": "schema_drift",
        },
    ]


def test_round48_renewal_word_renders_banner_when_warnings_present(tmp_path):
    """End-to-end: pass two schema_drift warnings into the renewal
    Word renderer and assert the rendered .docx body contains the
    Partial Data Warning heading + each warning line.
    """

    base = tmp_path / "renewal_banner_test"
    word_path = app_simple._create_simple_renewal_report(
        base_path=str(base),
        customer_name="Test Portfolio",
        technology="Contact Center",
        days=90,
        renewal_analysis={
            "renewal_risk_score": 12.0,
            "renewal_risk_category": "LOW",
            "adoption_barriers_count": 0,
            "support_cases_count": 0,
            "bems_escalations_count": 0,
        },
        customer_ab=pd.DataFrame(),
        customer_csone=pd.DataFrame(),
        portfolio_mode=True,
        all_customers=["Cust A"],
        partial_data_warnings=_build_sample_pdw(),
    )

    from docx import Document

    doc = Document(word_path)
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Partial Data Warning" in body_text, (
        "Renewal Word body missing 'Partial Data Warning' heading"
    )
    assert "customer_pulse" in body_text, (
        "Renewal Word body missing customer_pulse warning entry"
    )
    assert "adoption_barriers" in body_text, (
        "Renewal Word body missing adoption_barriers warning entry"
    )
    assert "schema_drift" in body_text, (
        "Renewal Word body missing kind=schema_drift in warning entry"
    )


def test_round48_renewal_word_no_banner_when_no_warnings(tmp_path):
    """Inverse safety: when no warnings are passed, the Word body
    must NOT contain the banner heading.  A false-positive banner
    would itself be a defect (calls into question artifacts that
    are clean).
    """

    base = tmp_path / "renewal_no_banner_test"
    word_path = app_simple._create_simple_renewal_report(
        base_path=str(base),
        customer_name="Test Portfolio",
        technology="Contact Center",
        days=90,
        renewal_analysis={
            "renewal_risk_score": 12.0,
            "renewal_risk_category": "LOW",
            "adoption_barriers_count": 0,
            "support_cases_count": 0,
            "bems_escalations_count": 0,
        },
        customer_ab=pd.DataFrame(),
        customer_csone=pd.DataFrame(),
        portfolio_mode=True,
        all_customers=["Cust A"],
        partial_data_warnings=None,
    )

    from docx import Document

    doc = Document(word_path)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert all("Partial Data Warning" not in h for h in headings), (
        "Renewal Word body shows Partial Data Warning heading even "
        "though no warnings were passed -- false-positive banner"
    )
