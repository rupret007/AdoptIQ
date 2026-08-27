"""Round 146 Flask contracts for the Manager Decision Workspace."""

from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
import types
from importlib.machinery import ModuleSpec
from pathlib import Path

from openpyxl import Workbook


# Route tests exercise guarded adapters only; never require a Snowflake wheel
# or connection on the local development machine.
if importlib.util.find_spec("snowflake") is None:
    snowflake_package = types.ModuleType("snowflake")
    snowflake_package.__path__ = []  # type: ignore[attr-defined]
    snowflake_package.__spec__ = ModuleSpec(
        "snowflake", loader=None, is_package=True
    )
    snowflake_connector = types.ModuleType("snowflake.connector")
    snowflake_connector.__spec__ = ModuleSpec(
        "snowflake.connector", loader=None
    )
    snowflake_connector.DictCursor = object

    def _no_snowflake_connect(*_args, **_kwargs):
        raise RuntimeError("Snowflake is unavailable in Round 146 route tests")

    snowflake_connector.connect = _no_snowflake_connect
    snowflake_package.connector = snowflake_connector
    sys.modules["snowflake"] = snowflake_package
    sys.modules["snowflake.connector"] = snowflake_connector


def _source_workbook(path: Path, *, fingerprint: str, open_plans: int) -> Path:
    workbook = Workbook()
    info = workbook.active
    info.title = "Report_Info"
    info.append(["Item", "Value"])
    for row in (
        ("Report_Type", "leader"),
        ("Manager", "Manager One"),
        ("Technology", "All"),
        ("Scope_Type", "team"),
        ("Days", 90),
        ("Data_As_Of_UTC", "2026-08-03T12:00:00Z"),
        ("Fact_Contract_SHA256", fingerprint),
        ("Source_State:Action_Plans", "available"),
    ):
        info.append(row)
    lineage = workbook.create_sheet("Metric_Lineage")
    lineage.append([
        "Metric_Key", "Display_Label", "Metric_Value", "Unit",
        "Source_State", "Source_Sheet",
    ])
    lineage.append([
        "kpi.action_plans_open", "Open Action Plans", open_plans,
        "records", "available", "Action_Plans",
    ])
    chart = workbook.create_sheet("Chart_Data")
    chart.append([
        "Chart_ID", "Metric_Key", "Display_Label", "Series", "Category",
        "Period_Start", "Value", "Unit", "Source_State",
    ])
    chart.append([
        "action_plan_status_aging", "chart.action_plan_status.open",
        "Action Plan status and aging — Open", "Status", "Open", "",
        open_plans, "records", "available",
    ])
    actions = workbook.create_sheet("Action_Plans")
    actions.append([
        "ID", "BU_NAME", "SUBJECT_C", "STATUS_C", "OWNER_NAME_C",
        "DUE_DATE_C", "PRIORITY_C", "NEXT_ACTION_C",
    ])
    actions.append([
        "AP-1", "Acme", "Adoption plan", "On Track", "Owner",
        "2026-08-20", "High", "Confirm enablement",
    ])
    accounts = workbook.create_sheet("Account_Summary")
    accounts.append(["Account", "Risk_Band", "Risk_Score_0_100"])
    accounts.append(["Acme", "HIGH", 75])
    workbook.save(path)
    workbook.close()
    return path


def test_scope_preview_discloses_guarded_fixture_and_roster(client, monkeypatch):
    import app_simple

    manager = app_simple.TEAM_ROSTER[0][0]
    monkeypatch.setenv("ADOPTIQ_LOCAL_ACCEPTANCE_ACTIVE", "1")
    response = client.get(
        "/api/decision-workspace/scope-preview",
        query_string={
            "report_type": "leader",
            "manager": manager,
            "technology": "All",
            "days": 90,
            "scope_type": "team",
        },
    )

    assert response.status_code == 200
    preview = response.get_json()["preview"]
    assert preview["scope_type"] == "team"
    assert preview["member_count"] >= 1
    assert preview["source_mode"] == "local_fixture_guarded"
    assert preview["live_validation_performed"] is False
    assert any("no live validation" in item.lower() for item in preview["limitations"])


def test_scope_preview_rejects_customer_not_in_authorized_options(client, monkeypatch):
    import app_simple

    manager = app_simple.TEAM_ROSTER[0][0]
    monkeypatch.setattr(
        app_simple,
        "_r146_leader_scope_options_payload",
        lambda *_args, **_kwargs: ({
            "success": True,
            "members": [],
            "customers": [{"value": "Authorized Customer"}],
            "customers_available": True,
            "warning": "",
        }, 200),
    )
    response = client.get(
        "/api/decision-workspace/scope-preview",
        query_string={
            "report_type": "leader",
            "manager": manager,
            "technology": "All",
            "days": 90,
            "scope_type": "customer",
            "scope_value": "Outside Customer",
        },
    )

    assert response.status_code == 400
    assert "authorized" in response.get_json()["error"].lower()


def test_report_view_is_path_free_and_uses_existing_download_contract(
    client, monkeypatch, tmp_path
):
    import app_simple
    from offline_validation_receipt import OfflineValidationReceiptStatus

    workbook_path = _source_workbook(
        tmp_path / "AdoptIQ_Source_Data_Test.xlsx",
        fingerprint="report-fingerprint",
        open_plans=1,
    )
    word_path = tmp_path / "AdoptIQ_Report_Test.docx"
    word_path.write_bytes(b"test")
    status = {
        "analysis_id": "round146-report",
        "status": "completed",
        "report_type": "leader",
        "manager": "Manager One",
        "technology": "All",
        "days": 90,
        "word_report": str(word_path),
        "excel_report": str(workbook_path),
        "excel_hash": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
    }
    monkeypatch.setattr(app_simple, "_r146_status_for_workspace", lambda _aid: status)
    monkeypatch.setattr(
        app_simple,
        "_r92_resolve_output_artifact",
        lambda raw, **_kwargs: str(raw) if raw else None,
    )
    receipt_call = {}

    def _receipt_loader(analysis_id, **kwargs):
        receipt_call["analysis_id"] = analysis_id
        receipt_call["projection"] = kwargs["expected_web_projection_sha256"]
        return OfflineValidationReceiptStatus(state="invalid")

    monkeypatch.setattr(app_simple, "load_offline_validation_receipt", _receipt_loader)

    response = client.get("/api/decision-workspace/report/round146-report")

    assert response.status_code == 200
    report = response.get_json()["report"]
    assert report["workbook_loaded"] is True
    assert report["fact_fingerprint"] == "report-fingerprint"
    assert report["downloads"]["word"].endswith("/docx")
    assert report["downloads"]["source_data"].endswith("/xlsx")
    assert report["ask_ai_binding"]["analysis_id"] == "round146-report"
    assert receipt_call["analysis_id"] == "round146-report"
    assert len(receipt_call["projection"]) == 64
    int(receipt_call["projection"], 16)
    assert report["offline_validation_receipt"]["state"] == "invalid"
    assert report["offline_validation_receipt"]["fixture_validation_performed"] is False
    assert report["customer_share_readiness"]["live_validation_performed"] is False
    assert report["customer_share_readiness"]["production_accuracy_claimed"] is False
    assert report["customer_share_readiness"]["release_ready"] is False
    assert report["customer_share_readiness"]["customer_shareable"] is False
    assert "signature" not in json.dumps(report)
    rendered = response.get_data(as_text=True)
    assert str(tmp_path) not in rendered


def test_hash_mismatch_blocks_workspace_and_report_bound_ask_ai(
    client, monkeypatch, tmp_path
):
    import app_simple

    analysis_id = "round146-integrity-mismatch"
    workbook_path = _source_workbook(
        tmp_path / "private-source-data.xlsx",
        fingerprint="integrity-fingerprint",
        open_plans=1,
    )
    recorded_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    with workbook_path.open("ab") as workbook_handle:
        workbook_handle.write(b"post-history-tamper")
    status = {
        "analysis_id": analysis_id,
        "status": "completed",
        "report_type": "leader",
        "manager": "Manager One",
        "technology": "All",
        "days": 90,
        "excel_report": str(workbook_path),
    }
    persisted = {**status, "excel_hash": recorded_hash, "_rehydrated_from_audit": True}
    monkeypatch.setattr(
        app_simple, "_r146_status_for_workspace", lambda _analysis_id: status
    )
    monkeypatch.setattr(
        app_simple, "_build_status_from_report_history", lambda _analysis_id: persisted
    )
    monkeypatch.setattr(
        app_simple,
        "_r92_resolve_output_artifact",
        lambda raw, **_kwargs: str(raw) if raw else None,
    )
    monkeypatch.setattr(app_simple, "_check_ask_ai_throttle", lambda: None)

    workspace_response = client.get(
        f"/api/decision-workspace/report/{analysis_id}"
    )
    ask_ai_response = client.post(
        "/api/ask-ai-portfolio",
        json={
            "question": "What changed in this report?",
            "report_analysis_id": analysis_id,
        },
    )
    ask_ai_stream_response = client.post(
        "/api/ask-ai-portfolio/stream",
        json={
            "question": "What changed in this report?",
            "report_analysis_id": analysis_id,
        },
    )

    for response in (
        workspace_response,
        ask_ai_response,
        ask_ai_stream_response,
    ):
        assert response.status_code == 409
        payload = response.get_json()
        assert payload["ok"] is False
        assert "integrity" in payload["error"].lower()
        rendered = response.get_data(as_text=True)
        assert str(tmp_path) not in rendered
        assert recorded_hash not in rendered


def test_recorded_hash_fails_closed_when_workbook_cannot_be_read(
    client, monkeypatch, tmp_path
):
    import app_simple

    analysis_id = "round146-unreadable-verified-workbook"
    workbook_path = tmp_path / "private-unreadable-source.xlsx"
    workbook_path.write_bytes(b"not-an-xlsx-archive")
    status = {
        "analysis_id": analysis_id,
        "status": "completed",
        "report_type": "leader",
        "manager": "Manager One",
        "excel_report": str(workbook_path),
        "excel_hash": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
    }
    monkeypatch.setattr(
        app_simple, "_r146_status_for_workspace", lambda _analysis_id: status
    )
    monkeypatch.setattr(
        app_simple,
        "_r92_resolve_output_artifact",
        lambda raw, **_kwargs: str(raw) if raw else None,
    )

    response = client.get(f"/api/decision-workspace/report/{analysis_id}")

    assert response.status_code == 409
    assert "integrity" in response.get_json()["error"].lower()
    assert str(tmp_path) not in response.get_data(as_text=True)


def test_history_api_filters_and_never_exposes_artifact_paths(client, monkeypatch):
    import app_simple

    monkeypatch.setattr(
        app_simple,
        "get_report_history",
        lambda: [{
            "request_id": "history-one",
            "report_type": "leader",
            "manager": "Manager One",
            "technology": "All",
            "customer_name": "Acme",
            "status": "completed",
            "created_at": "2026-08-03T12:00:00Z",
            "word_path": "/private/customer/report.docx",
            "excel_path": "/private/customer/source.xlsx",
        }],
    )
    monkeypatch.setattr(
        app_simple,
        "_r146_status_for_workspace",
        lambda _aid: {
            "word_path": "/private/customer/report.docx",
            "excel_path": "/private/customer/source.xlsx",
        },
    )
    monkeypatch.setattr(
        app_simple,
        "_r92_resolve_output_artifact",
        lambda raw, **_kwargs: str(raw) if raw else None,
    )

    response = client.get(
        "/api/decision-workspace/history",
        query_string={"manager": "Manager One", "report_type": "leader"},
    )

    assert response.status_code == 200
    report = response.get_json()["reports"][0]
    assert report["word_url"].endswith("/docx")
    assert report["source_data_url"].endswith("/xlsx")
    assert "word_path" not in report
    assert "/private/customer" not in response.get_data(as_text=True)


def test_compare_route_uses_canonical_server_workbooks(client, monkeypatch, tmp_path):
    import app_simple

    before = _source_workbook(
        tmp_path / "before.xlsx", fingerprint="before", open_plans=3
    )
    after = _source_workbook(
        tmp_path / "after.xlsx", fingerprint="after", open_plans=1
    )
    statuses = {
        "before-run": {
            "analysis_id": "before-run", "status": "completed",
            "report_type": "leader", "manager": "Manager One",
            "technology": "All", "days": 90, "excel_report": str(before),
        },
        "after-run": {
            "analysis_id": "after-run", "status": "completed",
            "report_type": "leader", "manager": "Manager One",
            "technology": "All", "days": 90, "excel_report": str(after),
        },
    }
    monkeypatch.setattr(
        app_simple, "_r146_status_for_workspace", lambda aid: statuses.get(aid)
    )
    monkeypatch.setattr(
        app_simple,
        "_r92_resolve_output_artifact",
        lambda raw, **_kwargs: str(raw) if raw else None,
    )

    response = client.post(
        "/api/decision-workspace/compare",
        json={"before_analysis_id": "before-run", "after_analysis_id": "after-run"},
    )

    assert response.status_code == 200
    comparison = response.get_json()["comparison"]
    assert comparison["before_fingerprint"] == "before"
    assert comparison["after_fingerprint"] == "after"
    assert comparison["metric_changes"][0]["delta"] == -2


def test_compare_route_rejects_unlike_scopes_without_returning_deltas(
    client, monkeypatch
):
    import app_simple

    base = {
        "workbook_loaded": True,
        "canonical_snapshot": True,
        "report_type": "leader",
        "technology": "All",
        "scope_type": "team",
        "scope_value": "",
        "days": 90,
        "source_states": {"Action_Plans": "available"},
        "metrics": [{
            "metric_key": "kpi.action_plans_open",
            "label": "Open Action Plans",
            "value": 1,
            "source_state": "available",
            "source_sheet": "Action_Plans",
        }],
        "action_plans": [],
    }
    snapshots = {
        "before-run": dict(base, manager="Manager One", fact_fingerprint="before"),
        "after-run": dict(
            base,
            manager="Manager Two",
            fact_fingerprint="after",
            metrics=[dict(base["metrics"][0], value=999)],
        ),
    }
    monkeypatch.setattr(
        app_simple,
        "_r146_workspace_snapshot",
        lambda analysis_id: (snapshots[analysis_id], 200),
    )

    response = client.post(
        "/api/decision-workspace/compare",
        json={"before_analysis_id": "before-run", "after_analysis_id": "after-run"},
    )

    assert response.status_code == 409
    payload = response.get_json()
    assert payload["ok"] is False
    assert payload["comparison"]["same_scope"] is False
    assert payload["comparison"]["metric_changes"] == []
    assert payload["comparison"]["action_plan_changes"] == []
    assert payload["comparison"]["business_change_count"] == 0


def test_download_rehydrates_evicted_completed_report_from_history(
    client, monkeypatch, tmp_path
):
    import app_simple

    report_path = tmp_path / "archived.docx"
    report_path.write_bytes(b"archived-report")
    analysis_id = "evicted-round146"
    with app_simple.analysis_status_lock:
        app_simple.analysis_status.pop(analysis_id, None)
    monkeypatch.setattr(app_simple, "STATUS_FILE", str(tmp_path / "missing-status.json"))
    monkeypatch.setattr(
        app_simple,
        "_build_status_from_report_history",
        lambda _aid: {
            "analysis_id": analysis_id,
            "status": "completed",
            "word_path": str(report_path),
        },
    )
    monkeypatch.setattr(
        app_simple,
        "_r92_resolve_output_artifact",
        lambda raw, **_kwargs: str(raw) if raw else None,
    )

    try:
        response = client.get(f"/download/{analysis_id}/docx")
        assert response.status_code == 200
        assert response.data == b"archived-report"
    finally:
        with app_simple.analysis_status_lock:
            app_simple.analysis_status.pop(analysis_id, None)
