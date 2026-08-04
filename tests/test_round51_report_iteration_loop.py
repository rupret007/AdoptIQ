"""Round 51 regression tests for live report iteration harness."""

from __future__ import annotations

import zipfile
from argparse import Namespace
from pathlib import Path

import openpyxl

from report_iteration_loop import (
    LiveReportRunner,
    build_debug_filename,
    build_runner_config,
    build_scenario_map,
    compare_docx_against_baseline,
    compare_kpi_parity,
    compare_xlsx_against_baseline,
    extract_csrf_token,
    extract_docx_kpis,
    extract_xlsx_kpis,
    evaluate_app_health,
    evaluate_report_quality,
    select_latest_baseline,
    validate_docx_structure,
    validate_xlsx_structure,
)


def _write_docx_like(path: Path, text: str) -> None:
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
    <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:body>
        <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
      </w:body>
    </w:document>
    """
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", xml)


def _write_xlsx(path: Path, sheets: dict[str, list[str]], data_rows: int = 1) -> None:
    workbook = openpyxl.Workbook()
    first = True
    for name, headers in sheets.items():
        if first:
            sheet = workbook.active
            sheet.title = name
            first = False
        else:
            sheet = workbook.create_sheet(title=name)
        for idx, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=idx, value=header)
        for row_idx in range(2, data_rows + 2):
            for col_idx, _header in enumerate(headers, start=1):
                sheet.cell(row=row_idx, column=col_idx, value=f"r{row_idx}v{col_idx}")
    workbook.save(path)
    workbook.close()


def _args(**overrides):
    defaults = {
        "base_url": "http://127.0.0.1:5151",
        "downloads_dir": "~/Downloads",
        "iterations": 1,
        "scenarios": "all",
        "poll_interval": 5.0,
        "timeout": 1800,
        "run_id": "test",
        "stop_on_failure": False,
        "baseline_mode": "latest",
        "strict": False,
        "min_docx_similarity": 0.35,
        "min_sheet_overlap": 0.5,
        "min_header_similarity": 0.3,
        "min_docx_chars": 200,
        "min_docx_numeric_similarity": 0.8,
        # Round 52.1: noise-immune table-only numeric drift gate; defaults
        # to 0.95 in the live CLI parser.
        "min_docx_table_numeric_similarity": 0.95,
        "max_xlsx_row_delta_ratio": 0.2,
        "max_xlsx_row_delta_abs": 25,
    }
    defaults.update(overrides)
    return Namespace(**defaults)


class _FakeResponse:
    def __init__(self, status_code=200, text="", json_payload=None):
        self.status_code = status_code
        self.text = text
        self._json_payload = json_payload

    def json(self):
        if isinstance(self._json_payload, Exception):
            raise self._json_payload
        return self._json_payload

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class _FakeSession:
    def __init__(self, responses):
        self.responses = responses
        self.calls = []
        self.cookies = {}

    def get(self, url, timeout=30):
        self.calls.append((url, timeout))
        for path, response in self.responses.items():
            if url.endswith(path):
                return response
        raise AssertionError(f"unexpected URL {url}")


def test_round51_scenario_map_matches_requested_matrix():
    scenarios = build_scenario_map()
    assert set(scenarios) == {"comprehensive", "compact", "renewal", "leader"}
    assert scenarios["comprehensive"].payload["manager"] == "Brian Frazier"
    assert scenarios["comprehensive"].payload["technology"] == "All Contact Center"
    assert scenarios["compact"].payload["manager"] == "All Managers"
    assert scenarios["renewal"].payload["report_type"] == "renewal_portfolio"
    assert scenarios["leader"].endpoint == "/start_leader_report"


def test_round133_exhaustive_matrix_builder_importable_from_iteration_loop():
    from report_iteration_loop import build_exhaustive_option_matrix, parse_matrix_blocks

    matrix = build_exhaustive_option_matrix(days=90)
    assert len(matrix) >= 40
    keys = [k for k in matrix if k.startswith("e_")]
    assert keys
    assert parse_matrix_blocks("E") == ["E"]


def test_round51_extract_csrf_and_debug_filename():
    html = '<meta name="csrf-token" content="abc123" />'
    assert extract_csrf_token(html) == "abc123"

    debug_name = build_debug_filename(
        "AdoptIQ_Report_Leader_Brian_Frazier_90d_123.docx",
        run_id="run-1",
        scenario_key="leader",
        timestamp="20260429T170000Z",
    )
    assert debug_name.startswith("AdoptIQ_Report_Leader_Brian_Frazier_90d_123__data-loop-run-1__scenario-leader")
    assert debug_name.endswith(".docx")


def test_debug_filename_bounds_long_renewal_artifact_component() -> None:
    original = (
        "AdoptIQ_Source_Data_Renewal_Portfolio_Local_Fixture_Manager_"
        "Webex_Contact_Center_Enterprise_90d_1785806837632433000_fb0491ba.xlsx"
    )
    debug_name = build_debug_filename(
        original,
        run_id="data-loop-r145-final-current-v2-pass1",
        scenario_key="f_renewal_local_Webex_Contact_Center_Enterprise",
        timestamp="20260804T012719Z",
    )

    assert len(debug_name.encode("utf-8")) <= 240
    assert "__scenario-f_renewal_local_Webex_Contact_Center_Enterprise" in debug_name
    assert "__h-" in debug_name
    assert debug_name.endswith(".xlsx")


def test_debug_filename_digest_keeps_long_names_distinct() -> None:
    common = "AdoptIQ_Report_" + ("Very_Long_Customer_Name_" * 20)
    first = build_debug_filename(common + "A.docx", "run", "scenario", "20260804T012719Z")
    second = build_debug_filename(common + "B.docx", "run", "scenario", "20260804T012719Z")

    assert first != second
    assert len(first.encode("utf-8")) <= 240
    assert len(second.encode("utf-8")) <= 240


def test_round51_baseline_selection_uses_latest_matching_file(tmp_path: Path):
    newest = tmp_path / "AdoptIQ_Report_Compact_All_Managers_All_Contact_Center_90d_222.docx"
    older = tmp_path / "AdoptIQ_Report_Compact_All_Managers_All_Contact_Center_90d_111.docx"
    ignored = tmp_path / "AdoptIQ_Report_Compact_All_Managers_All_Contact_Center_90d_333__data-loop-current.docx"
    other = tmp_path / "AdoptIQ_Report_Leader_Brian_Frazier_90d_444.docx"

    for path in (older, newest, ignored, other):
        path.write_text("x", encoding="utf-8")

    older.touch()
    newest.touch()

    selected = select_latest_baseline(
        downloads_dir=tmp_path,
        scenario_key="compact",
        extension="docx",
        run_id="current",
        current_debug_name="current.docx",
    )
    assert selected == newest


def test_round51_docx_structural_and_baseline_diff(tmp_path: Path):
    current = tmp_path / "current.docx"
    baseline = tmp_path / "baseline.docx"
    _write_docx_like(current, "AdoptIQ summary customer risk and support health overview for contact center adoption.")
    _write_docx_like(baseline, "AdoptIQ summary customer risk and support health overview for contact center adoption.")

    structural = validate_docx_structure(current, min_chars=30)
    assert structural.passed is True

    diff = compare_docx_against_baseline(current, baseline, min_similarity=0.4)
    assert diff.passed is True
    assert diff.details["similarity"] >= 0.4


def test_round51_xlsx_structural_and_baseline_diff(tmp_path: Path):
    current = tmp_path / "current.xlsx"
    baseline = tmp_path / "baseline.xlsx"

    _write_xlsx(
        current,
        {
            "Summary": ["Metric", "Value"],
            "Risk_Overview": ["Customer", "Risk Score"],
        },
    )
    _write_xlsx(
        baseline,
        {
            "Summary": ["Metric", "Value"],
            "Risk_Overview": ["Customer", "Risk Score"],
        },
    )

    structural = validate_xlsx_structure(current)
    assert structural.passed is True

    diff = compare_xlsx_against_baseline(
        current,
        baseline,
        min_sheet_overlap=0.5,
        min_header_similarity=0.3,
    )
    assert diff.passed is True
    assert diff.details["sheet_overlap"] >= 0.5


def test_round51_strict_docx_numeric_drift_fails_even_when_words_match(tmp_path: Path):
    current = tmp_path / "current.docx"
    baseline = tmp_path / "baseline.docx"
    _write_docx_like(current, "Total Customers 12 Support Cases 99 Renewal Risk Score 17.5")
    _write_docx_like(baseline, "Total Customers 52 Support Cases 176 Renewal Risk Score 14.8")

    diff = compare_docx_against_baseline(
        current,
        baseline,
        min_similarity=0.5,
        strict=True,
        min_numeric_similarity=0.8,
    )

    assert diff.passed is False
    assert diff.details["similarity"] >= 0.5
    assert diff.details["numeric_similarity"] < 0.8
    assert "12" in diff.details["numeric_only_in_current"]


def test_round51_strict_xlsx_row_count_drift_is_reported(tmp_path: Path):
    current = tmp_path / "current.xlsx"
    baseline = tmp_path / "baseline.xlsx"
    sheets = {"Summary": ["Metric", "Value"], "Report_Info": ["Field", "Value"]}
    _write_xlsx(current, sheets, data_rows=80)
    _write_xlsx(baseline, sheets, data_rows=2)

    diff = compare_xlsx_against_baseline(
        current,
        baseline,
        min_sheet_overlap=0.8,
        min_header_similarity=0.8,
        strict=True,
        max_row_delta_ratio=0.2,
        max_row_delta_abs=5,
    )

    assert diff.passed is False
    assert diff.details["row_deltas"]["summary"]["delta"] > 5


def test_round97_3_docx_table_numeric_gate_tolerates_tiny_decimal_wobble(tmp_path: Path):
    from docx import Document

    baseline = tmp_path / "baseline.docx"
    current = tmp_path / "current.docx"
    for path, value in ((baseline, "175.8"), (current, "175.9")):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Category"
        table.rows[0].cells[1].text = "Priority Score"
        table.rows[1].cells[0].text = "Uncategorized"
        table.rows[1].cells[1].text = value
        doc.save(path)

    diff = compare_docx_against_baseline(
        current,
        baseline,
        min_similarity=0.0,
        strict=True,
        min_numeric_similarity=0.0,
        min_table_numeric_similarity=0.95,
    )

    assert diff.passed is True
    assert diff.details["table_numeric_similarity"] == 1.0


def test_round51_xlsx_structure_requires_expected_sheets(tmp_path: Path):
    current = tmp_path / "current.xlsx"
    _write_xlsx(current, {"Summary": ["Metric", "Value"]})

    structural = validate_xlsx_structure(current, expected_sheets=("summary", "report_info"))

    assert structural.passed is False
    assert structural.details["missing_expected_sheets"] == ["report_info"]


def test_round51_kpi_sidecar_extracts_and_compares_docx_xlsx_values(tmp_path: Path):
    docx_path = tmp_path / "report.docx"
    xlsx_path = tmp_path / "data.xlsx"
    _write_docx_like(docx_path, "placeholder")

    # Replace the tiny XML docx with a python-docx table so KPI extraction
    # covers the same table shape as real reports.
    from docx import Document

    # Round 52 (Phase 2): use a 3-column header table so the harness'
    # "headers + values row" heuristic applies. 2-column tables are now
    # treated as label/value because the renewal Customer Health Dashboard
    # and leader Team Performance Metrics both use that 2-column shape.
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Total Customers"
    table.rows[0].cells[1].text = "Support Cases"
    table.rows[0].cells[2].text = "Critical (P1)"
    table.rows[1].cells[0].text = "52"
    table.rows[1].cells[1].text = "176"
    table.rows[1].cells[2].text = "5"
    doc.save(docx_path)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.append(["Metric", "Value"])
    sheet.append(["Customers in portfolio", "52"])
    sheet.append(["Total Support Cases", "176"])
    workbook.create_sheet("Report_Info").append(["Field", "Value"])
    workbook.save(xlsx_path)
    workbook.close()

    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(docx_kpis, xlsx_kpis, strict=True)

    assert parity.passed is True
    assert "total_customers" in parity.details["common_kpis"]
    assert "support_cases" in parity.details["common_kpis"]


def test_round97_3_kpi_scan_ignores_partial_data_warning_denominator(tmp_path: Path):
    from docx import Document

    docx_path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph(
        "• adoption_barriers (tech_filter_scope_excluded): AB tech filter "
        "'All Contact Center' kept 14 of 183 adoption barriers and excluded "
        "169 without Contact Center evidence."
    )
    doc.add_paragraph("• Adoption Barriers: 14 row(s) (each row = one in-scope task)")
    doc.save(docx_path)

    values = extract_docx_kpis(docx_path)["values"]

    assert values["adoption_barriers"] == "14"
    assert values["adoption_barriers"] != "183"


def test_round97_3_kpi_scan_ignores_per_customer_action_plan_prose(tmp_path: Path):
    from docx import Document

    docx_path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("Relationship Health: The relationship is managed through 13 action plans.")
    doc.add_paragraph("Total Action Plans: 251")
    doc.save(docx_path)

    values = extract_docx_kpis(docx_path)["values"]

    assert values["action_plans"] == "251"


def test_round97_3_kpi_scan_ignores_open_active_ab_prefix_prose(tmp_path: Path):
    from docx import Document

    docx_path = tmp_path / "ab_prefix.docx"
    doc = Document()
    doc.add_paragraph("Problem Statement: 5 active adoption barriers related to product bugs.")
    doc.add_paragraph("Barrier Categorization: categorize all 14 open adoption barriers.")
    doc.add_paragraph("Open Adoption Barriers: 13")
    doc.save(docx_path)

    values = extract_docx_kpis(docx_path)["values"]

    assert values["open_adoption_barriers"] == "13"
    assert "5" not in values.values()
    assert "14" not in values.values()


def test_round97_3_leader_xlsx_prefers_team_summary_rollup_over_detail_ledger(tmp_path: Path):
    xlsx_path = tmp_path / "leader.xlsx"
    workbook = openpyxl.Workbook()
    action_plans = workbook.active
    action_plans.title = "Action_Plans"
    action_plans.append(["ID", "Status"])
    for idx in range(251):
        action_plans.append([f"AP-{idx}", "Open"])
    team_summary = workbook.create_sheet("Team_Summary")
    team_summary.append(["Team_Member", "Num_Action_Plans"])
    team_summary.append(["One", 130])
    team_summary.append(["Two", 135])
    adoption_barriers = workbook.create_sheet("Adoption_Barriers")
    adoption_barriers.append(["ID", "Status"])
    for idx in range(32):
        adoption_barriers.append([f"AB-{idx}", "Open"])
    workbook.save(xlsx_path)
    workbook.close()

    values = extract_xlsx_kpis(xlsx_path)["values"]

    assert values["action_plans"] == "265"
    assert values["adoption_barriers"] == "32"


def test_round97_3_quality_gate_ignores_numeric_heading_labels(tmp_path: Path):
    from docx import Document

    docx_path = tmp_path / "report.docx"
    xlsx_path = tmp_path / "data.xlsx"
    doc = Document()
    doc.add_heading("Top 10 Focus Accounts by Risk", level=1)
    doc.add_paragraph("Total Customers: 10 [Source: Snowflake EDW Sales subscriptions]")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Total Customers"
    table.rows[1].cells[1].text = "10 [Source: Snowflake EDW Sales subscriptions]"
    doc.save(docx_path)
    _write_xlsx(xlsx_path, {"Summary": ["Metric", "Value"], "Report_Info": ["Item", "Value"]})

    payload, gate = evaluate_report_quality(docx_path, xlsx_path, scenario_key="renewal", strict=True)

    assert gate.passed is True
    assert payload["uncited_numeric_paragraph_count"] == 0


def test_round51_build_runner_config_strict_raises_thresholds():
    config = build_runner_config(_args(strict=True, min_docx_similarity=0.1))

    assert config.strict is True
    assert config.min_docx_similarity >= 0.55
    assert config.min_sheet_overlap >= 0.85
    assert config.min_header_similarity >= 0.8


def test_round51_local_http_headers_forward_secure_session_cookie(tmp_path: Path):
    config = build_runner_config(_args(downloads_dir=str(tmp_path)))
    runner = LiveReportRunner(config)
    runner.csrf_token = "csrf"
    runner.session_cookie_value = "session-value"

    headers = runner._headers(include_json_content_type=True)

    assert headers["X-CSRFToken"] == "csrf"
    assert headers["Cookie"] == "session=session-value"
    assert headers["Content-Type"] == "application/json"


def test_round97_2_app_health_preflight_passes_with_required_endpoints():
    session = _FakeSession(
        {
            "/ping": _FakeResponse(text="OK"),
            "/api/version": _FakeResponse(
                json_payload={
                    "ok": True,
                    "version": "1.0.4",
                    "build": "69",
                    "process_started_at_utc": "2026-05-22T00:00:00Z",
                    "restart_required": False,
                }
            ),
            "/api/status/all": _FakeResponse(json_payload={"analyses": []}),
            "/api/corpus/status": _FakeResponse(json_payload={"boot": {"source": "fresh"}}),
            "/api/intel/status": _FakeResponse(json_payload={"boot": {"source": "fresh"}}),
        }
    )

    payload, gate = evaluate_app_health(session, "http://127.0.0.1:5151")

    assert gate.passed is True
    assert payload["failed_required_paths"] == []
    assert [probe["path"] for probe in payload["probes"]] == [
        "/ping",
        "/api/version",
        "/api/status/all",
        "/api/corpus/status",
        "/api/intel/status",
    ]


def test_round97_2_app_health_preflight_fails_closed_on_bad_ping():
    session = _FakeSession(
        {
            "/ping": _FakeResponse(text="NOT READY"),
            "/api/version": _FakeResponse(
                json_payload={
                    "ok": True,
                    "version": "1.0.4",
                    "build": "69",
                    "process_started_at_utc": "2026-05-22T00:00:00Z",
                    "restart_required": False,
                }
            ),
            "/api/status/all": _FakeResponse(json_payload={"analyses": []}),
            "/api/corpus/status": _FakeResponse(json_payload={"boot": {}}),
            "/api/intel/status": _FakeResponse(json_payload={"boot": {}}),
        }
    )

    payload, gate = evaluate_app_health(session, "http://127.0.0.1:5151")

    assert gate.passed is False
    assert payload["failed_required_paths"] == ["/ping"]


def test_round97_2_bootstrap_session_runs_health_before_homepage(tmp_path: Path):
    session = _FakeSession(
        {
            "/ping": _FakeResponse(text="OK"),
            "/api/version": _FakeResponse(
                json_payload={
                    "ok": True,
                    "version": "1.0.4",
                    "build": "69",
                    "process_started_at_utc": "2026-05-22T00:00:00Z",
                    "restart_required": False,
                }
            ),
            "/api/status/all": _FakeResponse(json_payload={"analyses": []}),
            "/api/corpus/status": _FakeResponse(json_payload={"boot": {"source": "fresh"}}),
            "/api/intel/status": _FakeResponse(json_payload={"boot": {"source": "fresh"}}),
            "/": _FakeResponse(text='<meta name="csrf-token" content="token123" />'),
        }
    )
    config = build_runner_config(_args(downloads_dir=str(tmp_path)))
    runner = LiveReportRunner(config)
    runner.session = session

    runner.bootstrap_session()

    assert runner.csrf_token == "token123"
    assert runner.app_health_gate.passed is True
    assert session.calls[-1][0] == "http://127.0.0.1:5151/"
