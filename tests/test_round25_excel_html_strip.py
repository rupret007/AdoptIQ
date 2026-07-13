"""Round 25 / Phase F.2: Excel writer must strip HTML from object columns.

Pre-Round 25 the Excel writer emitted whatever string lived in object
columns -- including raw HTML markup leaking out of Snowflake views
that store rich-text descriptions.  The reference Brian Frazier /
All Contact Center / 90d report had cells like::

    <a href="/" target="_blank"> </a>
    <img src="/lightning/r/Account/0014..." />
    <p>Customer asked about <strong>renewal</strong> options.</p>

Excel renders ``<`` literally so the recipient saw raw markup instead
of the intended text.  Phase F.2 introduces
``data_normalization.strip_html_from_dataframe`` and applies it to
``AB_Detail_All`` and ``CSConsole_Customer_Pulse`` immediately before
``df.to_excel(...)``.

This test enforces:

1.  The standalone helper drops well-formed tags and unescapes HTML
    entities while leaving plain text untouched.
2.  Cells without ``<`` are returned unchanged (no churn on clean
    exports).
3.  An end-to-end ``write_excel_workbook`` smoke produces an
    ``AB_Detail_All`` sheet with no ``<`` characters in object cells.
4.  An end-to-end ``write_excel_workbook`` smoke produces a
    ``CSConsole_Customer_Pulse`` sheet with no ``<`` characters.
"""

from __future__ import annotations

import os
import tempfile

import pandas as pd
import pytest

from openpyxl import load_workbook

import adoptiq_backend
from data_normalization import (
    strip_html_from_dataframe,
    strip_html_from_string,
)


def test_strip_html_helper_drops_tags_and_unescapes_entities() -> None:
    """Single-cell helper is the truth source for the dataframe pass."""

    raw = '<a href="/" target="_blank">click here</a>'
    assert strip_html_from_string(raw) == "click here", (
        "Round 25 / Phase F.2: anchor tags must be stripped, leaving "
        "the inner text content intact."
    )

    raw_img = '<img src="/lightning/r/Account/0014..." />'
    assert strip_html_from_string(raw_img) == "", (
        "Round 25 / Phase F.2: self-closing image tags must collapse "
        "to an empty string."
    )

    raw_compound = (
        "<p>Customer asked about <strong>renewal</strong> options "
        "&amp; pricing.</p>"
    )
    cleaned = strip_html_from_string(raw_compound)
    assert "<" not in cleaned and ">" not in cleaned, (
        f"Round 25 / Phase F.2: stripped output must not contain "
        f"any angle brackets.  Got: {cleaned!r}"
    )
    assert "&amp;" not in cleaned, (
        f"Round 25 / Phase F.2: HTML entities must be unescaped.  "
        f"Got: {cleaned!r}"
    )
    assert "Customer asked about renewal options & pricing." in cleaned


def test_strip_html_helper_passes_plain_text_through_unchanged() -> None:
    """Strings without ``<`` are returned identical (no churn)."""

    assert strip_html_from_string("plain text") == "plain text"
    assert strip_html_from_string("") == ""
    assert strip_html_from_string(42) == 42  # non-string passes through


def test_strip_html_helper_handles_dataframe_in_place_safe_copy() -> None:
    """Dataframe pass must rewrite object columns without mutating input."""

    df = pd.DataFrame(
        {
            "ID": ["AB-1", "AB-2", "AB-3"],
            "Account": [
                '<a href="/">ACME</a>',
                "Plain ACME 2",
                '<img src="/foo.png">ACME 3',
            ],
            "Numeric": [1, 2, 3],
        }
    )
    out = strip_html_from_dataframe(df)
    # Original frame untouched (we copy before rewrite).
    assert df.loc[0, "Account"] == '<a href="/">ACME</a>'
    # Rewritten frame has clean text.
    assert out.loc[0, "Account"] == "ACME"
    assert out.loc[1, "Account"] == "Plain ACME 2"  # untouched
    assert out.loc[2, "Account"] == "ACME 3"
    # Numeric columns are untouched.
    assert list(out["Numeric"]) == [1, 2, 3]


def test_excel_write_strips_html_from_ab_detail_all() -> None:
    """End-to-end: workbook ``AB_Detail_All`` sheet has no HTML markup."""

    # Use canonical AB_Detail_All column names so the Round 15 export
    # schema keeps them in the curated output.  The HTML strip pass
    # runs before the schema pass, so the test still exercises the
    # Phase F.2 code path even when the output schema narrows columns.
    ab_df = pd.DataFrame(
        [
            {
                "ID": "AB-001",
                "NAME": '<a href="/" target="_blank">ACME Corp</a>',
                "DESCRIPTION_C": "<p>Customer is <strong>blocked</strong> on rollout.</p>",
            },
            {
                "ID": "AB-002",
                "NAME": '<img src="/lightning/r/Account/0014..." />NATIONAL GRID',
                "DESCRIPTION_C": "Plain text -- no markup here.",
            },
            {
                "ID": "AB-003",
                "NAME": "FARMERS",
                "DESCRIPTION_C": 'Reference: <a href="/sf">SF link</a>',
            },
        ]
    )
    sheets = {"AB_Detail_All": ab_df}

    with tempfile.TemporaryDirectory() as tmp:
        base = os.path.join(tmp, "round25_phaseF_html_strip_ab")
        out = adoptiq_backend.write_excel_workbook(base, sheets)
        assert os.path.exists(out)

        wb = load_workbook(out, data_only=True)
        assert "AB_Detail_All" in wb.sheetnames, (
            f"AB_Detail_All sheet missing.  Got sheets: {wb.sheetnames!r}"
        )
        ws = wb["AB_Detail_All"]

        # Sweep every object cell looking for residual ``<`` -- the
        # writer must not ship any HTML markup on this sheet.
        for row in ws.iter_rows(min_row=2, values_only=True):
            for v in row:
                if isinstance(v, str):
                    assert "<" not in v, (
                        f"Round 25 / Phase F.2: AB_Detail_All cell "
                        f"still contains HTML markup: {v!r}.  Strip "
                        f"helper failed to fire on the writer's hot "
                        f"path."
                    )

        # Spot-check the cleaned values render the intended text --
        # we walk every cell rather than indexing by column position
        # because the Round 15 export schema may re-order / drop
        # columns on the way out.
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        joined = " | ".join(str(v) for r in rows for v in r if v is not None)
        assert "ACME Corp" in joined, (
            f"Round 25 / Phase F.2: cleaned ACME Corp text missing "
            f"from AB_Detail_All export.  Got: {joined!r}"
        )
        assert "NATIONAL GRID" in joined, (
            f"Round 25 / Phase F.2: cleaned NATIONAL GRID text "
            f"missing from AB_Detail_All export.  Got: {joined!r}"
        )


def test_excel_write_strips_html_from_csconsole_customer_pulse() -> None:
    """End-to-end: workbook ``CSConsole_Customer_Pulse`` sheet has no HTML."""

    pulse_df = pd.DataFrame(
        [
            {
                "ACCOUNT_NAME": '<a href="/">Pulse ACME</a>',
                "PULSE_HEALTH": "<p>Healthy</p>",
            },
            {
                "ACCOUNT_NAME": "Pulse FARMERS",
                "PULSE_HEALTH": '<img src="/foo.png" />Critical',
            },
        ]
    )

    with tempfile.TemporaryDirectory() as tmp:
        base = os.path.join(tmp, "round25_phaseF_html_strip_pulse")
        out = adoptiq_backend.write_excel_workbook(
            base,
            {"AB_Detail_All": pd.DataFrame([{"ID": "AB-1"}])},  # placeholder
            csconsole_data={"customer_pulse": pulse_df},
        )
        assert os.path.exists(out)

        wb = load_workbook(out, data_only=True)
        assert "CSConsole_Customer_Pulse" in wb.sheetnames, (
            f"CSConsole_Customer_Pulse sheet missing.  Got sheets: "
            f"{wb.sheetnames!r}"
        )
        ws = wb["CSConsole_Customer_Pulse"]

        for row in ws.iter_rows(min_row=2, values_only=True):
            for v in row:
                if isinstance(v, str):
                    assert "<" not in v, (
                        f"Round 25 / Phase F.2: "
                        f"CSConsole_Customer_Pulse cell still contains "
                        f"HTML markup: {v!r}."
                    )
