"""Round 52 (Phase 1): manifest-backed baselines + --init-baseline tests."""

from __future__ import annotations

import json
import zipfile
from argparse import Namespace
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from report_iteration_loop import (
    BaselineEntry,
    LiveReportRunner,
    MANIFEST_SCHEMA_VERSION,
    build_arg_parser,
    build_runner_config,
    capture_baseline_artifact,
    load_baseline_manifest,
    verify_baseline_entry,
    write_baseline_manifest,
)


def _write_docx_like(path: Path, text: str) -> None:
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
    <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:body>
        <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
      </w:body>
    </w:document>
    """
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", xml)


def _write_xlsx(path: Path, sheets: dict[str, list[str]]) -> None:
    workbook = openpyxl.Workbook()
    first = True
    for name, headers in sheets.items():
        if first:
            sheet = workbook.active
            sheet.title = name
            first = False
        else:
            sheet = workbook.create_sheet(title=name)
        for idx, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=idx, value=header)
        sheet.cell(row=2, column=1, value="rowval1")
        sheet.cell(row=2, column=2, value="rowval2")
    workbook.save(path)
    workbook.close()


def _args(**overrides: Any) -> Namespace:
    defaults: dict[str, Any] = {
        "base_url": "http://127.0.0.1:5151",
        "downloads_dir": "~/Downloads",
        "iterations": 1,
        "scenarios": "all",
        "poll_interval": 5.0,
        "timeout": 1800,
        "run_id": "test",
        "stop_on_failure": False,
        "baseline_mode": "latest",
        "baseline_manifest": "",
        "init_baseline": False,
        "init_baseline_dir": "",
        "init_baseline_label": "",
        "strict": False,
        "min_docx_similarity": 0.35,
        "min_sheet_overlap": 0.5,
        "min_header_similarity": 0.3,
        "min_docx_chars": 200,
        "min_docx_numeric_similarity": 0.8,
        "max_xlsx_row_delta_ratio": 0.2,
        "max_xlsx_row_delta_abs": 25,
    }
    defaults.update(overrides)
    return Namespace(**defaults)


def test_round52_argparser_supports_off_latest_and_manifest_modes():
    parser = build_arg_parser()
    parsed = parser.parse_args(
        ["--baseline-mode", "manifest", "--baseline-manifest", "/tmp/m.json"]
    )
    assert parsed.baseline_mode == "manifest"
    assert parsed.baseline_manifest == "/tmp/m.json"

    parsed_off = parser.parse_args(["--baseline-mode", "off"])
    assert parsed_off.baseline_mode == "off"

    with pytest.raises(SystemExit):
        parser.parse_args(["--baseline-mode", "bogus"])


def test_round52_build_runner_config_requires_manifest_path_when_mode_manifest():
    with pytest.raises(ValueError, match="--baseline-manifest"):
        build_runner_config(_args(baseline_mode="manifest"))


def test_round52_build_runner_config_requires_dir_when_init_baseline_set():
    with pytest.raises(ValueError, match="--init-baseline-dir"):
        build_runner_config(_args(init_baseline=True))


def test_round52_write_and_load_manifest_round_trip(tmp_path: Path):
    docx_path = tmp_path / "captured" / "comprehensive" / "report.docx"
    docx_path.parent.mkdir(parents=True)
    _write_docx_like(docx_path, "Round 52 baseline content for manifest round trip.")
    xlsx_path = tmp_path / "captured" / "comprehensive" / "data.xlsx"
    _write_xlsx(xlsx_path, {"Summary": ["Metric", "Value"]})

    manifest_path = tmp_path / "captured" / "baseline_manifest.json"
    payload = write_baseline_manifest(
        manifest_path,
        {"comprehensive": {"docx": docx_path, "xlsx": xlsx_path}},
        label="round-52-test",
    )

    assert payload["version"] == MANIFEST_SCHEMA_VERSION
    assert payload["label"] == "round-52-test"
    assert "comprehensive" in payload["scenarios"]
    docx_entry = payload["scenarios"]["comprehensive"]["docx"]
    assert docx_entry["sha256"]
    assert docx_entry["size_bytes"] == docx_path.stat().st_size

    loaded = load_baseline_manifest(manifest_path)
    assert "comprehensive" in loaded
    assert "docx" in loaded["comprehensive"]
    assert loaded["comprehensive"]["docx"].sha256 == docx_entry["sha256"]
    assert loaded["comprehensive"]["docx"].path.resolve() == docx_path.resolve()


def test_round52_load_manifest_rejects_unsupported_version(tmp_path: Path):
    manifest = tmp_path / "bad.json"
    manifest.write_text(json.dumps({"version": 99, "scenarios": {}}), encoding="utf-8")
    with pytest.raises(ValueError, match="version"):
        load_baseline_manifest(manifest)


def test_round52_load_manifest_rejects_invalid_sha(tmp_path: Path):
    docx_path = tmp_path / "report.docx"
    _write_docx_like(docx_path, "x")
    manifest = tmp_path / "bad.json"
    manifest.write_text(
        json.dumps(
            {
                "version": MANIFEST_SCHEMA_VERSION,
                "scenarios": {
                    "comprehensive": {
                        "docx": {
                            "path": "report.docx",
                            "sha256": "tooshort",
                            "size_bytes": 1,
                        }
                    }
                },
            }
        ),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="sha256"):
        load_baseline_manifest(manifest)


def test_round52_verify_baseline_entry_detects_size_mismatch(tmp_path: Path):
    docx_path = tmp_path / "report.docx"
    _write_docx_like(docx_path, "Round 52 size guard.")
    entry = BaselineEntry(
        scenario_key="comprehensive",
        file_type="docx",
        path=docx_path,
        sha256="0" * 64,
        size_bytes=1,
        captured_at_utc=None,
    )
    result = verify_baseline_entry(entry)
    assert result.passed is False
    assert result.details["reason"] == "baseline_size_mismatch"


def test_round52_verify_baseline_entry_detects_sha_mismatch(tmp_path: Path):
    docx_path = tmp_path / "report.docx"
    _write_docx_like(docx_path, "Round 52 sha guard.")
    entry = BaselineEntry(
        scenario_key="comprehensive",
        file_type="docx",
        path=docx_path,
        sha256="0" * 64,
        size_bytes=docx_path.stat().st_size,
        captured_at_utc=None,
    )
    result = verify_baseline_entry(entry)
    assert result.passed is False
    assert result.details["reason"] == "baseline_sha256_mismatch"


def test_round52_verify_baseline_entry_passes_when_file_matches(tmp_path: Path):
    docx_path = tmp_path / "report.docx"
    _write_docx_like(docx_path, "Round 52 happy path.")
    from report_iteration_loop import _file_sha256

    entry = BaselineEntry(
        scenario_key="comprehensive",
        file_type="docx",
        path=docx_path,
        sha256=_file_sha256(docx_path),
        size_bytes=docx_path.stat().st_size,
        captured_at_utc=None,
    )
    result = verify_baseline_entry(entry)
    assert result.passed is True
    assert result.details["reason"] == "verified"


def test_round52_runner_off_mode_passes_baseline_diff_without_lookup(tmp_path: Path):
    config = build_runner_config(_args(downloads_dir=str(tmp_path), baseline_mode="off"))
    assert config.baseline_mode == "off"
    runner = LiveReportRunner(config)
    assert runner.manifest == {}


def test_round52_runner_manifest_mode_loads_pinned_files(tmp_path: Path):
    docx_path = tmp_path / "comprehensive" / "report.docx"
    docx_path.parent.mkdir(parents=True)
    _write_docx_like(docx_path, "Round 52 manifest fixture.")
    xlsx_path = tmp_path / "comprehensive" / "data.xlsx"
    _write_xlsx(xlsx_path, {"Summary": ["Metric", "Value"]})

    manifest_path = tmp_path / "baseline_manifest.json"
    write_baseline_manifest(
        manifest_path,
        {"comprehensive": {"docx": docx_path, "xlsx": xlsx_path}},
    )

    config = build_runner_config(
        _args(
            downloads_dir=str(tmp_path),
            baseline_mode="manifest",
            baseline_manifest=str(manifest_path),
        )
    )
    runner = LiveReportRunner(config)
    assert "comprehensive" in runner.manifest
    assert "docx" in runner.manifest["comprehensive"]
    assert "xlsx" in runner.manifest["comprehensive"]


def test_round52_capture_baseline_artifact_copies_into_scenario_subdir(tmp_path: Path):
    src = tmp_path / "AdoptIQ_Report_x.docx"
    _write_docx_like(src, "Round 52 capture fixture.")
    dest_root = tmp_path / "captured"
    target = capture_baseline_artifact(dest_root, "comprehensive", src)
    assert target.parent == dest_root / "comprehensive"
    assert target.exists()
    assert target.read_bytes() == src.read_bytes()


def test_round52_thresholds_summary_records_baseline_mode_and_manifest_path(tmp_path: Path):
    docx_path = tmp_path / "comprehensive" / "report.docx"
    docx_path.parent.mkdir(parents=True)
    _write_docx_like(docx_path, "Round 52 thresholds fixture.")
    xlsx_path = tmp_path / "comprehensive" / "data.xlsx"
    _write_xlsx(xlsx_path, {"Summary": ["Metric", "Value"]})
    manifest_path = tmp_path / "baseline_manifest.json"
    write_baseline_manifest(
        manifest_path,
        {"comprehensive": {"docx": docx_path, "xlsx": xlsx_path}},
    )

    config = build_runner_config(
        _args(
            downloads_dir=str(tmp_path),
            baseline_mode="manifest",
            baseline_manifest=str(manifest_path),
        )
    )

    from report_iteration_loop import thresholds_summary

    summary = thresholds_summary(config)
    assert summary["baseline_mode"] == "manifest"
    assert summary["baseline_manifest_path"] == str(manifest_path)
    assert summary["init_baseline"] is False
