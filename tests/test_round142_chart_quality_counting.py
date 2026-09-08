"""Round 142: raster chart images count in DOCX quality inspection."""

from __future__ import annotations

import base64
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook

from decision_report_delivery import SOURCE_DATA_SHEET_NAMES
from report_iteration_loop import (
    _count_docx_chart_parts,
    _inspect_docx_visible_charts,
    evaluate_report_quality,
)


_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _write_png(path: Path) -> None:
    path.write_bytes(_ONE_PIXEL_PNG)


def _write_quality_ready_docx(path: Path) -> Document:
    doc = Document()
    doc.add_heading("Portfolio Overview", level=1)
    doc.add_heading("Operational Trends", level=2)
    doc.add_paragraph("Current portfolio observations are summarized below.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Signal"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Adoption"
    table.cell(1, 1).text = "Stable"
    doc.save(path)
    return doc


def _write_canonical_chart_workbook(
    path: Path,
    *,
    renderable: tuple[str, ...],
    expose_partial_value: bool = False,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in SOURCE_DATA_SHEET_NAMES:
        workbook.create_sheet(sheet_name)
    chart = workbook["Chart_Data"]
    chart.append(
        ["Metric_Key", "Chart_ID", "Value", "Source_State", "Period_Start"]
    )
    metric_keys = {
        "activity_mix": "chart.activity_mix.action_plans",
        "action_plan_status_aging": "chart.action_plan_status.open",
        "risk_distribution": "chart.risk_distribution.low",
        "activity_trend": "chart.activity_trend.action_plans.2026_08",
    }
    for chart_id, metric_key in metric_keys.items():
        available = chart_id in renderable
        value = 1 if available else None
        state = "available" if available else "partial"
        if expose_partial_value and chart_id == "risk_distribution":
            value = 1
            state = "available"
        chart.append(
            [
                metric_key,
                chart_id,
                value,
                state,
                "2026-08-01T00:00:00Z" if chart_id == "activity_trend" else None,
            ]
        )
    if expose_partial_value:
        # A hostile row must not expose a value while declaring itself partial.
        # Mixed groups are valid when incomplete sibling rows are withheld.
        chart.append(
            [
                "chart.risk_distribution.high",
                "risk_distribution",
                99,
                "partial",
                None,
            ]
        )
    workbook.save(path)
    workbook.close()


def _add_accessible_charts(doc: Document, png_path: Path, count: int) -> None:
    for index in range(count):
        picture = doc.add_picture(str(png_path), width=Inches(2.5))
        picture._inline.docPr.set("title", f"Canonical chart {index + 1}")
        picture._inline.docPr.set(
            "descr",
            f"Canonical source-backed chart {index + 1}.",
        )


def test_embedded_png_with_dimensions_and_alt_metadata_counts_as_visible_chart(
    tmp_path: Path,
) -> None:
    png_path = tmp_path / "action-plan-status.png"
    docx_path = tmp_path / "report-with-raster-chart.docx"
    _write_png(png_path)

    doc = _write_quality_ready_docx(docx_path)
    picture = doc.add_picture(str(png_path), width=Inches(2.5))
    picture._inline.docPr.set("title", "Action Plan Status")
    picture._inline.docPr.set(
        "descr",
        "Action plan status distribution chart based on tracked plan records.",
    )
    doc.save(docx_path)

    metadata = _inspect_docx_visible_charts(docx_path)

    assert metadata["visible_chart_count"] == 1
    assert metadata["native_chart_count"] == 0
    assert metadata["embedded_raster_chart_count"] == 1
    assert metadata["inspection_errors"] == []
    embedded = metadata["embedded_raster_charts"][0]
    assert embedded["width_emu"] > 0
    assert embedded["height_emu"] > 0
    assert embedded["title"] == "Action Plan Status"
    assert embedded["alt_text"].startswith("Action plan status distribution chart")
    assert embedded["media_part"].endswith(".png")
    assert _count_docx_chart_parts(docx_path) == 1

    payload, _gate = evaluate_report_quality(
        docx_path,
        None,
        scenario_key="comprehensive",
        strict=False,
    )
    assert payload["chart_count"] == 1
    assert payload["chart_metadata"]["embedded_raster_chart_count"] == 1


def test_document_without_chart_has_zero_visible_charts(tmp_path: Path) -> None:
    docx_path = tmp_path / "report-without-chart.docx"
    _write_quality_ready_docx(docx_path)

    metadata = _inspect_docx_visible_charts(docx_path)

    assert metadata["visible_chart_count"] == 0
    assert metadata["native_chart_count"] == 0
    assert metadata["embedded_raster_chart_count"] == 0
    assert metadata["embedded_raster_charts"] == []
    assert _count_docx_chart_parts(docx_path) == 0


def test_partial_canonical_workbook_requires_only_renderable_charts(
    tmp_path: Path,
) -> None:
    png_path = tmp_path / "chart.png"
    docx_path = tmp_path / "partial.docx"
    xlsx_path = tmp_path / "partial.xlsx"
    _write_png(png_path)
    doc = _write_quality_ready_docx(docx_path)
    _add_accessible_charts(doc, png_path, 1)
    doc.save(docx_path)
    _write_canonical_chart_workbook(
        xlsx_path,
        renderable=("action_plan_status_aging",),
    )

    payload, gate = evaluate_report_quality(
        docx_path,
        xlsx_path,
        scenario_key="comprehensive",
        strict=True,
        expected_min_charts=4,
    )

    assert gate.passed is True, gate.details
    assert payload["configured_min_charts"] == 4
    assert payload["expected_min_charts"] == 1
    assert payload["canonical_chart_contract"]["contract_valid"] is True
    assert payload["canonical_chart_contract"]["renderable_chart_count"] == 1
    assert payload["canonical_chart_contract"]["withheld_chart_count"] == 3


def test_complete_canonical_workbook_retains_full_chart_minimum(
    tmp_path: Path,
) -> None:
    png_path = tmp_path / "chart.png"
    docx_path = tmp_path / "complete.docx"
    xlsx_path = tmp_path / "complete.xlsx"
    _write_png(png_path)
    doc = _write_quality_ready_docx(docx_path)
    _add_accessible_charts(doc, png_path, 1)
    doc.save(docx_path)
    _write_canonical_chart_workbook(
        xlsx_path,
        renderable=(
            "activity_mix",
            "action_plan_status_aging",
            "risk_distribution",
            "activity_trend",
        ),
    )

    payload, gate = evaluate_report_quality(
        docx_path,
        xlsx_path,
        scenario_key="comprehensive",
        strict=True,
        expected_min_charts=4,
    )

    assert gate.passed is False
    assert payload["expected_min_charts"] == 4
    assert payload["canonical_chart_contract"]["renderable_chart_count"] == 4
    assert any("requires at least 4" in error for error in payload["errors"])


def test_malformed_canonical_chart_contract_never_lowers_quality_floor(
    tmp_path: Path,
) -> None:
    png_path = tmp_path / "chart.png"
    docx_path = tmp_path / "malformed.docx"
    xlsx_path = tmp_path / "malformed.xlsx"
    _write_png(png_path)
    doc = _write_quality_ready_docx(docx_path)
    _add_accessible_charts(doc, png_path, 4)
    doc.save(docx_path)
    _write_canonical_chart_workbook(
        xlsx_path,
        renderable=(
            "activity_mix",
            "action_plan_status_aging",
            "activity_trend",
        ),
        expose_partial_value=True,
    )

    payload, gate = evaluate_report_quality(
        docx_path,
        xlsx_path,
        scenario_key="comprehensive",
        strict=True,
        expected_min_charts=4,
    )

    assert gate.passed is False
    assert payload["configured_min_charts"] == 4
    assert payload["expected_min_charts"] == 4
    assert payload["canonical_chart_contract"]["canonical_workbook_detected"] is True
    assert payload["canonical_chart_contract"]["contract_valid"] is False
    assert any("renderability contract is invalid" in error for error in payload["errors"])


def test_unlabelled_embedded_png_does_not_count_as_chart(tmp_path: Path) -> None:
    png_path = tmp_path / "decorative-image.png"
    docx_path = tmp_path / "report-with-decorative-image.docx"
    _write_png(png_path)

    doc = Document()
    doc.add_picture(str(png_path), width=Inches(1))
    doc.save(docx_path)

    metadata = _inspect_docx_visible_charts(docx_path)

    assert metadata["visible_chart_count"] == 0
    assert metadata["embedded_raster_chart_count"] == 0


def test_native_chart_parts_remain_supported(tmp_path: Path) -> None:
    docx_path = tmp_path / "report-with-native-chart-part.docx"
    _write_quality_ready_docx(docx_path)
    with zipfile.ZipFile(docx_path, mode="a") as archive:
        archive.writestr(
            "word/charts/chart1.xml",
            "<c:chartSpace xmlns:c=\"http://schemas.openxmlformats.org/drawingml/2006/chart\"/>",
        )

    metadata = _inspect_docx_visible_charts(docx_path)

    assert metadata["visible_chart_count"] == 1
    assert metadata["native_chart_count"] == 1
    assert metadata["native_chart_parts"] == ["word/charts/chart1.xml"]
    assert metadata["embedded_raster_chart_count"] == 0
    assert _count_docx_chart_parts(docx_path) == 1
