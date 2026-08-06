"""Round 94 full-sweep regressions.

These tests pin the actionable findings from the read-only sweep before
Build 67: citation regex parity, phase timing fidelity, vector bake
failure semantics, corpus quality-gate scope, hybrid retrieval fallback,
formatter markdown safety, CSConsole technology strictness, and frontend
null/failure handling.
"""

from __future__ import annotations
from source_shape_utils import assert_in_source

import json
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
import pytest


ROOT = Path(__file__).resolve().parents[1]


def _read(rel_path: str) -> str:
    return (ROOT / rel_path).read_text(encoding="utf-8")


def test_round94_injector_regex_matches_gate_and_rejects_case_ids() -> None:
    import report_iteration_loop as gate
    import report_source_injector as injector

    assert injector._PARAGRAPH_KPI_NUMERIC_RE.pattern == gate._PARAGRAPH_KPI_NUMERIC_RE.pattern
    assert gate._PARAGRAPH_KPI_NUMERIC_RE.search("Case: 700356476") is None
    assert injector._PARAGRAPH_KPI_NUMERIC_RE.search("Case: 700356476") is None
    assert injector._PARAGRAPH_KPI_NUMERIC_RE.search("Total Customers: 52")


def test_round94_update_analysis_status_records_phase_timing_on_step_change() -> None:
    import app_simple

    aid = "Round94_Phase_Timing"
    with app_simple.analysis_status_lock:
        app_simple.analysis_status[aid] = {
            "status": "running",
            "progress": 8,
            "message": "Connecting",
            "current_step": "Database Connection",
            "step_start_time": datetime.now(timezone.utc),
            "completed_steps": [],
            "phase_timings": {},
        }
    try:
        app_simple.update_analysis_status(
            aid,
            {
                "progress": 20,
                "message": "Fetching team subscriptions",
                "current_step": "Team Data Retrieval",
            },
            save=False,
        )
        with app_simple.analysis_status_lock:
            status = dict(app_simple.analysis_status[aid])
        assert "Database Connection" in status["phase_timings"]
        assert status["completed_steps"] == ["Database Connection"]
        assert status["current_step"] == "Team Data Retrieval"
    finally:
        with app_simple.analysis_status_lock:
            app_simple.analysis_status.pop(aid, None)


def test_round94_bake_vector_failure_is_nonzero_and_does_not_persist(monkeypatch, tmp_path) -> None:
    import corpus_crypto
    import corpus_indexer
    import scripts.bake_corpus as bake

    downloads = tmp_path / "downloads"
    downloads.mkdir()
    bake_dir = tmp_path / "bake"
    bake_dir.mkdir()
    onedrive_root = tmp_path / "onedrive"
    onedrive_root.mkdir()
    sentinel = onedrive_root / corpus_crypto.DEFAULT_SENTINEL_NAME
    sentinel.write_text("{}", encoding="utf-8")

    class FakeHandle:
        conn = object()

        def __init__(self) -> None:
            self.persist_calls: list[bool] = []

        def commit_to_disk(self) -> None:  # pragma: no cover - should not be called
            raise AssertionError("vector failure must fail before commit")

        def close(self, *, persist: bool = True) -> None:
            self.persist_calls.append(persist)

    handle = FakeHandle()
    monkeypatch.setattr(corpus_crypto, "resolve_sentinel_path", lambda _root: sentinel)
    monkeypatch.setattr(corpus_crypto, "open_corpus_for_user", lambda **_kwargs: handle)
    monkeypatch.setattr(
        corpus_indexer,
        "index_folder",
        lambda _conn, _root: SimpleNamespace(files_seen=1, files_skipped=0, files_oversized=0, errors=[]),
    )
    monkeypatch.setattr(bake, "_bake_chunk_vectors", lambda _conn: (_ for _ in ()).throw(RuntimeError("missing embedder")))

    rc = bake._index_into_encrypted_corpus(downloads, bake_dir, onedrive_root)

    assert rc == 7
    assert handle.persist_calls == [False]


def test_round94_generated_reports_need_quality_sidecar_in_any_corpus_source(tmp_path) -> None:
    from corpus_indexer import enumerate_corpus_files

    approved = tmp_path / "AdoptIQ_Report_Approved.docx"
    missing = tmp_path / "AdoptIQ_Report_Missing.docx"
    csone_dump = tmp_path / "AdoptIQ Enhanced Premium Collab Summary.xlsx"
    for path in (approved, missing, csone_dump):
        path.write_text("placeholder", encoding="utf-8")
    approved.with_name(approved.name + ".adoptiq_corpus.json").write_text(
        json.dumps({"corpus_eligible": True}),
        encoding="utf-8",
    )

    files = enumerate_corpus_files(tmp_path, recursive=False, require_adoptiq_quality_gate=True)
    names = [f.filename for f in files]

    assert approved.name in names
    assert csone_dump.name in names
    assert missing.name not in names


def test_round94_rank_evidence_falls_back_when_hybrid_raises(monkeypatch) -> None:
    from ask_ai_grounded import EvidenceRecord, rank_evidence
    from config import Config

    records = [
        EvidenceRecord(
            source_id="REC-1",
            source_type="snowflake",
            customer="Acme",
            timestamp="2026-01-01",
            text="critical renewal risk from support cases",
            confidence=0.9,
        )
    ]
    original = Config.ASK_AI_RETRIEVAL_METHOD
    try:
        Config.ASK_AI_RETRIEVAL_METHOD = "hybrid"
        monkeypatch.setattr("ask_ai_embeddings.embed_query", lambda _q: (_ for _ in ()).throw(RuntimeError("onnx blew up")))
        ranked = rank_evidence(records, "support cases", domains=["snowflake"])
    finally:
        Config.ASK_AI_RETRIEVAL_METHOD = original
    assert [r.source_id for r in ranked] == ["REC-1"]
    assert ranked[0].rrf_score is None


def test_round94_csconsole_no_tech_match_does_not_account_fallback() -> None:
    from adoptiq_backend import _filter_csconsole_data_by_technology

    df = pd.DataFrame(
        [
            {
                "ID": "AP-1",
                "ACCOUNT_ID_C": "0011",
                "SUBJECT_C": "Webex Meetings rollout",
                "DESCRIPTION_C": "Meetings adoption item",
                "BU_NAME": "Acme Corp",
            }
        ]
    )
    out = _filter_csconsole_data_by_technology(
        df,
        "All Contact Center",
        customer_names=["Acme Corp"],
        account_ids=["0011"],
    )
    assert out.empty


def test_round94_leader_detailed_ab_handles_none_and_strips_markdown() -> None:
    pytest.importorskip("docx")
    from docx import Document
    from leader_report_generator import LeaderReportGenerator

    gen = LeaderReportGenerator.__new__(LeaderReportGenerator)
    gen.doc = Document()
    gen._create_detailed_ab_list({"No Data CSSM": {"adoption_barriers": None}})
    assert "No adoption barriers found" in "\n".join(p.text for p in gen.doc.paragraphs)

    gen = LeaderReportGenerator.__new__(LeaderReportGenerator)
    gen.doc = Document()
    gen._create_detailed_ab_list(
        {
            "CSSM One": {
                "adoption_barriers": pd.DataFrame(
                    [
                        {
                            "ID": "AB-1",
                            "BU_NAME": "Acme Corp",
                            "SUBJECT_C": "**Blocked rollout**",
                            "SUB_TECHNOLOGY_C": "Webex Contact Center",
                            "AB_CATEGORY_C": "Training",
                            "SEVERITY_C": "High",
                        }
                    ]
                )
            }
        }
    )
    table_text = "\n".join(cell.text for table in gen.doc.tables for row in table.rows for cell in row.cells)
    assert "**" not in table_text
    assert "Blocked rollout" in table_text


def test_round94_compact_voice_of_customer_strips_markdown_subjects() -> None:
    pytest.importorskip("docx")
    from compact_report_formatter import CompactReportFormatter

    fmt = CompactReportFormatter()
    fmt.add_adoption_barriers_voice_section(
        pd.DataFrame(
            [
                {
                    "customer_name": "Acme Corp",
                    "SUBJECT_C": "**Login blocker**",
                    "SEVERITY_C": "High",
                }
            ]
        )
    )
    text = "\n".join(p.text for p in fmt.doc.paragraphs)
    assert "**" not in text
    assert_in_source(text, "Login blocker", label='text')


def test_round94_frontend_source_shape_guards() -> None:
    leader = _read("templates/leader_report_form.html")
    analyze = _read("templates/analyze.html")
    progress = _read("templates/progress.html")
    jobs = _read("static/js/report_jobs_dashboard.js")

    assert "if (!form || !submitBtn)" in leader
    assert "document.getElementById('csrfToken')" in leader
    assert "csrfInput ? csrfInput.value : ''" in leader
    assert "if (fileInput)" in leader
    assert "document.querySelector('input[name=\"report_type\"]:checked')?.value || 'leader'" in analyze
    assert "var allowedCsoneStatuses" in progress
    assert "report-jobs-open failed" in jobs
    assert "report-jobs-cancel failed" in jobs


def test_round94_word_build_label_text_carries_canonical_audit_fields() -> None:
    import _r68_build_label

    label = _r68_build_label.get_build_label_text()
    assert label.startswith("AdoptIQ v")
    assert "App_Version:" in label
    assert "App_Build:" in label
    assert "Process_Started_At_UTC:" in label
    assert "Report_Generated_At_UTC:" in label


def test_round94_comprehensive_xlsx_report_info_carries_partial_warnings(tmp_path) -> None:
    from openpyxl import load_workbook

    from adoptiq_backend import write_excel_workbook

    base = tmp_path / "round94_partial_warning"
    warning = {
        "dataset": "adoption_barriers",
        "kind": "tech_filter_scope_excluded",
        "error": "AB tech filter 'All Contact Center' kept 13 of 65 adoption barriers.",
    }

    write_excel_workbook(
        str(base),
        {"AB_Detail_All": pd.DataFrame([{"Customer": "Acme", "Risk": 7}])},
        manager="Brian Frazier",
        technology="All Contact Center",
        days=90,
        partial_data_warnings=[warning],
    )

    wb = load_workbook(base.with_suffix(".xlsx"), read_only=True, data_only=True)
    rows = list(wb["Report_Info"].iter_rows(values_only=True))
    assert (
        "Partial_Data_Warning",
        "adoption_barriers (tech_filter_scope_excluded): AB tech filter 'All Contact Center' kept 13 of 65 adoption barriers.",
    ) in rows
