"""Round 52 (Phase 3): fixture-based DOCX/XLSX parity tests for all 4 reports.

These tests generate small in-memory fixtures, write them to ``.docx`` / ``.xlsx``
via the same code paths the production formatters use (or close-mirror shapes
for paths that require live Snowflake), then run the live-harness KPI
extractors over the produced files. The goal is to catch the class of bug
where the Word and Excel outputs drift apart silently — without requiring a
live Snowflake/CSOne environment.

Coverage:
- ``comprehensive`` -- reuses the Round 16 cross-format helpers (canonical
  ``build_summary_rows`` + ``add_executive_summary_table``) to write a real
  workbook and document, then runs the harness extractors.
- ``compact`` -- mirrors the Executive_Dashboard / At-a-Glance shapes.
- ``renewal`` -- mirrors the Customer Health Dashboard + Report_Info shapes.
- ``leader`` -- mirrors the Team Performance Metrics + Team_Summary shapes.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest
from docx import Document

import report_export_styling as styling

try:
    import report_word_styling as word_styling
except Exception:  # pragma: no cover - dep missing only in degraded env
    word_styling = None  # type: ignore[assignment]

from report_iteration_loop import (
    SCENARIO_REQUIRED_KPIS,
    compare_kpi_parity,
    extract_docx_kpis,
    extract_xlsx_kpis,
    validate_docx_structure,
    validate_xlsx_structure,
)


def _build_comprehensive_fixture() -> dict[str, Any]:
    ab_df = pd.DataFrame(
        {
            "ACCOUNT_NAME": ["Acme", "Beta", "Gamma", "Acme", "Delta"],
            "STATUS": ["Open", "Open", "Closed", "Open", "Open"],
            "PRIORITY": ["Critical", "High", "Medium", "Critical", "Low"],
        }
    )
    csone_df = pd.DataFrame(
        {
            "Customer Name": ["Acme", "Beta", "Beta", "Gamma", "Delta", "Epsilon"],
            "Severity": ["1", "2", "1", "3", "1", "2"],
            "Status": ["Open", "Open", "Closed", "Open", "Open", "Open"],
            "TAC Case Number": ["T1", "T2", "T3", "T4", "T5", "T6"],
            "Sub Case Type": [
                "Break/Fix",
                "Configuration",
                "Provisioning",
                "Configuration",
                "Break/Fix",
                "Break/Fix",
            ],
            "Escalated": ["Y", "N", "N", "N", "Y", "N"],
        }
    )
    cs_pulse = pd.DataFrame(
        {
            "ACCOUNT_NAME": ["Acme", "Beta", "Gamma", "Zeta"],
            "HEALTH_SCORE": [3.0, 5.0, 8.0, 9.0],
        }
    )
    return {
        "sheets": {
            "AB_Detail_All": ab_df,
            "CSOne_Detail_All": csone_df,
        },
        "csconsole_data": {"customer_pulse": cs_pulse},
    }


def test_round52_comprehensive_fixture_parity_via_canonical_helpers(tmp_path: Path):
    """Round 52 / Phase 3 -- comprehensive scenario fixture parity.

    Uses the same helpers as Round 16's Phase 1.2 to write the Excel
    Summary sheet and the Word executive-summary table from one shared
    fixture, then runs the harness's KPI extractors and asserts the
    common KPIs agree value-for-value.
    """

    if word_styling is None:
        pytest.skip("report_word_styling not importable")

    fixture = _build_comprehensive_fixture()

    xlsx_path = tmp_path / "comprehensive.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="xlsxwriter") as writer:
        pd.DataFrame({"placeholder": [None]}).to_excel(
            writer, sheet_name="Placeholder", index=False
        )
        styling.write_summary_sheet(
            writer,
            fixture["sheets"],
            fixture["csconsole_data"],
            manager="Brian Frazier",
            tech="All Contact Center",
            days=90,
            generated_at_utc_iso_z="2026-04-29T00:00:00Z",
        )

    docx_path = tmp_path / "comprehensive.docx"
    document = Document()
    rows = word_styling.build_executive_summary_rows(
        fixture["sheets"],
        fixture["csconsole_data"],
        manager="Brian Frazier",
        tech="All Contact Center",
        days=90,
        generated_at_utc_iso_z="2026-04-29T00:00:00Z",
    )
    word_styling.add_executive_summary_table(document, rows)
    document.add_paragraph("Manager: Brian Frazier")
    document.add_paragraph("Technology: All Contact Center")
    document.add_paragraph("Analysis Period: 90 days")
    document.save(str(docx_path))

    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(
        docx_kpis,
        xlsx_kpis,
        strict=True,
        required_keys=SCENARIO_REQUIRED_KPIS["comprehensive"],
    )
    assert parity.passed, parity.details
    assert parity.details["mismatches"] == {}
    assert "total_customers" in parity.details["common_kpis"]
    assert "manager" in (set(docx_kpis["values"]) | set(xlsx_kpis["values"]))


def _write_compact_xlsx_fixture(path: Path) -> None:
    """Mirror app_simple compact Excel writer: Executive_Dashboard + Risk_Summary."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Executive_Dashboard"
    sheet.append(["Metric", "Value"])
    rows = (
        ("Total Customers Analyzed", 52),
        ("High-Risk Customers", 8),
        ("Critical Adoption Barriers", 3),
        ("Escalated Support Cases", 176),
        ("Overall Risk Score", 14.8),
        ("Analysis Period (Days)", 90),
        ("Technology Focus", "All Contact Center"),
        ("Manager", "All Managers"),
        ("Report Generated", "2026-04-29T00:00:00Z"),
    )
    for row in rows:
        sheet.append(list(row))
    risk = workbook.create_sheet("Risk_Summary")
    risk.append(
        [
            "Customer",
            "Risk_Score",
            "Risk_Level",
            "Risk_Band",
            "Adoption_Barriers",
            "Support_Cases",
        ]
    )
    risk.append(["AcmeCorp", 14.8, "High", "High", 3, 176])
    workbook.save(path)
    workbook.close()


def _write_compact_docx_fixture(path: Path) -> None:
    """Mirror executive_intelligence_formatter At-a-Glance Dashboard table."""
    doc = Document()
    table = doc.add_table(rows=2, cols=7)
    headers = (
        "Total Customers",
        "Support Cases",
        "Critical (P1)",
        "High (P2)",
        "BEMS Escalations",
        "Software Defects",
        "Security Vulnerabilities",
    )
    values = ("52", "176", "5", "12", "3", "8", "2")
    for col_idx, header in enumerate(headers):
        table.rows[0].cells[col_idx].text = header
    for col_idx, value in enumerate(values):
        table.rows[1].cells[col_idx].text = value
    doc.add_paragraph("Risk Summary: Overall Risk Score: 14.8")
    doc.add_paragraph("Manager: All Managers")
    doc.add_paragraph("Technology: All Contact Center")
    doc.add_paragraph("Analysis Period: 90 days")
    doc.save(path)


def test_round52_compact_fixture_parity_passes_strict_required_keys(tmp_path: Path):
    docx_path = tmp_path / "compact.docx"
    xlsx_path = tmp_path / "compact.xlsx"
    _write_compact_docx_fixture(docx_path)
    _write_compact_xlsx_fixture(xlsx_path)

    structural_docx = validate_docx_structure(docx_path, min_chars=10)
    structural_xlsx = validate_xlsx_structure(
        xlsx_path,
        expected_sheets=("executive_dashboard", "risk_summary"),
    )
    assert structural_docx.passed, structural_docx.details
    assert structural_xlsx.passed, structural_xlsx.details

    parity = compare_kpi_parity(
        extract_docx_kpis(docx_path),
        extract_xlsx_kpis(xlsx_path),
        strict=True,
        required_keys=SCENARIO_REQUIRED_KPIS["compact"],
    )
    assert parity.passed, parity.details
    assert "total_customers" in parity.details["common_kpis"]
    assert parity.details["mismatches"] == {}


def _write_renewal_xlsx_fixture(path: Path) -> None:
    """Mirror renewal Excel writer: Renewal_Summary + Key_Metrics + Report_Info."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Renewal_Summary"
    sheet.append(
        [
            "Customer",
            "Overall_Risk_Score",
            "Risk_Level",
            "Analysis_Date",
            "Next_Review_Date",
        ]
    )
    customers = [
        ("Cust A", 15, "High", "2026-04-29", "2026-05-29"),
        ("Cust B", 12, "Medium", "2026-04-29", "2026-05-29"),
        ("Cust C", 18, "High", "2026-04-29", "2026-05-29"),
    ]
    for row in customers:
        sheet.append(list(row))

    key_metrics = workbook.create_sheet("Key_Metrics")
    key_metrics.append(
        ["Risk_Score", "Risk_Category", "Analysis_Period", "Key_Findings"]
    )
    key_metrics.append([15, "High", 90, "Sample"])

    info = workbook.create_sheet("Report_Info")
    info.append(["Field", "Value"])
    info.append(["Report_Type", "renewal_portfolio"])
    info.append(["Customer_Name", "Portfolio Wide"])
    info.append(["Technology", "All Contact Center"])
    info.append(["Manager", "All Managers"])
    info.append(["Days", 90])
    info.append(["Generated_At_UTC", "2026-04-29T00:00:00Z"])
    workbook.save(path)
    workbook.close()


def _write_renewal_docx_fixture(path: Path) -> None:
    """Mirror renewal Customer Health Dashboard + executive summary."""
    doc = Document()
    doc.add_paragraph("Manager: All Managers")
    doc.add_paragraph("Technology: All Contact Center")
    doc.add_paragraph("Analysis Period: 90 days")
    table = doc.add_table(rows=8, cols=2)
    rows = (
        ("Support Cases (Last 90 Days)", "176"),
        ("Active Adoption Barriers", "3"),
        ("BEMS Escalations", "5"),
        ("Service Incidents (status.webex.com)", "1"),
        ("High-Impact Incidents", "0"),
        ("Correlated Service Incidents", "0"),
        ("Overall Risk Score", "15"),
        ("Risk Category", "High"),
    )
    for idx, (label, value) in enumerate(rows):
        table.rows[idx].cells[0].text = label
        table.rows[idx].cells[1].text = value
    doc.save(path)


def test_round52_renewal_fixture_parity_passes_strict_required_keys(tmp_path: Path):
    docx_path = tmp_path / "renewal.docx"
    xlsx_path = tmp_path / "renewal.xlsx"
    _write_renewal_docx_fixture(docx_path)
    _write_renewal_xlsx_fixture(xlsx_path)

    structural_xlsx = validate_xlsx_structure(
        xlsx_path,
        expected_sheets=("report_info", "renewal_summary", "key_metrics"),
    )
    assert structural_xlsx.passed, structural_xlsx.details

    parity = compare_kpi_parity(
        extract_docx_kpis(docx_path),
        extract_xlsx_kpis(xlsx_path),
        strict=True,
        required_keys=SCENARIO_REQUIRED_KPIS["renewal"],
    )
    assert parity.passed, parity.details
    # Renewal exercises the cross-format alias chain:
    # docx "Support Cases (Last 90 Days)" → support_cases
    # xlsx Report_Info "Days" → window_days
    # xlsx Report_Info "Manager"/"Technology" → manager/technology
    union_keys = set(extract_docx_kpis(docx_path)["values"]) | set(
        extract_xlsx_kpis(xlsx_path)["values"]
    )
    assert "window_days" in union_keys
    assert "technology" in union_keys


def _write_leader_xlsx_fixture(path: Path) -> None:
    """Mirror leader Excel writer: Team_Summary + Report_Info."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Team_Summary"
    sheet.append(
        [
            "Team_Member",
            "Num_Customers",
            "Num_Subscriptions",
            "Num_Action_Plans",
            "Num_Adoption_Barriers",
            "Num_Customer_Pulse",
            "Num_Success_Priorities",
            "Num_TAC_Cases",
            "Customers",
        ]
    )
    members = [
        ("Member A", 5, 8, 4, 3, 2, 1, 6, "Acme,Beta"),
        ("Member B", 4, 6, 3, 2, 1, 0, 4, "Gamma"),
        ("Member C", 6, 9, 5, 3, 2, 1, 8, "Delta,Epsilon"),
    ]
    for row in members:
        sheet.append(list(row))

    info = workbook.create_sheet("Report_Info")
    info.append(["Field", "Value", "Detail", "Generated_At"])
    info.append(["Status", "completed", "", "2026-04-29T00:00:00Z"])
    info.append(["Manager", "Brian Frazier", "", ""])
    info.append(["Days", 90, "", ""])
    info.append(["Sheets_Written", 5, "", ""])
    workbook.save(path)
    workbook.close()


def _write_leader_docx_fixture(path: Path) -> None:
    """Mirror leader Team Performance Metrics two-column table."""
    doc = Document()
    doc.add_paragraph("Manager: Brian Frazier")
    doc.add_paragraph("Analysis Period: 90 days")
    table = doc.add_table(rows=6, cols=2)
    rows = (
        ("Team Members", "3"),
        ("Action Plans", "12"),
        ("Adoption Barriers", "8"),
        ("Customer Pulse Records", "5"),
        ("TAC Cases", "18"),
        ("BEMS Escalations", "3"),
    )
    for idx, (label, value) in enumerate(rows):
        table.rows[idx].cells[0].text = label
        table.rows[idx].cells[1].text = value
    doc.save(path)


def test_round52_leader_fixture_parity_passes_strict_required_keys(tmp_path: Path):
    docx_path = tmp_path / "leader.docx"
    xlsx_path = tmp_path / "leader.xlsx"
    _write_leader_docx_fixture(docx_path)
    _write_leader_xlsx_fixture(xlsx_path)

    structural_xlsx = validate_xlsx_structure(
        xlsx_path,
        expected_sheets=("report_info", "team_summary"),
    )
    assert structural_xlsx.passed, structural_xlsx.details

    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(
        docx_kpis,
        xlsx_kpis,
        strict=True,
        required_keys=SCENARIO_REQUIRED_KPIS["leader"],
    )
    assert parity.passed, parity.details
    # Cross-format coverage: DOCX "Team Members" / XLSX team_summary row count
    # both feed the total_customers KPI; both must be present.
    docx_values = docx_kpis["values"]
    xlsx_values = xlsx_kpis["values"]
    assert docx_values.get("total_customers") == xlsx_values.get("total_customers")
    assert docx_values.get("action_plans") == xlsx_values.get("action_plans")
    assert docx_values.get("adoption_barriers") == xlsx_values.get("adoption_barriers")


def test_round52_fixture_parity_detects_intentional_mismatch(tmp_path: Path):
    """Negative test: when DOCX and XLSX disagree on a common KPI in strict mode,
    the parity gate must fail and identify the offender."""

    docx_path = tmp_path / "drift.docx"
    xlsx_path = tmp_path / "drift.xlsx"
    _write_compact_docx_fixture(docx_path)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Executive_Dashboard"
    sheet.append(["Metric", "Value"])
    sheet.append(["Total Customers Analyzed", 9999])  # mismatch vs DOCX 52
    sheet.append(["Escalated Support Cases", 176])
    sheet.append(["Overall Risk Score", 14.8])
    sheet.append(["Analysis Period (Days)", 90])
    sheet.append(["Technology Focus", "All Contact Center"])
    sheet.append(["Manager", "All Managers"])
    workbook.save(xlsx_path)
    workbook.close()

    parity = compare_kpi_parity(
        extract_docx_kpis(docx_path),
        extract_xlsx_kpis(xlsx_path),
        strict=True,
        required_keys=SCENARIO_REQUIRED_KPIS["compact"],
    )
    assert parity.passed is False
    assert "total_customers" in parity.details["mismatches"]
    assert parity.details["mismatches"]["total_customers"]["docx"] == "52"
    assert parity.details["mismatches"]["total_customers"]["xlsx"] == "9999"
