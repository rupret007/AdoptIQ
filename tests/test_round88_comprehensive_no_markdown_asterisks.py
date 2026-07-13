"""Round 88 / F1 — strip markdown asterisks from Comprehensive narrative.

The Build 63 acceptance audit found 25 paragraphs in the Comprehensive
DOCX with literal ``*xxx*`` italic markers leaked verbatim from the LLM
prompt template into the rendered Word document.

Two-pronged fix lives in ``adoptiq_backend.py``:

1. **Sanitizer** — ``append_to_word_report.clean_and_format_text`` is now
   italic-aware: it strips single-star italic markers (``*text*``) using
   ``_r88_italic_re`` and converts the inner text to a Word italic run
   (no literal asterisks reach the docx).  Bold (``**text**``) handling
   is unchanged so the SSoT is preserved.

2. **Prompt cleanup** — the comprehensive prompt templates
   (``PROMPT_COMPREHENSIVE_TEMPLATE`` + ``PROMPT_CUSTOMER_TEMPLATE`` +
   ``PROMPT_COMPACT_EXECUTIVE_TEMPLATE``) had instructional/meta lines
   wrapped in ``*xxx*`` that the LLM mirrored back as output.  Those
   wrappers were removed so the model has nothing to copy.

This module pins both contracts via source-shape AND a synthetic
``append_to_word_report`` round-trip.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest


_REPO_ROOT = Path(__file__).resolve().parent.parent
_BACKEND_PATH = _REPO_ROOT / "adoptiq_backend.py"


def test_r88_f1_italic_regex_helper_present_in_clean_and_format_text() -> None:
    """The ``_r88_italic_re`` helper MUST be declared in the same scope
    as ``clean_and_format_text`` inside ``append_to_word_report`` so the
    sanitizer is applied to every line that flows through the writer.
    """

    text = _BACKEND_PATH.read_text(encoding="utf-8")
    assert "_r88_italic_re = re.compile(" in text, (
        "Round 88 / F1: _r88_italic_re must be declared in adoptiq_backend "
        "so clean_and_format_text can strip leftover italic chrome."
    )
    # Lookbehind/lookahead for word/star MUST be present so the regex is
    # safe against bold pairs and word-internal stars.
    assert r"(?<![*\w])\*([^*\n]+?)\*(?![*\w])" in text, (
        "Round 88 / F1: italic regex must use the negative lookbehind/"
        "lookahead pattern to avoid double-stripping bold or mid-word stars."
    )


def test_r88_f1_clean_and_format_text_handles_italic_in_no_bold_path() -> None:
    """When a line has zero ``**`` markers but contains ``*xxx*`` italic
    chrome (the exact pattern Brian's Build 63 reports leaked), the
    sanitizer MUST strip the asterisks and emit an italic run.
    """

    text = _BACKEND_PATH.read_text(encoding="utf-8")
    # Find clean_and_format_text definition body
    start = text.find("def clean_and_format_text(text, paragraph):")
    assert start > 0, "clean_and_format_text must exist"
    # Look for the italic-handling branch in the count == 0 path
    body_window = text[start : start + 3500]
    assert "italic_segments = _r88_italic_re.split" in body_window, (
        "Round 88 / F1: clean_and_format_text must split on _r88_italic_re "
        "and emit italic runs (not literal ``*`` asterisks) when no bold "
        "markers are present."
    )
    assert "run.italic = True" in body_window, (
        "Round 88 / F1: clean_and_format_text must set italic=True on the "
        "italic-marked run."
    )


def test_r88_f1_prompt_template_has_no_italic_wrapped_instruction_lines() -> None:
    """The comprehensive prompt template MUST NOT contain instructional
    lines wrapped in ``*xxx*`` markdown italic markers.  Pre-R88 it had
    ~19 such lines that the LLM mirrored back into the rendered docx.
    """

    text = _BACKEND_PATH.read_text(encoding="utf-8")
    # Match lines that START with ``*`` followed immediately by a non-
    # whitespace, non-asterisk character — exactly the pattern of the
    # instruction-wrap chrome (``*State the grade...*``).  Bullet
    # markers (``* ``) are EXCLUDED by the regex (the immediate next
    # char must not be whitespace).
    pattern = re.compile(r"^\*[^*\s]", re.MULTILINE)
    matches = pattern.findall(text)
    assert not matches, (
        f"Round 88 / F1: adoptiq_backend.py still contains "
        f"{len(matches)} italic-wrapped instructional line(s) starting with "
        "``*x`` — these will be mirrored by the LLM into the Comprehensive "
        "docx.  Remove the outer ``*...*`` wrappers."
    )


def test_r88_f1_round_trip_strips_italic_chrome_from_docx() -> None:
    """End-to-end: feed ``append_to_word_report`` a synthetic narrative
    with the Build 63 leakage pattern and assert the resulting docx
    paragraphs contain ZERO literal ``*`` characters.
    """

    pytest.importorskip("docx", reason="python-docx required for round-trip test")

    from docx import Document  # noqa: PLC0415

    from adoptiq_backend import append_to_word_report  # noqa: PLC0415

    synthetic_narrative = (
        "## Critical Trouble Spots\n"
        "\n"
        "*BEMS escalations indicate problems requiring back-end engineering:*\n"
        "\n"
        "*P1/P2 cases requiring immediate attention:*\n"
        "\n"
        "*Known software defects from help.webex.com that are impacting your customers:*\n"
        "\n"
        "*The health score is currently C due to a high volume of complex technical escalations.*\n"
        "\n"
        "Total Customers: 39\n"
        "Total Adoption Barriers: 161\n"
    )

    doc = Document()
    append_to_word_report(doc, synthetic_narrative)

    # Inspect every paragraph's text + run-level text; neither should
    # carry a literal ``*`` from our synthetic input.
    leaks = []
    for para in doc.paragraphs:
        if "*" in para.text:
            leaks.append(("paragraph", para.text))
        for run in para.runs:
            if "*" in run.text:
                leaks.append(("run", run.text))

    assert not leaks, (
        f"Round 88 / F1: append_to_word_report leaked {len(leaks)} literal "
        "``*`` markers into the rendered docx.  The italic sanitizer "
        "extension is not catching the Build 63 leakage pattern.  Leaks: "
        f"{leaks!r}"
    )


def test_r88_f1_round_trip_preserves_italic_run_for_stripped_chrome() -> None:
    """When the sanitizer strips ``*xxx*`` chrome, the inner text MUST
    still appear in the docx (just as italic, not as literal asterisks).
    Pin the positive contract: data is preserved, only the markers go.
    """

    pytest.importorskip("docx", reason="python-docx required for round-trip test")

    from docx import Document  # noqa: PLC0415

    from adoptiq_backend import append_to_word_report  # noqa: PLC0415

    doc = Document()
    append_to_word_report(
        doc,
        "*P1/P2 cases requiring immediate attention:*\n",
    )

    # Find the rendered paragraph and assert content preserved + italic
    found_italic_with_text = False
    for para in doc.paragraphs:
        for run in para.runs:
            if (
                "P1/P2 cases requiring immediate attention" in run.text
                and run.italic
            ):
                found_italic_with_text = True
                break

    assert found_italic_with_text, (
        "Round 88 / F1: when the sanitizer strips ``*xxx*`` chrome, the "
        "inner text must be preserved as an italic run.  This regression "
        "test ensures content fidelity is maintained."
    )


def test_r88_f1_bold_markers_still_processed_correctly() -> None:
    """The italic-handling extension MUST NOT regress the existing
    ``**bold**`` handling (the SSoT lineage from R44/R48/R49).  Pin the
    bold contract by feeding a synthetic narrative with mixed bold +
    italic and asserting both render correctly.
    """

    pytest.importorskip("docx", reason="python-docx required for round-trip test")

    from docx import Document  # noqa: PLC0415

    from adoptiq_backend import append_to_word_report  # noqa: PLC0415

    doc = Document()
    append_to_word_report(
        doc,
        "**Total Customers:** 39 with *escalation patterns* indicating risk.\n",
    )

    found_bold = False
    found_italic = False
    leaked = False
    for para in doc.paragraphs:
        if "*" in para.text:
            leaked = True
        for run in para.runs:
            if "Total Customers:" in run.text and run.bold:
                found_bold = True
            if "escalation patterns" in run.text and run.italic:
                found_italic = True

    assert not leaked, (
        "Round 88 / F1: mixed bold + italic input must not leak literal "
        "``*`` markers into the docx."
    )
    assert found_bold, (
        "Round 88 / F1: ``**Total Customers:**`` must render as a bold "
        "run.  The italic extension regressed bold handling."
    )
    assert found_italic, (
        "Round 88 / F1: ``*escalation patterns*`` must render as an "
        "italic run after sanitization."
    )
