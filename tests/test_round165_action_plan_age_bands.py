"""Round 165 unresolved Action Plan age-band and evidence contracts."""

from __future__ import annotations

import base64
from collections import OrderedDict
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import canonical_metrics as cm
import decision_report_delivery as delivery


AS_OF = pd.Timestamp("2026-08-03T12:00:00Z")
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _boundary_plans() -> pd.DataFrame:
    rows = []
    for index, age in enumerate((0, 14, 15, 30, 31, 60, 61, 90, 91), 1):
        rows.append(
            {
                "ID": f"AP-{index:03d}",
                "BU_NAME": "Acme Corporation",
                "SUBJECT_C": f"Unresolved plan aged {age} days",
                "STATUS_C": "Open",
                "CREATED_DATE_C": (AS_OF - pd.Timedelta(age, unit="D")).isoformat(),
            }
        )
    rows.extend(
        [
            {
                "ID": "AP-UNKNOWN-AGE",
                "BU_NAME": "Acme Corporation",
                "SUBJECT_C": "Unresolved plan without created date",
                "STATUS_C": "On Hold",
            },
            {
                "ID": "AP-COMPLETED",
                "BU_NAME": "Acme Corporation",
                "SUBJECT_C": "Completed plan",
                "STATUS_C": "Completed - Successful",
                "CREATED_DATE_C": (
                    AS_OF - pd.Timedelta(365, unit="D")
                ).isoformat(),
            },
        ]
    )
    return pd.DataFrame(rows)


def _facts() -> dict:
    return delivery.build_report_facts(
        {
            "Alex Rivera": {
                "subscriptions": pd.DataFrame(
                    [
                        {
                            "SUBSCRIPTION_ID": "SUB-001",
                            "ACCOUNT_ID_C": "ACC-001",
                            "BU_NAME": "Acme Corporation",
                        }
                    ]
                ),
                "action_plans": _boundary_plans(),
                "adoption_barriers": pd.DataFrame(),
                "customer_pulse": pd.DataFrame(),
                "tac_cases": pd.DataFrame(),
                "success_priorities": pd.DataFrame(),
            }
        },
        report_type="Leader",
        scope_type="team",
        scope_value="Dana Manager team",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
        data_as_of_utc=AS_OF,
        data_as_of_state="available",
        external_incidents=[],
        external_bugs=[],
    )


def test_age_bands_partition_unresolved_plans_at_exact_boundaries() -> None:
    lifecycle = cm.build_action_plan_lifecycle(_boundary_plans(), as_of=AS_OF)

    assert lifecycle["total"] == 11
    assert lifecycle["completed"] == 1
    assert lifecycle["unresolved_total"] == 10
    assert lifecycle["age_band_counts"] == {
        "0–14 days": 2,
        "15–30 days": 2,
        "31–60 days": 2,
        "61–90 days": 2,
        ">90 days": 1,
        "Unknown": 1,
    }
    assert sum(lifecycle["age_band_counts"].values()) == lifecycle["unresolved_total"]
    completed = lifecycle["records"].set_index("AdoptIQ_Record_ID").loc[
        "AP-COMPLETED"
    ]
    assert completed["AdoptIQ_Age_Days"] == 365
    assert completed["AdoptIQ_Age_Band"] == cm.ACTION_PLAN_COMPLETED_AGE_LABEL

    series = cm.action_plan_chart_series(lifecycle)
    totals = series.groupby("Series")["Value"].sum().to_dict()
    assert totals == {
        cm.ACTION_PLAN_STATUS_SERIES: 11,
        cm.ACTION_PLAN_AGE_SERIES: 10,
    }


def test_age_series_fails_closed_for_incomplete_action_plan_source() -> None:
    plans = _boundary_plans()
    plans.attrs["partial"] = True
    plans.attrs["source_mode_detail"] = "later source page failed"

    lifecycle = cm.build_action_plan_lifecycle(plans, as_of=AS_OF)
    series = cm.action_plan_chart_series(lifecycle)

    assert lifecycle["source_state"] == "partial"
    assert set(series["Source_State"]) == {"partial"}
    assert series["Value"].isna().all()
    assert set(series["Series"]) == {
        cm.ACTION_PLAN_STATUS_SERIES,
        cm.ACTION_PLAN_AGE_SERIES,
    }


def test_age_chart_lineage_evidence_and_tamper_contracts(tmp_path: Path) -> None:
    facts = _facts()
    sheets = delivery.build_source_data_sheets(facts)
    contract = delivery.validate_cross_artifact_contract(facts, sheets)
    assert contract["ok"], contract["errors"]

    age_chart = facts["chart_data"].loc[
        facts["chart_data"]["Series"] == cm.ACTION_PLAN_AGE_SERIES
    ]
    assert set(age_chart["Metric_Key"]) == {
        "chart.action_plan_age.0_14_days",
        "chart.action_plan_age.15_30_days",
        "chart.action_plan_age.31_60_days",
        "chart.action_plan_age.61_90_days",
        "chart.action_plan_age.90_days",
        "chart.action_plan_age.unknown",
    }
    assert int(age_chart["Value"].sum()) == facts["action_plan_lifecycle"][
        "unresolved_total"
    ]
    evidence = sheets["Evidence_Links"].loc[
        sheets["Evidence_Links"]["Evidence_Key"]
        == "chart.action_plan_age.0_14_days"
    ]
    assert len(evidence) == 2
    for row_number in evidence["Source_Row_Number"]:
        source_row = sheets["Action_Plans"].iloc[int(row_number) - 2]
        assert source_row["AdoptIQ_Age_Band"] == "0–14 days"

    tampered = OrderedDict((name, frame.copy()) for name, frame in sheets.items())
    action_rows = tampered["Action_Plans"]
    target_index = action_rows.index[action_rows["AdoptIQ_Age_Band"].eq("0–14 days")][0]
    action_rows.loc[target_index, "AdoptIQ_Age_Band"] = ">90 days"
    tampered_contract = delivery.validate_cross_artifact_contract(facts, tampered)
    assert not tampered_contract["ok"]
    assert any(
        "unresolved age bands" in error for error in tampered_contract["errors"]
    )

    workbook_path = tmp_path / "age-band-contract.xlsx"
    delivery.write_source_data_workbook(workbook_path, sheets)
    workbook = load_workbook(workbook_path)
    worksheet = workbook["Action_Plans"]
    headers = {cell.value: cell.column for cell in worksheet[1]}
    for row_number in range(2, worksheet.max_row + 1):
        if worksheet.cell(row_number, headers["AdoptIQ_Age_Band"]).value == "0–14 days":
            worksheet.cell(row_number, headers["AdoptIQ_Age_Band"]).value = ">90 days"
            break
    workbook.save(workbook_path)
    workbook.close()
    written_contract = delivery.validate_written_source_workbook(workbook_path, facts)
    assert not written_contract["ok"]
    assert any(
        "unresolved age bands" in error for error in written_contract["errors"]
    )


def test_action_plan_chart_renders_status_and_age_as_two_panels(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    matplotlib = pytest.importorskip("matplotlib")
    matplotlib.use("Agg")
    pyplot = pytest.importorskip("matplotlib.pyplot")
    original_subplots = pyplot.subplots
    captured: dict[str, object] = {}

    def tracking_subplots(*args, **kwargs):
        figure, axes = original_subplots(*args, **kwargs)
        captured["axes"] = axes
        return figure, axes

    monkeypatch.setattr(pyplot, "subplots", tracking_subplots)
    facts = _facts()
    rows = facts["chart_data"].loc[
        facts["chart_data"]["Chart_ID"] == "action_plan_status_aging"
    ]
    target = tmp_path / "action-plan-status-aging.png"

    assert delivery._render_chart_image(  # noqa: SLF001 - focused renderer contract
        "action_plan_status_aging", rows, target
    )
    axes = list(captured["axes"])
    assert len(axes) == 2
    assert [axis.get_title() for axis in axes] == [
        "Lifecycle status · all plans",
        "Age · unresolved only",
    ]
    assert target.stat().st_size > 0
    assert set(facts["chart_data"]["Chart_ID"]) == {
        "activity_mix",
        "action_plan_status_aging",
        "risk_distribution",
        "activity_trend",
    }


def test_word_age_table_is_canonical_and_tamper_detected(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fake_renderer(_chart_id: str, _rows: pd.DataFrame, target: Path) -> bool:
        target.write_bytes(_TINY_PNG)
        return True

    monkeypatch.setattr(delivery, "_render_chart_image", fake_renderer)
    facts = _facts()
    document = delivery.build_concise_word_document(facts)
    semantic = delivery.validate_word_semantics(facts, document)
    assert semantic["ok"], semantic["errors"]
    age_tables = [
        table
        for table in document.tables
        if tuple(cell.text for cell in table.rows[0].cells)
        == ("Unresolved-plan age", "Distinct plans")
    ]
    assert len(age_tables) == 1
    assert len(document.inline_shapes) == 4

    age_tables[0].rows[1].cells[1].text = "999"
    tampered = delivery.validate_word_semantics(facts, document)
    assert not tampered["ok"]
    assert any(
        "Unresolved-plan age | Distinct plans" in error
        for error in tampered["errors"]
    )
