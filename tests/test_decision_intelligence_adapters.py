from __future__ import annotations

import copy

from docx import Document
import pandas as pd
import pytest

from decision_intelligence import (
    ANALYSIS_SCHEMA_VERSION,
    AnalysisBundle,
    AnalysisContext,
    AnalysisDiagnostics,
    AnalysisRequest,
    CustomerAnalysis,
    DataQualitySummary,
    DecisionBrief,
    EvidenceReference,
    Finding,
    MetricValue,
    PortfolioAnalysis,
    RecommendedAction,
)
from decision_intelligence_adapters import (
    ask_ai_safe_projection,
    customer_decision_brief_frame,
    decision_brief_excel_frames,
    portfolio_decision_brief_frame,
    project_canonical_portfolio_metrics,
    project_legacy_risk_profiles,
    render_decision_brief_word,
    stamp_projection,
    validate_cross_output_reconciliation,
)


AS_OF = "2026-07-13T12:00:00Z"


def _quality(confidence: str = "HIGH", coverage: float = 1.0) -> DataQualitySummary:
    return DataQualitySummary(
        coverage_ratio=coverage,
        confidence=confidence,
        source_states={"adoption_barriers": "available", "support_cases": "available"},
        source_freshness={"adoption_barriers": "current", "support_cases": "current"},
    )


def _action(
    action_id: str,
    customer_id: str,
    finding_id: str,
    evidence_id: str,
    text: str,
) -> RecommendedAction:
    return RecommendedAction(
        action_id=action_id,
        scope_kind="customer",
        scope_id=customer_id,
        action_type="retention",
        specific_action=text,
        rationale="The cited signal requires a bounded follow-up.",
        triggering_finding_ids=(finding_id,),
        evidence_ids=(evidence_id,),
        proposed_owner="Customer Success Manager",
        owner_confidence="HIGH",
        urgency="Within 7 days",
        rank=1,
        priority_score=88.0,
        ranking_factors={"risk": 72.0},
        dependencies=(),
        expected_outcome="Risk is explicitly addressed.",
        measurable_success_signal="Owner records the resolution plan.",
        timing_window="7 days",
        effort="Medium",
        confidence="HIGH",
    )


@pytest.fixture
def bundle() -> AnalysisBundle:
    request = AnalysisRequest(
        portfolio_scope="All Managers",
        technology_scope=("Webex Contact Center",),
        time_range_start="2026-04-14T12:00:00Z",
        time_range_end=AS_OF,
        as_of_time=AS_OF,
    )
    acme_id = "customer:acme"
    beta_id = "customer:beta"
    acme_evidence = EvidenceReference(
        evidence_id="evidence:acme-ab",
        source_type="adoption_barriers",
        source_record="ID",
        stable_source_identifier="AB-1",
        customer_identity=acme_id,
        relevant_field="SEVERITY_C",
        observed_value="High",
        observation_timestamp="2026-07-10T12:00:00Z",
        ingestion_timestamp=AS_OF,
        source_freshness="current",
        safe_excerpt="ID=AB-1; SEVERITY_C=High",
    )
    beta_evidence = EvidenceReference(
        evidence_id="evidence:beta-state",
        source_type="support_cases",
        source_record="source-state",
        stable_source_identifier="support_cases:source-state",
        customer_identity=beta_id,
        relevant_field="availability",
        observed_value="missing",
        ingestion_timestamp=AS_OF,
        source_freshness="unknown",
        safe_excerpt="Support evidence is unavailable.",
    )

    acme_finding = Finding(
        finding_id="finding:acme-risk",
        scope_kind="customer",
        scope_id=acme_id,
        kind="risk_signal",
        category="adoption",
        title="A high-severity adoption barrier is open.",
        structured_value={"barrier_count": 2},
        severity="high",
        importance=88,
        direction="worsening",
        evidence_ids=(acme_evidence.evidence_id,),
        source_freshness="current",
        confidence="HIGH",
        data_coverage=1.0,
    )
    acme_action = _action(
        "action:acme-plan",
        acme_id,
        acme_finding.finding_id,
        acme_evidence.evidence_id,
        "Create a dated adoption-barrier resolution plan.",
    )
    acme_brief = DecisionBrief(
        scope_kind="customer",
        scope_id=acme_id,
        what_changed=("A high-severity barrier is now active.",),
        why_it_matters=("The barrier raises near-term renewal risk.",),
        next_action_ids=(acme_action.action_id,),
        what_remains_uncertain=("Resolution timing is not yet recorded.",),
        evidence_ids=(acme_evidence.evidence_id,),
        confidence="HIGH",
        synthesis="Acme needs a bounded retention intervention.",
    )
    beta_brief = DecisionBrief(
        scope_kind="customer",
        scope_id=beta_id,
        what_changed=(),
        why_it_matters=("Risk cannot be characterized without current support evidence.",),
        next_action_ids=(),
        what_remains_uncertain=("Current support activity is unavailable.",),
        evidence_ids=(beta_evidence.evidence_id,),
        confidence="LOW",
        synthesis="Beta remains unscored because evidence is insufficient.",
    )

    acme = CustomerAnalysis(
        customer_id=acme_id,
        customer_name="Acme Corp",
        aliases=("Acme",),
        subscriptions=("SUB-1",),
        technologies=("Webex Contact Center",),
        responsible_teams=("Central",),
        metrics=(
            MetricValue("action_plans_count", 1, as_of_time=AS_OF),
            MetricValue("adoption_barriers_count", 2, as_of_time=AS_OF),
            MetricValue("customer_pulse_count", 1, as_of_time=AS_OF),
            MetricValue("support_cases_count", 3, as_of_time=AS_OF),
        ),
        risk_profile={
            "risk_score_0_100": 72.0,
            "risk_score_0_10": 7.2,
            "risk_band": "HIGH",
            "risk_assessment_state": "AVAILABLE",
            "components": {
                "support_cases": {"score": 70.0, "details": {"case_count": 3}},
                "engagement": {"score": 40.0, "details": {"total_activity": 7}},
            },
            "risk_factors": ["High-severity adoption barrier"],
            "key_findings": ["Two adoption barriers are in scope"],
            "recommendations": [acme_action.specific_action],
            "evidence_quality": {"coverage_ratio": 1.0},
        },
        adoption_state={"open_barriers": 2},
        pulse_state={"latest": "negative"},
        renewal_context={"days_to_renewal": 45},
        engagement_state={"activity_count": 7},
        findings=(acme_finding,),
        temporal_changes=(),
        recommended_actions=(acme_action,),
        evidence_ids=(acme_evidence.evidence_id,),
        data_quality=_quality(),
        unresolved_conflicts=(),
        executive_synthesis=acme_brief.synthesis,
        decision_brief=acme_brief,
    )
    beta = CustomerAnalysis(
        customer_id=beta_id,
        customer_name="Beta Inc",
        aliases=(),
        subscriptions=("SUB-2",),
        technologies=("Webex Contact Center",),
        responsible_teams=("Central",),
        metrics=(
            MetricValue("action_plans_count", 0, as_of_time=AS_OF),
            MetricValue("adoption_barriers_count", 1, as_of_time=AS_OF),
            MetricValue("customer_pulse_count", 0, as_of_time=AS_OF),
            MetricValue("support_cases_count", 1, as_of_time=AS_OF),
        ),
        risk_profile={
            "risk_score_0_100": None,
            "risk_score_0_10": None,
            "risk_band": "UNKNOWN",
            "risk_assessment_state": "INSUFFICIENT_EVIDENCE",
            "components": {
                "activity_volume": {"score": 15.0, "details": {"total_activity": 1}}
            },
            "risk_factors": [],
            "key_findings": [],
            "recommendations": [],
        },
        adoption_state={"open_barriers": 1},
        pulse_state={},
        renewal_context={},
        engagement_state={"activity_count": 1},
        findings=(),
        temporal_changes=(),
        recommended_actions=(),
        evidence_ids=(beta_evidence.evidence_id,),
        data_quality=_quality("LOW", 0.25),
        unresolved_conflicts=(),
        executive_synthesis=beta_brief.synthesis,
        decision_brief=beta_brief,
    )

    portfolio_action = _action(
        "action:portfolio-plan",
        "portfolio:all",
        acme_finding.finding_id,
        acme_evidence.evidence_id,
        "Review the highest-risk account in the weekly portfolio meeting.",
    )
    portfolio_brief = DecisionBrief(
        scope_kind="portfolio",
        scope_id="portfolio:all",
        what_changed=("One customer is high risk.",),
        why_it_matters=("Half of the portfolio requires intervention or evidence repair.",),
        next_action_ids=(portfolio_action.action_id,),
        what_remains_uncertain=("One customer has an UNKNOWN risk band.",),
        evidence_ids=(acme_evidence.evidence_id, beta_evidence.evidence_id),
        confidence="MEDIUM",
        synthesis="The portfolio has one high-risk and one unscored customer.",
    )
    portfolio = PortfolioAnalysis(
        portfolio_id="portfolio:all",
        portfolio_scope="All Managers",
        customer_count=2,
        customer_ids=(acme_id, beta_id),
        metrics=(
            MetricValue("average_risk_score_0_100", 72.0, unit="score", as_of_time=AS_OF),
            MetricValue("bems_count", 1, as_of_time=AS_OF),
            MetricValue("break_fix_cases", 2, as_of_time=AS_OF),
            MetricValue("p1_cases", 1, as_of_time=AS_OF),
            MetricValue("p2_cases", 1, as_of_time=AS_OF),
            MetricValue("p3_cases", 1, as_of_time=AS_OF),
            MetricValue("p4_cases", 1, as_of_time=AS_OF),
            MetricValue("provisioning_cases", 1, as_of_time=AS_OF),
            MetricValue("total_barriers", 3, as_of_time=AS_OF),
            MetricValue("total_cases", 4, as_of_time=AS_OF),
            MetricValue("total_customers", 2, as_of_time=AS_OF),
        ),
        risk_distribution={
            "CRITICAL": 0,
            "HIGH": 1,
            "MEDIUM": 0,
            "LOW": 0,
            "HEALTHY": 0,
            "UNKNOWN": 1,
        },
        opportunity_distribution={"EXPANSION": 1, "UNKNOWN": 1},
        temporal_changes=(),
        emerging_patterns=("Evidence coverage is uneven.",),
        recurring_barriers=("Adoption readiness",),
        concentration_risks=("One of two customers is high risk.",),
        cross_customer_themes=("Adoption",),
        ranked_customer_ids=(acme_id, beta_id),
        recommended_actions=(portfolio_action,),
        data_quality=_quality("MEDIUM", 0.625),
        evidence_gaps=("Beta support evidence",),
        executive_synthesis=portfolio_brief.synthesis,
        decision_brief=portfolio_brief,
    )
    context = AnalysisContext(
        request_fingerprint=request.request_fingerprint,
        comparison_scope_fingerprint=request.comparison_scope_fingerprint,
        analysis_fingerprint="analysis:fixture-v2",
        generated_time=AS_OF,
        as_of_time=AS_OF,
        source_cutoff_time=AS_OF,
        selected_customers=(acme_id, beta_id),
        selected_subscriptions=("SUB-1", "SUB-2"),
        selected_technologies=("Webex Contact Center",),
        selected_teams=("Central",),
        source_availability={"adoption_barriers": "available", "support_cases": "available"},
        source_freshness={"adoption_barriers": "current", "support_cases": "current"},
        configuration_version="fixture-1",
        analysis_engine_version="decision-intelligence-v2",
        warnings=("Beta support evidence is unavailable.",),
        degraded_mode_indicators=("partial_support_coverage",),
        data_quality_summary=_quality("MEDIUM", 0.625),
    )
    return AnalysisBundle(
        schema_version=ANALYSIS_SCHEMA_VERSION,
        request=request,
        context=context,
        customers=(acme, beta),
        portfolio=portfolio,
        evidence=(acme_evidence, beta_evidence),
        diagnostics=AnalysisDiagnostics(missing_evidence=("Beta support evidence",)),
    )


def test_legacy_profiles_preserve_dual_scales_unknown_and_activity_aliases(bundle):
    profiles = project_legacy_risk_profiles(bundle)

    acme = profiles["Acme Corp"]
    assert acme["score"] == 7.2
    assert acme["risk_score_0_10"] == 7.2
    assert acme["risk_score_0_100"] == 72.0
    assert acme["risk_band"] == "HIGH"
    assert acme["color"] == "Red"
    assert acme["components"]["engagement"] == acme["components"]["activity_volume"]

    beta = profiles["Beta Inc"]
    assert beta["score"] is None
    assert beta["risk_score_0_10"] is None
    assert beta["risk_score_0_100"] is None
    assert beta["risk_band"] == "UNKNOWN"
    assert beta["risk_assessment_state"] == "INSUFFICIENT_EVIDENCE"
    assert beta["components"]["engagement"] == beta["components"]["activity_volume"]


def test_canonical_portfolio_metrics_are_frozen_bundle_projections(bundle):
    metrics = project_canonical_portfolio_metrics(bundle)

    assert metrics["total_customers"] == 2
    assert metrics["total_barriers"] == 3
    assert metrics["total_cases"] == 4
    assert metrics["bems_count"] == 1
    assert metrics["high_risk_customers"] == 1
    assert metrics["unknown_risk_customers"] == 1
    assert metrics["scored_customers"] == 1
    assert metrics["risk_band_counts"] == {
        "CRITICAL": 0,
        "HIGH": 1,
        "MEDIUM": 0,
        "LOW": 0,
        "HEALTHY": 0,
        "UNKNOWN": 1,
    }


def test_decision_brief_frames_are_deterministic_and_excel_safe(bundle):
    customers_a = customer_decision_brief_frame(bundle)
    customers_b = customer_decision_brief_frame(bundle)
    portfolio = portfolio_decision_brief_frame(bundle)

    pd.testing.assert_frame_equal(customers_a, customers_b)
    assert customers_a["Customer_ID"].tolist() == ["customer:acme", "customer:beta"]
    assert customers_a.loc[0, "Risk_Score_0_10"] == 7.2
    assert pd.isna(customers_a.loc[1, "Risk_Score_0_10"])
    assert "\n" not in customers_a.loc[0, "Synthesis"]
    assert portfolio.loc[0, "Total_Customers"] == 2
    assert portfolio.loc[0, "High_Risk_Customers"] == 1
    assert customers_a.attrs["_decision_intelligence"]["payload_digest"].startswith("sha256:")
    assert portfolio.attrs["_decision_intelligence"]["analysis_fingerprint"] == bundle.analysis_fingerprint


def test_word_helper_renders_in_memory_and_returns_reconcilable_content(bundle):
    document = Document()

    report = render_decision_brief_word(document, bundle)

    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Portfolio Decision Brief" in text
    assert "Customer Decision Brief: Acme Corp" in text
    assert "Risk: UNKNOWN (n/a)" in text
    assert len(document.tables) == 1
    assert report["_decision_intelligence"]["projection_kind"] == "report"
    assert report["_decision_intelligence"]["payload_digest"].startswith("sha256:")


def test_ask_projection_is_bounded_and_only_emits_whitelisted_evidence(bundle):
    ask = ask_ai_safe_projection(
        bundle,
        allowed_evidence_ids=("evidence:acme-ab",),
        max_customers=1,
        max_evidence=1,
        max_findings_per_customer=1,
        max_actions_per_customer=1,
    )

    assert [row["customer_id"] for row in ask["customers"]] == ["customer:acme"]
    assert ask["evidence_whitelist"] == ["evidence:acme-ab"]
    assert [row["evidence_id"] for row in ask["evidence"]] == ["evidence:acme-ab"]
    assert ask["truncation"]["customers_truncated"] is True
    assert ask["_decision_intelligence"]["projected_customer_ids"] == ["customer:acme"]
    assert ask["customers"][0]["findings"][0]["evidence_ids"] == ["evidence:acme-ab"]


def test_ask_projection_honors_zero_bounds(bundle):
    ask = ask_ai_safe_projection(
        bundle,
        max_customers=1,
        max_evidence=1,
        max_findings_per_customer=0,
        max_actions_per_customer=0,
    )

    assert ask["customers"][0]["findings"] == []
    assert ask["customers"][0]["actions"] == []
    assert ask["customers"][0]["risk_profile"]["next_best_actions"] == []
    assert ask["portfolio"]["actions"] == []


def test_cross_output_reconciliation_proves_all_outputs_share_bundle(bundle):
    metrics = project_canonical_portfolio_metrics(bundle)
    risks = project_legacy_risk_profiles(bundle)
    report = render_decision_brief_word(Document(), bundle)
    export = decision_brief_excel_frames(bundle)
    renewal = stamp_projection(
        bundle,
        "renewal",
        {
            "analysis_fingerprint": bundle.analysis_fingerprint,
            "canonical_metrics": metrics,
            "risk_profiles": risks,
        },
        evidence_ids=("evidence:acme-ab", "evidence:beta-state"),
    )
    dashboard = stamp_projection(
        bundle,
        "dashboard",
        {
            "analysis_fingerprint": bundle.analysis_fingerprint,
            "canonical_metrics": metrics,
            "risk_profiles": risks,
        },
    )
    ask = ask_ai_safe_projection(bundle)

    result = validate_cross_output_reconciliation(
        bundle,
        report=report,
        export=export,
        renewal=renewal,
        dashboard=dashboard,
        ask_ai=ask,
    )

    assert result.ok is True
    assert result.errors == ()
    assert result.checked_projections == (
        "report",
        "export",
        "renewal",
        "dashboard",
        "ask_ai",
    )


def test_cross_output_reconciliation_detects_post_projection_drift(bundle):
    metrics = project_canonical_portfolio_metrics(bundle)
    base = {
        "analysis_fingerprint": bundle.analysis_fingerprint,
        "canonical_metrics": metrics,
    }
    projections = {
        kind: stamp_projection(bundle, kind, base)
        for kind in ("report", "export", "renewal", "dashboard")
    }
    projections["dashboard"] = copy.deepcopy(projections["dashboard"])
    projections["dashboard"]["canonical_metrics"]["total_cases"] = 999

    result = validate_cross_output_reconciliation(
        bundle,
        report=projections["report"],
        export=projections["export"],
        renewal=projections["renewal"],
        dashboard=projections["dashboard"],
        ask_ai=ask_ai_safe_projection(bundle),
    )

    assert result.ok is False
    assert any("dashboard: payload changed after projection" in error for error in result.errors)
    assert any("dashboard: payload canonical_metrics does not match bundle" in error for error in result.errors)


def test_cross_output_reconciliation_requires_all_five_surfaces_by_default(bundle):
    report = stamp_projection(
        bundle,
        "report",
        {"canonical_metrics": project_canonical_portfolio_metrics(bundle)},
    )

    result = validate_cross_output_reconciliation(bundle, report=report)

    assert result.ok is False
    assert "export: required projection was not supplied" in result.errors
    with pytest.raises(ValueError, match="Cross-output reconciliation failed"):
        validate_cross_output_reconciliation(
            bundle,
            report=report,
            raise_on_error=True,
        )
