"""Round 52 (Phase 2): expanded KPI extraction + required-key parity tests."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from docx import Document

from report_iteration_loop import (
    SCENARIO_REQUIRED_KPIS,
    compare_kpi_parity,
    extract_docx_kpis,
    extract_xlsx_kpis,
)


def _write_compact_docx(path: Path) -> None:
    """Mirror the executive_intelligence_formatter At-a-Glance dashboard."""
    doc = Document()
    table = doc.add_table(rows=2, cols=7)
    headers = (
        "Total Customers",
        "Support Cases",
        "Critical (P1)",
        "High (P2)",
        "BEMS Escalations",
        "Software Defects",
        "Security Vulnerabilities",
    )
    values = ("52", "176", "5", "12", "3", "8", "2")
    for col_idx, header in enumerate(headers):
        table.rows[0].cells[col_idx].text = header
    for col_idx, value in enumerate(values):
        table.rows[1].cells[col_idx].text = value
    doc.add_paragraph("Risk Summary: Overall Risk Score: 14.8")
    doc.add_paragraph("Manager: All Managers")
    doc.add_paragraph("Technology: All Contact Center")
    doc.add_paragraph("Analysis Period: 90 days")
    doc.save(path)


def _write_compact_xlsx(path: Path) -> None:
    """Mirror app_simple compact Executive_Dashboard sheet shape."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Executive_Dashboard"
    sheet.append(["Metric", "Value"])
    sheet.append(["Total Customers Analyzed", "52"])
    sheet.append(["High-Risk Customers", "8"])
    sheet.append(["Critical Adoption Barriers", "3"])
    sheet.append(["Escalated Support Cases", "176"])
    sheet.append(["Overall Risk Score", "14.8"])
    sheet.append(["Analysis Period (Days)", "90"])
    sheet.append(["Technology Focus", "All Contact Center"])
    sheet.append(["Manager", "All Managers"])
    risk = workbook.create_sheet("Risk_Summary")
    risk.append(["Customer", "Risk_Score", "Risk_Level", "Risk_Band"])
    risk.append(["AcmeCorp", "14.8", "High", "High"])
    workbook.save(path)
    workbook.close()


def _write_renewal_docx(path: Path) -> None:
    """Mirror renewal Customer Health Dashboard table."""
    doc = Document()
    doc.add_paragraph("Technology: All Contact Center")
    doc.add_paragraph("Analysis Period: 90 days")
    table = doc.add_table(rows=8, cols=2)
    rows = (
        ("Support Cases (Last 90 Days)", "176"),
        ("Active Adoption Barriers", "3"),
        ("BEMS Escalations", "5"),
        ("Service Incidents (status.webex.com)", "1"),
        ("High-Impact Incidents", "0"),
        ("Correlated Service Incidents", "0"),
        ("Overall Risk Score", "14.8"),
        ("Risk Category", "High"),
    )
    for idx, (label, value) in enumerate(rows):
        table.rows[idx].cells[0].text = label
        table.rows[idx].cells[1].text = value
    doc.save(path)


def _write_renewal_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Renewal_Summary"
    sheet.append(["Customer", "Overall_Risk_Score", "Risk_Level"])
    sheet.append(["Cust A", 14, "High"])
    sheet.append(["Cust B", 16, "High"])
    sheet.append(["Cust C", 14, "Medium"])
    info = workbook.create_sheet("Report_Info")
    info.append(["Field", "Value"])
    info.append(["Technology", "All Contact Center"])
    info.append(["Manager", "All Managers"])
    info.append(["Days", "90"])
    workbook.save(path)
    workbook.close()


def _write_leader_docx(path: Path) -> None:
    """Mirror leader Team Performance Metrics two-column table."""
    doc = Document()
    doc.add_paragraph("Manager: Brian Frazier")
    doc.add_paragraph("Analysis Period: 90 days")
    table = doc.add_table(rows=7, cols=2)
    rows = (
        ("Customers", "15"),
        ("Team Members", "3"),
        ("Action Plans", "12"),
        ("Adoption Barriers", "8"),
        ("Customer Pulse Records", "5"),
        ("TAC Cases", "9"),
        ("BEMS Escalations", "3"),
    )
    for idx, (label, value) in enumerate(rows):
        table.rows[idx].cells[0].text = label
        table.rows[idx].cells[1].text = value
    doc.save(path)


def _write_leader_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Team_Summary"
    sheet.append(
        [
            "Team_Member",
            "Num_Customers",
            "Num_Action_Plans",
            "Num_Adoption_Barriers",
            "Num_Customer_Pulse",
            "Num_TAC_Cases",
        ]
    )
    sheet.append(["Member 1", 5, 4, 2, 1, 3])
    sheet.append(["Member 2", 4, 3, 3, 2, 2])
    sheet.append(["Member 3", 6, 5, 3, 2, 4])
    info = workbook.create_sheet("Report_Info")
    info.append(["Field", "Value"])
    info.append(["Manager", "Brian Frazier"])
    info.append(["Days", "90"])
    workbook.save(path)
    workbook.close()


def test_round52_compact_docx_kpis_extracts_dashboard_and_paragraph_signals(tmp_path: Path):
    docx_path = tmp_path / "compact.docx"
    _write_compact_docx(docx_path)
    kpis = extract_docx_kpis(docx_path)
    values = kpis["values"]
    assert values["total_customers"] == "52"
    assert values["support_cases"] == "176"
    assert values["critical_cases"] == "5"
    assert values["high_cases"] == "12"
    assert values["bems"] == "3"
    assert values["risk_score"] == "14.8"
    assert values["manager"] == "All Managers"
    assert values["technology"] == "All Contact Center"
    assert values["window_days"] == "90"


def test_round52_compact_xlsx_kpis_extracts_executive_dashboard(tmp_path: Path):
    xlsx_path = tmp_path / "compact.xlsx"
    _write_compact_xlsx(xlsx_path)
    kpis = extract_xlsx_kpis(xlsx_path)
    values = kpis["values"]
    assert values["total_customers"] == "52"
    # Round 52 / partial-data-warning Phase 5: ``Escalated Support Cases``
    # is now its own canonical, distinct from the total ``support_cases``
    # canonical (which the DOCX renders).  The compact Executive_Dashboard
    # only ships the escalation count; the parity gate now compares both
    # metrics on their own keys instead of letting "escalated" overwrite
    # "total".
    assert values["escalated_support_cases"] == "176"
    assert "support_cases" not in values, (
        "Executive_Dashboard does not carry a separate Total-Support-Cases "
        "metric; only Escalated Support Cases is rendered there"
    )
    assert values["risk_score"] == "14.8"
    assert values["window_days"] == "90"
    assert values["manager"] == "All Managers"
    assert values["technology"] == "All Contact Center"


def test_round52_compact_strict_parity_passes_with_required_keys(tmp_path: Path):
    docx_path = tmp_path / "compact.docx"
    xlsx_path = tmp_path / "compact.xlsx"
    _write_compact_docx(docx_path)
    _write_compact_xlsx(xlsx_path)
    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(
        docx_kpis,
        xlsx_kpis,
        strict=True,
        required_keys=SCENARIO_REQUIRED_KPIS["compact"],
    )
    assert parity.passed is True, parity.details
    assert parity.details["missing_required_keys"] == []
    assert "total_customers" in parity.details["common_kpis"]
    assert parity.details["mismatches"] == {}


def test_round52_renewal_docx_extracts_health_dashboard(tmp_path: Path):
    docx_path = tmp_path / "renewal.docx"
    _write_renewal_docx(docx_path)
    kpis = extract_docx_kpis(docx_path)
    values = kpis["values"]
    assert values["support_cases"] == "176"
    assert values["open_adoption_barriers"] == "3"
    assert values["bems"] == "5"
    assert values["risk_score"] == "14.8"
    assert values["technology"] == "All Contact Center"
    assert values["window_days"] == "90"


def test_round52_renewal_xlsx_aggregates_renewal_summary_and_report_info(tmp_path: Path):
    xlsx_path = tmp_path / "renewal.xlsx"
    _write_renewal_xlsx(xlsx_path)
    kpis = extract_xlsx_kpis(xlsx_path)
    values = kpis["values"]
    assert values["total_customers"] == "3"
    # Round 52 (Phase 2): risk_score is intentionally NOT aggregated from the
    # per-customer Renewal_Summary sheet; the canonical KPI must come from
    # an authoritative summary cell. Renewal Excel does not currently expose
    # one, so this is recorded as a coverage gap rather than an averaged
    # value that would diverge from the DOCX headline.
    assert "risk_score" not in values
    assert values["technology"] == "All Contact Center"
    assert values["manager"] == "All Managers"
    assert values["window_days"] == "90"


def test_round52_renewal_strict_parity_uses_renewal_required_keys(tmp_path: Path):
    docx_path = tmp_path / "renewal.docx"
    xlsx_path = tmp_path / "renewal.xlsx"
    _write_renewal_docx(docx_path)
    _write_renewal_xlsx(xlsx_path)
    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(
        docx_kpis,
        xlsx_kpis,
        strict=True,
        required_keys=SCENARIO_REQUIRED_KPIS["renewal"],
    )
    assert parity.passed is True, parity.details
    assert parity.details["missing_required_keys"] == []


def test_round52_leader_team_summary_sheet_sums_columns(tmp_path: Path):
    xlsx_path = tmp_path / "leader.xlsx"
    _write_leader_xlsx(xlsx_path)
    kpis = extract_xlsx_kpis(xlsx_path)
    values = kpis["values"]
    assert values["team_members"] == "3"
    assert values["total_customers"] == "15"
    assert values["action_plans"] == "12"
    assert values["adoption_barriers"] == "8"
    assert values["customer_pulse"] == "5"
    assert values["support_cases"] == "9"
    assert values["window_days"] == "90"
    assert values["manager"] == "Brian Frazier"


def test_round52_leader_strict_parity_uses_leader_required_keys(tmp_path: Path):
    docx_path = tmp_path / "leader.docx"
    xlsx_path = tmp_path / "leader.xlsx"
    _write_leader_docx(docx_path)
    _write_leader_xlsx(xlsx_path)
    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(
        docx_kpis,
        xlsx_kpis,
        strict=True,
        required_keys=SCENARIO_REQUIRED_KPIS["leader"],
    )
    assert parity.passed is True, parity.details
    assert parity.details["missing_required_keys"] == []


def test_round52_strict_required_key_missing_fails_parity(tmp_path: Path):
    docx_path = tmp_path / "bare.docx"
    xlsx_path = tmp_path / "bare.xlsx"
    doc = Document()
    doc.add_paragraph("Empty body with no KPIs.")
    doc.save(docx_path)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Other"
    sheet.append(["Foo", "Bar"])
    workbook.save(xlsx_path)
    workbook.close()

    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(
        docx_kpis,
        xlsx_kpis,
        strict=True,
        required_keys=("manager", "window_days"),
    )
    assert parity.passed is False
    assert "manager" in parity.details["missing_required_keys"]
    assert "window_days" in parity.details["missing_required_keys"]
    assert parity.details["reason"] == "missing_required_kpis"


def test_round52_kpi_value_mismatch_still_fails_in_strict(tmp_path: Path):
    docx_path = tmp_path / "mismatch.docx"
    xlsx_path = tmp_path / "mismatch.xlsx"
    # Round 52 (Phase 2): use 3-column header table (mirrors compact
    # At-a-Glance Dashboard's 7-column shape) so the header heuristic kicks in
    # and pairs ``Total Customers`` (header) with ``52`` (data row).
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Total Customers"
    table.rows[0].cells[1].text = "Support Cases"
    table.rows[0].cells[2].text = "Critical (P1)"
    table.rows[1].cells[0].text = "52"
    table.rows[1].cells[1].text = "176"
    table.rows[1].cells[2].text = "5"
    doc.save(docx_path)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Executive_Dashboard"
    sheet.append(["Metric", "Value"])
    sheet.append(["Total Customers Analyzed", "12"])
    sheet.append(["Escalated Support Cases", "99"])
    workbook.save(xlsx_path)
    workbook.close()

    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(docx_kpis, xlsx_kpis, strict=True)
    assert parity.passed is False
    assert "total_customers" in parity.details["mismatches"]
    assert parity.details["mismatches"]["total_customers"]["docx"] == "52"
    assert parity.details["mismatches"]["total_customers"]["xlsx"] == "12"
