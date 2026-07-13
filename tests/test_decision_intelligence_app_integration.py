"""Focused app-worker Decision Intelligence V2 seam coverage."""

from __future__ import annotations

import inspect
from pathlib import Path
import sys
import types
from unittest import mock

import pandas as pd
import pytest

# The focused suite is offline.  app_simple imports connector clients, but the
# seam under test never calls them.
if "hvac" not in sys.modules:
    hvac_stub = types.ModuleType("hvac")
    hvac_stub.Client = mock.MagicMock
    sys.modules["hvac"] = hvac_stub
if "snowflake.connector" not in sys.modules:
    snowflake_stub = types.ModuleType("snowflake")
    snowflake_stub.__path__ = []
    connector_stub = types.ModuleType("snowflake.connector")
    connector_stub.DictCursor = object()
    connector_stub.connect = mock.MagicMock()
    snowflake_stub.connector = connector_stub
    sys.modules["snowflake"] = snowflake_stub
    sys.modules["snowflake.connector"] = connector_stub
if "feedparser" not in sys.modules:
    feedparser_stub = types.ModuleType("feedparser")
    feedparser_stub.parse = mock.MagicMock(return_value={})
    sys.modules["feedparser"] = feedparser_stub

import app_simple
import decision_intelligence


def _sources() -> dict[str, object]:
    return {
        "subscriptions": pd.DataFrame(
            [
                {
                    "BU_NAME": "Acme Corp",
                    "ACCOUNT_ID_C": "ACC-1",
                    "SUBSCRIPTION_ID": "SUB-1",
                    "RENEWAL_RISK_CATEGORY": "High",
                    "STATUS": "Active",
                }
            ]
        ),
        "adoption_barriers": [
            pd.DataFrame(
                [
                    {
                        "BU_NAME": "Acme Corp",
                        "ACCOUNT_ID_C": "ACC-1",
                        "ID": "AB-1",
                        "SEVERITY_C": "Critical",
                        "AB_STATUS_C": "Open",
                    }
                ]
            )
        ],
        "support_cases": pd.DataFrame(
            [
                {
                    "Customer Name": "Acme Corp",
                    "SR Number": "SR-1",
                    "Severity": "P1",
                    "Status": "Open",
                    "case_type_class": "break_fix_technical",
                }
            ]
        ),
        "customer_pulse": pd.DataFrame(
            [{"BU_NAME": "Acme Corp", "ID": "CP-1", "SCORE__C": 3.0}]
        ),
        "action_plans": pd.DataFrame(
            [{"BU_NAME": "Acme Corp", "ID": "AP-1", "STATUS_C": "Open"}]
        ),
        "success_priorities": pd.DataFrame(
            [
                {
                    "RELATED_CUSTOMER__C": "Acme Corp",
                    "ID": "SP-1",
                    "STATUS_C": "Open",
                }
            ]
        ),
        "external_incidents": [],
    }


def test_shared_boundary_builds_once_persists_and_projects(
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setenv("ADOPTIQ_ANALYSIS_SNAPSHOT_DIR", str(tmp_path / "snapshots"))
    original_build = decision_intelligence.build_analysis_bundle
    calls = []

    def _spy_build(request, sources, **kwargs):
        calls.append((request, sources, kwargs))
        return original_build(request, sources, **kwargs)

    monkeypatch.setattr(decision_intelligence, "build_analysis_bundle", _spy_build)
    status = {}
    warnings = []
    state = app_simple._decision_intelligence_v2_prepare(
        report_mode="compact",
        status=status,
        manager="Manager One",
        technology="All",
        days=90,
        customer_name="Acme Corp",
        data_retrieved_at="2026-07-13T12:00:00Z",
        partial_data_warnings=warnings,
        **_sources(),
    )

    assert len(calls) == 1
    assert state["bundle"] is not None
    assert state["warning"] == ""
    assert set(state["risk_profiles"]) == {"Acme Corp"}
    assert state["portfolio_metrics"]["total_customers"] == 1
    assert state["renewal_analyses"]["Acme Corp"]["break_fix_cases_count"] == 1
    assert state["renewal_analyses"]["Acme Corp"]["provisioning_cases_count"] == 0
    assert set(state["excel_sheets"]) == {
        "Decision_Brief",
        "Customer_Decision_Briefs",
    }
    assert status["decision_intelligence_v2_status"] == "canonical"
    assert status["analysis_fingerprint"] == state["bundle"].analysis_fingerprint
    assert status["analysis_schema_version"] == state["bundle"].schema_version
    assert Path(status["analysis_snapshot_path"]).is_file()
    assert warnings == []

    legacy_called = False

    def _legacy():
        nonlocal legacy_called
        legacy_called = True
        return {}

    projected = app_simple._decision_intelligence_risk_profiles_or_legacy(
        state, _legacy
    )
    assert projected == state["risk_profiles"]
    assert legacy_called is False


def test_failure_is_visible_and_only_then_uses_legacy(
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setenv("ADOPTIQ_ANALYSIS_SNAPSHOT_DIR", str(tmp_path / "snapshots"))

    def _fail(*_args, **_kwargs):
        raise RuntimeError("synthetic canonical failure")

    monkeypatch.setattr(decision_intelligence, "build_analysis_bundle", _fail)
    status = {}
    warnings = []
    state = app_simple._decision_intelligence_v2_prepare(
        report_mode="renewal_single",
        status=status,
        manager="Manager One",
        technology="All",
        days=30,
        customer_name="Acme Corp",
        partial_data_warnings=warnings,
        **_sources(),
    )

    assert state["bundle"] is None
    assert "legacy compatibility calculators" in state["warning"]
    assert status["decision_intelligence_v2_status"] == "legacy_fallback"
    assert warnings[0]["kind"] == "canonical_bundle_failed"
    assert list(state["excel_sheets"]) == ["Decision_Brief"]

    expected = {"Acme Corp": {"score": 4.2}}
    assert app_simple._decision_intelligence_risk_profiles_or_legacy(
        state, lambda: expected
    ) == expected


def test_external_fetch_failure_is_missing_while_successful_empty_is_observed(
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setenv("ADOPTIQ_ANALYSIS_SNAPSHOT_DIR", str(tmp_path / "snapshots"))
    expected_bug = {"id": "BUG-1"}
    monkeypatch.setattr(
        app_simple,
        "fetch_help_webex_bugs",
        mock.Mock(return_value=[expected_bug]),
    )
    monkeypatch.setattr(
        app_simple,
        "fetch_status_incidents",
        mock.Mock(side_effect=RuntimeError("synthetic incident outage")),
    )
    failed_warnings = []

    bugs, failed_incidents = (
        app_simple._fetch_external_intelligence_preserving_availability(
            days=90,
            partial_data_warnings=failed_warnings,
            report_label="Compact report",
        )
    )

    assert bugs == [expected_bug]
    assert failed_incidents is None
    assert [warning["dataset"] for warning in failed_warnings] == [
        "ext_incidents"
    ]
    assert failed_warnings[0]["kind"] == "fetch_failed"

    failed_sources = _sources()
    failed_sources["external_incidents"] = failed_incidents
    failed_state = app_simple._decision_intelligence_v2_prepare(
        report_mode="compact",
        status={},
        manager="Manager One",
        technology="All",
        days=90,
        customer_name="Acme Corp",
        data_retrieved_at="2026-07-13T12:00:00Z",
        **failed_sources,
    )
    assert (
        failed_state["bundle"].context.source_availability["external_incidents"]
        == "missing"
    )

    monkeypatch.setattr(
        app_simple,
        "fetch_help_webex_bugs",
        mock.Mock(side_effect=RuntimeError("synthetic bug outage")),
    )
    incident_fetch = mock.Mock(return_value=[])
    monkeypatch.setattr(app_simple, "fetch_status_incidents", incident_fetch)
    empty_warnings = []
    empty_bugs, empty_incidents = (
        app_simple._fetch_external_intelligence_preserving_availability(
            days=30,
            partial_data_warnings=empty_warnings,
            report_label="Comprehensive report",
        )
    )

    assert empty_bugs == []
    assert empty_incidents == []
    incident_fetch.assert_called_once_with(days_back=30)
    assert [warning["dataset"] for warning in empty_warnings] == ["ext_bugs"]

    empty_sources = _sources()
    empty_sources["external_incidents"] = empty_incidents
    empty_state = app_simple._decision_intelligence_v2_prepare(
        report_mode="comprehensive",
        status={},
        manager="Manager One",
        technology="All",
        days=30,
        customer_name="Acme Corp",
        data_retrieved_at="2026-07-13T12:00:00Z",
        **empty_sources,
    )
    assert (
        empty_state["bundle"].context.source_availability["external_incidents"]
        == "observed_empty"
    )


@pytest.mark.parametrize(
    "worker",
    (
        app_simple.run_compact_analysis,
        app_simple.run_comprehensive_analysis,
        app_simple.run_leader_report_generation,
    ),
)
def test_app_workers_share_incident_availability_fetch_boundary(worker) -> None:
    source = inspect.getsource(worker)

    assert source.count(
        "_fetch_external_intelligence_preserving_availability("
    ) == 1
    assert "fetch_status_incidents(" not in source


def test_compact_scope_failures_never_feed_unfiltered_csconsole_to_v2(
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setenv("ADOPTIQ_ANALYSIS_SNAPSHOT_DIR", str(tmp_path / "snapshots"))
    alias_rows = pd.DataFrame(
        [
            {
                "ID": "CUSTOMER-GOOD",
                "customer_name": "Acme Corp",
                "BU_NAME": "ACME CORP",
            },
            {
                "ID": "CUSTOMER-CONFLICT",
                "customer_name": "Acme Corp",
                "BU_NAME": "Beta LLC",
            },
            {
                "ID": "CUSTOMER-NONSCALAR",
                "customer_name": "Acme Corp",
                "BU_NAME": ["Acme Corp"],
            },
        ]
    )
    monkeypatch.setattr(
        app_simple,
        "_filter_csconsole_data_by_technology",
        lambda frame, *_args, **_kwargs: frame.copy(deep=True),
    )
    monkeypatch.setattr(
        app_simple,
        "_apply_scope_filter_ab",
        lambda frame, *_args, **_kwargs: frame.copy(deep=True),
    )
    alias_warnings = []
    alias_prepared = app_simple._decision_intelligence_prepare_compact_sources(
        technology="Webex Contact Center",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["ACC-1"],
        action_plans=alias_rows,
        customer_pulse=alias_rows,
        success_priorities=alias_rows,
        adoption_barriers=alias_rows,
        partial_data_warnings=alias_warnings,
    )
    assert all(
        frame["ID"].tolist() == ["CUSTOMER-GOOD"]
        for frame in alias_prepared.values()
    )
    assert all(
        frame.attrs["customer_alias_scope_excluded_rows"] == 2
        and frame.attrs["customer_alias_conflicting_rows"] == 1
        and frame.attrs["customer_alias_invalid_rows"] == 1
        for frame in alias_prepared.values()
    )
    assert {warning["dataset"] for warning in alias_warnings} == {
        "csconsole_action_plans",
        "csconsole_customer_pulse",
        "csconsole_success_priorities",
        "csconsole_adoption_barriers",
    }
    assert {warning["kind"] for warning in alias_warnings} == {
        "customer_scope_excluded"
    }

    raw = {
        "action_plans": pd.DataFrame([{"ID": "UNFILTERED-AP"}]),
        "customer_pulse": pd.DataFrame([{"ID": "UNFILTERED-CP"}]),
        "success_priorities": pd.DataFrame([{"ID": "UNFILTERED-SP"}]),
        "adoption_barriers": pd.DataFrame([{"ID": "UNFILTERED-AB"}]),
    }

    def _scope_failure(*_args, **_kwargs):
        raise RuntimeError("synthetic technology filter failure")

    monkeypatch.setattr(
        app_simple,
        "_filter_csconsole_data_by_technology",
        _scope_failure,
    )
    monkeypatch.setattr(app_simple, "_apply_scope_filter_ab", _scope_failure)
    warnings = []
    prepared = app_simple._decision_intelligence_prepare_compact_sources(
        technology="Webex Contact Center",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["ACC-1"],
        partial_data_warnings=warnings,
        **raw,
    )

    assert prepared == {
        "action_plans": None,
        "customer_pulse": None,
        "success_priorities": None,
        "adoption_barriers": None,
    }
    assert {warning["dataset"] for warning in warnings} == {
        "csconsole_action_plans",
        "csconsole_customer_pulse",
        "csconsole_success_priorities",
        "csconsole_adoption_barriers",
    }
    assert {warning["kind"] for warning in warnings} == {
        "technology_scope_failed"
    }

    def _unverifiable_scope(*_args, **_kwargs):
        unavailable = pd.DataFrame()
        unavailable.attrs.update(
            {
                "fetch_error": "technology evidence columns unavailable",
                "fetch_error_kind": "technology_scope_unverifiable",
            }
        )
        return unavailable

    monkeypatch.setattr(
        app_simple,
        "_filter_csconsole_data_by_technology",
        _unverifiable_scope,
    )
    monkeypatch.setattr(app_simple, "_apply_scope_filter_ab", _unverifiable_scope)
    diagnostic_warnings = []
    diagnostic_prepared = app_simple._decision_intelligence_prepare_compact_sources(
        technology="Webex Contact Center",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["ACC-1"],
        partial_data_warnings=diagnostic_warnings,
        **raw,
    )
    assert all(
        frame.empty
        and frame.attrs["fetch_error_kind"]
        == "technology_scope_unverifiable"
        for frame in diagnostic_prepared.values()
    )
    assert {warning["dataset"] for warning in diagnostic_warnings} == {
        "csconsole_action_plans",
        "csconsole_customer_pulse",
        "csconsole_success_priorities",
        "csconsole_adoption_barriers",
    }
    assert {warning["kind"] for warning in diagnostic_warnings} == {
        "technology_scope_unverifiable"
    }

    base = _sources()
    state = app_simple._decision_intelligence_v2_prepare(
        report_mode="compact",
        status={},
        manager="Manager One",
        technology="Webex Contact Center",
        days=90,
        customer_name="Acme Corp",
        subscriptions=base["subscriptions"],
        adoption_barriers=[],
        support_cases=base["support_cases"],
        customer_pulse=prepared["customer_pulse"],
        action_plans=prepared["action_plans"],
        success_priorities=prepared["success_priorities"],
        external_incidents=[],
        data_retrieved_at="2026-07-13T12:00:00Z",
        partial_data_warnings=warnings,
    )
    availability = state["bundle"].context.source_availability
    assert availability["action_plans"] == "missing"
    assert availability["customer_pulse"] == "missing"
    assert availability["success_priorities"] == "missing"
    assert availability["adoption_barriers"] == "missing"

    compact_source = inspect.getsource(app_simple.run_compact_analysis)
    assert compact_source.count(
        "_decision_intelligence_prepare_compact_sources("
    ) == 1
    assert "success_priorities=_compact_v2_success_priorities" in compact_source
    assert "_compact_v2_action_plans = csconsole_action_plans" not in compact_source
    assert "_compact_v2_customer_pulse = csconsole_customer_pulse" not in compact_source
    assert "_compact_v2_csconsole_ab = csconsole_adoption_barriers" not in compact_source


def test_compact_renewal_and_leader_history_receive_v2_identity() -> None:
    status = {
        "analysis_schema_version": "2.0.0",
        "analysis_fingerprint": "analysis:abc",
        "analysis_request_fingerprint": "request:def",
        "analysis_comparison_scope_fingerprint": "scope:ghi",
        "analysis_snapshot_path": "/synthetic/snapshot.json",
    }
    assert app_simple._decision_intelligence_history_kwargs(status) == {
        "analysis_schema_version": "2.0.0",
        "analysis_fingerprint": "analysis:abc",
        "analysis_request_fingerprint": "request:def",
        "analysis_comparison_scope_fingerprint": "scope:ghi",
        "analysis_snapshot_path": "/synthetic/snapshot.json",
    }

    for worker in (
        app_simple.run_compact_analysis,
        app_simple.run_customer_renewal_analysis,
        app_simple.run_leader_report_generation,
    ):
        source = inspect.getsource(worker)
        assert source.count(
            "**_decision_intelligence_history_kwargs(status)"
        ) == 1


def test_word_and_report_info_reuse_the_existing_bundle(
    monkeypatch,
    tmp_path,
) -> None:
    from docx import Document

    monkeypatch.setenv("ADOPTIQ_ANALYSIS_SNAPSHOT_DIR", str(tmp_path / "snapshots"))
    state = app_simple._decision_intelligence_v2_prepare(
        report_mode="renewal_single",
        status={},
        manager="Manager One",
        technology="All",
        days=30,
        customer_name="Acme Corp",
        data_retrieved_at="2026-07-13T12:00:00Z",
        **_sources(),
    )
    path = tmp_path / "report.docx"
    document = Document()
    document.add_heading("Legacy Renewal Report", 0)
    document.save(path)

    assert app_simple._decision_intelligence_append_word(str(path), state) is True
    assert app_simple._decision_intelligence_append_word(str(path), state) is True
    rendered = Document(path)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert text.count("Portfolio Decision Brief") == 1
    assert "Customer Decision Brief: Acme Corp" in text

    rows = app_simple._decision_intelligence_report_info_rows(state)
    items = {row["Item"]: row["Value"] for row in rows}
    assert items["Analysis_Fingerprint"] == state["bundle"].analysis_fingerprint
    assert items["Analysis_Schema_Version"] == state["bundle"].schema_version
    assert Path(items["Analysis_Snapshot_Path"]).is_file()

    xlsx_path = tmp_path / "report.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        pd.DataFrame(
            [{"Item": "Export type", "Value": "Standard"}]
        ).to_excel(writer, sheet_name="Report_Info", index=False)
    assert app_simple._decision_intelligence_append_excel_report_info(
        str(xlsx_path), state
    ) is True
    assert app_simple._decision_intelligence_append_excel_report_info(
        str(xlsx_path), state
    ) is True
    report_info = pd.read_excel(xlsx_path, sheet_name="Report_Info")
    assert report_info["Item"].tolist().count("Analysis_Fingerprint") == 1
    workbook_items = dict(zip(report_info["Item"], report_info["Value"]))
    assert workbook_items["Analysis_Fingerprint"] == (
        state["bundle"].analysis_fingerprint
    )


def test_comprehensive_source_seam_merges_final_action_plans_once(
    monkeypatch,
) -> None:
    csconsole = pd.DataFrame(
        [
            {
                "ID": "AP-1",
                "BU_NAME": "Acme Corp",
                "STATUS_C": "Open",
                "LastModifiedDate": "2026-07-01T00:00:00Z",
                "Source": "CSConsole",
            },
            {
                "ID": "AP-CUSTOMER-CONFLICT",
                "customer_name": "Acme Corp",
                "BU_NAME": "Beta LLC",
                "STATUS_C": "Open",
                "Source": "CSConsole",
            },
            {
                "ID": "AP-CUSTOMER-NONSCALAR",
                "customer_name": "Acme Corp",
                "BU_NAME": {"name": "Acme Corp"},
                "STATUS_C": "Open",
                "Source": "CSConsole",
            },
        ]
    )
    snowflake = pd.DataFrame(
        [
            {
                "ID": "AP-1",
                "BU_NAME": "Acme Corp",
                "ACCOUNT_ID_C": "ACC-1",
                "STATUS_C": "Closed",
                "LastModifiedDate": "2026-07-12T00:00:00Z",
                "Source": "Snowflake",
            },
            {
                "ID": "AP-2",
                "BU_NAME": "Acme Corp",
                "ACCOUNT_ID_C": "ACC-1",
                "STATUS_C": "Open",
                "LastModifiedDate": "2026-07-11T00:00:00Z",
                "Source": "Snowflake",
            },
        ]
    )
    fetch = mock.Mock(return_value=snowflake)
    monkeypatch.setattr(app_simple, "_r65_fetch_aps_snowflake", fetch)
    monkeypatch.setattr(
        app_simple,
        "_filter_csconsole_data_by_technology",
        lambda frame, *_args, **_kwargs: frame.copy(deep=True),
    )
    pulse = pd.DataFrame([{"BU_NAME": "Acme Corp", "ID": "P-1"}])
    priorities = pd.DataFrame(
        [{"RELATED_CUSTOMER__C": "Acme Corp", "ID": "SP-1"}]
    )
    barriers = pd.DataFrame([{"BU_NAME": "Acme Corp", "ID": "AB-1"}])
    warnings = []

    prepared = app_simple._decision_intelligence_prepare_comprehensive_sources(
        ctx=object(),
        technology="All",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["ACC-1"],
        owner_emails=["owner@example.com"],
        action_plans=csconsole,
        customer_pulse=pulse,
        success_priorities=priorities,
        adoption_barriers=barriers,
        partial_data_warnings=warnings,
    )

    fetch.assert_called_once_with(
        mock.ANY,
        ["ACC-1"],
        90,
        owner_emails=["owner@example.com"],
    )
    merged = prepared["action_plans"].set_index("ID")
    assert sorted(merged.index.tolist()) == ["AP-1", "AP-2"]
    assert merged.loc["AP-1", "Source"] == "Snowflake"
    assert prepared["action_plan_provenance"] == "csconsole+snowflake"
    assert prepared["success_priorities"].equals(priorities)
    assert warnings == [
        {
            "dataset": "csconsole_action_plans",
            "error": "conflicting_or_invalid_customer_aliases",
            "kind": "customer_scope_excluded",
            "effect": (
                "2 row(s) were omitted because populated customer identity "
                "aliases conflicted or were non-scalar."
            ),
        }
    ]


def test_comprehensive_owner_widening_cannot_escape_selected_accounts(
    monkeypatch,
) -> None:
    snowflake = pd.DataFrame(
        [
            {
                "ID": "AP-IN",
                "BU_NAME": "Acme Corp",
                "ACCOUNT_ID_C": " 001ABCDEFGHIJKL ",
                "TECHNOLOGY_C": "Webex Contact Center",
                "STATUS_C": "Open",
            },
            {
                "ID": "AP-CASE-COLLISION",
                "BU_NAME": "Other Corp",
                "ACCOUNT_ID_C": "001abcdefghijklAAA",
                "TECHNOLOGY_C": "Webex Contact Center",
                "STATUS_C": "Open",
            },
            {
                "ID": "AP-SUFFIX-COLLISION",
                "BU_NAME": "Other Corp",
                "ACCOUNT_ID_C": "001ABCDEFGHIJKLBBB",
                "TECHNOLOGY_C": "Webex Contact Center",
                "STATUS_C": "Open",
            },
            {
                "ID": "AP-OUT",
                "BU_NAME": "Other Corp",
                "ACCOUNT_ID_C": "ACC-OUT",
                "TECHNOLOGY_C": "Webex Contact Center",
                "STATUS_C": "Open",
            },
            {
                "ID": "AP-CONFLICT",
                "BU_NAME": "Acme Corp",
                "ACCOUNT_ID_C": "001ABCDEFGHIJKL",
                "account_id_c": "ACC-OUT",
                "TECHNOLOGY_C": "Webex Contact Center",
                "STATUS_C": "Open",
            },
        ]
    )
    fetch = mock.Mock(return_value=snowflake)
    monkeypatch.setattr(
        app_simple,
        "_r65_fetch_aps_snowflake",
        fetch,
    )
    monkeypatch.setattr(
        app_simple,
        "_filter_csconsole_data_by_technology",
        lambda frame, *_args, **_kwargs: frame.copy(deep=True),
    )
    warnings = []

    prepared = app_simple._decision_intelligence_prepare_comprehensive_sources(
        ctx=object(),
        technology="Webex Contact Center",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["001ABCDEFGHIJKLY55"],
        owner_emails=["owner@example.com"],
        action_plans=pd.DataFrame(),
        customer_pulse=pd.DataFrame(),
        success_priorities=pd.DataFrame(),
        adoption_barriers=pd.DataFrame(),
        partial_data_warnings=warnings,
    )

    assert prepared["snowflake_action_plans"]["ID"].tolist() == ["AP-IN"]
    assert prepared["action_plans"]["ID"].tolist() == ["AP-IN"]
    assert prepared["action_plans"].attrs["account_scope_filter"] == (
        "strict_membership"
    )
    assert prepared["action_plans"].attrs["scope_excluded_rows"] == 4
    assert prepared["action_plans"].attrs["technology_scope_filter"] == (
        "explicit_columns"
    )
    assert [warning["kind"] for warning in warnings] == [
        "account_scope_excluded"
    ]

    fetch.return_value = pd.DataFrame(
        [{"ID": "AP-UNVERIFIABLE", "STATUS_C": "Open"}]
    )
    unverified_warnings = []
    unverified = app_simple._decision_intelligence_prepare_comprehensive_sources(
        ctx=object(),
        technology="All",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["001ABCDEFGHIJKLY55"],
        owner_emails=["owner@example.com"],
        action_plans=pd.DataFrame(),
        customer_pulse=pd.DataFrame(),
        success_priorities=pd.DataFrame(),
        adoption_barriers=pd.DataFrame(),
        partial_data_warnings=unverified_warnings,
    )
    assert unverified["action_plans"].empty
    assert unverified["action_plans"].attrs["fetch_error_kind"] == (
        "partial_action_plan_source_failure"
    )
    assert unverified["snowflake_action_plans"].attrs[
        "fetch_error_kind"
    ] == "account_scope_unverifiable"
    assert [warning["kind"] for warning in unverified_warnings] == [
        "account_scope_unverifiable"
    ]

    fetch.return_value = pd.DataFrame(
        [
            {
                "ID": "AP-NO-TECH-EVIDENCE",
                "ACCOUNT_ID_C": "001ABCDEFGHIJKL",
                "STATUS_C": "Open",
            }
        ]
    )
    technology_warnings = []
    technology_unverified = (
        app_simple._decision_intelligence_prepare_comprehensive_sources(
            ctx=object(),
            technology="Webex Contact Center",
            days=90,
            customer_names=["Acme Corp"],
            account_ids=["001ABCDEFGHIJKLY55"],
            owner_emails=["owner@example.com"],
            action_plans=pd.DataFrame(),
            customer_pulse=pd.DataFrame(),
            success_priorities=pd.DataFrame(),
            adoption_barriers=pd.DataFrame(),
            partial_data_warnings=technology_warnings,
        )
    )
    assert technology_unverified["action_plans"].empty
    assert technology_unverified["snowflake_action_plans"].attrs[
        "fetch_error_kind"
    ] == "technology_scope_unverifiable"
    assert [warning["kind"] for warning in technology_warnings] == [
        "technology_scope_unverifiable"
    ]


def test_comprehensive_explicit_technology_cannot_escape_same_account(
    monkeypatch,
) -> None:
    snowflake = pd.DataFrame(
        [
            {
                "ID": "AP-CONTACT-CENTER",
                "BU_NAME": "Acme Corp",
                "ACCOUNT_ID_C": "ACC-IN",
                "TECHNOLOGY_C": "Webex Contact Center",
                "STATUS_C": "Open",
            },
            {
                "ID": "AP-MEETINGS",
                "BU_NAME": "Acme Corp",
                "ACCOUNT_ID_C": "ACC-IN",
                "technology_c": "Webex Meetings",
                "STATUS_C": "Open",
            },
            {
                "ID": "AP-TECH-CONFLICT",
                "BU_NAME": "Acme Corp",
                "ACCOUNT_ID_C": "ACC-IN",
                "TECHNOLOGY_C": "Webex Contact Center",
                "sub_technology_c": "Webex Meetings",
                "STATUS_C": "Open",
            },
        ]
    )
    fetch = mock.Mock(return_value=snowflake)
    monkeypatch.setattr(
        app_simple,
        "_r65_fetch_aps_snowflake",
        fetch,
    )
    warnings = []

    prepared = app_simple._decision_intelligence_prepare_comprehensive_sources(
        ctx=object(),
        technology="Webex Contact Center",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["ACC-IN"],
        owner_emails=["owner@example.com"],
        action_plans=pd.DataFrame(),
        customer_pulse=pd.DataFrame(),
        success_priorities=pd.DataFrame(),
        adoption_barriers=pd.DataFrame(),
        partial_data_warnings=warnings,
    )

    assert prepared["snowflake_action_plans"]["ID"].tolist() == [
        "AP-CONTACT-CENTER"
    ]
    assert prepared["action_plans"]["ID"].tolist() == [
        "AP-CONTACT-CENTER"
    ]
    assert prepared["action_plans"].attrs["technology_scope_filter"] == (
        "explicit_columns"
    )
    assert prepared["action_plans"].attrs[
        "technology_scope_excluded_rows"
    ] == 2
    assert [warning["kind"] for warning in warnings] == [
        "technology_scope_excluded"
    ]

    all_contact_center_warnings = []
    all_contact_center = (
        app_simple._decision_intelligence_prepare_comprehensive_sources(
            ctx=object(),
            technology="All Contact Center",
            days=90,
            customer_names=["Acme Corp"],
            account_ids=["ACC-IN"],
            owner_emails=["owner@example.com"],
            action_plans=pd.DataFrame(),
            customer_pulse=pd.DataFrame(),
            success_priorities=pd.DataFrame(),
            adoption_barriers=pd.DataFrame(),
            partial_data_warnings=all_contact_center_warnings,
        )
    )
    assert all_contact_center["action_plans"]["ID"].tolist() == [
        "AP-CONTACT-CENTER"
    ]
    assert [warning["kind"] for warning in all_contact_center_warnings] == [
        "technology_scope_excluded"
    ]

    invalid_warnings = []
    invalid = app_simple._decision_intelligence_prepare_comprehensive_sources(
        ctx=object(),
        technology="not-a-real-technology",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["ACC-IN"],
        owner_emails=["owner@example.com"],
        action_plans=pd.DataFrame([{"ID": "CS-AP"}]),
        customer_pulse=pd.DataFrame([{"ID": "PULSE"}]),
        success_priorities=pd.DataFrame([{"ID": "SP"}]),
        adoption_barriers=pd.DataFrame([{"ID": "AB"}]),
        partial_data_warnings=invalid_warnings,
    )
    assert fetch.call_count == 2
    assert invalid["action_plan_provenance"] == "invalid_technology_scope"
    assert all(
        invalid[source].empty
        for source in (
            "action_plans",
            "customer_pulse",
            "success_priorities",
            "adoption_barriers",
            "snowflake_action_plans",
        )
    )
    assert invalid["action_plans"].attrs["fetch_error_kind"] == (
        "invalid_technology_scope"
    )
    assert [warning["kind"] for warning in invalid_warnings] == [
        "invalid_technology_scope"
    ]

    monkeypatch.setattr(app_simple, "MANAGERS", ["Manager One"])
    monkeypatch.setitem(app_simple.app.config, "WTF_CSRF_ENABLED", False)
    response = app_simple.app.test_client().post(
        "/start_analysis",
        json={
            "manager": "Manager One",
            "technology": "not-a-real-technology",
            "report_type": "comprehensive",
            "days": 90,
        },
        headers={"X-Requested-With": "XMLHttpRequest"},
    )
    assert response.status_code == 400
    assert response.get_json()["error"] == "Unsupported technology"

    fetch.return_value = pd.DataFrame()
    scope_warnings = []
    unverifiable = app_simple._decision_intelligence_prepare_comprehensive_sources(
        ctx=object(),
        technology="Webex Contact Center",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["ACC-IN"],
        owner_emails=["owner@example.com"],
        action_plans=pd.DataFrame(
            [{"ID": "NO-TECH-EVIDENCE", "ACCOUNT_ID_C": "ACC-IN"}]
        ),
        customer_pulse=pd.DataFrame(),
        success_priorities=pd.DataFrame(),
        adoption_barriers=pd.DataFrame(),
        partial_data_warnings=scope_warnings,
    )
    assert unverifiable["action_plans"].empty
    assert unverifiable["action_plans"].attrs["fetch_error_kind"] == (
        "technology_scope_unverifiable"
    )
    assert [warning["kind"] for warning in scope_warnings] == [
        "technology_scope_unverifiable"
    ]

    fetch.return_value = pd.DataFrame(
        [
            ["AP-GOOD", "ACC-IN", "ACC-IN", "Webex Contact Center", ""],
            ["AP-ACCOUNT-CONFLICT", "ACC-IN", "ACC-OUT", "Webex Contact Center", ""],
            ["AP-TECH-CONFLICT", "ACC-IN", "ACC-IN", "Webex Contact Center", "Webex Meetings"],
        ],
        columns=[
            "ID",
            "ACCOUNT_ID_C",
            "ACCOUNT_ID_C",
            "TECHNOLOGY_C",
            "TECHNOLOGY_C",
        ],
    )
    duplicate_alias_warnings = []
    duplicate_aliases = (
        app_simple._decision_intelligence_prepare_comprehensive_sources(
            ctx=object(),
            technology="Webex Contact Center",
            days=90,
            customer_names=["Acme Corp"],
            account_ids=["ACC-IN"],
            owner_emails=["owner@example.com"],
            action_plans=pd.DataFrame(),
            customer_pulse=pd.DataFrame(),
            success_priorities=pd.DataFrame(),
            adoption_barriers=pd.DataFrame(),
            partial_data_warnings=duplicate_alias_warnings,
        )
    )
    assert duplicate_aliases["action_plans"]["ID"].tolist() == ["AP-GOOD"]
    assert {warning["kind"] for warning in duplicate_alias_warnings} == {
        "account_scope_excluded",
        "technology_scope_excluded",
    }

    fetch.return_value = pd.DataFrame(
        [
            {
                "ID": "AP-COMPOUND",
                "ACCOUNT_ID_C": "ACC-IN",
                "TECHNOLOGY_C": "Webex Meetings / Webex Calling",
            },
            {
                "ID": "AP-MAPPING",
                "ACCOUNT_ID_C": "ACC-IN",
                "TECHNOLOGY_C": {"product": "Webex Meetings"},
            },
            {
                "ID": "AP-LIST",
                "ACCOUNT_ID_C": "ACC-IN",
                "TECHNOLOGY_C": ["Webex Meetings"],
            },
        ]
    )
    malformed_warnings = []
    malformed = app_simple._decision_intelligence_prepare_comprehensive_sources(
        ctx=object(),
        technology="Webex Meetings & Messaging",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["ACC-IN"],
        owner_emails=["owner@example.com"],
        action_plans=pd.DataFrame(),
        customer_pulse=pd.DataFrame(),
        success_priorities=pd.DataFrame(),
        adoption_barriers=pd.DataFrame(),
        partial_data_warnings=malformed_warnings,
    )
    assert malformed["action_plans"].empty
    assert malformed["snowflake_action_plans"].empty
    assert [warning["kind"] for warning in malformed_warnings] == [
        "technology_scope_excluded"
    ]


def test_comprehensive_action_plan_failure_is_not_observed_empty(
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setenv("ADOPTIQ_ANALYSIS_SNAPSHOT_DIR", str(tmp_path / "snapshots"))
    monkeypatch.setattr(
        app_simple,
        "_r65_fetch_aps_snowflake",
        mock.Mock(side_effect=RuntimeError("synthetic Snowflake AP outage")),
    )
    monkeypatch.setattr(
        app_simple,
        "_filter_csconsole_data_by_technology",
        lambda frame, *_args, **_kwargs: frame.copy(deep=True),
    )
    warnings = []
    prepared = app_simple._decision_intelligence_prepare_comprehensive_sources(
        ctx=object(),
        technology="All",
        days=90,
        customer_names=["Acme Corp"],
        account_ids=["ACC-1"],
        owner_emails=["owner@example.com"],
        action_plans=pd.DataFrame(),
        customer_pulse=pd.DataFrame(),
        success_priorities=pd.DataFrame(),
        adoption_barriers=pd.DataFrame(),
        partial_data_warnings=warnings,
    )

    plans = prepared["action_plans"]
    assert plans.empty
    assert plans.attrs["fetch_error"] == "synthetic Snowflake AP outage"
    assert plans.attrs["fetch_error_kind"] == (
        "partial_action_plan_source_failure"
    )
    assert prepared["action_plan_provenance"] == (
        "empty+snowflake_unavailable"
    )
    assert warnings[0]["dataset"] == "snowflake_action_plans"

    base = _sources()
    state = app_simple._decision_intelligence_v2_prepare(
        report_mode="comprehensive",
        status={},
        manager="Manager One",
        technology="All",
        days=90,
        customer_name="Acme Corp",
        subscriptions=base["subscriptions"],
        adoption_barriers=base["adoption_barriers"],
        support_cases=base["support_cases"],
        customer_pulse=base["customer_pulse"],
        action_plans=plans,
        success_priorities=base["success_priorities"],
        external_incidents=[],
        data_retrieved_at="2026-07-13T12:00:00Z",
        partial_data_warnings=warnings,
    )
    assert state["bundle"].context.source_availability["action_plans"] == (
        "fetch_failed"
    )


def test_comprehensive_worker_uses_one_canonical_boundary_before_legacy() -> None:
    source = inspect.getsource(app_simple.run_comprehensive_analysis)

    assert source.count("_decision_intelligence_v2_prepare(") == 1
    assert "build_analysis_bundle" not in source
    source_prepare = source.index(
        "_decision_intelligence_prepare_comprehensive_sources("
    )
    bundle_prepare = source.index("_decision_intelligence_v2_prepare(")
    legacy_risk = source.index("compute_customer_risk_profile(")
    assert source_prepare < bundle_prepare < legacy_risk
    assert source.count("_r65_fetch_aps_snowflake(") == 1
    assert "not _comprehensive_di_sources_prepared" in source
    assert "action_plans=filtered_action_plans" in source
    assert "success_priorities=filtered_success_priorities" in source
    assert "risk_profiles = dict(_decision_v2.get(\"risk_profiles\") or {})" in source
    assert "portfolio_metrics = dict(_canonical_projection)" in source
    assert "all_sheets.update(_decision_v2.get(\"excel_sheets\") or {})" in source
    assert "_decision_intelligence_append_word(docx_path, _decision_v2)" in source
    assert "_decision_intelligence_append_excel_report_info(" in source
    assert "analysis_fingerprint=status.get('analysis_fingerprint', '')" in source
    assert '"portfolio_health_score": "B"' not in source


def test_active_workers_have_one_shared_boundary_and_leader_only_merges_handoff() -> None:
    compact_source = inspect.getsource(app_simple.run_compact_analysis)
    renewal_source = inspect.getsource(app_simple.run_customer_renewal_analysis)
    leader_source = inspect.getsource(app_simple.run_leader_report_generation)
    subscription_source = inspect.getsource(app_simple.run_subscription_analysis)

    assert compact_source.count("_decision_intelligence_v2_prepare(") == 1
    assert renewal_source.count("_decision_intelligence_v2_prepare(") == 1
    assert "build_analysis_bundle" not in compact_source
    assert "build_analysis_bundle" not in renewal_source
    assert "decision_intelligence_excel_frames" in leader_source
    assert "_decision_intelligence_v2_prepare(" not in leader_source
    assert subscription_source.count("_decision_intelligence_v2_prepare(") == 1
    assert subscription_source.count("fetch_subscription_data(") == 1
    assert subscription_source.count("get_subscription_renewal_risk(") == 1
    assert "render_decision_brief_word" in subscription_source
    assert "Report_Info" in subscription_source


def test_report_modes_share_a_temporal_comparison_scope(
    monkeypatch,
    tmp_path,
) -> None:
    """Presentation mode must not split one factual portfolio timeline."""

    monkeypatch.setenv("ADOPTIQ_ANALYSIS_SNAPSHOT_DIR", str(tmp_path / "snapshots"))
    common = {
        "status": {},
        "manager": "Manager One",
        "technology": "All",
        "days": 90,
        "data_retrieved_at": "2026-07-13T12:00:00Z",
        **_sources(),
    }
    compact = app_simple._decision_intelligence_v2_prepare(
        report_mode="compact",
        **common,
    )
    renewal = app_simple._decision_intelligence_v2_prepare(
        report_mode="renewal_portfolio",
        **{**common, "status": {}},
    )

    assert compact["bundle"].request.request_fingerprint != (
        renewal["bundle"].request.request_fingerprint
    )
    assert compact["bundle"].context.comparison_scope_fingerprint == (
        renewal["bundle"].context.comparison_scope_fingerprint
    )
