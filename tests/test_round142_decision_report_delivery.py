"""Round 142 shared concise Word + Source Data delivery contract."""

from __future__ import annotations

import base64
from pathlib import Path

import pandas as pd
import pytest
from docx import Document

import decision_report_delivery as delivery


AS_OF = "2026-08-03T12:00:00Z"
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _team_fixture() -> dict:
    plans = []
    for index in range(1, 16):
        plans.append(
            {
                "ID": f"AP-{index:03d}",
                "BU_NAME": "Acme Corporation" if index <= 8 else "Beta Industries",
                "SUBJECT_C": "" if index == 2 else f"Synthetic action {index}",
                "STATUS_C": (
                    "Open"
                    if index in {1, 3, 4, 5, 6, 7, 8, 9, 10, 15}
                    else "On Hold"
                    if index == 2
                    else "Completed - Successful"
                    if index in {11, 12, 13}
                    else "custom state"
                ),
                "DUE_DATE_C": "2026-07-20" if index <= 3 else "2026-08-10",
                "CREATED_DATE_C": f"2026-07-{min(index, 28):02d}",
                "PRIORITY_C": "P1" if index <= 3 else "P3",
                "NEXT_ACTION_C": f"Owner step {index}",
                "NEXT_ACTION_OWNER_C": "Alex Rivera",
            }
        )
    # Stable-ID duplicate should stay out of all counts and Source Data rows.
    plans.append(dict(plans[0], SUBJECT_C="Duplicate fan-out row"))
    return {
        "Alex Rivera": {
            "subscriptions": pd.DataFrame(
                [
                    {"SUBSCRIPTION_ID": "SUB-001", "ACCOUNT_ID_C": "ACC-001", "BU_NAME": "Acme Corporation"},
                    {"SUBSCRIPTION_ID": "SUB-002", "ACCOUNT_ID_C": "ACC-002", "BU_NAME": "Beta Industries"},
                ]
            ),
            "action_plans": pd.DataFrame(plans),
            "adoption_barriers": pd.DataFrame(
                [
                    {
                        "ID": "AB-001",
                        "BU_NAME": "Acme Corporation",
                        "SUBJECT_C": "Synthetic adoption blocker",
                        "SEVERITY_C": "High",
                        "STATUS_C": "Open",
                        "OPEN_DATE_C": "2026-07-05",
                    }
                ]
            ),
            "customer_pulse": pd.DataFrame(
                [
                    {
                        "ID": "CP-001",
                        "BU_NAME": "Acme Corporation",
                        "PULSE_RATING__C": "Poor",
                        "CREATED_DATE_C": "2026-07-10",
                    }
                ]
            ),
            "tac_cases": pd.DataFrame(
                [
                    {
                        "SR Number": "700000001",
                        "Customer": "Acme Corporation",
                        "Severity": "1",
                        "Case Status": "Open",
                        "Date/Time Opened": "2026-07-15",
                        "Problem Description": "BEMS CSC escalation fixture",
                        "Transaction ID": "BEMS000001",
                    }
                ]
            ),
            "success_priorities": pd.DataFrame(
                [{"ID": "SP-001", "RELATED_CUSTOMER__C": "Acme Corporation"}]
            ),
        },
        "Morgan Lee": {
            "subscriptions": pd.DataFrame(
                [{"SUBSCRIPTION_ID": "SUB-003", "ACCOUNT_ID_C": "ACC-003", "BU_NAME": "Gamma Public"}]
            ),
            "action_plans": pd.DataFrame(),
            "adoption_barriers": pd.DataFrame(),
            "customer_pulse": pd.DataFrame(),
            "tac_cases": pd.DataFrame(),
            "success_priorities": pd.DataFrame(),
        },
    }


def _facts() -> dict:
    return delivery.build_report_facts(
        _team_fixture(),
        report_type="Leader",
        scope_type="team",
        scope_value="Dana Manager team",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
        external_incidents=[{"id": "INC-001", "title": "Synthetic incident", "status": "resolved"}],
        external_bugs=[{"bug_id": "BUG-001", "title": "Synthetic known issue"}],
    )


def _fake_chart_renderer(_chart_id: str, _rows: pd.DataFrame, target: Path) -> bool:
    target.write_bytes(_TINY_PNG)
    return True


def test_shared_facts_drive_source_data_and_complete_lineage() -> None:
    facts = _facts()
    sheets = delivery.build_source_data_sheets(facts)
    contract = delivery.validate_cross_artifact_contract(facts, sheets)

    assert contract["ok"], contract["errors"]
    assert facts["kpis"]["team_members"] == 2
    assert facts["kpis"]["customers"] == 3
    assert facts["action_plan_lifecycle"]["total"] == 15
    assert len(sheets["Action_Plans"]) == 15
    assert set(delivery._FRAME_KEYS.values()).issuperset(  # noqa: SLF001 - contract inventory assertion
        {"Action_Plans", "Adoption_Barriers", "Customer_Pulse", "TAC_Cases"}
    )
    assert {
        "Report_Info",
        "Metric_Lineage",
        "Chart_Data",
        "Action_Plans",
        "Adoption_Barriers",
        "Customer_Pulse",
        "TAC_Cases",
        "BEMS",
    }.issubset(sheets)
    assert sheets["Action_Plans"]["Record_ID"].is_unique
    assert "Title unavailable" in set(sheets["Action_Plans"]["AdoptIQ_Title"])
    assert sheets["Metric_Lineage"]["Metric_Key"].is_unique


def test_concise_word_embeds_charts_and_excludes_raw_record_dump(monkeypatch) -> None:
    facts = _facts()
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    doc = delivery.build_concise_word_document(facts)
    sheets = delivery.build_source_data_sheets(facts)
    contract = delivery.validate_cross_artifact_contract(facts, sheets, doc)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

    assert contract["ok"], contract["errors"]
    assert contract["word"]["word_count"] < 5000
    assert len(doc.inline_shapes) == 4
    assert "All Action Plans" not in text
    assert "All Customer Pulse" not in text
    # This lowest-priority plan is outside the top-10 Word list but remains in Source Data.
    assert "Synthetic action 15" not in text
    assert "Synthetic action 15" in set(sheets["Action_Plans"]["AdoptIQ_Title"])


def test_source_data_file_is_separately_named_and_readable(tmp_path: Path) -> None:
    facts = _facts()
    sheets = delivery.build_source_data_sheets(facts)
    word_path = tmp_path / "AdoptIQ_Report_Leader_Dana_Manager_Team_20260803.docx"
    source_path = delivery.source_data_path_for_word(word_path)
    delivery.write_source_data_workbook(source_path, sheets)

    assert source_path.name == "AdoptIQ_Source_Data_Leader_Dana_Manager_Team_20260803.xlsx"
    written_contract = delivery.validate_written_source_workbook(source_path, facts)
    assert written_contract["ok"], written_contract["errors"]
    workbook = pd.ExcelFile(source_path)
    assert set(delivery.validate_cross_artifact_contract(facts, sheets)["required_sheets"]).issubset(
        workbook.sheet_names
    )
    action_plans = pd.read_excel(source_path, sheet_name="Action_Plans")
    assert "Record ID" in action_plans.columns or "Record_ID" in action_plans.columns
    assert len(action_plans) == 15


def test_source_writer_wraps_long_audit_fields_without_clipping(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    facts = _facts()
    source_path = tmp_path / "AdoptIQ_Source_Data_readable_audit_fields.xlsx"
    delivery.write_source_data_workbook(
        source_path,
        delivery.build_source_data_sheets(facts),
    )

    workbook = load_workbook(source_path, read_only=False, data_only=False)
    report_info = workbook["Report_Info"]
    fingerprint_row = next(
        row[0].row
        for row in report_info.iter_rows(min_row=2)
        if row[0].value == "Fact_Contract_SHA256"
    )
    assert report_info.cell(fingerprint_row, 2).alignment.wrap_text
    assert (report_info.row_dimensions[fingerprint_row].height or 0) > 15

    risk = workbook["Risk_Components"]
    detail_column = next(
        cell.column
        for cell in risk[1]
        if cell.value == "Component_Details_JSON"
    )
    long_detail_row = max(
        range(2, risk.max_row + 1),
        key=lambda row_number: len(str(risk.cell(row_number, detail_column).value or "")),
    )
    assert risk.cell(long_detail_row, detail_column).alignment.wrap_text
    assert (risk.row_dimensions[long_detail_row].height or 0) > 15


def test_source_builder_rejects_undigested_additional_sheets() -> None:
    with pytest.raises(ValueError, match="content-digested contract"):
        delivery.build_source_data_sheets(
            _facts(),
            additional_sheets={"Legacy_Raw": pd.DataFrame([{"value": 1}])},
        )


def test_fact_fingerprint_rejects_mismatched_word_or_source_data() -> None:
    facts = _facts()
    sheets = delivery.build_source_data_sheets(facts)
    doc = delivery.build_concise_word_document(facts)

    tampered_sheets = {name: frame.copy() for name, frame in sheets.items()}
    marker = tampered_sheets["Report_Info"]["Item"].eq("Fact_Contract_SHA256")
    tampered_sheets["Report_Info"].loc[marker, "Value"] = "0" * 64
    source_contract = delivery.validate_cross_artifact_contract(facts, tampered_sheets)
    assert not source_contract["ok"]
    assert any("fingerprint" in error for error in source_contract["errors"])

    tampered_records = {name: frame.copy() for name, frame in sheets.items()}
    tampered_records["Action_Plans"].loc[0, "Record_ID"] = "FABRICATED-AP"
    tampered_records["Action_Plans"].loc[0, "AdoptIQ_Title"] = "Fabricated title"
    tampered_records["TAC_Cases"].loc[0, "Record_ID"] = "FABRICATED-TAC"
    record_contract = delivery.validate_cross_artifact_contract(facts, tampered_records)
    assert not record_contract["ok"]
    assert any("content digest" in error for error in record_contract["errors"])

    tampered_info = {name: frame.copy() for name, frame in sheets.items()}
    scope_marker = tampered_info["Report_Info"]["Item"].eq("Scope_Value")
    tampered_info["Report_Info"].loc[scope_marker, "Value"] = "Unvalidated portfolio"
    info_contract = delivery.validate_cross_artifact_contract(facts, tampered_info)
    assert not info_contract["ok"]
    assert any("Report_Info scope" in error for error in info_contract["errors"])

    doc.core_properties.identifier = "0" * 64
    word_contract = delivery.validate_cross_artifact_contract(facts, sheets, doc)
    assert not word_contract["ok"]
    assert any("Word fact fingerprint" in error for error in word_contract["errors"])


def test_source_writer_defangs_formula_like_record_text(tmp_path: Path) -> None:
    facts = _facts()
    facts["external_incidents"].loc[0, "title"] = "=2+2"
    sheets = delivery.build_source_data_sheets(facts)
    source_path = tmp_path / "AdoptIQ_Source_Data_formula_safety.xlsx"

    delivery.write_source_data_workbook(source_path, sheets)

    written_contract = delivery.validate_written_source_workbook(source_path, facts)
    assert written_contract["ok"], written_contract["errors"]
    assert written_contract["formula_cells"] == 0
    serialized = pd.read_excel(source_path, sheet_name="External_Incidents")
    assert serialized.loc[0, "title"] == "=2+2"


def test_written_validator_rejects_report_info_scope_tampering(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    facts = _facts()
    source_path = tmp_path / "AdoptIQ_Source_Data_tampered_scope.xlsx"
    delivery.write_source_data_workbook(
        source_path,
        delivery.build_source_data_sheets(facts),
    )
    workbook = load_workbook(source_path)
    worksheet = workbook["Report_Info"]
    for row in worksheet.iter_rows(min_row=2):
        if row[0].value == "Scope_Value":
            row[1].value = "Fabricated scope"
            break
    workbook.save(source_path)

    contract = delivery.validate_written_source_workbook(source_path, facts)

    assert not contract["ok"]
    assert any("Report_Info scope" in error for error in contract["errors"])


def test_failed_source_is_unavailable_not_zero() -> None:
    fixture = _team_fixture()
    failed = pd.DataFrame()
    failed.attrs["fetch_error"] = "sanitized fixture outage"
    fixture["Alex Rivera"]["customer_pulse"] = failed
    facts = delivery.build_report_facts(
        fixture,
        report_type="Leader",
        scope_type="member",
        scope_value="Alex Rivera",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
        partial_data_warnings=[{"dataset": "Customer_Pulse", "kind": "fetch_failed"}],
    )

    coverage = facts["source_coverage"].set_index("Source_Sheet")
    assert coverage.loc["Customer_Pulse", "Source_State"] == "failed"
    assert pd.isna(coverage.loc["Customer_Pulse", "Record_Count"])
    activity = facts["chart_data"].loc[
        facts["chart_data"]["Metric_Key"] == "chart.activity_mix.customer_pulse"
    ].iloc[0]
    assert pd.isna(activity["Value"])
    assert activity["Source_State"] == "failed"


def test_docx_round_trip_preserves_chart_alt_metadata(tmp_path: Path, monkeypatch) -> None:
    facts = _facts()
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    doc = delivery.build_concise_word_document(facts)
    path = tmp_path / "report.docx"
    doc.save(path)

    reopened = Document(path)
    assert len(reopened.inline_shapes) == 4
    titles = []
    for shape in reopened.inline_shapes:
        titles.append(shape._inline.docPr.get("title"))
        assert "Source Data File" in shape._inline.docPr.get("descr")
    assert "Activity Mix by Source" in titles
    assert "Activity Trend" in titles


def test_customer_identity_keeps_legal_suffix_accounts_separate() -> None:
    team_data = {
        "Jordan Owner": {
            "subscriptions": pd.DataFrame(
                [
                    {"SUBSCRIPTION_ID": "SUB-I", "ACCOUNT_ID_C": "ACC-I", "BU_NAME": "Acme Inc"},
                    {"SUBSCRIPTION_ID": "SUB-L", "ACCOUNT_ID_C": "ACC-L", "BU_NAME": "Acme LLC"},
                ]
            ),
            # These ID-less records must join by exact normalized label.  The
            # old suffix-folding path put both rows into one risk profile.
            "action_plans": pd.DataFrame(
                [
                    {"ID": "AP-I", "BU_NAME": "ACME INC", "STATUS_C": "Open"},
                    {"ID": "AP-L", "BU_NAME": "Acme LLC", "STATUS_C": "Open"},
                ]
            ),
            "adoption_barriers": pd.DataFrame(),
            "customer_pulse": pd.DataFrame(),
            "tac_cases": pd.DataFrame(),
            "success_priorities": pd.DataFrame(),
        }
    }

    facts = delivery.build_report_facts(
        team_data,
        report_type="Leader",
        scope_type="team",
        scope_value="Suffix-safe team",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
    )

    assert facts["kpis"]["customers"] == 2
    assert set(facts["risk_profiles"]) == {"Acme Inc", "Acme LLC"}
    assert {
        customer: profile["components"]["action_plans"]["details"]["count"]
        for customer, profile in facts["risk_profiles"].items()
    } == {"Acme Inc": 1, "Acme LLC": 1}
    assert {row[0] for row in facts["account_summary_all"]} == {"Acme Inc", "Acme LLC"}


def test_customer_identity_prefers_stable_account_id_for_guarded_alias_join() -> None:
    team_data = {
        "Jordan Owner": {
            "subscriptions": pd.DataFrame(
                [
                    {
                        "SUBSCRIPTION_ID": "SUB-1",
                        "ACCOUNT_ID_C": "ACC-1",
                        "BU_NAME": "Example Incorporated",
                    }
                ]
            ),
            "action_plans": pd.DataFrame(
                [
                    {
                        "ID": "AP-1",
                        "ACCOUNT_ID_C": "ACC-1",
                        "BU_NAME": "Example Inc",
                        "STATUS_C": "Open",
                    },
                    {
                        "ID": "AP-2",
                        "BU_NAME": "EXAMPLE INC.",
                        "STATUS_C": "Open",
                    },
                ]
            ),
            "adoption_barriers": pd.DataFrame(),
            "customer_pulse": pd.DataFrame(),
            "tac_cases": pd.DataFrame(),
            "success_priorities": pd.DataFrame(),
        }
    }

    facts = delivery.build_report_facts(
        team_data,
        report_type="Leader",
        scope_type="customer",
        scope_value="Example Incorporated",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
    )

    assert facts["kpis"]["customers"] == 1
    profile = facts["risk_profiles"]["Example Incorporated"]
    assert profile["canonical_identity"]["account_ids"] == ("acc-1",)
    assert profile["components"]["action_plans"]["details"]["count"] == 2


def test_partition_uses_email_identity_for_duplicate_display_names() -> None:
    subscriptions = pd.DataFrame(
        [
            {
                "SUBSCRIPTION_ID": "SUB-1",
                "ACCOUNT_ID_C": "ACC-1",
                "BU_NAME": "Alpha Public",
                "CSSM_NAME": "Alex Rivera",
                "CSSM_EMAIL": "alex.one@example.com",
            },
            {
                "SUBSCRIPTION_ID": "SUB-2",
                "ACCOUNT_ID_C": "ACC-2",
                "BU_NAME": "Beta Public",
                "CSSM_NAME": "Alex Rivera",
                "CSSM_EMAIL": "alex.two@example.com",
            },
        ]
    )
    action_plans = pd.DataFrame(
        [
            {
                "ID": "AP-1",
                "ACCOUNT_ID_C": "ACC-1",
                "BU_NAME": "Alpha Public",
                "OWNER_NAME": "Alex Rivera",
                "STATUS_C": "Open",
            },
            {
                "ID": "AP-2",
                "ACCOUNT_ID_C": "ACC-2",
                "BU_NAME": "Beta Public",
                "OWNER_NAME": "Alex Rivera",
                "STATUS_C": "Open",
            },
        ]
    )

    partitioned = delivery.partition_portfolio_by_member(
        subscriptions=subscriptions,
        action_plans=action_plans,
        adoption_barriers=pd.DataFrame(),
        customer_pulse=pd.DataFrame(),
        tac_cases=pd.DataFrame(),
    )

    member_keys = sorted(key for key in partitioned if not key.startswith("__"))
    assert member_keys == [
        "Alex Rivera <alex.one@example.com>",
        "Alex Rivera <alex.two@example.com>",
    ]
    assert {
        key: partitioned[key]["action_plans"]["ID"].tolist()
        for key in member_keys
    } == {
        "Alex Rivera <alex.one@example.com>": ["AP-1"],
        "Alex Rivera <alex.two@example.com>": ["AP-2"],
    }

    facts = delivery.build_report_facts(
        partitioned,
        report_type="Comprehensive",
        scope_type="team",
        scope_value="Duplicate-name team",
        manager_name="Dana Manager",
        days=90,
        as_of=AS_OF,
    )
    assert facts["kpis"]["team_members"] == 2
    assert {row[0] for row in facts["member_summary_all"]} == set(member_keys)
