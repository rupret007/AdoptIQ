"""Round 124 / F9: Comprehensive partial-data banner scope-vs-load wording.

``executive_report_builder.ExecutiveReportBuilder.add_partial_data_warning_banner``
used to always say "failed to load" even when the only warning was a scope
exclusion (``tech_filter_scope_excluded``) -- a scope decision, not a load
failure.  This ports the R112/F5 branch into the comprehensive Word builder.

Made-with: Cursor.
"""

from docx import Document

from executive_report_builder import ExecutiveReportBuilder


def _body_text(rb: ExecutiveReportBuilder, tmp_path) -> str:
    target = tmp_path / "banner.docx"
    rb.save(str(target))
    doc = Document(str(target))
    return "\n".join(p.text for p in doc.paragraphs)


def test_all_scope_warnings_use_filtered_wording(tmp_path):
    rb = ExecutiveReportBuilder()
    rb.add_heading("Test", level=0)
    rb.add_partial_data_warning_banner(
        [
            {
                "dataset": "adoption_barriers",
                "error": "169 of 183 barriers excluded as non-Contact-Center",
                "kind": "tech_filter_scope_excluded",
            }
        ]
    )
    text = _body_text(rb, tmp_path)
    assert "Partial Data Warning" in text
    assert "filtered out by the requested scope" in text
    assert "failed to load" not in text


def test_load_failure_warning_keeps_failed_to_load_wording(tmp_path):
    rb = ExecutiveReportBuilder()
    rb.add_heading("Test", level=0)
    rb.add_partial_data_warning_banner(
        [
            {
                "dataset": "customer_pulse",
                "error": "ESA_C360_CS_TASK__C blocked by column policy",
                "kind": "schema_drift",
            }
        ]
    )
    text = _body_text(rb, tmp_path)
    assert "failed to load" in text
    assert "filtered out by the requested scope" not in text


def test_mixed_warnings_fall_back_to_load_failure_wording(tmp_path):
    rb = ExecutiveReportBuilder()
    rb.add_heading("Test", level=0)
    rb.add_partial_data_warning_banner(
        [
            {"dataset": "ab", "error": "scoped out", "kind": "tech_filter_scope_excluded"},
            {"dataset": "pulse", "error": "policy block", "kind": "schema_drift"},
        ]
    )
    text = _body_text(rb, tmp_path)
    # Not all-scope -> generic load-failure preamble.
    assert "failed to load" in text


def test_empty_input_renders_no_banner(tmp_path):
    rb = ExecutiveReportBuilder()
    rb.add_heading("Test", level=0)
    rb.add_partial_data_warning_banner(None)
    rb.add_partial_data_warning_banner([])
    text = _body_text(rb, tmp_path)
    assert "Partial Data Warning" not in text
