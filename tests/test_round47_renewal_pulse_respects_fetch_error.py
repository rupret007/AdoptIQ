"""Round 47 / R47-RP-PULSE-PARITY (F-RP-PULSE-DUAL-TRUTH) -- regression
test that the renewal Word writer honors ``df.attrs['fetch_error']``
on the customer pulse frame the same way the Excel writer does.

Build23 audit (run ``1777445600``) caught Word saying
``Total Customer Pulse Records: 186`` with 10 sample rows in the body,
while the matching Excel sheet ``Customer_Customer_Pulse`` was a
single-row Data Unavailable envelope (``schema_drift: missing slot(s)
rating on a non-empty result (186 row(s))``).  Word/Excel dual-truth
is a demo-blocking parity defect.

This test renders the renewal Word block in isolation with a pulse
frame that carries ``fetch_error`` in ``df.attrs`` and asserts:

1. The Word body does NOT cite a row count (``Total Customer Pulse
   Records: <n>`` is suppressed).
2. The Word body explicitly says the data is unavailable and points
   the reader at the Excel envelope.
3. The Word body still renders the heading so an operator scrolling
   the report doesn't think the section was silently skipped.
"""

from __future__ import annotations

import importlib
from unittest.mock import patch

import pandas as pd
import pytest


@pytest.fixture(scope="module")
def app_simple():
    return importlib.import_module("app_simple")


def _build_pulse_frame_with_fetch_error() -> pd.DataFrame:
    """Mirror the upstream condition the Build23 audit captured: a
    non-empty CSConsole pulse pull whose schema validation failed
    because the contract slot ``rating`` did not resolve to a source
    column.  ``data_source_validator`` writes the failure into
    ``df.attrs['fetch_error']`` / ``fetch_error_kind`` so the
    downstream writers can branch on a single signal.
    """

    df = pd.DataFrame(
        [
            {"BU_NAME": "ACME CORP US", "PULSE_RATING__C": "Green", "COMMENTS__C": "ok"},
            {"BU_NAME": "ACME CORP US", "PULSE_RATING__C": "Yellow", "COMMENTS__C": "watch"},
        ]
    )
    df.attrs["fetch_error"] = (
        "schema_drift:customer_pulse: missing slot(s) rating on a "
        "non-empty result (186 row(s))"
    )
    df.attrs["fetch_error_kind"] = "schema_drift"
    return df


def _read_docx_body(path: str) -> str:
    """Return the concatenated paragraph text of a .docx so the test
    can assert on the operator-visible narrative without depending on
    paragraph structure."""

    from docx import Document

    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def _read_docx_headings(path: str) -> list[str]:
    from docx import Document

    doc = Document(path)
    return [
        p.text
        for p in doc.paragraphs
        if p.style and p.style.name.startswith("Heading")
    ]


def test_renewal_word_pulse_section_honors_fetch_error(tmp_path, app_simple) -> None:
    """When the pulse frame carries ``fetch_error`` in ``attrs``, the
    renewal Word block must NOT cite the raw row count and MUST say
    'unavailable' with a reason and a pointer to the Excel envelope."""

    pulse_frame = _build_pulse_frame_with_fetch_error()

    renewal_analysis = {
        "customer_name": "ACME CORP US",
        "renewal_risk_score": 5.0,
        "renewal_risk_category": "HEALTHY",
        "renewal_recommendations": [],
        "adoption_barriers_count": 0,
        "support_cases_count": 0,
        "customer_analyses": {},
    }
    base_path = str(tmp_path / "renewal_test")
    out_path = app_simple._create_simple_renewal_report(
        base_path,
        "ACME CORP US",  # customer_name
        "All Contact Center",  # technology
        90,  # days
        renewal_analysis,  # renewal_analysis
        pd.DataFrame(),  # customer_ab
        pd.DataFrame(),  # customer_csone
        ext_bugs=[],
        ext_incidents=[],
        software_defects=None,
        psirt_vulns=None,
        portfolio_mode=False,
        all_customers=None,
        chart_paths=None,
        customer_action_plans=pd.DataFrame(),
        customer_customer_pulse=pulse_frame,
        customer_success_priorities=pd.DataFrame(),
    )
    assert out_path and isinstance(out_path, str), "renewal report writer returned no path"

    body_text = _read_docx_body(out_path)

    # The dual-truth count must be suppressed.
    assert "Total Customer Pulse Records: 2" not in body_text, (
        "Word renewal report cited a pulse row count for a fetch_error frame; "
        "this re-introduces the F-RP-PULSE-DUAL-TRUTH parity defect."
    )
    # The unavailable disclosure must be present and must reference
    # the Excel envelope so a reader sees both artifacts agree.
    assert "Customer Pulse data unavailable" in body_text, (
        "Word renewal report did not surface the Customer Pulse "
        "fetch_error in the operator-visible body."
    )
    assert "schema_drift" in body_text, (
        "Word renewal report did not name the fetch_error_kind so the "
        "operator can correlate against the Excel envelope reason."
    )
    assert "Customer_Customer_Pulse" in body_text, (
        "Word renewal report did not point the reader at the matching "
        "Excel sheet name -- breaks parity guidance."
    )
    # Heading still rendered so the section is not silently skipped.
    headings = _read_docx_headings(out_path)
    assert any("Customer Pulse" in h for h in headings), (
        "Word renewal report skipped the Customer Pulse heading entirely; "
        "operators may think pulse was not analyzed at all."
    )


def test_renewal_word_pulse_section_renders_normally_when_no_fetch_error(tmp_path, app_simple) -> None:
    """Negative control: the gate must NOT engage when the pulse frame
    is healthy, otherwise we would hide pulse data on every run."""

    healthy = pd.DataFrame(
        [
            {"BU_NAME": "ACME CORP US", "PULSE_RATING__C": "Green", "COMMENTS__C": "ok"},
        ]
    )
    # No attrs set -> healthy frame -> Word should cite the count.
    renewal_analysis = {
        "customer_name": "ACME CORP US",
        "renewal_risk_score": 5.0,
        "renewal_risk_category": "HEALTHY",
        "renewal_recommendations": [],
        "adoption_barriers_count": 0,
        "support_cases_count": 0,
        "customer_analyses": {},
    }
    base_path = str(tmp_path / "renewal_healthy")
    out_path = app_simple._create_simple_renewal_report(
        base_path,
        "ACME CORP US",
        "All Contact Center",
        90,
        renewal_analysis,
        pd.DataFrame(),
        pd.DataFrame(),
        ext_bugs=[],
        ext_incidents=[],
        software_defects=None,
        psirt_vulns=None,
        portfolio_mode=False,
        all_customers=None,
        chart_paths=None,
        customer_action_plans=pd.DataFrame(),
        customer_customer_pulse=healthy,
        customer_success_priorities=pd.DataFrame(),
    )
    body_text = _read_docx_body(out_path)
    assert "Total Customer Pulse Records: 1" in body_text, (
        "Healthy pulse frame must still render its row count; gate was "
        "too aggressive."
    )
    assert "Customer Pulse data unavailable" not in body_text, (
        "Healthy pulse frame must NOT trigger the unavailable disclosure."
    )
