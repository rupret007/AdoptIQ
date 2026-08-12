"""Round 146 Subscription report/source-data accuracy regressions."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Callable

import pandas as pd
from docx import Document
from openpyxl import load_workbook

import decision_report_delivery as delivery
import local_acceptance_lab as lab
import local_acceptance_runtime as runtime
from report_iteration_loop import (
    compare_kpi_parity,
    evaluate_report_quality,
    extract_docx_kpis,
    extract_xlsx_kpis,
)


def _document_text(path: Path) -> str:
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _run_subscription(
    tmp_path: Path,
    *,
    analysis_id: str,
    support_fetch: Callable[..., pd.DataFrame] | None = None,
    payload_mutator: Callable[[dict], dict] | None = None,
) -> tuple[Path, Path]:
    import app_simple

    installation = runtime.install_runtime_adapters(
        lab.build_scenario_bundle("healthy"), app_simple
    )
    original_support_fetch = app_simple.fetch_support_cases_snowflake
    original_subscription_fetch = app_simple.fetch_subscription_data
    if support_fetch is not None:
        app_simple.fetch_support_cases_snowflake = support_fetch
    if payload_mutator is not None:
        def _mutated_subscription_fetch(subscription_id: str, days: int = 90) -> dict:
            payload = original_subscription_fetch(subscription_id, days)
            return payload_mutator(payload)

        app_simple.fetch_subscription_data = _mutated_subscription_fetch

    replacements = {
        "_r81_resolve_report_output_dir": lambda *_args, **_kwargs: tmp_path,
        "save_analysis_status": lambda: None,
        "record_report_completion": lambda *_args, **_kwargs: None,
        "store_report_insights": lambda *_args, **_kwargs: None,
        "_r92_write_corpus_sidecars": lambda *_args, **_kwargs: None,
        "auto_audit_report": lambda *_args, **_kwargs: None,
    }
    originals = {name: getattr(app_simple, name) for name in replacements}
    for name, replacement in replacements.items():
        setattr(app_simple, name, replacement)

    app_simple.analysis_status[analysis_id] = {
        "subscription_id": "SUB-001",
        "days": 90,
        "report_type": "subscription",
        "start_time": "2026-08-03T21:00:00Z",
        "status": "initializing",
        "phase_timings": {},
        "completed_steps": [],
    }
    try:
        app_simple.run_subscription_analysis(analysis_id)
        status = dict(app_simple.analysis_status[analysis_id])
        assert status["status"] == "completed"
        return Path(status["word_report"]), Path(status["excel_report"])
    finally:
        app_simple.analysis_status.pop(analysis_id, None)
        app_simple.fetch_support_cases_snowflake = original_support_fetch
        app_simple.fetch_subscription_data = original_subscription_fetch
        installation.restore()
        for name, original in originals.items():
            setattr(app_simple, name, original)


def _assert_formula_free(path: Path) -> None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        formulas = [
            (sheet.title, cell.coordinate, cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
    finally:
        workbook.close()
    assert formulas == []


def test_subscription_source_data_headers_match_values_and_deep_dive_scope(
    tmp_path: Path,
) -> None:
    word_path, source_path = _run_subscription(
        tmp_path,
        analysis_id="sub_SUB-001_1460000100",
    )

    with pd.ExcelFile(source_path) as workbook:
        assert tuple(workbook.sheet_names) == delivery.SOURCE_DATA_SHEET_NAMES
        action_plans = pd.read_excel(workbook, sheet_name="Action_Plans")
        tac_cases = pd.read_excel(workbook, sheet_name="TAC_Cases")
        subscriptions = pd.read_excel(workbook, sheet_name="Subscriptions")
        risk_components = pd.read_excel(workbook, sheet_name="Risk_Components")
        report_info = pd.read_excel(workbook, sheet_name="Report_Info")
        lineage = pd.read_excel(workbook, sheet_name="Metric_Lineage")
        evidence_links = pd.read_excel(workbook, sheet_name="Evidence_Links")
        chart_data = pd.read_excel(workbook, sheet_name="Chart_Data")
        success_priorities = pd.read_excel(
            workbook, sheet_name="Success_Priorities"
        )

    # Regression for the pre-R146 shifted-row defect: the curated header must
    # describe the value beneath it, including the manager-action fields.
    assert {"Record_ID", "Customer Name", "Status", "Due Date"}.issubset(
        action_plans.columns
    )
    ap_001 = action_plans.set_index("Record_ID").loc["AP-001"]
    assert ap_001["Customer Name"] == "Acme Corporation"
    assert ap_001["Status"] == "Open"
    assert str(ap_001["Due Date"]).startswith("2026-08-01")
    assert ap_001["Next Action Owner"] == "Alex Rivera"

    source_subscriptions = subscriptions.loc[
        subscriptions["Legacy_Record_Type"].fillna("")
        != "Family-specific reported fact"
    ].reset_index(drop=True)
    reported_subscription_facts = subscriptions.loc[
        subscriptions["Legacy_Record_Type"].fillna("")
        == "Family-specific reported fact"
    ]
    assert len(source_subscriptions) == 1
    assert source_subscriptions.loc[0, "Subscription ID"] == "SUB-001"
    assert source_subscriptions.loc[0, "Account ID"] == "ACC-001"
    assert source_subscriptions.loc[0, "Customer Name"] == "Acme Corporation"
    assert not reported_subscription_facts.empty
    assert set(reported_subscription_facts["Metric_Key"].dropna()).issubset(
        set(evidence_links["Evidence_Key"].dropna())
    )

    assert tac_cases["SR Number"].tolist() == ["CASE-001"]
    assert tac_cases["Account ID"].tolist() == ["ACC-001"]
    assert tac_cases["Subscription Reference Id"].tolist() == ["SUB-001"]
    assert tac_cases["Priority (Normalized)"].tolist() == ["P1"]
    assert tac_cases["Case Status (Normalized)"].tolist() == ["Open"]
    assert tac_cases["Is Open"].tolist() == [True]
    support_component = risk_components.loc[
        risk_components["Component"].eq("support_cases")
    ].iloc[0]
    support_detail = json.loads(support_component["Component_Details_JSON"])
    assert support_detail["count"] == 1
    assert 0 <= support_component["Component_Score_0_100"] <= 100
    activity_component = risk_components.loc[
        risk_components["Component"].eq("activity_volume")
    ].iloc[0]
    activity_detail = json.loads(activity_component["Component_Details_JSON"])
    assert activity_detail["total_activity"] == (
        len(action_plans) + len(tac_cases) + 1 + 1
    )
    incident_component = risk_components.loc[
        risk_components["Component"].eq("incidents")
    ].iloc[0]
    incident_detail = json.loads(incident_component["Component_Details_JSON"])
    assert pd.isna(incident_component["Component_Score_0_100"])
    assert incident_detail["data_state"] == "missing"
    assert incident_detail["source_state"] == "unavailable"
    assert incident_detail["count"] is None

    # External incidents are not fetched by the Subscription path. The risk
    # distribution is therefore honestly withheld while the three charts
    # backed by complete selected-scope sources remain renderable.
    risk_chart = chart_data.loc[
        chart_data["Chart_ID"].eq("risk_distribution")
    ]
    assert not risk_chart.empty
    assert set(risk_chart["Source_State"]) == {"partial"}
    assert risk_chart["Value"].isna().all()

    info = dict(zip(report_info["Item"], report_info["Value"], strict=True))
    assert info["Report_Type"] == "Subscription"
    assert info["Scope_Type"] == "subscription"
    assert info["Scope_Value"] == "SUB-001"
    assert info["Technology"] == "Webex Calling"
    assert info["Source_State:TAC_Cases"] == "available"
    assert info["Partial_Data_Warning_Count"] >= 1
    assert len(str(info["Fact_Contract_SHA256"])) == 64
    tac_metric = lineage.set_index("Metric_Key").loc["kpi.tac_cases"]
    assert tac_metric["Metric_Value"] == 1
    tac_links = evidence_links.loc[
        evidence_links["Evidence_Key"].eq("kpi.tac_cases")
    ]
    assert tac_links["Record_ID"].tolist() == ["CASE-001"]
    assert tac_links["Source_Row_Number"].tolist() == [2]
    assert all(
        not str(column).startswith(("FIXTURE_", "LOCAL_ACCEPTANCE_"))
        for column in success_priorities.columns
    )

    word_text = _document_text(word_path)
    assert len(word_text.split()) < 1_500
    assert "Complete selected-scope records" in word_text
    assert "guarded offline test data" in word_text
    assert "The selected scope covers 1 customer." in word_text
    assert "team member" not in word_text.casefold()
    # Round 162.3 renders one exact criteria-scoped signal per applicable
    # source. The selected account's case is therefore visible and linked;
    # scope safety is pinned below with an explicit out-of-account fixture.
    assert "CASE-001" in word_text
    assert "CASE-OUTSIDE" not in word_text
    assert "TAC Cases\n1\nkpi.tac_cases" in word_text
    document = Document(word_path)
    assert len(document.inline_shapes) == 3
    assert all(
        any(
            paragraph.text.startswith("AdoptIQ v")
            for paragraph in section.footer.paragraphs
        )
        for section in document.sections
    )
    action_table = next(
        table
        for table in document.tables
        if [cell.text for cell in table.rows[0].cells][:3]
        == ["Record ID", "Account", "Owner"]
    )
    assert action_table.rows[1].cells[2].text == "Alex Rivera"
    legacy_workbook = word_path.with_suffix(".xlsx")
    assert legacy_workbook != source_path
    assert not legacy_workbook.exists()
    assert sorted(tmp_path.glob("*.xlsx")) == [source_path]
    parity = compare_kpi_parity(
        extract_docx_kpis(word_path),
        extract_xlsx_kpis(source_path),
        strict=True,
    )
    assert parity.passed, parity.details
    quality_payload, quality = evaluate_report_quality(
        word_path,
        source_path,
        scenario_key="g_subscription_analysis",
        strict=True,
    )
    assert quality.passed, quality_payload
    assert quality_payload["unbacked_metric_claim_count"] == 0
    assert all(
        unsafe not in (word_text + source_path.read_bytes().decode("latin1", "ignore"))
        for unsafe in ("fetch_error", "Traceback", "/Users/")
    )
    _assert_formula_free(source_path)


def test_subscription_withholds_out_of_account_rows_from_every_source(
    tmp_path: Path,
) -> None:
    def broaden_payload(payload: dict) -> dict:
        broadened = dict(payload)
        additions = {
            "adoption_barriers": {
                "ID": "AB-OUTSIDE",
                "ACCOUNT_ID_C": "ACC-OUTSIDE",
                "BU_NAME": "Outside Customer",
                "SUBJECT_C": "Outside barrier",
                "STATUS_C": "Open",
            },
            "action_plans": {
                "ID": "AP-OUTSIDE",
                "ACCOUNT_ID_C": "ACC-OUTSIDE",
                "BU_NAME": "Outside Customer",
                "SUBJECT_C": "Outside action",
                "STATUS_C": "Open",
            },
            "customer_pulse": {
                "ID": "CP-OUTSIDE",
                "ACCOUNT__C": "ACC-OUTSIDE",
                "BU_NAME": "Outside Customer",
                "PULSE_RATING__C": "Poor",
            },
            "success_priorities": {
                "ID": "SP-OUTSIDE",
                "ACCOUNT_ID_C": "ACC-OUTSIDE",
                "RELATED_CUSTOMER__C": "Outside Customer",
                "STATUS_C": "Active",
            },
        }
        for dataset, row in additions.items():
            broadened[dataset] = [*(payload.get(dataset) or []), row]
        return broadened

    def broad_support_cases(*_args, **_kwargs) -> pd.DataFrame:
        return pd.DataFrame(
            [
                {
                    "CASE_ID": "CASE-001",
                    "ACCOUNT_ID_C": "ACC-001",
                    "BU_NAME": "Acme Corporation",
                    "SUBJECT": "Registration authentication outage",
                    "SEVERITY": "P1",
                    "STATUS": "Open",
                    "DATE_OPENED": "2026-07-15T12:00:00Z",
                },
                {
                    "CASE_ID": "CASE-OUTSIDE",
                    "ACCOUNT_ID_C": "ACC-OUTSIDE",
                    "BU_NAME": "Outside Customer",
                    "SUBJECT": "Outside account case",
                    "SEVERITY": "P1",
                    "STATUS": "Open",
                    "DATE_OPENED": "2026-07-20T12:00:00Z",
                },
            ]
        )

    word_path, source_path = _run_subscription(
        tmp_path,
        analysis_id="sub_SUB-001_1460000102",
        support_fetch=broad_support_cases,
        payload_mutator=broaden_payload,
    )

    outside_ids = {
        "AB-OUTSIDE",
        "AP-OUTSIDE",
        "CP-OUTSIDE",
        "SP-OUTSIDE",
        "CASE-OUTSIDE",
    }
    with pd.ExcelFile(source_path) as workbook:
        for sheet in (
            "Adoption_Barriers",
            "Action_Plans",
            "Customer_Pulse",
            "Success_Priorities",
            "TAC_Cases",
        ):
            frame = pd.read_excel(workbook, sheet_name=sheet)
            public_values = set(frame.astype(str).to_numpy().ravel().tolist())
            assert not public_values & outside_ids, (sheet, public_values & outside_ids)
        tac_cases = pd.read_excel(workbook, sheet_name="TAC_Cases")
        report_info = pd.read_excel(workbook, sheet_name="Report_Info")

    assert tac_cases["SR Number"].tolist() == ["CASE-001"]
    info = dict(zip(report_info["Item"], report_info["Value"], strict=True))
    assert info["Source_State:TAC_Cases"] == "partial"
    assert info["Partial_Data_Warning_Count"] >= 5
    public_text = _document_text(word_path) + source_path.read_bytes().decode(
        "latin1", "ignore"
    )
    assert "CASE-001" in public_text
    for outside_id in outside_ids:
        assert outside_id not in public_text


def test_subscription_tac_unavailable_is_public_and_not_a_false_zero(
    tmp_path: Path,
) -> None:
    def unavailable_support_cases(*_args, **_kwargs) -> pd.DataFrame:
        frame = pd.DataFrame()
        frame.attrs["fetch_error"] = (
            "ProgrammingError: secret connection detail at /Users/internal/path"
        )
        frame.attrs["fetch_error_kind"] = "access_or_schema"
        return frame

    word_path, source_path = _run_subscription(
        tmp_path,
        analysis_id="sub_SUB-001_1460000101",
        support_fetch=unavailable_support_cases,
    )
    tac_cases = pd.read_excel(source_path, sheet_name="TAC_Cases")
    report_info = pd.read_excel(source_path, sheet_name="Report_Info")
    evidence_links = pd.read_excel(source_path, sheet_name="Evidence_Links")
    word_text = _document_text(word_path)
    public_text = "\n".join(
        [
            word_text,
            tac_cases.to_csv(index=False),
            report_info.to_csv(index=False),
        ]
    )

    info = dict(zip(report_info["Item"], report_info["Value"], strict=True))
    tac_links = evidence_links.loc[
        evidence_links["Evidence_Key"].eq("kpi.tac_cases")
    ]
    assert tac_cases.empty
    assert info["Source_State:TAC_Cases"] == "unavailable"
    assert info["Partial_Data_Warning_Count"] >= 2
    assert len(tac_links) == 1
    assert tac_links.iloc[0]["Evidence_Role"] == "unavailable_state"
    assert pd.isna(tac_links.iloc[0]["Source_Row_Number"])
    assert "unavailable data is not shown as zero" in word_text.casefold()
    assert "TAC Cases\nUnavailable" in word_text
    for unsafe in (
        "ProgrammingError",
        "access_or_schema",
        "fetch_error",
        "/Users/",
        "Traceback",
    ):
        assert unsafe not in public_text
    _assert_formula_free(source_path)
