"""Round 52 / Phase 5: harness KPI-extractor polish tests.

The strict live loop's parity gate kept reporting cosmetic-only
mismatches that masked real signal:

  * ``technology``: DOCX value was ``"All Contact Center | Analysis
    Period: 90 days"`` because the renewal report's header line packs
    multiple ``Label: Value`` pairs onto one row separated by `` | ``.
    The XLSX side has only ``"All Contact Center"``. The greedy
    paragraph-text regex consumed the trailing pair as part of the
    ``Technology`` value.

  * ``risk_score``: DOCX value was ``"0.7"`` and XLSX value was
    ``"0.7/10"``. The underlying score is identical -- the DOCX strips
    the ``/10`` suffix for presentation while the XLSX keeps it.

  * ``window_days`` and similar: one side writes ``"90"`` and the
    other writes ``"90 days"``.

Both fixes are guarded here so future refactors of the extractors
cannot silently regress and re-introduce the false-positive mismatches.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from report_iteration_loop import (
    _normalize_kpi_value,
    extract_docx_kpis,
)


def test_round52_paragraph_text_kpi_stops_at_pipe_separator(tmp_path: Path):
    """Renewal-style header line: Technology + Analysis Period packed on
    one row separated by ' | '.  The Technology value must NOT swallow
    the trailing 'Analysis Period: 90 days' suffix."""
    doc = Document()
    doc.add_paragraph("Renewal Risk Portfolio")
    doc.add_paragraph(
        "Technology: All Contact Center | Analysis Period: 90 days"
    )
    docx_path = tmp_path / "renewal_pipe_header.docx"
    doc.save(docx_path)

    kpis = extract_docx_kpis(docx_path)
    values = kpis.get("values", {})

    assert values.get("technology") == "All Contact Center", (
        f"technology must stop at the ' | ' separator; got {values.get('technology')!r}"
    )
    # Bonus: the trailing pair should also be captured under
    # ``window_days`` via the numeric paragraph regex.
    assert values.get("window_days") == "90", (
        f"window_days from 'Analysis Period: 90 days' must be captured "
        f"(numeric extractor path); got {values.get('window_days')!r}"
    )


def test_round52_paragraph_text_kpi_handles_no_pipe(tmp_path: Path):
    """Sanity guard: when there is no ' | ' separator, the regex must
    still capture the full single value."""
    doc = Document()
    doc.add_paragraph("Technology: All Contact Center")
    docx_path = tmp_path / "no_pipe.docx"
    doc.save(docx_path)

    kpis = extract_docx_kpis(docx_path)
    values = kpis.get("values", {})
    assert values.get("technology") == "All Contact Center"


def test_round52_normalize_value_strips_trailing_slash_max_suffix():
    """``"0.7/10"`` and ``"0.7"`` represent the same score.  Strip the
    ``/<max>`` presentational suffix so the parity gate does not fire."""
    assert _normalize_kpi_value("0.7/10") == "0.7"
    assert _normalize_kpi_value("8.5 / 10") == "8.5"
    # untouched when there is no suffix
    assert _normalize_kpi_value("0.7") == "0.7"


def test_round52_normalize_value_strips_trailing_days_unit():
    """``"90"`` and ``"90 days"`` represent the same window."""
    assert _normalize_kpi_value("90 days") == "90"
    assert _normalize_kpi_value("30 day") == "30"
    assert _normalize_kpi_value("90") == "90"


def test_round52_normalize_value_strips_trailing_direct_reports_unit():
    """Leader-report Team Size: ``"7 Direct Reports"`` vs ``"7"``."""
    assert _normalize_kpi_value("7 Direct Reports") == "7"
    assert _normalize_kpi_value("7 direct report") == "7"
    assert _normalize_kpi_value("7") == "7"


def test_round52_normalize_value_does_not_strip_substantive_text():
    """The suffix-strip MUST NOT eat substantive trailing words.
    Only `/<num>`, `days?`, `direct reports?` are allow-listed."""
    assert _normalize_kpi_value("All Contact Center") == "All Contact Center"
    assert _normalize_kpi_value("Brian Frazier") == "Brian Frazier"
    assert _normalize_kpi_value("High Risk") == "High Risk"


def test_round52_normalize_value_does_not_eat_pure_numeric_units_in_middle():
    """Defense in depth: a value like ``"3 of 10 stages"`` must not be
    truncated -- the ``/<num>`` strip is anchored to the END of the
    string only."""
    assert _normalize_kpi_value("3 of 10 stages") == "3 of 10 stages"


def test_round52_team_summary_skips_title_row_to_find_real_header(tmp_path: Path):
    """Real leader Team_Summary sheets have a TITLE row above the
    actual header row.  Before this fix the harness counted the
    header row itself as a team member (so a team of 11 reported
    team_members=12).  Detect the real header by scanning for
    the marker columns ("Team_Member" / "Num_Customers")."""
    import openpyxl
    from report_iteration_loop import _extract_team_summary_sheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Team_Summary"
    # Row 1 -- title row (no KPI markers).
    ws.append(["Team Summary - Brian Frazier Team Report"])
    # Row 2 -- the real header.
    ws.append([
        "Team_Member", "Num_Customers", "Num_Subscriptions",
        "Num_Action_Plans", "Num_Adoption_Barriers", "Num_TAC_Cases",
    ])
    # Rows 3-13 -- 11 team members.
    for i in range(11):
        ws.append([f"Member{i+1}", i + 1, 2, 5, 1, 3])

    xlsx_path = tmp_path / "leader_with_title.xlsx"
    wb.save(xlsx_path)

    # Re-open in read-only / data_only mode to mirror the harness path.
    wb2 = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws2 = wb2["Team_Summary"]
    values: dict[str, str] = {}
    _extract_team_summary_sheet(ws2, values)

    assert values.get("team_members") == "11", (
        f"team_summary handler must skip title row and count 11 team "
        f"members; got {values.get('team_members')!r}"
    )


def test_round52_team_summary_handles_no_title_row(tmp_path: Path):
    """Sanity guard: when there is no title row, the handler must still
    detect the header at row 0 and produce the correct team size."""
    import openpyxl
    from report_iteration_loop import _extract_team_summary_sheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Team_Summary"
    ws.append([
        "Team_Member", "Num_Customers", "Num_TAC_Cases",
    ])
    for i in range(7):
        ws.append([f"Member{i+1}", i + 1, 3])

    xlsx_path = tmp_path / "leader_no_title.xlsx"
    wb.save(xlsx_path)
    wb2 = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws2 = wb2["Team_Summary"]
    values: dict[str, str] = {}
    _extract_team_summary_sheet(ws2, values)

    assert values.get("team_members") == "7"


def test_round52_escalated_support_cases_canonical_distinct_from_support_cases():
    """``Escalated Support Cases`` is a STRICT SUBSET of total
    ``Support Cases`` and must not collide on the same canonical key.
    Otherwise the strict parity gate fires when DOCX renders the total
    (e.g. 293) and the XLSX Executive_Dashboard renders only the
    escalation count (e.g. 2)."""
    from report_iteration_loop import _canonical_kpi_label

    assert _canonical_kpi_label("Escalated Support Cases") == "escalated_support_cases"
    assert _canonical_kpi_label("Support Cases") == "support_cases"
    assert _canonical_kpi_label("Total Support Cases") == "support_cases"


def test_round52_data_sources_table_does_not_pollute_count_canonicals(tmp_path: Path):
    """Leader Data Sources table looks like::

        Action Plans                  | CSConsole
        Adoption Barriers             | CSConsole / Snowflake C360_CS_TASK_C_VW
        Customer Pulse                | CSConsole

    The right-hand column is the SOURCE LABEL, not a count.  Pre-Phase
    5 the harness canonicalised those rows as
    ``action_plans=CSConsole`` etc., guaranteeing a parity mismatch
    against the XLSX which carries the actual row count.  The
    numeric-required guard must drop those rows."""
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "Data Source"
    table.rows[0].cells[1].text = "Provider"
    pairs = (
        ("Action Plans", "CSConsole"),
        ("Adoption Barriers", "CSConsole / Snowflake C360_CS_TASK_C_VW"),
        ("Customer Pulse", "CSConsole"),
    )
    for idx, (label, value) in enumerate(pairs, start=1):
        table.rows[idx].cells[0].text = label
        table.rows[idx].cells[1].text = value
    docx_path = tmp_path / "leader_data_sources.docx"
    doc.save(docx_path)

    kpis = extract_docx_kpis(docx_path)
    values = kpis.get("values", {})
    for forbidden in ("action_plans", "adoption_barriers", "customer_pulse"):
        assert forbidden not in values, (
            f"{forbidden!r} must NOT be canonicalised from a non-numeric "
            f"source label; got values={values!r}"
        )


def test_round52_text_valued_canonicals_still_extracted(tmp_path: Path):
    """Sanity guard: text-valued canonicals (manager / technology /
    risk_category) must still be extracted from text cells."""
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    pairs = (
        ("Manager", "Brian Frazier"),
        ("Technology", "All Contact Center"),
        ("Risk Category", "High"),
    )
    for idx, (label, value) in enumerate(pairs):
        table.rows[idx].cells[0].text = label
        table.rows[idx].cells[1].text = value
    docx_path = tmp_path / "text_canonicals.docx"
    doc.save(docx_path)

    kpis = extract_docx_kpis(docx_path)
    values = kpis.get("values", {})
    assert values.get("manager") == "Brian Frazier"
    assert values.get("technology") == "All Contact Center"
    assert values.get("risk_category") == "High"


def test_round52_numeric_kpi_value_predicate_handles_common_shapes():
    """The numeric guard must accept the common shapes the report
    builders produce (counts, percentages, currency, decimals) but
    reject text-only labels."""
    from report_iteration_loop import _is_numeric_kpi_value

    for value in ("0", "12", "1,234", "14.8", "100%", "$5,000", "-3", "0.7"):
        assert _is_numeric_kpi_value(value), f"expected {value!r} to count as numeric"
    for value in ("CSConsole", "All Contact Center", "Brian Frazier", "", None, "n/a"):
        assert not _is_numeric_kpi_value(value), f"expected {value!r} to be rejected"


def test_round52_multi_row_breakdown_table_uses_total_footer(tmp_path: Path):
    """Round 52 / accuracy-fix-loop regression.

    The leader DOCX's "Team Member Activity Breakdown" is a 13x6 table
    with header [Team Member, Action Plans, Adoption Barriers, Customer
    Pulse, TAC Cases, BEMS], one row per CSSM, and a TOTAL footer.  The
    pre-fix harness paired headers with row 1 (the FIRST team member's
    counts) instead of the TOTAL row, so it extracted the per-person
    metrics (action_plans=50) while the XLSX side reported the team
    portfolio totals (action_plans=354). Pin the new behavior:
    multi-row 3+ column tables must prefer a "TOTAL"/"TEAM TOTAL" row
    when one is present, and fall through (no auto-pair) when none is.
    """
    doc = Document()
    table = doc.add_table(rows=4, cols=6)
    headers = (
        "Team Member",
        "Action Plans",
        "Adoption Barriers",
        "Customer Pulse",
        "TAC Cases",
        "BEMS",
    )
    for col, label in enumerate(headers):
        table.rows[0].cells[col].text = label
    rows = (
        ("Angelica Hernandez Becerra", "50", "5", "8", "14", "1"),
        ("Arpit Patel", "26", "2", "1", "27", "4"),
        ("TOTAL", "354", "69", "87", "381", "81"),
    )
    for ri, row in enumerate(rows, start=1):
        for col, val in enumerate(row):
            table.rows[ri].cells[col].text = val
    docx_path = tmp_path / "team_breakdown_total.docx"
    doc.save(docx_path)

    kpis = extract_docx_kpis(docx_path)
    values = kpis.get("values", {})
    assert values.get("action_plans") == "354", values
    assert values.get("adoption_barriers") == "69", values
    assert values.get("customer_pulse") == "87", values
    assert values.get("support_cases") == "381", values
    assert values.get("bems") == "81", values


def test_round52_multi_row_breakdown_table_without_total_skips_auto_pair(
    tmp_path: Path,
):
    """If a multi-row 3+ column table has no TOTAL row, the harness
    must NOT pair headers with row 1.  Otherwise per-person counts
    masquerade as portfolio totals (the original leader bug)."""
    doc = Document()
    table = doc.add_table(rows=3, cols=4)
    for col, label in enumerate(
        ("Team Member", "Action Plans", "Adoption Barriers", "TAC Cases")
    ):
        table.rows[0].cells[col].text = label
    for col, val in enumerate(
        ("Angelica Hernandez Becerra", "50", "5", "14")
    ):
        table.rows[1].cells[col].text = val
    for col, val in enumerate(("Arpit Patel", "26", "2", "27")):
        table.rows[2].cells[col].text = val
    docx_path = tmp_path / "team_breakdown_no_total.docx"
    doc.save(docx_path)

    kpis = extract_docx_kpis(docx_path)
    values = kpis.get("values", {})
    for forbidden in ("action_plans", "adoption_barriers", "support_cases"):
        assert forbidden not in values, (
            f"{forbidden!r} must NOT be auto-paired with the first team "
            f"member's row when no TOTAL footer exists; got {values!r}"
        )


def test_round52_total_label_aliases_canonicalize(tmp_path: Path):
    """The leader executive-summary bullets render team-wide totals as
    `Total Action Plans: 354`, `Total TAC Cases: 381`, `Total BEMS
    Escalations: 81`, `Total Customer Pulse records: 87`. All four
    label variants must canonicalize so the parity gate has a
    paragraph-side source of truth even when the table heuristic
    fails."""
    from report_iteration_loop import _canonical_kpi_label

    assert _canonical_kpi_label("Total Action Plans") == "action_plans"
    assert _canonical_kpi_label("Total TAC Cases") == "support_cases"
    assert _canonical_kpi_label("Total BEMS Escalations") == "bems"
    assert _canonical_kpi_label("Total Customer Pulse") == "customer_pulse"
    assert _canonical_kpi_label("Total Customer Pulse records") == "customer_pulse"
    assert _canonical_kpi_label("Total Adoption Barriers") == "adoption_barriers"


def test_round52_two_row_three_column_table_still_pairs(tmp_path: Path):
    """Sanity guard: the 2-row (header + single data row) 3+ column
    pattern (used by comprehensive Title Page metrics) must continue
    to pair row 0 with row 1.  Only multi-row tables require a totals
    footer to opt into auto-pair."""
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Total Customers"
    table.rows[0].cells[1].text = "Support Cases"
    table.rows[0].cells[2].text = "Critical (P1)"
    table.rows[1].cells[0].text = "52"
    table.rows[1].cells[1].text = "176"
    table.rows[1].cells[2].text = "5"
    docx_path = tmp_path / "title_page.docx"
    doc.save(docx_path)

    kpis = extract_docx_kpis(docx_path)
    values = kpis.get("values", {})
    assert values.get("total_customers") == "52"
    assert values.get("support_cases") == "176"
    assert values.get("critical_cases") == "5"


def test_round52_corpus_context_paragraph_does_not_leak_into_technology(
    tmp_path: Path,
):
    """Round 52 / ship regression.

    ``report_corpus_context.py`` renders per-customer narrative blocks
    whose lines start with ``Technology: <tech>; Observed: <iso> ->
    <iso>; Prior occurrences: N`` (and sometimes ``; Sentiment
    direction: ...``).  Pre-fix, these lines satisfied the harness's
    ``_PARAGRAPH_KPI_TEXT_RE`` and the value greedily consumed the
    trailing telemetry, polluting the canonical ``technology`` KPI
    with a string like ``"Cloud and Hybrid Products; Observed: ...;
    Prior occurrences: 879"``.  Pin the new behavior: the corpus-
    context line must be REJECTED so the only ``technology`` source
    is the report-scope label (not an embedded narrative tag).
    """
    doc = Document()
    doc.add_paragraph(
        "Technology: Cloud and Hybrid Products; Observed: 2026-04-29T14:16:44Z "
        "-> 2026-04-29T18:33:33Z; Prior occurrences: 879"
    )
    doc.add_paragraph(
        "Technology: Other/Unknown; Observed: 2026-04-29T14:24:34Z -> "
        "2026-04-29T18:33:33Z; Prior occurrences: 162; Sentiment direction: flat"
    )
    docx_path = tmp_path / "corpus_context_only.docx"
    doc.save(docx_path)

    kpis = extract_docx_kpis(docx_path)
    values = kpis.get("values", {})
    assert "technology" not in values, (
        "corpus-context paragraphs (Observed:/Prior occurrences:/Sentiment "
        f"direction:) must NOT leak into the technology canonical; "
        f"got values={values!r}"
    )


def test_round52_legitimate_technology_paragraph_still_extracted(tmp_path: Path):
    """Sanity guard: a clean ``Technology: All Contact Center``
    paragraph (no corpus-context markers) must still extract."""
    doc = Document()
    doc.add_paragraph("Technology: All Contact Center")
    docx_path = tmp_path / "clean_technology.docx"
    doc.save(docx_path)

    kpis = extract_docx_kpis(docx_path)
    values = kpis.get("values", {})
    assert values.get("technology") == "All Contact Center", values


def test_round52_corpus_context_value_class_stops_at_semicolon(tmp_path: Path):
    """Defense in depth: even if a corpus-context paragraph somehow
    bypasses the ``_looks_like_corpus_context_line`` guard (e.g. the
    tells are renamed in a future round), the value class must STILL
    stop at ``;`` so the bleed cannot exceed the technology name
    itself.  This pins the value-class tightening separately from the
    corpus-context guard."""
    from report_iteration_loop import _PARAGRAPH_KPI_TEXT_RE

    text = "Technology: Cloud and Hybrid Products; Some Trailing Junk"
    m = _PARAGRAPH_KPI_TEXT_RE.match(text)
    assert m is not None, "regex should still match the start"
    assert m.group("value").strip() == "Cloud and Hybrid Products", (
        f"value must stop at the first ';'; got {m.group('value')!r}"
    )


def test_round52_per_scenario_docx_threshold_overrides_lower_comprehensive():
    """Round 52 / ship regression -- updated for Round 52.1 contract.

    The comprehensive report embeds an AI-generated insights section
    that regenerates run-to-run (paragraph counts shift by ~200 and
    diff produces 1614+ narrative-only changes between two runs against
    the same scope).  A single global text-similarity floor cannot
    cover both narrative-templated and narrative-generated reports, so
    comprehensive opts into a per-scenario lower text-floor.

    Round 52.1 adjustment: ``effective_docx_thresholds`` now returns a
    3-tuple including the table-only numeric drift gate.  Comprehensive
    also relaxes its overall numeric gate to 0.55 (informational) and
    relies on the table-only gate as the new binding signal for real
    Snowflake data drift.  Compact / renewal / leader keep all three
    floors at the config defaults.
    """
    from report_iteration_loop import (
        SCENARIO_DOCX_THRESHOLD_OVERRIDES,
        RunnerConfig,
        effective_docx_thresholds,
    )

    cfg = RunnerConfig(
        base_url="",
        downloads_dir=Path("."),
        iterations=1,
        poll_interval_seconds=1.0,
        scenario_timeout_seconds=60,
        run_id="x",
        stop_on_failure=False,
        scenario_keys=[],
        baseline_mode="manifest",
        min_docx_similarity=0.55,
        min_sheet_overlap=0.85,
        min_header_similarity=0.8,
        min_docx_chars=200,
        strict=True,
        min_docx_numeric_similarity=0.80,
        min_docx_table_numeric_similarity=0.95,
        max_xlsx_row_delta_ratio=0.2,
        max_xlsx_row_delta_abs=25,
    )

    text, numeric, table_numeric = effective_docx_thresholds(cfg, "comprehensive")
    assert text == 0.40, f"comprehensive must use the text override; got {text!r}"
    # Round 52.1: comprehensive numeric is now overridden to 0.55
    # (informational); the binding accuracy signal is the table-only gate.
    assert numeric == 0.55, (
        f"comprehensive numeric must use the Round 52.1 override; got {numeric!r}"
    )
    assert table_numeric == 0.95, (
        "comprehensive table-only numeric must inherit the config default"
        f"; got {table_numeric!r}"
    )

    for s in ("compact", "renewal", "leader"):
        text, numeric, table_numeric = effective_docx_thresholds(cfg, s)
        assert text == 0.55, f"{s} must use the global text floor; got {text!r}"
        assert numeric == 0.80, f"{s} numeric must remain at the config floor"
        assert table_numeric == 0.95, (
            f"{s} table-only numeric must remain at the config floor"
        )

    assert "comprehensive" in SCENARIO_DOCX_THRESHOLD_OVERRIDES
    assert "compact" not in SCENARIO_DOCX_THRESHOLD_OVERRIDES
    assert "renewal" not in SCENARIO_DOCX_THRESHOLD_OVERRIDES
    assert "leader" not in SCENARIO_DOCX_THRESHOLD_OVERRIDES
