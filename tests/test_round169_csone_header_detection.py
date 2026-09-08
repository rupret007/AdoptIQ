"""Round 169 / P2: CSOne header detection skips title rows and malformed placeholders."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from adoptiq_backend import load_csone_excel


def _write_misaligned_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Export"
    sheet.append(["Synthetic CSOne export title", None, None, None])
    sheet.append(["Unnamed: 0", "Unnamed: 1", "Unnamed: 2", "Unnamed: 3"])
    sheet.append(["700000001", "Synthetic case title", "P2", "Fixture Customer"])
    sheet.append(["700000002", "Second case", "P3", "Fixture Customer"])
    wb.save(path)
    wb.close()


def _write_real_header_after_title_rows(path: Path) -> None:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "TAC"
    sheet.append(["Weekly export", "Do not use as header", ""])
    sheet.append(["", "", ""])
    sheet.append(["SR Number", "Title", "Severity", "Customer Name"])
    sheet.append(["700000010", "Resolved header row", "P3", "Fixture Org"])
    wb.save(path)
    wb.close()


def test_misaligned_unnamed_header_row_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "misaligned_csone.xlsx"
    _write_misaligned_workbook(path)

    loaded = load_csone_excel(path)

    assert loaded.empty


def test_header_candidate_skips_title_rows_and_finds_case_columns(tmp_path: Path) -> None:
    path = tmp_path / "offset_header_csone.xlsx"
    _write_real_header_after_title_rows(path)

    loaded = load_csone_excel(path)

    assert len(loaded) == 1
    assert "SR Number" in loaded.columns
    assert loaded.iloc[0]["SR Number"] == "700000010"


def test_r169_header_helpers_are_wired_in_backend() -> None:
    import adoptiq_backend as backend

    source = Path(backend.__file__).read_text(encoding="utf-8")
    assert "_r169_header_candidate" in source
    assert "_r169_headers_are_malformed" in source
