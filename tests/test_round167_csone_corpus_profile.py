"""Round 167 metadata-only profiling for local CSOne workbook corpora."""

from __future__ import annotations

import json

import openpyxl

from adoptiq_backend import load_csone_excel
from scripts import profile_csone_corpus as profiler


def _write_csone(path, rows, *, include_warning=True) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Enhanced Premium Collab"
    if include_warning:
        sheet.cell(row=1, column=1, value="Warning:")
        sheet.cell(row=1, column=2, value="Truncated export notice")
        header_row = 15
    else:
        header_row = 1
    headers = [
        "Customer Name: Customer Name",
        "Subscription Reference Id",
        "Severity",
        "SR Number",
        "Case Number",
        "Case Status",
        "Transaction ID",
        "Date/Time Opened",
        "Problem Description",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column, value=header)
    for row_number, values in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=column, value=value)
    workbook.save(path)


def test_profile_is_metadata_only_and_counts_placeholders(tmp_path) -> None:
    path = tmp_path / "AdoptIQ Enhanced Premium Collab Summary-2026-08-11-07-00-08.xlsx"
    _write_csone(
        path,
        [
            ["Private Customer", "SUB-1", "P1", "SR-1", "CASE-1", "Open", "", "2026-08-01", "Private narrative"],
            ["Private Customer", "SUB-1", "Unknown", "SR-1", "CASE-1", "Open", "BEMS-1", "2026-08-02", "Undefined"],
        ],
    )

    profile = profiler.profile_workbook(path)
    summary = profiler.summarize_profiles([profile])
    serialized = json.dumps(summary, sort_keys=True)

    assert summary["workbook_count"] == 1
    assert summary["header_row_distribution"] == {"15": 1}
    assert summary["columns"]["Severity"]["placeholder_cells"] == 1
    assert summary["columns"]["Problem Description"]["placeholder_cells"] == 1
    assert summary["identifier_quality"]["SR Number"]["duplicate"] == 1
    assert "Private Customer" not in serialized
    assert "Private narrative" not in serialized
    assert "SR-1" not in serialized


def test_cli_profiles_multiple_schema_shapes_without_copying_rows(tmp_path, capsys) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    _write_csone(
        corpus / "AdoptIQ Enhanced Premium Collab Summary-2026-08-10-07-00-08.xlsx",
        [["Customer A", "SUB-1", "P2", "SR-1", "CASE-1", "Open", "", "2026-08-01", "Text"]],
    )
    _write_csone(
        corpus / "AdoptIQ Enhanced Premium Collab Summary-2026-08-11-07-00-08.xlsx",
        [["Customer B", "SUB-2", "P3", "SR-2", "CASE-2", "Closed", "", "2026-08-02", "Text"]],
        include_warning=False,
    )
    summary_path = tmp_path / "summary.json"

    assert profiler.main(
        ["--input-dir", str(corpus), "--summary", str(summary_path)]
    ) == 0
    payload = json.loads(summary_path.read_text(encoding="utf-8"))

    assert payload["workbook_count"] == 2
    assert payload["total_profiled_rows"] == 2
    assert payload["header_row_distribution"] == {"1": 1, "15": 1}
    assert payload["source_rows_exported"] is False
    assert "Customer A" not in summary_path.read_text(encoding="utf-8")
    assert "workbook_count" in capsys.readouterr().out


def test_production_loader_excludes_footer_rows_but_retains_substantive_missing_id(
    tmp_path,
) -> None:
    path = tmp_path / "corpus-shaped.xlsx"
    _write_csone(
        path,
        [
            [
                "Customer A",
                "SUB-1",
                "P2",
                "SR-1",
                "CASE-1",
                "Open",
                "",
                "2026-08-01",
                "Case narrative",
            ],
            [
                "Customer B",
                "SUB-2",
                "P1",
                "",
                "",
                "Open",
                "",
                "2026-08-02",
                "Substantive missing-ID case",
            ],
            ["Export summary", "", "", "", "", "", "", "", ""],
        ],
    )

    loaded = load_csone_excel(path)

    assert len(loaded) == 2
    assert loaded.attrs["excluded_non_record_rows"] == 1
    assert any(
        item.get("kind") == "non_record_rows_excluded"
        for item in loaded.attrs.get("partial_data_warnings", [])
    )
