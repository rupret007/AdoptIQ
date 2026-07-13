"""Regression coverage for unavailable risk scores across report surfaces."""

from __future__ import annotations

import inspect
from pathlib import Path

import pandas as pd
import pytest

import canonical_metrics as cm


def _document_text(document) -> str:
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def test_portfolio_metrics_unknown_bucket_reconciles_all_profiles() -> None:
    profiles = {
        "Critical": {"risk_band": "CRITICAL", "risk_score_0_100": 80.0},
        "Healthy": {"risk_band": "HEALTHY", "risk_score_0_100": 0.0},
        "Explicit unknown": {"risk_band": "UNKNOWN", "risk_score_0_100": None},
        "Missing score": {},
        "Non-finite": {"risk_score_0_100": float("nan")},
        "Malformed": None,
        "Infinite": {"risk_score_0_100": float("inf")},
        "Stale unknown": {"risk_band": "UNKNOWN", "risk_score_0_100": 100.0},
        "Declared healthy": {"risk_band": "HEALTHY", "risk_score_0_100": 100.0},
        "Unavailable state": {
            "risk_assessment_state": "UNAVAILABLE",
            "risk_score_0_100": 100.0,
        },
        "Unavailable band": {
            "risk_band": "UNAVAILABLE",
            "risk_score_0_100": 100.0,
        },
        "Insufficient band": {
            "risk_band": "INSUFFICIENT_EVIDENCE",
            "risk_score_0_100": 100.0,
        },
        "N/A band": {"risk_band": "N/A", "risk_score_0_100": 100.0},
    }

    payload = cm.build_portfolio_metrics(
        ab_df=pd.DataFrame(),
        csone_df=pd.DataFrame(),
        risk_profiles=profiles,
    )

    assert payload["unknown_risk_customers"] == 10
    assert payload["healthy_customers"] == 2
    assert payload["scored_risk_customers"] == 3
    assert payload["risk_profile_customers"] == len(profiles)
    assert sum(payload["risk_band_counts"].values()) == len(profiles)
    assert cm.compute_high_risk_count(profiles) == 1
    assert cm.compute_high_risk_count(
        profiles, scale=cm.RISK_SCALE_0_TO_10
    ) == 1
    assert not cm.is_high_risk_profile(profiles["Stale unknown"])
    assert not cm.is_high_risk_profile(profiles["Declared healthy"])
    assert not cm.is_high_risk_profile(profiles["Unavailable state"])
    assert not cm.is_high_risk_profile(profiles["Unavailable band"])
    assert not cm.is_high_risk_profile(profiles["Insufficient band"])
    assert not cm.is_high_risk_profile(profiles["N/A band"])

    legacy_red = {"Legacy": {"score": 5.4, "color": "Red"}}
    legacy_payload = cm.build_portfolio_metrics(
        ab_df=pd.DataFrame(),
        csone_df=pd.DataFrame(),
        risk_profiles=legacy_red,
        risk_scale=cm.RISK_SCALE_0_TO_10,
    )
    assert legacy_payload["high_risk_customers"] == 1
    assert legacy_payload["high_only_risk_customers"] == 1
    assert legacy_payload["unknown_risk_customers"] == 0


def test_renewal_portfolio_rollup_excludes_unknown_scores() -> None:
    from app_simple import (
        _r124_deterministic_grade_line,
        _roll_up_portfolio_renewal_risk,
    )

    mixed = _roll_up_portfolio_renewal_risk(
        {
            "High": {"renewal_risk_score": 60.0},
            "Healthy": {"renewal_risk_score": 0.0},
            "Unavailable": {
                "renewal_risk_score": None,
                "renewal_risk_category": "UNKNOWN",
            },
        }
    )
    assert mixed["average_risk_score"] == pytest.approx(30.0)
    assert mixed["scored_customers"] == 2
    assert mixed["unknown_risk_customers"] == ["Unavailable"]
    assert mixed["high_risk_customers"] == ["High"]
    assert mixed["low_risk_customers"] == ["Healthy"]

    unavailable = _roll_up_portfolio_renewal_risk(
        {
            "A": {"renewal_risk_score": None},
            "B": {"renewal_risk_score": float("nan")},
        }
    )
    assert unavailable["average_risk_score"] is None
    assert unavailable["risk_category"] == "UNKNOWN"
    assert unavailable["scored_customers"] == 0
    assert unavailable["low_risk_customers"] == []

    grade_line = _r124_deterministic_grade_line(
        {"risk_score_0_10": None, "risk_band": "UNKNOWN"}
    )
    assert grade_line is not None
    assert "canonical risk N/A" in grade_line
    assert "0.0/10" not in grade_line


def test_renewal_customer_slicing_preserves_fetch_failure_metadata() -> None:
    from app_simple import (
        _calculate_simple_renewal_risk,
        _r98_slice_customer_frame,
        run_customer_renewal_analysis,
    )

    failed = pd.DataFrame()
    failed.attrs.update(
        {
            "fetch_error": "upstream timeout",
            "fetch_error_kind": "timeout",
            "fetch_error_dataset": "synthetic",
        }
    )

    sliced = _r98_slice_customer_frame(failed, "Acme", ("customer_name",))
    assert sliced is not None
    assert sliced.empty
    assert sliced.attrs["fetch_error"] == "upstream timeout"
    assert _r98_slice_customer_frame(None, "Acme", ("customer_name",)) is None

    result = _calculate_simple_renewal_risk(
        customer_name="Acme",
        customer_ab=failed.copy(),
        customer_csone=failed.copy(),
        team_subs_df=failed.copy(),
        days=90,
        ext_incidents=None,
        customer_pulse=failed.copy(),
        customer_action_plans=failed.copy(),
    )
    assert result["renewal_risk_score"] is None
    assert result["renewal_risk_category"] == "UNKNOWN"
    assert result["service_incidents_count"] is None
    assert result["high_impact_incidents_count"] is None

    orchestration_source = inspect.getsource(run_customer_renewal_analysis)
    assert "ext_incidents = None" in orchestration_source
    assert "Incident correlation and risk evidence are unavailable." in orchestration_source
    assert "for _external_warning in _renewal_external_warnings" in orchestration_source


def test_renewal_word_renders_unavailable_score_as_na(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    from app_simple import _create_simple_renewal_report

    output = _create_simple_renewal_report(
        base_path=str(tmp_path / "unknown-renewal"),
        customer_name="Acme",
        technology="Webex",
        days=90,
        renewal_analysis={
            "renewal_risk_score": None,
            "renewal_risk_score_10": None,
            "renewal_risk_category": "UNKNOWN",
            "key_findings": ["Risk evidence could not be loaded."],
            "risk_factors": [],
            "recommendations": ["Validate source coverage."],
            "adoption_barriers_count": 0,
            "support_cases_count": 0,
            "bems_escalations_count": 0,
            "support_cases_from_snowflake": False,
        },
        customer_ab=pd.DataFrame(),
        customer_csone=pd.DataFrame(),
        portfolio_mode=False,
    )

    text = _document_text(Document(output))
    assert "N/A (UNKNOWN; insufficient evidence)" in text
    assert "Renewal Risk Score: N/A" in text
    assert "0.0/10  (UNKNOWN" not in text
    assert "0.0/10 (UNKNOWN" not in text


def test_compact_sections_do_not_treat_unknown_as_standard_monitoring() -> None:
    from compact_report_formatter import CompactReportFormatter

    formatter = CompactReportFormatter()
    formatter.add_portfolio_health_dashboard(
        pd.DataFrame(),
        pd.DataFrame(),
        {
            "overall_risk_score": None,
            "critical_adoption_barriers": 0,
            "escalated_cases": 0,
        },
    )
    formatter.add_predictive_risk_section(
        pd.DataFrame(),
        pd.DataFrame(),
        {
            "Acme": {
                "score": None,
                "risk_score_0_100": None,
                "risk_band": "UNKNOWN",
                "color": "Gray",
            }
        },
    )
    formatter.add_high_risk_customers(
        pd.DataFrame(),
        pd.DataFrame(),
        {
            "Acme": {
                "score": None,
                "risk_score_0_100": None,
                "risk_band": "UNKNOWN",
                "color": "Gray",
                "category": "Risk Unavailable - Validate Evidence",
            }
        },
    )

    text = _document_text(formatter.doc)
    assert "Renewal Risk\nUnknown" in text
    assert "High-Risk Customers: N/A - customer risk evidence is unavailable" in text
    assert "do not infer healthy status from unavailable scores" in text
    assert "UNKNOWN - Risk Evidence Unavailable" in text
    assert "Acme" in text
    assert "GRAY - No Renewal Risk" not in text


def test_compact_and_wxcc_slices_preserve_source_failure_state(monkeypatch) -> None:
    import compact_report_formatter as compact
    from wxcc_health_input_exporter import slice_df_by_customer

    failed = pd.DataFrame()
    failed.attrs["fetch_error"] = "synthetic upstream failure"
    captured = []

    def fake_profile(**kwargs):
        captured.append(kwargs)
        return {
            "risk_score_0_10": None,
            "risk_score_0_100": None,
            "risk_band": "UNKNOWN",
            "risk_factors": [],
            "components": {
                "adoption_barriers": {"details": {"aging_open_count": 0}}
            },
        }

    monkeypatch.setattr(compact, "compute_customer_risk_profile", fake_profile)
    risk_data = compact.calculate_renewal_risk_scores(
        pd.DataFrame({"customer_name": ["Acme"]}),
        pd.DataFrame(),
        pulse_df=failed.copy(),
        action_plans_df=failed.copy(),
        subs_df=failed.copy(),
        ext_incidents=None,
    )

    assert len(captured) == 1
    assert risk_data["Acme"]["risk_band"] == "UNKNOWN"
    for key in ("customer_pulse", "customer_action_plans", "customer_subs"):
        assert captured[0][key].attrs["fetch_error"] == "synthetic upstream failure"
    assert captured[0]["ext_incidents"] is None

    wxcc_slice = slice_df_by_customer(failed, "Acme", pd.DataFrame())
    assert wxcc_slice.attrs["fetch_error"] == "synthetic upstream failure"
