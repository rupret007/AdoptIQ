"""Round 15 / Phase 3 -- Word visual polish tests.

Covers:

* ``report_word_styling.markdown_heading_level`` -- markdown ``#``
  count to Word heading level mapping (the Phase-3.1 demote rule).
* ``risk_band_fill_color`` -- canonical band -> hex resolution.
* ``set_cell_shading`` / ``get_cell_fill_hex`` -- round-trip cell
  fill color set + read.
* ``add_banded_top_n_table`` -- bold header row, banded rows,
  defensive on empty inputs.
* ``add_risk_callout`` -- single-cell colored table whose fill
  matches the canonical band color.
* ``add_executive_summary_table`` -- prepends a Heading-2 with a
  2-column KPI table; tolerant of missing inputs.
* End-to-end: ``adoptiq_backend.append_to_word_report`` no longer
  drops markdown ``#`` straight onto Heading 1 (Round 15 / Phase
  3.1 regression guard).
* Marker presence so the Round-15 Phase-3 wiring can't silently
  regress.
"""

from __future__ import annotations
from source_shape_utils import assert_in_source

from pathlib import Path

import pytest

import report_word_styling as wstyle


# ---------------------------------------------------------------------------
# Heading discipline
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "hash_count,expected_level",
    [
        (1, 2),
        (2, 3),
        (3, 4),
        (4, 4),
        (5, 4),
    ],
)
def test_phase_3_markdown_heading_level_demotes_one_step(hash_count, expected_level):
    """``# Title`` should never collapse to Heading 1.

    Round 15 / Phase 3.1 -- before this round every markdown ``#``
    in LLM-generated content rendered as Heading 1, producing the
    38-H1 gold docx we shipped.  Reserve H1 for the document title
    only and demote ``# / ## / ### / ####`` by one level.
    """
    assert wstyle.markdown_heading_level(hash_count) == expected_level


@pytest.mark.parametrize("bad_input", [None, "abc", -3, 0])
def test_phase_3_markdown_heading_level_handles_garbage(bad_input):
    """Invalid hash counts should never raise; default to H2 / H4."""
    level = wstyle.markdown_heading_level(bad_input)
    assert wstyle.MIN_BODY_HEADING_LEVEL <= level <= wstyle.MAX_HEADING_LEVEL


def test_phase_3_min_body_heading_level_is_two():
    """Heading 1 stays reserved for the title page only."""
    assert wstyle.MIN_BODY_HEADING_LEVEL == 2


def test_phase_3_max_heading_level_caps_at_four():
    """Deeply nested LLM markdown should never balloon past H4."""
    assert wstyle.MAX_HEADING_LEVEL == 4


# ---------------------------------------------------------------------------
# Color resolution
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "band,expected_substring",
    [
        ("Critical Risk", "C0392B"),
        ("High Risk", "FF6B6B"),
        ("Medium Risk", "FFB81C"),
        ("Low Risk", "5DBCD2"),
        ("Healthy", "28B463"),
        ("CRITICAL", "D62728"),
        ("HIGH", "FF7F0E"),
        ("MEDIUM", "FFD700"),
    ],
)
def test_phase_3_risk_band_fill_color_uses_canonical(band, expected_substring):
    """Risk bands must resolve to the canonical_metrics palette."""
    hex_value = wstyle.risk_band_fill_color(band)
    assert hex_value.upper() == expected_substring.upper()
    assert len(hex_value) == 6


@pytest.mark.parametrize("band", [None, "", "bogus-band-name", 42])
def test_phase_3_risk_band_fill_color_falls_back_for_unknown(band):
    """Unknown / missing bands return a 6-char hex (no crash)."""
    hex_value = wstyle.risk_band_fill_color(band)
    assert isinstance(hex_value, str)
    assert len(hex_value) == 6


def test_phase_3_strip_hex_normalizes_inputs():
    """``_strip_hex`` is the gate every cell-shading call goes through."""
    assert wstyle._strip_hex("#A1B2C3") == "A1B2C3"
    assert wstyle._strip_hex("a1b2c3") == "A1B2C3"
    assert wstyle._strip_hex("#FFA1B2C3") == "A1B2C3"  # ARGB stripped
    # Malformed inputs collapse to the neutral grey, never raise.
    assert wstyle._strip_hex("zzzz") == "BFBFBF"
    assert wstyle._strip_hex(None) == "BFBFBF"


# ---------------------------------------------------------------------------
# Cell shading round-trip
# ---------------------------------------------------------------------------


def test_phase_3_set_and_get_cell_fill_round_trips():
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    assert wstyle.get_cell_fill_hex(cell) is None
    assert wstyle.set_cell_shading(cell, "#1F77B4") is True
    assert wstyle.get_cell_fill_hex(cell) == "1F77B4"


def test_phase_3_set_cell_shading_handles_none_cell():
    """``None`` cell input is safe -- returns False, no crash."""
    assert wstyle.set_cell_shading(None, "#ABCDEF") is False
    assert wstyle.get_cell_fill_hex(None) is None


# ---------------------------------------------------------------------------
# Banded top-N table
# ---------------------------------------------------------------------------


def test_phase_3_add_banded_top_n_table_creates_header_and_rows():
    from docx import Document

    doc = Document()
    table = wstyle.add_banded_top_n_table(
        doc,
        headers=["Customer", "ARR", "Severity"],
        rows=[
            ["Acme Corp", "$1,250,000", "P1"],
            ["Beta Industries", "$890,000", "P2"],
            ["Cisco Systems", "$540,000", "P3"],
        ],
    )
    assert table is not None
    assert len(table.rows) == 4  # header + 3 data rows
    assert len(table.columns) == 3
    # Header row should be filled with the Cisco-blue header color
    # and every header run should be bolded.
    header_cells = table.rows[0].cells
    for cell in header_cells:
        assert wstyle.get_cell_fill_hex(cell) == wstyle.HEADER_FILL_HEX
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                assert run.bold is True
    # Banded rows: row 1 (data row index 0) un-banded, row 2 (data
    # row index 1) banded, row 3 (data row index 2) un-banded.
    second_data_row = table.rows[2].cells
    assert wstyle.get_cell_fill_hex(second_data_row[0]) == wstyle.ALT_ROW_FILL_HEX
    third_data_row = table.rows[3].cells
    assert wstyle.get_cell_fill_hex(third_data_row[0]) is None


def test_phase_3_add_banded_top_n_table_returns_none_for_empty():
    from docx import Document

    doc = Document()
    assert wstyle.add_banded_top_n_table(doc, headers=[], rows=[]) is None
    assert wstyle.add_banded_top_n_table(None, headers=["A"], rows=[["x"]]) is None


# ---------------------------------------------------------------------------
# Risk callout
# ---------------------------------------------------------------------------


def test_phase_3_add_risk_callout_uses_band_color():
    from docx import Document

    doc = Document()
    table = wstyle.add_risk_callout(
        doc,
        "Top Critical Risk",
        "Critical Risk",
        "Acme Corp has 3 unresolved P1 cases > 14d old.",
    )
    assert table is not None
    assert len(table.rows) == 1
    assert len(table.columns) == 1
    cell = table.rows[0].cells[0]
    expected = wstyle.risk_band_fill_color("Critical Risk")
    assert wstyle.get_cell_fill_hex(cell) == expected
    text = "\n".join(p.text for p in cell.paragraphs)
    assert_in_source(text, "Top Critical Risk", label='text')
    assert_in_source(text, "Acme Corp", label='text')


def test_phase_3_add_risk_callout_handles_missing_band():
    from docx import Document

    doc = Document()
    table = wstyle.add_risk_callout(doc, "Notice", None, "Body text.")
    assert table is not None
    cell = table.rows[0].cells[0]
    fill = wstyle.get_cell_fill_hex(cell)
    assert isinstance(fill, str)
    assert len(fill) == 6


def test_phase_3_add_risk_callout_returns_none_for_none_doc():
    assert wstyle.add_risk_callout(None, "x", "Critical Risk", "body") is None


# ---------------------------------------------------------------------------
# Executive summary table
# ---------------------------------------------------------------------------


def test_phase_3_add_executive_summary_table_renders_heading_and_table():
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE  # noqa: F401 - keeps test deps explicit

    doc = Document()
    rows = [
        ("Customers in portfolio", "27"),
        ("Adoption barriers (total)", "143"),
        ("TAC cases (P1)", "5"),
    ]
    table = wstyle.add_executive_summary_table(doc, rows, title="Executive Summary", title_level=2)
    assert table is not None
    # Heading should be Heading 2.
    headings = [p for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
    assert any(p.style.name == "Heading 2" and "Executive Summary" in p.text for p in headings)
    # Table layout: 1 header row + 3 KPI rows, 2 columns each.
    assert len(table.rows) == 4
    assert len(table.columns) == 2
    # Verify the KPI label/value cells contain the data.
    for r_idx, (label, value) in enumerate(rows, start=1):
        assert table.rows[r_idx].cells[0].text == label
        assert table.rows[r_idx].cells[1].text == value


def test_phase_3_add_executive_summary_table_skips_empty_rows():
    from docx import Document

    doc = Document()
    assert wstyle.add_executive_summary_table(doc, [], title="X") is None
    assert wstyle.add_executive_summary_table(doc, [(None, "v")], title="X") is None


def test_phase_3_add_executive_summary_table_caps_title_level():
    from docx import Document

    doc = Document()
    table = wstyle.add_executive_summary_table(
        doc,
        [("a", "1")],
        title="Deep",
        title_level=99,
    )
    assert table is not None
    headings = [p for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
    # Cap kicks in at MAX_HEADING_LEVEL.
    assert any(p.style.name == f"Heading {wstyle.MAX_HEADING_LEVEL}" for p in headings)


# ---------------------------------------------------------------------------
# End-to-end: append_to_word_report respects the Phase-3.1 demote rule.
# ---------------------------------------------------------------------------


def test_phase_3_append_to_word_report_demotes_markdown_headings(tmp_path):
    """The single-hash markdown heading must NOT land on Heading 1.

    Round 15 / Phase 3.1 regression guard: if anyone re-introduces
    ``add_heading(... level=1)`` for ``# md`` content, this test
    fails immediately.
    """
    from docx import Document

    import adoptiq_backend

    out = tmp_path / "round15_phase3_demote.docx"
    doc = Document()
    # Title page anchor: this *must* keep being H1.
    title = doc.add_heading("Manager Portfolio", level=1)
    assert title.style.name == "Heading 1"
    doc.save(str(out))

    md = (
        "# Top Section\n"
        "Some intro paragraph.\n\n"
        "## Subsection\n"
        "Subsection body.\n\n"
        "### Sub-subsection\n"
        "More detail.\n\n"
        "* bullet one\n"
        "* bullet two\n"
    )
    adoptiq_backend.append_to_word_report(str(out), md, heading="Section From Caller")

    reloaded = Document(str(out))
    body_headings = [
        (p.style.name, p.text)
        for p in reloaded.paragraphs
        if p.style and p.style.name.startswith("Heading")
    ]
    assert body_headings, "expected at least one heading paragraph"

    h1_paragraphs = [t for s, t in body_headings if s == "Heading 1"]
    # Only the original title page H1 should survive.
    assert h1_paragraphs == ["Manager Portfolio"], (
        "Round 15 / Phase 3.1 regression: markdown `# heading` collapsed "
        f"onto Heading 1: {h1_paragraphs}"
    )

    # The ``heading`` arg should be Heading 2.
    assert any(s == "Heading 2" and "Section From Caller" in t for s, t in body_headings)
    # Markdown ``# Top Section`` should be Heading 2.
    assert any(s == "Heading 2" and "Top Section" in t for s, t in body_headings)
    # Markdown ``## Subsection`` should be Heading 3.
    assert any(s == "Heading 3" and "Subsection" in t for s, t in body_headings)
    # Markdown ``### Sub-subsection`` should be Heading 4.
    assert any(s == "Heading 4" and "Sub-subsection" in t for s, t in body_headings)


def test_phase_3_create_executive_title_page_adds_summary_table(tmp_path):
    """The title page should now lead with an Executive Summary table.

    Round 15 / Phase 3.6 -- the gold docx had no executive summary
    table at the top; only a buried 3x4 dashboard.  After Phase 3
    the title page must end with a Heading 2 ``Executive Summary``
    followed by a 2-column KPI table.
    """
    from docx import Document

    import adoptiq_backend

    doc = Document()
    metrics = {
        "total_customers": 27,
        "total_barriers": 143,
        "total_cases": 88,
        "bems_count": 4,
        "high_risk_customers": 5,
        "medium_risk_customers": 12,
        "low_risk_customers": 7,
        "healthy_customers": 3,
        "critical_p1": 6,
        "high_p2": 14,
    }
    adoptiq_backend.create_executive_title_page(
        doc, "Brian Frazier", "Contact Center", 90, metrics
    )

    headings = [
        (p.style.name, p.text)
        for p in doc.paragraphs
        if p.style and p.style.name.startswith("Heading")
    ]
    assert any(s == "Heading 2" and "Executive Summary" in t for s, t in headings), (
        f"Expected a Heading-2 'Executive Summary' on the title page; got {headings}"
    )

    # There should be at least one table that matches the executive
    # summary signature: 2 columns, header + >= 3 data rows, header
    # text containing "Metric" and "Value".
    summary_tables = []
    for tbl in doc.tables:
        try:
            if len(tbl.columns) != 2:
                continue
            if len(tbl.rows) < 4:
                continue
            header_cells = tbl.rows[0].cells
            if header_cells[0].text == "Metric" and header_cells[1].text == "Value":
                summary_tables.append(tbl)
        except Exception:
            continue
    assert summary_tables, "Round 15 / Phase 3.6: no Executive Summary table on title page"

    # Verify the summary table contains the Phase-3.6 KPIs.
    body_text = "\n".join(
        c.text for tbl in summary_tables for r in tbl.rows for c in r.cells
    )
    assert_in_source(body_text, "Customers in portfolio", label='body_text')
    assert_in_source(body_text, "27", label='body_text')


# ---------------------------------------------------------------------------
# Marker-presence regression guards (mirror the Round-14 pattern).
# ---------------------------------------------------------------------------


_REPO_ROOT = Path(__file__).resolve().parent.parent


def _read(path: str) -> str:
    return (_REPO_ROOT / path).read_text(encoding="utf-8")


def test_phase_3_marker_in_report_word_styling():
    contents = _read("report_word_styling.py")
    assert "Round 15 / Phase 3" in contents


def test_phase_3_marker_in_adoptiq_backend_demote():
    contents = _read("adoptiq_backend.py")
    # Phase 3.1 marker for the heading-demote rule must be present.
    assert "Round 15 / Phase 3.1" in contents


def test_phase_3_marker_in_adoptiq_backend_summary_table():
    contents = _read("adoptiq_backend.py")
    # Phase 3.6 marker for the executive summary table must be present.
    assert "Round 15 / Phase 3.6" in contents


def test_phase_3_word_styling_module_imports_clean():
    """The new helper module imports without side effects."""
    import importlib

    import report_word_styling

    importlib.reload(report_word_styling)
    assert hasattr(report_word_styling, "markdown_heading_level")
    assert hasattr(report_word_styling, "risk_band_fill_color")
    assert hasattr(report_word_styling, "set_cell_shading")
    assert hasattr(report_word_styling, "add_banded_top_n_table")
    assert hasattr(report_word_styling, "add_risk_callout")
    assert hasattr(report_word_styling, "add_executive_summary_table")
    assert hasattr(report_word_styling, "build_executive_summary_rows")
