"""Round 25 / Phase D: Excel writer must not advertise a fallback path.

Pre-Round 25 ``adoptiq_backend.write_excel_workbook`` wrapped its body in
a ``try: from enhanced_excel_formatter ... except: ...`` block whose
``try:`` branch could never succeed -- the ``enhanced_excel_formatter``
module has never existed anywhere in the repo (no source file, no
entry-point, no historical commit).  Every generated workbook therefore
fell through to the ``except`` path and the ``Report_Info`` metadata
sheet announced ``Export type: Standard (fallback)`` followed by a
reassuring note that ``the enhanced formatter was not available``.

That framing was dishonest: there was no formatter to fall back *from*,
so the user was being told the workbook was a degraded copy of a
non-existent premium output.  Phase D promotes the inline body to the
canonical writer and rewrites ``Report_Info`` to reflect reality.

This test enforces both halves of that contract by:

1.  Refusing any live import of ``enhanced_excel_formatter`` in the
    backend module (the only acceptable references are the explanatory
    Round 25 / Phase D comments left behind as a tombstone).
2.  Confirming the generated workbook's ``Report_Info`` sheet contains
    ``Standard`` (without the ``(fallback)`` suffix) and that neither
    the legacy ``fallback`` nor ``enhanced formatter was not available``
    strings survive into the sheet.
"""

from __future__ import annotations

import os
import re
import tempfile

import pandas as pd
import pytest

import adoptiq_backend
from openpyxl import load_workbook


def test_no_live_import_of_enhanced_excel_formatter() -> None:
    """The backend source must not contain a live ``enhanced_excel_formatter`` import.

    Round 25 / Phase D specifically deletes the ``try: from
    enhanced_excel_formatter import create_enhanced_excel_report``
    import.  We allow the module name to appear *only* inside comments
    that document the removal -- any line that begins with ``from`` or
    ``import`` and references the module is a regression of Phase D.
    """

    src_path = os.path.join(os.path.dirname(adoptiq_backend.__file__), "adoptiq_backend.py")
    with open(src_path, "r", encoding="utf-8") as fh:
        source = fh.read()

    # Live import patterns -- ``enhanced_excel_formatter`` must not appear
    # on any non-comment ``import`` / ``from ... import`` line.
    live_import_re = re.compile(
        r"^\s*(?:from\s+enhanced_excel_formatter|import\s+enhanced_excel_formatter)",
        re.MULTILINE,
    )
    matches = live_import_re.findall(source)
    assert matches == [], (
        "Round 25 / Phase D regression: ``enhanced_excel_formatter`` is "
        "being imported again.  The module does not exist in the repo "
        "and the historical try/except wrapper that papered over its "
        "absence was deleted on purpose.  Remove the import."
    )


def test_report_info_sheet_says_standard_without_fallback_suffix() -> None:
    """Report_Info ``Export type`` row must read ``Standard`` (no suffix)."""

    sheets = {"AB_Detail_All": pd.DataFrame([{"ID": "AB-1", "Account": "ACME"}])}

    with tempfile.TemporaryDirectory() as tmp:
        base = os.path.join(tmp, "round25_phaseD")
        out = adoptiq_backend.write_excel_workbook(base, sheets)
        assert os.path.exists(out), f"Workbook was not written: {out}"

        wb = load_workbook(out, data_only=True)
        assert "Report_Info" in wb.sheetnames, (
            "Report_Info metadata sheet missing -- Phase D must keep "
            "writing it, just with honest content."
        )
        ws = wb["Report_Info"]

        rows = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and r[0] is not None:
                rows[str(r[0]).strip()] = ("" if r[1] is None else str(r[1]).strip())

        assert "Export type" in rows, (
            f"Report_Info missing ``Export type`` row.  Saw: {rows!r}"
        )
        assert rows["Export type"] == "Standard", (
            "Round 25 / Phase D: ``Export type`` must read ``Standard`` "
            "exactly -- the (fallback) suffix and the apologetic note "
            "are gone because there is no enhanced formatter to fall "
            f"back from.  Got: {rows['Export type']!r}"
        )

        # The legacy honesty-bug strings must be absent everywhere on
        # this sheet -- not just on the ``Export type`` row.
        joined = " | ".join(f"{k}={v}" for k, v in rows.items())
        assert "fallback" not in joined.lower(), (
            f"Round 25 / Phase D: legacy ``fallback`` framing leaked "
            f"into Report_Info.  Saw: {joined!r}"
        )
        assert "enhanced formatter" not in joined.lower(), (
            f"Round 25 / Phase D: legacy ``enhanced formatter`` "
            f"apology leaked into Report_Info.  Saw: {joined!r}"
        )


def test_report_info_sheet_records_generation_timestamp() -> None:
    """The replacement ``Generated at (UTC)`` row must be present and ISO-Z formatted."""

    sheets = {"AB_Detail_All": pd.DataFrame([{"ID": "AB-1"}])}

    with tempfile.TemporaryDirectory() as tmp:
        base = os.path.join(tmp, "round25_phaseD_ts")
        out = adoptiq_backend.write_excel_workbook(base, sheets)
        wb = load_workbook(out, data_only=True)
        ws = wb["Report_Info"]
        rows = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and r[0] is not None:
                rows[str(r[0]).strip()] = ("" if r[1] is None else str(r[1]).strip())

        assert "Generated at (UTC)" in rows, (
            "Round 25 / Phase D: ``Generated at (UTC)`` neutral note "
            "must replace the misleading enhanced-formatter apology.  "
            f"Saw: {rows!r}"
        )
        ts = rows["Generated at (UTC)"]
        # Either ISO-Z formatted or empty (defensive fallback when host
        # clock raises).  An empty string is acceptable; a junk value is
        # not.
        if ts:
            assert re.match(
                r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$", ts
            ), (
                f"Round 25 / Phase D: ``Generated at (UTC)`` must use "
                f"the UTC ISO-Z format Round 5 / Phase 6.3 standardized "
                f"on for cross-format timestamp parity.  Got: {ts!r}"
            )
