"""Round 167 CSConsole drill-through link contract."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

import decision_report_delivery as delivery
from source_record_links import (
    build_source_record_url,
    is_allowed_source_record_url,
)
from tests.test_round142_decision_report_delivery import _facts, _fake_chart_renderer


@pytest.mark.parametrize(
    ("source", "object_name"),
    [
        ("action_plans", "C360_CS_Task__c"),
        ("Adoption_Barriers", "C360_CS_Task__c"),
        ("Customer_Pulse", "ESA_C360_Customer_Pulse__c"),
        ("Success_Priorities", "ESA_C360_SUCCESS_PRIORITY__C"),
    ],
)
def test_supported_sources_get_allowlisted_csconsole_link(source, object_name) -> None:
    url = build_source_record_url(source, "a0A000000000001AAA")
    assert f"/lightning/r/{object_name}/a0A000000000001AAA/view" in url
    assert is_allowed_source_record_url(url) is True


@pytest.mark.parametrize(
    "record_id",
    ["", "../../escape", "abc?next=evil", "abc#fragment", "abc def", "//evil.test"],
)
def test_unsafe_or_missing_record_ids_never_produce_links(record_id) -> None:
    assert build_source_record_url("action_plans", record_id) == ""


def test_unknown_source_never_invents_a_link() -> None:
    assert build_source_record_url("tac_cases", "SR-001") == ""
    assert is_allowed_source_record_url("https://evil.test/lightning/r/X/abc/view") is False
    assert (
        is_allowed_source_record_url(
            "https://ciscosales.lightning.force.com:not-a-port/"
            "lightning/r/C360_CS_Task__c/a0A000000000001AAA/view"
        )
        is False
    )


def test_canonical_workbook_contains_matching_clickable_csconsole_links(
    tmp_path: Path,
) -> None:
    facts = _facts()
    sheets = delivery.build_source_data_sheets(facts)
    action_plans = sheets["Action_Plans"]
    assert "Source_Record_URL" in action_plans.columns
    linked = action_plans.loc[action_plans["Record_ID"].astype(str).ne("")]
    assert linked["Source_Record_URL"].map(is_allowed_source_record_url).all()
    assert action_plans.loc[
        action_plans["Record_ID"].astype(str).eq(""), "Source_Record_URL"
    ].fillna("").eq("").all()

    target = tmp_path / "source-links.xlsx"
    delivery.write_source_data_workbook(target, sheets)
    workbook = load_workbook(target, read_only=False, data_only=False)
    worksheet = workbook["Action_Plans"]
    headers = {cell.value: cell.column for cell in worksheet[1]}
    record_id_column = headers["Record_ID"]
    url_column = headers["Source_Record_URL"]
    for row_number in range(2, worksheet.max_row + 1):
        record_id = worksheet.cell(row_number, record_id_column).value
        url_cell = worksheet.cell(row_number, url_column)
        expected_url = build_source_record_url("action_plans", record_id)
        if expected_url:
            assert url_cell.value == expected_url
            assert url_cell.hyperlink is not None
            assert url_cell.hyperlink.target == expected_url
        else:
            assert url_cell.value in (None, "")
            assert url_cell.hyperlink is None
    workbook.close()

    validated = delivery.validate_written_source_workbook(target, facts)
    assert validated["ok"], validated["errors"]
    assert validated["source_url_cells"] == validated["source_url_hyperlinks"]
    assert validated["source_url_cells"] > 0


def test_canonical_word_links_selected_action_plan_records(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    facts = _facts()
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    document = delivery.build_concise_word_document(facts)
    result = delivery.validate_word_semantics(facts, document)

    assert result["ok"], result["errors"]
    assert result["source_record_hyperlink_count"] > 0
