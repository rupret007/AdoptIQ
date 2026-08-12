"""Round 165 direct-download artifact-integrity regressions."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook

import app_simple


_PUBLIC_ARTIFACT_ROUTES = ("analysis-download", "filename-download", "open-report")


def _request_artifact(client, route: str, analysis_id: str, artifact: Path, file_type: str):
    if route == "analysis-download":
        return client.get(f"/download/{analysis_id}/{file_type}")
    if route == "filename-download":
        return client.get(f"/download-file/{artifact.name}")
    if route == "open-report":
        return client.post(f"/open-report/{analysis_id}/{file_type}")
    raise AssertionError(f"unsupported test route: {route}")


@pytest.fixture
def direct_download_state(monkeypatch: pytest.MonkeyPatch):
    """Install isolated completed statuses and their durable audit rows."""

    with app_simple.analysis_status_lock:
        original_status = dict(app_simple.analysis_status)
        app_simple.analysis_status.clear()
    audit_rows: dict[str, dict[str, Any] | None] = {}
    monkeypatch.setattr(
        app_simple,
        "_build_status_from_report_history",
        lambda analysis_id: audit_rows.get(analysis_id),
    )
    monkeypatch.setattr(
        app_simple,
        "_r92_resolve_output_artifact",
        lambda raw, **_kwargs: str(raw) if raw and Path(raw).is_file() else None,
    )

    def install(
        analysis_id: str,
        artifact: Path,
        *,
        file_type: str,
        canonical: bool = True,
        audit_hash: str = "",
        status_hash: str = "",
    ) -> None:
        path_key = "word_report" if file_type == "docx" else "excel_report"
        hash_key = "word_hash" if file_type == "docx" else "excel_hash"
        status = {
            "analysis_id": analysis_id,
            "status": "completed",
            path_key: str(artifact),
            hash_key: status_hash,
        }
        if canonical:
            status["fact_fingerprint"] = "f" * 64
        with app_simple.analysis_status_lock:
            app_simple.analysis_status[analysis_id] = status
        audit_rows[analysis_id] = {hash_key: audit_hash}

    try:
        yield install
    finally:
        with app_simple.analysis_status_lock:
            app_simple.analysis_status.clear()
            app_simple.analysis_status.update(original_status)


@pytest.mark.flask
@pytest.mark.parametrize(
    ("file_type", "payload"),
    (("docx", b"verified-word"), ("xlsx", b"verified-source-data")),
)
def test_direct_download_requires_matching_audit_owned_hash(
    client,
    tmp_path: Path,
    direct_download_state,
    file_type: str,
    payload: bytes,
) -> None:
    artifact = tmp_path / f"canonical.{file_type}"
    artifact.write_bytes(payload)
    expected = hashlib.sha256(payload).hexdigest()
    analysis_id = f"round165-verified-{file_type}"
    direct_download_state(
        analysis_id,
        artifact,
        file_type=file_type,
        audit_hash=expected,
    )

    response = client.get(f"/download/{analysis_id}/{file_type}")

    assert response.status_code == 200
    assert response.data == payload
    assert response.headers["X-AdoptIQ-Artifact-Integrity"] == "sha256-verified"
    assert "Warning" not in response.headers


@pytest.mark.flask
@pytest.mark.parametrize("file_type", ("docx", "xlsx"))
def test_direct_download_blocks_changed_canonical_artifact(
    client,
    tmp_path: Path,
    direct_download_state,
    file_type: str,
) -> None:
    artifact = tmp_path / f"changed.{file_type}"
    artifact.write_bytes(b"changed-after-audit")
    analysis_id = f"round165-changed-{file_type}"
    direct_download_state(
        analysis_id,
        artifact,
        file_type=file_type,
        audit_hash=hashlib.sha256(b"original-audited-bytes").hexdigest(),
    )

    response = client.get(f"/download/{analysis_id}/{file_type}")

    assert response.status_code == 409
    assert response.get_json()["integrity_state"] == "hash-mismatch"
    assert b"changed-after-audit" not in response.data


@pytest.mark.flask
def test_canonical_download_without_persisted_hash_fails_closed(
    client,
    tmp_path: Path,
    direct_download_state,
) -> None:
    artifact = tmp_path / "missing-hash.docx"
    artifact.write_bytes(b"canonical-but-unhashed")
    analysis_id = "round165-canonical-missing-hash"
    direct_download_state(
        analysis_id,
        artifact,
        file_type="docx",
        audit_hash="",
    )

    response = client.get(f"/download/{analysis_id}/docx")

    assert response.status_code == 409
    assert response.get_json()["integrity_state"] == "missing-persisted-hash"
    assert b"canonical-but-unhashed" not in response.data


@pytest.mark.flask
def test_legacy_hashless_download_is_explicitly_marked_unverified(
    client,
    tmp_path: Path,
    direct_download_state,
) -> None:
    artifact = tmp_path / "legacy.docx"
    artifact.write_bytes(b"pre-canonical-legacy-report")
    analysis_id = "round165-legacy-hashless"
    direct_download_state(
        analysis_id,
        artifact,
        file_type="docx",
        canonical=False,
        audit_hash="",
    )

    response = client.get(f"/download/{analysis_id}/docx")

    assert response.status_code == 200
    assert response.data == b"pre-canonical-legacy-report"
    assert response.headers["X-AdoptIQ-Artifact-Integrity"] == "legacy-unverified"
    assert "no persisted SHA-256" in response.headers["Warning"]


@pytest.mark.flask
def test_malformed_persisted_hash_is_not_treated_as_legacy(
    client,
    tmp_path: Path,
    direct_download_state,
) -> None:
    artifact = tmp_path / "malformed.xlsx"
    artifact.write_bytes(b"source-data")
    analysis_id = "round165-malformed-hash"
    direct_download_state(
        analysis_id,
        artifact,
        file_type="xlsx",
        canonical=False,
        audit_hash="not-a-sha256",
    )

    response = client.get(f"/download/{analysis_id}/xlsx")

    assert response.status_code == 409
    assert response.get_json()["integrity_state"] == "invalid-persisted-hash"


@pytest.mark.flask
def test_durable_audit_hash_overrides_volatile_status_hash(
    client,
    tmp_path: Path,
    direct_download_state,
) -> None:
    artifact = tmp_path / "audit-authoritative.docx"
    payload = b"volatile-status-would-accept-this"
    artifact.write_bytes(payload)
    analysis_id = "round165-audit-authoritative"
    direct_download_state(
        analysis_id,
        artifact,
        file_type="docx",
        audit_hash=hashlib.sha256(b"durable-audited-bytes").hexdigest(),
        status_hash=hashlib.sha256(payload).hexdigest(),
    )

    response = client.get(f"/download/{analysis_id}/docx")

    assert response.status_code == 409
    assert response.get_json()["integrity_state"] == "hash-mismatch"


@pytest.mark.flask
@pytest.mark.parametrize("route", _PUBLIC_ARTIFACT_ROUTES)
@pytest.mark.parametrize("file_type", ("docx", "xlsx"))
def test_tampered_current_artifact_is_blocked_across_every_public_route(
    client,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    direct_download_state,
    route: str,
    file_type: str,
) -> None:
    artifact = tmp_path / f"round165-tampered-{route}.{file_type}"
    tampered_payload = b"tampered-current-report-bytes"
    artifact.write_bytes(tampered_payload)
    analysis_id = f"round165-tampered-{route}-{file_type}"
    direct_download_state(
        analysis_id,
        artifact,
        file_type=file_type,
        audit_hash=hashlib.sha256(b"original-audited-report-bytes").hexdigest(),
    )
    opened: list[str] = []
    monkeypatch.setattr(
        app_simple,
        "_r92_open_path_in_default_app",
        lambda path: opened.append(path),
    )

    response = _request_artifact(
        client,
        route,
        analysis_id,
        artifact,
        file_type,
    )

    assert response.status_code == 409
    assert response.get_json()["integrity_state"] == "hash-mismatch"
    assert tampered_payload not in response.data
    assert str(tmp_path) not in response.get_data(as_text=True)
    assert analysis_id not in response.get_data(as_text=True)
    assert opened == []


@pytest.mark.flask
@pytest.mark.parametrize("route", _PUBLIC_ARTIFACT_ROUTES)
@pytest.mark.parametrize("file_type", ("docx", "xlsx"))
def test_hashless_legacy_artifact_stays_available_with_explicit_warning(
    client,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    direct_download_state,
    route: str,
    file_type: str,
) -> None:
    artifact = tmp_path / f"round165-legacy-{route}.{file_type}"
    payload = b"pre-canonical-compatible-artifact"
    artifact.write_bytes(payload)
    analysis_id = f"round165-legacy-{route}-{file_type}"
    direct_download_state(
        analysis_id,
        artifact,
        file_type=file_type,
        canonical=False,
        audit_hash="",
    )
    opened: list[str] = []
    monkeypatch.setattr(
        app_simple,
        "_r92_open_path_in_default_app",
        lambda path: opened.append(path),
    )

    response = _request_artifact(
        client,
        route,
        analysis_id,
        artifact,
        file_type,
    )

    assert response.status_code == 200
    assert response.headers["X-AdoptIQ-Artifact-Integrity"] == "legacy-unverified"
    assert "no persisted SHA-256" in response.headers["Warning"]
    if route == "open-report":
        assert response.get_json()["ok"] is True
        assert response.get_json()["target"] == file_type
        assert opened == [str(artifact)]
    else:
        assert response.data == payload
        assert opened == []


@pytest.mark.flask
@pytest.mark.parametrize("file_type", ("docx", "xlsx"))
def test_unattributed_canonical_artifact_cannot_fall_back_to_legacy_mode(
    client,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    file_type: str,
) -> None:
    artifact = tmp_path / f"unattributed-current.{file_type}"
    if file_type == "docx":
        document = Document()
        document.core_properties.identifier = "f" * 64
        document.add_paragraph("current canonical report")
        document.save(artifact)
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Report_Info"
        worksheet.append(["Item", "Value"])
        worksheet.append(["Fact_Contract_SHA256", "f" * 64])
        workbook.save(artifact)
        workbook.close()

    with app_simple.analysis_status_lock:
        original_status = dict(app_simple.analysis_status)
        app_simple.analysis_status.clear()
    monkeypatch.setattr(
        app_simple,
        "_r98_resolve_download_filename",
        lambda _filename: str(artifact),
    )
    monkeypatch.setattr(app_simple, "get_report_history", lambda **_kwargs: [])
    try:
        response = client.get(f"/download-file/{artifact.name}")
    finally:
        with app_simple.analysis_status_lock:
            app_simple.analysis_status.clear()
            app_simple.analysis_status.update(original_status)

    assert response.status_code == 409
    assert response.get_json()["integrity_state"] == "missing-audit-context"
