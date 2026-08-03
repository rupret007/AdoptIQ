"""Round 142 deterministic, Snowflake-free acceptance artifact harness."""

from __future__ import annotations

import base64
import hashlib
import importlib.util
import json
import zipfile
from pathlib import Path

import pandas as pd
import pytest
from docx import Document

import decision_report_delivery as delivery
from scripts import generate_offline_acceptance_artifacts as harness


AS_OF = "2026-08-03T12:00:00Z"
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)
_REQUIRED_SOURCE_SHEETS = {
    "Report_Info",
    "Metric_Lineage",
    "Chart_Data",
    "Action_Plans",
    "Adoption_Barriers",
    "Customer_Pulse",
    "TAC_Cases",
    "BEMS",
    "Subscriptions",
    "Success_Priorities",
    "External_Incidents",
    "External_Bugs",
    "Risk_Components",
    "Member_Summary",
    "Account_Summary",
}
_EXPECTED_SCOPE = {
    "team": ("Leader", "team", 2, 3),
    "member": ("Leader", "member", 1, 2),
    "customer": ("Leader", "customer", 2, 1),
    "comprehensive": ("Comprehensive", "team", 2, 3),
}


def _fake_chart_renderer(_chart_id: str, _rows: pd.DataFrame, target: Path) -> bool:
    target.write_bytes(_TINY_PNG)
    return True


def _load_json(path: str | Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def test_sanitized_fixture_covers_acceptance_edge_cases() -> None:
    payload = harness.load_sanitized_fixture()
    assert set(payload["team_data"]) == {"Alex Rivera", "Morgan Lee"}

    bundles = list(payload["team_data"].values())
    customers = {
        row["BU_NAME"]
        for bundle in bundles
        for row in bundle["subscriptions"]
    }
    assert customers == {"Acme Corporation", "Beta Industries", "Gamma Public Sector"}

    for source, id_column in (
        ("action_plans", "ID"),
        ("adoption_barriers", "ID"),
        ("customer_pulse", "ID"),
        ("tac_cases", "SR Number"),
    ):
        ids = [
            row.get(id_column)
            for bundle in bundles
            for row in bundle[source]
            if row.get(id_column)
        ]
        assert len(ids) > len(set(ids)), f"{source} must include a shared stable ID"

    action_plans = [row for bundle in bundles for row in bundle["action_plans"]]
    assert any(not row.get("ID") for row in action_plans)
    assert any(not str(row.get("SUBJECT_C") or "").strip() for row in action_plans)
    assert {row.get("STATUS_C") for row in action_plans}.issuperset(
        {"Open", "On Hold", "Completed - Successful", "Awaiting business validation"}
    )

    bems_ids = {
        row.get("Transaction ID")
        for bundle in bundles
        for row in bundle["tac_cases"]
        if str(row.get("Transaction ID") or "").startswith("BEMS-")
    }
    assert bems_ids == {"BEMS-0001", "BEMS-0002"}
    assert payload["external_incidents"]
    assert payload["external_bugs"]


@pytest.mark.parametrize("scope", harness.SUPPORTED_SCOPES)
def test_all_scopes_generate_paired_artifacts_and_manifests(
    scope: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    result = harness.generate_acceptance_artifacts(
        scope=scope,
        as_of=AS_OF,
        output_dir=tmp_path / scope,
    )

    paths = {
        key: Path(result[key])
        for key in (
            "word_path",
            "source_data_path",
            "measurement_manifest_path",
            "parity_manifest_path",
            "chart_manifest_path",
        )
    }
    assert all(path.is_file() for path in paths.values())
    assert paths["word_path"].name.startswith("AdoptIQ_Report_")
    assert paths["source_data_path"].name.startswith("AdoptIQ_Source_Data_")
    assert (
        paths["source_data_path"].stem.removeprefix("AdoptIQ_Source_Data_")
        == paths["word_path"].stem.removeprefix("AdoptIQ_Report_")
    )

    report_type, scope_type, team_members, customers = _EXPECTED_SCOPE[scope]
    measurement = _load_json(paths["measurement_manifest_path"])
    assert measurement["requested_scope"] == scope
    assert measurement["report_type"] == report_type
    assert measurement["scope_type"] == scope_type
    assert measurement["as_of_utc"] == "2026-08-03T12:00:00+00:00"
    assert measurement["kpis"]["team_members"] == team_members
    assert measurement["kpis"]["customers"] == customers
    assert measurement["document"]["embedded_chart_count"] == 4
    assert measurement["document"]["word_count"] <= measurement["document"]["word_budget"]

    parity = _load_json(paths["parity_manifest_path"])
    assert parity["requested_scope"] == scope
    assert parity["ok"] is True
    assert all(parity["checks"].values())
    assert parity["fixture_oracle"]["actual"] == parity["fixture_oracle"]["expected"]
    assert parity["shared_attribution"]["expected"] is (scope != "member")
    assert parity["shared_attribution"]["observed"] is (scope != "member")

    charts = _load_json(paths["chart_manifest_path"])
    assert charts["requested_scope"] == scope
    assert {item["chart_id"] for item in charts["charts"]} == {
        "activity_mix",
        "action_plan_status_aging",
        "risk_distribution",
        "activity_trend",
    }
    assert all(item["available_value_rows"] > 0 for item in charts["charts"])

    document = Document(paths["word_path"])
    assert len(document.inline_shapes) == 4
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Source and Lineage Note" in document_text
    assert "Data Coverage Warning" in document_text
    assert "All Action Plans" not in document_text

    workbook = pd.ExcelFile(paths["source_data_path"])
    assert _REQUIRED_SOURCE_SHEETS.issubset(workbook.sheet_names)
    action_plans = pd.read_excel(paths["source_data_path"], sheet_name="Action_Plans")
    measured_sheets = {
        item["sheet"]: item for item in measurement["source_sheets"]
    }
    assert measured_sheets["Action_Plans"]["rows"] == len(action_plans)
    assert measured_sheets["Action_Plans"]["columns"] == len(action_plans.columns)
    assert {
        "Record_ID",
        "AdoptIQ_Title",
        "AdoptIQ_Status_Bucket",
        "AdoptIQ_Data_Quality",
        "Scope_Type",
        "Scope_Value",
        "Source_System",
        "Attributed_Team_Members",
    }.issubset(action_plans.columns)
    assert "ETL_ID" not in action_plans.columns
    assert set(action_plans["Source_System"]) == {harness.OFFLINE_SOURCE_SYSTEM}
    assert len(action_plans) == measurement["action_plan_lifecycle"]["total"]
    bems = pd.read_excel(paths["source_data_path"], sheet_name="BEMS")
    incidents = pd.read_excel(paths["source_data_path"], sheet_name="External_Incidents")
    bugs = pd.read_excel(paths["source_data_path"], sheet_name="External_Bugs")
    assert len(bems) > 0
    assert {"Record_ID", "SR Number", "Transaction ID", "Source_System"}.issubset(
        bems.columns
    )
    assert len(incidents) == 2
    assert {"Record_ID", "id", "Scope_Type", "Source_System"}.issubset(
        incidents.columns
    )
    assert len(bugs) == 2
    assert {"Record_ID", "bug_id", "Scope_Type", "Source_System"}.issubset(bugs.columns)
    report_info = pd.read_excel(paths["source_data_path"], sheet_name="Report_Info")
    warning_rows = report_info.loc[
        report_info["Item"].fillna("").astype(str).str.match(r"Partial_Data_Warning_\d+$")
    ]
    assert len(warning_rows) == 1
    assert "no live Snowflake" in str(warning_rows.iloc[0]["Detail"])
    source_state_rows = report_info.loc[
        report_info["Item"].fillna("").astype(str).str.startswith("Source_State:")
    ]
    assert "partial" in set(source_state_rows["Value"].astype(str))


def test_cli_rejects_invalid_scope_without_writing_artifacts(tmp_path: Path) -> None:
    output_dir = tmp_path / "invalid"
    with pytest.raises(SystemExit) as exc_info:
        harness.main(
            [
                "--scope",
                "invalid",
                "--as-of",
                AS_OF,
                "--output-dir",
                str(output_dir),
            ]
        )

    assert exc_info.value.code == 2
    assert not output_dir.exists()


def test_same_fixture_and_as_of_produce_identical_manifests(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    first = harness.generate_acceptance_artifacts(
        scope="team",
        as_of=AS_OF,
        output_dir=tmp_path / "first",
    )
    second = harness.generate_acceptance_artifacts(
        scope="team",
        as_of=AS_OF,
        output_dir=tmp_path / "second",
    )

    for key in (
        "measurement_manifest_path",
        "parity_manifest_path",
        "chart_manifest_path",
    ):
        first_payload = _load_json(first[key])
        second_payload = _load_json(second[key])
        assert first_payload == second_payload

    for key in ("word_path", "source_data_path"):
        first_path = Path(first[key])
        second_path = Path(second[key])
        assert hashlib.sha256(first_path.read_bytes()).digest() == hashlib.sha256(
            second_path.read_bytes()
        ).digest()
        with zipfile.ZipFile(first_path) as archive:
            assert {item.date_time for item in archive.infolist()} == {
                (1980, 1, 1, 0, 0, 0)
            }
            core = archive.read("docProps/core.xml")
            assert b"2026-08-03T12:00:00Z" in core


def test_team_and_comprehensive_share_the_same_portfolio_facts(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(delivery, "_render_chart_image", _fake_chart_renderer)
    team = harness.generate_acceptance_artifacts(
        scope="team",
        as_of=AS_OF,
        output_dir=tmp_path / "team",
    )
    comprehensive = harness.generate_acceptance_artifacts(
        scope="comprehensive",
        as_of=AS_OF,
        output_dir=tmp_path / "comprehensive",
    )

    team_measurement = _load_json(team["measurement_manifest_path"])
    comprehensive_measurement = _load_json(
        comprehensive["measurement_manifest_path"]
    )
    assert team_measurement["kpis"] == comprehensive_measurement["kpis"]
    assert (
        team_measurement["action_plan_lifecycle"]
        == comprehensive_measurement["action_plan_lifecycle"]
    )
    team_charts = _load_json(team["chart_manifest_path"])
    comprehensive_charts = _load_json(comprehensive["chart_manifest_path"])
    assert team_charts["series"] == comprehensive_charts["series"]

    for sheet_name in (
        "Action_Plans",
        "Adoption_Barriers",
        "Customer_Pulse",
        "TAC_Cases",
        "BEMS",
        "Subscriptions",
        "Success_Priorities",
    ):
        team_sheet = pd.read_excel(team["source_data_path"], sheet_name=sheet_name)
        comprehensive_sheet = pd.read_excel(
            comprehensive["source_data_path"],
            sheet_name=sheet_name,
        )
        assert list(team_sheet.columns) == list(comprehensive_sheet.columns)
        team_sheet = team_sheet.drop(columns=["Scope_Value"], errors="ignore")
        comprehensive_sheet = comprehensive_sheet.drop(
            columns=["Scope_Value"],
            errors="ignore",
        )
        sort_columns = [
            column
            for column in ("Record_ID", "CSSM", "Attributed_Team_Members")
            if column in team_sheet.columns
        ]
        if sort_columns:
            team_sheet = team_sheet.sort_values(
                sort_columns,
                kind="stable",
                na_position="last",
            )
            comprehensive_sheet = comprehensive_sheet.sort_values(
                sort_columns,
                kind="stable",
                na_position="last",
            )
        pd.testing.assert_frame_equal(
            team_sheet.reset_index(drop=True),
            comprehensive_sheet.reset_index(drop=True),
            check_dtype=False,
        )


@pytest.mark.skipif(
    importlib.util.find_spec("matplotlib") is None,
    reason="matplotlib is installed by the application requirements, not the base test runtime",
)
def test_installed_runtime_embeds_real_charts(tmp_path: Path) -> None:
    result = harness.generate_acceptance_artifacts(
        scope="team",
        as_of=AS_OF,
        output_dir=tmp_path,
    )

    document = Document(result["word_path"])
    assert len(document.inline_shapes) == 4
    parity = _load_json(result["parity_manifest_path"])
    assert parity["ok"] is True


def test_harness_has_no_app_or_snowflake_import_boundary() -> None:
    source = Path(harness.__file__).read_text(encoding="utf-8")
    for forbidden in ("import app_simple", "import adoptiq_backend", "import snowflake"):
        assert forbidden not in source
    for public_function in (
        "partition_portfolio_by_member",
        "build_report_facts",
        "build_source_data_sheets",
        "build_concise_word_document",
        "validate_cross_artifact_contract",
        "source_data_path_for_word",
        "write_source_data_workbook",
        "validate_written_source_workbook",
    ):
        assert f"delivery.{public_function}(" in source
