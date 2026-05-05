"""Round 78 / B1 — Comprehensive narrative stub-bullet filter.

Build 53 acceptance audit on the Comprehensive DOCX
(``AdoptIQ_Report_Brian_Frazier_All_Contact_Center_90d_*.docx``) found
the per-customer storyboard LLM emitting **170 stub bullets** of the
form ``"<Category>: Data unavailable."`` -- pure empty-bucket
acknowledgements with no narrative content.  Pattern frequency from
the audit::

    21x "Industry Benchmarking: Data unavailable."
    15x "Technical Competency: Data unavailable."
    14x "Adoption Velocity: Data unavailable."
    13x "Competitive Positioning: Data unavailable."
    13x "Integration Complexity: Data unavailable."
    10x "Operational Disruption: Data unavailable."
     9x "Sentiment Indicators: Data unavailable."
     8x "Communication Patterns: Data unavailable."

The bullets are technically accurate (the LLM honestly admits a gap)
but they're noise -- per-customer narratives become 50%+ empty
bullets that drown the substantive ones.

Round 78 / B1 adds a post-filter in
``adoptiq_backend.append_to_word_report`` (the bullet branch) that
drops the line entirely when the entire bullet text matches
``_R78_STUB_RE`` -- a regex specifically tuned to fire ONLY when the
ENTIRE bullet is the stub.  Bullets with any substantive narrative
after the marker are preserved unchanged.

Critical contract:
    DROP:    ``"Industry Benchmarking: Data unavailable."``
    DROP:    ``"**Industry Benchmarking:** Data unavailable."``
    DROP:    ``"Operational Disruption: Data unavailable"`` (no period)
    PRESERVE: ``"Operational Disruption: Data unavailable. No active incidents."``
    PRESERVE: ``"Strategic Headwinds: ... (SP-ID: data unavailable) suggests..."``
    PRESERVE: ``"Financial Impact: Data unavailable regarding ARR..."``

Made-with: Cursor.  Round 78.
"""

from __future__ import annotations

import inspect
import re

import pytest


# ---------------------------------------------------------------------------
# Source-shape pin
# ---------------------------------------------------------------------------


def test_r78_stub_re_module_constant_is_defined() -> None:
    """``_R78_STUB_RE`` must be defined as a module-level compiled
    regex constant in ``adoptiq_backend.py``.  If a future refactor
    inlines or renames it, downstream regression tests below will
    break loudly -- this is the canonical name future code must use.
    """
    import adoptiq_backend

    assert hasattr(adoptiq_backend, "_R78_STUB_RE"), (
        "Round 78 / B1: _R78_STUB_RE module constant is missing from "
        "adoptiq_backend.py -- the post-filter has been removed."
    )
    pat = adoptiq_backend._R78_STUB_RE
    assert isinstance(pat, re.Pattern), (
        "Round 78 / B1: _R78_STUB_RE must be a compiled regex Pattern."
    )


def test_r78_stub_re_pattern_anchored_to_full_line() -> None:
    """The regex MUST be anchored ``^...$`` so partial-string matches
    can never silently drop substantive bullets that happen to
    contain the literal ``Data unavailable`` somewhere mid-text.
    """
    import adoptiq_backend

    pattern_str = adoptiq_backend._R78_STUB_RE.pattern
    assert pattern_str.startswith("^"), (
        "Round 78 / B1: _R78_STUB_RE pattern must start with ^ "
        "(full-line anchor) so substantive bullets are never dropped."
    )
    assert pattern_str.rstrip("$").endswith(r"\s*"), (
        "Round 78 / B1: _R78_STUB_RE pattern must end with \\s*$ so "
        "trailing whitespace and EOL are tolerated but substantive "
        "follow-on text is rejected."
    )


def test_r78_stub_filter_wired_into_append_to_word_report() -> None:
    """The filter must be wired into the bullet branch of
    ``append_to_word_report``.  We assert the source-shape so a
    refactor that moves the filter to a different site still has to
    update this test.

    Note: ``adoptiq_backend.append_to_word_report`` is rebound at
    module-load time to ``enhanced_append_to_word_report`` (a pass-
    through wrapper), so we inspect the underlying
    ``original_append_to_word_report`` for the source-shape pin.
    """
    import adoptiq_backend

    real = getattr(
        adoptiq_backend,
        "original_append_to_word_report",
        adoptiq_backend.append_to_word_report,
    )
    src = inspect.getsource(real)
    assert "_R78_STUB_RE" in src, (
        "Round 78 / B1: the stub-bullet filter is no longer wired "
        "into append_to_word_report.  The Comprehensive DOCX will "
        "regress to ~170 empty bullets."
    )
    # Mirror the empty-line-skip convention (i += 1 then continue).
    # The order in the source MUST be: regex check, then i += 1,
    # then continue -- otherwise the while loop will infinite-loop.
    block = src[src.index("_R78_STUB_RE"):]
    block = block[: 400]
    assert "i += 1" in block, (
        "Round 78 / B1: the stub-skip path must increment i before "
        "continue so the while-loop doesn't infinite-loop on the "
        "unincremented index."
    )
    assert "continue" in block, (
        "Round 78 / B1: the stub-skip path must use continue to "
        "fall through to the next markdown line."
    )


# ---------------------------------------------------------------------------
# Regex behavior — pure stubs MATCH (drop)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "stub",
    [
        "Industry Benchmarking: Data unavailable.",
        "Technical Competency: Data unavailable.",
        "Adoption Velocity: Data unavailable.",
        "Competitive Positioning: Data unavailable.",
        "Integration Complexity: Data unavailable.",
        "Operational Disruption: Data unavailable.",
        "Sentiment Indicators: Data unavailable.",
        "Communication Patterns: Data unavailable.",
    ],
)
def test_pure_stubs_from_build53_audit_all_match(stub: str) -> None:
    """The 8 categories most-frequently empty in Build 53 audit must
    ALL match the regex.  This is the canonical drop set.
    """
    import adoptiq_backend

    assert adoptiq_backend._R78_STUB_RE.match(stub), (
        f"Round 78 / B1: pure stub {stub!r} should match the filter "
        f"but did not -- the Build 53 noise will return."
    )


@pytest.mark.parametrize(
    "stub",
    [
        "**Industry Benchmarking:** Data unavailable.",
        "**Industry Benchmarking**: Data unavailable.",
        "Industry Benchmarking: **Data unavailable.**",
        "**Industry Benchmarking: Data unavailable.**",
    ],
)
def test_bold_wrapped_pure_stubs_match(stub: str) -> None:
    """Various LLM bold-wrap permutations are all stubs and must all
    drop.  ``**Cat:** Data unavailable.``, ``**Cat**: Data
    unavailable.``, ``Cat: **Data unavailable.**``, and the
    fully-bold ``**Cat: Data unavailable.**`` forms are equally
    empty -- the user gains nothing from any of them.
    """
    import adoptiq_backend

    assert adoptiq_backend._R78_STUB_RE.match(stub), (
        f"Round 78 / B1: bold-wrapped stub {stub!r} should match "
        f"the filter but did not."
    )


def test_stub_without_terminal_period_matches() -> None:
    """LLMs occasionally drop the trailing period.  The filter must
    still catch the stub.
    """
    import adoptiq_backend

    assert adoptiq_backend._R78_STUB_RE.match(
        "Industry Benchmarking: Data unavailable"
    ), (
        "Round 78 / B1: stub without trailing period should still "
        "match -- LLM does not always emit clean punctuation."
    )


def test_stub_with_lowercase_data_unavailable_matches() -> None:
    """The category prefix is Title Case (LLM convention) but
    ``data unavailable`` itself can be any case.  We use [Dd]
    [Uu] character classes for that portion only -- so we do NOT
    inadvertently match generic lowercase prefixes like 'random
    text: data unavailable.'.
    """
    import adoptiq_backend

    # Lowercase 'data unavailable' but Title Case prefix - matches.
    assert adoptiq_backend._R78_STUB_RE.match(
        "Industry Benchmarking: data unavailable."
    )
    # Lowercase prefix - must NOT match (avoids false-positives).
    assert not adoptiq_backend._R78_STUB_RE.match(
        "industry benchmarking: data unavailable."
    ), (
        "Round 78 / B1: prefix MUST start with uppercase letter -- "
        "this avoids matching arbitrary lowercase user text."
    )


# ---------------------------------------------------------------------------
# Regex behavior — substantive bullets MUST NOT match (preserve)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "preserved",
    [
        # From Build 53 audit: stub marker + substantive narrative
        "Operational Disruption: Data unavailable. No active incidents or cases reported.",
        # Stub marker + reasoning continuation
        "Technical Competency: Data unavailable; however, the reliance on external playcards suggests a need for significant enablement.",
        # Stub marker as preamble to ARR commentary
        "Financial Impact: Data unavailable regarding specific ARR/dollar figures. However, the 'Completed - Unsuccessful' status of a major strategic initiative represents a significant risk.",
        # Mid-sentence parenthetical, not a stub
        "Strategic Headwinds: The lack of defined success priorities (SP-ID: data unavailable) suggests that the customer is not currently leveraging the full strategic potential.",
        # Generic narrative without 'Data unavailable.' as the entire content
        "Industry Benchmarking: Customer ranks in the top quartile of comparable Webex deployments.",
        # 'unavailable' but with a different phrase
        "Operational Disruption: System unavailable for the past 48 hours.",
    ],
)
def test_substantive_bullets_are_preserved(preserved: str) -> None:
    """Bullets with any substantive narrative content after the
    'Data unavailable' marker (or where 'data unavailable' is
    parenthetical inside other narrative) MUST be preserved.
    """
    import adoptiq_backend

    assert not adoptiq_backend._R78_STUB_RE.match(preserved), (
        f"Round 78 / B1: substantive bullet {preserved!r} should "
        f"NOT match the stub filter -- the user would lose real "
        f"narrative content."
    )


def test_negative_control_unrelated_text() -> None:
    """Bullets that have nothing to do with stubs must never match.
    """
    import adoptiq_backend

    for text in [
        "Customer reported high satisfaction scores in Q2.",
        "Top adoption barrier: Webex App login issues affecting 12 users.",
        "ARR at risk: $1.2M across 3 critical accounts.",
        "Recommendation: Schedule executive engagement within 14 days.",
        "Status: Completed - Successful",
    ]:
        assert not adoptiq_backend._R78_STUB_RE.match(text), (
            f"Round 78 / B1: unrelated text {text!r} matched the "
            f"stub filter -- false positive risk."
        )


# ---------------------------------------------------------------------------
# End-to-end — append_to_word_report drops only the stub bullets
# ---------------------------------------------------------------------------


def test_end_to_end_mixed_markdown_drops_only_stubs() -> None:
    """Feed a markdown body with 5 pure stubs interleaved with 5
    substantive bullets through ``append_to_word_report``.  The
    resulting DOCX must carry exactly 5 List Bullet paragraphs --
    the substantive ones -- and zero of the stubs.

    The function signature is ``append_to_word_report(doc_or_path,
    markdown_content, heading=None)`` -- when we pass a Document
    object it operates in-place and we can inspect it directly
    (no need to save+reload).
    """
    import docx

    import adoptiq_backend

    markdown = """## Risks

* Industry Benchmarking: Data unavailable.
* Technical Competency: The customer's CSM team is highly skilled but understaffed.
* Adoption Velocity: Data unavailable.
* Competitive Positioning: Strong incumbent position vs Microsoft Teams.
* Integration Complexity: Data unavailable.
* Operational Disruption: Data unavailable. No active incidents reported.
* Sentiment Indicators: Data unavailable.
* Communication Patterns: Weekly Webex Connect cadence with the assigned CSM.
* Industry Benchmarking: Customer ranks in the top quartile of comparable deployments.
* Strategic Headwinds: Data unavailable.
"""

    doc = docx.Document()
    adoptiq_backend.append_to_word_report(
        doc, markdown, heading="R78 / B1 mixed input"
    )

    bullet_paras = [
        p
        for p in doc.paragraphs
        if p.style and p.style.name == "List Bullet"
    ]
    assert len(bullet_paras) == 5, (
        f"Round 78 / B1: expected exactly 5 List Bullet paragraphs "
        f"after stub filter, got {len(bullet_paras)}: "
        f"{[p.text for p in bullet_paras]!r}"
    )
    bullet_texts = {p.text for p in bullet_paras}
    # No bullet should be a pure stub.  Pure stubs have nothing
    # after 'Data unavailable.' (or have nothing else at all).
    for text in bullet_texts:
        if "Data unavailable" not in text:
            continue
        idx = text.index("Data unavailable")
        tail = text[idx + len("Data unavailable"):].strip(".* ")
        assert tail, (
            f"Round 78 / B1: bullet {text!r} appears to be a pure "
            f"stub but survived the filter."
        )
    # Specific substantive bullets must be present
    assert any("understaffed" in t for t in bullet_texts)
    assert any("Microsoft Teams" in t for t in bullet_texts)
    assert any("No active incidents" in t for t in bullet_texts)


def test_end_to_end_all_stubs_yields_zero_bullets() -> None:
    """If a section consists entirely of stubs, the DOCX has zero
    List Bullet paragraphs and the heading still renders -- the
    filter never produces an exception or empty paragraph artifact.
    """
    import docx

    import adoptiq_backend

    markdown = """## All-empty section

* Industry Benchmarking: Data unavailable.
* Technical Competency: Data unavailable.
* Adoption Velocity: Data unavailable.
"""

    doc = docx.Document()
    adoptiq_backend.append_to_word_report(
        doc, markdown, heading="R78 / B1 all stubs"
    )

    bullet_paras = [
        p
        for p in doc.paragraphs
        if p.style and p.style.name == "List Bullet"
    ]
    assert len(bullet_paras) == 0, (
        f"Round 78 / B1: all-stub section should produce 0 List "
        f"Bullet paragraphs, got {len(bullet_paras)}: "
        f"{[p.text for p in bullet_paras]!r}"
    )
    # Section heading still rendered
    headings = [
        p.text
        for p in doc.paragraphs
        if p.style and p.style.name and "Heading" in p.style.name
    ]
    assert any("All-empty section" in h for h in headings), (
        "Round 78 / B1: section heading should survive even when all "
        "bullets are filtered."
    )


def test_end_to_end_bullet_prefix_variants_all_filter() -> None:
    """``* ``, ``- ``, and ``• `` markdown prefixes all enter the
    same bullet branch.  Each must drop the stub equally.
    """
    import docx

    import adoptiq_backend

    markdown = """## Prefix variants

* Industry Benchmarking: Data unavailable.
- Technical Competency: Data unavailable.
• Adoption Velocity: Data unavailable.
* Real bullet stays.
"""

    doc = docx.Document()
    adoptiq_backend.append_to_word_report(
        doc, markdown, heading="R78 / B1 prefix variants"
    )

    bullet_paras = [
        p
        for p in doc.paragraphs
        if p.style and p.style.name == "List Bullet"
    ]
    assert len(bullet_paras) == 1, (
        f"Round 78 / B1: 3 stub variants and 1 real bullet -> "
        f"expected 1 surviving List Bullet, got {len(bullet_paras)}: "
        f"{[p.text for p in bullet_paras]!r}"
    )
    assert "Real bullet stays" in bullet_paras[0].text


# ---------------------------------------------------------------------------
# Defense-in-depth — make sure infinite-loop-on-stub regression is
# pinned (the while-loop body without i += 1 would hang the writer).
# ---------------------------------------------------------------------------


def test_filter_does_not_hang_when_first_line_is_stub() -> None:
    """The filter increments ``i`` before ``continue`` so the
    writer never spins on the first stub.  This test would hang
    indefinitely if the increment were dropped -- we add an explicit
    timeout via signal on POSIX (skip on Windows) for belt-and-
    braces against the regression class.
    """
    import signal
    import sys

    import docx

    import adoptiq_backend

    markdown = "* Industry Benchmarking: Data unavailable.\n* Real bullet."
    doc = docx.Document()

    if sys.platform != "win32":
        def _handler(_signum, _frame):
            raise TimeoutError(
                "Round 78 / B1: append_to_word_report hung -- the "
                "i += 1 increment is missing from the stub-skip path."
            )

        old = signal.signal(signal.SIGALRM, _handler)
        signal.alarm(5)
        try:
            adoptiq_backend.append_to_word_report(
                doc, markdown, heading="R78 / B1 hang guard"
            )
        finally:
            signal.alarm(0)
            signal.signal(signal.SIGALRM, old)
    else:
        adoptiq_backend.append_to_word_report(
            doc, markdown, heading="R78 / B1 hang guard"
        )

    bullet_paras = [
        p
        for p in doc.paragraphs
        if p.style and p.style.name == "List Bullet"
    ]
    assert len(bullet_paras) == 1
    assert "Real bullet" in bullet_paras[0].text
