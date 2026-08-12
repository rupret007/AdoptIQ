"""Round 148 integration repair regressions."""

from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from docx import Document

import ask_ai_grounded as ask_ai
import canonical_report_adapter as adapter
import decision_report_delivery as delivery
from canonical_report_adapter import CanonicalReportAdapterError
from compact_report_formatter import calculate_renewal_risk_scores
from decision_report_delivery import _build_risk_profiles, _customer_incidents
from report_iteration_loop import compare_kpi_parity, extract_docx_kpis, extract_xlsx_kpis
from risk_scoring import _score_customer_pulse


def test_mixed_case_evidence_ids_share_one_public_citation_identity() -> None:
    # Round 148: live CSConsole IDs are often mixed case while the validator
    # canonicalizes model citations to uppercase.
    records = [
        {
            "source_id": "action-plan:mixedCase-1",
            "source_type": "ActionPlan",
            "text": "Status Open",
        }
    ]

    allowed_ids = ask_ai._r148_exact_allowed_source_ids(records)  # noqa: SLF001
    answer, rejected = ask_ai.compose_grounded_answer(
        {
            "executive_summary": "",
            "claims": [
                {
                    "statement": "Status Open",
                    "citations": ["ACTION-PLAN:MIXEDCASE-1"],
                }
            ],
            "actions": [],
            "unknowns": [],
        },
        allowed_ids,
        evidence_records=records,
    )

    assert allowed_ids == {"ACTION-PLAN:MIXEDCASE-1"}
    assert rejected == 0
    assert "[Sources: ACTION-PLAN:MIXEDCASE-1]" in answer


def test_public_evidence_rows_use_validator_citation_identity() -> None:
    import app_simple as public_ai

    rows = public_ai._r147_public_ai_evidence_rows(  # noqa: SLF001
        [{"source_id": "action-plan:mixedCase-1", "text": "Status Open"}]
    )

    assert rows[0]["source_id"] == "ACTION-PLAN:MIXEDCASE-1"


def test_evidence_lookup_matches_canonicalized_source_id(
    client,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import app_simple as public_ai

    monkeypatch.setattr(
        public_ai,
        "_get_ask_ai_query_diag",
        lambda _query_id: {
            "_r74_evidence_records": [
                {
                    "source_id": "action-plan:mixedCase-1",
                    "source_type": "ActionPlan",
                    "text": "Status Open",
                }
            ]
        },
    )
    monkeypatch.setattr(
        public_ai,
        "_r71_diag_rate_limit_check",
        lambda _client_ip: (True, 0),
    )

    response = client.get(
        "/api/ask-ai/evidence/query-1/ACTION-PLAN:MIXEDCASE-1"
    )

    assert response.status_code == 200
    assert response.get_json()["record"]["source_id"] == "ACTION-PLAN:MIXEDCASE-1"


def test_action_plan_evidence_carries_decision_fields() -> None:
    frame = pd.DataFrame(
        [
            {
                "ID": "mixedCase-1",
                "BU_NAME": "Acme Corporation",
                "SUBJECT_C": "Complete rollout review",
                "STATUS_C": "Open",
                "PRIORITY_C": "High",
                "DUE_DATE_C": "2026-08-12",
                "OWNER_NAME_C": "Assigned Owner",
            }
        ]
    )

    records, _ids = ask_ai._portfolio_records_from_payload(  # noqa: SLF001
        {"csconsole_action_plans": frame}
    )
    action_plan = next(record for record in records if record.source_type == "ActionPlan")

    assert action_plan.customer == "Acme Corporation"
    assert "Priority: High" in action_plan.text
    assert "Due Date: 2026-08-12" in action_plan.text
    assert "Owner: Assigned Owner" in action_plan.text


def test_entity_detection_does_not_treat_first_word_alias_as_identity() -> None:
    records = [
        {
            "source_id": "AP-1",
            "source_type": "ActionPlan",
            "customer": "Acme Corporation",
            "text": "All action plans are Open",
        },
        {
            "source_id": "SUB-1",
            "source_type": "Subscription",
            "technology": "All Contact Center",
            "text": "Subscription record",
        },
    ]

    assert ask_ai._r147_claim_supported_by_citations(  # noqa: SLF001
        {"statement": "All action plans are Open", "citations": ["AP-1"]},
        records,
        set(),
    )


def test_rejected_model_paraphrase_falls_back_to_exact_cited_row() -> None:
    records = [
        {
            "source_id": "AP-mixedCase-1",
            "source_type": "ActionPlan",
            "customer": "Acme Corporation",
            "timestamp": "2026-08-05",
            "text": "Status: Open | Priority: High",
        }
    ]
    answer, rejected = ask_ai.compose_grounded_answer(
        {
            "executive_summary": "",
            "claims": [
                {
                    "statement": "Acme Corporation currently needs urgent attention",
                    "citations": ["AP-MIXEDCASE-1"],
                }
            ],
            "actions": [],
            "unknowns": [],
        },
        {"AP-MIXEDCASE-1"},
        evidence_records=records,
    )

    assert rejected == 0
    assert "currently needs urgent attention" not in answer
    assert "[ActionPlan] Customer: Acme Corporation" in answer
    assert "[Sources: AP-MIXEDCASE-1]" in answer


def test_compose_bootstraps_citations_when_all_claims_rejected() -> None:
    records = [
        {
            "source_id": "AP-1",
            "source_type": "ActionPlan",
            "customer": "Acme Corporation",
            "text": "Status: Open",
        },
        {
            "source_id": "AB-2",
            "source_type": "AdoptionBarrier",
            "customer": "Beta Industries",
            "text": "Severity: High",
        },
    ]
    answer, _rejected = ask_ai.compose_grounded_answer(
        {
            "executive_summary": "Portfolio totals are elevated.",
            "claims": [
                {
                    "statement": "There are exactly 44 customers in scope",
                    "citations": ["AP-1"],
                }
            ],
            "actions": [],
            "unknowns": [],
        },
        {"AP-1", "AB-2"},
        evidence_records=records,
        evidence_bootstrap=True,
    )

    assert "[Sources: AP-1]" in answer
    assert "Supported Findings" in answer


def test_compose_skips_bootstrap_for_forced_gap_only_payloads() -> None:
    gap_answer, _rejected = ask_ai.compose_grounded_answer(
        {
            "executive_summary": "",
            "claims": [],
            "actions": [],
            "unknowns": [
                "The requested source and period are absent from the retrieved "
                "evidence; no estimate or substitute was used."
            ],
        },
        {"AP-1"},
        evidence_records=[
            {
                "source_id": "AP-1",
                "source_type": "ActionPlan",
                "text": "Status: Open",
            }
        ],
        evidence_bootstrap=False,
    )

    assert "Supported Findings" not in gap_answer
    assert "Evidence Gaps" in gap_answer
    assert "[Sources:" not in gap_answer


def test_forced_gap_uses_turn_question_not_conversation_prefix() -> None:
    follow_up = (
        "Using the prior answer only as conversational context, identify the one "
        "decision that should be made first and support it with current source IDs."
    )
    history_blob = (
        "Conversation context\n"
        "Q: State the exact customer, subscription, open Action Plan, open adoption "
        "barrier, and TAC case counts for this scope.\n"
        "A: Canonical metric total_customers: 44."
    )
    polluted = f"{history_blob}\n\nCurrent question: {follow_up}"
    records = [{"source_type": "ActionPlan", "source_id": "AP-1"}]

    assert ask_ai._detect_case_search_intent(polluted)
    assert not ask_ai._detect_case_search_intent(follow_up)
    assert ask_ai.build_retrieval_plan(follow_up)["intent"] == "default"
    assert ask_ai._r148_forced_evidence_gap(
        follow_up,
        records,
        retrieval_intent="default",
    ) is None
    assert ask_ai._r148_forced_evidence_gap(
        "List every support case for this scope.",
        records,
        retrieval_intent="case_search_enumeration",
    ) is not None


def test_missing_requested_source_forces_gap_without_substitution() -> None:
    evidence = [
        {
            "source_id": "AP-1",
            "source_type": "ActionPlan",
            "text": "Status: Open",
        }
    ]

    assert ask_ai._r148_forced_evidence_gap(  # noqa: SLF001
        "Find support cases that mention authentication.",
        evidence,
        retrieval_intent="case_search_enumeration",
    ) == (
        "support_case_source_absent",
        "Support-case evidence is absent from the retrieved evidence for "
        "this scope; no substitute source was used.",
    )
    assert ask_ai._r148_forced_evidence_gap(  # noqa: SLF001
        "Find support cases that mention authentication.",
        [*evidence, {"source_id": "CASE-1", "source_type": "SupportCase"}],
        retrieval_intent="case_search_enumeration",
    ) is None


def test_explicit_unanswerable_request_forces_gap() -> None:
    gap = ask_ai._r148_forced_evidence_gap(  # noqa: SLF001
        "Give an exact score for a source not present in the retrieved evidence.",
        [],
    )

    assert gap is not None
    assert gap[0] == "explicit_missing_evidence"
    assert "no estimate or substitute" in gap[1]


def test_tac_count_question_does_not_force_missing_support_case_gap() -> None:
    gap = ask_ai._r148_forced_evidence_gap(  # noqa: SLF001
        "State the exact customer, subscription, and TAC case counts.",
        [{"source_id": "TAC-1", "source_type": "CSOne"}],
        retrieval_intent="case_search_enumeration",
    )

    assert gap is None


def _renewal_facts() -> dict[str, list[list[object]]]:
    return {
        "account_summary_all": [
            ["Acme Corporation", "LOW", 27.2, 0, 0, 0, 0],
            ["Beta Industries", "LOW", 23.9, 0, 0, 0, 0],
            ["Gamma Public Sector", "HIGH", 59.8, 0, 0, 0, 0],
        ]
    }


def test_current_renewal_display_and_canonical_scores_reconcile() -> None:
    # Round 148: 0-10 display aliases may differ by up to 0.5 after scaling
    # because the workbook publishes them at one decimal beside the 0-100 fact.
    frame = pd.DataFrame(
        [
            {
                "Customer": "Acme Corporation",
                "Overall_Risk_Score": 2.7,
                "Risk_Score_0_10": 2.7,
                "Risk_Score_0_100": 27.2,
                "Risk_Band": "LOW",
            },
            {
                "Customer": "Beta Industries",
                "Overall_Risk_Score": 2.4,
                "Risk_Score_0_10": 2.4,
                "Risk_Score_0_100": 23.9,
                "Risk_Band": "LOW",
            },
            {
                "Customer": "Gamma Public Sector",
                "Overall_Risk_Score": 6.0,
                "Risk_Score_0_10": 6.0,
                "Risk_Score_0_100": 59.8,
                "Risk_Band": "HIGH",
            },
        ]
    )

    claims = adapter._validate_family_risk_claims(
        {"Renewal_Summary": frame},
        _renewal_facts(),
        info={},
        scope_type="team",
        scope_value="Entire team",
    )

    acme_scores = {
        claim["value"]
        for claim in claims
        if claim["customer"] == "Acme Corporation" and claim["kind"] == "score"
    }
    assert acme_scores == {27.0, 27.2}


def test_current_renewal_explicit_0_100_contradiction_fails_closed() -> None:
    frame = pd.DataFrame(
        [
            {
                "Customer": "Acme Corporation",
                "Overall_Risk_Score": 2.7,
                "Risk_Score_0_10": 2.7,
                "Risk_Score_0_100": 82.0,
                "Risk_Band": "LOW",
            }
        ]
    )

    with pytest.raises(
        CanonicalReportAdapterError,
        match=r"Risk_Score_0_100=82\.0.*canonical Account_Summary",
    ):
        adapter._validate_family_risk_claims(
            {"Renewal_Summary": frame},
            _renewal_facts(),
            info={},
            scope_type="team",
            scope_value="Entire team",
        )


def test_legacy_renewal_overall_score_without_scale_fields_stays_0_100() -> None:
    assert (
        adapter._risk_score(
            2.7,
            "Overall_Risk_Score",
            sheet_name="Renewal_Summary",
            record={"Overall_Risk_Score": 2.7},
        )
        == 2.7
    )


@pytest.mark.parametrize("explicit_ten", ["Unavailable (Partial)", float("nan")])
def test_unavailable_renewal_ten_point_sibling_does_not_force_scale(
    explicit_ten: object,
) -> None:
    assert (
        adapter._risk_score(
            27.2,
            "Overall_Risk_Score",
            sheet_name="Renewal_Summary",
            record={
                "Overall_Risk_Score": 27.2,
                "Risk_Score_0_10": explicit_ten,
                "Risk_Score_0_100": 27.2,
            },
        )
        == 27.2
    )


def test_friendly_pulse_headers_preserve_numeric_risk_signal() -> None:
    raw = pd.DataFrame({"SCORE__C": [2], "PULSE_RATING__C": ["Fair"]})
    exported = pd.DataFrame({"Pulse Score": [2], "Pulse Rating": ["Fair"]})

    raw_result = _score_customer_pulse(raw)
    exported_result = _score_customer_pulse(exported)
    assert exported_result["score"] == raw_result["score"]
    assert exported_result["details"]["count"] == raw_result["details"]["count"]
    assert exported_result["details"]["poor_bad_count"] == raw_result["details"]["poor_bad_count"]


def test_compact_risk_collapses_cross_subscription_tac_fanout() -> None:
    unique_case = pd.DataFrame(
        [
            {
                "customer_name": "Acme Corporation",
                "SR Number": "700000001",
                "Severity": "P1",
                "Case Status": "Open",
                "Date/Time Opened": "2026-07-01",
            }
        ]
    )
    fanned_cases = pd.concat([unique_case, unique_case], ignore_index=True)

    unique = calculate_renewal_risk_scores(
        pd.DataFrame(),
        unique_case,
        recent_window_days=90,
    )
    fanned = calculate_renewal_risk_scores(
        pd.DataFrame(),
        fanned_cases,
        recent_window_days=90,
    )

    assert fanned["Acme Corporation"]["score"] == unique["Acme Corporation"]["score"]
    assert fanned["Acme Corporation"]["risk_band"] == unique["Acme Corporation"]["risk_band"]


def test_decision_risk_keeps_untagged_portfolio_incidents_context_only() -> None:
    frames = {
        "subscriptions": pd.DataFrame([{"Customer Name": "Acme Corporation"}]),
        "action_plans": pd.DataFrame(),
        "adoption_barriers": pd.DataFrame(),
        "customer_pulse": pd.DataFrame(),
        "tac_cases": pd.DataFrame(),
        "success_priorities": pd.DataFrame(),
    }
    incident = {
        "id": "INC-1",
        "status": "active",
        "severity": "high",
        "title": "Portfolio service incident",
    }

    without_incident = _build_risk_profiles(
        frames,
        days=90,
        as_of="2026-08-05T00:00:00Z",
    )
    with_incident = _build_risk_profiles(
        frames,
        days=90,
        as_of="2026-08-05T00:00:00Z",
        external_incidents=[incident],
    )

    missing_component = without_incident["Acme Corporation"]["components"]["incidents"]
    context_only_component = with_incident["Acme Corporation"]["components"]["incidents"]
    assert missing_component["score"] is None
    assert missing_component["details"]["data_state"] == "missing"
    assert context_only_component["score"] == 0.0
    assert context_only_component["details"]["count"] == 0


def test_customer_incident_filter_honors_alias_registry() -> None:
    incidents = [
        {
            "id": "INC-NYU",
            "customer_name": "NYU MEDICAL CENTER",
            "status": "active",
        },
        {
            "id": "INC-OTHER",
            "customer_name": "Unrelated Customer",
            "status": "active",
        },
    ]

    selected = _customer_incidents(incidents, "NYU LANGONE HEALTH SYSTEMS")

    assert [record["id"] for record in selected] == ["INC-NYU"]


def _decision_team(
    *,
    subscriptions: pd.DataFrame,
    action_plans: pd.DataFrame | None = None,
    adoption_barriers: pd.DataFrame | None = None,
) -> dict[str, dict[str, pd.DataFrame]]:
    return {
        "Synthetic Member": {
            "subscriptions": subscriptions,
            "action_plans": action_plans if action_plans is not None else pd.DataFrame(),
            "adoption_barriers": (
                adoption_barriers if adoption_barriers is not None else pd.DataFrame()
            ),
            "customer_pulse": pd.DataFrame(),
            "tac_cases": pd.DataFrame(),
            "success_priorities": pd.DataFrame(),
        }
    }


def test_unknown_subscription_id_is_disclosed_missing_not_unpublished() -> None:
    facts = delivery.build_report_facts(
        _decision_team(
            subscriptions=pd.DataFrame(
                [{"SUBSCRIPTION_ID": "Unknown", "BU_NAME": "Acme Corporation"}]
            )
        ),
        report_type="Leader",
        scope_type="team",
        scope_value="Synthetic team",
        manager_name="Synthetic Manager",
        days=90,
        as_of="2026-08-05T00:00:00Z",
    )
    sheets = delivery.build_source_data_sheets(facts)

    subscriptions = sheets["Subscriptions"]
    contract = delivery.validate_cross_artifact_contract(facts, sheets)

    assert subscriptions.loc[0, "Record_ID"] == ""
    assert subscriptions.loc[0, "Record_ID_Data_Quality"] == "Missing stable source ID"
    assert "Subscriptions contains a source ID without public Record_ID" not in contract["errors"]
    assert contract["ok"], contract["errors"]


def test_activity_trend_evidence_links_only_in_window_part_of_boundary_week() -> None:
    facts = delivery.build_report_facts(
        _decision_team(
            subscriptions=pd.DataFrame(
                [{"SUBSCRIPTION_ID": "SUB-1", "BU_NAME": "Acme Corporation"}]
            ),
            action_plans=pd.DataFrame(
                [
                    {"ID": "AP-BEFORE", "BU_NAME": "Acme Corporation", "CREATED_DATE_C": "2026-05-06"},
                    {"ID": "AP-IN-1", "BU_NAME": "Acme Corporation", "CREATED_DATE_C": "2026-05-08"},
                    {"ID": "AP-IN-2", "BU_NAME": "Acme Corporation", "CREATED_DATE_C": "2026-05-09"},
                ]
            ),
        ),
        report_type="Leader",
        scope_type="team",
        scope_value="Synthetic team",
        manager_name="Synthetic Manager",
        days=90,
        as_of="2026-08-05T00:00:00Z",
    )
    sheets = delivery.build_source_data_sheets(facts)
    metric_key = "chart.activity_trend.action_plans.2026-05-04"

    chart_row = facts["chart_data"].loc[
        facts["chart_data"]["Metric_Key"].eq(metric_key)
    ].iloc[0]
    evidence = sheets["Evidence_Links"].loc[
        sheets["Evidence_Links"]["Evidence_Key"].eq(metric_key)
    ]
    contract = delivery.validate_cross_artifact_contract(facts, sheets)

    assert chart_row["Value"] == 2
    assert len(evidence) == 2
    assert contract["ok"], contract["errors"]


def test_decimal_export_cells_keep_evidence_fingerprint_after_xlsx_roundtrip(
    tmp_path,
) -> None:
    facts = delivery.build_report_facts(
        _decision_team(
            subscriptions=pd.DataFrame(
                [{"SUBSCRIPTION_ID": "SUB-1", "BU_NAME": "Acme Corporation"}]
            ),
            adoption_barriers=pd.DataFrame(
                [
                    {
                        "ID": "AB-1",
                        "BU_NAME": "Acme Corporation",
                        "AOV_C": Decimal("123456.789012345678"),
                    }
                ]
            ),
        ),
        report_type="Leader",
        scope_type="team",
        scope_value="Synthetic team",
        manager_name="Synthetic Manager",
        days=90,
        as_of="2026-08-05T00:00:00Z",
    )
    profile = next(iter(facts["risk_profiles"].values()))
    profile["components"]["action_plans"]["score"] = 25.333333333333336
    sheets = delivery.build_source_data_sheets(facts)
    output_path = delivery.write_source_data_workbook(
        tmp_path / "leader_source.xlsx",
        sheets,
    )

    contract = delivery.validate_written_source_workbook(output_path, facts)

    assert contract["ok"], contract["errors"]
    workbook = openpyxl.load_workbook(output_path, read_only=True)
    try:
        worksheet = workbook["Adoption_Barriers"]
        headers = [cell.value for cell in next(worksheet.iter_rows())]
        value_index = headers.index("Annual Order Value")
        exported_value = next(worksheet.iter_rows(min_row=2))[value_index].value
        risk_worksheet = workbook["Risk_Components"]
        risk_headers = [cell.value for cell in next(risk_worksheet.iter_rows())]
        component_index = risk_headers.index("Component")
        score_index = risk_headers.index("Component_Score_0_100")
        action_plan_score = next(
            row[score_index].value
            for row in risk_worksheet.iter_rows(min_row=2)
            if row[component_index].value == "action_plans"
        )
    finally:
        workbook.close()
    assert exported_value == "123456.789012345678"
    assert action_plan_score == pytest.approx(25.333333333333336)


def test_source_data_export_strips_external_incident_html_before_fingerprinting(
    tmp_path,
) -> None:
    facts = delivery.build_report_facts(
        _decision_team(
            subscriptions=pd.DataFrame(
                [{"SUBSCRIPTION_ID": "SUB-1", "BU_NAME": "Acme Corporation"}]
            )
        ),
        report_type="Leader",
        scope_type="team",
        scope_value="Synthetic team",
        manager_name="Synthetic Manager",
        days=90,
        as_of="2026-08-05T00:00:00Z",
        external_incidents=[
            {
                "id": "INC-1",
                "title": "Service incident",
                "description": "<p>Service <strong>degradation</strong></p>",
                "status": "active",
            }
        ],
    )
    sheets = delivery.build_source_data_sheets(facts)
    output_path = delivery.write_source_data_workbook(
        tmp_path / "leader_source.xlsx",
        sheets,
    )

    contract = delivery.validate_written_source_workbook(output_path, facts)
    assert contract["ok"], contract["errors"]
    workbook = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
    try:
        worksheet = workbook["External_Incidents"]
        headers = [cell.value for cell in next(worksheet.iter_rows())]
        description_index = headers.index("description")
        description = next(worksheet.iter_rows(min_row=2))[description_index].value
    finally:
        workbook.close()

    assert description == "Service degradation"
    assert "<" not in description


def test_concise_leader_flow_does_not_append_uncontracted_be_table() -> None:
    source = (Path(__file__).resolve().parents[1] / "app_simple.py").read_text(
        encoding="utf-8"
    )
    marker = "# Round 148: the canonical Leader artifact is the concise decision"
    region = source[source.index(marker) : source.index(marker) + 5000]

    assert "_r142_leader_concise_default = True" in region
    assert "not _r142_leader_concise_default" in region
    assert "add_be_priority_focus_areas_section(" in region
    assert "omitted from the concise Word contract" in region


def test_partial_decision_report_parity_compares_explicit_withholding(tmp_path) -> None:
    docx_path = tmp_path / "partial.docx"
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Customers"
    table.rows[1].cells[1].text = "Unavailable (Partial)"
    table.rows[2].cells[0].text = "Action Plans"
    table.rows[2].cells[1].text = "Unavailable (Partial)"
    doc.save(docx_path)

    xlsx_path = tmp_path / "partial.xlsx"
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Customers", 1])
    summary.append(["Action Plans", 3])
    lineage = workbook.create_sheet("Metric_Lineage")
    lineage.append(["Metric_Key", "Display_Label", "Metric_Value", "Source_State"])
    lineage.append(
        ["kpi.customers", "Customers", "Unavailable (Partial)", "partial"]
    )
    lineage.append(["kpi.action_plans_total", "Action Plans", None, "partial"])
    workbook.save(xlsx_path)
    workbook.close()

    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(docx_kpis, xlsx_kpis, strict=True)

    assert docx_kpis["withheld_kpis"] == ["action_plans", "total_customers"]
    assert xlsx_kpis["values"] == {}
    assert xlsx_kpis["withheld_kpis"] == ["action_plans", "total_customers"]
    assert parity.passed is True
    assert parity.details["reason"] == "compared_withheld_kpis"


def test_partial_decision_report_parity_rejects_value_availability_mismatch() -> None:
    parity = compare_kpi_parity(
        {"values": {"total_customers": "1"}, "withheld_kpis": []},
        {"values": {}, "withheld_kpis": ["total_customers"]},
        strict=True,
    )

    assert parity.passed is False
    assert parity.details["reason"] == "kpi_availability_mismatch"
    assert parity.details["availability_mismatches"] == ["total_customers"]


def test_partial_parity_allows_format_specific_withheld_kpis() -> None:
    parity = compare_kpi_parity(
        {
            "values": {},
            "withheld_kpis": ["action_plans", "incidents", "risk_score"],
        },
        {
            "values": {},
            "withheld_kpis": ["action_plans", "open_action_plans"],
        },
        strict=True,
    )

    assert parity.passed is True
    assert parity.details["reason"] == "compared_withheld_kpis"
    assert parity.details["common_withheld_kpis"] == ["action_plans"]
