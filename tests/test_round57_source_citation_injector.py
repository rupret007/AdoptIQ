"""Round 57 / Phase B -- pin the post-render Word source-citation injector.

The injector lives in ``report_source_injector.py`` and runs after each
of the four report-generation workers in ``app_simple.py`` finalizes its
.docx.  These tests exercise the injector against synthetic Word
documents that reproduce the gate's three failure modes (table KPI
without adjacent citation, paragraph KPI without inline citation,
narrative paragraph with short numeric tokens) and confirm:

1. The injector adds source chrome where the gate would otherwise
   complain: a generic ``[Source: AdoptIQ Report Data Sources]`` fallback
   for mixed/unknown KPI lines and, after R82, per-system chrome for
   canonical single-source KPI lines.
2. The injection is idempotent (running twice changes nothing on the
   second pass).
3. Already-cited content is preserved untouched (writers that called
   ``format_inline_source(...)`` directly keep their richer citation;
   the injector never overwrites or duplicates it).
4. Metadata paragraphs (``"Generated: ..."``, ``"Build 31"``, etc.) are
   skipped -- the gate's ``_numeric_tokens_requiring_source`` helper
   excludes them and the injector mirrors that behavior.
5. The supervisor-side gate (``evaluate_report_quality`` from
   ``report_iteration_loop``) flips ``quality.passed`` from ``False`` to
   ``True`` after the injector runs on a synthetic minimum-shape report.

The tests build the .docx in a tmp_path so they cost nothing in CI and
do not depend on the live Build31 .app.

Round 57 / Phase B.  Made-with: Cursor.
"""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402

from report_source_injector import (  # noqa: E402
    inject_source_citations_into_docx,
    _SOURCE_TOKEN_RE,
    _numeric_tokens_requiring_source,
)


def _build_minimum_shape_docx(path: Path) -> None:
    """Create a small .docx that triggers each of the gate's failure modes."""

    doc = Document()
    doc.add_heading("Test Report", level=1)
    # Metadata paragraph -- must NOT be touched.
    doc.add_paragraph("Generated: 2026-04-29 23:55:00 UTC")
    doc.add_paragraph("Build 31")
    # Narrative paragraph carrying short numeric tokens (should be cited).
    doc.add_paragraph("We saw 23% growth in Q3 with 14 escalations and 8 cases.")
    # Single label:value paragraph metric claim.
    doc.add_paragraph("Total Customers: 52")
    # Already-cited paragraph -- must NOT be re-cited.
    doc.add_paragraph(
        "Total Adoption Barriers: 68 [Source: CSConsole / Snowflake C360_CS_TASK_C_VW]"
    )
    # Two-column "label | value" KPI table.
    table_a = doc.add_table(rows=3, cols=2)
    table_a.rows[0].cells[0].text = "Metric"
    table_a.rows[0].cells[1].text = "Value"
    table_a.rows[1].cells[0].text = "Total Customers"
    table_a.rows[1].cells[1].text = "52"
    table_a.rows[2].cells[0].text = "Adoption Barriers"
    table_a.rows[2].cells[1].text = "68"
    # Multi-column header KPI table.
    table_b = doc.add_table(rows=2, cols=4)
    table_b.rows[0].cells[0].text = "Team Member"
    table_b.rows[0].cells[1].text = "Total Customers"
    table_b.rows[0].cells[2].text = "Adoption Barriers"
    table_b.rows[0].cells[3].text = "Support Cases"
    table_b.rows[1].cells[0].text = "Brian Frazier"
    table_b.rows[1].cells[1].text = "52"
    table_b.rows[1].cells[2].text = "68"
    table_b.rows[1].cells[3].text = "381"
    doc.save(str(path))


def _doc_full_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def test_injector_returns_structured_count_dict_on_minimum_shape(tmp_path: Path) -> None:
    """Sanity: the injector reports separate paragraph/table/skip/error counts."""

    docx = tmp_path / "min_shape.docx"
    _build_minimum_shape_docx(docx)

    counts = inject_source_citations_into_docx(docx)

    assert set(counts.keys()) == {
        "paragraphs_injected",
        "table_cells_injected",
        "table_captions_added",  # Round 114 / Build 83
        "skipped_already_cited",
        "skipped_no_numeric",
        "errors",
    }
    assert counts["errors"] == 0
    assert counts["paragraphs_injected"] >= 2, (
        "injector must cite the narrative + label:value paragraphs that "
        "carry short numeric KPI tokens; got "
        f"{counts['paragraphs_injected']}"
    )
    # Round 114 / Build 83: the two-column ``Metric | Value`` table still
    # gets per-row cell citations (2 rows), while the multi-column matrix
    # is now de-cluttered to a single trailing caption instead of one
    # citation per numeric cell.
    assert counts["table_cells_injected"] >= 2, (
        "injector must cite the metric value cells in the two-column KPI "
        f"table; got {counts['table_cells_injected']}"
    )
    assert counts["table_captions_added"] == 1, (
        "the multi-column matrix must receive exactly ONE trailing source "
        f"caption (not per-cell citations); got {counts['table_captions_added']}"
    )
    assert counts["skipped_already_cited"] >= 1, (
        "the pre-cited 'Total Adoption Barriers: 68 [Source: ...]' paragraph "
        "must register as skipped, not silently re-cited."
    )


def test_injector_is_idempotent_on_second_pass(tmp_path: Path) -> None:
    """Running twice must leave the document byte-stable on the second pass.

    The injector skips the save when no new citations would land, so the
    file's mtime is preserved across no-op re-invocations.  This is the
    contract that lets per-worker callers wrap the injector in a try/except
    and re-run it on retry without doubling citations.
    """

    docx = tmp_path / "idem.docx"
    _build_minimum_shape_docx(docx)

    first = inject_source_citations_into_docx(docx)
    after_first_text = _doc_full_text(docx)
    first_mtime = docx.stat().st_mtime

    second = inject_source_citations_into_docx(docx)
    after_second_text = _doc_full_text(docx)
    second_mtime = docx.stat().st_mtime

    assert first["paragraphs_injected"] > 0 or first["table_cells_injected"] > 0
    assert second["paragraphs_injected"] == 0
    assert second["table_cells_injected"] == 0
    assert after_first_text == after_second_text
    # Round 57: idempotent skip-save preserves mtime.  If this regresses,
    # downstream sha256 + size_bytes contracts in baseline manifests will
    # silently churn between identical runs.
    assert first_mtime == second_mtime, (
        "second pass must not re-save the document; mtime drifted from "
        f"{first_mtime} to {second_mtime}"
    )


def test_injector_preserves_pre_cited_paragraphs_unchanged(tmp_path: Path) -> None:
    """Writers that already emit format_inline_source must not see double citations."""

    docx = tmp_path / "pre_cited.docx"
    doc = Document()
    doc.add_paragraph(
        "Total Adoption Barriers: 68 [Source: CSConsole / Snowflake C360_CS_TASK_C_VW; "
        "Field(s): ID; Verification: Query by Record ID in CSConsole or Snowflake]"
    )
    doc.save(str(docx))

    counts = inject_source_citations_into_docx(docx)

    assert counts["paragraphs_injected"] == 0
    assert counts["table_cells_injected"] == 0
    assert counts["skipped_already_cited"] >= 1

    text = _doc_full_text(docx)
    # Exactly ONE citation in the paragraph (the writer's original).
    assert len(_SOURCE_TOKEN_RE.findall(text)) == 1, (
        "injector must not append a fallback citation when an inline citation already exists; "
        f"text now reads: {text!r}"
    )


def test_injector_skips_metadata_paragraphs(tmp_path: Path) -> None:
    """Generated/Build/Version/Page lines must never receive a citation."""

    docx = tmp_path / "metadata.docx"
    doc = Document()
    metadata_lines = [
        "Generated: 2026-04-29 23:55:00 UTC",
        "Build 31",
        "Version 1.0.4",
        "Analysis ID: brian_frazier_2026-04-29",
        "Page 1 of 17",
        "Report metadata: see appendix",
    ]
    for line in metadata_lines:
        doc.add_paragraph(line)
    doc.save(str(docx))

    counts = inject_source_citations_into_docx(docx)
    assert counts["paragraphs_injected"] == 0, (
        "metadata paragraphs must be skipped per the gate's "
        "_numeric_tokens_requiring_source filter"
    )
    text = _doc_full_text(docx)
    assert not _SOURCE_TOKEN_RE.search(text), (
        "no metadata paragraph should have received a [Source: ...] citation"
    )


def test_injector_cites_metric_match_in_metadata_paragraph(tmp_path: Path) -> None:
    """Regression: metadata-shaped header paragraphs that ALSO contain a
    ``Label: number`` metric claim must still receive an inline citation
    for the metric claim.

    Reproduces the renewal-report header line found in supervisor pass #2
    (``"Generated: ... | Analysis Period: 90 days | ..."``) where the
    paragraph trips the metadata short-circuit AND is extracted as a
    metric claim by ``_extract_docx_metric_claims``. The injector must
    bias toward the stricter check (cite every metric claim) so the
    ``unbacked_metric_claim_count`` gate stays at zero.
    """
    from report_iteration_loop import evaluate_report_quality

    docx = tmp_path / "metadata_metric.docx"
    doc = Document()
    doc.add_paragraph(
        "Report Type: Renewal Portfolio | Generated: April 30, 2026 at 05:21 UTC | "
        "Customer: All Managers's Portfolio | Technology: All Contact Center | "
        "Analysis Period: 90 days | AdoptIQ | Data sources: CSOne"
    )
    doc.save(str(docx))

    pre_payload, _ = evaluate_report_quality(
        docx_path=docx,
        xlsx_path=None,
        scenario_key="renewal",
        strict=True,
    )
    pre_unbacked = pre_payload.get("unbacked_metric_claim_count", 0)
    assert pre_unbacked >= 1, (
        "synthetic metadata+metric paragraph must surface as an unbacked "
        "metric claim BEFORE injection (otherwise this test isn't actually "
        "exercising the regression)"
    )

    counts = inject_source_citations_into_docx(docx, scenario_key="renewal")
    assert counts["paragraphs_injected"] == 1, (
        "the metadata+metric paragraph must receive an inline citation, "
        f"counts: {counts}"
    )

    text = _doc_full_text(docx)
    assert _SOURCE_TOKEN_RE.search(text), (
        "metadata+metric paragraph must contain a [Source: ...] citation "
        f"after injection; text: {text!r}"
    )

    post_payload, _ = evaluate_report_quality(
        docx_path=docx,
        xlsx_path=None,
        scenario_key="renewal",
        strict=True,
    )
    post_unbacked = post_payload.get("unbacked_metric_claim_count", 0)
    assert post_unbacked == 0, (
        "after injection no metric claim from the metadata paragraph should "
        f"remain unbacked; before={pre_unbacked} after={post_unbacked}"
    )


def test_injector_handles_canonical_match_with_noncanonical_followups(tmp_path: Path) -> None:
    """Regression: paragraph with one canonical KPI claim followed by
    several non-canonical ``Label: number`` matches (e.g. status
    breakdown bullets) must receive citations between every regex
    match.

    Reproduces the renewal-report Action Plans paragraph found in
    supervisor pass #3:

        Total Action Plans: 889
        Source: CSConsole.
        Status Breakdown:
          - Completed - Successful: 583
          - New Request: 132
          - On Track: 126
          ...

    The gate's ``_paragraph_claim_source_backed`` slices the segment
    from ``Total Action Plans: 889`` to the NEXT regex match
    (``Successful: 583``) -- a trailing-only citation lands AFTER that
    boundary and leaves ``Total Action Plans`` unbacked. The injector
    must interleave citations between every regex match when at least
    one match is canonical.
    """
    from report_iteration_loop import evaluate_report_quality

    docx = tmp_path / "action_plan_breakdown.docx"
    doc = Document()
    doc.add_paragraph(
        "Total Action Plans: 889\nSource: CSConsole.\nStatus Breakdown:\n"
        "  - Completed - Successful: 583\n  - New Request: 132\n"
        "  - On Track: 126\n  - Completed - Unsuccessful: 20\n"
        "  - On Hold: 11\n  - Closed - Cancelled: 11\n  - Off Trajectory: 6"
    )
    doc.save(str(docx))

    pre_payload, _ = evaluate_report_quality(
        docx_path=docx,
        xlsx_path=None,
        scenario_key="renewal",
        strict=True,
    )
    pre_unbacked = pre_payload.get("unbacked_metric_claim_count", 0)
    assert pre_unbacked >= 1, (
        "synthetic action-plans breakdown paragraph must surface as "
        "unbacked BEFORE injection (otherwise the regression isn't "
        "reachable)"
    )

    counts = inject_source_citations_into_docx(docx, scenario_key="renewal")
    assert counts["paragraphs_injected"] == 1, (
        "the multi-match paragraph must be rewritten, "
        f"counts: {counts}"
    )

    post_payload, _ = evaluate_report_quality(
        docx_path=docx,
        xlsx_path=None,
        scenario_key="renewal",
        strict=True,
    )
    post_unbacked = post_payload.get("unbacked_metric_claim_count", 0)
    assert post_unbacked == 0, (
        "after injection the canonical 'Total Action Plans' claim must "
        f"be source_backed; before={pre_unbacked} after={post_unbacked}"
    )


def test_injector_clears_quality_gate_on_minimum_shape_report(tmp_path: Path) -> None:
    """End-to-end: run the gate before + after, confirm quality.passed flips True."""

    pytest.importorskip("openpyxl")  # quality gate may inspect XLSX too
    from report_iteration_loop import evaluate_report_quality

    docx = tmp_path / "gate_check.docx"
    _build_minimum_shape_docx(docx)

    pre_payload, pre_gate = evaluate_report_quality(
        docx_path=docx,
        xlsx_path=None,
        scenario_key="comprehensive",
        strict=True,
    )
    assert pre_gate.passed is False, (
        "synthetic minimum-shape report must fail the gate BEFORE injection "
        "(otherwise this test isn't actually exercising the injector)"
    )
    pre_unbacked = pre_payload.get("details", {}).get("unbacked_metric_claim_count", 0)
    pre_uncited = pre_payload.get("details", {}).get("uncited_numeric_paragraph_count", 0)

    inject_source_citations_into_docx(docx, scenario_key="comprehensive")

    post_payload, post_gate = evaluate_report_quality(
        docx_path=docx,
        xlsx_path=None,
        scenario_key="comprehensive",
        strict=True,
    )
    post_details = post_payload.get("details", {})
    post_unbacked = post_details.get("unbacked_metric_claim_count", 0)
    post_uncited = post_details.get("uncited_numeric_paragraph_count", 0)

    assert post_unbacked == 0, (
        f"injector must source-back every metric claim the gate detects; "
        f"pre={pre_unbacked} post={post_unbacked}"
    )
    assert post_uncited == 0, (
        f"injector must source-back every paragraph carrying short numeric "
        f"tokens; pre={pre_uncited} post={post_uncited}"
    )
    assert post_gate.passed is True, (
        f"quality gate must pass after injection on the minimum-shape report; "
        f"errors={post_payload.get('details', {}).get('errors', [])}"
    )


def test_injector_handles_missing_file_without_raising(tmp_path: Path) -> None:
    """Wrapper contract: the injector NEVER raises into the caller."""

    nonexistent = tmp_path / "does_not_exist.docx"
    counts = inject_source_citations_into_docx(nonexistent)

    assert counts["errors"] == 1
    assert counts["paragraphs_injected"] == 0
    assert counts["table_cells_injected"] == 0


def test_app_simple_safe_wrapper_swallows_all_failures(tmp_path: Path, monkeypatch) -> None:
    """The safe wrapper in app_simple must NEVER propagate exceptions."""

    import app_simple

    # Force the underlying injector to raise.
    def _boom(*args, **kwargs):
        raise RuntimeError("synthetic boom")

    monkeypatch.setattr(app_simple, "_r57_inject_source_citations", _boom)

    # Call the safe wrapper -- must return None and not raise.
    result = app_simple._r57_inject_citations_safe(
        str(tmp_path / "fake.docx"), scenario_key="leader"
    )
    assert result is None


def test_multi_match_paragraph_back_cites_every_match(tmp_path: Path) -> None:
    """Paragraphs with multiple ``Label: number`` matches must back-cite each one.

    The gate's ``_paragraph_claim_source_backed`` checks the segment
    between THIS match and the NEXT match for ``[source:``. A
    trailing-only citation only satisfies the LAST match's segment;
    earlier matches stay unbacked. The injector must interleave
    citations between matches.
    """

    from report_iteration_loop import (
        evaluate_report_quality,
        _PARAGRAPH_KPI_NUMERIC_RE,
        _paragraph_claim_source_backed,
    )

    docx = tmp_path / "multi.docx"
    doc = Document()
    # Three metric matches in one paragraph -- previously only the last
    # one would have a [Source:] in its segment.
    doc.add_paragraph(
        "Total Customers: 52. Adoption Barriers: 68. Support Cases: 381."
    )
    doc.save(str(docx))

    inject_source_citations_into_docx(docx, scenario_key="leader")

    re_doc = Document(str(docx))
    text = re_doc.paragraphs[0].text
    matches = list(_PARAGRAPH_KPI_NUMERIC_RE.finditer(text))
    assert len(matches) == 3, (
        f"expected three Label:number matches, got {len(matches)} in {text!r}"
    )
    for idx in range(len(matches)):
        assert _paragraph_claim_source_backed(text, idx, matches), (
            f"match {idx} segment must contain [Source:] after injection; "
            f"text={text!r}"
        )

    # And the gate-level citation accounting must be clean (the only
    # remaining error in this synthetic doc is "no tables", which is a
    # separate scanability rule unrelated to source citations).
    _payload, gate = evaluate_report_quality(
        docx_path=docx, xlsx_path=None, scenario_key="leader", strict=True
    )
    details = gate.details if isinstance(gate.details, dict) else {}
    assert details.get("unbacked_metric_claim_count") == 0, (
        f"every match in the multi-match paragraph must be back-cited; "
        f"unbacked={details.get('unbacked_metric_claim_count')}"
    )
    assert details.get("uncited_numeric_paragraph_count") == 0, (
        f"the multi-match paragraph must satisfy the narrative-token gate; "
        f"uncited={details.get('uncited_numeric_paragraph_count')}"
    )
    assert details.get("source_citation_count", 0) >= 3, (
        f"injector must emit at least one citation per match; "
        f"source_citation_count={details.get('source_citation_count')}"
    )


def test_normalize_kpi_value_strips_injected_citation_chrome() -> None:
    """Parity gate must compare numeric tokens, not citation chrome."""

    from report_iteration_loop import _normalize_kpi_value, _is_numeric_kpi_value

    # Round 57: the injector appends ``[Source: ...]`` to KPI value cells
    # in the docx. The xlsx side has the bare number. Without strip, parity
    # gate would fire on every injected metric -- effectively reverting the
    # injector. The fix lives in _normalize_kpi_value + _is_numeric_kpi_value.
    cell_text = "68 [Source: AdoptIQ Report Data Sources]"
    assert _normalize_kpi_value(cell_text) == "68"
    assert _is_numeric_kpi_value(cell_text) is True
    # Sanity: bare values still normalize to themselves.
    assert _normalize_kpi_value("68") == "68"
    assert _is_numeric_kpi_value("68") is True
    # Multi-token chrome (the injector emits the same chrome multiple
    # times in multi-match paragraphs; cell-level injection only emits
    # once but be defensive against future writers).
    assert _normalize_kpi_value("68 [Source: A] [Source: B]") == "68"


def test_numeric_tokens_helper_matches_gate_semantics() -> None:
    """The injector's local mirror of _numeric_tokens_requiring_source must match the gate."""

    # Mirror cases pulled directly from the gate's own behavior.
    assert _numeric_tokens_requiring_source("") == []
    assert _numeric_tokens_requiring_source("   ") == []
    assert _numeric_tokens_requiring_source("Total Customers: 52") == ["52"]
    # Long numerics are treated as IDs, not KPI claims.
    assert _numeric_tokens_requiring_source("Customer ID: 12345678") == []
    # Mixed: short + long, only the short token survives.
    assert _numeric_tokens_requiring_source("Cohort 7 had 12345678 events") == ["7"]
    # Already source-backed paragraph short-circuits to empty.
    assert (
        _numeric_tokens_requiring_source("We saw 23% growth [Source: AdoptIQ]") == []
    )
    # Metadata markers short-circuit.
    assert _numeric_tokens_requiring_source("Build 31 of 100 deployed") == []
