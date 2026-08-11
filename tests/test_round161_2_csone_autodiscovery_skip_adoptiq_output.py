"""Round 161.2 — CSOne autodiscovery must skip AdoptIQ report XLSX artifacts."""
from __future__ import annotations

import os
import sys
import time
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import app_simple
from adoptiq_backend import load_csone_excel


def _make_xlsx(path: Path, *, size: int) -> None:
    path.write_bytes(b"\x00" * size)


def _make_readable_tac_workbook(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TAC"
    ws.append(["SR Number", "Title", "Severity", "Customer Name"])
    ws.append(["TAC001", "Case one", "P2", "ACME"])
    wb.save(str(path))


@pytest.fixture
def reset_csone_folder():
    saved = app_simple.app.config.get("CSONE_ONEDRIVE_FOLDER")
    yield
    app_simple.app.config["CSONE_ONEDRIVE_FOLDER"] = saved


def test_r161_2_helper_flags_adoptiq_leader_report_filename() -> None:
    name = "AdoptIQ_Report_Leader_Brian_Frazier_All_Contact_Center_90d_20260811.xlsx"
    assert app_simple._r161_2_is_adoptiq_output_csone_filename(name) is True


def test_r161_2_helper_flags_adoptiq_enhanced_collab_summary() -> None:
    name = "AdoptIQ Enhanced Premium Collab Summary-2026-07-16.xlsx"
    assert app_simple._r161_2_is_adoptiq_output_csone_filename(name) is True


def test_r161_2_helper_allows_real_csone_export_filename() -> None:
    assert app_simple._r161_2_is_adoptiq_output_csone_filename("CSOne_TAC_Export_2026-07-15.xlsx") is False


def test_diag_skips_adoptiq_output_and_picks_real_csone(reset_csone_folder, tmp_path) -> None:
    adoptiq = tmp_path / "AdoptIQ_Data_Brian_Frazier_90d.xlsx"
    _make_xlsx(adoptiq, size=8192)
    os.utime(adoptiq, (time.time(), time.time()))

    real = tmp_path / "CSOne_Export_July2026.xlsx"
    _make_readable_tac_workbook(real)
    os.utime(real, (time.time() - 3600, time.time() - 3600))

    app_simple.app.config["CSONE_ONEDRIVE_FOLDER"] = str(tmp_path)
    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path == str(real)
    assert status == "synced"
    assert count == 1


def test_diag_not_synced_when_only_adoptiq_outputs_present(reset_csone_folder, tmp_path) -> None:
    _make_xlsx(tmp_path / "AdoptIQ_Report_Renewal_scope.xlsx", size=4096)
    app_simple.app.config["CSONE_ONEDRIVE_FOLDER"] = str(tmp_path)
    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path is None
    assert status == "not_synced"
    assert count == 0


def test_load_csone_excel_uses_non_active_sheet_with_tac_rows(tmp_path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Report_Info"
    info.append(["Item", "Value"])
    info.append(["Export type", "AdoptIQ"])

    tac = wb.create_sheet("TAC_Cases")
    tac.append(["SR Number", "Title", "Severity", "Customer Name"])
    tac.append(["TAC001", "Case one", "P2", "ACME"])
    tac.append(["TAC002", "Case two", "P3", "BETA"])

    path = tmp_path / "AdoptIQ_Data_test.xlsx"
    wb.save(str(path))

    df = load_csone_excel(path)
    assert len(df) == 2
    assert "SR Number" in df.columns
