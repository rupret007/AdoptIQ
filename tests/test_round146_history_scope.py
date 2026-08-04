"""Round 146 durable history, scope, and safe workspace route regressions."""

from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook


def _use_legacy_history_db(monkeypatch, tmp_path: Path):
    import enhanced_admin_dashboard_v2 as admin

    database = tmp_path / "history.db"
    with sqlite3.connect(database) as connection:
        connection.execute(
            """
            CREATE TABLE report_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                request_id TEXT, report_type TEXT, manager TEXT,
                technology TEXT, customer_name TEXT, status TEXT,
                start_time TEXT, end_time TEXT, ip_address TEXT,
                user_agent TEXT, error_message TEXT, days INTEGER,
                word_path TEXT, excel_path TEXT, word_hash TEXT,
                excel_hash TEXT, partial_data_warnings_json TEXT,
                created_at TEXT
            )
            """
        )
    monkeypatch.setattr(admin, "DB_PATH", str(database))
    monkeypatch.setattr(admin, "_db_initialized_for_path", None)
    monkeypatch.setenv(
        "ADOPTIQ_AUDIT_MIRROR_PATH", str(tmp_path / "history.audit.jsonl")
    )
    admin.init_database()
    return admin, database


def _workspace_book(
    path: Path,
    *,
    canonical: bool,
    open_plans: int,
    scope_type: str = "team",
    scope_value: str = "Entire team",
) -> Path:
    workbook = Workbook()
    info = workbook.active
    info.title = "Report_Info"
    info.append(["Item", "Value"])
    info_rows = [
        ("Report_Type", "leader"),
        ("Manager", "Manager One"),
        ("Technology", "All"),
        ("Scope_Type", scope_type),
        ("Scope_Value", scope_value),
        ("Days", 90),
        ("Data_As_Of_UTC", "2026-08-03T12:00:00Z"),
        ("Source_State:Action_Plans", "available"),
    ]
    if canonical:
        info_rows.append(("Fact_Contract_SHA256", path.stem))
    for row in info_rows:
        info.append(row)

    lineage = workbook.create_sheet("Metric_Lineage")
    lineage.append(
        [
            "Metric_Key",
            "Display_Label",
            "Metric_Value",
            "Unit",
            "Source_State",
            "Source_Sheet",
        ]
    )
    lineage.append(
        [
            "kpi.action_plans_open",
            "Open Action Plans",
            open_plans,
            "records",
            "available",
            "Action_Plans",
        ]
    )
    if canonical:
        chart = workbook.create_sheet("Chart_Data")
        chart.append(
            [
                "Chart_ID",
                "Metric_Key",
                "Display_Label",
                "Series",
                "Category",
                "Period_Start",
                "Value",
                "Unit",
                "Source_State",
            ]
        )
        chart.append(
            [
                "action_plan_status_aging",
                "chart.action_plan_status.open",
                "Action Plan status and aging — Open",
                "Status",
                "Open",
                "",
                open_plans,
                "records",
                "available",
            ]
        )
    actions = workbook.create_sheet("Action_Plans")
    actions.append(
        [
            "ID",
            "BU_NAME",
            "SUBJECT_C",
            "STATUS_C",
            "OWNER_NAME_C",
            "DUE_DATE_C",
            "PRIORITY_C",
            "NEXT_ACTION_C",
        ]
    )
    actions.append(
        [
            "AP-1",
            "Acme",
            "Adoption plan",
            "On Track",
            "Owner",
            "2026-08-20",
            "High",
            "Confirm enablement",
        ]
    )
    accounts = workbook.create_sheet("Account_Summary")
    accounts.append(["Account", "Risk_Band", "Risk_Score_0_100"])
    accounts.append(["Acme", "HIGH", 75])
    workbook.save(path)
    workbook.close()
    return path


def test_legacy_history_schema_migrates_and_round_trips_canonical_scope(
    monkeypatch, tmp_path
):
    admin, database = _use_legacy_history_db(monkeypatch, tmp_path)

    with sqlite3.connect(database) as connection:
        columns = {
            row[1] for row in connection.execute("PRAGMA table_info(report_history)")
        }
    assert {
        "scope_type",
        "scope_value",
        "scope_member",
        "data_as_of_utc",
        "fact_fingerprint",
    }.issubset(columns)

    admin.record_report_completion(
        "leader-member-history",
        "leader",
        "Manager One",
        "All",
        "",
        "completed",
        "2026-08-03T11:00:00Z",
        "2026-08-03T12:00:00Z",
        days=90,
        scope_type="member",
        scope_value="MEMBER@EXAMPLE.COM",
        scope_member="MEMBER@EXAMPLE.COM",
        data_as_of_utc="2026-08-03T10:30:00Z",
        fact_fingerprint="sha256:member-facts",
    )
    record = next(
        row
        for row in admin.get_report_history(limit=1_000)
        if row.get("request_id") == "leader-member-history"
    )
    assert record["scope_type"] == "member"
    assert record["scope_value"] == "member@example.com"
    assert record["scope_member"] == "member@example.com"
    assert record["data_as_of_utc"] == "2026-08-03T10:30:00Z"
    assert record["fact_fingerprint"] == "sha256:member-facts"


def test_history_filters_and_report_scope_survive_real_rehydration(
    client, monkeypatch, tmp_path
):
    import app_simple

    admin, database = _use_legacy_history_db(monkeypatch, tmp_path)
    common = {
        "report_type": "leader",
        "manager": "Manager One",
        "technology": "All",
        "customer_name": "",
        "status": "completed",
        "start_time": "2026-08-03T11:00:00Z",
        "end_time": "2026-08-03T12:00:00Z",
        "days": 90,
        "data_as_of_utc": "2026-08-03T10:30:00Z",
    }
    admin.record_report_completion(
        "leader-member-rehydrated",
        common["report_type"],
        common["manager"],
        common["technology"],
        common["customer_name"],
        common["status"],
        common["start_time"],
        common["end_time"],
        days=common["days"],
        scope_type="member",
        scope_value="member@example.com",
        scope_member="member@example.com",
        data_as_of_utc=common["data_as_of_utc"],
        fact_fingerprint="member-fingerprint",
    )
    admin.record_report_completion(
        "leader-customer-rehydrated",
        common["report_type"],
        common["manager"],
        common["technology"],
        common["customer_name"],
        common["status"],
        common["start_time"],
        common["end_time"],
        days=common["days"],
        scope_type="customer",
        scope_value="Acme Corporation",
        scope_member="member@example.com",
        data_as_of_utc=common["data_as_of_utc"],
        fact_fingerprint="customer-fingerprint",
    )
    # Simulate a full report matrix completed after the two Leader deep dives.
    # The admin page still defaults to 50 rows, while decision history asks for
    # a larger bounded window so these durable scoped records remain findable.
    with sqlite3.connect(database) as connection:
        connection.executemany(
            """
            INSERT INTO report_history
                (request_id, report_type, manager, technology, status,
                 created_at, scope_type)
            VALUES (?, 'compact', 'Other Manager', 'All', 'completed', ?, 'team')
            """,
            [
                (f"later-matrix-{index:02d}", f"2099-01-01T00:{index:02d}:00Z")
                for index in range(60)
            ],
        )

    monkeypatch.setattr(app_simple, "get_report_history", admin.get_report_history)
    monkeypatch.setattr(
        app_simple, "STATUS_FILE", str(tmp_path / "missing-status.json")
    )
    with app_simple.analysis_status_lock:
        app_simple.analysis_status.pop("leader-member-rehydrated", None)
        app_simple.analysis_status.pop("leader-customer-rehydrated", None)

    member_response = client.get(
        "/api/decision-workspace/history", query_string={"scope_type": "member"}
    )
    assert member_response.status_code == 200
    assert [
        item["analysis_id"] for item in member_response.get_json()["reports"]
    ] == ["leader-member-rehydrated"]

    customer_response = client.get(
        "/api/decision-workspace/history",
        query_string={"scope_type": "customer", "customer": "Acme"},
    )
    assert customer_response.status_code == 200
    assert [
        item["analysis_id"] for item in customer_response.get_json()["reports"]
    ] == ["leader-customer-rehydrated"]

    member_report = client.get(
        "/api/decision-workspace/report/leader-member-rehydrated"
    ).get_json()["report"]
    assert member_report["scope_type"] == "member"
    assert member_report["scope_value"] == "member@example.com"
    assert member_report["scope_member"] == "member@example.com"
    assert member_report["fact_fingerprint"] == "member-fingerprint"

    customer_report_response = client.get(
        "/api/decision-workspace/report/leader-customer-rehydrated"
    )
    customer_report = customer_report_response.get_json()["report"]
    assert customer_report["scope_type"] == "customer"
    assert customer_report["scope_value"] == "Acme Corporation"
    assert customer_report["scope_member"] == "member@example.com"
    assert str(tmp_path) not in customer_report_response.get_data(as_text=True)


def test_history_does_not_advertise_stale_artifacts(client, monkeypatch, tmp_path):
    import app_simple

    admin, _database = _use_legacy_history_db(monkeypatch, tmp_path)
    stale_root = tmp_path / "deleted-private-customer-folder"
    admin.record_report_completion(
        "stale-history-artifacts",
        "leader",
        "Manager One",
        "All",
        "",
        "completed",
        "2026-08-03T11:00:00Z",
        "2026-08-03T12:00:00Z",
        days=90,
        word_path=str(stale_root / "report.docx"),
        excel_path=str(stale_root / "source.xlsx"),
        scope_type="team",
    )
    monkeypatch.setattr(app_simple, "get_report_history", admin.get_report_history)
    monkeypatch.setattr(
        app_simple, "STATUS_FILE", str(tmp_path / "missing-status.json")
    )
    with app_simple.analysis_status_lock:
        app_simple.analysis_status.pop("stale-history-artifacts", None)

    response = client.get(
        "/api/decision-workspace/history",
        query_string={"manager": "Manager One"},
    )
    assert response.status_code == 200
    record = response.get_json()["reports"][0]
    assert record["word_available"] is False
    assert record["excel_available"] is False
    assert record["word_url"] == ""
    assert record["source_data_url"] == ""
    assert str(stale_root) not in response.get_data(as_text=True)


def test_report_projection_uses_public_warnings_and_disables_unbound_ask_ai(
    client, monkeypatch, tmp_path
):
    import app_simple

    unreadable = tmp_path / "unreadable.xlsx"
    unreadable.write_text("not an xlsx archive", encoding="utf-8")
    raw_warning = {
        "dataset": "csone",
        "kind": "fetch_error",
        "error": "internal Snowflake connector detail",
    }
    status = {
        "analysis_id": "warning-projection",
        "status": "completed",
        "report_type": "leader",
        "manager": "Manager One",
        "scope_type": "team",
        "excel_report": str(unreadable),
        "partial_data_warnings": [raw_warning],
    }
    monkeypatch.setattr(
        app_simple, "_r146_status_for_workspace", lambda _analysis_id: status
    )
    monkeypatch.setattr(
        app_simple,
        "_r92_resolve_output_artifact",
        lambda raw, **_kwargs: str(raw) if raw else None,
    )

    warning_response = client.get(
        "/api/decision-workspace/report/warning-projection"
    )
    assert warning_response.status_code == 200
    warning_report = warning_response.get_json()["report"]
    assert "partial_data_warnings" not in warning_report
    assert any(
        "could not be read" in warning.lower()
        for warning in warning_report["source_warnings"]
    )
    warning_json = warning_response.get_data(as_text=True).lower()
    assert "internal snowflake connector detail" not in warning_json
    assert '"kind"' not in warning_json

    old_member_book = _workspace_book(
        tmp_path / "old-member.xlsx",
        canonical=True,
        open_plans=1,
        scope_type="member",
        scope_value="Member Name Without Canonical Email",
    )
    old_status = {
        "analysis_id": "old-member-unbound",
        "status": "completed",
        "report_type": "leader",
        "manager": "Manager One",
        "excel_report": str(old_member_book),
    }
    monkeypatch.setattr(
        app_simple, "_r146_status_for_workspace", lambda _analysis_id: old_status
    )
    old_response = client.get(
        "/api/decision-workspace/report/old-member-unbound"
    )
    old_report = old_response.get_json()["report"]
    assert old_report["ask_ai_binding_available"] is False
    assert old_report["ask_ai_url"] == ""


def test_compare_route_allows_canonical_and_rejects_legacy_snapshot(
    client, monkeypatch, tmp_path
):
    import app_simple

    before = _workspace_book(
        tmp_path / "canonical-before.xlsx", canonical=True, open_plans=3
    )
    after = _workspace_book(
        tmp_path / "canonical-after.xlsx", canonical=True, open_plans=1
    )
    legacy = _workspace_book(
        tmp_path / "legacy-after.xlsx", canonical=False, open_plans=1
    )
    statuses = {
        "canonical-before": {
            "analysis_id": "canonical-before",
            "status": "completed",
            "report_type": "leader",
            "manager": "Manager One",
            "technology": "All",
            "days": 90,
            "scope_type": "team",
            "excel_report": str(before),
        },
        "canonical-after": {
            "analysis_id": "canonical-after",
            "status": "completed",
            "report_type": "leader",
            "manager": "Manager One",
            "technology": "All",
            "days": 90,
            "scope_type": "team",
            "excel_report": str(after),
        },
        "legacy-after": {
            "analysis_id": "legacy-after",
            "status": "completed",
            "report_type": "leader",
            "manager": "Manager One",
            "technology": "All",
            "days": 90,
            "scope_type": "team",
            "excel_report": str(legacy),
        },
    }
    monkeypatch.setattr(
        app_simple, "_r146_status_for_workspace", lambda analysis_id: statuses[analysis_id]
    )
    monkeypatch.setattr(
        app_simple,
        "_r92_resolve_output_artifact",
        lambda raw, **_kwargs: str(raw) if raw and Path(str(raw)).is_file() else None,
    )

    allowed = client.post(
        "/api/decision-workspace/compare",
        json={
            "before_analysis_id": "canonical-before",
            "after_analysis_id": "canonical-after",
        },
    )
    assert allowed.status_code == 200
    assert allowed.get_json()["comparison"]["metric_changes"][0]["delta"] == -2

    rejected = client.post(
        "/api/decision-workspace/compare",
        json={
            "before_analysis_id": "canonical-before",
            "after_analysis_id": "legacy-after",
        },
    )
    assert rejected.status_code == 409
    assert "canonical source data contract" in rejected.get_json()["error"].lower()
