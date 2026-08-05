"""Round 147 HTTP contracts for exact report evidence drill-down."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

import decision_report_delivery as delivery
import manager_decision_workspace as workspace
from tests.test_round142_decision_report_delivery import _facts


@pytest.fixture
def evidence_route_state(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    import app_simple

    source_path = tmp_path / "private-customer-source-data.xlsx"
    delivery.write_source_data_workbook(
        source_path,
        delivery.build_source_data_sheets(_facts()),
    )
    snapshot = workspace.load_workbook_snapshot(source_path)
    snapshot.update({"analysis_id": "round147-evidence-route", "workbook_loaded": True})
    artifact_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    status = {
        "analysis_id": "round147-evidence-route",
        "excel_report": str(source_path),
        "excel_hash": artifact_hash,
    }
    monkeypatch.setattr(
        app_simple,
        "_r146_workspace_snapshot",
        lambda _analysis_id: (snapshot, 200),
    )
    monkeypatch.setattr(app_simple, "_r146_status_for_workspace", lambda _analysis_id: status)
    monkeypatch.setattr(
        app_simple,
        "_r146_workspace_artifact",
        lambda _status, kind: str(source_path) if kind == "excel" else None,
    )
    return app_simple, source_path, snapshot


def test_evidence_route_returns_exact_rows_without_paths(client, evidence_route_state) -> None:
    _app, source_path, _snapshot = evidence_route_state
    response = client.get(
        "/api/decision-workspace/report/round147-evidence-route/evidence",
        query_string={"evidence_key": "kpi.action_plans_overdue", "limit": 1},
    )

    assert response.status_code == 200
    assert response.headers["Cache-Control"] == "no-store"
    payload = response.get_json()
    assert payload["evidence"]["total_records"] == 2
    assert len(payload["evidence"]["records"]) == 1
    assert payload["evidence"]["records"][0]["source_sheet"] == "Action_Plans"
    assert payload["evidence"]["source_data_url"].endswith("/xlsx")
    assert str(source_path) not in response.get_data(as_text=True)
    assert str(source_path.parent) not in response.get_data(as_text=True)


@pytest.mark.parametrize(
    "persisted_hash",
    ("", "not-a-sha256", "0" * 64),
    ids=("missing", "malformed", "mismatch"),
)
def test_evidence_route_requires_matching_persisted_artifact_hash(
    client,
    monkeypatch: pytest.MonkeyPatch,
    evidence_route_state,
    persisted_hash: str,
) -> None:
    app_simple, _source_path, _snapshot = evidence_route_state
    monkeypatch.setattr(
        app_simple,
        "_r146_persisted_excel_hash",
        lambda *_args: persisted_hash,
    )

    response = client.get(
        "/api/decision-workspace/report/round147-evidence-route/evidence",
        query_string={"evidence_key": "kpi.action_plans_overdue"},
    )

    assert response.status_code == 409
    assert "persisted artifact integrity proof" in response.get_json()["error"].lower()


def test_evidence_route_distinguishes_invalid_and_missing_keys(
    client, evidence_route_state
) -> None:
    invalid = client.get(
        "/api/decision-workspace/report/round147-evidence-route/evidence",
        query_string={"evidence_key": "../private.xlsx"},
    )
    missing = client.get(
        "/api/decision-workspace/report/round147-evidence-route/evidence",
        query_string={"evidence_key": "kpi.not_present"},
    )

    assert invalid.status_code == 400
    assert missing.status_code == 404


def test_evidence_route_fails_closed_on_fingerprint_mismatch(
    client, monkeypatch: pytest.MonkeyPatch, evidence_route_state
) -> None:
    app_simple, _source_path, snapshot = evidence_route_state
    mismatched = dict(snapshot, fact_fingerprint="not-the-workbook-fingerprint")
    monkeypatch.setattr(
        app_simple,
        "_r146_workspace_snapshot",
        lambda _analysis_id: (mismatched, 200),
    )

    response = client.get(
        "/api/decision-workspace/report/round147-evidence-route/evidence",
        query_string={"evidence_key": "kpi.action_plans_overdue"},
    )
    assert response.status_code == 409
    assert "integrity" in response.get_json()["error"].lower()


def test_evidence_route_fails_closed_after_row_tamper(
    client, evidence_route_state
) -> None:
    _app, source_path, _snapshot = evidence_route_state
    workbook = load_workbook(source_path)
    sheet = workbook["Action_Plans"]
    title_column = next(cell.column for cell in sheet[1] if cell.value == "AdoptIQ_Title")
    sheet.cell(2, title_column).value = "Tampered after report generation"
    workbook.save(source_path)
    workbook.close()

    response = client.get(
        "/api/decision-workspace/report/round147-evidence-route/evidence",
        query_string={"evidence_key": "kpi.action_plans_overdue"},
    )
    assert response.status_code == 409
    assert "integrity" in response.get_json()["error"].lower()


def test_evidence_route_rejects_legacy_snapshot(
    client, monkeypatch: pytest.MonkeyPatch, evidence_route_state
) -> None:
    app_simple, _source_path, snapshot = evidence_route_state
    legacy = dict(
        snapshot,
        evidence_available=False,
        evidence_notice="Generate a current canonical report.",
    )
    monkeypatch.setattr(
        app_simple,
        "_r146_workspace_snapshot",
        lambda _analysis_id: (legacy, 200),
    )

    response = client.get(
        "/api/decision-workspace/report/round147-evidence-route/evidence",
        query_string={"evidence_key": "kpi.action_plans_overdue"},
    )
    assert response.status_code == 409
    assert "current canonical report" in response.get_json()["error"].lower()


def test_evidence_route_enforces_server_record_cap(client, evidence_route_state) -> None:
    response = client.get(
        "/api/decision-workspace/report/round147-evidence-route/evidence",
        query_string={
            "evidence_key": "chart.activity_trend.action_plans.2026-07-06",
            "limit": 999_999,
        },
    )
    assert response.status_code == 200
    evidence = response.get_json()["evidence"]
    assert len(evidence["records"]) <= workspace.MAX_EVIDENCE_RECORDS
