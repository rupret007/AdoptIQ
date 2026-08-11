"""Round 68 / Build 42 (B1): pin behavior of get_latest_csone_from_folder
and the new get_latest_csone_from_folder_diag sibling helper that skips
0-byte placeholder files (OneDrive Files-On-Demand stubs) and reports a
``not_synced`` sync_status when the autodiscovery folder is
present-but-empty.

Why this suite exists: pre-Round-68 the autodiscovery helper happily
returned a 0-byte placeholder as "the latest CSOne report" because
``listdir`` shows Files-On-Demand stubs at zero bytes on disk until the
OneDrive client pulls them.  The leader-report worker then loaded the
zero-byte xlsx, got an empty DataFrame, and the validator tripped the
required-but-empty branch -- producing the Round 38 "csone missing or
empty" abort.  This suite pins the corpus-gate-style filter so the
regression cannot return.

Mirrors the corpus gate semantics in
``corpus_bootstrap._check_onedrive_sync_status`` so the same
``synced`` / ``not_synced`` / ``unknown`` taxonomy applies to BOTH
the corpus refresh path and the CSOne autodiscovery path.
"""
from __future__ import annotations

import os
import sys
import time
from pathlib import Path

import pytest

# Ensure project root is importable so tests can import the Flask app.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import app_simple


def _make_xlsx(path: Path, *, size: int) -> None:
    """Create a fake xlsx at ``path`` of exactly ``size`` bytes."""
    path.write_bytes(b"\x00" * size)


def _make_readable_xlsx(path: Path) -> None:
    """Round 162: minimal valid workbook for autodiscovery readability probe."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TAC"
    ws.append(["SR Number", "Title", "Severity", "Customer Name"])
    ws.append(["TAC001", "Case one", "P2", "ACME"])
    wb.save(str(path))


@pytest.fixture
def reset_csone_folder():
    """Save / restore CSONE_ONEDRIVE_FOLDER on the live Flask app config."""
    saved = app_simple.app.config.get('CSONE_ONEDRIVE_FOLDER')
    yield
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = saved


def test_diag_returns_unknown_when_env_unset(reset_csone_folder, tmp_path):
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = None
    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path is None
    assert status == "unknown"
    assert count == 0


def test_diag_returns_not_synced_when_folder_missing(reset_csone_folder, tmp_path):
    missing = tmp_path / "does-not-exist"
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(missing)
    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path is None
    assert status == "not_synced"
    assert count == 0


def test_diag_returns_not_synced_when_folder_empty(reset_csone_folder, tmp_path):
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(tmp_path)
    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path is None
    assert status == "not_synced"
    assert count == 0


def test_diag_skips_zero_byte_placeholder_files(reset_csone_folder, tmp_path):
    """The classic Round-38 trigger: OneDrive Files-On-Demand placeholders."""
    _make_xlsx(tmp_path / "report1.xlsx", size=0)
    _make_xlsx(tmp_path / "report2.xlsx", size=0)
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(tmp_path)

    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path is None, "0-byte placeholders must NOT be returned as candidates"
    assert status == "not_synced"
    assert count == 0


def test_diag_returns_synced_with_real_file(reset_csone_folder, tmp_path):
    real = tmp_path / "report-real.xlsx"
    _make_readable_xlsx(real)
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(tmp_path)

    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path == str(real)
    assert status == "synced"
    assert count == 1


def test_diag_picks_newest_real_file_only(reset_csone_folder, tmp_path):
    placeholder = tmp_path / "old-placeholder.xlsx"
    _make_xlsx(placeholder, size=0)

    older_real = tmp_path / "older-real.xlsx"
    _make_readable_xlsx(older_real)
    os.utime(older_real, (time.time() - 3600, time.time() - 3600))

    newer_real = tmp_path / "newer-real.xlsx"
    _make_readable_xlsx(newer_real)

    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(tmp_path)
    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path == str(newer_real), "must skip placeholder, must pick newest of real files"
    assert status == "synced"
    assert count == 2


def test_legacy_get_latest_csone_from_folder_returns_just_path(reset_csone_folder, tmp_path):
    """Back-compat: existing callers that ignore sync_status must still work."""
    real = tmp_path / "report.xlsx"
    _make_readable_xlsx(real)
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(tmp_path)
    assert app_simple.get_latest_csone_from_folder() == str(real)


def test_legacy_helper_returns_none_for_only_placeholders(reset_csone_folder, tmp_path):
    _make_xlsx(tmp_path / "stub1.xlsx", size=0)
    _make_xlsx(tmp_path / "stub2.xlsx", size=0)
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(tmp_path)
    assert app_simple.get_latest_csone_from_folder() is None


def test_diag_ignores_lock_files(reset_csone_folder, tmp_path):
    """``~$report.xlsx`` Office lock files MUST NOT be candidates even when sized."""
    _make_xlsx(tmp_path / "~$report.xlsx", size=8192)
    real = tmp_path / "report.xlsx"
    _make_readable_xlsx(real)
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(tmp_path)

    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path == str(real)
    assert status == "synced"
    assert count == 1


def test_diag_ignores_non_xlsx_extensions(reset_csone_folder, tmp_path):
    _make_xlsx(tmp_path / "notes.txt", size=8192)
    _make_xlsx(tmp_path / "scratch.csv", size=4096)
    real = tmp_path / "report.xlsx"
    _make_readable_xlsx(real)
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(tmp_path)

    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path == str(real)
    assert status == "synced"
    assert count == 1


def test_diag_accepts_xls_extension(reset_csone_folder, tmp_path):
    real = tmp_path / "legacy-report.xls"
    _make_xlsx(real, size=2048)
    app_simple.app.config['CSONE_ONEDRIVE_FOLDER'] = str(tmp_path)

    path, status, count = app_simple.get_latest_csone_from_folder_diag()
    assert path == str(real)
    assert status == "synced"
    assert count == 1


def test_source_shape_get_latest_csone_from_folder_uses_diag_helper():
    """The legacy wrapper MUST delegate to the diag variant -- if a future
    refactor reverts the wrapper to walk listdir directly, this guard
    fires before the regression ships.
    """
    import inspect
    src = inspect.getsource(app_simple.get_latest_csone_from_folder)
    assert "get_latest_csone_from_folder_diag" in src, (
        "Round 68 / Build 42 (B1): get_latest_csone_from_folder MUST "
        "delegate to get_latest_csone_from_folder_diag so the 0-byte "
        "skip is honored on every call site."
    )
