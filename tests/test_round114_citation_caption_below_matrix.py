"""Round 114 / Build 83 -- pin the matrix citation de-clutter contract.

Build 82 acceptance surfaced ``[Source: AdoptIQ Report Data Sources]``
appended to EVERY numeric cell of EVERY data row in the Leader report's
multi-column count matrices (Team Activity Summary, Individual Team
Member Performance) -- one citation per member per KPI column, which
cluttered the rendered tables and quoted AdoptIQ as the source inline.

Round 114 replaces that with ONE compact italic ``Sources: ...`` caption
directly below each matrix, aggregating the matrix's numeric KPI columns
through the R82 source taxonomy.  These tests pin:

1. A multi-column matrix receives exactly ONE trailing caption paragraph
   and ZERO ``[Source`` markers in its data-row cells.
2. The caption aggregates the columns by their R82 source system.
3. The injector is idempotent on the matrix path (second pass adds no
   second caption and leaves the document byte-stable).
4. Two-column ``label | value`` tables keep their per-row citation
   (they were never the clutter complaint; preserves the R82 contract).
5. The quality scorer (``evaluate_report_quality``) reports ZERO unbacked
   matrix claims when the caption is present -- the caption backs the
   matrix's aggregated multi-column claims (contract evolution, not a
   weakened check).
6. Columns that do not resolve through the R82 taxonomy degrade to the
   generic ``AdoptIQ Report Data Sources`` string so a matrix never
   loses its citation entirely.

The .docx fixtures are built in tmp_path so the suite costs nothing in
CI and does not depend on a live build.

Round 114 / Build 83.  Made-with: Cursor.
"""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402

from report_iteration_loop import (  # noqa: E402
    _matrix_has_following_source_caption,
    evaluate_report_quality,
)
from report_source_injector import (  # noqa: E402
    _MATRIX_SOURCE_CAPTION_PREFIX,
    _build_matrix_source_caption,
    inject_source_citations_into_docx,
)


def _build_matrix_docx(path: Path) -> None:
    """A doc whose only KPI table is a multi-column count matrix."""
    doc = Document()
    doc.add_heading("Leader Report", level=1)
    doc.add_paragraph("Team Activity Summary")
    table = doc.add_table(rows=4, cols=5)
    header = ["Team Member", "Action Plans", "Adoption Barriers", "Customer Pulse", "TAC Cases"]
    for col_idx, label in enumerate(header):
        table.rows[0].cells[col_idx].text = label
    members = [
        ("Brandon Doan", "12", "3", "1", "4"),
        ("Greg Dolberry", "7", "5", "2", "0"),
        ("TOTAL", "19", "8", "3", "4"),
    ]
    for row_idx, member in enumerate(members, start=1):
        for col_idx, value in enumerate(member):
            table.rows[row_idx].cells[col_idx].text = value
    doc.save(str(path))


def _build_two_column_docx(path: Path) -> None:
    """A doc whose only KPI table is a two-column ``label | value`` table."""
    doc = Document()
    doc.add_heading("Summary", level=1)
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Total Customers"
    table.rows[1].cells[1].text = "52"
    table.rows[2].cells[0].text = "Adoption Barriers"
    table.rows[2].cells[1].text = "68"
    doc.save(str(path))


def _caption_paragraphs(path: Path) -> list[str]:
    doc = Document(str(path))
    return [
        p.text.strip()
        for p in doc.paragraphs
        if (p.text or "").strip().startswith(_MATRIX_SOURCE_CAPTION_PREFIX)
    ]


def _matrix_data_cell_texts(path: Path) -> list[str]:
    """All data-row (non-header) cell texts across every table in the doc."""
    doc = Document(str(path))
    out: list[str] = []
    for table in doc.tables:
        for row in list(table.rows)[1:]:
            for cell in row.cells:
                out.append(cell.text)
    return out


def test_matrix_gets_one_caption_and_zero_per_cell_citations(tmp_path: Path) -> None:
    docx = tmp_path / "matrix.docx"
    _build_matrix_docx(docx)

    counts = inject_source_citations_into_docx(docx)

    assert counts["table_captions_added"] == 1, (
        "a multi-column count matrix must receive exactly ONE trailing "
        f"caption; got {counts['table_captions_added']}"
    )
    captions = _caption_paragraphs(docx)
    assert len(captions) == 1, f"expected exactly one caption paragraph, got {captions}"

    # No data-row cell may carry the per-cell ``[Source`` chrome anymore.
    for cell_text in _matrix_data_cell_texts(docx):
        assert "[source:" not in cell_text.lower(), (
            f"matrix data cell must NOT carry an inline citation; got {cell_text!r}"
        )


def test_caption_aggregates_r82_source_systems(tmp_path: Path) -> None:
    docx = tmp_path / "matrix.docx"
    _build_matrix_docx(docx)
    inject_source_citations_into_docx(docx)

    caption = _caption_paragraphs(docx)[0]
    # Action Plans / Adoption Barriers / Customer Pulse -> Snowflake CSConsole.
    assert "Snowflake CSConsole" in caption
    # TAC Cases -> Snowflake CSOne.
    assert "Snowflake CSOne" in caption
    assert caption.startswith(_MATRIX_SOURCE_CAPTION_PREFIX)


def test_caption_idempotent_on_second_pass(tmp_path: Path) -> None:
    docx = tmp_path / "matrix.docx"
    _build_matrix_docx(docx)

    first = inject_source_citations_into_docx(docx)
    first_text = "\n".join(p.text for p in Document(str(docx)).paragraphs)
    first_mtime = docx.stat().st_mtime

    second = inject_source_citations_into_docx(docx)
    second_text = "\n".join(p.text for p in Document(str(docx)).paragraphs)
    second_mtime = docx.stat().st_mtime

    assert first["table_captions_added"] == 1
    assert second["table_captions_added"] == 0, "second pass must not add a duplicate caption"
    assert len(_caption_paragraphs(docx)) == 1, "exactly one caption must survive a re-run"
    assert first_text == second_text
    assert first_mtime == second_mtime, "idempotent re-run must not re-save the document"


def test_two_column_table_keeps_per_row_citation(tmp_path: Path) -> None:
    docx = tmp_path / "twocol.docx"
    _build_two_column_docx(docx)

    counts = inject_source_citations_into_docx(docx)

    assert counts["table_cells_injected"] >= 2, (
        "two-column label|value rows must keep their per-row cell citation"
    )
    assert counts["table_captions_added"] == 0, (
        "a two-column table is not a matrix and must NOT receive a caption"
    )
    assert _caption_paragraphs(docx) == []
    # The value cells DO carry inline citations (unchanged behaviour).
    joined = "\n".join(_matrix_data_cell_texts(docx))
    assert "[source:" in joined.lower()


def test_quality_gate_zero_unbacked_matrix_claims_with_caption(tmp_path: Path) -> None:
    docx = tmp_path / "matrix.docx"
    _build_matrix_docx(docx)

    pre_payload, _pre_gate = evaluate_report_quality(
        docx_path=docx, xlsx_path=None, scenario_key="leader", strict=True
    )
    pre_unbacked = pre_payload.get("unbacked_metric_claim_count", 0)
    assert pre_unbacked > 0, (
        "fixture must start with unbacked matrix claims so the test exercises the caption path"
    )

    inject_source_citations_into_docx(docx, scenario_key="leader")

    post_payload, _post_gate = evaluate_report_quality(
        docx_path=docx, xlsx_path=None, scenario_key="leader", strict=True
    )
    post_unbacked = post_payload.get("unbacked_metric_claim_count", 0)
    assert post_unbacked == 0, (
        "the trailing caption must source-back the matrix's multi-column claims; "
        f"pre={pre_unbacked} post={post_unbacked}"
    )


def test_caption_generic_fallback_when_columns_unresolved(tmp_path: Path) -> None:
    """Numeric columns that do not map through R82 fall back to the generic source."""
    doc = Document()
    doc.add_heading("Custom Matrix", level=1)
    table = doc.add_table(rows=3, cols=4)
    header = ["Region", "Widgets", "Gizmos", "Doohickeys"]
    for col_idx, label in enumerate(header):
        table.rows[0].cells[col_idx].text = label
    rows = [("North", "11", "22", "33"), ("South", "1", "2", "3")]
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            table.rows[row_idx].cells[col_idx].text = value
    docx = tmp_path / "custom.docx"
    doc.save(str(docx))

    counts = inject_source_citations_into_docx(docx)

    assert counts["table_captions_added"] == 1
    caption = _caption_paragraphs(docx)[0]
    assert "AdoptIQ Report Data Sources" in caption, (
        "unresolved KPI columns must degrade to the generic source string, "
        f"not silently drop the citation; got {caption!r}"
    )


def test_build_matrix_source_caption_groups_by_tag() -> None:
    header = ["Member", "Action Plans", "Adoption Barriers", "TAC Cases", "Support Cases"]
    # Numeric columns 1..4 (column 0 is the name/label column).
    caption = _build_matrix_source_caption(header, [1, 2, 3, 4])
    assert caption.startswith("Sources:")
    # CSConsole group lists Action Plans + Adoption Barriers together.
    assert "Action Plans, Adoption Barriers - Snowflake CSConsole" in caption
    # CSOne group lists TAC Cases + Support Cases together.
    assert "Snowflake CSOne" in caption


def test_build_matrix_source_caption_empty_columns_returns_generic() -> None:
    caption = _build_matrix_source_caption(["A", "B"], [])
    assert caption == "Sources: AdoptIQ Report Data Sources"


def test_matrix_has_following_source_caption_detects_and_rejects(tmp_path: Path) -> None:
    docx = tmp_path / "matrix.docx"
    _build_matrix_docx(docx)
    # Before injection: no caption following the matrix.
    doc = Document(str(docx))
    assert _matrix_has_following_source_caption(doc.tables[0]) is False

    inject_source_citations_into_docx(docx)

    # After injection: caption is detected.
    doc2 = Document(str(docx))
    assert _matrix_has_following_source_caption(doc2.tables[0]) is True


def test_caption_only_document_is_saved(tmp_path: Path) -> None:
    """A doc whose ONLY change is a caption must still be persisted to disk."""
    docx = tmp_path / "matrix_only.docx"
    _build_matrix_docx(docx)

    counts = inject_source_citations_into_docx(docx)

    assert counts["paragraphs_injected"] == 0, (
        "this fixture's paragraphs carry no KPI claims; only the caption should change"
    )
    assert counts["table_cells_injected"] == 0
    assert counts["table_captions_added"] == 1
    # Reopen from disk to prove the save happened despite zero paragraph/cell injections.
    assert len(_caption_paragraphs(docx)) == 1
