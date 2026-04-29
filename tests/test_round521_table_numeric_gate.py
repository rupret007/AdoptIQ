"""Round 52.1 / Phase 2: ``compare_docx_against_baseline`` table gate.

Pins the contract that ``compare_docx_against_baseline`` now computes
and reports ``table_numeric_similarity`` and ``table_numeric_threshold``
in ``details``, and that the gate ``passed`` flag is ANDed with the new
table-only similarity meeting its threshold in strict mode.

The most important test is the iter2 reproducer: when paragraph
numerics drift but table numerics stay stable, the new gate must PASS
under per-scenario overrides (comprehensive numeric override = 0.55;
table-only stays at 0.95) -- this is the exact failure mode we shipped
this round to fix.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from report_iteration_loop import compare_docx_against_baseline


def _build_docx(
    tmp_path: Path,
    name: str,
    paragraphs: list[str],
    table_rows: list[list[str]],
) -> Path:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        cols = max(len(row) for row in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=cols)
        for r_idx, row in enumerate(table_rows):
            for c_idx, cell_text in enumerate(row):
                table.rows[r_idx].cells[c_idx].text = cell_text
    out = tmp_path / name
    doc.save(out)
    return out


def test_round521_gate_reports_table_numeric_similarity(tmp_path: Path):
    """The gate must surface ``table_numeric_similarity`` and threshold
    in its ``details``, and the field must be a 0..1 float in strict mode."""
    rows = [
        ["Metric", "Value"],
        ["Total Customers", "42"],
        ["Open Action Plans", "381"],
    ]
    a = _build_docx(tmp_path, "a.docx", ["Pattern 1: noise"], rows)
    b = _build_docx(tmp_path, "b.docx", ["Pattern 1: noise"], rows)

    result = compare_docx_against_baseline(
        a,
        b,
        min_similarity=0.40,
        strict=True,
        min_numeric_similarity=0.55,
        min_table_numeric_similarity=0.95,
    )

    details = result.details
    assert "table_numeric_similarity" in details, (
        "gate must report table_numeric_similarity in details"
    )
    assert "table_numeric_threshold" in details
    assert details["table_numeric_threshold"] == 0.95
    assert isinstance(details["table_numeric_similarity"], float)
    assert 0.0 <= details["table_numeric_similarity"] <= 1.0


def test_round521_gate_passes_when_only_paragraphs_drift_and_tables_stable(
    tmp_path: Path,
):
    """The iter2 reproducer.

    Same tables on both sides, but the AI-narrative paragraphs cite
    different bug IDs / percentages / case counts.  Under the per-
    scenario override (numeric=0.55, table_numeric=0.95) the gate must
    PASS -- the table-only signal proves real Snowflake data did not
    drift, even though the overall numeric signal is noisy.
    """
    rows = [
        ["Metric", "Value"],
        ["Total Customers", "42"],
        ["Open Action Plans", "381"],
        ["Risk Score", "0.7"],
        ["Customer Pulse", "118"],
        ["TAC Cases", "237"],
    ]
    iter1 = _build_docx(
        tmp_path,
        "iter1.docx",
        paragraphs=[
            "Pattern 1 (17% of cases): BEMS-12345 reproduced 4 times",
            "Sentiment direction: down; prior occurrences: 9",
            "Evidence: TAC-99812, TAC-99813",
        ],
        table_rows=rows,
    )
    iter2 = _build_docx(
        tmp_path,
        "iter2.docx",
        paragraphs=[
            "Pattern 1 (43% of cases): BEMS-99999 reproduced 7 times",
            "Sentiment direction: up; prior occurrences: 12",
            "Evidence: TAC-77711, TAC-77712, TAC-77713",
        ],
        table_rows=rows,
    )

    result = compare_docx_against_baseline(
        iter1,
        iter2,
        # Comprehensive's effective text floor.
        min_similarity=0.40,
        strict=True,
        # Comprehensive's per-scenario informational numeric override.
        min_numeric_similarity=0.55,
        min_table_numeric_similarity=0.95,
    )

    assert result.details["table_numeric_similarity"] == 1.0, (
        f"tables identical -> table-only sim must be 1.0; got {result.details}"
    )
    assert result.passed, (
        "with table_numeric == 1.0 and overall numeric >= 0.55 override, "
        f"gate must PASS; got {result.details}"
    )


def test_round521_gate_fails_when_table_numerics_drift_below_threshold(
    tmp_path: Path,
):
    """Inverse safety net: if real Snowflake data DOES drift in a way
    that mutates table cells, the new gate must FAIL even if the overall
    numeric similarity happened to coincide.  This guarantees the gate
    cannot be silently disabled by pathological coincidences."""
    rows_baseline = [
        ["Metric", "Value"],
        # Numerics that will all be different on the current side.
        ["A", "10"],
        ["B", "20"],
        ["C", "30"],
        ["D", "40"],
    ]
    rows_current = [
        ["Metric", "Value"],
        ["A", "111"],
        ["B", "222"],
        ["C", "333"],
        ["D", "444"],
    ]

    baseline = _build_docx(tmp_path, "baseline.docx", [], rows_baseline)
    current = _build_docx(tmp_path, "current.docx", [], rows_current)

    result = compare_docx_against_baseline(
        current,
        baseline,
        min_similarity=0.0,  # neutralize text gate
        strict=True,
        min_numeric_similarity=0.0,  # neutralize overall numeric gate
        min_table_numeric_similarity=0.95,
    )

    assert result.details["table_numeric_similarity"] < 0.95, (
        f"divergent table numerics must drop sim below 0.95; got {result.details}"
    )
    assert not result.passed, (
        "table-only gate must be the binding signal here; got "
        f"{result.details}"
    )


def test_round521_gate_table_threshold_disabled_outside_strict_mode(
    tmp_path: Path,
):
    """In non-strict mode, the table-only gate is informational only --
    ``table_numeric_threshold`` reports None and the gate does not block
    on the new signal.  Mirrors the existing overall-numeric semantics."""
    rows_a = [["Metric", "Value"], ["A", "10"]]
    rows_b = [["Metric", "Value"], ["A", "999"]]
    a = _build_docx(tmp_path, "a.docx", [], rows_a)
    b = _build_docx(tmp_path, "b.docx", [], rows_b)

    result = compare_docx_against_baseline(
        a,
        b,
        min_similarity=0.0,
        strict=False,
        min_numeric_similarity=0.0,
        min_table_numeric_similarity=0.95,
    )

    assert "table_numeric_similarity" in result.details
    assert result.details["table_numeric_threshold"] is None, (
        "non-strict must report a None threshold (informational only); "
        f"got {result.details['table_numeric_threshold']!r}"
    )
    assert result.passed, (
        "non-strict mode must not block on table-only drift; got "
        f"{result.details}"
    )
