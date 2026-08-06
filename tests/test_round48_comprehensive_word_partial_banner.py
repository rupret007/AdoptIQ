"""Round 48 / F-COMP-PARTIAL-BANNER-MISSING regression tests.

Pin that ``ExecutiveReportBuilder`` (the comprehensive Word
renderer) surfaces a Partial Data Warning banner identical in
content/style to the renewal / compact / EI banners.  Pre-Round 48
the comprehensive Word silently consumed a "successful" report
built on partial data while ``analysis_status[*][
'partial_data_warnings']`` correctly carried the warnings -- a
human reading just the .docx had no way to know.
"""

from __future__ import annotations
from source_shape_utils import assert_in_source, index_in_source

import inspect
from pathlib import Path

import pytest

from executive_report_builder import ExecutiveReportBuilder


# ---------------------------------------------------------------------------
# Method must exist with the expected signature
# ---------------------------------------------------------------------------


def test_round48_builder_exposes_partial_banner_method():
    assert hasattr(ExecutiveReportBuilder, "add_partial_data_warning_banner"), (
        "ExecutiveReportBuilder.add_partial_data_warning_banner missing"
    )
    sig = inspect.signature(
        ExecutiveReportBuilder.add_partial_data_warning_banner
    )
    # self + partial_data_warnings = 2 params.
    assert len(sig.parameters) == 2, (
        f"Expected method signature (self, partial_data_warnings); "
        f"got {list(sig.parameters)}"
    )


def test_round48_builder_fix_anchor_present():
    src = Path("executive_report_builder.py").read_text(encoding="utf-8")
    assert_in_source(src, "F-COMP-PARTIAL-BANNER-MISSING", label='src')


def test_round48_app_simple_calls_banner_after_title():
    """The comprehensive code path must call
    ``report_builder.add_partial_data_warning_banner`` immediately
    after ``report_builder.add_title_page`` so the banner sits on
    page 2 (after the cover) and before any KPI dashboard.
    """

    src = Path("app_simple.py").read_text(encoding="utf-8")
    title_pos = index_in_source(src, "report_builder.add_title_page(status")
    banner_pos = src.find(
        "report_builder.add_partial_data_warning_banner(partial_data_warnings)",
        title_pos,
    )
    dashboard_pos = src.find(
        "add_executive_visual_dashboard(report_builder.doc",
        title_pos,
    )
    assert title_pos != -1, "Comprehensive title-page call site missing"
    assert banner_pos != -1, "Comprehensive banner call site missing"
    assert dashboard_pos != -1, "Comprehensive dashboard call site missing"
    assert title_pos < banner_pos < dashboard_pos, (
        f"Banner must sit between title (idx {title_pos}) and "
        f"dashboard (idx {dashboard_pos}); got banner at {banner_pos}"
    )


# ---------------------------------------------------------------------------
# End-to-end render
# ---------------------------------------------------------------------------


def _sample_warnings():
    return [
        {
            "dataset": "customer_pulse",
            "error": "schema_drift:customer_pulse: missing slot(s) rating",
            "kind": "schema_drift",
        },
        {
            "dataset": "ESA_C360_CS_TASK__C",
            "error": "Snowflake table blocked by policy",
            "kind": "policy",
        },
    ]


def test_round48_banner_renders_into_docx_body(tmp_path):
    """Build a minimal report via ``ExecutiveReportBuilder``, call
    the banner method, save the .docx, then re-read it and assert
    each warning entry is present in the body text.
    """

    rb = ExecutiveReportBuilder()
    rb.add_heading("Test Comprehensive Report", level=0)
    rb.add_partial_data_warning_banner(_sample_warnings())

    target = tmp_path / "comp_banner_test.docx"
    rb.save(str(target))
    assert target.exists()

    from docx import Document

    doc = Document(str(target))
    body_text = "\n".join(p.text for p in doc.paragraphs)

    assert_in_source(body_text, "Partial Data Warning", label='body_text')
    assert_in_source(body_text, "customer_pulse", label='body_text')
    assert_in_source(body_text, "schema_drift", label='body_text')
    assert_in_source(body_text, "ESA_C360_CS_TASK__C", label='body_text')
    assert_in_source(body_text, "policy", label='body_text')


def test_round48_banner_no_op_on_empty_input(tmp_path):
    """Calling the banner method with None / [] must produce no
    visible heading (false-positive banner is itself a defect).
    """

    rb = ExecutiveReportBuilder()
    rb.add_heading("Test Comprehensive Report", level=0)
    rb.add_partial_data_warning_banner(None)
    rb.add_partial_data_warning_banner([])

    target = tmp_path / "comp_banner_empty_test.docx"
    rb.save(str(target))

    from docx import Document

    doc = Document(str(target))
    headings = [
        p.text for p in doc.paragraphs if p.style.name.startswith("Heading")
    ]
    assert all("Partial Data Warning" not in h for h in headings), (
        "Comprehensive Word body shows banner heading despite empty "
        "warning input -- false-positive banner"
    )
