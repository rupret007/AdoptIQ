"""Round 169 exact, privacy-safe 17-sheet cross-family parity."""

from __future__ import annotations

import ast
import json
from collections import Counter
from pathlib import Path

import pandas as pd
import pytest

import decision_report_delivery as delivery
import report_iteration_loop as iteration
import report_source_parity as parity
from scripts import run_report_option_matrix as matrix_runner


REPORT_TYPES = {
    "compact": "Compact",
    "comprehensive": "Comprehensive",
    "leader": "Leader",
    "renewal": "Renewal Portfolio",
}
SOURCE_STATE_SHEETS = {
    "Subscriptions",
    "Action_Plans",
    "Adoption_Barriers",
    "Customer_Pulse",
    "TAC_Cases",
    "BEMS",
    "Success_Priorities",
    "External_Incidents",
    "External_Bugs",
    "Defect_Correlations",
}


def _source_context(record_id: str, source_system: str) -> dict[str, object]:
    return {
        "Record_ID": record_id,
        "Record_ID_Data_Quality": "OK",
        "CSSM": "Fixture Owner",
        "Scope_Type": "customer",
        "Scope_Value": "Acme Corporation",
        "Source_System": source_system,
        "Attributed_Team_Members": "Fixture Owner",
    }


def test_fixed_source_export_projection_normalizes_route_aliases() -> None:
    ab_base = {
        **_source_context("AB-1", "Snowflake C360 Adoption Barriers"),
        "Source_Record_URL": (
            "https://ciscosales.lightning.force.com/lightning/r/"
            "C360_CS_Task__c/AB-1/view"
        ),
        "ID": "AB-1",
        "BU_NAME": "Acme Corporation",
        "ACCOUNT_ID_C": "A1",
        "SUBJECT_C": "Calling dependency",
        "DESCRIPTION_C": "A deterministic source description.",
        "STATUS_C": "Open",
        "SEVERITY_C": "P1",
        "OPEN_DATE_C": "2026-07-05",
        "CLOSED_DATE_C": "",
    }
    legacy_enriched = pd.DataFrame(
        [
            {
                **ab_base,
                "title": "Calling dependency",
                "description_2": "A deterministic source description.",
                "ab_category_final": "Uncategorized",
                "sub_technology": "Other / Unclassified",
                "status_norm": "route-only value",
                "severity_norm": "route-only value",
                "open_age_days": 999,
                "open_date": pd.Timestamp("2026-07-05"),
                "_AdoptIQ_Subtechnology_Derived": True,
            }
        ]
    )
    direct_raw = pd.DataFrame([ab_base])

    legacy_public = delivery._prepare_export_frame(  # noqa: SLF001
        legacy_enriched,
        "Adoption_Barriers",
    )
    direct_public = delivery._prepare_export_frame(  # noqa: SLF001
        direct_raw,
        "Adoption_Barriers",
    )

    pd.testing.assert_frame_equal(legacy_public, direct_public)
    assert direct_public.iloc[0]["Status (Normalized)"] == "Open"
    assert direct_public.iloc[0]["Severity (Normalized)"] == "Critical"
    assert "Open Age (Days)" not in direct_public.columns
    assert "description_2" not in direct_public.columns
    assert "_AdoptIQ_Subtechnology_Derived" not in direct_public.columns

    classified = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame(
            [
                {
                    **ab_base,
                    "AB_CATEGORY_C": "Raw category alias",
                    "PRODUCT_NAME_C": "Raw product alias",
                    "ab_category_final": "Canonical category",
                    "sub_technology": "Canonical technology",
                    "bemscsc_refs": "CJPIM-123; not-a-reference; free text",
                    "DESCRIPTION_C": "Also references CSCabc123 and CJPIM-123.",
                }
            ]
        ),
        "Adoption_Barriers",
    )
    assert classified.iloc[0]["Barrier Category (Final)"] == "Canonical category"
    assert classified.iloc[0]["sub_technology"] == "Canonical technology"
    assert classified.iloc[0]["bemscsc_refs"] == "CJPIM-123, CSCABC123"

    pulse_common = {
        **_source_context("CP-1", "Snowflake C360 Customer Pulse"),
        "Source_Record_URL": (
            "https://ciscosales.lightning.force.com/lightning/r/"
            "ESA_C360_Customer_Pulse__c/CP-1/view"
        ),
        "ID": "CP-1",
        "BU_NAME": "Acme Corporation",
        "CUSTOMER_PULSE__C": "Red",
        "PULSE_RATING__C": "Poor",
        "SCORE__C": 1,
        "PULSE_DATE_C": "2026-07-10",
        "COMMENTS__C": "Recovery plan requested.",
    }
    pulse_account_alias = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame([{**pulse_common, "ACCOUNT__C": "A1"}]),
        "Customer_Pulse",
    )
    pulse_normalized_alias = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame([{**pulse_common, "ACCOUNT_ID_C": "A1"}]),
        "Customer_Pulse",
    )
    pd.testing.assert_frame_equal(pulse_account_alias, pulse_normalized_alias)
    assert pulse_account_alias.iloc[0]["Account ID"] == "A1"


def test_fixed_source_export_projection_removes_route_only_attribution_aliases() -> None:
    subscription = {
        **_source_context("SUB-1", "Snowflake subscriptions"),
        "SUBSCRIPTION_ID": "SUB-1",
        "ACCOUNT_ID_C": "A1",
        "BU_NAME": "Acme Corporation",
        "PRODUCT_NAME": "Webex Calling",
        "CSSM_EMAIL": "owner@example.invalid",
        "TECHNOLOGY_C": "Webex Calling",
        "SUB_TECHNOLOGY_C": "Calling",
        "SUBSCRIPTION_STATUS": "Active",
        "RENEWAL_DATE": "2027-01-31",
        "START_DATE": "2026-02-01",
        "CONTRACT_END_DATE": "2027-01-31",
        "ARR": 125000,
        "CURRENCY": "USD",
    }
    direct = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame(
            [
                {
                    **subscription,
                    "manager_name": "Fixture Manager",
                    "cssm_name": "Fixture Owner",
                    "cssm_email": "owner@example.invalid",
                }
            ]
        ),
        "Subscriptions",
    )
    canonical = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame([subscription]),
        "Subscriptions",
    )
    pd.testing.assert_frame_equal(direct, canonical)
    assert not {"manager_name", "cssm_name", "cssm_email"} & set(direct.columns)
    assert not {"FIXTURE_MEMBER", "LOCAL_ACCEPTANCE_RECORD_ID"} & set(direct.columns)
    assert direct.iloc[0]["SUB_TECHNOLOGY_C"] == "Calling"
    assert direct.iloc[0]["Status"] == "Active"
    assert direct.iloc[0]["Renewal Date"] == "2027-01-31T00:00:00Z"
    assert direct.iloc[0]["ARR"] == "125000"
    assert direct.iloc[0]["Currency"] == "USD"

    priority = {
        **_source_context("SP-1", "Snowflake C360 Success Priorities"),
        "Source_Record_URL": (
            "https://ciscosales.lightning.force.com/lightning/r/"
            "ESA_C360_SUCCESS_PRIORITY__C/SP-1/view"
        ),
        "ID": "SP-1",
        "ACCOUNT_ID_C": "A1",
        "RELATED_CUSTOMER__C": "Acme Corporation",
        "STATUS_C": "Active",
        "SUCCESS_PRIORITY_TITLE__C": "Complete migration readiness review",
        "DESCRIPTION__C": "Verify the final outcome with the customer.",
        "PRIORITY_LEVEL__C": "High",
        "OPEN_DATE_C": "2026-07-01",
        "CLOSED_DATE_C": "",
        "OWNERID": "owner-1",
        "COMMENTS__C": "Dated follow-up required.",
        "CSSM_EMAIL": "owner@example.invalid",
    }
    legacy = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame(
            [
                {
                    **priority,
                    "Source_Reported_Attribution": "Fixture Owner",
                    "Attribution_Basis": "route-local adapter metadata",
                }
            ]
        ),
        "Success_Priorities",
    )
    direct_priority = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame([priority]),
        "Success_Priorities",
    )
    pd.testing.assert_frame_equal(legacy, direct_priority)
    assert not {"FIXTURE_MEMBER", "LOCAL_ACCEPTANCE_RECORD_ID"} & set(legacy.columns)
    assert legacy.iloc[0]["Related Customer"] == "Acme Corporation"
    assert legacy.iloc[0]["Success Priority Title"] == (
        "Complete migration readiness review"
    )
    assert legacy.iloc[0]["Description"] == (
        "Verify the final outcome with the customer."
    )
    assert legacy.iloc[0]["Priority Level"] == "High"
    assert legacy.iloc[0]["Open Date"] == "2026-07-01T00:00:00Z"
    assert legacy.iloc[0]["Owner"] == "owner-1"
    assert legacy.iloc[0]["Comments"] == "Dated follow-up required."


def test_fixed_tac_projection_normalizes_date_and_boolean_cell_types() -> None:
    common = {
        **_source_context("SR-1", "CSOne"),
        "Customer": "Acme Corporation",
        "ACCOUNT_ID_C": "A1",
        "SR Number": "SR-1",
        "Title": "Support evidence",
        "Severity": "P2",
        "Case Status": "Open",
        "Date/Time Closed": "",
    }
    text_route = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame(
            [
                {
                    **common,
                    "Date/Time Opened": "2026-07-10T12:30:00Z",
                    "is_open": "1",
                    "is_closed": "0",
                    "is_bems": "false",
                    "open_age_days": "37",
                }
            ]
        ),
        "TAC_Cases",
    )
    typed_route = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame(
            [
                {
                    **common,
                    "Date/Time Opened": pd.Timestamp(
                        "2026-07-10T12:30:00Z"
                    ),
                    "is_open": True,
                    "is_closed": False,
                    "is_bems": False,
                    "open_age_days": 37.0,
                }
            ]
        ),
        "TAC_Cases",
    )

    pd.testing.assert_frame_equal(text_route, typed_route)
    assert text_route.iloc[0]["Date/Time Opened"] == "2026-07-10T12:30:00Z"
    assert bool(text_route.iloc[0]["Is Open"]) is True
    assert text_route.iloc[0]["Open Age (Days)"] == 37

    fractional_age = delivery._prepare_export_frame(  # noqa: SLF001
        pd.DataFrame(
            [
                {
                    **common,
                    "Date/Time Opened": "2026-07-10T12:30:00Z",
                    "open_age_days": 1.5,
                }
            ]
        ),
        "TAC_Cases",
    )
    assert fractional_age.iloc[0]["Open Age (Days)"] == "1.5"


def test_barrier_technology_enrichment_requires_one_authoritative_value() -> None:
    barriers = pd.DataFrame(
        [
            {
                "ID": "AB-1",
                "ACCOUNT_ID_C": "A1",
                "sub_technology": "Other / Unclassified",
                "_AdoptIQ_Subtechnology_Derived": True,
            },
            {
                "ID": "AB-2",
                "ACCOUNT_ID_C": "A2",
                "sub_technology": "Other / Unclassified",
                "_AdoptIQ_Subtechnology_Derived": True,
            },
            {"ID": "AB-3", "ACCOUNT_ID_C": "A3"},
            {
                "ID": "AB-4",
                "ACCOUNT_ID_C": "A2",
                "sub_technology": "Webex Calling",
            },
        ]
    )
    subscriptions = pd.DataFrame(
        [
            {"ACCOUNT_ID_C": "A1", "TECHNOLOGY_C": "Webex Calling"},
            {"ACCOUNT_ID_C": "A1", "TECHNOLOGY_C": "Webex Calling"},
            {"ACCOUNT_ID_C": "A2", "TECHNOLOGY_C": "Cisco UCCE"},
            {"ACCOUNT_ID_C": "A2", "TECHNOLOGY_C": "Cisco UCCX"},
        ]
    )

    enriched = delivery._enrich_barrier_technology_from_subscriptions(  # noqa: SLF001
        barriers,
        subscriptions,
    ).set_index("ID")

    assert enriched.loc["AB-1", "sub_technology"] == "Webex Calling"
    assert enriched.loc["AB-2", "sub_technology"] == "Other / Unclassified"
    assert enriched.loc["AB-3", "sub_technology"] == "Other / Unclassified"
    assert enriched.loc["AB-4", "sub_technology"] == "Webex Calling"
    assert enriched.attrs["ab_technology_unique_subscription_enrichment_count"] == 1
    assert enriched.attrs["ab_technology_ambiguous_subscription_count"] == 1
    assert enriched.attrs["ab_technology_missing_subscription_count"] == 1


def test_prepare_ab_marks_inferred_technology_but_preserves_source_native() -> None:
    import adoptiq_backend as backend

    source = pd.DataFrame(
        [
            {
                "ID": "AB-DERIVED",
                "ACCOUNT_ID_C": "A1",
                "SUBJECT_C": "Webex Calling dependency",
            },
            {
                "ID": "AB-NATIVE",
                "ACCOUNT_ID_C": "A2",
                "SUBJECT_C": "Generic dependency",
                "sub_technology": "Cisco UCCE",
            },
        ]
    )
    subscriptions = pd.DataFrame(
        [
            {"ACCOUNT_ID_C": "A1", "BU_NAME": "Acme Corporation"},
            {"ACCOUNT_ID_C": "A2", "BU_NAME": "Beta Industries"},
        ]
    )

    prepared = backend._prepare_ab(source, subscriptions).set_index("ID")  # noqa: SLF001

    assert bool(prepared.loc["AB-DERIVED", "_AdoptIQ_Subtechnology_Derived"])
    assert not bool(prepared.loc["AB-NATIVE", "_AdoptIQ_Subtechnology_Derived"])
    assert prepared.loc["AB-NATIVE", "sub_technology"] == "Cisco UCCE"


def test_legacy_report_writers_use_canonical_fixture_and_live_data_modes() -> None:
    source = (Path(__file__).parents[1] / "app_simple.py").read_text(
        encoding="utf-8"
    )
    functions = {
        node.name: node
        for node in ast.parse(source).body
        if isinstance(node, ast.FunctionDef)
    }

    for name in (
        "run_compact_analysis",
        "run_customer_renewal_analysis",
        "run_subscription_analysis",
    ):
        literals = {
            node.value
            for node in ast.walk(functions[name])
            if isinstance(node, ast.Constant) and isinstance(node.value, str)
        }
        assert "Guarded offline fixture" in literals
        assert "Application source path" in literals
        assert "Guarded local acceptance snapshot" not in literals


def _base_frame(sheet_name: str) -> pd.DataFrame:
    identities = [f"{sheet_name}-1", f"{sheet_name}-2"]
    if sheet_name in {
        "Metric_Lineage",
        "Risk_Components",
        "Member_Summary",
        "Account_Summary",
    }:
        data = {
            "Metric_Key": identities,
            "Semantic_Value": ["alpha", "beta"],
            "Source_State": ["available", "available"],
        }
        if sheet_name == "Metric_Lineage":
            data["Grouping"] = ["", ""]
        return pd.DataFrame(data)
    if sheet_name == "Chart_Data":
        return pd.DataFrame(
            {
                "Chart_ID": ["activity_mix", "activity_mix"],
                "Metric_Key": identities,
                "Series": ["Activities", "Activities"],
                "Category": ["Action Plans", "TAC Cases"],
                "Value": [2, 3],
                "Source_State": ["available", "available"],
                "Semantic_Value": ["alpha", "beta"],
            }
        )
    if sheet_name == "Evidence_Links":
        return pd.DataFrame(
            {
                "Evidence_Key": identities,
                "Evidence_Role": ["supporting_record", "supporting_record"],
                "Source_Sheet": ["Action_Plans", "TAC_Cases"],
                "Source_Row_SHA256": ["1" * 64, "2" * 64],
                "Record_ID": identities,
                "Source_State": ["available", "available"],
                "Semantic_Value": ["alpha", "beta"],
            }
        )
    return pd.DataFrame(
        {
            "Record_ID": identities,
            "Attributed_Team_Members": ["Manager-safe A", "Manager-safe B"],
            "Semantic_Value": ["alpha", "beta"],
        }
    )


def _with_typed_family_facts(frame: pd.DataFrame, sheet_name: str, family: str) -> pd.DataFrame:
    shared = frame.copy()
    key = f"legacy.family.{family}.fixture.r1.value.00000000"
    if sheet_name == "Subscriptions":
        family_row = {
            "Record_ID": f"LEGACY-{family.upper()}",
            "Metric_Key": key,
            "Legacy_Record_Type": "Family-specific reported fact",
            "Legacy_Report_Family": family,
            "Legacy_Source_Sheet": "Summary",
            "Legacy_Source_Row_Number": 1,
            "Legacy_Source_Field": "Value",
            "Legacy_Fact_Label": "Family value",
            "Legacy_Fact_Value": family,
            "Semantic_Value": f"family-only-{family}",
        }
    elif sheet_name == "Metric_Lineage":
        family_row = {
            "Metric_Key": key,
            "Grouping": "legacy family-specific reported fact",
            "Semantic_Value": f"family-only-{family}",
            "Source_State": "available",
        }
    elif sheet_name == "Evidence_Links":
        family_row = {
            "Evidence_Key": key,
            "Evidence_Role": "legacy_family_reported_fact",
            "Source_Sheet": "Subscriptions",
            "Record_ID": f"LEGACY-{family.upper()}",
            "Semantic_Value": f"family-only-{family}",
            "Source_State": "available",
        }
    else:
        raise AssertionError(sheet_name)
    return pd.concat([shared, pd.DataFrame([family_row])], ignore_index=True, sort=False)


def _write_workbook(
    path: Path,
    *,
    family: str,
    mutate_sheet: str = "",
    duplicate_identity_sheet: str = "",
    missing_identity_sheet: str = "",
    typed_family_facts: bool = False,
    typed_fact_declared_family: str = "",
    drop_typed_fact_sheet: str = "",
    presentation_fact_kind: str = "",
    untyped_family_fact_sheet: str = "",
    row_scope_sheet: str = "",
    row_scope_type: str = "team",
    row_scope_value: str = "Privacy Safe Manager team",
    info_overrides: dict[str, object] | None = None,
    duplicate_info_item: str = "",
    reverse_rows: bool = False,
    reverse_columns: bool = False,
    omit_sheet: str = "",
    add_extra_sheet: bool = False,
    reverse_sheet_order: bool = False,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    info: dict[str, object] = {
        "Report_Type": REPORT_TYPES[family],
        "Manager": "Privacy Safe Manager",
        "Technology": "All",
        "Scope_Type": "team",
        "Scope_Value": "Privacy Safe Manager team",
        "Days": 90,
        "Data_As_Of_UTC": "2026-08-03T20:00:00Z",
        "Data_As_Of_State": "available",
        "Retrieval_Attempted_At_UTC": "2026-08-03T20:00:00Z",
        "Evaluation_As_Of_UTC": "2026-08-03T21:00:00Z",
        "Data_Mode": "Guarded offline fixture",
        "Live_Source_Validation": "No",
        "Due_Soon_Days": 14,
        "Action_Plan_Age_Bands": "0-30; 31-60; 61-90; 90+; Unknown",
        "Activity_Total_State": "available",
        "TAC_Case_Type_Classified": 2,
        "TAC_Case_Type_Not_Derivable": 0,
        "TAC_Case_Type_Coverage_Pct": 100,
        "Partial_Data_Warning_Count": 0,
        "Fact_Contract_SHA256": "0" * 64,
    }
    info.update(
        {
            f"Sheet_SHA256:{sheet_name}": "0" * 64
            for sheet_name in parity.PARITY_SHEET_NAMES[1:]
        }
    )
    info.update(
        {
            f"Source_State:{sheet_name}": "available"
            for sheet_name in SOURCE_STATE_SHEETS
        }
    )
    info.update(info_overrides or {})
    for warning_index in range(int(info["Partial_Data_Warning_Count"])):
        info[f"Partial_Data_Warning_{warning_index + 1}"] = (
            "privacy-safe warning"
        )
    info_rows = [
        {"Item": key, "Value": value, "Detail": "privacy-safe fixture"}
        for key, value in info.items()
    ]
    if duplicate_info_item:
        info_rows.append(
            {
                "Item": duplicate_info_item,
                "Value": info[duplicate_info_item],
                "Detail": "duplicate",
            }
        )

    frames: dict[str, pd.DataFrame] = {
        "Report_Info": pd.DataFrame(
            info_rows,
            columns=["Item", "Value", "Detail"],
        )
    }
    for sheet_name in parity.PARITY_SHEET_NAMES[1:]:
        frame = _base_frame(sheet_name)
        if typed_family_facts and family in {"compact", "renewal"} and sheet_name in {
            "Subscriptions",
            "Metric_Lineage",
            "Evidence_Links",
        }:
            frame = _with_typed_family_facts(frame, sheet_name, family)
            if (
                typed_fact_declared_family
                and sheet_name == "Subscriptions"
            ):
                frame.loc[
                    frame["Legacy_Record_Type"]
                    .fillna("")
                    .astype(str)
                    .str.casefold()
                    .eq("family-specific reported fact"),
                    "Legacy_Report_Family",
                ] = typed_fact_declared_family
            if drop_typed_fact_sheet == sheet_name:
                frame = frame.iloc[:-1].copy()
        if presentation_fact_kind and sheet_name == "Metric_Lineage":
            if presentation_fact_kind == "account":
                presentation_row = {
                    "Metric_Key": "summary.account.fixture.review_focus",
                    "Grouping": "account",
                    "Canonical_Function": (
                        "decision_report_delivery._comprehensive_account_evidence_rows"
                    ),
                    "Unit": "review focus",
                    "Cross_Family_Parity": "family_specific_presentation_fact",
                    "Source_State": "available",
                }
            else:
                presentation_row = {
                    "Metric_Key": "summary.member.fixture.manager_intervention",
                    "Grouping": "team member",
                    "Canonical_Function": (
                        "decision_report_delivery._leader_intervention_rows"
                    ),
                    "Unit": "manager intervention",
                    "Cross_Family_Parity": "family_specific_presentation_fact",
                    "Source_State": "available",
                }
            frame = pd.concat(
                [frame, pd.DataFrame([presentation_row])],
                ignore_index=True,
                sort=False,
            )
        if presentation_fact_kind and sheet_name == "Evidence_Links":
            account = presentation_fact_kind == "account"
            presentation_row = {
                "Evidence_Key": (
                    "summary.account.fixture.review_focus"
                    if account
                    else "summary.member.fixture.manager_intervention"
                ),
                "Evidence_Role": "derived_summary_row",
                "Evidence_Type": "summary",
                "Source_Sheet": "Account_Summary" if account else "Member_Summary",
                "Unit": "review focus" if account else "manager intervention",
                "Cross_Family_Parity": "family_specific_presentation_fact",
                "Source_State": "available",
            }
            frame = pd.concat(
                [frame, pd.DataFrame([presentation_row])],
                ignore_index=True,
                sort=False,
            )
        if sheet_name == row_scope_sheet:
            frame["Scope_Type"] = row_scope_type
            frame["Scope_Value"] = row_scope_value
        if sheet_name == untyped_family_fact_sheet:
            key_column = "Evidence_Key" if sheet_name == "Evidence_Links" else "Metric_Key"
            if key_column not in frame.columns:
                frame[key_column] = ""
            frame.loc[0, key_column] = "legacy.family.compact.untyped"
        if sheet_name == mutate_sheet:
            frame.loc[0, "Semantic_Value"] = "route-specific drift"
        if sheet_name == duplicate_identity_sheet:
            identity_column = (
                "Evidence_Key"
                if sheet_name == "Evidence_Links"
                else "Metric_Key"
                if sheet_name in {
                    "Metric_Lineage",
                    "Chart_Data",
                    "Risk_Components",
                    "Member_Summary",
                    "Account_Summary",
                }
                else "Record_ID"
            )
            frame.loc[1, identity_column] = frame.loc[0, identity_column]
            if sheet_name == "Evidence_Links":
                for column in (
                    "Evidence_Role",
                    "Source_Sheet",
                    "Source_Row_SHA256",
                    "Record_ID",
                ):
                    frame.loc[1, column] = frame.loc[0, column]
        if sheet_name == missing_identity_sheet:
            identity_column = (
                "Evidence_Key"
                if sheet_name == "Evidence_Links"
                else "Metric_Key"
                if sheet_name in {
                    "Metric_Lineage",
                    "Chart_Data",
                    "Risk_Components",
                    "Member_Summary",
                    "Account_Summary",
                }
                else "Record_ID"
            )
            frame.loc[0, identity_column] = ""
            if sheet_name == "Chart_Data":
                for column in ("Chart_ID", "Series", "Category"):
                    frame.loc[0, column] = ""
            if sheet_name == "Evidence_Links":
                for column in (
                    "Evidence_Role",
                    "Source_Sheet",
                    "Source_Row_SHA256",
                    "Record_ID",
                ):
                    frame.loc[0, column] = ""
        if reverse_rows:
            frame = frame.iloc[::-1].reset_index(drop=True)
        if reverse_columns:
            frame = frame.loc[:, list(reversed(frame.columns))]
        frames[sheet_name] = frame

    digest_values = {
        f"Sheet_SHA256:{sheet_name}": parity.sheet_content_sha256(
            frames[sheet_name]
        )
        for sheet_name in parity.PARITY_SHEET_NAMES[1:]
    }
    for row in info_rows:
        if row["Item"] in digest_values:
            row["Value"] = digest_values[row["Item"]]
    frames["Report_Info"] = pd.DataFrame(
        info_rows,
        columns=["Item", "Value", "Detail"],
    )

    sheet_order = list(parity.PARITY_SHEET_NAMES)
    if reverse_sheet_order:
        sheet_order = list(reversed(sheet_order))
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name in sheet_order:
            if sheet_name == omit_sheet:
                continue
            frames[sheet_name].to_excel(writer, sheet_name=sheet_name, index=False)
        if add_extra_sheet:
            pd.DataFrame({"unexpected": [1]}).to_excel(
                writer,
                sheet_name="Unexpected",
                index=False,
            )


def _quartet(
    tmp_path: Path,
    *,
    family_kwargs: dict[str, dict[str, object]] | None = None,
) -> tuple[dict[str, object], dict[str, dict[str, str]]]:
    results: list[dict[str, object]] = []
    declarations: dict[str, dict[str, str]] = {}
    for family in REPORT_TYPES:
        path = tmp_path / f"{family}.xlsx"
        _write_workbook(
            path,
            family=family,
            **(family_kwargs or {}).get(family, {}),
        )
        scenario = f"fixture_{family}"
        results.append(
            {
                "scenario": scenario,
                "all_passed": True,
                "artifacts": [{"file_type": "xlsx", "debug_path": str(path)}],
            }
        )
        declarations[scenario] = {
            "cohort": "privacy-safe-quartet",
            "family": family,
        }
    return {"results": results}, declarations


def _audit(
    summary: dict[str, object],
    declarations: dict[str, dict[str, str]],
) -> dict[str, object]:
    return matrix_runner._cross_report_source_consistency(
        summary,
        declared_cohorts=declarations,
    )


def test_exact_seventeen_sheet_ssot_and_clean_quartet(tmp_path: Path) -> None:
    summary, declarations = _quartet(tmp_path)

    audit = _audit(summary, declarations)

    assert parity.PARITY_SHEET_NAMES == tuple(delivery.SOURCE_DATA_SHEET_NAMES)
    assert len(parity.PARITY_SHEET_NAMES) == 17
    assert audit["ok"] is True
    assert audit["required_sheets"] == list(delivery.SOURCE_DATA_SHEET_NAMES)
    assert audit["comparisons"] == 17
    assert audit["comparisons_expected"] == 17


def test_real_canonical_quartet_preserves_only_typed_presentation_facts(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    from tests.test_round142_decision_report_delivery import _team_fixture

    results: list[dict[str, object]] = []
    declarations: dict[str, dict[str, str]] = {}
    paths: dict[str, Path] = {}
    for family, report_type in REPORT_TYPES.items():
        facts = delivery.build_report_facts(
            _team_fixture(),
            report_type=report_type,
            scope_type="team",
            scope_value="Dana Manager team",
            manager_name="Dana Manager",
            days=90,
            as_of="2026-08-03T12:00:00Z",
            external_incidents=[
                {
                    "id": "INC-001",
                    "title": "Synthetic incident",
                    "status": "resolved",
                }
            ],
            external_bugs=[
                {"bug_id": "BUG-001", "title": "Synthetic known issue"}
            ],
        )
        sheets = delivery.build_source_data_sheets(facts)
        path = tmp_path / f"canonical-{family}.xlsx"
        delivery.write_source_data_workbook(path, sheets)
        paths[family] = path
        scenario = f"canonical_{family}"
        results.append(
            {
                "scenario": scenario,
                "all_passed": True,
                "artifacts": [{"file_type": "xlsx", "debug_path": str(path)}],
            }
        )
        declarations[scenario] = {
            "cohort": "canonical-real-quartet",
            "family": family,
        }

    summary = {"results": results}
    assert _audit(summary, declarations)["ok"] is True

    comprehensive = load_workbook(paths["comprehensive"])
    lineage = comprehensive["Metric_Lineage"]
    lineage_headers = {
        str(cell.value): cell.column for cell in next(lineage.iter_rows())
    }
    marker_column = lineage_headers["Cross_Family_Parity"]
    marker_rows = [
        row
        for row in range(2, lineage.max_row + 1)
        if lineage.cell(row, marker_column).value
        == "family_specific_presentation_fact"
    ]
    assert len(marker_rows) == 3
    comprehensive.close()

    leader_original = paths["leader"].read_bytes()
    leader = load_workbook(paths["leader"])
    evidence = leader["Evidence_Links"]
    evidence_headers = {
        str(cell.value): cell.column for cell in next(evidence.iter_rows())
    }
    marker_column = evidence_headers["Cross_Family_Parity"]
    marker_row = next(
        row
        for row in range(2, evidence.max_row + 1)
        if evidence.cell(row, marker_column).value
        == "family_specific_presentation_fact"
    )
    evidence.cell(marker_row, marker_column).value = None
    leader.save(paths["leader"])
    leader.close()
    untyped = _audit(summary, declarations)
    assert untyped["ok"] is False
    assert untyped["read_errors"][0]["kind"] == (
        "untyped_family_presentation_fact"
    )

    paths["leader"].write_bytes(leader_original)
    leader = load_workbook(paths["leader"])
    evidence = leader["Evidence_Links"]
    evidence_headers = {
        str(cell.value): cell.column for cell in next(evidence.iter_rows())
    }
    marker_column = evidence_headers["Cross_Family_Parity"]
    role_column = evidence_headers["Evidence_Role"]
    marker_row = next(
        row
        for row in range(2, evidence.max_row + 1)
        if evidence.cell(row, marker_column).value
        == "family_specific_presentation_fact"
    )
    evidence.cell(marker_row, role_column).value = "supporting_record"
    leader.save(paths["leader"])
    leader.close()
    malformed = _audit(summary, declarations)
    assert malformed["ok"] is False
    assert malformed["read_errors"][0]["kind"] == (
        "family_presentation_marker_contract_mismatch"
    )


@pytest.mark.parametrize("sheet_name", parity.PARITY_SHEET_NAMES[1:])
def test_same_identity_semantic_drift_is_detected_on_every_content_sheet(
    tmp_path: Path,
    sheet_name: str,
) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={"leader": {"mutate_sheet": sheet_name}},
    )

    audit = _audit(summary, declarations)

    assert audit["ok"] is False
    assert [item["source_sheet"] for item in audit["mismatches"]] == [
        sheet_name
    ]


def test_row_and_column_order_do_not_change_semantic_identity(tmp_path: Path) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={
            "leader": {"reverse_rows": True, "reverse_columns": True}
        },
    )

    assert _audit(summary, declarations)["ok"] is True


def test_clock_skew_is_bounded_only_by_freshness_and_state_stays_semantic(
    tmp_path: Path,
) -> None:
    clock_skew = {
        "Data_As_Of_UTC": "2026-08-03T20:05:00Z",
        "Retrieval_Attempted_At_UTC": "2026-08-03T20:05:00Z",
        "Evaluation_As_Of_UTC": "2026-08-03T21:05:00Z",
    }
    summary, declarations = _quartet(
        tmp_path / "bounded",
        family_kwargs={"compact": {"info_overrides": clock_skew}},
    )

    bounded = matrix_runner._cross_report_source_consistency(
        summary,
        max_freshness_skew_seconds=300,
        declared_cohorts=declarations,
    )
    beyond = matrix_runner._cross_report_source_consistency(
        summary,
        max_freshness_skew_seconds=299,
        declared_cohorts=declarations,
    )

    assert bounded["ok"] is True
    assert bounded["mismatches"] == []
    assert beyond["ok"] is False
    assert beyond["mismatches"] == []
    assert len(beyond["freshness_mismatches"]) == 1

    state_summary, state_declarations = _quartet(
        tmp_path / "state",
        family_kwargs={
            "leader": {"info_overrides": {"Data_As_Of_State": "partial"}}
        },
    )
    state_audit = _audit(state_summary, state_declarations)
    assert state_audit["ok"] is False
    assert "Report_Info" in {
        item["source_sheet"] for item in state_audit["mismatches"]
    }
    assert len(state_audit["freshness_mismatches"]) == 1


@pytest.mark.parametrize(
    "item",
    [
        "Due_Soon_Days",
        "Action_Plan_Age_Bands",
        "Activity_Total_State",
        "TAC_Case_Type_Classified",
        "TAC_Case_Type_Not_Derivable",
        "TAC_Case_Type_Coverage_Pct",
        "Partial_Data_Warning_Count",
    ],
)
def test_manager_visible_report_info_fact_drift_is_detected(
    tmp_path: Path,
    item: str,
) -> None:
    value: object = (
        1
        if item != "Action_Plan_Age_Bands"
        else "different age-band contract"
    )
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={
            "leader": {"info_overrides": {item: value}},
        },
    )

    audit = _audit(summary, declarations)

    assert audit["ok"] is False
    assert {item["source_sheet"] for item in audit["mismatches"]} == {
        "Report_Info"
    }


@pytest.mark.parametrize(
    ("mutation", "field"),
    [
        ("duplicate_identity_sheet", "duplicate_identity_count"),
        ("missing_identity_sheet", "missing_identity_count"),
    ],
)
def test_duplicate_and_missing_identities_are_explicit_failures(
    tmp_path: Path,
    mutation: str,
    field: str,
) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={"leader": {mutation: "BEMS"}},
    )

    audit = _audit(summary, declarations)

    assert audit["ok"] is False
    mismatch = next(
        item for item in audit["mismatches"] if item["source_sheet"] == "BEMS"
    )
    assert mismatch["observed"]["leader"][field] == 1
    assert audit["identity_quality_errors"][0][field] == 1


@pytest.mark.parametrize(
    "mutation",
    ["duplicate_identity_sheet", "missing_identity_sheet"],
)
def test_identical_invalid_identities_cannot_make_all_families_false_green(
    tmp_path: Path,
    mutation: str,
) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={
            family: {mutation: "BEMS"} for family in REPORT_TYPES
        },
    )

    audit = _audit(summary, declarations)

    assert audit["ok"] is False
    assert audit["mismatches"] == []
    assert len(audit["identity_quality_errors"]) == 4
    assert {
        item["source_sheet"] for item in audit["identity_quality_errors"]
    } == {"BEMS"}


def test_only_explicitly_typed_family_facts_are_excluded(tmp_path: Path) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={
            family: {"typed_family_facts": True} for family in REPORT_TYPES
        },
    )

    audit = _audit(summary, declarations)

    assert audit["ok"] is True
    assert audit["mismatches"] == []


def test_row_scope_must_match_report_info_before_normalization(
    tmp_path: Path,
) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={
            "leader": {
                "row_scope_sheet": "BEMS",
                "row_scope_value": "Different Manager team",
            }
        },
    )

    audit = _audit(summary, declarations)

    assert audit["ok"] is False
    assert audit["read_errors"][0]["kind"] == "row_scope_value_mismatch"


def test_team_scope_value_is_exact_and_public_metadata_is_digest_only(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "compact.xlsx"
    _write_workbook(workbook_path, family="compact")

    signature = parity.build_workbook_parity_signature(workbook_path)
    serialized = json.dumps(signature, sort_keys=True)

    assert signature["metadata"]["report_family"] == "compact"
    assert len(signature["metadata"]["scope_contract_sha256"]) == 64
    assert "Privacy Safe Manager" not in serialized
    assert "Privacy Safe Manager team" not in serialized
    assert "manager" not in signature["metadata"]
    assert "scope_value" not in signature["metadata"]


def test_typed_legacy_fact_is_bound_to_report_family_and_all_three_sheets(
    tmp_path: Path,
) -> None:
    wrong_family, wrong_declarations = _quartet(
        tmp_path / "family",
        family_kwargs={
            "compact": {
                "typed_family_facts": True,
                "typed_fact_declared_family": "renewal",
            },
            "renewal": {"typed_family_facts": True},
        },
    )
    missing_evidence, missing_declarations = _quartet(
        tmp_path / "referential",
        family_kwargs={
            "compact": {
                "typed_family_facts": True,
                "drop_typed_fact_sheet": "Evidence_Links",
            },
            "renewal": {"typed_family_facts": True},
        },
    )

    wrong_audit = _audit(wrong_family, wrong_declarations)
    missing_audit = _audit(missing_evidence, missing_declarations)

    assert wrong_audit["ok"] is False
    assert wrong_audit["read_errors"][0]["kind"] == (
        "subscriptions_family_fact_report_family_mismatch"
    )
    assert missing_audit["ok"] is False
    assert missing_audit["read_errors"][0]["kind"] == (
        "legacy_family_fact_cross_sheet_mismatch"
    )


@pytest.mark.parametrize(
    ("family", "kind", "expected_error"),
    [
        ("leader", "account", "account_presentation_fact_wrong_report_family"),
        ("comprehensive", "member", "member_presentation_fact_wrong_report_family"),
    ],
)
def test_presentation_fact_marker_is_bound_to_its_only_allowed_family(
    tmp_path: Path,
    family: str,
    kind: str,
    expected_error: str,
) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={family: {"presentation_fact_kind": kind}},
    )

    audit = _audit(summary, declarations)

    assert audit["ok"] is False
    assert audit["read_errors"][0]["kind"] == expected_error


def test_declared_sheet_digest_must_match_written_content(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    summary, declarations = _quartet(tmp_path)
    leader_result = next(
        result
        for result in summary["results"]
        if result["scenario"] == "fixture_leader"
    )
    leader_path = Path(leader_result["artifacts"][0]["debug_path"])
    workbook = load_workbook(leader_path)
    worksheet = workbook["BEMS"]
    headers = {
        str(cell.value): cell.column for cell in next(worksheet.iter_rows())
    }
    worksheet.cell(2, headers["Semantic_Value"]).value = "post-write tamper"
    workbook.save(leader_path)
    workbook.close()

    audit = _audit(summary, declarations)

    assert audit["ok"] is False
    assert audit["mismatches"] == []
    assert audit["read_errors"][0]["kind"] == "sheet_sha256_mismatch"


@pytest.mark.parametrize(
    "sheet_name",
    ["Subscriptions", "Metric_Lineage", "Evidence_Links"],
)
def test_untyped_family_fact_fails_closed(tmp_path: Path, sheet_name: str) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={"compact": {"untyped_family_fact_sheet": sheet_name}},
    )

    audit = _audit(summary, declarations)

    assert audit["ok"] is False
    assert len(audit["read_errors"]) == 1
    assert "untyped_family_fact" in audit["read_errors"][0]["kind"]


@pytest.mark.parametrize(
    "kwargs",
    [
        {"duplicate_info_item": "Manager"},
        {"info_overrides": {"Source_State:TAC_Cases": "maybe"}},
        {
            "info_overrides": {
                "Data_As_Of_UTC": "2026-08-04T00:00:00Z",
                "Evaluation_As_Of_UTC": "2026-08-03T21:00:00Z",
            }
        },
        {"omit_sheet": "Risk_Components"},
        {"add_extra_sheet": True},
        {"reverse_sheet_order": True},
    ],
)
def test_strict_report_info_and_sheet_inventory_fail_closed(
    tmp_path: Path,
    kwargs: dict[str, object],
) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={"leader": kwargs},
    )

    audit = _audit(summary, declarations)

    assert audit["ok"] is False
    assert len(audit["read_errors"]) == 1


def test_declared_cohort_rejects_missing_and_duplicate_family(tmp_path: Path) -> None:
    summary, declarations = _quartet(tmp_path)
    missing = dict(declarations)
    missing.pop("fixture_leader")
    duplicate = dict(declarations)
    duplicate["fixture_renewal"] = {
        "cohort": "privacy-safe-quartet",
        "family": "leader",
    }

    missing_audit = _audit(summary, missing)
    duplicate_audit = _audit(summary, duplicate)

    assert missing_audit["ok"] is False
    assert missing_audit["incomplete_equivalent_scope_groups"]
    assert duplicate_audit["ok"] is False
    assert duplicate_audit["duplicate_family_groups"]


def test_no_declared_cohort_fails_closed(tmp_path: Path) -> None:
    summary, _declarations = _quartet(tmp_path)

    audit = matrix_runner._cross_report_source_consistency(
        summary,
        declared_cohorts={},
    )

    assert audit["ok"] is False
    assert audit["comparison_requirement_met"] is False
    assert audit["cohort_declaration_errors"] == [
        {"kind": "no_declared_parity_cohort"}
    ]


def test_declared_cohort_rejects_missing_result_and_scope_drift(tmp_path: Path) -> None:
    summary, declarations = _quartet(tmp_path)
    missing_summary = {"results": summary["results"][:-1]}
    drift_summary, drift_declarations = _quartet(
        tmp_path / "scope-drift",
        family_kwargs={
            "leader": {"info_overrides": {"Technology": "Webex Calling"}}
        },
    )

    missing_audit = _audit(missing_summary, declarations)
    drift_audit = _audit(drift_summary, drift_declarations)

    assert missing_audit["ok"] is False
    assert missing_audit["cohort_membership_errors"][0]["kind"] == (
        "scenario_result_missing"
    )
    assert drift_audit["ok"] is False
    assert len(drift_audit["scope_mismatches"]) == 1


def test_audit_output_never_retains_source_values_or_paths(tmp_path: Path) -> None:
    summary, declarations = _quartet(
        tmp_path,
        family_kwargs={"leader": {"mutate_sheet": "External_Bugs"}},
    )

    serialized = json.dumps(_audit(summary, declarations), sort_keys=True)

    assert "Privacy Safe Manager" not in serialized
    assert "External_Bugs-1" not in serialized
    assert str(tmp_path) not in serialized
    assert "route-specific drift" not in serialized


def test_matrix_declarations_are_exact_and_non_equivalent_sweeps_are_unmarked() -> None:
    local = iteration.build_local_acceptance_option_matrix()
    local_declarations = matrix_runner._declared_source_parity_cohorts(
        local,
        list(local),
    )
    local_counts = Counter(item["cohort"] for item in local_declarations.values())
    assert local_counts == {
        "local-primary-team": 4,
        "local-primary-customer": 4,
    }
    primary_team = [
        local[key]
        for key, item in local_declarations.items()
        if item["cohort"] == "local-primary-team"
    ]
    assert {
        str(scenario.payload.get("technology") or "All")
        for scenario in primary_team
    } == {"All"}
    assert all(
        not scenario.source_parity_cohort
        for key, scenario in local.items()
        if key.startswith(("b_", "e_", "f_"))
    )

    multi = iteration.build_local_acceptance_multi_manager_matrix()
    multi_declarations = matrix_runner._declared_source_parity_cohorts(
        multi,
        list(multi),
    )
    multi_counts = Counter(item["cohort"] for item in multi_declarations.values())
    assert multi_counts == {
        "local-all-managers-team": 4,
        "local-primary-manager-team": 4,
        "local-secondary-manager-team": 4,
    }
    assert all(
        set(item["family"] for item in multi_declarations.values() if item["cohort"] == cohort)
        == set(matrix_runner._CROSS_REPORT_REQUIRED_FAMILIES)
        for cohort in multi_counts
    )
