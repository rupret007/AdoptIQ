from __future__ import annotations

from pathlib import Path

import pandas as pd

import app_simple
import canonical_metrics as cm
import decision_report_delivery as delivery
from compact_report_formatter import CompactReportFormatter
from executive_intelligence_formatter import ExecutiveIntelligenceFormatter
from report_consistency import validate_report_consistency


def _unavailable_frame(detail: str = "fixture source unavailable") -> pd.DataFrame:
    frame = pd.DataFrame()
    frame.attrs["source_unavailable"] = True
    frame.attrs["source_unavailable_detail"] = detail
    return frame


def _tile_values(formatter) -> list[str]:
    table = formatter.doc.tables[0]
    return [cell.text.strip() for cell in table.rows[1].cells]


def test_customer_universe_is_id_first_across_renamed_sources() -> None:
    subscriptions = pd.DataFrame(
        {"ACCOUNT_ID_C": ["001"], "BU_NAME": ["Legacy Brand"]}
    )
    action_plans = pd.DataFrame(
        {"ACCOUNT_ID_C": ["001"], "BU_NAME": ["New Brand"]}
    )

    assert cm.list_customers(
        subs_df=subscriptions,
        action_plans_df=action_plans,
    ) == ["Legacy Brand"]
    assert cm.count_customers(
        subs_df=subscriptions,
        action_plans_df=action_plans,
    ) == 1


def test_adoption_barrier_sources_are_joined_with_id_dedupe_and_provenance() -> None:
    snowflake = pd.DataFrame(
        {
            "ID": ["AB-1", "AB-2"],
            "BU_NAME": ["Acme", "Beta"],
            "title": ["Shared barrier", "Snowflake only"],
            "description": ["", "Primary detail"],
        }
    )
    csconsole = pd.DataFrame(
        {
            "ID": ["AB-1", "AB-3"],
            "BU_NAME": ["Acme", "Gamma"],
            "title": ["Shared barrier", "CSConsole only"],
            "description": ["Secondary detail", "Additional record"],
        }
    )

    merged = cm.merge_adoption_barrier_sources(
        [snowflake, csconsole],
        source_labels=["Snowflake Adoption Barriers", "CSConsole Adoption Barriers"],
    )

    assert merged["ID"].tolist() == ["AB-1", "AB-2", "AB-3"]
    shared = merged.loc[merged["ID"].eq("AB-1")].iloc[0]
    assert shared["description"] == "Secondary detail"
    assert shared["Source_System"] == (
        "Snowflake Adoption Barriers + CSConsole Adoption Barriers"
    )
    assert cm.count_total_barriers(merged) == 3


def test_consistency_gate_can_enforce_all_applicable_customer_sources() -> None:
    ab = pd.DataFrame({"ID": ["AB1"], "BU_NAME": ["AB Only"]})
    tac = pd.DataFrame({"Case #": ["SR1"], "Customer Name": ["TAC Only"]})
    pulse = pd.DataFrame({"BU_NAME": ["Pulse Only"]})
    subscriptions = pd.DataFrame({"ACCOUNT_ID_C": ["S1"], "BU_NAME": ["Sub Only"]})
    success_priorities = pd.DataFrame({"ACCOUNT_ID_C": ["SP1"], "BU_NAME": ["SP Only"]})
    extras = [subscriptions, success_priorities]
    portfolio_metrics = cm.build_portfolio_metrics(ab_df=ab, csone_df=tac)
    portfolio_metrics["total_customers"] = cm.count_customers(
        ab_df=ab,
        csone_df=tac,
        pulse_df=pulse,
        extra_frames=extras,
    )

    result = validate_report_consistency(
        ab,
        tac,
        portfolio_metrics=portfolio_metrics,
        pulse_df=pulse,
        extra_frames=extras,
        include_all_customer_sources=True,
    )

    assert result["is_valid"] is True
    assert result["metrics"]["total_customers"] == 5


def test_decision_report_analyzes_every_customer_bearing_source() -> None:
    team_data = {
        "Owner": {
            "subscriptions": pd.DataFrame(
                [{"ACCOUNT_ID_C": "S1", "BU_NAME": "Subscription Only"}]
            ),
            "action_plans": pd.DataFrame(
                [{"ID": "AP1", "ACCOUNT_ID_C": "A1", "BU_NAME": "Plan Only", "STATUS_C": "Open"}]
            ),
            "adoption_barriers": pd.DataFrame(
                [{"ID": "AB1", "ACCOUNT_ID_C": "B1", "BU_NAME": "Barrier Only", "STATUS_C": "Open"}]
            ),
            "customer_pulse": pd.DataFrame(
                [{"ID": "CP1", "ACCOUNT_ID_C": "P1", "BU_NAME": "Pulse Only", "PULSE_RATING__C": "Poor"}]
            ),
            "tac_cases": pd.DataFrame(
                [
                    {
                        "SR Number": "7001",
                        "ACCOUNT_ID_C": "T1",
                        "Customer": "TAC Only",
                        "Severity": "2",
                        "Date/Time Opened": "2026-08-01T12:00:00Z",
                    }
                ]
            ),
            "success_priorities": pd.DataFrame(
                [{"ID": "SP1", "ACCOUNT_ID_C": "SP1", "RELATED_CUSTOMER__C": "Priority Only"}]
            ),
        }
    }

    facts = delivery.build_report_facts(
        team_data,
        report_type="Comprehensive",
        scope_type="team",
        scope_value="Owner team",
        manager_name="Owner",
        technology="All",
        days=90,
        as_of="2026-08-11T12:00:00Z",
    )
    sheets = delivery.build_source_data_sheets(facts)
    document = delivery.build_concise_word_document(facts)
    contract = delivery.validate_cross_artifact_contract(facts, sheets, document)

    assert facts["kpis"]["customers"] == 6
    assert set(facts["risk_profiles"]) == {
        "Barrier Only",
        "Plan Only",
        "Priority Only",
        "Pulse Only",
        "Subscription Only",
        "TAC Only",
    }
    assert contract["ok"], contract["errors"]


def test_compact_dashboard_uses_all_sources_and_withholds_unavailable_tac() -> None:
    formatter = CompactReportFormatter()
    unavailable = _unavailable_frame()
    formatter.add_at_a_glance_dashboard(
        pd.DataFrame({"ACCOUNT_ID_C": ["A"], "BU_NAME": ["AB Only"]}),
        unavailable,
        extra_customer_frames=[
            pd.DataFrame({"ACCOUNT_ID_C": ["S"], "BU_NAME": ["Sub Only"]}),
            pd.DataFrame({"ACCOUNT_ID_C": ["P"], "BU_NAME": ["Plan Only"]}),
        ],
        pulse_df=pd.DataFrame({"ACCOUNT_ID_C": ["C"], "BU_NAME": ["Pulse Only"]}),
    )

    assert _tile_values(formatter) == [
        "4",
        "Unavailable",
        "Unavailable",
        "Unavailable",
        "Unavailable",
    ]
    text = "\n".join(paragraph.text for paragraph in formatter.doc.paragraphs)
    assert "must not be interpreted as zero" in text


def test_ei_dashboard_uses_all_sources_and_withholds_unavailable_tac() -> None:
    formatter = ExecutiveIntelligenceFormatter()
    unavailable = _unavailable_frame()
    formatter.add_executive_dashboard(
        pd.DataFrame({"ACCOUNT_ID_C": ["A"], "BU_NAME": ["AB Only"]}),
        unavailable,
        {},
        {},
        team_subs_df=pd.DataFrame(
            {"ACCOUNT_ID_C": ["S"], "BU_NAME": ["Sub Only"]}
        ),
        csconsole_action_plans=pd.DataFrame(
            {"ACCOUNT_ID_C": ["P"], "BU_NAME": ["Plan Only"]}
        ),
        csconsole_customer_pulse=pd.DataFrame(
            {"ACCOUNT_ID_C": ["C"], "BU_NAME": ["Pulse Only"]}
        ),
        csconsole_success_priorities=pd.DataFrame(
            {"ACCOUNT_ID_C": ["SP"], "BU_NAME": ["Priority Only"]}
        ),
        csconsole_adoption_barriers=pd.DataFrame(
            {"ACCOUNT_ID_C": ["CAB"], "BU_NAME": ["CSConsole AB Only"]}
        ),
    )

    assert _tile_values(formatter)[:5] == [
        "6",
        "Unavailable",
        "Unavailable",
        "Unavailable",
        "Unavailable",
    ]
    formatter.add_bems_escalation_section(unavailable)
    text = "\n".join(paragraph.text for paragraph in formatter.doc.paragraphs)
    assert "must not be interpreted as zero" in text
    assert "Total BEMS Escalations: 0" not in text


def test_named_technology_subscription_scope_fails_closed() -> None:
    subscriptions = pd.DataFrame(
        {
            "ACCOUNT_ID_C": ["CALL", "MEET"],
            "BU_NAME": ["Calling Customer", "Meetings Customer"],
            "TECHNOLOGY_C": ["Webex Calling", "Webex Meetings"],
            "SUB_TECHNOLOGY_C": ["Calling", "Meetings"],
        }
    )
    scoped = app_simple._r162_scope_subscription_customers(
        subscriptions,
        "Webex Calling",
    )
    assert scoped["BU_NAME"].tolist() == ["Calling Customer"]

    contact_center = subscriptions.iloc[[0]].copy()
    contact_center.loc[:, "BU_NAME"] = "Contact Center Customer"
    contact_center.loc[:, "TECHNOLOGY_C"] = "Webex Contact Center"
    mixed = pd.concat([contact_center, subscriptions.iloc[[1]]], ignore_index=True)
    acc_scoped = app_simple._r162_scope_subscription_customers(
        mixed,
        "All Contact Center",
    )
    assert acc_scoped["BU_NAME"].tolist() == ["Contact Center Customer"]

    unscopable = app_simple._r162_scope_subscription_customers(
        subscriptions[["ACCOUNT_ID_C", "BU_NAME"]],
        "Webex Calling",
    )
    assert unscopable.empty
    assert cm.source_data_state(unscopable)["state"] == "unavailable"


def test_comprehensive_threads_scoped_joined_sources_to_headline_and_gate() -> None:
    source = Path(app_simple.__file__).read_text(encoding="utf-8")

    assert "_r47_comp_total_narrow = len(all_customers_comprehensive)" in source
    assert "customer_pulse_df=_count_customer_pulse" in source
    assert "subscriptions_df=team_subs_for_customer_counting" in source
    assert "extra_frames=_comp_extra_customer_frames or None" in source
    assert "include_all_customer_sources=True" in source
    assert "adoption_barriers=_r142_all_adoption_barriers" in source


def test_single_customer_subscription_search_retains_technology_fields() -> None:
    import adoptiq_backend

    source = Path(adoptiq_backend.__file__).read_text(encoding="utf-8")
    search_block = source.split("def search_subscriptions_by_customer", 1)[1].split(
        "def get_subscription_details", 1
    )[0]

    assert '_column_or_default_expr(dsm_columns, "TECHNOLOGY_C"' in search_block
    assert '_column_or_default_expr(dsm_columns, "SUB_TECHNOLOGY_C"' in search_block
    assert "_dsm_attribution_email_columns" in search_block


def test_csone_discovery_excludes_case_variants_and_unsupported_xls() -> None:
    assert app_simple._r161_2_is_adoptiq_output_csone_filename(
        "ADOPTIQ_REPORT_X.xlsx"
    )
    assert app_simple._r161_2_is_adoptiq_output_csone_filename(
        "ADOPTIQ ENHANCED PREMIUM COLLAB SUMMARY-2026.xlsx"
    )
    assert app_simple._r162_csone_workbook_is_readable("legacy.xls") is False


def test_csone_discovery_survives_per_file_mtime_race(
    tmp_path: Path,
    monkeypatch,
) -> None:
    from openpyxl import Workbook

    older = tmp_path / "csone_older.xlsx"
    raced = tmp_path / "csone_newer.xlsx"
    for path in (older, raced):
        workbook = Workbook()
        workbook.save(path)

    real_getmtime = app_simple.os.path.getmtime

    def race_getmtime(path: str) -> float:
        if Path(path) == raced:
            raise OSError("OneDrive hydration race")
        return real_getmtime(path)

    monkeypatch.setitem(
        app_simple.app.config,
        "CSONE_ONEDRIVE_FOLDER",
        str(tmp_path),
    )
    monkeypatch.setattr(app_simple.os.path, "getmtime", race_getmtime)

    selected, state, count = app_simple.get_latest_csone_from_folder_diag()

    assert selected == str(older)
    assert state == "synced"
    assert count == 2
