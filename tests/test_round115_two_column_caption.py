"""Round 115 / Build 84 -- two-column KPI card citation de-clutter.

Build 83 live acceptance (Leader report) still carried **125 in-cell
``[Source: ...]`` citations**, ALL in two-column per-CSSM "Metric | Value"
KPI cards (the R82 Phase B3 per-row inline path).  These visually matched
the user's "throughout the report... clutters up" complaint just like the
multi-column matrices R114 fixed.

Round 115 extends the R114 caption-below-table treatment to two-column
cards: the value cells go clean and ONE aggregated ``Sources: ...`` caption
is appended below the table, preserving the R82 per-source taxonomy by
grouping the cited rows' labels by their source system.

These tests pin:

1. A two-column card receives exactly ONE caption + ZERO in-cell citations.
2. The caption aggregates DISTINCT R82 source systems (taxonomy preserved).
3. A single-source card lists only that one system.
4. Idempotency: a second pass adds no duplicate caption and does not re-save.
5. Mixed-source aggregation groups labels under their own systems.
6. The strict quality gate reports ZERO unbacked two-column claims when the
   caption is present (caption-backing, not a weakened check).
7. Rows already cited in-cell by an older build are left alone (back-compat).

Round 115 / Build 84.  Made-with: Cursor.
"""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402

from report_iteration_loop import evaluate_report_quality  # noqa: E402
from report_source_injector import (  # noqa: E402
    _MATRIX_SOURCE_CAPTION_PREFIX,
    _build_two_column_source_caption,
    inject_source_citations_into_docx,
)


def _build_two_column_card(path: Path) -> None:
    """A per-CSSM "Metric | Value" KPI card with mixed-source rows."""
    doc = Document()
    doc.add_heading("Greg Dolberry", level=2)
    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Adoption Barriers"
    table.rows[1].cells[1].text = "3"
    table.rows[2].cells[0].text = "Action Plans"
    table.rows[2].cells[1].text = "16"
    table.rows[3].cells[0].text = "Total Support Cases"
    table.rows[3].cells[1].text = "4"
    doc.save(str(path))


def _build_single_source_card(path: Path) -> None:
    doc = Document()
    doc.add_heading("Summary", level=2)
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Adoption Barriers"
    table.rows[1].cells[1].text = "68"
    table.rows[2].cells[0].text = "Customer Pulse"
    table.rows[2].cells[1].text = "12"
    doc.save(str(path))


def _captions(path: Path) -> list[str]:
    doc = Document(str(path))
    return [
        p.text.strip()
        for p in doc.paragraphs
        if (p.text or "").strip().startswith(_MATRIX_SOURCE_CAPTION_PREFIX)
    ]


def _all_data_cells(path: Path) -> list[str]:
    doc = Document(str(path))
    out: list[str] = []
    for table in doc.tables:
        for row in list(table.rows)[1:]:
            for cell in row.cells:
                out.append(cell.text)
    return out


def test_two_column_card_gets_one_caption_zero_in_cell(tmp_path: Path) -> None:
    docx = tmp_path / "card.docx"
    _build_two_column_card(docx)

    counts = inject_source_citations_into_docx(docx)

    assert counts["table_captions_added"] == 1, counts
    assert counts.get("table_cells_injected", 0) == 0, counts
    for cell in _all_data_cells(docx):
        assert "[source:" not in cell.lower(), cell
    assert len(_captions(docx)) == 1


def test_caption_aggregates_distinct_source_systems(tmp_path: Path) -> None:
    docx = tmp_path / "card.docx"
    _build_two_column_card(docx)
    inject_source_citations_into_docx(docx)

    caption = _captions(docx)[0]
    # AB + AP -> CSConsole; Total Support Cases -> CSOne.
    assert "Snowflake CSConsole" in caption, caption
    assert "Snowflake CSOne" in caption, caption
    assert caption.startswith(_MATRIX_SOURCE_CAPTION_PREFIX)


def test_single_source_card_lists_only_that_system(tmp_path: Path) -> None:
    docx = tmp_path / "single.docx"
    _build_single_source_card(docx)
    inject_source_citations_into_docx(docx)

    caption = _captions(docx)[0]
    assert "Snowflake CSConsole" in caption, caption
    assert "Snowflake CSOne" not in caption, caption


def test_two_column_caption_idempotent(tmp_path: Path) -> None:
    docx = tmp_path / "card.docx"
    _build_two_column_card(docx)

    first = inject_source_citations_into_docx(docx)
    first_text = "\n".join(p.text for p in Document(str(docx)).paragraphs)
    first_mtime = docx.stat().st_mtime

    second = inject_source_citations_into_docx(docx)
    second_text = "\n".join(p.text for p in Document(str(docx)).paragraphs)
    second_mtime = docx.stat().st_mtime

    assert first["table_captions_added"] == 1
    assert second["table_captions_added"] == 0, "second pass must add no duplicate caption"
    assert len(_captions(docx)) == 1
    assert first_text == second_text
    assert first_mtime == second_mtime, "idempotent re-run must not re-save"


def test_build_two_column_source_caption_groups_by_tag() -> None:
    rows = [
        ["Adoption Barriers", "3"],
        ["Action Plans", "16"],
        ["Total Support Cases", "4"],
    ]
    caption = _build_two_column_source_caption(rows, [0, 1, 2])
    assert caption.startswith("Sources:")
    assert "Adoption Barriers, Action Plans - Snowflake CSConsole" in caption, caption
    assert "Total Support Cases - Snowflake CSOne" in caption, caption


def test_build_two_column_source_caption_empty_returns_generic() -> None:
    caption = _build_two_column_source_caption([["A", "1"]], [])
    assert caption == "Sources: AdoptIQ Report Data Sources"


def test_build_two_column_source_caption_unresolved_label_is_generic() -> None:
    rows = [["Widgets Shipped", "99"]]
    caption = _build_two_column_source_caption(rows, [0])
    assert "AdoptIQ Report Data Sources" in caption, caption


def test_quality_gate_zero_unbacked_two_column_claims_with_caption(tmp_path: Path) -> None:
    docx = tmp_path / "card.docx"
    _build_two_column_card(docx)

    pre_payload, _pre = evaluate_report_quality(
        docx_path=docx, xlsx_path=None, scenario_key="leader", strict=True
    )
    pre_unbacked = pre_payload.get("unbacked_metric_claim_count", 0)
    assert pre_unbacked > 0, "fixture must start with unbacked two-column claims"

    inject_source_citations_into_docx(docx, scenario_key="leader")

    post_payload, _post = evaluate_report_quality(
        docx_path=docx, xlsx_path=None, scenario_key="leader", strict=True
    )
    post_unbacked = post_payload.get("unbacked_metric_claim_count", 0)
    assert post_unbacked == 0, (
        "the trailing caption must source-back the two-column claims; "
        f"pre={pre_unbacked} post={post_unbacked}"
    )


def test_legacy_in_cell_cited_rows_excluded_from_caption(tmp_path: Path) -> None:
    """A row already cited in-cell by an OLDER build keeps its placement and is
    not double-counted into the caption (back-compat with pre-R115 reports)."""
    doc = Document()
    doc.add_heading("Legacy Card", level=2)
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Adoption Barriers"
    # Simulate an older build's in-cell citation already present.
    table.rows[1].cells[1].text = "68 [Source: Snowflake CSConsole]"
    table.rows[2].cells[0].text = "Total Support Cases"
    table.rows[2].cells[1].text = "381"
    docx = tmp_path / "legacy.docx"
    doc.save(str(docx))

    counts = inject_source_citations_into_docx(docx)

    # The AB row's value cell ("68 [Source: ...]") is no longer a clean
    # numeric value, so it never qualifies as a fresh metric claim -- it
    # keeps its existing in-cell placement and is excluded from the new
    # caption.  Only the clean Total Support Cases row drives the caption.
    assert counts["table_captions_added"] == 1, counts
    assert counts.get("table_cells_injected", 0) == 0, counts
    caption = _captions(docx)[0]
    assert "Snowflake CSOne" in caption, caption
    # The AB label (already cited in-cell) must not be aggregated into the
    # caption -- it keeps its existing in-cell placement.
    assert "Adoption Barriers" not in caption, caption
    # The pre-existing in-cell citation survives untouched.
    ab_value = Document(str(docx)).tables[0].cell(1, 1).text
    assert "Snowflake CSConsole" in ab_value, ab_value
