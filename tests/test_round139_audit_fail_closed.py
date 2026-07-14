"""Round 139 / Build 109 — fail-closed CLI + honest admin audit."""

import importlib.util
import sqlite3
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]


def _load_r114():
    path = REPO / "scripts" / "r114_audit_reports.py"
    spec = importlib.util.spec_from_file_location("r114_audit_reports", path)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(mod)
    return mod


def test_r114_returns_nonzero_when_docx_missing(tmp_path):
    r114 = _load_r114()
    base = tmp_path / "AdoptIQ_Report_Compact_Test"
    base.with_suffix(".docx")  # intentionally absent
    base.with_suffix(".xlsx").write_bytes(b"not-a-real-xlsx")
    rc = r114.main(["--target", f"Compact={base}"])
    assert rc == 1


def test_r114_detects_tac_case_na_in_docx(tmp_path):
    r114 = _load_r114()
    from docx import Document

    docx = tmp_path / "leader.docx"
    doc = Document()
    doc.add_paragraph("Open items include TAC Case: N/A for customer.")
    doc.save(docx)
    findings = r114.audit_docx(docx)
    assert findings.get("tac_case_na", 0) >= 1


def test_r114_dup_case_ids_on_sr_number(tmp_path):
    r114 = _load_r114()
    from openpyxl import Workbook

    xlsx = tmp_path / "cases.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["SR Number", "Severity"])
    ws.append(["SR-1", "P2"])
    ws.append(["SR-1", "P2"])
    wb.save(xlsx)
    findings = r114.audit_xlsx(xlsx)
    assert findings.get("dup_case_ids")


def test_admin_audit_skips_unimplemented_checks(monkeypatch, tmp_path):
    sys.path.insert(0, str(REPO))
    import enhanced_admin_dashboard_v2 as admin

    monkeypatch.setattr(admin, "init_database", lambda: None)

    class _Cursor:
        def execute(self, *args, **kwargs):
            return self

        def fetchone(self):
            return None

    class _Conn:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def cursor(self):
            return _Cursor()

        def commit(self):
            return None

    monkeypatch.setattr(admin, "db_connection", lambda: _Conn())
    monkeypatch.setattr(admin, "log_security_event", lambda *a, **k: None)

    result = admin.audit_report("test-analysis-id-0001")
    checks = {c["check"]: c for c in result["checks"]}
    assert checks["data_sources"]["status"] == "skipped"
    assert checks["bems_detection"]["status"] == "skipped"
    assert checks["references"]["status"] == "skipped"
    assert result["status"] != "excellent"
    assert result["max_score"] == 70
