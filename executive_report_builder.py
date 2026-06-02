"""
ExecutiveReportBuilder - Clean Word report builder for AdoptIQ comprehensive reports.
Wraps python-docx Document with executive-ready formatting and delegates to adoptiq_backend.
"""
import logging
from typing import Optional

from docx import Document
from docx.shared import Inches

from adoptiq_backend import (
    create_executive_title_page,
    append_to_word_report,
    _sanitize_llm_grade_brackets,
)


logger = logging.getLogger(__name__)


# Round 7 / Phase 1.11: heuristic guards applied at the
# parse_ai_output_and_add boundary so prompt-injection attempts that
# arrive *inside* the model output are flagged and the offending
# segment is dropped before being written into the executive Word
# document.  These mirror the negative constraints that the prompt
# layer already enforces (no fabrication, no contact info, no system
# override).
_PROMPT_INJECTION_GUARDS = (
    "ignore previous instructions",
    "ignore the previous instructions",
    "disregard prior",
    "system:",
    "<<sys>>",
    "<|im_start|>",
    "you are now",
    "act as if you",
    "reveal the system prompt",
    "prompt injection",
)
_AI_OUTPUT_DELIMITER = "=== BEGIN AI_REPORT ==="
_AI_OUTPUT_END_DELIMITER = "=== END AI_REPORT ==="


class ExecutiveReportBuilder:
    """Builds executive Word reports with clean formatting (no markdown artifacts)."""

    def __init__(self):
        self.doc = Document()

    def add_title_page(self, manager: str, tech: str, days: int, portfolio_metrics: dict):
        """Add professional executive title page."""
        create_executive_title_page(self.doc, manager, tech, days, portfolio_metrics)

    def add_page_break(self):
        """Add page break."""
        self.doc.add_page_break()

    def add_partial_data_warning_banner(self, partial_data_warnings):
        """Round 48 / F-COMP-PARTIAL-BANNER-MISSING: surface upstream
        fetch failures (schema_drift, column-policy block, timeouts)
        directly in the comprehensive Word report.

        Mirrors the banner already emitted by:
          * ``executive_intelligence_formatter.py`` ~L1610
          * ``app_simple._create_simple_renewal_report`` (R48)
          * ``compact_report_formatter`` ~L2922 (R46)

        Without this banner the comprehensive Word document silently
        consumed a "successful" report built on partial data while
        ``analysis_status[*]['partial_data_warnings']`` correctly
        carried the warnings -- the banner is the user-visible
        contract that the two surfaces agree.

        Defensive: any rendering failure logs at WARNING level and
        is swallowed so the banner can never block the rest of the
        report from being produced.
        """

        if not partial_data_warnings:
            return
        try:
            self.doc.add_heading("\u26a0 Partial Data Warning", level=1)
            # Round 124 / F9: distinguish a scope exclusion (data loaded fine
            # but was filtered out by the requested technology / manager / time
            # window) from a genuine load failure.  Pre-R124 the builder always
            # said "failed to load", which misled the operator when the only
            # warning was ``tech_filter_scope_excluded`` (a scope decision).
            # Ports the R112/F5 branch already used by the compact + renewal
            # Word paths in app_simple.py.
            _r124_scope_kinds = {
                'tech_filter_scope_excluded',
                # Round 125 / A3: All-Managers Webex comprehensive emits
                # ``tech_filter_empty_after_scope`` (the AB set is empty *after*
                # the technology scope filter runs), a scope decision -- not a
                # load failure.  Pre-R125 it fell through to the "failed to
                # load" wording.  Treat it as a scope exclusion alongside the
                # ``tech_filter_scope_excluded`` kind.
                'tech_filter_empty_after_scope',
                'manager_filter_scope_excluded',
                'time_window_scope_excluded',
                'no_onedrive_sync',
                'autodiscovered_empty_after_scope',
            }
            _r124_all_scope = bool(partial_data_warnings) and all(
                str((w or {}).get('kind') or '') in _r124_scope_kinds
                or str((w or {}).get('kind') or '').startswith('tech_filter_scope')
                for w in partial_data_warnings
            )
            if _r124_all_scope:
                self.doc.add_paragraph(
                    "One or more upstream data sources returned data that was "
                    "filtered out by the requested scope (technology filter, "
                    "manager filter, or time window).  The data loaded "
                    "successfully; the bulleted list below names what was "
                    "excluded and why.  Sections affected by the filter are "
                    "reduced rather than missing.  The Excel workbook lists "
                    "the same warnings in its Report_Info / "
                    "Partial_Data_Warning_Count cells."
                )
            else:
                self.doc.add_paragraph(
                    "One or more upstream data sources failed to load for "
                    "this comprehensive report run.  Sections that depend "
                    "on the affected sources are marked \"unavailable\" "
                    "rather than rendered as zero.  The Excel workbook "
                    "lists the same warnings in its Report_Info / "
                    "Partial_Data_Warning_Count cells.  Rerun once the "
                    "source is reachable for a complete picture."
                )
            for _r48_w in partial_data_warnings:
                _r48_ds = str((_r48_w or {}).get('dataset') or 'unknown')
                _r48_err = str((_r48_w or {}).get('error') or 'unknown error')
                _r48_kind = str((_r48_w or {}).get('kind') or 'runtime')
                self.doc.add_paragraph(
                    f"\u2022 {_r48_ds} ({_r48_kind}): {_r48_err}",
                    style='List Bullet',
                )
            self.doc.add_paragraph("")
        except Exception as _r48_banner_err:  # noqa: BLE001
            logger.warning(
                "Round 48 / F-COMP-PARTIAL-BANNER-MISSING: "
                "comprehensive Word banner failed: %s",
                _r48_banner_err,
            )

    def add_heading(self, text: str, level: int = 1):
        """Add heading."""
        self.doc.add_heading(text, level=level)

    def add_paragraph(self, text: str = "", bold_sections=None):
        """Add paragraph.

        Round 7 / Phase 1.12: previously ``bold_sections`` was silently
        ignored, which let callers think they were producing bold runs
        when they weren't.  We now honor a simple, documented contract:
        ``bold_sections`` may be a list of substrings; each occurrence
        in ``text`` is rendered as a bold run, with the surrounding
        text rendered as plain runs.  Callers that don't care can omit
        the parameter (it still defaults to ``None``).
        """
        if not bold_sections:
            self.doc.add_paragraph(text)
            return

        try:
            para = self.doc.add_paragraph()
            remaining = str(text or "")
            patterns = [str(s) for s in bold_sections if s]
            while remaining:
                # Find the earliest match across all patterns; ties
                # break by longest substring so nested labels honor
                # the more specific pattern.
                next_idx = -1
                next_pat = ""
                for pat in patterns:
                    if not pat:
                        continue
                    idx = remaining.find(pat)
                    if idx < 0:
                        continue
                    if (next_idx < 0
                            or idx < next_idx
                            or (idx == next_idx and len(pat) > len(next_pat))):
                        next_idx = idx
                        next_pat = pat
                if next_idx < 0:
                    para.add_run(remaining)
                    break
                if next_idx > 0:
                    para.add_run(remaining[:next_idx])
                bold_run = para.add_run(next_pat)
                bold_run.bold = True
                remaining = remaining[next_idx + len(next_pat):]
        except Exception as _bold_err:
            logger.warning(
                "Round 7 / Phase 1.12: bold_sections rendering "
                "failed (%s); falling back to plain paragraph.",
                _bold_err,
            )
            self.doc.add_paragraph(text)

    def save(
        self,
        path: str,
        customer_name: Optional[str] = None,
        report_subject: Optional[str] = None,
    ):
        """Save document to path.

        Round 25 / Phase E: stamp ``doc.core_properties`` so the
        generated ``.docx`` identifies as AdoptIQ output rather than a
        python-docx-default file with a 2013 created date.  The stamp
        runs *before* ``doc.save(...)`` because python-docx serializes
        ``core.xml`` from the in-memory properties at save time.

        ``customer_name`` and ``report_subject`` are optional; when
        provided they enrich the ``Title`` and ``Subject`` properties
        with portfolio context (e.g. ``"Brian Frazier All Contact
        Center - Executive Analysis - 2026-04-25"``).  When omitted we
        still emit a non-empty title that names the report family and
        date so the Properties dialog never shows an empty title.
        """
        try:
            from datetime import datetime as _r25e_datetime, timezone as _r25e_tz

            _r25e_now = _r25e_datetime.now(_r25e_tz.utc)
            _r25e_customer = (customer_name or "").strip()
            _r25e_subject = (report_subject or "").strip()

            if _r25e_customer:
                _r25e_title = (
                    f"{_r25e_customer} - Executive Analysis - "
                    f"{_r25e_now.strftime('%Y-%m-%d')}"
                )
            else:
                _r25e_title = (
                    f"AdoptIQ Comprehensive Executive Report - "
                    f"{_r25e_now.strftime('%Y-%m-%d')}"
                )

            cp = self.doc.core_properties
            cp.author = "AdoptIQ Executive Report Generator"
            cp.last_modified_by = "AdoptIQ Executive Report Generator"
            cp.title = _r25e_title
            cp.subject = _r25e_subject or "AdoptIQ Comprehensive Executive Report"
            cp.comments = "AdoptIQ Comprehensive Executive Report"
            cp.category = "Executive Analytics"
            cp.created = _r25e_now
            cp.modified = _r25e_now
            try:
                cp.company = "AdoptIQ"
            except Exception:
                pass
        except Exception as _r25e_err:
            logger.debug(
                "Round 25 / Phase E: core-properties stamp skipped on "
                "ExecutiveReportBuilder.save: %s",
                _r25e_err,
            )

        # Round 70 / Phase 1 (#1): the ExecutiveReportBuilder.save path is
        # exercised by the Comprehensive report's CLI smoke harness and by
        # any future test that swaps in this thin builder for the
        # executive_intelligence_formatter. R68/A1 only wired the
        # formatter.save(); without this hook the same comprehensive
        # report from the alternate writer ships unstamped.
        try:
            from _r68_build_label import apply_word_footer as _r68_apply_word_footer
            _r68_apply_word_footer(self.doc)
        except Exception as _r68_err:
            # Round 73 / Phase 1 (F1): promoted to warning so the next
            # missing-footer regression surfaces in the admin error log
            # instead of hiding under the default debug threshold.
            logger.warning("Round 70 / Phase 1: ExecutiveReportBuilder word footer skipped: %s", _r68_err)

        self.doc.save(path)

    def parse_ai_output_and_add(self, ai_output: str):
        """Parse AI markdown output and add to document with clean formatting.

        Round 7 / Phase 1.11: enforce delimiter + negative-constraint
        screening at this layer instead of trusting the upstream
        prompt to have filtered everything.  We:

        1. Strip the ``=== BEGIN AI_REPORT === ... === END AI_REPORT ===``
           fence if present so the renderer never embeds the literal
           delimiters in the document.
        2. Drop any line containing a known prompt-injection pattern
           and emit a WARNING with a hashed digest of the dropped
           segment so we can audit drift over time.
        3. Hand the cleaned text to ``append_to_word_report``.
        """
        cleaned = self._enforce_ai_output_contract(ai_output or "")
        append_to_word_report(self.doc, cleaned)

    @staticmethod
    def _enforce_ai_output_contract(ai_output: str) -> str:
        """Apply Round 7 / Phase 1.11 delimiter + negative constraints.

        Public for testing.
        """
        text = str(ai_output or "")
        # Step 1: unwrap delimiters if present.
        if _AI_OUTPUT_DELIMITER in text:
            try:
                _, _, after = text.partition(_AI_OUTPUT_DELIMITER)
                inner, _, _ = after.partition(_AI_OUTPUT_END_DELIMITER)
                text = inner if inner else after
            except Exception as _delim_err:
                logger.warning(
                    "Round 7 / Phase 1.11: AI output delimiter "
                    "stripping failed (%s); using raw output.",
                    _delim_err,
                )
        # Step 2: drop prompt-injection lines.
        kept = []
        dropped = 0
        for line in text.splitlines():
            lower = line.lower()
            if any(g in lower for g in _PROMPT_INJECTION_GUARDS):
                dropped += 1
                try:
                    import hashlib as _hl
                    _digest = _hl.sha256(line.encode("utf-8")).hexdigest()[:8]
                except Exception:
                    _digest = "unavailable"
                logger.warning(
                    "Round 7 / Phase 1.11: dropped suspected prompt-"
                    "injection line from AI output (digest=%s, len=%d).",
                    _digest, len(line),
                )
                continue
            kept.append(line)
        if dropped:
            logger.warning(
                "Round 7 / Phase 1.11: dropped %d AI output line(s) "
                "matching prompt-injection guards.", dropped,
            )
        cleaned = "\n".join(kept)
        # Round 27: strip stray brackets around the customer health
        # grade letter (e.g., ``Customer Health Score: [C]`` ->
        # ``Customer Health Score: C``).  Defense-in-depth alongside
        # the same rewrite inside ``append_to_word_report``; running
        # it here too ensures any future caller of this contract
        # method also gets the sanitized text.
        cleaned = _sanitize_llm_grade_brackets(cleaned)
        # Round 49 / F-COMP-BEMS-MD-LEAK-R49: strip square brackets
        # around real BEMS / CSC / CSCxx ID patterns from the LLM
        # output before it reaches the comprehensive Word document.
        # Same rationale as the compact path's R49 wire in
        # app_simple._parse_markdown_for_fallback.
        try:
            from report_utils import strip_bems_brackets_from_llm_text as _r49_strip
            cleaned = _r49_strip(cleaned)
        except Exception:
            pass
        return cleaned

    def _add_customer_separator(self):
        """Add visual separator between customer sections (page break)."""
        self.doc.add_page_break()

    def add_historical_context_section(
        self,
        customer_names,
        technology=None,
        *,
        enabled: bool = None,
    ) -> dict:
        """Round 17 / Phase D.2 -- append the CSOne corpus historical
        context section.  Always returns a small status dict the
        caller can log; never raises.

        ``enabled`` defaults to ``Config.CORPUS_KNOWLEDGE_ENABLED``
        when not supplied, so the call site does not have to import
        Config when the feature flag is off.
        """
        try:
            if enabled is None:
                from config import Config as _r17_cfg
                enabled_resolved = bool(getattr(_r17_cfg, "CORPUS_KNOWLEDGE_ENABLED", False))
            else:
                enabled_resolved = bool(enabled)
            from report_corpus_context import (
                build_historical_context,
                render_to_word,
            )
            context = build_historical_context(
                customer_names or [],
                technology=technology,
                enabled=enabled_resolved,
            )
            try:
                render_to_word(self.doc, context)
            except Exception as render_err:  # noqa: BLE001
                logger.warning(
                    "Round 17 / Historical Context: Word render failed (%s); "
                    "skipping section.",
                    render_err,
                )
                return {
                    "rendered": False,
                    "available": context.available,
                    "reason": "render_failed",
                }
            return {
                "rendered": True,
                "available": context.available,
                "entries": len(context.entries),
                "unmatched": len(context.unmatched),
            }
        except Exception as err:  # noqa: BLE001
            logger.warning(
                "Round 17 / Historical Context: section build failed (%s); skipping.",
                err,
            )
            return {"rendered": False, "available": False, "reason": "build_failed"}
