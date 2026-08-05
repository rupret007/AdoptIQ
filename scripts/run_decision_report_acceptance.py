#!/usr/bin/env python3
"""Round 143 one-command acceptance for AdoptIQ decision reports.

``auto`` runs live only when the local app and Snowflake connectivity probe are
green; otherwise it executes the sanitized deterministic fixture and records
the live deferral.  ``live`` never falls back.  Both paths generate Team,
Member, Customer, and Comprehensive pairs twice and write one redacted JSON
summary suitable for a work-machine rollout decision.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import re
import subprocess
import sys
import time
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping
from urllib.parse import quote, urlparse

import openpyxl
import pandas as pd
import requests
from docx import Document


REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import decision_report_delivery as delivery  # noqa: E402
from leader_scope import (  # noqa: E402
    LeaderScopeValidationError,
    filter_leader_subscriptions,
    validate_leader_scope_request,
)
from report_iteration_loop import (  # noqa: E402
    _inspect_docx_visible_charts,
    extract_csrf_token,
    parse_download_name,
)
from scripts.generate_offline_acceptance_artifacts import (  # noqa: E402
    SUPPORTED_SCOPES,
    generate_acceptance_artifacts,
)


SUMMARY_SCHEMA = "decision-report-acceptance/v1"
REQUIRED_LIVE_SOURCE_SHEETS = (
    "Subscriptions",
    "Action_Plans",
    "Adoption_Barriers",
    "Customer_Pulse",
    "TAC_Cases",
)
ACCEPTABLE_LIVE_SOURCE_STATES = frozenset({"available", "zero", "filtered"})
CHART_TITLES = {
    "activity_mix": "Activity Mix by Source",
    "action_plan_status_aging": "Action Plan Status and Aging",
    "risk_distribution": "Customer Risk Distribution",
    "activity_trend": "Activity Trend",
}
ACTIVITY_SHEETS = (
    "Action_Plans",
    "Adoption_Barriers",
    "Customer_Pulse",
    "TAC_Cases",
)
TERMINAL_STATUSES = frozenset({"completed", "error", "cancelled"})


@dataclass(frozen=True)
class ScopeExpectation:
    requested_scope: str
    report_type: str
    scope_type: str
    scope_value: str


def _utc_now() -> str:
    return datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json_safe(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_safe(item) for item in value]
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, pd.Timestamp):
        return None if pd.isna(value) else value.isoformat()
    if value is pd.NA or value is pd.NaT:
        return None
    if hasattr(value, "item"):
        try:
            return _json_safe(value.item())
        except (TypeError, ValueError):
            pass
    if isinstance(value, float) and pd.isna(value):
        return None
    return value


def _write_json(path: Path, payload: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            _json_safe(payload),
            indent=2,
            sort_keys=True,
            allow_nan=False,
        )
        + "\n",
        encoding="utf-8",
    )


def _parse_as_of(raw: str) -> pd.Timestamp:
    parsed = pd.to_datetime(str(raw or "").strip(), errors="coerce", utc=True)
    if pd.isna(parsed):
        raise ValueError("--as-of must be an explicit valid ISO timestamp")
    return parsed


def _ensure_safe_output_dir(path: Path) -> Path:
    target = path.expanduser().resolve()
    git_metadata = (REPO_ROOT / ".git").resolve()
    if target == git_metadata or git_metadata in target.parents:
        raise ValueError("acceptance artifacts cannot be written inside .git")
    try:
        target.relative_to(REPO_ROOT)
    except ValueError:
        return target

    # Artifacts may live in the conventional ignored folder.  Any other
    # in-repo destination must already be ignored, preventing an accidental
    # customer-data commit after a live work-machine run.
    relative = target.relative_to(REPO_ROOT)
    check = subprocess.run(  # noqa: S603
        [
            "git",
            "check-ignore",
            "--quiet",
            "--",
            str(relative / ".adoptiq-artifact-probe"),
        ],
        cwd=REPO_ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    if check.returncode != 0:
        raise ValueError(
            "an output directory inside the repository must be Git-ignored; "
            "use .adoptiq-acceptance or a directory outside the repository"
        )
    return target


def _read_workbook(path: Path) -> dict[str, pd.DataFrame]:
    with pd.ExcelFile(path) as workbook:
        return {
            str(sheet): workbook.parse(sheet_name=sheet)
            for sheet in workbook.sheet_names
        }


def _report_info_map(frame: pd.DataFrame) -> dict[str, Any]:
    if not {"Item", "Value"}.issubset(frame.columns):
        return {}
    return {
        str(row["Item"]): row["Value"]
        for _, row in frame.iterrows()
        if str(row.get("Item") or "").strip()
    }


def _source_states(info: Mapping[str, Any]) -> dict[str, str]:
    prefix = "Source_State:"
    return {
        key[len(prefix) :]: str(value or "").strip().lower()
        for key, value in info.items()
        if key.startswith(prefix)
    }


def _workbook_style_and_formula_audit(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    formula_cells: list[str] = []
    formula_error_cells: list[str] = []
    missing_frozen_panes: list[str] = []
    missing_filters: list[str] = []
    invalid_widths: list[str] = []
    invalid_header_heights: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            if str(worksheet.freeze_panes or "") != "A2":
                missing_frozen_panes.append(worksheet.title)
            if not str(worksheet.auto_filter.ref or "").strip():
                missing_filters.append(worksheet.title)
            if (worksheet.row_dimensions[1].height or 0) < 20:
                invalid_header_heights.append(worksheet.title)
            for column_index in range(1, max(worksheet.max_column, 1) + 1):
                letter = openpyxl.utils.get_column_letter(column_index)
                width = worksheet.column_dimensions[letter].width
                if width is None or not 8 <= float(width) <= 52:
                    invalid_widths.append(f"{worksheet.title}:{letter}")
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula_cells.append(f"{worksheet.title}!{cell.coordinate}")
                    if cell.data_type == "e" or str(cell.value or "").strip() in {
                        "#REF!",
                        "#DIV/0!",
                        "#VALUE!",
                        "#NAME?",
                        "#N/A",
                    }:
                        formula_error_cells.append(
                            f"{worksheet.title}!{cell.coordinate}"
                        )
    finally:
        workbook.close()
    return {
        "formula_cells": formula_cells,
        "formula_error_cells": formula_error_cells,
        "missing_frozen_panes": missing_frozen_panes,
        "missing_filters": missing_filters,
        "invalid_widths": invalid_widths,
        "invalid_header_heights": invalid_header_heights,
    }


def _stable_record_id_audit(
    sheets: Mapping[str, pd.DataFrame],
) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for sheet_name in ACTIVITY_SHEETS:
        frame = sheets.get(sheet_name, pd.DataFrame())
        ids = frame.get("Record_ID", pd.Series(dtype="object")).fillna("").astype(str).str.strip()
        quality = frame.get(
            "Record_ID_Data_Quality",
            pd.Series("", index=frame.index, dtype="object"),
        ).fillna("").astype(str)
        missing = ids.eq("")
        duplicated = ids.loc[ids.ne("")].duplicated(keep=False)
        missing_disclosed = quality.loc[missing].eq("Missing stable source ID").all()
        result[sheet_name] = {
            "rows": int(len(frame)),
            "stable_id_rows": int(ids.ne("").sum()),
            "missing_id_rows": int(missing.sum()),
            "missing_ids_disclosed": bool(missing_disclosed),
            "duplicate_stable_id_rows": int(duplicated.sum()),
        }
    return result


def _lifecycle_audit(sheets: Mapping[str, pd.DataFrame]) -> dict[str, Any]:
    action_plans = sheets.get("Action_Plans", pd.DataFrame())
    buckets = action_plans.get(
        "AdoptIQ_Status_Bucket",
        pd.Series("Unknown", index=action_plans.index, dtype="object"),
    ).fillna("Unknown").astype(str).value_counts().to_dict()
    expected_bucket_names = (
        "Overdue",
        "Due Soon",
        "Open",
        "Blocked / On Hold",
        "Completed",
        "Unknown",
    )
    normalized_buckets = {
        name: int(buckets.get(name, 0)) for name in expected_bucket_names
    }
    chart = sheets.get("Chart_Data", pd.DataFrame())
    ap_chart = chart.loc[
        chart.get("Chart_ID", pd.Series(dtype="object"))
        .fillna("")
        .astype(str)
        .eq("action_plan_status_aging")
    ]
    numeric_chart_values = pd.to_numeric(ap_chart.get("Value"), errors="coerce")
    chart_total = int(numeric_chart_values.fillna(0).sum())
    chart_source_states = sorted(
        set(
            ap_chart.get("Source_State", pd.Series(dtype="object"))
            .fillna("unknown")
            .astype(str)
            .str.casefold()
        )
    )
    lineage = sheets.get("Metric_Lineage", pd.DataFrame())
    lineage_values = {}
    if {"Metric_Key", "Value"}.issubset(lineage.columns):
        lineage_values = {
            str(row["Metric_Key"]): row["Value"]
            for _, row in lineage.iterrows()
            if str(row.get("Metric_Key") or "").startswith("kpi.action_plans_")
        }
    total = int(len(action_plans))
    missing_title_mask = action_plans.get(
        "AdoptIQ_Data_Quality",
        pd.Series("", index=action_plans.index, dtype="object"),
    ).fillna("").astype(str).str.contains("Missing title", case=False, regex=False)
    titles = action_plans.get(
        "AdoptIQ_Title",
        pd.Series("", index=action_plans.index, dtype="object"),
    ).fillna("").astype(str)
    missing_title_fallback_ok = titles.loc[missing_title_mask].eq(
        "Title unavailable"
    ).all()
    return {
        "total": total,
        "bucket_counts": normalized_buckets,
        "partition_total": int(sum(normalized_buckets.values())),
        "chart_total": chart_total,
        "chart_value_rows": int(numeric_chart_values.notna().sum()),
        "chart_source_states": chart_source_states,
        "lineage_values": _json_safe(lineage_values),
        "missing_title_rows": int(missing_title_mask.sum()),
        "missing_title_fallback_ok": bool(missing_title_fallback_ok),
    }


def _chart_audit(
    word_path: Path,
    sheets: Mapping[str, pd.DataFrame],
    document: Document,
) -> dict[str, Any]:
    chart_data = sheets.get("Chart_Data", pd.DataFrame())
    chart_ids = set(
        chart_data.get("Chart_ID", pd.Series(dtype="object"))
        .dropna()
        .astype(str)
    )
    available_ids: list[str] = []
    for chart_id in CHART_TITLES:
        rows = chart_data.loc[
            chart_data.get("Chart_ID", pd.Series(dtype="object"))
            .fillna("")
            .astype(str)
            .eq(chart_id)
        ]
        if chart_id == "activity_trend":
            rows = rows.loc[
                rows.get("Period_Start", pd.Series(index=rows.index, dtype="object"))
                .notna()
            ]
        if pd.to_numeric(rows.get("Value"), errors="coerce").notna().any():
            available_ids.append(chart_id)
    visible = _inspect_docx_visible_charts(word_path, doc=document)
    title_set = {
        str(item.get("title") or "")
        for item in visible.get("embedded_raster_charts", [])
    }
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    explanations = {
        chart_id: (
            CHART_TITLES[chart_id] in text
            and (
                "Chart unavailable" in text
                or "Chart rendering was unavailable" in text
                or "Chart withheld" in text
            )
        )
        for chart_id in CHART_TITLES
        if chart_id not in available_ids
    }
    return {
        "required_chart_ids_present": sorted(set(CHART_TITLES) & chart_ids),
        "available_chart_ids": available_ids,
        "visible_chart_count": int(visible.get("visible_chart_count") or 0),
        "accessible_chart_titles": sorted(title_set),
        "missing_series_explanations": explanations,
        "inspection_errors": list(visible.get("inspection_errors") or []),
    }


def _tac_bems_audit(sheets: Mapping[str, pd.DataFrame]) -> dict[str, Any]:
    tac = sheets.get("TAC_Cases", pd.DataFrame())
    bems = sheets.get("BEMS", pd.DataFrame())
    tac_ids = set(
        tac.get("Record_ID", pd.Series(dtype="object"))
        .fillna("")
        .astype(str)
        .str.strip()
    ) - {""}
    bems_ids = set(
        bems.get("Record_ID", pd.Series(dtype="object"))
        .fillna("")
        .astype(str)
        .str.strip()
    ) - {""}
    association_columns = [
        column
        for column in (
            "Subscription Reference Id",
            "SUBSCRIPTION_ID",
            "ACCOUNT_ID_C",
            "Account ID",
        )
        if column in tac.columns
    ]
    stable_association_rows = pd.Series(False, index=tac.index)
    for column in association_columns:
        stable_association_rows = stable_association_rows | (
            tac[column]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.lower()
            .isin({"", "nan", "none", "null"})
            .eq(False)
        )
    unassigned = (
        tac.get("CSSM", pd.Series("", index=tac.index, dtype="object"))
        .fillna("")
        .astype(str)
        .eq("Unassigned / Portfolio")
    )
    return {
        "tac_rows": int(len(tac)),
        "bems_rows": int(len(bems)),
        "bems_is_non_additive_subset": bool(len(bems) <= len(tac) and bems_ids <= tac_ids),
        "stable_association_columns": association_columns,
        "assigned_rows_with_stable_association": int(
            stable_association_rows.loc[~unassigned].sum()
        ),
        "assigned_rows": int((~unassigned).sum()),
        "retained_unassigned_rows": int(unassigned.sum()),
    }


def validate_artifact_pair(
    *,
    word_path: Path,
    source_data_path: Path,
    expectation: ScopeExpectation,
    manager: str,
    days: int,
    requested_as_of: pd.Timestamp,
    status: Mapping[str, Any] | None = None,
    live: bool,
    as_of_skew_seconds: int = 14400,
) -> dict[str, Any]:
    """Validate one shipped pair without recalculating report metrics."""

    errors: list[str] = []
    checks: dict[str, bool] = {}
    status = status or {}
    sheets = _read_workbook(source_data_path)
    expected_sheet_names = list(delivery.SOURCE_DATA_SHEET_NAMES)
    checks["exact_canonical_sheet_inventory"] = list(sheets) == expected_sheet_names

    info = _report_info_map(sheets.get("Report_Info", pd.DataFrame()))
    checks["report_type_matches"] = str(info.get("Report_Type") or "") == expectation.report_type
    checks["manager_matches"] = str(info.get("Manager") or "") == manager
    checks["scope_type_matches"] = str(info.get("Scope_Type") or "") == expectation.scope_type
    checks["scope_value_matches"] = str(info.get("Scope_Value") or "") == expectation.scope_value
    try:
        info_days = int(float(info.get("Days")))
    except (TypeError, ValueError):
        info_days = -1
    checks["analysis_window_matches"] = info_days == int(days)

    artifact_as_of = pd.to_datetime(
        info.get("Data_As_Of_UTC"), errors="coerce", utc=True
    )
    checks["artifact_as_of_is_explicit"] = not pd.isna(artifact_as_of)
    if not pd.isna(artifact_as_of):
        skew = abs((artifact_as_of - requested_as_of).total_seconds())
        checks["as_of_matches_acceptance_clock"] = (
            skew <= as_of_skew_seconds if live else skew == 0
        )
    else:
        skew = None
        checks["as_of_matches_acceptance_clock"] = False

    status_as_of = pd.to_datetime(
        status.get("data_retrieved_at"), errors="coerce", utc=True
    )
    if live:
        checks["status_as_of_is_exposed"] = not pd.isna(status_as_of)
        checks["status_as_of_matches_artifact"] = (
            not pd.isna(status_as_of)
            and not pd.isna(artifact_as_of)
            and abs((status_as_of - artifact_as_of).total_seconds()) <= 1
        )
        checks["server_delivery_contract_passed"] = bool(
            (status.get("delivery_contract") or {}).get("ok")
        )
        checks["server_source_data_contract_passed"] = bool(
            (status.get("source_data_contract") or {}).get("ok")
        )

    expected_word_prefix = "AdoptIQ_Report_"
    expected_source_prefix = "AdoptIQ_Source_Data_"
    checks["word_filename_contract"] = word_path.name.startswith(expected_word_prefix)
    checks["source_data_filename_contract"] = source_data_path.name.startswith(
        expected_source_prefix
    )
    checks["paired_filename_contract"] = (
        word_path.stem.removeprefix(expected_word_prefix)
        == source_data_path.stem.removeprefix(expected_source_prefix)
    )

    document = Document(word_path)
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    checks["word_title_matches"] = (
        f"AdoptIQ {expectation.report_type} Decision Report" in paragraph_text
    )
    checks["word_scope_matches"] = expectation.scope_value in paragraph_text
    word_contract = delivery.validate_word_content(document)
    checks["word_content_budget"] = not word_contract["errors"]
    checks["word_excludes_raw_appendices"] = not word_contract["forbidden_headings"]

    fingerprint = str(info.get("Fact_Contract_SHA256") or "")
    checks["fact_fingerprint_present"] = bool(re.fullmatch(r"[0-9a-f]{64}", fingerprint))
    checks["word_fingerprint_matches_source_data"] = (
        str(document.core_properties.identifier or "") == fingerprint
    )
    sheet_hashes: dict[str, str] = {}
    for sheet_name in expected_sheet_names[1:]:
        frame = sheets.get(sheet_name, pd.DataFrame())
        digest = delivery._frame_content_digest(  # noqa: SLF001 - canonical artifact verification
            frame,
            sheet_name=sheet_name,
            already_exported=True,
        )
        sheet_hashes[sheet_name] = digest
        expected_digest = str(info.get(f"Sheet_SHA256:{sheet_name}") or "")
        checks[f"sheet_fingerprint:{sheet_name}"] = digest == expected_digest

    source_states = _source_states(info)
    source_counts = {name: int(len(frame)) for name, frame in sheets.items()}
    checks["all_source_states_present"] = all(
        sheet_name in source_states
        for sheet_name in (
            "Action_Plans",
            "Adoption_Barriers",
            "Customer_Pulse",
            "TAC_Cases",
            "Subscriptions",
            "Success_Priorities",
            "External_Incidents",
            "External_Bugs",
        )
    )
    state_count_consistency = True
    for sheet_name, state in source_states.items():
        row_count = source_counts.get(sheet_name, 0)
        if state == "zero" and row_count != 0:
            state_count_consistency = False
        if state in {"failed", "unavailable"}:
            chart = sheets.get("Chart_Data", pd.DataFrame())
            dependent = chart.loc[
                chart.get("Source_Sheet", pd.Series(dtype="object"))
                .fillna("")
                .astype(str)
                .str.replace(" ", "_", regex=False)
                .eq(sheet_name)
            ]
            if pd.to_numeric(dependent.get("Value"), errors="coerce").notna().any():
                state_count_consistency = False
    checks["source_state_count_semantics"] = state_count_consistency
    if live:
        checks["required_live_sources_not_degraded"] = all(
            source_states.get(sheet_name) in ACCEPTABLE_LIVE_SOURCE_STATES
            for sheet_name in REQUIRED_LIVE_SOURCE_SHEETS
        )

    record_ids = _stable_record_id_audit(sheets)
    checks["record_id_quality_is_disclosed"] = all(
        item["missing_ids_disclosed"] for item in record_ids.values()
    )
    checks["stable_ids_are_deduplicated"] = all(
        item["duplicate_stable_id_rows"] == 0 for item in record_ids.values()
    )

    lifecycle = _lifecycle_audit(sheets)
    checks["action_plan_lifecycle_is_partition"] = (
        lifecycle["partition_total"] == lifecycle["total"]
    )
    lifecycle_chart_complete = bool(lifecycle["chart_source_states"]) and set(
        lifecycle["chart_source_states"]
    ).issubset({"available", "zero"})
    checks["action_plan_chart_reconciles"] = (
        lifecycle["chart_total"] == lifecycle["total"]
        if lifecycle_chart_complete
        else lifecycle["chart_value_rows"] == 0
    )
    checks["incomplete_action_plan_chart_fails_closed"] = (
        True
        if lifecycle_chart_complete
        else lifecycle["chart_value_rows"] == 0
        and bool(lifecycle["chart_source_states"])
        and not set(lifecycle["chart_source_states"]).intersection(
            {"available", "zero"}
        )
    )
    checks["missing_title_fallback_is_honest"] = lifecycle[
        "missing_title_fallback_ok"
    ]

    charts = _chart_audit(word_path, sheets, document)
    checks["four_chart_contracts_present"] = set(
        charts["required_chart_ids_present"]
    ) == set(CHART_TITLES)
    checks["available_charts_are_embedded"] = (
        charts["visible_chart_count"] >= len(charts["available_chart_ids"])
    )
    checks["chart_alt_text_is_present"] = all(
        CHART_TITLES[chart_id] in charts["accessible_chart_titles"]
        for chart_id in charts["available_chart_ids"]
    )
    checks["unavailable_chart_series_are_explained"] = all(
        charts["missing_series_explanations"].values()
    )

    lineage = sheets.get("Metric_Lineage", pd.DataFrame())
    chart_data = sheets.get("Chart_Data", pd.DataFrame())
    lineage_keys = set(
        lineage.get("Metric_Key", pd.Series(dtype="object"))
        .dropna()
        .astype(str)
    )
    chart_keys = set(
        chart_data.get("Metric_Key", pd.Series(dtype="object"))
        .dropna()
        .astype(str)
    ) - {""}
    checks["chart_lineage_reconciles"] = chart_keys <= lineage_keys

    tac_bems = _tac_bems_audit(sheets)
    checks["bems_is_non_additive_tac_subset"] = tac_bems[
        "bems_is_non_additive_subset"
    ]
    checks["tac_uses_stable_association_or_quarantine"] = (
        tac_bems["assigned_rows_with_stable_association"]
        == tac_bems["assigned_rows"]
    )

    workbook_audit = _workbook_style_and_formula_audit(source_data_path)
    checks["workbook_is_formula_free"] = not workbook_audit["formula_cells"]
    checks["workbook_has_no_formula_errors"] = not workbook_audit[
        "formula_error_cells"
    ]
    checks["workbook_has_frozen_headers"] = not workbook_audit[
        "missing_frozen_panes"
    ]
    checks["workbook_has_filters"] = not workbook_audit["missing_filters"]
    checks["workbook_has_readable_dimensions"] = not (
        workbook_audit["invalid_widths"]
        or workbook_audit["invalid_header_heights"]
    )

    for name, passed in checks.items():
        if not passed:
            errors.append(name)
    return {
        "ok": not errors,
        "errors": errors,
        "checks": checks,
        "artifacts": {
            "word_path": str(word_path),
            "source_data_path": str(source_data_path),
            "word_sha256": _sha256(word_path),
            "source_data_sha256": _sha256(source_data_path),
        },
        "report_metadata": {
            "report_type": info.get("Report_Type"),
            "manager": info.get("Manager"),
            "scope_type": info.get("Scope_Type"),
            "scope_value": info.get("Scope_Value"),
            "days": info_days,
            "requested_as_of_utc": requested_as_of.isoformat(),
            "artifact_as_of_utc": (
                artifact_as_of.isoformat() if not pd.isna(artifact_as_of) else None
            ),
            "status_as_of_utc": (
                status_as_of.isoformat() if not pd.isna(status_as_of) else None
            ),
            "as_of_skew_seconds": skew,
            "fact_contract_sha256": fingerprint,
            "sheet_sha256": sheet_hashes,
        },
        "source_states": source_states,
        "source_counts": source_counts,
        "record_id_quality": record_ids,
        "action_plan_lifecycle": lifecycle,
        "charts": charts,
        "tac_bems": tac_bems,
        "word": word_contract,
        "workbook": workbook_audit,
    }


def run_scope_authorization_probes() -> dict[str, Any]:
    """Exercise the pure fail-closed authorization SSoT every run."""

    roster = [
        ("Acceptance Manager", "Inside Member", "inside@example.test"),
        ("Other Manager", "Outside Member", "outside@example.test"),
    ]
    outside_rejected = False
    ambiguous_rejected = False
    try:
        validate_leader_scope_request(
            "Acceptance Manager",
            "member",
            "outside@example.test",
            roster,
        )
    except LeaderScopeValidationError:
        outside_rejected = True

    ambiguous = pd.DataFrame(
        [
            {
                "CSSM_EMAIL": "inside@example.test",
                "BU_NAME": "Customer Alpha",
                "ACCOUNT_ID_C": "SHARED-1",
            },
            {
                "CSSM_EMAIL": "inside@example.test",
                "BU_NAME": "Customer Beta",
                "ACCOUNT_ID_C": "SHARED-1",
            },
        ]
    )
    selection = validate_leader_scope_request(
        "Acceptance Manager",
        "customer",
        "Customer Alpha",
        roster,
    )
    try:
        filter_leader_subscriptions(ambiguous, selection)
    except LeaderScopeValidationError:
        ambiguous_rejected = True
    return {
        "outside_manager_member_rejected": outside_rejected,
        "ambiguous_shared_account_customer_rejected": ambiguous_rejected,
        "ok": outside_rejected and ambiguous_rejected,
    }


def probe_live_sources(base_url: str, timeout: int = 20) -> dict[str, Any]:
    parsed = urlparse(base_url)
    if (parsed.hostname or "").lower() not in {"127.0.0.1", "localhost", "::1"}:
        return {
            "ok": False,
            "error": "base URL must target the local AdoptIQ app",
        }
    session = requests.Session()
    app_probe: dict[str, Any]
    snowflake_probe: dict[str, Any]
    try:
        response = session.get(f"{base_url.rstrip('/')}/ping", timeout=timeout)
        app_probe = {
            "ok": response.status_code == 200 and response.text.strip() == "OK",
            "status_code": response.status_code,
        }
    except requests.RequestException as exc:
        app_probe = {
            "ok": False,
            "error_kind": type(exc).__name__,
            "error": str(exc)[:240],
        }
    try:
        response = session.get(
            f"{base_url.rstrip('/')}/api/diag/connectivity",
            timeout=timeout,
        )
        try:
            payload = response.json()
        except ValueError:
            payload = {}
        snowflake_probe = {
            "ok": response.status_code == 200 and bool(payload.get("ok")),
            "status_code": response.status_code,
            "stage": payload.get("stage"),
            "error_kind": payload.get("error_kind"),
        }
    except requests.RequestException as exc:
        snowflake_probe = {
            "ok": False,
            "error_kind": type(exc).__name__,
            "error": str(exc)[:240],
        }
    return {
        "ok": bool(app_probe.get("ok") and snowflake_probe.get("ok")),
        "app": app_probe,
        "snowflake": snowflake_probe,
        "csconsole": {
            "state": "validated_by_report_contract",
            "detail": "No independent side-effect-free CSConsole route; canonical source states are checked in each artifact.",
        },
        "csone": {
            "state": "validated_by_report_contract",
            "detail": "Upload/autodiscovery and canonical TAC_Cases source state are checked in each artifact.",
        },
    }


class LiveDecisionReportClient:
    """Small HTTP client for the production report routes."""

    def __init__(
        self,
        *,
        base_url: str,
        output_dir: Path,
        poll_interval: float,
        scenario_timeout: int,
        request_timeout: int,
        download_timeout: int,
        csone_file: Path | None,
    ) -> None:
        self.base_url = base_url.rstrip("/")
        self.base_output_dir = output_dir
        self.output_dir = output_dir
        self.poll_interval = poll_interval
        self.scenario_timeout = scenario_timeout
        self.request_timeout = request_timeout
        self.download_timeout = download_timeout
        self.csone_file = csone_file
        self.session = requests.Session()
        self.csrf_token = ""
        self.session_cookie = ""

    def bootstrap(self) -> None:
        response = self.session.get(
            f"{self.base_url}/",
            timeout=self.request_timeout,
        )
        response.raise_for_status()
        self.csrf_token = extract_csrf_token(response.text)
        self.session_cookie = self.session.cookies.get("session", "") or ""

    def _headers(self) -> dict[str, str]:
        headers = {
            "X-Requested-With": "XMLHttpRequest",
            "X-CSRFToken": self.csrf_token,
        }
        if self.session_cookie:
            headers["Cookie"] = f"session={self.session_cookie}"
        return headers

    def scope_options(
        self,
        manager: str,
        *,
        member_email: str = "",
        include_customers: bool = False,
    ) -> dict[str, Any]:
        response = self.session.get(
            f"{self.base_url}/api/leader_scope_options",
            params={
                "manager": manager,
                "member_email": member_email,
                "include_customers": "1" if include_customers else "0",
            },
            timeout=self.request_timeout,
        )
        try:
            payload = response.json()
        except ValueError as exc:
            raise RuntimeError("scope-options route returned invalid JSON") from exc
        if response.status_code != 200 or not payload.get("success"):
            raise RuntimeError(f"scope-options request failed: HTTP {response.status_code}")
        return payload

    def negative_scope_probe(
        self,
        *,
        manager: str,
        ambiguous_customer_name: str = "",
    ) -> dict[str, Any]:
        invalid_member = self._post_start(
            "/start_leader_report",
            {
                "manager": manager,
                "days": "90",
                "scope_type": "member",
                "scope_value": "outside-manager@example.invalid",
            },
            include_csone=False,
            allow_error=True,
        )
        outside_rejected = (
            invalid_member["status_code"] == 400
            and not invalid_member["payload"].get("success")
        )
        ambiguous_result: dict[str, Any] = {
            "attempted": False,
            "rejected": None,
        }
        if ambiguous_customer_name:
            response = self._post_start(
                "/start_leader_report",
                {
                    "manager": manager,
                    "days": "90",
                    "scope_type": "customer",
                    "scope_value": ambiguous_customer_name,
                },
                include_csone=False,
                allow_error=True,
            )
            ambiguous_result = {
                "attempted": True,
                "rejected": (
                    response["status_code"] == 400
                    and "shares an account" in str(response["payload"].get("error") or "")
                ),
            }
        return {
            "outside_manager_member_rejected": outside_rejected,
            "ambiguous_customer": ambiguous_result,
            "ok": outside_rejected and (
                not ambiguous_result["attempted"]
                or bool(ambiguous_result["rejected"])
            ),
        }

    def _post_start(
        self,
        endpoint: str,
        payload: Mapping[str, Any],
        *,
        include_csone: bool = True,
        allow_error: bool = False,
    ) -> dict[str, Any]:
        files = None
        handle = None
        try:
            if include_csone and self.csone_file is not None:
                handle = self.csone_file.open("rb")
                files = {
                    "csone_file": (
                        self.csone_file.name,
                        handle,
                        mimetypes.guess_type(self.csone_file.name)[0]
                        or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                }
            response = self.session.post(
                f"{self.base_url}{endpoint}",
                data=dict(payload),
                files=files,
                headers=self._headers(),
                timeout=self.request_timeout,
            )
        finally:
            if handle is not None:
                handle.close()
        try:
            response_payload = response.json()
        except ValueError:
            response_payload = {"error": "invalid_json_response"}
        result = {
            "status_code": response.status_code,
            "payload": response_payload,
        }
        if not allow_error and (
            response.status_code >= 400 or not response_payload.get("success")
        ):
            error = str(response_payload.get("error") or "report start failed")
            raise RuntimeError(
                f"{endpoint} start failed: HTTP {response.status_code}: {error[:240]}"
            )
        return result

    def run(
        self,
        *,
        endpoint: str,
        payload: Mapping[str, Any],
        scope_key: str,
    ) -> tuple[dict[str, Any], Path, Path]:
        started = self._post_start(endpoint, payload)
        analysis_id = str(started["payload"].get("analysis_id") or "")
        if not analysis_id:
            raise RuntimeError("report start response omitted analysis_id")
        deadline = time.monotonic() + self.scenario_timeout
        while True:
            response = self.session.get(
                f"{self.base_url}/status/{quote(analysis_id)}",
                timeout=self.request_timeout,
            )
            response.raise_for_status()
            status = response.json()
            if status.get("status") in TERMINAL_STATUSES:
                break
            if time.monotonic() >= deadline:
                raise TimeoutError(
                    f"timed out waiting for {scope_key} after {self.scenario_timeout}s"
                )
            time.sleep(self.poll_interval)
        if status.get("status") != "completed" or status.get("error"):
            raise RuntimeError(
                f"{scope_key} report failed: {status.get('error') or status.get('message')}"
            )
        if not status.get("word_available") or not status.get("excel_available"):
            raise RuntimeError(f"{scope_key} did not publish both required artifacts")

        artifact_paths = []
        for file_type, prefix in (
            ("docx", "AdoptIQ_Report_"),
            ("xlsx", "AdoptIQ_Source_Data_"),
        ):
            response = self.session.get(
                f"{self.base_url}/download/{quote(analysis_id)}/{file_type}",
                timeout=self.download_timeout,
            )
            response.raise_for_status()
            raw_name = parse_download_name(
                response.headers.get("Content-Disposition", ""),
                f"{prefix}{analysis_id}.{file_type}",
            )
            safe_name = Path(raw_name).name
            if not safe_name.startswith(prefix):
                safe_name = f"{prefix}{analysis_id}.{file_type}"
            path = self.output_dir / safe_name
            if path.exists():
                raise RuntimeError(f"refusing to overwrite acceptance artifact: {path.name}")
            path.write_bytes(response.content)
            artifact_paths.append(path)
        return status, artifact_paths[0], artifact_paths[1]


def _extract_query_identifiers(value: Any) -> list[str]:
    found: set[str] = set()

    def visit(item: Any, key: str = "") -> None:
        if isinstance(item, Mapping):
            for child_key, child_value in item.items():
                visit(child_value, str(child_key))
        elif isinstance(item, (list, tuple)):
            for child in item:
                visit(child, key)
        elif "query_id" in key.casefold() and item not in (None, ""):
            found.add(str(item))

    visit(value)
    return sorted(found)


def _offline_expectation(result: Mapping[str, Any]) -> ScopeExpectation:
    return ScopeExpectation(
        requested_scope=str(result["scope"]),
        report_type=str(result["report_type"]),
        scope_type=str(result["scope_type"]),
        scope_value=str(result["scope_value"]),
    )


def run_offline_pass(
    *,
    pass_number: int,
    output_dir: Path,
    manager: str,
    days: int,
    as_of: pd.Timestamp,
) -> dict[str, Any]:
    pass_dir = output_dir / f"pass-{pass_number}"
    scope_results: dict[str, Any] = {}
    for scope in SUPPORTED_SCOPES:
        scope_dir = pass_dir / scope
        try:
            generated = generate_acceptance_artifacts(
                scope=scope,
                as_of=as_of.isoformat(),
                output_dir=scope_dir,
                manager_name=manager,
                days=days,
            )
            validation = validate_artifact_pair(
                word_path=Path(generated["word_path"]),
                source_data_path=Path(generated["source_data_path"]),
                expectation=_offline_expectation(generated),
                manager=manager,
                days=days,
                requested_as_of=as_of,
                live=False,
            )
            validation["manifests"] = {
                key: generated[key]
                for key in (
                    "measurement_manifest_path",
                    "parity_manifest_path",
                    "chart_manifest_path",
                )
            }
            scope_results[scope] = validation
        except Exception as exc:  # noqa: BLE001 - summary must retain all failures
            scope_results[scope] = {
                "ok": False,
                "errors": [f"{type(exc).__name__}: {exc}"],
            }
    return {
        "pass_number": pass_number,
        "mode": "offline",
        "scopes": scope_results,
        "ok": all(result.get("ok") for result in scope_results.values()),
    }


def _choose_live_scopes(
    client: LiveDecisionReportClient,
    *,
    manager: str,
    member_email: str,
    customer_name: str,
    customer_member_email: str,
) -> dict[str, ScopeExpectation]:
    roster = client.scope_options(manager)
    members = list(roster.get("members") or [])
    if not members:
        raise RuntimeError("selected manager has no live Leader roster members")
    if member_email:
        member = next(
            (
                item
                for item in members
                if str(item.get("email") or "").casefold()
                == member_email.casefold()
            ),
            None,
        )
        if member is None:
            raise RuntimeError("--member-email is not in the selected manager roster")
    else:
        member = members[0]
    selected_member_email = str(member.get("email") or "")
    selected_member_name = str(member.get("name") or selected_member_email)

    customer_member = customer_member_email.strip()
    if customer_member:
        if not any(
            str(item.get("email") or "").casefold() == customer_member.casefold()
            for item in members
        ):
            raise RuntimeError(
                "--customer-member-email is not in the selected manager roster"
            )
    customers_payload = client.scope_options(
        manager,
        member_email=customer_member,
        include_customers=True,
    )
    customers = list(customers_payload.get("customers") or [])
    if not customers_payload.get("customers_available"):
        raise RuntimeError(
            "live customer options are unavailable; Snowflake scope authorization cannot be proven"
        )
    if customer_name:
        customer = next(
            (
                item
                for item in customers
                if str(item.get("value") or "").casefold()
                == customer_name.casefold()
            ),
            None,
        )
        if customer is None:
            raise RuntimeError(
                "--customer-name is not an unambiguous authorized live customer option"
            )
    elif customers:
        customer = customers[0]
    else:
        raise RuntimeError("selected live manager/member has no safe customer scope")
    selected_customer = str(customer.get("value") or "")
    customer_display = selected_customer
    if customer_member:
        customer_member_name = next(
            str(item.get("name") or customer_member)
            for item in members
            if str(item.get("email") or "").casefold()
            == customer_member.casefold()
        )
        customer_display = f"{selected_customer} ({customer_member_name})"
    return {
        "team": ScopeExpectation("team", "Leader", "team", "Entire team"),
        "member": ScopeExpectation(
            "member",
            "Leader",
            "member",
            f"{selected_member_name} ({selected_member_email.lower()})",
        ),
        "customer": ScopeExpectation(
            "customer",
            "Leader",
            "customer",
            customer_display,
        ),
        "comprehensive": ScopeExpectation(
            "comprehensive",
            "Comprehensive",
            "team",
            f"{manager} team",
        ),
    }


def run_live_pass(
    *,
    pass_number: int,
    client: LiveDecisionReportClient,
    expectations: Mapping[str, ScopeExpectation],
    manager: str,
    days: int,
    as_of: pd.Timestamp,
    member_email: str,
    customer_name: str,
    customer_member_email: str,
    as_of_skew_seconds: int,
) -> dict[str, Any]:
    pass_dir = client.base_output_dir / f"pass-{pass_number}"
    pass_dir.mkdir(parents=True, exist_ok=True)
    client.output_dir = pass_dir
    scenarios = {
        "team": (
            "/start_leader_report",
            {
                "manager": manager,
                "days": str(days),
                "scope_type": "team",
                "scope_value": "",
                "scope_member": "",
            },
        ),
        "member": (
            "/start_leader_report",
            {
                "manager": manager,
                "days": str(days),
                "scope_type": "member",
                "scope_value": member_email,
                "scope_member": member_email,
            },
        ),
        "customer": (
            "/start_leader_report",
            {
                "manager": manager,
                "days": str(days),
                "scope_type": "customer",
                "scope_value": customer_name,
                "scope_member": customer_member_email,
            },
        ),
        "comprehensive": (
            "/start_analysis",
            {
                "report_type": "comprehensive",
                "manager": manager,
                "technology": "All Contact Center",
                "days": str(days),
                "subscription_id": "",
                "customer_name": "",
            },
        ),
    }
    results: dict[str, Any] = {}
    for scope, (endpoint, payload) in scenarios.items():
        try:
            status, word_path, workbook_path = client.run(
                endpoint=endpoint,
                payload=payload,
                scope_key=scope,
            )
            validation = validate_artifact_pair(
                word_path=word_path,
                source_data_path=workbook_path,
                expectation=expectations[scope],
                manager=manager,
                days=days,
                requested_as_of=as_of,
                status=status,
                live=True,
                as_of_skew_seconds=as_of_skew_seconds,
            )
            validation["analysis_id"] = status.get("analysis_id")
            validation["query_identifiers"] = _extract_query_identifiers(status)
            results[scope] = validation
        except Exception as exc:  # noqa: BLE001 - keep running all four scopes
            results[scope] = {
                "ok": False,
                "errors": [f"{type(exc).__name__}: {exc}"],
            }
    return {
        "pass_number": pass_number,
        "mode": "live",
        "scopes": results,
        "ok": all(result.get("ok") for result in results.values()),
    }


def compare_passes(passes: Iterable[Mapping[str, Any]], *, live: bool) -> dict[str, Any]:
    pass_list = list(passes)
    results: dict[str, Any] = {}
    if len(pass_list) < 2:
        return {"ok": False, "errors": ["at least two acceptance passes are required"]}
    for scope in SUPPORTED_SCOPES:
        first = (pass_list[0].get("scopes") or {}).get(scope) or {}
        second = (pass_list[1].get("scopes") or {}).get(scope) or {}
        if not first.get("ok") or not second.get("ok"):
            results[scope] = {
                "ok": False,
                "errors": ["one or both scope passes failed"],
            }
            continue
        first_meta = first.get("report_metadata") or {}
        second_meta = second.get("report_metadata") or {}
        sheet_hashes_identical = (
            first_meta.get("sheet_sha256") == second_meta.get("sheet_sha256")
        )
        semantic_metrics_identical = all(
            first.get(key) == second.get(key)
            for key in (
                "source_states",
                "source_counts",
                "record_id_quality",
                "action_plan_lifecycle",
                "tac_bems",
            )
        )
        byte_identical = all(
            (first.get("artifacts") or {}).get(key)
            == (second.get("artifacts") or {}).get(key)
            for key in ("word_sha256", "source_data_sha256")
        )
        # Live facts intentionally carry the real per-query prefetch clock, so
        # byte/fact hashes can differ even with stable source records.  The
        # canonical per-sheet hashes and normalized metrics are the binding
        # no-source-drift comparison in live mode.
        passed = sheet_hashes_identical and semantic_metrics_identical
        if not live:
            passed = passed and byte_identical and (
                first_meta.get("fact_contract_sha256")
                == second_meta.get("fact_contract_sha256")
            )
        results[scope] = {
            "ok": passed,
            "byte_identical": byte_identical,
            "fact_fingerprint_identical": (
                first_meta.get("fact_contract_sha256")
                == second_meta.get("fact_contract_sha256")
            ),
            "sheet_hashes_identical": sheet_hashes_identical,
            "semantic_metrics_identical": semantic_metrics_identical,
            "live_clock_difference_expected": live,
        }
    return {
        "ok": all(result.get("ok") for result in results.values()),
        "scopes": results,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Generate and validate Team, Member, Customer, and Comprehensive "
            "AdoptIQ decision reports twice."
        )
    )
    parser.add_argument("--mode", choices=("auto", "offline", "live"), default="auto")
    parser.add_argument("--manager", required=True, help="Recorded report manager")
    parser.add_argument("--days", required=True, type=int, help="Analysis window, 1-365")
    parser.add_argument("--as-of", required=True, help="Explicit ISO acceptance clock")
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--base-url", default="http://127.0.0.1:5151")
    parser.add_argument("--member-email", default="")
    parser.add_argument("--customer-name", default="")
    parser.add_argument("--customer-member-email", default="")
    parser.add_argument(
        "--ambiguous-customer-name",
        default="",
        help="Optional known ambiguous live customer used for a fail-closed probe",
    )
    parser.add_argument("--csone-file", type=Path)
    parser.add_argument("--poll-interval", type=float, default=5.0)
    parser.add_argument("--timeout", type=int, default=2400)
    parser.add_argument("--request-timeout", type=int, default=120)
    parser.add_argument("--download-timeout", type=int, default=300)
    parser.add_argument(
        "--as-of-skew-seconds",
        type=int,
        default=14400,
        help=(
            "Maximum seconds between the explicit acceptance-run clock and a "
            "live report's actual prefetch clock (default 4 hours for two full passes)"
        ),
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    manager = str(args.manager or "").strip()
    if not manager:
        parser.error("--manager cannot be blank")
    if not 1 <= int(args.days) <= 365:
        parser.error("--days must be between 1 and 365")
    try:
        as_of = _parse_as_of(args.as_of)
        output_dir = _ensure_safe_output_dir(args.output_dir)
    except ValueError as exc:
        parser.error(str(exc))
    csone_file = args.csone_file.expanduser().resolve() if args.csone_file else None
    if csone_file is not None and (
        not csone_file.is_file()
        or csone_file.stat().st_size <= 0
        or csone_file.suffix.casefold() != ".xlsx"
    ):
        parser.error("--csone-file must be a non-empty .xlsx file")
    output_dir.mkdir(parents=True, exist_ok=True)
    summary_path = output_dir / "decision_report_acceptance_summary.json"
    connectivity = probe_live_sources(args.base_url)
    mode_executed = args.mode
    limitations: list[str] = []
    if args.mode == "auto" and not connectivity.get("ok"):
        mode_executed = "offline"
        limitations.append(
            "Live app/Snowflake connectivity was unavailable; auto mode ran only sanitized offline fixtures."
        )
    elif args.mode == "auto":
        mode_executed = "live"

    summary: dict[str, Any] = {
        "schema_version": SUMMARY_SCHEMA,
        "started_at_utc": _utc_now(),
        "mode_requested": args.mode,
        "mode_executed": mode_executed,
        "manager": manager,
        "days": int(args.days),
        "requested_as_of_utc": as_of.isoformat(),
        "base_url": args.base_url if mode_executed == "live" else None,
        "output_dir": str(output_dir),
        "connectivity_preflight": connectivity,
        "live_validation_performed": mode_executed == "live",
        "local_scope_authorization_probes": run_scope_authorization_probes(),
        "limitations": limitations,
        "passes": [],
        "repeatability": {"ok": False, "errors": ["not run"]},
        "failures_requiring_review": [],
        "all_passed": False,
    }

    if args.mode == "live" and not connectivity.get("ok"):
        summary["failures_requiring_review"].append(
            "Explicit live mode requires a running local app plus successful Snowflake connectivity preflight."
        )
        summary["completed_at_utc"] = _utc_now()
        _write_json(summary_path, summary)
        print(json.dumps({"all_passed": False, "summary_path": str(summary_path)}))
        return 5

    if mode_executed == "offline":
        for pass_number in (1, 2):
            result = run_offline_pass(
                pass_number=pass_number,
                output_dir=output_dir,
                manager=manager,
                days=int(args.days),
                as_of=as_of,
            )
            summary["passes"].append(result)
        summary["repeatability"] = compare_passes(summary["passes"], live=False)
    else:
        client = LiveDecisionReportClient(
            base_url=args.base_url,
            output_dir=output_dir,
            poll_interval=max(float(args.poll_interval), 0.25),
            scenario_timeout=max(int(args.timeout), 60),
            request_timeout=max(int(args.request_timeout), 10),
            download_timeout=max(int(args.download_timeout), 30),
            csone_file=csone_file,
        )
        try:
            client.bootstrap()
            expectations = _choose_live_scopes(
                client,
                manager=manager,
                member_email=args.member_email,
                customer_name=args.customer_name,
                customer_member_email=args.customer_member_email,
            )
            live_member_email = expectations["member"].scope_value.rsplit("(", 1)[-1].rstrip(")")
            live_customer_name = expectations["customer"].scope_value
            if args.customer_member_email and live_customer_name.endswith(")"):
                live_customer_name = live_customer_name.rsplit(" (", 1)[0]
            summary["selected_live_scopes"] = {
                key: {
                    "report_type": value.report_type,
                    "scope_type": value.scope_type,
                    "scope_value": value.scope_value,
                }
                for key, value in expectations.items()
            }
            summary["live_scope_authorization_probes"] = client.negative_scope_probe(
                manager=manager,
                ambiguous_customer_name=args.ambiguous_customer_name,
            )
            for pass_number in (1, 2):
                result = run_live_pass(
                    pass_number=pass_number,
                    client=client,
                    expectations=expectations,
                    manager=manager,
                    days=int(args.days),
                    as_of=as_of,
                    member_email=live_member_email,
                    customer_name=live_customer_name,
                    customer_member_email=args.customer_member_email,
                    as_of_skew_seconds=max(int(args.as_of_skew_seconds), 1),
                )
                summary["passes"].append(result)
            summary["repeatability"] = compare_passes(
                summary["passes"], live=True
            )
        except Exception as exc:  # noqa: BLE001 - always emit a useful summary
            summary["failures_requiring_review"].append(
                f"{type(exc).__name__}: {exc}"
            )

    for pass_result in summary["passes"]:
        for scope, result in (pass_result.get("scopes") or {}).items():
            for error in result.get("errors") or []:
                summary["failures_requiring_review"].append(
                    f"pass {pass_result.get('pass_number')} {scope}: {error}"
                )
    summary["all_passed"] = bool(
        len(summary["passes"]) == 2
        and all(item.get("ok") for item in summary["passes"])
        and summary["repeatability"].get("ok")
        and summary["local_scope_authorization_probes"].get("ok")
        and not summary["failures_requiring_review"]
        and (
            mode_executed != "live"
            or (summary.get("live_scope_authorization_probes") or {}).get("ok")
        )
    )
    summary["completed_at_utc"] = _utc_now()
    _write_json(summary_path, summary)
    print(
        json.dumps(
            {
                "all_passed": summary["all_passed"],
                "mode_executed": mode_executed,
                "summary_path": str(summary_path),
            },
            sort_keys=True,
        )
    )
    return 0 if summary["all_passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
