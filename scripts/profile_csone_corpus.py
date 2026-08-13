#!/usr/bin/env python3
"""Profile a local CSOne workbook corpus without exporting source rows.

The profiler is intentionally metadata-only.  It records schemas, missingness,
placeholder rates, identifier quality, workbook chronology, and cell-shape
statistics.  Customer names, case text, email addresses, record identifiers,
and raw cell values are never written to the summary or stdout.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import statistics
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import openpyxl


REPO_ROOT = Path(__file__).resolve().parents[1]
SUMMARY_SCHEMA = "csone-corpus-profile/v1"
PLACEHOLDER_TOKENS = frozenset(
    {
        "--",
        "n/a",
        "na",
        "none",
        "not available",
        "null",
        "undefined",
        "unknown",
        "unavailable",
    }
)
SENSITIVE_HEADER_MARKERS = (
    "customer",
    "contact",
    "email",
    "owner",
    "description",
    "details",
    "summary",
    "activity",
    "action plan",
    "title",
)
IDENTIFIER_COLUMNS = (
    "SR Number",
    "Case Number",
    "Transaction ID",
    "Subscription Reference Id",
)
FILENAME_DATE_RE = re.compile(r"-(\d{4}-\d{2}-\d{2})-\d{2}-\d{2}-\d{2}\.xlsx$", re.I)


@dataclass(frozen=True)
class WorkbookProfile:
    digest: str
    observed_date: str
    sheet_count: int
    selected_sheet_index: int
    header_row_one_based: int
    row_count: int
    column_count: int
    schema_sha256: str
    headers: tuple[str, ...]
    nonblank: Mapping[str, int]
    placeholders: Mapping[str, int]
    max_lengths: Mapping[str, int]
    type_counts: Mapping[str, Mapping[str, int]]
    identifier_counts: Mapping[str, Mapping[str, int]]
    hyperlink_count: int
    formula_count: int


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _normalized_token(value: object) -> str:
    return _clean(value).casefold()


def _value_kind(value: object) -> str:
    if value is None or _clean(value) == "":
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)):
        return "number"
    return "text"


def _header_candidate(sheet: Any) -> tuple[int, list[str]] | None:
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        populated = sum(1 for value in row if _clean(value))
        if populated < 3:
            if row_number >= 40:
                return None
            continue
        headers = [
            _clean(value) if _clean(value) else f"col_{column_number}"
            for column_number, value in enumerate(row)
        ]
        return row_number, headers
    return None


def _workbook_date(path: Path) -> str:
    match = FILENAME_DATE_RE.search(path.name)
    return match.group(1) if match else ""


def _safe_header_name(header: str, column_index: int, seen: Counter[str]) -> str:
    base = _clean(header) or f"col_{column_index}"
    seen[base] += 1
    return base if seen[base] == 1 else f"{base} [{seen[base]}]"


def profile_workbook(path: Path) -> WorkbookProfile:
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        candidates: list[tuple[int, int, list[str], Any]] = []
        for sheet_index, sheet in enumerate(workbook.worksheets):
            candidate = _header_candidate(sheet)
            if candidate is None:
                continue
            header_row, headers = candidate
            body_rows = max(int(sheet.max_row or 0) - header_row, 0)
            candidates.append((body_rows, sheet_index, headers, sheet))
        if not candidates:
            raise ValueError("workbook has no parseable CSOne worksheet")
        _, selected_index, raw_headers, sheet = max(
            candidates,
            key=lambda item: (item[0], -item[1]),
        )
        header_row, _ = _header_candidate(sheet) or (0, [])
        seen_headers: Counter[str] = Counter()
        headers = [
            _safe_header_name(header, index, seen_headers)
            for index, header in enumerate(raw_headers)
        ]
        nonblank: Counter[str] = Counter()
        placeholders: Counter[str] = Counter()
        max_lengths: Counter[str] = Counter()
        type_counts: dict[str, Counter[str]] = defaultdict(Counter)
        identifiers: dict[str, dict[str, Any]] = {
            column: {"populated": 0, "blank": 0, "duplicate": 0, "seen": set()}
            for column in IDENTIFIER_COLUMNS
            if column in headers
        }
        row_count = 0
        hyperlink_count = 0
        formula_count = 0
        for row in sheet.iter_rows(min_row=header_row + 1):
            values = [cell.value for cell in row[: len(headers)]]
            if not any(_clean(value) for value in values):
                continue
            row_count += 1
            for column_index, header in enumerate(headers):
                value = values[column_index] if column_index < len(values) else None
                token = _clean(value)
                kind = _value_kind(value)
                type_counts[header][kind] += 1
                if token:
                    nonblank[header] += 1
                    max_lengths[header] = max(max_lengths[header], len(token))
                    if _normalized_token(value) in PLACEHOLDER_TOKENS:
                        placeholders[header] += 1
                cell = row[column_index] if column_index < len(row) else None
                if cell is not None:
                    if getattr(cell, "hyperlink", None):
                        hyperlink_count += 1
                    if getattr(cell, "data_type", "") == "f":
                        formula_count += 1
                identifier = identifiers.get(header)
                if identifier is not None:
                    if not token:
                        identifier["blank"] += 1
                    else:
                        identifier["populated"] += 1
                        token_digest = hashlib.sha256(token.encode("utf-8")).digest()
                        if token_digest in identifier["seen"]:
                            identifier["duplicate"] += 1
                        else:
                            identifier["seen"].add(token_digest)

        identifier_counts = {
            column: {
                "populated": int(values["populated"]),
                "blank": int(values["blank"]),
                "duplicate": int(values["duplicate"]),
                "distinct": len(values["seen"]),
            }
            for column, values in identifiers.items()
        }
        schema_payload = json.dumps(headers, separators=(",", ":"), ensure_ascii=False)
        return WorkbookProfile(
            digest=hashlib.sha256(path.read_bytes()).hexdigest(),
            observed_date=_workbook_date(path),
            sheet_count=len(workbook.worksheets),
            selected_sheet_index=selected_index,
            header_row_one_based=header_row,
            row_count=row_count,
            column_count=len(headers),
            schema_sha256=hashlib.sha256(schema_payload.encode("utf-8")).hexdigest(),
            headers=tuple(headers),
            nonblank=dict(nonblank),
            placeholders=dict(placeholders),
            max_lengths=dict(max_lengths),
            type_counts={key: dict(value) for key, value in type_counts.items()},
            identifier_counts=identifier_counts,
            hyperlink_count=hyperlink_count,
            formula_count=formula_count,
        )
    finally:
        workbook.close()


def _percentile(values: Sequence[int], fraction: float) -> int:
    if not values:
        return 0
    ordered = sorted(int(value) for value in values)
    index = round((len(ordered) - 1) * min(max(fraction, 0.0), 1.0))
    return ordered[index]


def summarize_profiles(profiles: Sequence[WorkbookProfile]) -> dict[str, Any]:
    if not profiles:
        raise ValueError("no CSOne workbooks were profiled")
    schema_counts = Counter(profile.schema_sha256 for profile in profiles)
    header_rows = Counter(profile.header_row_one_based for profile in profiles)
    column_files: Counter[str] = Counter()
    column_nonblank: Counter[str] = Counter()
    column_placeholders: Counter[str] = Counter()
    column_max_lengths: Counter[str] = Counter()
    column_type_counts: dict[str, Counter[str]] = defaultdict(Counter)
    identifier_totals: dict[str, Counter[str]] = defaultdict(Counter)
    for profile in profiles:
        column_files.update(profile.headers)
        column_nonblank.update(profile.nonblank)
        column_placeholders.update(profile.placeholders)
        for column, length in profile.max_lengths.items():
            column_max_lengths[column] = max(column_max_lengths[column], length)
        for column, counts in profile.type_counts.items():
            column_type_counts[column].update(counts)
        for column, counts in profile.identifier_counts.items():
            identifier_totals[column].update(counts)

    total_rows = sum(profile.row_count for profile in profiles)
    columns: dict[str, Any] = {}
    for column in sorted(column_files):
        populated = int(column_nonblank[column])
        placeholder_count = int(column_placeholders[column])
        columns[column] = {
            "workbooks_present": int(column_files[column]),
            "workbook_coverage_ratio": round(column_files[column] / len(profiles), 6),
            "populated_cells": populated,
            "blank_cells": max(total_rows - populated, 0),
            "populated_ratio": round(populated / total_rows, 6) if total_rows else 0.0,
            "placeholder_cells": placeholder_count,
            "placeholder_ratio_of_populated": (
                round(placeholder_count / populated, 6) if populated else 0.0
            ),
            "max_text_length": int(column_max_lengths[column]),
            "type_counts": dict(sorted(column_type_counts[column].items())),
            "sensitive_value_class": any(
                marker in column.casefold() for marker in SENSITIVE_HEADER_MARKERS
            ),
        }

    observed_dates = sorted(
        profile.observed_date for profile in profiles if profile.observed_date
    )
    row_counts = [profile.row_count for profile in profiles]
    dominant_schema, dominant_schema_count = schema_counts.most_common(1)[0]
    return {
        "schema_version": SUMMARY_SCHEMA,
        "sanitized": True,
        "source_rows_exported": False,
        "source_values_exported": False,
        "do_not_commit": True,
        "live_snowflake_validation_performed": False,
        "workbook_count": len(profiles),
        "observed_date_min": observed_dates[0] if observed_dates else "",
        "observed_date_max": observed_dates[-1] if observed_dates else "",
        "total_profiled_rows": total_rows,
        "row_count_distribution": {
            "min": min(row_counts),
            "median": int(statistics.median(row_counts)),
            "p95": _percentile(row_counts, 0.95),
            "max": max(row_counts),
        },
        "header_row_distribution": {
            str(key): value for key, value in sorted(header_rows.items())
        },
        "distinct_schema_count": len(schema_counts),
        "dominant_schema_sha256": dominant_schema,
        "dominant_schema_workbook_count": dominant_schema_count,
        "schema_fingerprint_counts": dict(sorted(schema_counts.items())),
        "hyperlink_cell_count": sum(profile.hyperlink_count for profile in profiles),
        "formula_cell_count": sum(profile.formula_count for profile in profiles),
        "identifier_quality": {
            column: dict(sorted(counts.items()))
            for column, counts in sorted(identifier_totals.items())
        },
        "columns": columns,
        "workbook_fingerprints_sha256": hashlib.sha256(
            "\n".join(sorted(profile.digest for profile in profiles)).encode("ascii")
        ).hexdigest(),
    }


def discover_workbooks(input_dir: Path) -> list[Path]:
    return sorted(
        path
        for path in input_dir.iterdir()
        if path.is_file()
        and path.suffix.casefold() == ".xlsx"
        and not path.name.startswith("~$")
    )


def _safe_summary_path(path: Path) -> Path:
    target = path.expanduser().resolve()
    try:
        relative = target.relative_to(REPO_ROOT)
    except ValueError:
        return target
    if not relative.parts or relative.parts[0] != ".adoptiq-acceptance":
        raise ValueError(
            "in-repository CSOne corpus summaries must stay under .adoptiq-acceptance"
        )
    return target


def _write_json(path: Path, payload: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    temporary.write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    os.replace(temporary, path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument(
        "--summary",
        type=Path,
        default=REPO_ROOT / ".adoptiq-acceptance" / "csone-corpus-profile.json",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Profile only the newest N workbooks; 0 profiles all workbooks.",
    )
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    args = build_parser().parse_args(list(argv) if argv is not None else None)
    input_dir = args.input_dir.expanduser().resolve()
    if not input_dir.is_dir():
        print("CSOne corpus input directory is unavailable", file=sys.stderr)
        return 2
    workbooks = discover_workbooks(input_dir)
    if args.limit:
        workbooks = workbooks[-max(int(args.limit), 0) :]
    if not workbooks:
        print("CSOne corpus contains no .xlsx workbooks", file=sys.stderr)
        return 2
    try:
        profiles = [profile_workbook(path) for path in workbooks]
        summary = summarize_profiles(profiles)
        summary_path = _safe_summary_path(args.summary)
        _write_json(summary_path, summary)
    except (OSError, ValueError) as exc:
        print(f"CSOne corpus profiling failed: {type(exc).__name__}", file=sys.stderr)
        return 5
    print(
        json.dumps(
            {
                "workbook_count": summary["workbook_count"],
                "distinct_schema_count": summary["distinct_schema_count"],
                "total_profiled_rows": summary["total_profiled_rows"],
                "summary": str(summary_path),
            },
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
