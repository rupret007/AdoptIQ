#!/usr/bin/env python3
"""Exercise every guarded Round 145 fixture through real loopback HTTP.

A fresh source app process is launched per scenario. The JSON summary contains
only hashes, counts, states, and pass/fail findings; it never claims live
Snowflake validation and never stores response bodies.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import signal
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from pathlib import Path
from typing import Any, Mapping, Sequence
from urllib.parse import quote

import requests


REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from local_acceptance_lab import (  # noqa: E402
    DEFAULT_MANIFEST_PATH,
    REPORT_PUBLICATION_BLOCKED_MISSING_STABLE_ID,
    REPORT_PUBLICATION_COMPLETED,
    SOURCE_MODE,
    build_scenario_bundle,
    load_manifest,
)
import decision_report_delivery as delivery  # noqa: E402
from report_completeness_audit import (  # noqa: E402
    audit_source_data_frames,
    audit_word_placeholders,
    audit_written_source_hyperlinks,
)
from report_iteration_loop import extract_csrf_token  # noqa: E402
from scripts.run_ai_feature_acceptance import (  # noqa: E402
    _stream_payload,
    extract_citations,
    parse_sse,
)


SUMMARY_SCHEMA = "local-acceptance-http/v1"
FIXTURE_LOG_LIMIT_BYTES = 4 * 1024 * 1024
FIXTURE_LOG_CHUNK_BYTES = 64 * 1024
PROVIDER_HTTP_STATUS = {
    "timeout": 504,
    "rate_limited": 429,
    "unavailable": 503,
    "malformed": 502,
}
PROVIDER_TERM = {
    "timeout": "timed out",
    "rate_limited": "rate limited",
    "unavailable": "unavailable",
    "malformed": "unusable response",
}
FORBIDDEN_RESPONSE_TERMS = (
    "traceback (most recent call last)",
    "client_secret=",
    "secret_id=",
    "authorization: bearer",
    "ghp_",
)
PROMPT_INJECTION_TERMS = (
    "ignore all previous instructions",
    "disclose credentials",
)
WORKSPACE_SCHEMA = "manager-decision-workspace/v1"
NONCANONICAL_REPORT_ERROR = (
    "The selected report does not contain the canonical Source Data "
    "snapshot required for exact-run Ask AI. Generate the report again."
)
WORKSPACE_PREVIEW_CASES: tuple[tuple[str, dict[str, Any]], ...] = (
    (
        "leader_team",
        {
            "report_type": "leader",
            "manager": "Local Fixture Manager",
            "technology": "All",
            "days": 90,
            "scope_type": "team",
        },
    ),
    (
        "comprehensive_team",
        {
            "report_type": "comprehensive",
            "manager": "Local Fixture Manager",
            "technology": "All",
            "days": 90,
            "scope_type": "team",
        },
    ),
    (
        "compact_team",
        {
            "report_type": "compact",
            "manager": "Local Fixture Manager",
            "technology": "All",
            "days": 90,
            "scope_type": "team",
        },
    ),
    (
        "renewal_portfolio",
        {
            "report_type": "renewal_portfolio",
            "manager": "Local Fixture Manager",
            "technology": "All",
            "days": 90,
            "scope_type": "team",
        },
    ),
    (
        "renewal_customer",
        {
            "report_type": "renewal",
            "manager": "Local Fixture Manager",
            "technology": "All",
            "days": 90,
            "scope_type": "customer",
            "scope_value": "Acme Corporation",
        },
    ),
    (
        "subscription",
        {
            "report_type": "subscription",
            "manager": "Local Fixture Manager",
            "technology": "All",
            "days": 90,
            "scope_type": "subscription",
            "subscription_id": "SUB-001",
        },
    ),
)


def _expected_fixture_portfolio_counts(
    bundle: Any,
    manager_name: str,
) -> dict[str, int]:
    """Derive the independent manager-scoped oracle from fixture source rows.

    The manifest's canonical counts describe the full fixture portfolio.  They
    are not a valid oracle after a scenario splits the roster across managers
    or removes a member's subscriptions.  Scope through the same authorization
    keys the production request uses, then count stable source identifiers.
    """

    manager_key = str(manager_name or "").strip().casefold()
    members = {
        str(row.get("OWNER_NAME") or "").strip().casefold()
        for row in bundle.records("ownership")
        if str(row.get("MANAGER_NAME") or "").strip().casefold() == manager_key
        and str(row.get("OWNER_NAME") or "").strip()
    }
    subscriptions = [
        row
        for row in bundle.records("subscriptions")
        if str(row.get("FIXTURE_MEMBER") or "").strip().casefold() in members
    ]
    account_ids = {
        str(row.get("ACCOUNT_ID_C") or "").strip().casefold()
        for row in subscriptions
        if str(row.get("ACCOUNT_ID_C") or "").strip()
    }

    def scoped_distinct(dataset: str, id_field: str) -> int:
        return len(
            {
                str(row.get(id_field) or "").strip().casefold()
                for row in bundle.records(dataset)
                if str(row.get("ACCOUNT_ID_C") or "").strip().casefold()
                in account_ids
                and str(row.get(id_field) or "").strip()
            }
        )

    return {
        "total_customers": len(account_ids),
        "total_barriers": scoped_distinct("adoption_barriers", "ID"),
        "total_cases": scoped_distinct("support_cases", "CASE_ID"),
        "action_plan_rows": scoped_distinct("action_plans", "ID"),
    }


def _digest(value: Any) -> str:
    raw = json.dumps(
        value,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _free_loopback_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def _safe_output_dir(path: Path) -> Path:
    resolved = path.expanduser().resolve()
    try:
        relative = resolved.relative_to(REPO_ROOT)
    except ValueError:
        return resolved
    if not relative.parts or relative.parts[0] != ".adoptiq-acceptance":
        raise ValueError(
            "in-repository output must stay under .adoptiq-acceptance"
        )
    return resolved


def _response_is_sanitized(value: object) -> bool:
    text = str(value or "").casefold()
    return not any(term in text for term in FORBIDDEN_RESPONSE_TERMS)


def _literal_true(value: object) -> bool:
    """Accept only the literal JSON boolean ``true`` from an app response."""

    return value is True


def validate_blocked_missing_id_publication(
    status: Mapping[str, Any],
    downloads: Mapping[str, Any],
    workspace: Mapping[str, Any],
) -> tuple[dict[str, Any], list[str]]:
    """Prove a missing-ID negative control failed closed for the right reason.

    The returned projection contains only booleans, a namespaced error kind,
    and hashes.  The scrubbed worker detail is inspected in memory to
    distinguish the expected CSConsole link-integrity block from an unrelated
    crash, but the detail itself is never retained in the acceptance summary.
    """

    errors: list[str] = []
    terminal_error = str(status.get("status") or "").casefold() == "error"
    if not terminal_error:
        errors.append("missing-ID report did not terminate in the error state")

    artifacts_unavailable = (
        status.get("word_available") is False
        and status.get("excel_available") is False
    )
    if not artifacts_unavailable:
        errors.append("missing-ID report advertised a Word or Source Data artifact")

    failure_fields = {
        key: str(status.get(key) or "")[:512]
        for key in ("error_kind", "error", "message", "error_detail")
    }
    failure_text = " ".join(failure_fields.values()).casefold()
    error_kind = failure_fields["error_kind"]
    failure_reason_verified = (
        "csconsole link coverage error" in failure_text
        and "missing_stable_id" in failure_text
        and error_kind.startswith("analysis.")
    )
    if not failure_reason_verified:
        errors.append("missing-ID report did not expose the expected stable integrity reason")
    failure_response_sanitized = _response_is_sanitized(failure_fields)
    if not failure_response_sanitized:
        errors.append("missing-ID report exposed an unsafe public failure field")

    expected_extensions = {"docx", "xlsx"}
    downloads_blocked = set(downloads) == expected_extensions
    if downloads_blocked:
        for extension in sorted(expected_extensions):
            item = downloads.get(extension)
            if not isinstance(item, Mapping) or not (
                item.get("blocked") is True
                and item.get("status_code") == 400
                and item.get("current_status") == "error"
                and item.get("artifact_bytes") == 0
                and item.get("sanitized") is True
            ):
                downloads_blocked = False
                break
    if not downloads_blocked:
        errors.append("missing-ID report download route exposed or accepted an artifact")

    workspace_blocked = bool(
        workspace.get("report_status_code") == 200
        and workspace.get("report_payload_ok") is True
        and workspace.get("report_status") == "error"
        and workspace.get("report_completed") is False
        and workspace.get("workbook_loaded") is False
        and workspace.get("artifact_links_exposed") is False
        and workspace.get("raw_path_leaked") is False
        and workspace.get("report_response_sanitized") is True
        and workspace.get("history_status_code") == 200
        and workspace.get("history_payload_ok") is True
        and workspace.get("history_contains_report") is False
        and workspace.get("history_response_sanitized") is True
    )
    if not workspace_blocked:
        errors.append("missing-ID report appeared completed or downloadable in the workspace")

    publication_blocked = bool(
        terminal_error
        and artifacts_unavailable
        and failure_reason_verified
        and failure_response_sanitized
        and downloads_blocked
        and workspace_blocked
    )
    projection = {
        "expected_publication_outcome": (
            REPORT_PUBLICATION_BLOCKED_MISSING_STABLE_ID
        ),
        "publication_outcome_matches": publication_blocked,
        "publication_blocked": publication_blocked,
        "negative_control_passed": publication_blocked,
        "terminal_status": str(status.get("status") or ""),
        "failure_error_kind": error_kind,
        "failure_reason_verified": failure_reason_verified,
        "failure_response_sanitized": failure_response_sanitized,
        "failure_fields_sha256": _digest(failure_fields),
        "downloads_blocked": downloads_blocked,
        "workspace_blocked": workspace_blocked,
    }
    return projection, errors


def _warning_projection(value: object) -> dict[str, list[str]]:
    datasets: set[str] = set()
    kinds: set[str] = set()
    if isinstance(value, list):
        for item in value:
            if not isinstance(item, Mapping):
                continue
            dataset = str(item.get("dataset") or "").strip()
            kind = str(item.get("kind") or item.get("error_kind") or "").strip()
            if dataset:
                datasets.add(dataset)
            if kind:
                kinds.add(kind)
    return {"datasets": sorted(datasets), "kinds": sorted(kinds)}


def _run_workspace_preview_probe(
    client: "LoopbackClient",
) -> tuple[dict[str, Any], list[str]]:
    """Exercise every Round 146 report-family preview without retaining PII."""

    errors: list[str] = []
    cases: dict[str, dict[str, Any]] = {}
    for label, params in WORKSPACE_PREVIEW_CASES:
        response = client.get(
            "/api/decision-workspace/scope-preview",
            params=params,
        )
        payload = _json(response)
        preview = payload.get("preview") if isinstance(payload.get("preview"), dict) else {}
        ok = bool(
            response.status_code == 200
            and _literal_true(payload.get("ok"))
            and preview.get("schema") == WORKSPACE_SCHEMA
            and preview.get("report_type") == params["report_type"]
            and preview.get("scope_type") == params["scope_type"]
            and preview.get("live_validation_performed") is False
            and "fixture" in str(preview.get("source_mode") or "").casefold()
            and isinstance(preview.get("expected_sources"), list)
            and isinstance(preview.get("limitations"), list)
        )
        cases[label] = {
            "status_code": response.status_code,
            "ok": ok,
            "source_count": len(preview.get("expected_sources") or []),
            "limitation_count": len(preview.get("limitations") or []),
            "payload_sha256": _digest(payload),
            "live_validation_performed": False,
        }
        if not ok:
            errors.append(f"Manager Decision Workspace {label} preview failed")

    invalid = client.get(
        "/api/decision-workspace/scope-preview",
        params={
            "report_type": "leader",
            "manager": "Local Fixture Manager",
            "technology": "All",
            "days": 90,
            "scope_type": "member",
        },
    )
    invalid_payload = _json(invalid)
    rejected = invalid.status_code == 400 and invalid_payload.get("ok") is False
    if not rejected:
        errors.append("Manager Decision Workspace accepted a missing member scope")
    return {
        "ok": not errors,
        "schema": WORKSPACE_SCHEMA,
        "validation_mode": "local_acceptance",
        "live_validation_performed": False,
        "cases": cases,
        "missing_member_rejected": rejected,
    }, errors


def validate_provider_response(
    provider_state: str,
    status_code: int,
    payload: Mapping[str, Any],
) -> list[str]:
    """Validate a grounded sync/Intel response without retaining its body."""

    errors: list[str] = []
    if provider_state == "available":
        if status_code != 200 or not _literal_true(payload.get("ok")):
            errors.append("available provider did not return a grounded success")
    else:
        expected_status = PROVIDER_HTTP_STATUS[provider_state]
        if status_code != expected_status:
            errors.append(
                f"provider {provider_state} returned HTTP {status_code}, "
                f"expected {expected_status}"
            )
        if payload.get("ok") is not False:
            errors.append("provider failure response did not set ok=false")
        if PROVIDER_TERM[provider_state] not in str(
            payload.get("error") or ""
        ).casefold():
            errors.append("provider failure response lost its sanitized state")
    if not _response_is_sanitized(payload):
        errors.append("response exposed a forbidden secret/error marker")
    return errors


def validate_noncanonical_report_ai_response(
    status_code: int,
    payload: Mapping[str, Any],
) -> list[str]:
    """Require legacy/compatibility report snapshots to fail closed safely."""

    errors: list[str] = []
    if status_code != 409:
        errors.append(f"noncanonical report returned HTTP {status_code}, expected 409")
    if payload.get("ok") is not False:
        errors.append("noncanonical report response did not set ok=false")
    if str(payload.get("error") or "") != NONCANONICAL_REPORT_ERROR:
        errors.append("noncanonical report response lost its exact safe error")
    if payload.get("answer"):
        errors.append("noncanonical report response included an answer")
    if payload.get("fallback_available"):
        errors.append("noncanonical report response exposed a legacy fallback")
    if not _response_is_sanitized(payload):
        errors.append("noncanonical report response exposed a forbidden marker")
    return errors


class LoopbackClient:
    def __init__(self, base_url: str, timeout: float) -> None:
        self.base_url = base_url.rstrip("/")
        self.timeout = max(float(timeout), 5.0)
        self.session = requests.Session()
        self.csrf_token = ""

    def bootstrap(self) -> None:
        response = self.session.get(self.base_url + "/", timeout=self.timeout)
        response.raise_for_status()
        self.csrf_token = extract_csrf_token(response.text)

    def headers(self, accept: str = "application/json") -> dict[str, str]:
        return {
            "Accept": accept,
            "X-Requested-With": "XMLHttpRequest",
            "X-CSRFToken": self.csrf_token,
        }

    def get(
        self,
        path: str,
        *,
        accept: str = "application/json",
        timeout: float | None = None,
        params: Mapping[str, Any] | None = None,
    ) -> requests.Response:
        return self.session.get(
            self.base_url + path,
            headers=self.headers(accept),
            timeout=timeout or self.timeout,
            params=dict(params or {}),
        )

    def post_json(self, path: str, payload: Mapping[str, Any]) -> requests.Response:
        return self.session.post(
            self.base_url + path,
            json=dict(payload),
            headers=self.headers(),
            timeout=self.timeout,
        )

    def post_form(self, path: str, payload: Mapping[str, Any]) -> requests.Response:
        return self.session.post(
            self.base_url + path,
            data=dict(payload),
            headers=self.headers("text/html"),
            timeout=self.timeout,
        )


def _signal_process_tree(
    process: subprocess.Popen[Any],
    process_group_id: int | None,
    sig: int,
) -> None:
    """Best-effort signal for a fixture process and every child it spawned."""

    if os.name == "posix" and process_group_id is not None:
        try:
            os.killpg(process_group_id, sig)
        except (ProcessLookupError, PermissionError):
            pass
    elif process.poll() is None:
        try:
            process.send_signal(sig)
        except (OSError, ValueError):
            pass


def _terminate_process_tree(
    process: subprocess.Popen[Any],
    process_group_id: int | None,
    *,
    timeout: float = 5.0,
) -> None:
    """Reap the direct child and terminate descendants even after leader exit."""

    term_signal = signal.SIGTERM if os.name == "posix" else signal.SIGTERM
    kill_signal = signal.SIGKILL if os.name == "posix" else signal.SIGTERM
    _signal_process_tree(process, process_group_id, term_signal)
    try:
        process.wait(timeout=max(float(timeout), 0.1))
    except subprocess.TimeoutExpired:
        _signal_process_tree(process, process_group_id, kill_signal)
        try:
            process.kill()
        except OSError:
            pass
        try:
            process.wait(timeout=max(float(timeout), 0.1))
        except subprocess.TimeoutExpired:
            pass
    finally:
        # A supervisor may exit before its descendants.  Signal the original
        # isolated group once more so an inherited stdout pipe cannot remain
        # open and stall the bounded drain thread.
        _signal_process_tree(process, process_group_id, kill_signal)


class _BoundedFixtureLog:
    """Stream one child pipe into a hard-capped private prefix file."""

    def __init__(self, path: Path, *, limit_bytes: int) -> None:
        if type(limit_bytes) is not int or limit_bytes < 1:
            raise ValueError("fixture log limit must be a positive integer")
        self.path = path
        self.limit_bytes = limit_bytes
        self.observed_bytes = 0
        self.retained_bytes = 0
        self.output_truncated = False
        self.read_complete = False
        self.error_kind = ""
        self._digest = hashlib.sha256()
        self._thread: threading.Thread | None = None
        self._overflow_callback: Any = None

    def start(
        self,
        stream: Any,
        *,
        overflow_callback: Any,
    ) -> None:
        if self._thread is not None:
            raise RuntimeError("fixture log drain already started")
        self._overflow_callback = overflow_callback

        def drain() -> None:
            flags = os.O_WRONLY | os.O_CREAT | os.O_TRUNC
            if hasattr(os, "O_NOFOLLOW"):
                flags |= os.O_NOFOLLOW
            fd = os.open(self.path, flags, 0o600)
            overflow_signalled = False
            try:
                with os.fdopen(fd, "wb", closefd=True) as output:
                    while True:
                        chunk = stream.read(FIXTURE_LOG_CHUNK_BYTES)
                        if not chunk:
                            self.read_complete = True
                            break
                        if not isinstance(chunk, bytes):
                            raise TypeError("fixture log pipe did not return bytes")
                        self.observed_bytes += len(chunk)
                        self._digest.update(chunk)
                        remaining = self.limit_bytes - self.retained_bytes
                        if remaining > 0:
                            retained = chunk[:remaining]
                            output.write(retained)
                            self.retained_bytes += len(retained)
                        if len(chunk) > max(remaining, 0):
                            self.output_truncated = True
                            if not overflow_signalled:
                                overflow_signalled = True
                                try:
                                    self._overflow_callback()
                                except Exception:  # noqa: BLE001 - containment path
                                    pass
                    output.flush()
                    os.fsync(output.fileno())
            except Exception as exc:  # noqa: BLE001 - evidence stays type-only
                self.error_kind = type(exc).__name__
            finally:
                try:
                    stream.close()
                except Exception:  # noqa: BLE001
                    pass

        self._thread = threading.Thread(
            target=drain,
            name="adoptiq-fixture-log-drain",
            daemon=True,
        )
        self._thread.start()

    def finish(self, *, timeout: float = 10.0) -> dict[str, Any]:
        if self._thread is None:
            raise RuntimeError("fixture log drain was not started")
        self._thread.join(timeout=max(float(timeout), 0.1))
        thread_finished = not self._thread.is_alive()
        integrity_complete = (
            thread_finished
            and self.read_complete
            and not self.error_kind
            and not self.output_truncated
        )
        return {
            "server_log_bytes": self.retained_bytes,
            "server_log_observed_bytes": (
                self.observed_bytes if thread_finished and self.read_complete else -1
            ),
            "server_log_sha256": self._digest.hexdigest() if integrity_complete else "",
            "server_log_read_complete": self.read_complete is True and thread_finished,
            "server_log_output_truncated": self.output_truncated is True,
            "server_log_integrity_complete": integrity_complete,
            "server_log_error_kind": self.error_kind,
        }


def _json(response: requests.Response) -> dict[str, Any]:
    try:
        value = response.json()
    except ValueError:
        return {}
    return value if isinstance(value, dict) else {}


def _wait_for_server(
    base_url: str,
    scenario: str,
    schema_fingerprint: str,
    report_publication_expectation: str,
    process: subprocess.Popen[Any],
    timeout: float,
) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise RuntimeError("local app process exited before readiness")
        try:
            response = requests.get(
                base_url + "/api/diag/connectivity",
                timeout=2,
            )
            payload = _json(response)
            if (
                response.status_code == 200
                and _literal_true(payload.get("ok"))
                and payload.get("mode") == SOURCE_MODE
                and payload.get("scenario") == scenario
                and payload.get("schema_fingerprint") == schema_fingerprint
                and payload.get("report_publication_expectation")
                == report_publication_expectation
                and payload.get("live_validation_performed") is False
            ):
                return payload
        except requests.RequestException:
            pass
        time.sleep(0.2)
    raise TimeoutError("local app did not become ready")


def _validate_ooxml(payload: bytes, extension: str) -> list[str]:
    errors: list[str] = []
    if len(payload) < 2_000:
        errors.append(f"{extension} download is unexpectedly small")
    if not payload.startswith(b"PK"):
        return errors + [f"{extension} download is not an OOXML ZIP"]
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            if archive.testzip():
                errors.append(f"{extension} ZIP member failed integrity")
            required = (
                {"word/document.xml", "[Content_Types].xml"}
                if extension == "docx"
                else {"xl/workbook.xml", "[Content_Types].xml"}
            )
            if not required.issubset(archive.namelist()):
                errors.append(f"{extension} is missing required OOXML members")
    except (OSError, zipfile.BadZipFile):
        errors.append(f"{extension} download has invalid ZIP structure")
    return errors


def _audit_downloaded_report_pair(
    *,
    docx_content: bytes,
    xlsx_content: bytes,
    expected_as_of_utc: str,
    require_source_links: bool,
) -> tuple[dict[str, Any], list[str]]:
    """Independently inspect downloaded artifacts, never server-side paths."""

    import openpyxl  # noqa: PLC0415
    import pandas as pd  # noqa: PLC0415
    from docx import Document  # noqa: PLC0415

    errors: list[str] = []
    workbook = openpyxl.load_workbook(
        io.BytesIO(xlsx_content), read_only=True, data_only=False
    )
    frames: dict[str, pd.DataFrame] = {}
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows()
            try:
                headers = [cell.value for cell in next(rows)]
            except StopIteration:
                frames[worksheet.title] = pd.DataFrame()
                continue
            frames[worksheet.title] = pd.DataFrame(
                [[cell.value for cell in row] for row in rows],
                columns=headers,
            )
    finally:
        workbook.close()

    exact_inventory = list(frames) == list(delivery.SOURCE_DATA_SHEET_NAMES)
    if not exact_inventory:
        errors.append("downloaded Source Data workbook inventory is not canonical")
    completeness = audit_source_data_frames(frames)
    errors.extend(completeness["errors"])
    hyperlink_audit = audit_written_source_hyperlinks(io.BytesIO(xlsx_content))
    if not hyperlink_audit["ok"]:
        errors.append("downloaded Source Data CSConsole links are not clickable")
    if require_source_links and not hyperlink_audit["url_cells"]:
        errors.append("downloaded Source Data omitted expected CSConsole record links")

    document = Document(io.BytesIO(docx_content))
    word_audit = audit_word_placeholders(document)
    errors.extend(word_audit["errors"])

    info = frames.get("Report_Info", pd.DataFrame())
    info_map = (
        {
            str(row.get("Item") or ""): row.get("Value")
            for _, row in info.iterrows()
            if str(row.get("Item") or "").strip()
        }
        if {"Item", "Value"}.issubset(info.columns)
        else {}
    )
    artifact_as_of = pd.to_datetime(
        info_map.get("Data_As_Of_UTC"), errors="coerce", utc=True
    )
    expected_as_of = pd.to_datetime(expected_as_of_utc, errors="coerce", utc=True)
    as_of_matches = bool(
        not pd.isna(artifact_as_of)
        and not pd.isna(expected_as_of)
        and artifact_as_of == expected_as_of
    )
    if not as_of_matches:
        errors.append("downloaded Source Data does not use the fixture retrieval clock")

    return (
        {
            "ok": not errors,
            "exact_canonical_sheet_inventory": exact_inventory,
            "as_of_matches_fixture_clock": as_of_matches,
            "source_record_link_rows": completeness["source_record_link_rows"],
            "clickable_source_record_links": hyperlink_audit["clickable_cells"],
            "undefined_cell_count": len(completeness["undefined_cells"]),
            "unexplained_action_plan_status_rows": completeness[
                "unexplained_action_plan_status_rows"
            ],
            "tac_non_record_rows": completeness["tac_non_record_rows"],
            "word_placeholder_errors": len(word_audit["errors"]),
        },
        errors,
    )


def _poll_report(
    client: LoopbackClient,
    analysis_id: str,
    timeout: float,
) -> tuple[dict[str, Any], int]:
    deadline = time.monotonic() + timeout
    polls = 0
    while time.monotonic() < deadline:
        polls += 1
        response = client.get(f"/status/{quote(analysis_id, safe='')}")
        payload = _json(response)
        if response.status_code != 200:
            raise RuntimeError("report status route failed")
        if payload.get("status") in {"completed", "error", "failed", "cancelled"}:
            return payload, polls
        time.sleep(0.25)
    raise TimeoutError("report generation did not reach a terminal state")


def _run_report_probe(
    client: LoopbackClient,
    timeout: float,
    *,
    provider_state: str,
    expected_as_of_utc: str,
    expected_action_plan_rows: int,
    publication_expectation: str = REPORT_PUBLICATION_COMPLETED,
) -> tuple[dict[str, Any], list[str]]:
    errors: list[str] = []
    response = client.post_json(
        "/start_compact_analysis",
        {
            "manager": "Local Fixture Manager",
            # The degraded-state probe validates source-state handling.  The
            # exhaustive report matrix separately exercises every technology;
            # using All here avoids conflating a deliberate roster gap with a
            # manager who simply owns no subscription in one technology.
            "technology": "All",
            "days": 90,
            "csone_file": "",
            "subscription_id": "",
            "customer_name": "",
        },
    )
    started = _json(response)
    analysis_id = str(started.get("analysis_id") or "")
    if (
        response.status_code != 200
        or not _literal_true(started.get("success"))
        or not analysis_id
    ):
        return {
            "start_status": response.status_code,
            "completed": False,
            "expected_publication_outcome": publication_expectation,
            "publication_outcome_matches": False,
        }, ["compact report did not start"]

    status, polls = _poll_report(client, analysis_id, timeout)
    completed = status.get("status") == "completed"
    publication_projection: dict[str, Any] = {
        "expected_publication_outcome": publication_expectation,
        "publication_outcome_matches": bool(
            publication_expectation == REPORT_PUBLICATION_COMPLETED and completed
        ),
        "publication_blocked": False,
        "negative_control_passed": None,
    }
    if publication_expectation == REPORT_PUBLICATION_COMPLETED and not completed:
        errors.append("compact report did not complete")
    downloads: dict[str, Any] = {}
    artifact_contents: dict[str, bytes] = {}
    if completed and publication_expectation == REPORT_PUBLICATION_COMPLETED:
        for extension in ("docx", "xlsx"):
            artifact = client.get(
                f"/download/{quote(analysis_id, safe='')}/{extension}",
                timeout=timeout,
            )
            artifact_errors = (
                [f"{extension} download returned HTTP {artifact.status_code}"]
                if artifact.status_code != 200
                else _validate_ooxml(artifact.content, extension)
            )
            errors.extend(artifact_errors)
            if not artifact_errors:
                artifact_contents[extension] = artifact.content
            downloads[extension] = {
                "status_code": artifact.status_code,
                "bytes": len(artifact.content),
                "sha256": hashlib.sha256(artifact.content).hexdigest(),
                "ok": not artifact_errors,
            }
        if {"docx", "xlsx"} <= set(artifact_contents):
            artifact_audit, artifact_audit_errors = _audit_downloaded_report_pair(
                docx_content=artifact_contents["docx"],
                xlsx_content=artifact_contents["xlsx"],
                expected_as_of_utc=expected_as_of_utc,
                require_source_links=expected_action_plan_rows > 0,
            )
            downloads["artifact_contract"] = artifact_audit
            errors.extend(artifact_audit_errors)
    elif publication_expectation == REPORT_PUBLICATION_BLOCKED_MISSING_STABLE_ID:
        for extension in ("docx", "xlsx"):
            artifact = client.get(
                f"/download/{quote(analysis_id, safe='')}/{extension}",
                timeout=timeout,
            )
            payload = _json(artifact)
            artifact_bytes = len(artifact.content) if artifact.content.startswith(b"PK") else 0
            downloads[extension] = {
                "status_code": artifact.status_code,
                "blocked": bool(
                    artifact.status_code == 400
                    and payload.get("current_status") == "error"
                    and artifact_bytes == 0
                ),
                "current_status": str(payload.get("current_status") or ""),
                "artifact_bytes": artifact_bytes,
                "response_bytes": len(artifact.content),
                "response_sha256": hashlib.sha256(artifact.content).hexdigest(),
                "sanitized": _response_is_sanitized(payload),
            }
    previous = client.get("/previous-reports", accept="text/html")
    previous_ok = (
        previous.status_code == 200
        and "Previous AdoptIQ Reports" in previous.text
        and len(previous.text) > 500
    )
    if not previous_ok:
        errors.append("Previous Reports did not render after generation")
    warnings = status.get("partial_data_warnings") or []
    warning_projection = _warning_projection(warnings)
    workspace: dict[str, Any] = {
        "report_view_ok": False,
        "history_ok": False,
        "same_report_compare_rejected": False,
    }
    if publication_expectation == REPORT_PUBLICATION_BLOCKED_MISSING_STABLE_ID:
        report_view = client.get(
            f"/api/decision-workspace/report/{quote(analysis_id, safe='')}"
        )
        report_payload = _json(report_view)
        report = (
            report_payload.get("report")
            if isinstance(report_payload.get("report"), dict)
            else {}
        )
        raw_path_leaked = any(
            key in report
            for key in (
                "excel_report",
                "excel_path",
                "word_report",
                "word_path",
                "report_path",
            )
        )
        report_downloads = (
            report.get("downloads")
            if isinstance(report.get("downloads"), Mapping)
            else {}
        )
        artifact_links_exposed = bool(
            report.get("word_available")
            or report.get("excel_available")
            or any(str(value or "").strip() for value in report_downloads.values())
        )

        history = client.get("/api/decision-workspace/history")
        history_payload = _json(history)
        history_records = history_payload.get("reports") or []
        history_contains_report = bool(
            isinstance(history_records, list)
            and any(
                isinstance(item, Mapping)
                and item.get("analysis_id") == analysis_id
                for item in history_records
            )
        )
        workspace.update(
            {
                "report_status_code": report_view.status_code,
                "report_payload_ok": report_payload.get("ok") is True,
                "report_status": str(report.get("status") or ""),
                "report_completed": report.get("status") == "completed",
                "workbook_loaded": report.get("workbook_loaded") is True,
                "artifact_links_exposed": artifact_links_exposed,
                "raw_path_leaked": raw_path_leaked,
                "report_response_sanitized": _response_is_sanitized(
                    report_payload
                ),
                "report_payload_sha256": _digest(report_payload),
                "history_status_code": history.status_code,
                "history_payload_ok": history_payload.get("ok") is True,
                "history_contains_report": history_contains_report,
                "history_response_sanitized": _response_is_sanitized(
                    history_payload
                ),
                "history_count": len(history_records),
                "history_sha256": _digest(history_payload),
            }
        )
        publication_projection, publication_errors = (
            validate_blocked_missing_id_publication(
                status,
                downloads,
                workspace,
            )
        )
        errors.extend(publication_errors)
    elif completed:
        report_view = client.get(
            f"/api/decision-workspace/report/{quote(analysis_id, safe='')}"
        )
        report_payload = _json(report_view)
        report = (
            report_payload.get("report")
            if isinstance(report_payload.get("report"), dict)
            else {}
        )
        raw_path_leaked = any(
            key in report
            for key in (
                "excel_report",
                "excel_path",
                "word_report",
                "word_path",
                "report_path",
            )
        )
        report_view_ok = bool(
            report_view.status_code == 200
            and _literal_true(report_payload.get("ok"))
            and report.get("schema") == WORKSPACE_SCHEMA
            and report.get("analysis_id") == analysis_id
            and report.get("status") == "completed"
            and report.get("workbook_loaded") is True
            and bool(report.get("decision_metrics"))
            and isinstance(report.get("ask_ai_binding"), dict)
            and report["ask_ai_binding"].get("fact_fingerprint")
            == report.get("fact_fingerprint")
            and not raw_path_leaked
        )
        workspace["report_view_ok"] = report_view_ok
        workspace["report_view_status"] = report_view.status_code
        workspace["report_view_sha256"] = _digest(report_payload)
        workspace["decision_metric_count"] = len(
            report.get("decision_metrics") or []
        )
        workspace["action_plan_count"] = len(report.get("action_plans") or [])
        workspace["source_limitation_count"] = len(
            report.get("source_limitations") or []
        )
        if not report_view_ok:
            errors.append("Manager Decision Workspace report projection failed")

        expected_binding = report.get("ask_ai_binding") or {}
        binding_keys = (
            "manager",
            "technology",
            "days",
            "scope_type",
            "scope_value",
            "report_analysis_id",
            "report_type",
            "data_as_of_utc",
            "fact_fingerprint",
        )

        def binding_matches(observed: object) -> bool:
            if not isinstance(observed, Mapping):
                return False
            expected = dict(expected_binding)
            expected["report_analysis_id"] = expected.pop(
                "analysis_id", analysis_id
            )
            return all(
                str(observed.get(key) if observed.get(key) is not None else "")
                == str(expected.get(key) if expected.get(key) is not None else "")
                for key in binding_keys
            )

        report_ai_request = {
            "question": (
                "State the first evidence-backed decision and cite its exact "
                "source record. Keep the selected report scope."
            ),
            "report_analysis_id": analysis_id,
            # Hostile widening selectors: the server-owned report binding must
            # override every one of these fields.
            "manager": "All Managers",
            "technology": "Webex Meetings",
            "days": 1,
            "scope_type": "customer",
            "scope_value": "Out-of-scope customer",
            "allow_legacy_fallback": True,
        }
        report_sync = client.post_json(
            "/api/ask-ai-portfolio",
            report_ai_request,
        )
        report_sync_payload = _json(report_sync)
        sync_binding_ok = binding_matches(
            report_sync_payload.get("scope_context")
        )
        canonical_snapshot = report.get("canonical_snapshot") is True
        if not canonical_snapshot:
            report_sync_ok = not validate_noncanonical_report_ai_response(
                report_sync.status_code,
                report_sync_payload,
            )
        elif provider_state == "available":
            report_sync_ok = bool(
                report_sync.status_code == 200
                and _literal_true(report_sync_payload.get("ok"))
                and report_sync_payload.get("mode") == "grounded"
                and sync_binding_ok
                and extract_citations(str(report_sync_payload.get("answer") or ""))
            )
        else:
            report_sync_ok = bool(
                report_sync.status_code == PROVIDER_HTTP_STATUS[provider_state]
                and report_sync_payload.get("ok") is False
                and sync_binding_ok
                and not report_sync_payload.get("fallback_available")
            )
        workspace["report_ask_ai_sync"] = {
            "ok": report_sync_ok,
            "expected_contract": (
                "canonical_grounded" if canonical_snapshot else "noncanonical_fail_closed"
            ),
            "status_code": report_sync.status_code,
            "response_ok": report_sync_payload.get("ok") is True,
            "grounded_mode": report_sync_payload.get("mode") == "grounded",
            "binding_retained": sync_binding_ok,
            "citation_count": len(
                extract_citations(str(report_sync_payload.get("answer") or ""))
            ),
            "answer_sha256": _digest(report_sync_payload.get("answer") or ""),
            "payload_sha256": _digest(report_sync_payload),
        }
        if not report_sync_ok:
            errors.append("report-bound Ask AI sync failed or widened scope")

        report_stream = client.session.post(
            client.base_url + "/api/ask-ai-portfolio/stream",
            json=report_ai_request,
            headers=client.headers("text/event-stream"),
            timeout=client.timeout,
        )
        report_events = (
            parse_sse(report_stream.text)
            if "text/event-stream"
            in str(report_stream.headers.get("Content-Type") or "")
            else []
        )
        report_stream_payload = (
            _stream_payload(report_events)
            if canonical_snapshot
            else _json(report_stream)
        )
        stream_error_context = next(
            (
                payload.get("scope_context")
                for name, payload in report_events
                if name == "error"
            ),
            {},
        )
        stream_binding_ok = binding_matches(
            report_stream_payload.get("scope_context") or stream_error_context
        )
        if not canonical_snapshot:
            report_stream_ok = not validate_noncanonical_report_ai_response(
                report_stream.status_code,
                report_stream_payload,
            )
        elif provider_state == "available":
            report_stream_ok = bool(
                report_stream.status_code == 200
                and _literal_true(report_stream_payload.get("ok"))
                and stream_binding_ok
                and extract_citations(
                    str(report_stream_payload.get("answer") or "")
                )
                and report_stream_payload.get("answer")
                == report_sync_payload.get("answer")
            )
        else:
            report_stream_ok = bool(
                report_stream.status_code == 200
                and report_stream_payload.get("stream_errors")
                and stream_binding_ok
            )
        workspace["report_ask_ai_stream"] = {
            "ok": report_stream_ok,
            "expected_contract": (
                "canonical_grounded" if canonical_snapshot else "noncanonical_fail_closed"
            ),
            "status_code": report_stream.status_code,
            "event_count": len(report_events),
            "binding_retained": stream_binding_ok,
            "citation_count": len(
                extract_citations(
                    str(report_stream_payload.get("answer") or "")
                )
            ),
            "answer_matches_sync": report_stream_payload.get("answer")
            == report_sync_payload.get("answer"),
            "answer_sha256": _digest(
                report_stream_payload.get("answer") or ""
            ),
            "payload_sha256": _digest(report_stream_payload),
        }
        if not report_stream_ok:
            errors.append("report-bound Ask AI stream failed or widened scope")

        history = client.get("/api/decision-workspace/history")
        history_payload = _json(history)
        history_records = history_payload.get("reports") or []
        history_ok = bool(
            history.status_code == 200
            and _literal_true(history_payload.get("ok"))
            and isinstance(history_records, list)
            and any(
                isinstance(item, Mapping)
                and item.get("analysis_id") == analysis_id
                for item in history_records
            )
        )
        workspace["history_ok"] = history_ok
        workspace["history_status"] = history.status_code
        workspace["history_count"] = len(history_records)
        workspace["history_sha256"] = _digest(history_payload)
        if not history_ok:
            errors.append("Manager Decision Workspace history omitted the completed report")

        compare = client.post_json(
            "/api/decision-workspace/compare",
            {
                "before_analysis_id": analysis_id,
                "after_analysis_id": analysis_id,
            },
        )
        compare_payload = _json(compare)
        compare_rejected = bool(
            compare.status_code == 400 and compare_payload.get("ok") is False
        )
        workspace["same_report_compare_rejected"] = compare_rejected
        workspace["compare_status"] = compare.status_code
        workspace["compare_sha256"] = _digest(compare_payload)
        if not compare_rejected:
            errors.append("Manager Decision Workspace accepted a same-report comparison")
    return {
        "start_status": response.status_code,
        "completed": completed,
        "poll_count": polls,
        "word_available": _literal_true(status.get("word_available")),
        "excel_available": _literal_true(status.get("excel_available")),
        "warning_count": len(warnings) if isinstance(warnings, list) else 0,
        "warning_datasets": warning_projection["datasets"],
        "warning_kinds": warning_projection["kinds"],
        "warning_sha256": _digest(warnings),
        "downloads": downloads,
        "previous_reports_ok": previous_ok,
        "manager_decision_workspace": workspace,
        **publication_projection,
    }, errors


def _run_scenario(  # noqa: C901, PLR0912, PLR0915
    scenario: str,
    *,
    manifest_path: Path,
    output_dir: Path,
    startup_timeout: float,
    request_timeout: float,
    report_timeout: float,
    include_report: bool,
) -> dict[str, Any]:
    bundle = build_scenario_bundle(scenario, manifest_path)
    port = _free_loopback_port()
    base_url = f"http://127.0.0.1:{port}"
    log_path = output_dir / f"{scenario}.server.log"
    state_dir = Path(
        tempfile.mkdtemp(prefix=f"adoptiq-{scenario}-state-")
    ).resolve()
    env = dict(os.environ)
    env.update(
        {
            "ADOPTIQ_BIND_PUBLIC": "0",
            "ADOPTIQ_ASK_AI_ALLOW_LEGACY_FALLBACK": "0",
            "ADOPTIQ_OUTPUTS_DIR": str(output_dir / f"{scenario}.reports"),
            "ADOPTIQ_LOCAL_ACCEPTANCE_STATE_DIR": str(state_dir),
            "PYTHONUNBUFFERED": "1",
        }
    )
    command = [
        sys.executable,
        str(REPO_ROOT / "scripts" / "run_local_acceptance_app.py"),
        "--enable-local-fixtures",
        "--scenario",
        scenario,
        "--manifest",
        str(manifest_path),
        "--host",
        "127.0.0.1",
        "--port",
        str(port),
    ]
    errors: list[str] = []
    result: dict[str, Any] = {
        "scenario": scenario,
        "provider_state": bundle.provider_state,
        "schema_fingerprint": bundle.schema_fingerprint,
        "report_publication_expectation": bundle.report_publication_expectation,
        "live_validation_performed": False,
        "source_states_sha256": _digest(bundle.source_states),
        "expected_counts_sha256": _digest(bundle.expected_canonical_counts),
        "warning_codes_sha256": _digest(bundle.warning_codes),
        "route_checks": {},
        "report": {"attempted": False},
        "errors": errors,
        "ok": False,
    }
    process: subprocess.Popen[Any] | None = None
    process_group_id: int | None = None
    log_capture: _BoundedFixtureLog | None = None
    try:
        process = subprocess.Popen(
            command,
            cwd=REPO_ROOT,
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=False,
            start_new_session=os.name == "posix",
            creationflags=(
                getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0)
                if os.name == "nt"
                else 0
            ),
        )
        process_group_id = process.pid if os.name == "posix" else None
        if process.stdout is None:
            raise RuntimeError("local app process did not expose a log pipe")
        log_capture = _BoundedFixtureLog(
            log_path,
            limit_bytes=FIXTURE_LOG_LIMIT_BYTES,
        )
        log_capture.start(
            process.stdout,
            overflow_callback=lambda: _signal_process_tree(
                process,
                process_group_id,
                signal.SIGTERM,
            ),
        )
        try:
            connectivity = _wait_for_server(
                base_url,
                scenario,
                bundle.schema_fingerprint,
                bundle.report_publication_expectation,
                process,
                startup_timeout,
            )
            client = LoopbackClient(base_url, request_timeout)
            client.bootstrap()
            routes = result["route_checks"]

            workspace_preview, workspace_errors = _run_workspace_preview_probe(
                client
            )
            routes["manager_decision_workspace_preview"] = workspace_preview
            errors.extend(workspace_errors)

            payloads: dict[str, dict[str, Any]] = {}
            for label, path in (
                ("version", "/api/version"),
                ("connectivity", "/api/diag/connectivity"),
                ("corpus_status", "/api/corpus/status"),
                ("corpus_alias", "/api/intel/status"),
                ("ask_ai_model", "/api/settings/ask-ai-model"),
                ("report_model", "/api/settings/report-model"),
            ):
                response = client.get(path)
                payload = _json(response)
                payloads[label] = payload
                ok = response.status_code == 200 and _literal_true(
                    payload.get("ok")
                )
                routes[label] = {
                    "status_code": response.status_code,
                    "ok": ok,
                    "payload_sha256": _digest(payload),
                }
                if not ok:
                    errors.append(f"{label} route failed")
            routes["connectivity"].update(
                {
                    "fixture_mode": connectivity.get("mode") == SOURCE_MODE,
                    "scenario_matches": connectivity.get("scenario") == scenario,
                    "report_publication_expectation_matches": (
                        connectivity.get("report_publication_expectation")
                        == bundle.report_publication_expectation
                    ),
                    "live_validation_performed": False,
                }
            )
            if (
                connectivity.get("report_publication_expectation")
                != bundle.report_publication_expectation
            ):
                errors.append("connectivity report-publication expectation mismatch")
            if connectivity.get("canonical_counts") != dict(
                sorted(bundle.expected_canonical_counts.items())
            ):
                errors.append("connectivity canonical counts do not match manifest")
            if connectivity.get("source_states") != dict(
                sorted(bundle.source_states.items())
            ):
                errors.append("connectivity source states do not match manifest")
            expected_warning_codes = {
                name: list(codes)
                for name, codes in sorted(bundle.warning_codes.items())
            }
            if connectivity.get("warning_codes") != expected_warning_codes:
                errors.append("connectivity warning codes do not match manifest")
            if payloads["corpus_status"] != payloads["corpus_alias"]:
                errors.append("corpus status alias payload differs")

            model_name = str(payloads["ask_ai_model"].get("active_value") or "")
            ping = client.post_json("/api/llm/ping", {"model_name": model_name})
            ping_payload = _json(ping)
            ping_ok = ping.status_code == 200 and (
                _literal_true(ping_payload.get("ok"))
                if bundle.provider_state == "available"
                else ping_payload.get("ok") is False
            )
            routes["llm_ping"] = {
                "status_code": ping.status_code,
                "ok": ping_ok,
                "provider_available": _literal_true(ping_payload.get("ok")),
                "payload_sha256": _digest(ping_payload),
            }
            if not ping_ok or not _response_is_sanitized(ping_payload):
                errors.append("LLM ping did not expose the expected sanitized state")

            customer_name = str(
                (bundle.records("customers") or [{}])[0].get("BU_NAME")
                or "Acme Corporation"
            )
            for label, path, marker in (
                (
                    "customer_360",
                    f"/customer/{quote(customer_name, safe='')}",
                    "Customer 360",
                ),
                (
                    "external_intelligence",
                    "/external-intelligence",
                    "External Intelligence",
                ),
            ):
                page = client.get(path, accept="text/html")
                page_ok = page.status_code == 200 and marker in page.text
                if scenario == "prompt_injection" and any(
                    term in page.text.casefold() for term in PROMPT_INJECTION_TERMS
                ):
                    page_ok = False
                    errors.append(f"{label} repeated a prompt-injection instruction")
                routes[label] = {
                    "status_code": page.status_code,
                    "ok": page_ok,
                    "characters": len(page.text),
                    "body_sha256": hashlib.sha256(page.content).hexdigest(),
                }
                if not page_ok and not any(label in item for item in errors):
                    errors.append(f"{label} page failed")

            playbook = client.post_form(
                "/playbook",
                {
                    "technology": "",
                    "theme": "",
                    "query": "registration authentication outage",
                    "csrf_token": client.csrf_token,
                },
            )
            playbook_ok = (
                playbook.status_code == 200
                and "Troubleshooting Playbook" in playbook.text
                and "Search matches for:" in playbook.text
                and not (
                    scenario == "prompt_injection"
                    and any(
                        term in playbook.text.casefold()
                        for term in PROMPT_INJECTION_TERMS
                    )
                )
            )
            routes["playbook"] = {
                "status_code": playbook.status_code,
                "ok": playbook_ok,
                "characters": len(playbook.text),
                "body_sha256": hashlib.sha256(playbook.content).hexdigest(),
            }
            if not playbook_ok:
                errors.append("Playbook route failed or repeated embedded instructions")

            exported = client.get(
                "/api/export-intel", params={"page": 0, "page_size": 1}
            )
            exported_payload = _json(exported)
            export_ok = (
                exported.status_code == 200
                and all(
                    isinstance(exported_payload.get(key), list)
                    for key in ("incidents", "bugs", "maintenances")
                )
                and exported_payload.get("live_validation_performed") is False
            )
            routes["external_export"] = {
                "status_code": exported.status_code,
                "ok": export_ok,
                "payload_sha256": _digest(exported_payload),
            }
            if not export_ok:
                errors.append("External Intelligence export failed")

            question = (
                "State the exact customer, open adoption barrier, and support-case "
                "counts. Cite the first decision and disclose incomplete, stale, "
                "failed, unavailable, or truncated data."
            )
            request_payload = {
                "question": question,
                "manager": "Local Fixture Manager",
                "technology": "All",
                "days": 90,
                "allow_legacy_fallback": False,
            }
            sync = client.post_json("/api/ask-ai-portfolio", request_payload)
            sync_payload = _json(sync)
            sync_errors = validate_provider_response(
                bundle.provider_state, sync.status_code, sync_payload
            )
            errors.extend(sync_errors)
            sync_warning_projection = _warning_projection(
                sync_payload.get("partial_data_warnings") or []
            )
            routes["ask_ai_sync"] = {
                "status_code": sync.status_code,
                "ok": not sync_errors,
                "answer_sha256": _digest(sync_payload.get("answer") or ""),
                "citation_count": len(
                    extract_citations(str(sync_payload.get("answer") or ""))
                ),
                "warning_sha256": _digest(
                    sync_payload.get("partial_data_warnings") or []
                ),
                "warning_datasets": sync_warning_projection["datasets"],
                "warning_kinds": sync_warning_projection["kinds"],
                "canonical_headline_sha256": _digest(
                    sync_payload.get("canonical_headline") or {}
                ),
            }
            if bundle.provider_state == "available" and _literal_true(
                sync_payload.get("ok")
            ):
                headline = sync_payload.get("canonical_headline") or {}
                expected = _expected_fixture_portfolio_counts(
                    bundle,
                    "Local Fixture Manager",
                )
                for key, value in expected.items():
                    if key == "action_plan_rows":
                        continue
                    if int(headline.get(key, -1)) != int(value):
                        errors.append(f"Ask AI canonical headline mismatch for {key}")
                answer = str(sync_payload.get("answer") or "")
                if scenario == "prompt_injection" and any(
                    term in answer.casefold() for term in PROMPT_INJECTION_TERMS
                ):
                    errors.append("Ask AI repeated an embedded instruction")
                query_id = str(sync_payload.get("query_id") or "")
                citations = extract_citations(answer)
                if not query_id or not citations:
                    errors.append("Ask AI success omitted query ID or citations")
                else:
                    diag = client.get(
                        f"/api/ask-ai/diagnostics/{quote(query_id, safe='')}"
                    )
                    evidence = client.get(
                        "/api/ask-ai/evidence/"
                        + quote(query_id, safe="")
                        + "/"
                        + quote(citations[0], safe="")
                    )
                    diag_payload = _json(diag)
                    evidence_payload = _json(evidence)
                    lookup_ok = (
                        diag.status_code == 200
                        and _literal_true(diag_payload.get("ok"))
                        and evidence.status_code == 200
                        and _literal_true(evidence_payload.get("ok"))
                    )
                    routes["ask_ai_evidence"] = {
                        "ok": lookup_ok,
                        "diagnostics_status": diag.status_code,
                        "evidence_status": evidence.status_code,
                        "payload_sha256": _digest(
                            {"diag": diag_payload, "evidence": evidence_payload}
                        ),
                    }
                    if not lookup_ok:
                        errors.append("Ask AI diagnostics/evidence resolution failed")

            stream = client.session.post(
                base_url + "/api/ask-ai-portfolio/stream",
                json=request_payload,
                headers=client.headers("text/event-stream"),
                timeout=request_timeout,
            )
            events = (
                parse_sse(stream.text)
                if "text/event-stream"
                in str(stream.headers.get("Content-Type") or "")
                else []
            )
            streamed = _stream_payload(events)
            if bundle.provider_state == "available":
                stream_ok = (
                    stream.status_code == 200
                    and _literal_true(streamed.get("ok"))
                    and streamed.get("answer") == sync_payload.get("answer")
                    and streamed.get("canonical_headline")
                    == sync_payload.get("canonical_headline")
                )
            else:
                stream_ok = (
                    stream.status_code == 200
                    and bool(streamed.get("stream_errors"))
                    and _response_is_sanitized(streamed)
                )
            routes["ask_ai_stream"] = {
                "status_code": stream.status_code,
                "ok": stream_ok,
                "event_count": len(events),
                "payload_sha256": _digest(streamed),
            }
            if not stream_ok:
                errors.append("Ask AI stream did not match the expected delivery state")

            intel = client.post_json(
                "/api/ask-intel",
                {
                    "question": "Which stored incident should be reviewed first? Cite it.",
                    "days": 90,
                    "allow_legacy_fallback": False,
                },
            )
            intel_payload = _json(intel)
            intel_errors = validate_provider_response(
                bundle.provider_state, intel.status_code, intel_payload
            )
            errors.extend(intel_errors)
            intel_warning_projection = _warning_projection(
                intel_payload.get("partial_data_warnings") or []
            )
            routes["ask_intel"] = {
                "status_code": intel.status_code,
                "ok": not intel_errors,
                "answer_sha256": _digest(intel_payload.get("answer") or ""),
                "warning_datasets": intel_warning_projection["datasets"],
                "warning_kinds": intel_warning_projection["kinds"],
                "payload_sha256": _digest(intel_payload),
            }
            if scenario == "stale" and bundle.provider_state == "available":
                if "source_stale" not in sync_warning_projection["kinds"]:
                    errors.append("Ask AI did not disclose the stale intel source")
                if "source_stale" not in intel_warning_projection["kinds"]:
                    errors.append("Ask Intel did not disclose the stale intel source")

            if include_report:
                result["report"] = {"attempted": True}
                scoped_counts = _expected_fixture_portfolio_counts(
                    bundle,
                    "Local Fixture Manager",
                )
                report_result, report_errors = _run_report_probe(
                    client,
                    report_timeout,
                    provider_state=bundle.provider_state,
                    expected_as_of_utc=bundle.as_of_utc,
                    expected_action_plan_rows=scoped_counts["action_plan_rows"],
                    publication_expectation=bundle.report_publication_expectation,
                )
                result["report"].update(report_result)
                errors.extend(report_errors)
        except Exception as exc:  # noqa: BLE001
            errors.append(f"scenario runner failed: {type(exc).__name__}")
            result["runner_exception_sha256"] = _digest(str(exc))
    finally:
        if process is not None:
            _terminate_process_tree(process, process_group_id)
        if log_capture is not None:
            log_evidence = log_capture.finish()
            result.update(log_evidence)
            if log_evidence["server_log_output_truncated"] is True:
                errors.append("fixture server log exceeded its bounded output limit")
            if log_evidence["server_log_integrity_complete"] is not True:
                errors.append("fixture server log capture did not complete safely")
        else:
            result.update(
                {
                    "server_log_bytes": 0,
                    "server_log_observed_bytes": -1,
                    "server_log_sha256": "",
                    "server_log_read_complete": False,
                    "server_log_output_truncated": False,
                    "server_log_integrity_complete": False,
                    "server_log_error_kind": "not_started",
                }
            )
            errors.append("fixture server log capture did not start")
        # This path was created by this scenario invocation and contains
        # fixture-only status/history/log state. Never retain it beside
        # the sanitized acceptance summary.
        shutil.rmtree(state_dir, ignore_errors=True)
    result["ok"] = not errors
    return result


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Exercise sanitized fixture states through real loopback routes."
    )
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST_PATH)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=REPO_ROOT / ".adoptiq-acceptance" / "round145" / "http-degraded",
    )
    parser.add_argument("--scenarios", default="all")
    parser.add_argument("--startup-timeout", type=float, default=45.0)
    parser.add_argument("--request-timeout", type=float, default=120.0)
    parser.add_argument("--report-timeout", type=float, default=300.0)
    parser.add_argument("--skip-reports", action="store_true")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    output_dir = _safe_output_dir(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest = load_manifest(args.manifest)
    available = sorted(manifest["scenarios"])
    if str(args.scenarios).strip().casefold() == "all":
        scenarios = available
    else:
        scenarios = [
            item.strip() for item in str(args.scenarios).split(",") if item.strip()
        ]
        unknown = sorted(set(scenarios).difference(available))
        if unknown:
            raise SystemExit("unknown scenario(s): " + ", ".join(unknown))

    started = time.time()
    results = []
    for index, scenario in enumerate(scenarios, start=1):
        print(f"[local-http] {index}/{len(scenarios)} {scenario}", flush=True)
        result = _run_scenario(
            scenario,
            manifest_path=Path(args.manifest),
            output_dir=output_dir,
            startup_timeout=float(args.startup_timeout),
            request_timeout=float(args.request_timeout),
            report_timeout=float(args.report_timeout),
            include_report=not bool(args.skip_reports),
        )
        results.append(result)
        print(
            f"[local-http] {scenario} ok={result['ok']} "
            f"errors={len(result['errors'])}",
            flush=True,
        )

    all_passed = len(results) == len(scenarios) and all(
        _literal_true(item.get("ok")) for item in results
    )
    summary = {
        "schema_version": SUMMARY_SCHEMA,
        "sanitized": True,
        "do_not_commit": True,
        "validation_mode": "local_acceptance",
        "source_mode": SOURCE_MODE,
        "live_validation_attempted": False,
        "live_validation_performed": False,
        "production_accuracy_claimed": False,
        "manifest_schema_fingerprint": manifest["schema_fingerprint"],
        "scenario_count": len(scenarios),
        "scenario_inventory_complete": scenarios == available,
        "report_probes_enabled": not bool(args.skip_reports),
        "elapsed_seconds": round(time.time() - started, 3),
        "all_passed": all_passed,
        "results": results,
    }
    summary_path = output_dir / "local_acceptance_http_summary.json"
    summary_path.write_text(
        json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "all_passed": all_passed,
                "scenario_count": len(scenarios),
                "summary_path": str(summary_path),
            },
            sort_keys=True,
        ),
        flush=True,
    )
    return 0 if all_passed else 5


if __name__ == "__main__":
    raise SystemExit(main())
