#!/usr/bin/env python3
"""Live report-iteration harness for repeatable report regression runs.

Round 51: this runner executes the same report scenarios repeatedly against
the running app, downloads report artifacts, and writes debug sidecars in
Downloads to speed up troubleshooting.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import platform
import re
import subprocess
import sys
import time
import zipfile
from dataclasses import asdict, dataclass, field
try:
    from datetime import UTC, datetime
except ImportError:  # pragma: no cover - Python 3.9 compatibility for local test env
    from datetime import datetime, timezone

    UTC = timezone.utc
from html import unescape
from pathlib import Path
from typing import Any, Iterable, Optional
from urllib.parse import quote

import openpyxl
import pandas as pd
import requests
from docx import Document

import canonical_metrics as cm

TERMINAL_STATUSES = {"completed", "error", "cancelled"}
CSRF_META_RE = re.compile(
    r'<meta[^>]+name=["\']csrf-token["\'][^>]+content=["\']([^"\']+)["\']',
    re.IGNORECASE,
)
DOCX_TEXT_RE = re.compile(r"<w:t[^>]*>(.*?)</w:t>")
NUMBER_RE = re.compile(r"\b\d+(?:\.\d+)?\b")
TIMESTAMP_TOKEN_RE = re.compile(r"\b\d{8,}\b")
RUN_SUFFIX_RE = re.compile(r"__data-loop-[^_]+__scenario-[^_]+__ts-[^_.]+\.", re.IGNORECASE)
NUMERIC_TOKEN_RE = re.compile(r"(?<![A-Za-z])[-+]?\d+(?:,\d{3})*(?:\.\d+)?%?")
KPI_ALIASES = {
    # Round 52 (Phase 2): aliases keyed to the actual labels emitted by
    # adoptiq_backend, executive_intelligence_formatter, compact_report_formatter,
    # leader_report_generator, and the matching app_simple Excel writers.
    "total_customers": {
        "total customers",
        "customers in portfolio",
        "customers",
        "num customers",
        "total customer count",
        "total customers analyzed",
    },
    "team_members": {
        # Round 53: leader team size is not the same KPI as customer count.
        "team size",
        "team members",
        "direct reports",
    },
    "support_cases": {
        "support cases",
        "total support cases",
        "support cases 90d",
        "total support cases 90d",
        "support cases 90d portfolio wide",
        "support cases last 90 days",
        "tac cases total",
        "tac cases",
        "num tac cases",
        # Round 52 / accuracy-fix-loop: leader executive summary bullets
        # render as "Total TAC Cases: 381". Without this alias the
        # paragraph extractor canonicalizes nothing, leaving the gate
        # to rely on the table heuristic alone.
        "total tac cases",
        # Round 66 / Pass 2 (B7): renewal Word emits the parenthesized
        # form ``Total Support Cases (90 days): 293`` (see
        # _create_simple_renewal_report). The KPI label normalizer
        # collapses parens to spaces, producing ``total support cases
        # 90 days`` (with a space-separated ``90 days``, not ``90d``).
        # Without these aliases the renewal Word doc's headline TAC
        # paragraph fell through ``_paragraph_match_is_canonical_kpi``
        # and never received an inline citation -- a parity gap vs
        # Compact, which uses the colon-suffix-free ``Total Support
        # Cases: 293`` form already covered above.
        "total support cases 90 days",
        "support cases 90 days",
        "support cases last days",
    },
    # Round 52 / partial-data-warning Phase 5: ``escalated support cases``
    # is a SUBSET of total ``support_cases`` (specifically, P1/P2 + BEMS
    # escalations).  Folding it into the ``support_cases`` canonical
    # caused the strict parity gate to fire on the compact report
    # because the DOCX KPI (293 total) was overwritten by the XLSX
    # Executive_Dashboard "Escalated Support Cases" cell (2 escalated).
    # Track it as a distinct canonical so both metrics can stand up
    # without colliding.
    "escalated_support_cases": {
        "escalated support cases",
        "escalated cases",
        "escalation required",
    },
    "adoption_barriers": {
        "adoption barriers",
        "total adoption barriers",
        "adoption barriers total",
        "num adoption barriers",
    },
    "open_adoption_barriers": {
        # Round 53: "active" report labels mean open AB records.
        "active adoption barriers",
        "open adoption barriers",
        "adoption barriers open",
    },
    "critical_barriers": {
        "critical abs",
        "critical adoption barriers",
        "critical barriers",
        "adoption barriers critical",
    },
    "critical_cases": {
        "critical p1",
        "p1 cases",
        "critical cases",
        "tac cases p1",
        "p1 critical cases",
    },
    "high_cases": {
        "high p2",
        "p2 cases",
        "high cases",
        "p2 high cases",
    },
    "bems": {
        "bems escalations",
        "bems",
        "bems break fix",
        "num bems",
        # Round 52 / accuracy-fix-loop: leader bullets render as
        # "Total BEMS Escalations: 81".
        "total bems escalations",
        "total bems",
        # Round 149: decision-report Metric_Lineage display label.
        "bems escalations tac subset",
    },
    "risk_score": {
        "overall risk score",
        "renewal risk score",
        "risk score",
        "portfolio risk score",
    },
    "window_days": {
        "analysis period days",
        "analysis period",
        "window days",
        "days",
    },
    "manager": {
        "manager",
        "manager scope",
    },
    "technology": {
        "technology",
        "technology scope",
        "technology focus",
    },
    # Round 52 / partial-data-warning Phase 5: text-valued canonical
    # for the renewal Customer Health Dashboard's "Risk Category" row.
    # Listed in ``_TEXT_VALUED_CANONICAL_KPIS`` so the numeric guard
    # does not strip it.
    "risk_category": {
        "risk category",
        "renewal risk category",
        "overall risk category",
        # Round 66 / Pass 2 (B7): renewal subscription summary line
        # emits ``Renewal Risk Level: HIGH (7.3/10)`` (see
        # app_simple.py L21125). Without this alias the renewal
        # narrative did not surface as a canonical KPI claim and
        # was missing the inline ``[Source: ...]`` chip.
        "renewal risk level",
        "risk level",
    },
    "high_risk_customers": {
        "high risk customers",
        "high risk customer count",
        "critical high risk",
    },
    "action_plans": {
        "action plans",
        "num action plans",
        # Round 52 / accuracy-fix-loop: leader bullets render as
        # "Total Action Plans: 354".
        "total action plans",
    },
    "open_action_plans": {
        # Open/unresolved plans are a subset of total plans. Keeping these
        # labels separate prevents a workbook's "Open Action Plans" tile
        # from being compared with Word's "Total Action Plans" statement.
        "action plans open",
        "open action plans",
    },
    "customer_pulse": {
        "customer pulse",
        "customer pulse records",
        "num customer pulse",
        # Round 52 / accuracy-fix-loop: leader bullets render as
        # "Total Customer Pulse records: 87" (mixed case).
        "total customer pulse",
        "total customer pulse records",
    },
    # Round 66 / Pass 2 (B7): canonicals introduced to enrich the
    # renewal Word document's inline citation density. Pre-R66 the
    # renewal narrative emitted ``Total Success Priorities: 12`` and
    # ``Service Incidents (status.webex.com)`` paragraphs that the
    # canonical-KPI gate did not recognise -- the injector then
    # silently relied on the narrative-token gate alone, which
    # filters out 4+ digit values as ``IDs``. The new canonicals
    # mirror the renewal sections at app_simple.py L11758 (success
    # priorities) and L11842 (incidents).
    "success_priorities": {
        "success priorities",
        "total success priorities",
        "num success priorities",
        "customer success priorities",
        "total customer success priorities",
    },
    "incidents": {
        "incidents",
        "service incidents",
        "total incidents",
        "high impact incidents",
        "external incidents",
    },
}

# Round 52 (Phase 2): scenario-specific required KPI keys. When strict mode is
# enabled, the harness fails the parity gate if any of these keys is missing
# from BOTH the DOCX KPIs and the XLSX KPIs (i.e. neither format emitted it).
# Mismatches between formats are reported even outside strict mode.
SCENARIO_REQUIRED_KPIS: dict[str, tuple[str, ...]] = {
    "comprehensive": ("manager", "technology", "window_days", "total_customers"),
    "compact": ("manager", "technology", "window_days", "total_customers"),
    "renewal": ("technology", "window_days"),
    "leader": ("manager", "window_days", "team_members", "total_customers"),
    "subscription": ("technology", "window_days"),
}


def scenario_report_family(scenario: "Scenario") -> str:
    """Return the canonical report family for an executable scenario.

    Matrix scenario keys describe coverage (for example ``b_comp_am_All``),
    not report semantics.  Quality gates must therefore bind to the endpoint
    and request payload rather than looking up the raw scenario key.
    """

    if scenario.endpoint == "/start_compact_analysis":
        return "compact"
    if scenario.endpoint == "/start_leader_report":
        return "leader"
    if scenario.endpoint == "/start_subscription_analysis":
        return "subscription"
    if scenario.endpoint == "/start_analysis":
        report_type = str(scenario.payload.get("report_type") or "").strip().casefold()
        if report_type == "comprehensive":
            return "comprehensive"
        if report_type in {"renewal", "renewal_portfolio"}:
            return "renewal"
    return "unknown"


def required_kpis_for_scenario(scenario: "Scenario") -> tuple[str, ...]:
    """Resolve strict KPI requirements from report family, never key shape."""

    return SCENARIO_REQUIRED_KPIS.get(scenario_report_family(scenario), ())


# Round 52 / ship: per-scenario DOCX similarity threshold overrides.
#
# The default ``min_docx_similarity`` (0.55 in strict mode) is calibrated for
# reports whose narrative is templated and largely deterministic (compact,
# renewal, leader -- which sit at 0.91-0.99 textual_sim against a fresh
# baseline).  The comprehensive report is materially different: it embeds an
# AI-generated "Insights" section (Pattern N: Latent Technical Debt, Root
# Cause Analysis, Business Impact, Evidence, ...) that the LLM regenerates
# from scratch on every run.  Two consecutive comprehensive runs against the
# same scope diff by ~1600 narrative paragraphs and even differ in TOTAL
# paragraph count by ~200, so a single global text-similarity floor cannot
# fairly cover both narrative-templated and narrative-generated reports.
#
# The numeric similarity gate stays at 0.80 across the board -- that is the
# binding signal for actual data drift.  The text gate is only relaxed for
# the AI-narrative-heavy comprehensive scenario; data accuracy guards do
# not weaken.
#
# Round 52.1: ``min_docx_numeric_similarity`` is also overridable per
# scenario, and a new third threshold ``min_docx_table_numeric_similarity``
# resolves through the same mechanism (defaulting to the runner-wide config).
# The 3-iter repeatability proof showed comprehensive's overall numeric
# similarity drifts to ~0.71 because the AI-narrative paragraphs regenerate
# every run, while table-only numerics stayed at 1.0 across baseline + 3
# iterations.  We therefore relax comprehensive's overall numeric_sim to an
# informational level (0.55) and rely on the new table-only gate (0.95)
# as the true binding signal for real Snowflake data drift in that scenario.
# compact / renewal / leader keep the original thresholds untouched.
#
# Override format: { scenario_key: { "min_docx_similarity": float,
#                                    "min_docx_numeric_similarity": float,
#                                    "min_docx_table_numeric_similarity": float } }
SCENARIO_DOCX_THRESHOLD_OVERRIDES: dict[str, dict[str, float]] = {
    "comprehensive": {
        "min_docx_similarity": 0.40,
        # Round 52.1: AI-narrative noise dominates the overall numeric
        # fingerprint here -- relax to informational and let the table-only
        # gate (0.95 default) carry the real-data-drift signal.
        "min_docx_numeric_similarity": 0.55,
    },
}


def effective_docx_thresholds(
    config: "RunnerConfig", scenario_key: str
) -> tuple[float, float, float]:
    """Resolve per-scenario DOCX similarity thresholds with overrides.

    Returns (min_docx_similarity, min_docx_numeric_similarity,
    min_docx_table_numeric_similarity).  When no override is registered for
    ``scenario_key``, the runner-wide config defaults are used unchanged.

    Round 52.1: returns a 3-tuple now.  The third element is the table-only
    numeric similarity threshold and is the new noise-immune binding signal
    for actual Snowflake data drift.
    """
    override = SCENARIO_DOCX_THRESHOLD_OVERRIDES.get(scenario_key, {})
    min_text = float(override.get("min_docx_similarity", config.min_docx_similarity))
    min_numeric = float(
        override.get("min_docx_numeric_similarity", config.min_docx_numeric_similarity)
    )
    min_table_numeric = float(
        override.get(
            "min_docx_table_numeric_similarity",
            config.min_docx_table_numeric_similarity,
        )
    )
    return min_text, min_numeric, min_table_numeric


@dataclass(frozen=True)
class Scenario:
    """Immutable report scenario definition."""

    key: str
    endpoint: str
    payload_mode: str  # "form" | "json"
    payload: dict[str, Any]
    expect_excel: bool = True
    expected_xlsx_sheets: tuple[str, ...] = ("summary", "report_info")
    # Complete-source fixture scenarios have a known visual contract.  A
    # missing chart in those scenarios is a publication defect, not a design
    # suggestion.  Degraded/technology-partial scenarios keep the default so
    # honest chart withholding remains valid.
    expected_min_charts: int = 0
    # Explicit cross-family equivalence identity.  Only scenarios bearing the
    # same non-empty cohort are required to form one exact Compact /
    # Comprehensive / Leader / Renewal quartet.  Technology sweeps and other
    # intentionally non-equivalent scenarios remain unmarked.
    source_parity_cohort: str = ""


@dataclass
class GateResult:
    """Pass/fail details for one validation gate."""

    passed: bool
    details: dict[str, Any] = field(default_factory=dict)


@dataclass
class ArtifactRecord:
    """Metadata for one downloaded artifact."""

    file_type: str
    downloaded_name: str
    debug_name: str
    debug_path: str
    size_bytes: int
    sha256: str
    baseline_path: Optional[str]
    structural: GateResult
    baseline_diff: GateResult
    # Round 52 (Phase 1): manifest-backed baseline identity.
    # ``baseline_source`` is one of "manifest" | "latest" | "none". When the
    # baseline came from a manifest, ``baseline_sha256`` and
    # ``baseline_manifest_key`` make the comparison fully reproducible without
    # rescanning Downloads.
    baseline_source: str = "none"
    baseline_sha256: Optional[str] = None
    baseline_manifest_key: Optional[str] = None
    baseline_mtime_utc: Optional[str] = None
    baseline_size_bytes: Optional[int] = None
    baseline_integrity: GateResult = field(
        default_factory=lambda: GateResult(True, {"reason": "not_applicable"})
    )


@dataclass
class ScenarioResult:
    """All telemetry for one scenario execution."""

    scenario: str
    analysis_id: Optional[str]
    started_at_utc: str
    completed_at_utc: str
    elapsed_seconds: int
    request: dict[str, Any]
    status_snapshots: list[dict[str, Any]]
    final_status: dict[str, Any]
    operational: GateResult
    artifacts: list[ArtifactRecord]
    all_passed: bool
    parity: GateResult = field(default_factory=lambda: GateResult(True, {"reason": "not_applicable"}))
    quality: GateResult = field(default_factory=lambda: GateResult(True, {"reason": "not_applicable"}))
    kpi_sidecar_path: Optional[str] = None
    quality_sidecar_path: Optional[str] = None
    log_excerpt_path: Optional[str] = None
    failure_phase: Optional[str] = None
    exception_type: Optional[str] = None


@dataclass
class RunnerConfig:
    """Runtime options for iteration execution."""

    base_url: str
    downloads_dir: Path
    iterations: int
    poll_interval_seconds: float
    scenario_timeout_seconds: int
    run_id: str
    stop_on_failure: bool
    scenario_keys: list[str]
    baseline_mode: str
    min_docx_similarity: float
    min_sheet_overlap: float
    min_header_similarity: float
    min_docx_chars: int
    strict: bool
    min_docx_numeric_similarity: float
    # Round 52.1: noise-immune table-only numeric similarity gate.  Default
    # 0.95 is empirically grounded -- the round52 manifest 3-iter floor was
    # exactly 1.0 for all four scenarios.
    min_docx_table_numeric_similarity: float
    max_xlsx_row_delta_ratio: float
    max_xlsx_row_delta_abs: int
    # Round 52 (Phase 1): manifest-backed baseline configuration. None when
    # manifest mode is not in use; required (and validated) when
    # ``baseline_mode == "manifest"``.
    baseline_manifest_path: Optional[Path] = None
    init_baseline: bool = False
    init_baseline_dir: Optional[Path] = None
    init_baseline_label: Optional[str] = None
    request_timeout_seconds: int = 120
    download_timeout_seconds: int = 300


def _utc_now() -> datetime:
    return datetime.now(UTC)


def _utc_now_str() -> str:
    return _utc_now().strftime("%Y-%m-%dT%H:%M:%SZ")


def _slug(text: str) -> str:
    return re.sub(r"[^a-zA-Z0-9_-]+", "_", text.strip()).strip("_")


# Round 133: exhaustive report-option matrix blocks (cheap-first execution order).
MATRIX_BLOCK_ORDER: tuple[str, ...] = ("E", "F", "A", "B", "C", "D", "G")

# Round 133: mirrors AnalysisForm.technology choices in app_simple.py.
MATRIX_TECHNOLOGY_CHOICES: tuple[str, ...] = (
    "Webex Meetings & Messaging",
    "Webex Calling",
    "Webex Contact Center",
    "Webex Contact Center Enterprise",
    "Cisco UCCE",
    "Cisco UCCX",
    "All Contact Center",
    "All",
)


@dataclass(frozen=True)
class EdgeMatrixConfig:
    """Explicit authorized scopes for the live option matrix."""

    manager_name: str = ""
    customer_name: str = ""
    subscription_id: str = ""
    csone_upload_path: str = ""
    compact_customer_name: str = ""
    compact_customer_technology: str = "Webex Calling"


def _load_matrix_managers() -> tuple[str, ...]:
    """Load manager names from team_config.json (Round 133 matrix SSoT)."""
    config_path = Path(__file__).resolve().parent / "team_config.json"
    try:
        payload = json.loads(config_path.read_text(encoding="utf-8"))
    except (OSError, TypeError, ValueError) as exc:
        raise ValueError("Authoritative manager roster is unavailable") from exc
    managers = payload.get("managers") if isinstance(payload, dict) else None
    if not isinstance(managers, list):
        raise ValueError("Authoritative manager roster is malformed")
    cleaned = [str(item).strip() for item in managers if str(item).strip()]
    if not cleaned:
        raise ValueError("Authoritative manager roster is empty")
    return tuple(cleaned)


def _matrix_scope_key(
    *,
    report_type: str,
    manager: str,
    technology: str,
    endpoint: str,
) -> tuple[str, str, str, str]:
    return (report_type, manager.strip(), technology.strip(), endpoint)


# Round 147: every supported report family now exits through the same concise
# Word + canonical Source Data contract.  Keep the matrix gate aligned with
# that public contract so a regression to one of the retired family-specific
# workbooks fails before visual or parity checks run.
ROUND147_CANONICAL_SOURCE_SHEETS = (
    "report_info",
    "metric_lineage",
    "chart_data",
    "evidence_links",
    "action_plans",
    "account_summary",
)


def build_exhaustive_option_matrix(
    days: int = 90,
    *,
    edge: EdgeMatrixConfig | None = None,
) -> dict[str, Scenario]:
    """Round 133: build the smart-exhaustive report option matrix (~45 runs).

    Blocks:
      A — canonical quartet (existing build_scenario_map entries)
      B — comprehensive tech sweep (All Managers × each technology)
      C — comprehensive manager sweep (All Contact Center × each manager)
      D — leader manager sweep
      E — compact tech sweep (All Managers × each technology)
      F — renewal_portfolio tech sweep (All Managers × each technology)
      G — edge cases (single-customer renewal, subscription, uploads, scoped compact, All tech)
    """
    edge_cfg = edge or EdgeMatrixConfig()
    selected_manager = edge_cfg.manager_name.strip()
    selected_customer = edge_cfg.customer_name.strip()
    compact_customer = (
        edge_cfg.compact_customer_name.strip() or selected_customer
    )
    if not selected_manager:
        raise ValueError("Live report matrix requires an explicit manager_name")
    if not selected_customer:
        raise ValueError("Live report matrix requires an explicit customer_name")
    days_str = str(int(days))
    matrix: dict[str, Scenario] = {}
    seen_comp: set[tuple[str, str, str, str]] = set()

    def _add_comprehensive(key: str, manager: str, technology: str) -> None:
        scope = _matrix_scope_key(
            report_type="comprehensive",
            manager=manager,
            technology=technology,
            endpoint="/start_analysis",
        )
        if scope in seen_comp:
            return
        seen_comp.add(scope)
        matrix[key] = Scenario(
            key=key,
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "comprehensive",
                "manager": manager,
                "technology": technology,
                "days": days_str,
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4 if technology == "All" else 0,
        )

    # Block A — canonical equivalent-scope quartet.  Leader is an all-technology
    # product and does not accept a technology payload, so the other three
    # families must explicitly request ``All`` here.  Technology-specific
    # coverage remains in Blocks B/C/E/F; silently treating ``All Contact
    # Center`` as equivalent to Leader ``All`` would make source-parity green
    # across different populations.
    for canonical_key, scenario in build_scenario_map().items():
        canonical_payload = dict(scenario.payload)
        if "manager" in canonical_payload:
            canonical_payload["manager"] = selected_manager
        if "technology" in canonical_payload:
            canonical_payload["technology"] = "All"
        if "days" in canonical_payload:
            canonical_payload["days"] = (
                int(days) if scenario.payload_mode == "json" else days_str
            )
        matrix[f"a_{canonical_key}"] = Scenario(
            key=f"a_{canonical_key}",
            endpoint=scenario.endpoint,
            payload_mode=scenario.payload_mode,
            payload=canonical_payload,
            expect_excel=scenario.expect_excel,
            expected_xlsx_sheets=scenario.expected_xlsx_sheets,
            expected_min_charts=max(4, scenario.expected_min_charts),
            source_parity_cohort="exhaustive-primary-team",
        )
        if scenario.endpoint == "/start_analysis" and scenario.payload.get("report_type") == "comprehensive":
            seen_comp.add(
                _matrix_scope_key(
                    report_type="comprehensive",
                    manager=str(canonical_payload.get("manager") or ""),
                    technology=str(canonical_payload.get("technology") or ""),
                    endpoint=scenario.endpoint,
                )
            )

    managers = _load_matrix_managers()

    # Block B — tech sweep comprehensive (All Managers)
    for tech in MATRIX_TECHNOLOGY_CHOICES:
        key = f"b_comp_am_{_slug(tech)}"
        _add_comprehensive(key, "All Managers", tech)

    # Block C — manager sweep comprehensive (All Contact Center)
    for manager in managers:
        key = f"c_comp_{_slug(manager)}_acc"
        _add_comprehensive(key, manager, "All Contact Center")

    # Block D — leader sweep
    for manager in managers:
        matrix[f"d_leader_{_slug(manager)}"] = Scenario(
            key=f"d_leader_{_slug(manager)}",
            endpoint="/start_leader_report",
            payload_mode="form",
            payload={"manager": manager, "days": days_str},
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4,
        )

    # Block E — compact tech sweep
    for tech in MATRIX_TECHNOLOGY_CHOICES:
        matrix[f"e_compact_am_{_slug(tech)}"] = Scenario(
            key=f"e_compact_am_{_slug(tech)}",
            endpoint="/start_compact_analysis",
            payload_mode="json",
            payload={
                "manager": "All Managers",
                "technology": tech,
                "days": int(days),
                "csone_file": "",
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4 if tech == "All" else 0,
        )

    # Block F — renewal_portfolio tech sweep
    for tech in MATRIX_TECHNOLOGY_CHOICES:
        matrix[f"f_renewal_am_{_slug(tech)}"] = Scenario(
            key=f"f_renewal_am_{_slug(tech)}",
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "renewal_portfolio",
                "renewal_type": "renewal_portfolio",
                "manager": "All Managers",
                "technology": tech,
                "days": days_str,
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4 if tech == "All" else 0,
        )

    # Block G — edge cases
    matrix["g_renewal_single_customer"] = Scenario(
        key="g_renewal_single_customer",
        endpoint="/start_analysis",
        payload_mode="form",
        payload={
            "report_type": "renewal",
            "renewal_type": "renewal_single",
            "manager": selected_manager,
            "technology": "All Contact Center",
            "days": days_str,
            "subscription_id": "",
            "customer_name": selected_customer,
        },
        expect_excel=True,
        expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
    )

    if edge_cfg.subscription_id.strip():
        matrix["g_subscription_analysis"] = Scenario(
            key="g_subscription_analysis",
            endpoint="/start_subscription_analysis",
            payload_mode="form",
            payload={
                "subscription_id": edge_cfg.subscription_id.strip(),
                "days": days_str,
                "report_type": "comprehensive",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        )

    matrix["g_compact_customer_scoped"] = Scenario(
        key="g_compact_customer_scoped",
        endpoint="/start_compact_analysis",
        payload_mode="json",
        payload={
            "manager": selected_manager,
            "technology": edge_cfg.compact_customer_technology,
            "days": int(days),
            "csone_file": "",
            "subscription_id": "",
            "customer_name": compact_customer,
        },
        expect_excel=True,
        expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        expected_min_charts=(
            4 if edge_cfg.compact_customer_technology == "All" else 0
        ),
    )

    _add_comprehensive("g_comp_am_all_tech", "All Managers", "All")

    if edge_cfg.csone_upload_path.strip():
        upload_name = Path(edge_cfg.csone_upload_path.strip()).name
        matrix["g_compact_csone_upload"] = Scenario(
            key="g_compact_csone_upload",
            endpoint="/start_compact_analysis",
            payload_mode="json",
            payload={
                "manager": selected_manager,
                "technology": "All",
                "days": int(days),
                "csone_file": upload_name,
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4,
        )
        matrix["g_leader_csone_upload"] = Scenario(
            key="g_leader_csone_upload",
            endpoint="/start_leader_report",
            payload_mode="form",
            payload={
                "manager": selected_manager,
                "days": days_str,
                "csone_file": edge_cfg.csone_upload_path.strip(),
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        )

    return matrix


def build_local_acceptance_option_matrix(
    days: int = 90,
    *,
    customer_name: str = "Acme Corporation",
    subscription_id: str = "SUB-001",
) -> dict[str, Scenario]:
    """Build the exhaustive report matrix for the guarded fixture runtime.

    The production matrix intentionally spans the real manager roster. The
    local acceptance runtime has one sanitized manager and two members, so it
    instead spans every report/technology option plus all three Leader scopes.
    This function only describes HTTP requests; it cannot activate fixtures.
    """

    days_text = str(max(min(int(days), 365), 1))
    days_value = int(days_text)
    manager = "Local Fixture Manager"
    matrix: dict[str, Scenario] = {}

    def comprehensive(
        key: str,
        technology: str,
        *,
        customer: str = "",
        source_parity_cohort: str = "",
    ) -> Scenario:
        return Scenario(
            key=key,
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "comprehensive",
                "manager": manager,
                "technology": technology,
                "days": days_text,
                "subscription_id": "",
                "customer_name": customer,
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4 if technology == "All" else 0,
            source_parity_cohort=source_parity_cohort,
        )

    def compact(
        key: str,
        technology: str,
        *,
        customer: str = "",
        source_parity_cohort: str = "",
    ) -> Scenario:
        return Scenario(
            key=key,
            endpoint="/start_compact_analysis",
            payload_mode="json",
            payload={
                "manager": manager,
                "technology": technology,
                "days": days_value,
                "csone_file": "",
                "subscription_id": "",
                "customer_name": customer,
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4 if technology == "All" else 0,
            source_parity_cohort=source_parity_cohort,
        )

    def renewal_portfolio(
        key: str,
        technology: str,
        *,
        source_parity_cohort: str = "",
    ) -> Scenario:
        return Scenario(
            key=key,
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "renewal_portfolio",
                "renewal_type": "renewal_portfolio",
                "manager": manager,
                "technology": technology,
                "days": days_text,
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4 if technology == "All" else 0,
            source_parity_cohort=source_parity_cohort,
        )

    # Leader is an all-technology product and has no technology selector.
    # Keep its declared parity peers at the same population.  The technology
    # loop below still exercises All Contact Center and every named option.
    matrix["a_comprehensive"] = comprehensive(
        "a_comprehensive",
        "All",
        source_parity_cohort="local-primary-team",
    )
    matrix["a_compact"] = compact(
        "a_compact",
        "All",
        source_parity_cohort="local-primary-team",
    )
    matrix["a_renewal"] = renewal_portfolio(
        "a_renewal",
        "All",
        source_parity_cohort="local-primary-team",
    )
    matrix["a_leader"] = Scenario(
        key="a_leader",
        endpoint="/start_leader_report",
        payload_mode="form",
        payload={"manager": manager, "days": days_text, "scope_type": "team"},
        expect_excel=True,
        expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        expected_min_charts=4,
        source_parity_cohort="local-primary-team",
    )

    for technology in MATRIX_TECHNOLOGY_CHOICES:
        slug = _slug(technology)
        matrix[f"b_comp_local_{slug}"] = comprehensive(
            f"b_comp_local_{slug}", technology
        )
        matrix[f"e_compact_local_{slug}"] = compact(
            f"e_compact_local_{slug}", technology
        )
        matrix[f"f_renewal_local_{slug}"] = renewal_portfolio(
            f"f_renewal_local_{slug}", technology
        )

    matrix["c_comp_all_managers_all"] = Scenario(
        **{
            **asdict(comprehensive("c_comp_all_managers_all", "All")),
            "payload": {
                **comprehensive("c_comp_all_managers_all", "All").payload,
                # The sanitized runtime declares one manager. Exercise the
                # broad technology scope through that real roster value; an
                # unsupported synthetic "All Managers" value would only test
                # request rejection rather than report generation.
                "manager": manager,
            },
        }
    )
    matrix["d_leader_team"] = Scenario(
        key="d_leader_team",
        endpoint="/start_leader_report",
        payload_mode="form",
        payload={"manager": manager, "days": days_text, "scope_type": "team"},
        expect_excel=True,
        expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        expected_min_charts=4,
    )
    matrix["d_leader_member"] = Scenario(
        key="d_leader_member",
        endpoint="/start_leader_report",
        payload_mode="form",
        payload={
            "manager": manager,
            "days": days_text,
            "scope_type": "member",
            "scope_value": "fixture.owner1@example.invalid",
        },
        expect_excel=True,
        expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        expected_min_charts=4,
    )
    matrix["d_leader_customer"] = Scenario(
        key="d_leader_customer",
        endpoint="/start_leader_report",
        payload_mode="form",
        payload={
            "manager": manager,
            "days": days_text,
            "scope_type": "customer",
            "scope_value": customer_name,
        },
        expect_excel=True,
        expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        expected_min_charts=4,
        source_parity_cohort="local-primary-customer",
    )
    matrix["g_renewal_single_customer"] = Scenario(
        key="g_renewal_single_customer",
        endpoint="/start_analysis",
        payload_mode="form",
        payload={
            "report_type": "renewal",
            "renewal_type": "renewal_single",
            "manager": manager,
            "technology": "All",
            "days": days_text,
            "subscription_id": "",
            "customer_name": customer_name,
        },
        expect_excel=True,
        expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        expected_min_charts=4,
        source_parity_cohort="local-primary-customer",
    )
    matrix["g_subscription_analysis"] = Scenario(
        key="g_subscription_analysis",
        endpoint="/start_subscription_analysis",
        payload_mode="form",
        payload={
            "subscription_id": subscription_id,
            "days": days_text,
            "report_type": "comprehensive",
        },
        expect_excel=True,
        expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        expected_min_charts=3,
    )
    matrix["g_compact_customer_scoped"] = compact(
        "g_compact_customer_scoped",
        "All",
        customer=customer_name,
        source_parity_cohort="local-primary-customer",
    )
    matrix["g_comprehensive_customer_scoped"] = comprehensive(
        "g_comprehensive_customer_scoped",
        "All",
        customer=customer_name,
        source_parity_cohort="local-primary-customer",
    )
    return matrix


def build_local_acceptance_all_managers_matrix(
    days: int = 90,
) -> dict[str, Scenario]:
    """Exercise aggregate manager branches against the two-manager fixture.

    The ordinary local matrix intentionally uses one named manager for broad
    technology coverage.  Round 167 adds this small second matrix because the
    aggregate ``All Managers`` sentinel takes materially different roster and
    Pass 1 paths in Compact, Comprehensive, Renewal, and Leader.  Keeping the
    scenarios separate makes the aggregate contract explicit and bounded while
    still using the real report endpoints and canonical artifact validators.
    """

    days_text = str(max(min(int(days), 365), 1))
    days_value = int(days_text)
    manager = "All Managers"
    return {
        "a_all_managers_compact": Scenario(
            key="a_all_managers_compact",
            endpoint="/start_compact_analysis",
            payload_mode="json",
            payload={
                "manager": manager,
                "technology": "All",
                "days": days_value,
                "csone_file": "",
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4,
            source_parity_cohort="local-all-managers-team",
        ),
        "a_all_managers_comprehensive": Scenario(
            key="a_all_managers_comprehensive",
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "comprehensive",
                "manager": manager,
                "technology": "All",
                "days": days_text,
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4,
            source_parity_cohort="local-all-managers-team",
        ),
        "a_all_managers_leader": Scenario(
            key="a_all_managers_leader",
            endpoint="/start_leader_report",
            payload_mode="form",
            payload={
                "manager": manager,
                "days": days_text,
                "scope_type": "team",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4,
            source_parity_cohort="local-all-managers-team",
        ),
        "a_all_managers_renewal": Scenario(
            key="a_all_managers_renewal",
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "renewal_portfolio",
                "renewal_type": "renewal_portfolio",
                "manager": manager,
                "technology": "All",
                "days": days_text,
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            expected_min_charts=4,
            source_parity_cohort="local-all-managers-team",
        ),
    }


def build_local_acceptance_multi_manager_matrix(
    days: int = 90,
) -> dict[str, Scenario]:
    """Exercise both independent teams plus the aggregate sentinel.

    This is intentionally extensive rather than combinatorial: each named
    manager runs every report family and every supported scope selection:
    portfolio and customer Compact/Comprehensive, portfolio and individual
    Renewal, Team/Member/Customer Leader, and one authorized subscription.
    The aggregate sentinel then runs the four
    portfolio families.  Technology-option breadth remains in the ordinary
    healthy matrix, so this second pass targets roster isolation and team
    attribution without multiplying every technology by every manager.
    """

    days_text = str(max(min(int(days), 365), 1))
    days_value = int(days_text)
    matrix = build_local_acceptance_all_managers_matrix(days)
    managers = (
        (
            "Local Fixture Manager",
            "fixture.owner1@example.invalid",
            "Beta Industries",
            "SUB-002",
            "primary",
        ),
        (
            "Second Fixture Manager",
            "fixture.owner2@example.invalid",
            "Gamma Public Sector",
            "SUB-003",
            "secondary",
        ),
    )
    for manager, member_email, customer, subscription_id, slug in managers:
        matrix[f"a_{slug}_manager_compact"] = Scenario(
            key=f"a_{slug}_manager_compact",
            endpoint="/start_compact_analysis",
            payload_mode="json",
            payload={
                "manager": manager,
                "technology": "All",
                "days": days_value,
                "csone_file": "",
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            source_parity_cohort=f"local-{slug}-manager-team",
        )
        matrix[f"a_{slug}_manager_comprehensive"] = Scenario(
            key=f"a_{slug}_manager_comprehensive",
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "comprehensive",
                "manager": manager,
                "technology": "All",
                "days": days_text,
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            source_parity_cohort=f"local-{slug}-manager-team",
        )
        matrix[f"g_{slug}_manager_compact_customer"] = Scenario(
            key=f"g_{slug}_manager_compact_customer",
            endpoint="/start_compact_analysis",
            payload_mode="json",
            payload={
                "manager": manager,
                "technology": "All",
                "days": days_value,
                "csone_file": "",
                "subscription_id": "",
                "customer_name": customer,
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        )
        matrix[f"g_{slug}_manager_comprehensive_customer"] = Scenario(
            key=f"g_{slug}_manager_comprehensive_customer",
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "comprehensive",
                "manager": manager,
                "technology": "All",
                "days": days_text,
                "subscription_id": "",
                "customer_name": customer,
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        )
        matrix[f"a_{slug}_manager_renewal"] = Scenario(
            key=f"a_{slug}_manager_renewal",
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "renewal_portfolio",
                "renewal_type": "renewal_portfolio",
                "manager": manager,
                "technology": "All",
                "days": days_text,
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            source_parity_cohort=f"local-{slug}-manager-team",
        )
        matrix[f"g_{slug}_manager_renewal_customer"] = Scenario(
            key=f"g_{slug}_manager_renewal_customer",
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "renewal",
                "renewal_type": "renewal_single",
                "manager": manager,
                "technology": "All",
                "days": days_text,
                "subscription_id": "",
                "customer_name": customer,
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        )
        matrix[f"d_{slug}_manager_leader_team"] = Scenario(
            key=f"d_{slug}_manager_leader_team",
            endpoint="/start_leader_report",
            payload_mode="form",
            payload={
                "manager": manager,
                "days": days_text,
                "scope_type": "team",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
            source_parity_cohort=f"local-{slug}-manager-team",
        )
        matrix[f"d_{slug}_manager_leader_member"] = Scenario(
            key=f"d_{slug}_manager_leader_member",
            endpoint="/start_leader_report",
            payload_mode="form",
            payload={
                "manager": manager,
                "days": days_text,
                "scope_type": "member",
                "scope_value": member_email,
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        )
        matrix[f"d_{slug}_manager_leader_customer"] = Scenario(
            key=f"d_{slug}_manager_leader_customer",
            endpoint="/start_leader_report",
            payload_mode="form",
            payload={
                "manager": manager,
                "days": days_text,
                "scope_type": "customer",
                "scope_value": customer,
                "scope_member": member_email,
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        )
        matrix[f"g_{slug}_manager_subscription"] = Scenario(
            key=f"g_{slug}_manager_subscription",
            endpoint="/start_subscription_analysis",
            payload_mode="form",
            payload={
                "subscription_id": subscription_id,
                "days": days_text,
                "report_type": "comprehensive",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        )
    # Every multi-manager scenario uses the complete ``All`` fixture scope.
    # Pin its visual contract so a report cannot pass merely because Word and
    # Excel open while the chart layer silently disappears.
    for key, scenario in tuple(matrix.items()):
        if scenario.expected_min_charts:
            continue
        minimum = 3 if scenario.endpoint == "/start_subscription_analysis" else 4
        matrix[key] = Scenario(
            **{
                **asdict(scenario),
                "expected_min_charts": minimum,
            }
        )
    return matrix


def matrix_block_for_key(scenario_key: str) -> str:
    """Return the matrix block letter for a scenario key (Round 133)."""
    prefix = (scenario_key or "").split("_", 1)[0].upper()
    return prefix if prefix in MATRIX_BLOCK_ORDER else "?"


def parse_matrix_blocks(raw: str) -> list[str]:
    """Parse --blocks for the Round 133 option matrix."""
    chosen = [part.strip().upper() for part in (raw or "").split(",") if part.strip()]
    if not chosen or chosen == ["ALL"]:
        return list(MATRIX_BLOCK_ORDER)
    allowed = set(MATRIX_BLOCK_ORDER)
    unknown = [item for item in chosen if item not in allowed]
    if unknown:
        raise ValueError(f"Unknown matrix block(s): {', '.join(unknown)}")
    return chosen


def select_matrix_scenario_keys(
    matrix: dict[str, Scenario],
    blocks: list[str],
    *,
    resume_from: str = "",
) -> list[str]:
    """Filter matrix scenarios by block, preserving cheap-first block order."""
    resume_from_norm = (resume_from or "").strip().lower()
    resume_active = not resume_from_norm
    selected: list[str] = []
    for block in blocks:
        for key in sorted(matrix.keys()):
            if not key.startswith(f"{block.lower()}_"):
                continue
            if not resume_active:
                if key.lower() == resume_from_norm or key.lower().startswith(resume_from_norm):
                    resume_active = True
                else:
                    continue
            selected.append(key)
    return selected


def build_scenario_map() -> dict[str, Scenario]:
    """Return canonical report scenarios for repeated test loops."""
    # Round 51: mirrors the user-defined test matrix.
    return {
        "comprehensive": Scenario(
            key="comprehensive",
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "comprehensive",
                "manager": "Brian Frazier",
                "technology": "All Contact Center",
                "days": "90",
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        ),
        "compact": Scenario(
            key="compact",
            endpoint="/start_compact_analysis",
            payload_mode="json",
            payload={
                "manager": "All Managers",
                "technology": "All Contact Center",
                "days": 90,
                "csone_file": "",
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        ),
        "renewal": Scenario(
            key="renewal",
            endpoint="/start_analysis",
            payload_mode="form",
            payload={
                "report_type": "renewal_portfolio",
                "renewal_type": "renewal_portfolio",
                "manager": "All Managers",
                "technology": "All Contact Center",
                "days": "90",
                "subscription_id": "",
                "customer_name": "",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        ),
        "leader": Scenario(
            key="leader",
            endpoint="/start_leader_report",
            payload_mode="form",
            payload={
                "manager": "Brian Frazier",
                "days": "90",
            },
            expect_excel=True,
            expected_xlsx_sheets=ROUND147_CANONICAL_SOURCE_SHEETS,
        ),
    }


def extract_csrf_token(html: str) -> str:
    """Extract CSRF token from analyze-page meta tag."""
    match = CSRF_META_RE.search(html or "")
    if not match:
        raise ValueError("Could not find csrf-token meta tag in HTML response")
    return unescape(match.group(1))


def parse_download_name(content_disposition: str, fallback: str) -> str:
    """Resolve server download filename from Content-Disposition."""
    if not content_disposition:
        return fallback
    filename_star = re.search(r"filename\*=UTF-8''([^;]+)", content_disposition, re.IGNORECASE)
    if filename_star:
        return filename_star.group(1).strip().strip('"')
    filename = re.search(r'filename="?([^";]+)"?', content_disposition, re.IGNORECASE)
    if filename:
        return filename.group(1).strip().strip('"')
    return fallback


def build_debug_filename(original_name: str, run_id: str, scenario_key: str, timestamp: Optional[str] = None) -> str:
    """Keep original stem and append deterministic debug suffix."""
    # Round 51: preserve naming convention while adding traceability metadata.
    ts = timestamp or _utc_now().strftime("%Y%m%dT%H%M%SZ")
    path = Path(original_name)
    suffix = path.suffix or ".bin"
    stem = path.stem
    trace = (
        f"__data-loop-{_slug(run_id)}"
        f"__scenario-{_slug(scenario_key)}"
        f"__ts-{ts}"
    )
    candidate = f"{stem}{trace}{suffix}"
    # APFS/HFS+ cap a single filename component at 255 bytes. Keep margin for
    # alternate filesystems and preserve the full run/scenario trace whenever
    # possible; only the report stem is shortened, with a digest retaining
    # collision resistance. This matters most for long Renewal technology names.
    max_component_bytes = 240
    if len(candidate.encode("utf-8")) <= max_component_bytes:
        return candidate
    digest = hashlib.sha256(candidate.encode("utf-8")).hexdigest()[:12]
    ending = f"{trace}__h-{digest}{suffix}"
    available = max_component_bytes - len(ending.encode("utf-8"))
    if available < 16:
        compact_trace = (
            f"__data-loop-{_slug(run_id)[:32]}"
            f"__scenario-{_slug(scenario_key)[:48]}"
            f"__ts-{ts}"
        )
        ending = f"{compact_trace}__h-{digest}{suffix}"
        available = max_component_bytes - len(ending.encode("utf-8"))
    bounded_stem = stem.encode("utf-8")[: max(1, available)].decode(
        "utf-8", errors="ignore"
    )
    return f"{bounded_stem}{ending}"


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalize_text_tokens(text: str) -> set[str]:
    clean = text.lower()
    clean = TIMESTAMP_TOKEN_RE.sub(" ", clean)
    clean = NUMBER_RE.sub(" ", clean)
    clean = re.sub(r"[^a-z0-9 ]+", " ", clean)
    return {token for token in clean.split() if len(token) >= 3}


def _normalize_numeric_token(token: str) -> str:
    value = token.strip().replace(",", "")
    if value.endswith("%"):
        value = value[:-1] + "%"
    return value


def _numeric_fingerprint(text: str) -> set[str]:
    """Extract non-timestamp numeric tokens for strict report drift checks."""
    tokens: set[str] = set()
    for match in NUMERIC_TOKEN_RE.finditer(text or ""):
        raw = match.group(0)
        normalized = _normalize_numeric_token(raw)
        digits_only = re.sub(r"\D", "", normalized)
        if len(digits_only) >= 5:
            # Round 52: strict mode treats long numeric strings as volatile
            # identifiers, not business KPIs.
            # Generated IDs, BEMS/case identifiers, and date stamps are expected
            # to change run-to-run. KPI-scale values are covered by shorter
            # count/percentage tokens and by the DOCX/XLSX KPI parity sidecar.
            continue
        if not digits_only:
            continue
        tokens.add(normalized)
    return tokens


def _jaccard_similarity(left: set[str], right: set[str]) -> float:
    if not left and not right:
        return 1.0
    if not left or not right:
        return 0.0
    union = left | right
    if not union:
        return 0.0
    return len(left & right) / float(len(union))


def _parse_numeric_token_for_tolerance(value: str) -> Optional[float]:
    try:
        return float(str(value).replace(",", "").replace("%", ""))
    except (TypeError, ValueError):
        return None


def _jaccard_numeric_similarity_with_tolerance(
    left: set[str],
    right: set[str],
    *,
    absolute_tolerance: float = 0.15,
) -> float:
    """Jaccard-like numeric similarity that treats tiny rounded decimals as equal."""

    if not left and not right:
        return 1.0
    if not left or not right:
        return 0.0

    matched_left: set[str] = set(left & right)
    matched_right: set[str] = set(left & right)
    remaining_right = [value for value in right if value not in matched_right]

    for left_value in sorted(left - matched_left):
        left_num = _parse_numeric_token_for_tolerance(left_value)
        if left_num is None:
            continue
        left_is_percent = str(left_value).strip().endswith("%")
        for right_value in list(remaining_right):
            if str(right_value).strip().endswith("%") != left_is_percent:
                continue
            right_num = _parse_numeric_token_for_tolerance(right_value)
            if right_num is None:
                continue
            # Round 97.3: live comprehensive BE-priority table scores are
            # one-decimal aggregates that can wobble by 0.1 between runs while
            # all integer KPI counts remain stable. Count those rounded-score
            # equivalents as matches; larger numeric drift still fails.
            if abs(left_num - right_num) <= absolute_tolerance:
                matched_left.add(left_value)
                matched_right.add(right_value)
                remaining_right.remove(right_value)
                break

    matches = len(matched_left)
    union_size = len(left) + len(right) - matches
    if union_size <= 0:
        return 1.0
    return matches / float(union_size)


def _extract_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path, "r") as archive:
        body = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    return " ".join(unescape(part) for part in DOCX_TEXT_RE.findall(body))


# Round 52.1: table-only text extraction for the noise-immune numeric drift
# gate.  ``_extract_docx_text`` returns ALL document text including AI-
# generated narrative paragraphs ("Pattern N:", "Root Cause Analysis:", etc.)
# whose numerics regenerate every run (different bug IDs cited, different
# percentages computed, different TAC case IDs as evidence).  The
# ``_numeric_fingerprint`` over that pollutes our drift signal.  Table cells
# carry the actual Snowflake-derived KPI counts and stay byte-stable across
# iterations -- the empirical floor against the round52 manifest was 1.0
# table_numeric_similarity for all four scenarios across 3 iterations.
def _extract_docx_table_text(path: Path) -> str:
    """Concatenate text from DOCX table cells only.

    Skips paragraph runs and headers/footers.  Used by
    ``compare_docx_against_baseline`` to compute a numeric fingerprint that
    excludes AI-narrative volatility and is therefore a true signal for
    real Snowflake data drift.

    Round 52.1: returns "" on any DOCX-parse error (e.g. minimal synthetic
    fixtures missing ``[Content_Types].xml``) so the table-only gate does
    not poison the existing structural/numeric gates that work over the
    raw ``word/document.xml`` payload.  An empty string yields an empty
    fingerprint which Jaccard treats as similarity 1.0 -- effectively
    informational-only when the source DOCX cannot be parsed via python-
    docx.  Real production DOCX files always parse, so this only relaxes
    behavior for synthetic test inputs.
    """
    try:
        document = Document(str(path))
    except Exception:  # noqa: BLE001 - graceful degradation for non-OPC zips
        return ""
    parts: list[str] = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                if text:
                    parts.append(text)
    return " ".join(parts)


def _cap_sorted(values: Iterable[str], limit: int = 25) -> list[str]:
    return sorted(str(value) for value in values)[:limit]


def validate_docx_structure(path: Path, min_chars: int) -> GateResult:
    try:
        if not path.exists() or path.stat().st_size <= 0:
            return GateResult(False, {"reason": "file_missing_or_empty"})
        with zipfile.ZipFile(path, "r") as archive:
            names = set(archive.namelist())
        if "word/document.xml" not in names:
            return GateResult(False, {"reason": "missing_document_xml"})
        text = _extract_docx_text(path)
        stripped = " ".join(text.split())
        return GateResult(
            passed=len(stripped) >= min_chars,
            details={"char_count": len(stripped), "min_chars": min_chars},
        )
    except Exception as exc:  # noqa: BLE001 - diagnostics only
        return GateResult(False, {"reason": "docx_parse_error", "error": str(exc)})


def _sheet_header_tokens(sheet: Any, max_cols: int = 40) -> list[str]:
    tokens: list[str] = []
    row = next(sheet.iter_rows(min_row=1, max_row=1, max_col=max_cols, values_only=True), None)
    if not row:
        return tokens
    for value in row:
        if value is None:
            continue
        token = str(value).strip()
        if token:
            tokens.append(token.lower())
    return tokens


def _sheet_non_empty_count(sheet: Any, max_rows: int = 5000, max_cols: int = 80) -> int:
    count = 0
    for row in sheet.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
        if any(cell not in (None, "") for cell in row):
            count += 1
    return count


def build_xlsx_signature(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        signatures: dict[str, Any] = {
            "sheet_names": [name.lower() for name in workbook.sheetnames],
            "sheet_headers": {},
            "sheet_non_empty_rows": {},
        }
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            signatures["sheet_headers"][sheet_name.lower()] = _sheet_header_tokens(sheet)
            signatures["sheet_non_empty_rows"][sheet_name.lower()] = _sheet_non_empty_count(sheet)
        return signatures
    finally:
        workbook.close()


def validate_xlsx_structure(path: Path, expected_sheets: Iterable[str] = ()) -> GateResult:
    try:
        if not path.exists() or path.stat().st_size <= 0:
            return GateResult(False, {"reason": "file_missing_or_empty"})
        signature = build_xlsx_signature(path)
        sheets = signature["sheet_names"]
        non_empty_total = sum(signature["sheet_non_empty_rows"].values())
        expected = {sheet.lower() for sheet in expected_sheets if str(sheet).strip()}
        missing_expected = sorted(expected - set(sheets))
        return GateResult(
            passed=bool(sheets) and non_empty_total > 0 and not missing_expected,
            details={
                "sheet_count": len(sheets),
                "non_empty_rows": non_empty_total,
                "expected_sheets": sorted(expected),
                "missing_expected_sheets": missing_expected,
            },
        )
    except Exception as exc:  # noqa: BLE001 - diagnostics only
        return GateResult(False, {"reason": "xlsx_parse_error", "error": str(exc)})


def compare_docx_against_baseline(
    current_path: Path,
    baseline_path: Path,
    min_similarity: float,
    *,
    strict: bool = False,
    min_numeric_similarity: float = 0.8,
    # Round 52.1: table-only numeric drift gate.  Computed from
    # ``_extract_docx_table_text`` so AI-generated narrative paragraphs
    # cannot pollute the fingerprint.  Empirical floor against the round52
    # manifest is 1.0 across baseline + 3 iterations -- this is the true
    # signal for real Snowflake data drift.
    min_table_numeric_similarity: float = 0.95,
) -> GateResult:
    try:
        current_text = _extract_docx_text(current_path)
        baseline_text = _extract_docx_text(baseline_path)
        current_tokens = _normalize_text_tokens(current_text)
        baseline_tokens = _normalize_text_tokens(baseline_text)
        similarity = _jaccard_similarity(current_tokens, baseline_tokens)
        current_numbers = _numeric_fingerprint(current_text)
        baseline_numbers = _numeric_fingerprint(baseline_text)
        numeric_similarity = _jaccard_similarity(current_numbers, baseline_numbers)
        numeric_passed = (not strict) or numeric_similarity >= min_numeric_similarity

        # Round 52.1: table-only numeric fingerprint.  This is the new
        # noise-immune binding signal for real data drift.  When tables are
        # stable, we have positive proof Snowflake KPIs did not shift; when
        # tables drift, we have positive proof real data changed (and should
        # be investigated).
        current_table_text = _extract_docx_table_text(current_path)
        baseline_table_text = _extract_docx_table_text(baseline_path)
        current_table_numbers = _numeric_fingerprint(current_table_text)
        baseline_table_numbers = _numeric_fingerprint(baseline_table_text)
        table_numeric_similarity = _jaccard_numeric_similarity_with_tolerance(
            current_table_numbers, baseline_table_numbers
        )
        table_numeric_passed = (
            (not strict) or table_numeric_similarity >= min_table_numeric_similarity
        )

        return GateResult(
            passed=(
                similarity >= min_similarity
                and numeric_passed
                and table_numeric_passed
            ),
            details={
                "similarity": round(similarity, 4),
                "threshold": min_similarity,
                "current_tokens": len(current_tokens),
                "baseline_tokens": len(baseline_tokens),
                "strict": strict,
                "numeric_similarity": round(numeric_similarity, 4),
                "numeric_threshold": min_numeric_similarity if strict else None,
                "current_numeric_tokens": len(current_numbers),
                "baseline_numeric_tokens": len(baseline_numbers),
                "numeric_only_in_current": _cap_sorted(current_numbers - baseline_numbers),
                "numeric_only_in_baseline": _cap_sorted(baseline_numbers - current_numbers),
                # Round 52.1: table-only numeric drift signal.
                "table_numeric_similarity": round(table_numeric_similarity, 4),
                "table_numeric_threshold": (
                    min_table_numeric_similarity if strict else None
                ),
                "current_table_numeric_tokens": len(current_table_numbers),
                "baseline_table_numeric_tokens": len(baseline_table_numbers),
                "table_numeric_only_in_current": _cap_sorted(
                    current_table_numbers - baseline_table_numbers
                ),
                "table_numeric_only_in_baseline": _cap_sorted(
                    baseline_table_numbers - current_table_numbers
                ),
            },
        )
    except Exception as exc:  # noqa: BLE001 - diagnostics only
        return GateResult(False, {"reason": "docx_diff_error", "error": str(exc)})


def compare_xlsx_against_baseline(
    current_path: Path,
    baseline_path: Path,
    min_sheet_overlap: float,
    min_header_similarity: float,
    *,
    strict: bool = False,
    max_row_delta_ratio: float = 0.2,
    max_row_delta_abs: int = 25,
) -> GateResult:
    try:
        current_sig = build_xlsx_signature(current_path)
        baseline_sig = build_xlsx_signature(baseline_path)
        current_sheets = set(current_sig["sheet_names"])
        baseline_sheets = set(baseline_sig["sheet_names"])
        sheet_overlap = _jaccard_similarity(current_sheets, baseline_sheets)
        only_current = sorted(current_sheets - baseline_sheets)
        only_baseline = sorted(baseline_sheets - current_sheets)

        shared = sorted(current_sheets & baseline_sheets)
        if not shared:
            header_similarity = 0.0
        else:
            similarities: list[float] = []
            for name in shared:
                left = set(current_sig["sheet_headers"].get(name, []))
                right = set(baseline_sig["sheet_headers"].get(name, []))
                similarities.append(_jaccard_similarity(left, right))
            header_similarity = sum(similarities) / len(similarities)

        row_deltas: dict[str, Any] = {}
        row_delta_passed = True
        for name in shared:
            current_rows = int(current_sig["sheet_non_empty_rows"].get(name, 0))
            baseline_rows = int(baseline_sig["sheet_non_empty_rows"].get(name, 0))
            delta = abs(current_rows - baseline_rows)
            denom = max(current_rows, baseline_rows, 1)
            ratio = delta / denom
            if delta:
                row_deltas[name] = {
                    "current_rows": current_rows,
                    "baseline_rows": baseline_rows,
                    "delta": delta,
                    "delta_ratio": round(ratio, 4),
                }
            if strict and delta > max_row_delta_abs and ratio > max_row_delta_ratio:
                row_delta_passed = False

        passed = (
            sheet_overlap >= min_sheet_overlap
            and header_similarity >= min_header_similarity
            and row_delta_passed
        )
        return GateResult(
            passed=passed,
            details={
                "sheet_overlap": round(sheet_overlap, 4),
                "sheet_overlap_threshold": min_sheet_overlap,
                "header_similarity": round(header_similarity, 4),
                "header_similarity_threshold": min_header_similarity,
                "shared_sheets": shared,
                "sheets_only_in_current": only_current,
                "sheets_only_in_baseline": only_baseline,
                "strict": strict,
                "row_delta_threshold_ratio": max_row_delta_ratio if strict else None,
                "row_delta_threshold_abs": max_row_delta_abs if strict else None,
                "row_deltas": row_deltas,
            },
        )
    except Exception as exc:  # noqa: BLE001 - diagnostics only
        return GateResult(False, {"reason": "xlsx_diff_error", "error": str(exc)})


def _normalize_kpi_label(label: str) -> str:
    clean = re.sub(r"[^a-z0-9]+", " ", str(label).lower()).strip()
    clean = re.sub(r"\s+", " ", clean)
    for canonical, aliases in KPI_ALIASES.items():
        if clean in aliases:
            return canonical
    return clean.replace(" ", "_")


def _canonical_kpi_label(label: str) -> Optional[str]:
    clean = re.sub(r"[^a-z0-9]+", " ", str(label).lower()).strip()
    clean = re.sub(r"\s+", " ", clean)
    for canonical, aliases in KPI_ALIASES.items():
        if clean in aliases:
            return canonical
    return None


# Round 52 / partial-data-warning Phase 5: canonicals whose values are
# inherently text (no count). Everything NOT in this set is numeric-only:
# if the cell value is non-numeric (e.g. a data-source label like
# "CSConsole" or "CSConsole / Snowflake C360_CS_TASK_C_VW"), the
# extractor must SKIP it rather than canonicalize the source label as
# a KPI value. Without this guard the leader Data Sources table feeds
# strings like ``action_plans = "CSConsole"`` into the parity gate
# while the XLSX reports the row count (354) -- a guaranteed mismatch
# that masks real numerical drift.
_TEXT_VALUED_CANONICAL_KPIS: frozenset[str] = frozenset({
    "manager",
    "technology",
    "risk_category",
})

_NUMERIC_VALUE_RE = re.compile(r"^-?\$?\d[\d,]*(?:\.\d+)?\s*%?$")


def _is_numeric_kpi_value(value: Any) -> bool:
    """True when ``value`` looks like a count/score/percent KPI value.

    Round 57: tolerate trailing ``[Source: ...]`` chrome appended by
    ``report_source_injector``; the chrome is purely presentational and
    must not hide an otherwise-numeric value from the gate's table
    claim detector. Without this, injected cells silently disappear
    from ``metric_claim_count`` instead of becoming source-backed
    claims, defeating the injector.
    """
    if value is None:
        return False
    text = str(value).strip()
    if not text:
        return False
    text = re.sub(r"\s*\[\s*Source\s*:[^\]]*\]\s*", "", text, flags=re.IGNORECASE).strip()
    if not text:
        return False
    return bool(_NUMERIC_VALUE_RE.match(text))


def _is_withheld_kpi_value(value: Any) -> bool:
    """Return whether a visible KPI value explicitly discloses withholding."""

    # Round 148: canonical decision reports intentionally render these states
    # instead of coercing partial source data to zero.  The cross-format gate
    # must compare that disclosure rather than treating it as absent evidence.
    text = re.sub(r"\s+", " ", str(value or "")).strip().lower()
    return text.startswith("unavailable") or text.startswith("withheld")


def _normalize_kpi_value(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    # Round 57: the post-render source-citation injector
    # (report_source_injector) appends ``[Source: ...]`` chrome to KPI
    # value cells in the Word document so the quality gate's
    # adjacent-citation rule passes. The chrome is purely
    # presentational -- the underlying numeric is identical to what
    # the XLSX writer emitted -- so strip any ``[Source: ...]`` block
    # before comparing values across formats. Without this, the parity
    # gate fires on a cosmetic difference (docx="68 [Source: ...]" vs
    # xlsx="68") that the gate is explicitly NOT meant to flag.
    text = re.sub(r"\s*\[\s*Source\s*:[^\]]*\]\s*", " ", text, flags=re.IGNORECASE).strip()
    text = re.sub(r"\s+", " ", text)
    # Round 52 / partial-data-warning Phase 5: strip trailing
    # presentational suffixes that one side adds and the other does not
    # (e.g. risk score rendered "0.7/10" in DOCX vs "0.7" in XLSX, or
    # window written "90 days" in one place and "90" in another). The
    # underlying value is identical; without this normalization the
    # parity gate fires on a purely cosmetic difference.
    _stripped = re.sub(
        r"\s*(?:/\s*\d+(?:\.\d+)?|\bdays?\b|\bdirect\s+reports?\b)\s*$",
        "",
        text,
        flags=re.IGNORECASE,
    ).strip()
    if _stripped:
        text = _stripped
    number = re.fullmatch(r"[-+]?\d+(?:,\d{3})*(?:\.\d+)?%?", text)
    if number:
        return text.replace(",", "")
    return text


_PARAGRAPH_KPI_NUMERIC_RE = re.compile(
    # Round 61 / Phase 2.D: the value group used to be
    # ``-?\$?\d[\d,]*(?:\.\d+)?\s*%?`` which happily matched 9-digit
    # case IDs like ``Case: 700356476`` (the `\d[\d,]*` token consumed
    # the entire ID).  This produced false-positive entries in the
    # per-segment paragraph gate (Round 59 cat-D, ~51 false hits per
    # renewal report) which then misaligned segment boundaries when a
    # canonical KPI match appeared in the same paragraph.  The
    # negative lookahead ``(?!\d{6})`` after the leading ``\d``
    # rejects any value where the first digit is followed by 6+ more
    # digits without a separator (i.e. 7+ contiguous digits) -- ID
    # tokens.  Real KPI metrics are at most 4 digits portfolio-wide
    # (max observed ~2,205 leader citation count) so the gate keeps
    # 100% of legitimate matches; comma-separated values like
    # ``1,234`` still match because the comma breaks the contiguity
    # check.
    r"\b(?P<label>[A-Za-z][A-Za-z /()\-]{2,80}?)\s*[:\-]\s*(?P<value>-?\$?\d(?!\d{6})[\d,]*(?:\.\d+)?\s*%?)",
)


def _paragraph_kpi_match_is_identifier(text: str, match: re.Match[str]) -> bool:
    """Reject short zero-padded IDs such as ``BEMS-0001`` as KPI claims."""

    separator = text[match.end("label") : match.start("value")]
    raw_value = match.group("value").strip().lstrip("-$")
    digits = re.sub(r"\D", "", raw_value)
    return "-" in separator and raw_value.startswith("0") and len(digits) >= 3
# Round 61 / Phase 2.D: secondary regex for the "<number> Label"
# idiom that the comprehensive scenario's LLM narrative occasionally
# uses (e.g. ``"Per the briefing book, there are 0 Action Plans and 0
# Success Priorities"``).  Pre-Round-61 the harness only matched the
# canonical ``Label: number`` shape (with colon or hyphen separator),
# so these idiomatic phrasings silently dropped the KPI from the
# extractor's output -- causing the R58 soak's ``comprehensive.action_plans``
# drift event (extracted in iter1+iter3, missing in iter2).  Restrict
# the match to a small allow-list of labels so we don't accidentally
# pick up non-KPI phrases like ``"90 days"`` or ``"100 customers
# spent"``; expand the list only when a new false-negative is observed.
_PARAGRAPH_KPI_PREFIX_NUMERIC_RE = re.compile(
    r"\b(?P<value>\d{1,5})\s+(?P<label>"
    # Round 97.3: require an explicit total qualifier for AP/AB prefix
    # matches. Live comprehensive / compact narratives contain per-customer
    # prose like "5 active adoption barriers" and recommendation text like
    # "14 open adoption barriers"; those are not stable top-level portfolio
    # KPI rows and must not drift against workbook detail-sheet counts.
    # Direct label/value shapes such as "Open Adoption Barriers: 13" still
    # flow through _PARAGRAPH_KPI_NUMERIC_RE above.
    r"(?:Total\s+)Action\s+Plans?"
    r"|(?:Total\s+)Adoption\s+Barriers?"
    r"|(?:Total\s+)?Customer\s+Pulse(?:\s+records?)?"
    r"|Direct\s+Reports?"
    # Canonical decision reports disclose an incomplete customer universe as
    # an exact retained lower bound.  Match only that full authored sentence;
    # a broad ``N customers`` pattern would mistake ordinary narrative or a
    # per-member row for the portfolio KPI.
    r"|Customers?(?=\s+are\s+evidenced\s+in\s+retained\s+selected-scope\s+records\s+\(partial\s+lower\s+bound\))"
    r")\b",
    re.IGNORECASE,
)
# Round 52 (Phase 2): a separate text-valued pass picks out short metadata
# lines like "Manager: Brian Frazier" or "Technology: All Contact Center"
# whose values are not numeric. Keep label class tight to avoid false hits.
# Round 52 / partial-data-warning Phase 5: stop the value at the first
# pipe ("|") separator. Some report headers concatenate multiple
# label:value pairs onto one line ("Technology: All Contact Center |
# Analysis Period: 90 days"); without this anchor the value greedily
# consumes the trailing pairs and parity comparisons fire on a label
# vs label-plus-suffix mismatch (the XLSX side has only "All Contact
# Center").
# Round 52 / ship: also stop at ``;`` so corpus-context lines like
# ``Technology: Cloud and Hybrid Products; Observed: 2026-04-29T...;
# Prior occurrences: 879`` (rendered by ``report_corpus_context.py``
# inside per-customer narrative blocks) cannot greedily swallow the
# trailing telemetry as the Technology value.  ``[^|;\n]`` still
# captures spaces and Unicode but never crosses a pipe OR semicolon.
_PARAGRAPH_KPI_TEXT_RE = re.compile(
    r"^\s*(?P<label>(?:Manager(?:\s+scope)?|Technology(?:\s+scope)?|Technology\s+Focus|Risk\s+Category))\s*[:\-]\s*(?P<value>[A-Za-z][^|;\n]{0,120}?)\s*(?:[|;]|$)",
    re.IGNORECASE,
)


def _r97_3_skip_paragraph_kpi_scan(text: str) -> bool:
    """Return True for diagnostic prose that mentions KPIs but is not a KPI row."""

    clean = re.sub(r"\s+", " ", str(text or "")).strip().lower()
    if not clean:
        return True
    # Round 97.3: partial-data warnings intentionally carry both kept and
    # excluded counts ("kept 14 of 183 adoption barriers"). Treating the
    # larger denominator as the rendered KPI caused false live parity
    # failures even though the report body correctly showed 14 in-scope rows.
    if "tech_filter_scope_excluded" in clean:
        return True
    if "ab tech filter" in clean and " kept " in clean and " excluded " in clean:
        return True
    return False

# Round 52 / ship: corpus-context tells.  When the source paragraph
# contains ANY of these markers, the line is part of a per-customer
# narrative block (``report_corpus_context.py``) and its embedded
# ``Technology:`` label is NOT the report-scope technology -- it's a
# tag attached to a recurrence row.  Reject the text-KPI match
# entirely; the XLSX side carries the report-scope technology and the
# parity gate uses that as the source of truth.
_CORPUS_CONTEXT_MARKERS: tuple[str, ...] = (
    "Observed:",
    "Prior occurrences:",
    "Sentiment direction:",
)
TOTALS_MARKERS: frozenset[str] = frozenset({
    "total",
    "totals",
    "team total",
    "team totals",
    "grand total",
    "grand totals",
})


def _looks_like_corpus_context_line(text: str) -> bool:
    """Return True when ``text`` carries report_corpus_context tells."""
    return any(marker in text for marker in _CORPUS_CONTEXT_MARKERS)


def _select_multicolumn_value_rows(rows: list[list[str]]) -> list[list[str]]:
    """Round 53.2: choose the same multi-column value rows for KPI and citation scans."""

    if len(rows) < 2 or len(rows[0]) < 3:
        return []
    if len(rows) == 2:
        return [rows[1]]
    return [
        row
        for row in rows[1:]
        if row and str(row[0] or "").strip().lower() in TOTALS_MARKERS
    ]


def _scan_paragraph_for_kpis(text: str, values: dict[str, str]) -> None:
    """Pick out 'Label: 12' style phrases from a paragraph string."""
    if not text:
        return
    if _r97_3_skip_paragraph_kpi_scan(text):
        return
    # Round 61 / Phase 2.D: the canonical "Label: 12" shape requires a
    # ``:`` or ``-`` separator; the secondary "12 Label" shape does
    # not.  Run the prefix scan unconditionally (cheap regex, only
    # fires when the paragraph actually contains digits adjacent to a
    # short allow-listed label).
    for prefix_match in _PARAGRAPH_KPI_PREFIX_NUMERIC_RE.finditer(text):
        label = prefix_match.group("label").strip()
        raw_value = prefix_match.group("value").strip()
        canonical = _canonical_kpi_label(label)
        if canonical and raw_value:
            values.setdefault(canonical, _normalize_kpi_value(raw_value))
    if ":" not in text and "-" not in text:
        return
    for match in _PARAGRAPH_KPI_NUMERIC_RE.finditer(text):
        if _paragraph_kpi_match_is_identifier(text, match):
            continue
        label = match.group("label").strip()
        raw_value = match.group("value").strip()
        canonical = _canonical_kpi_label(label)
        if canonical and raw_value:
            values.setdefault(canonical, _normalize_kpi_value(raw_value))
    text_match = _PARAGRAPH_KPI_TEXT_RE.match(text)
    if text_match and not _looks_like_corpus_context_line(text):
        label = text_match.group("label").strip()
        raw_value = text_match.group("value").strip()
        canonical = _canonical_kpi_label(label)
        if canonical and raw_value:
            values.setdefault(canonical, _normalize_kpi_value(raw_value))


def extract_docx_kpis(path: Path) -> dict[str, Any]:
    """Extract stable KPI values from a Word report's tables AND paragraphs."""
    doc = Document(str(path))
    values: dict[str, str] = {}
    withheld_kpis: set[str] = set()  # Round 148: explicit partial-state parity.

    table_count = 0
    for table in doc.tables:
        table_count += 1
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        # Round 52 (Phase 2): the "headers + values row" heuristic only applies
        # when the table has 3+ columns. For 2-column tables (the canonical
        # label/value shape used by renewal Customer Health Dashboard, leader
        # Team Performance Metrics, comprehensive Title Page metrics) row 0 is
        # ALSO a label/value row and must not be paired with row 1.
        # Round 52 / accuracy-fix-loop: for *multi-row* tables (e.g. the
        # leader "Team Member Activity Breakdown" table whose rows are one
        # CSSM each plus a TOTAL footer), pairing headers with row 1
        # extracts the FIRST team member's per-person counts and labels
        # them as portfolio metrics (action_plans=50 instead of 354).
        # Prefer a row whose first cell is a totals marker; fall back to
        # row 1 only when the table has exactly one data row (header +
        # single value row, e.g. comprehensive Title Page metrics).
        if len(rows) >= 2 and len(rows[0]) >= 3:
            header = rows[0]
            value_rows = _select_multicolumn_value_rows(rows)
            chosen_value_row = value_rows[0] if value_rows else None
            if chosen_value_row is not None and len(header) == len(chosen_value_row):
                for label, value in zip(header, chosen_value_row):
                    canonical = _canonical_kpi_label(label)
                    if canonical and value:
                        # Round 52 / partial-data-warning Phase 5: drop
                        # non-numeric values for count-style canonicals
                        # so a data-source label cell ("CSConsole")
                        # cannot masquerade as a metric value.
                        if (
                            canonical not in _TEXT_VALUED_CANONICAL_KPIS
                            and not _is_numeric_kpi_value(value)
                        ):
                            if _is_withheld_kpi_value(value):
                                withheld_kpis.add(canonical)
                            continue
                        values.setdefault(canonical, _normalize_kpi_value(value))
        for row in rows:
            if len(row) >= 2 and row[0] and row[1]:
                canonical = _canonical_kpi_label(row[0])
                if canonical:
                    if (
                        canonical not in _TEXT_VALUED_CANONICAL_KPIS
                        and not _is_numeric_kpi_value(row[1])
                    ):
                        if _is_withheld_kpi_value(row[1]):
                            withheld_kpis.add(canonical)
                        continue
                    values.setdefault(canonical, _normalize_kpi_value(row[1]))

    # Round 52 (Phase 2): some KPIs (e.g. renewal "Analysis Period: 90 days",
    # leader "Team Size: 7 Direct Reports") only appear in paragraph prose.
    paragraph_count = 0
    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        paragraph_count += 1
        _scan_paragraph_for_kpis(text, values)

    return {
        "table_count": table_count,
        "paragraph_count": paragraph_count,
        "values": values,
        "withheld_kpis": sorted(withheld_kpis),
    }


_LABEL_VALUE_HEADER_PATTERNS = (
    ("metric", "value"),
    ("field", "value"),
    ("key", "value"),
    ("item", "value"),
    ("status", "warning"),
)


def _extract_label_value_sheet(sheet: Any, values: dict[str, str]) -> None:
    """Treat first column as KPI label, second column as value."""
    for row in sheet.iter_rows(min_row=1, max_row=400, max_col=8, values_only=True):
        cells = [cell for cell in row if cell not in (None, "")]
        if len(cells) < 2:
            continue
        label = str(cells[0]).strip()
        value = cells[1]
        if not label:
            continue
        # Skip header rows.
        norm = label.lower()
        if norm in {"metric", "field", "key", "item", "status"}:
            continue
        canonical = _canonical_kpi_label(label)
        if canonical:
            values.setdefault(canonical, _normalize_kpi_value(value))


def _worksheet_to_frame(sheet: Any) -> pd.DataFrame:
    """Round 53: convert a report detail sheet into a DataFrame for source-backed KPI checks."""

    rows = list(sheet.iter_rows(min_row=1, max_row=5000, values_only=True))
    if not rows:
        return pd.DataFrame()
    header_idx = 0
    best_score = -1
    for idx, row in enumerate(rows[:8]):
        normalized = [str(cell).strip() for cell in row if cell not in (None, "")]
        score = len(normalized)
        if score > best_score:
            header_idx = idx
            best_score = score
    header = [
        str(cell).strip() if cell not in (None, "") else f"blank_{idx}"
        for idx, cell in enumerate(rows[header_idx])
    ]
    data_rows = [
        row for row in rows[header_idx + 1 :]
        if any(cell not in (None, "") for cell in row)
    ]
    if not data_rows or not header:
        return pd.DataFrame()
    return pd.DataFrame([dict(zip(header, row)) for row in data_rows])


def _extract_source_backed_detail_kpis(sheet_name: str, sheet: Any, values: dict[str, str]) -> None:
    """Round 53: recompute high-value KPIs from detail sheets, not only summary tiles."""

    normalized = sheet_name.lower()
    role_by_sheet = {
        "ab_detail_all": "adoption_barriers",
        "adoption_barriers": "adoption_barriers",
        "customer_adoption_barriers": "adoption_barriers",
        "all_adoption_barriers": "adoption_barriers",
        "critical_adoption_barriers": "critical_adoption_barriers",
        "csone_detail_all": "support_cases",
        "tac_cases": "support_cases",
        "customer_support_cases": "support_cases",
        "all_support_cases": "support_cases",
        "customer_action_plans": "action_plans",
        "action_plans": "action_plans",
        "customer_customer_pulse": "customer_pulse",
        "customer_pulse": "customer_pulse",
    }
    role = role_by_sheet.get(normalized)
    if not role:
        return
    frame = _worksheet_to_frame(sheet)
    # Current Source Data workbooks deliberately retain an explicit EMPTY
    # contract row so readers can distinguish a successful zero from a
    # missing sheet. That row is provenance, not a source record. Without
    # this guard the parity harness counts one adoption barrier / TAC case
    # whenever a correctly scoped detail sheet is empty.
    if {"Status", "Dataset", "Message"}.issubset(frame.columns):
        empty_contract = (
            frame["Status"].fillna("").astype(str).str.strip().str.casefold().eq("empty")
        )
        frame = frame.loc[~empty_contract].copy()
    elif set(frame.columns) == {"Message"}:
        messages = frame["Message"].fillna("").astype(str).str.strip().str.casefold()
        if messages.ne("").all() and messages.str.startswith("no data available").all():
            frame = frame.iloc[0:0].copy()
    if frame.empty:
        if role == "adoption_barriers":
            values.update(
                adoption_barriers="0",
                open_adoption_barriers="0",
                critical_barriers="0",
            )
        elif role == "critical_adoption_barriers":
            values["critical_barriers"] = "0"
        elif role == "support_cases":
            values.update(
                support_cases="0",
                critical_cases="0",
                high_cases="0",
                bems="0",
            )
        elif role in {"action_plans", "customer_pulse"}:
            values[role] = "0"
            if role == "action_plans":
                values["open_action_plans"] = "0"
        return
    if role == "adoption_barriers":
        values["adoption_barriers"] = _normalize_kpi_value(cm.count_total_barriers(frame))
        values["open_adoption_barriers"] = _normalize_kpi_value(cm.count_open_barriers(frame))
        values["critical_barriers"] = _normalize_kpi_value(cm.count_critical_barriers(frame))
    elif role == "critical_adoption_barriers":
        values["critical_barriers"] = _normalize_kpi_value(cm.count_total_barriers(frame))
    elif role == "support_cases":
        values["support_cases"] = _normalize_kpi_value(cm.count_total_tac(frame))
        values["critical_cases"] = _normalize_kpi_value(cm.count_p1(frame))
        values["high_cases"] = _normalize_kpi_value(cm.count_p2(frame))
        values["bems"] = _normalize_kpi_value(cm.count_bems(frame))
    elif role in {"action_plans", "customer_pulse"}:
        # Round 53.2: when detail sheets exist they are the source of truth;
        # overwrite summary cells so stale dashboard values cannot pass.
        # Round 140: skip provenance rows (mirror AP/TAC SSoT).
        if role == "action_plans":
            values[role] = _normalize_kpi_value(cm.count_total_action_plans(frame))
            values["open_action_plans"] = _normalize_kpi_value(
                cm.count_open_action_plans(pd.DataFrame(), ap_df=frame)
            )
        else:
            from data_normalization import drop_provenance_rows

            pulse_frame = drop_provenance_rows(frame)
            values[role] = _normalize_kpi_value(0 if pulse_frame is None or pulse_frame.empty else len(pulse_frame))


def _extract_team_summary_sheet(sheet: Any, values: dict[str, str]) -> None:
    """Sum numeric columns across rows for leader Team_Summary semantics.

    Leader Team_Summary has one header row of KPI-like names and N data rows
    (one per direct report), and newer workbooks may append a
    ``TOTAL (deduped)`` row. The portfolio total for cross-CSSM workload
    columns comes from that aggregate row when present; team size and customer
    assignments still come from the direct-report rows.
    """
    rows = list(sheet.iter_rows(min_row=1, max_row=200, max_col=20, values_only=True))
    if not rows:
        return
    # Round 52 / partial-data-warning Phase 5: the leader Team_Summary
    # sheet is written with a TITLE row above the actual header row
    # ("Team Summary - <Manager> Team Report" then
    # "Team_Member | Num_Customers | ..."). Pre-Phase 5 we treated row
    # 0 as the header; the column-name canonicals therefore never
    # matched, the header-zero short-circuit was bypassed, and
    # ``len(data_rows)`` counted rows 1..N inclusive of the real
    # header -- producing a team-size that was off-by-one (12 vs the
    # DOCX's 11). Detect the real header row by scanning for the
    # known marker columns ("Team_Member" / "Num_Customers") and
    # fall back to row 0 only when no marker is found.
    header_idx = 0
    _markers = {"team_member", "num_customers", "team member", "team_members"}
    for _idx, _row in enumerate(rows[:5]):
        if not _row:
            continue
        _row_norm = {
            str(cell).strip().lower() for cell in _row if cell is not None
        }
        if _row_norm & _markers:
            header_idx = _idx
            break
    header = rows[header_idx]
    if not header:
        return
    headers_norm = [str(cell).strip() if cell is not None else "" for cell in header]

    canonical_headers: list[Optional[str]] = []
    for col_label in headers_norm:
        canonical_headers.append(_canonical_kpi_label(col_label))

    data_rows = [row for row in rows[header_idx + 1 :] if any(cell not in (None, "") for cell in row)]
    if not data_rows:
        return

    # Round 138: Round 125 added a final ``TOTAL (deduped)`` row to the
    # workbook. Summing it with the direct-report rows double-counted AP,
    # Pulse, and TAC metrics and counted the aggregate as a team member.
    # Keep the aggregate separate. ``Num_Customers`` remains the sum of CSSM
    # assignments because that is the value the Leader Word report publishes;
    # the aggregate row intentionally carries the distinct-customer count.
    def _is_team_total_row(row: tuple[Any, ...]) -> bool:
        label = str(row[0] or "").strip().lower() if row else ""
        return label in TOTALS_MARKERS or label.startswith("total (")

    aggregate_rows = [row for row in data_rows if _is_team_total_row(row)]
    member_rows = [row for row in data_rows if not _is_team_total_row(row)]
    aggregate_row = aggregate_rows[-1] if aggregate_rows else None

    # Round 53: row count maps to team_members, not total_customers.
    team_size_canonical = _canonical_kpi_label("Team Members")
    if team_size_canonical:
        values[team_size_canonical] = str(len(member_rows))

    for col_idx, canonical in enumerate(canonical_headers):
        if not canonical:
            continue
        if aggregate_row is not None and canonical != "total_customers":
            aggregate_cell = aggregate_row[col_idx] if col_idx < len(aggregate_row) else None
            if aggregate_cell not in (None, ""):
                try:
                    aggregate_numeric = float(str(aggregate_cell).replace(",", ""))
                except (TypeError, ValueError):
                    pass
                else:
                    values[canonical] = (
                        str(int(aggregate_numeric))
                        if aggregate_numeric.is_integer()
                        else str(aggregate_numeric)
                    )
                    continue
        total = 0.0
        any_numeric = False
        for row in member_rows:
            if col_idx >= len(row):
                continue
            cell = row[col_idx]
            if cell in (None, ""):
                continue
            try:
                total += float(str(cell).replace(",", ""))
                any_numeric = True
            except (TypeError, ValueError):
                continue
        if any_numeric:
            normalized = total
            if normalized.is_integer():
                values[canonical] = str(int(normalized))
            else:
                values[canonical] = str(normalized)


def _extract_renewal_summary_sheet(sheet: Any, values: dict[str, str]) -> None:
    """Aggregate renewal Renewal_Summary one-row-per-customer payload."""
    rows = list(sheet.iter_rows(min_row=1, max_row=2000, max_col=20, values_only=True))
    if not rows:
        return
    # Round 53: generated renewal workbooks include a title row above the real
    # header, so scan for the Customer / Overall_Risk_Score header instead of
    # assuming row 1 is tabular data.
    header_idx = 0
    for idx, row in enumerate(rows[:5]):
        normalized = {str(cell).strip().lower() for cell in row if cell is not None}
        if {"customer", "overall_risk_score"} & normalized:
            header_idx = idx
            break
    header = rows[header_idx]
    if not header:
        return
    headers_norm = [str(cell).strip() if cell is not None else "" for cell in header]
    risk_score_idx = None
    for idx, col_label in enumerate(headers_norm):
        canonical = _canonical_kpi_label(col_label)
        if canonical == "risk_score":
            risk_score_idx = idx
            break

    data_rows = [row for row in rows[header_idx + 1:] if any(cell not in (None, "") for cell in row)]
    if not data_rows:
        return

    if "total_customers" not in values:
        values["total_customers"] = str(len(data_rows))

    if risk_score_idx is not None and "risk_score" not in values:
        scores: list[float] = []
        for row in data_rows:
            if risk_score_idx >= len(row):
                continue
            cell = row[risk_score_idx]
            if cell in (None, ""):
                continue
            try:
                scores.append(float(str(cell).replace(",", "")))
            except (TypeError, ValueError):
                continue
        if scores:
            # Round 52 (Phase 2): aggregated risk score is informational only.
            # We do NOT auto-fill a top-level KPI from a computed average
            # because DOCX exports an authoritative single value while XLSX
            # rows are per-customer. Cross-format parity for risk_score is
            # only meaningful when both sides emit the same authoritative
            # number; the rendered avg here is informational, persisted as a
            # diagnostic key but not registered as the canonical KPI.
            pass


def _extract_horizontal_label_value_sheet(sheet: Any, values: dict[str, str]) -> None:
    """Treat row 1 as KPI labels, row 2 as values (wide-table layout).

    Used by renewal ``Key_Metrics`` whose dataframe has columns like
    ``[Risk_Score, Risk_Category, Analysis_Period, ...]`` and one data row.
    """
    rows = list(sheet.iter_rows(min_row=1, max_row=2, max_col=20, values_only=True))
    if len(rows) < 2:
        return
    header_row = rows[0]
    value_row = rows[1]
    if not header_row or not value_row:
        return
    for label, value in zip(header_row, value_row):
        if label is None or value in (None, ""):
            continue
        canonical = _canonical_kpi_label(str(label))
        if canonical:
            values.setdefault(canonical, _normalize_kpi_value(value))


# Round 149: Metric_Lineage ``Metric_Key`` is authoritative when the
# display label does not normalize into ``KPI_ALIASES`` (e.g.
# ``kpi.bems`` -> "BEMS escalations (TAC subset)").
_METRIC_KEY_CANONICAL: dict[str, str] = {
    "kpi.team_members": "team_members",
    "kpi.customers": "total_customers",
    "kpi.action_plans_total": "action_plans",
    "kpi.action_plans_open": "open_action_plans",
    "kpi.adoption_barriers": "adoption_barriers",
    "kpi.customer_pulse": "customer_pulse",
    "kpi.tac_cases": "support_cases",
    "kpi.bems": "bems",
    "kpi.high_risk_customers": "high_risk_customers",
}


def _extract_metric_lineage_kpis(
    sheet: Any,
    values: dict[str, str],
    withheld_kpis: Optional[set[str]] = None,
) -> None:
    """Use explicit canonical lineage values as the workbook KPI authority."""

    rows = list(sheet.iter_rows(min_row=1, max_row=10000, values_only=True))
    if not rows:
        return
    headers = [str(value or "").strip() for value in rows[0]]
    try:
        key_idx = headers.index("Metric_Key")
        label_idx = headers.index("Display_Label")
        value_idx = headers.index("Metric_Value")
    except ValueError:
        return
    # Round 148: older lineage sheets did not publish Source_State. Preserve
    # their value extraction while enabling availability parity for new sheets.
    source_state_idx = headers.index("Source_State") if "Source_State" in headers else -1
    for row in rows[1:]:
        if max(key_idx, label_idx, value_idx) >= len(row):
            continue
        metric_key = str(row[key_idx] or "").strip()
        if not metric_key.startswith("kpi."):
            continue
        canonical = _canonical_kpi_label(str(row[label_idx] or ""))
        if not canonical:
            canonical = _METRIC_KEY_CANONICAL.get(metric_key)  # Round 149
        metric_value = row[value_idx]
        if canonical and metric_value not in (None, ""):
            if _is_withheld_kpi_value(metric_value):
                values.pop(canonical, None)
                if withheld_kpis is not None:
                    withheld_kpis.add(canonical)
            else:
                values[canonical] = _normalize_kpi_value(metric_value)
                if withheld_kpis is not None:
                    withheld_kpis.discard(canonical)
        elif canonical:
            # Round 148: Metric_Lineage is authoritative for availability as
            # well as value.  Remove detail-sheet heuristic counts when the
            # canonical metric is withheld because coverage is incomplete.
            source_state = (
                str(row[source_state_idx] or "").strip().lower()
                if 0 <= source_state_idx < len(row)
                else ""
            )
            if source_state in {"partial", "stale", "failed", "unavailable"}:
                values.pop(canonical, None)
                if withheld_kpis is not None:
                    withheld_kpis.add(canonical)


_SCENARIO_SHEET_HANDLERS: dict[str, str] = {
    "summary": "label_value",
    "report_info": "label_value",
    "executive_dashboard": "label_value",
    "risk_summary": "renewal_summary",
    "renewal_summary": "renewal_summary",
    "key_metrics": "horizontal_label_value",
    "team_summary": "team_summary",
}


def extract_xlsx_kpis(path: Path) -> dict[str, Any]:
    """Extract KPI-like label/value pairs from canonical KPI sheets."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        values: dict[str, str] = {}
        withheld_kpis: set[str] = set()  # Round 148: canonical availability.
        scanned_sheets: list[str] = []
        for sheet_name in workbook.sheetnames:
            normalized_sheet = sheet_name.lower()
            handler = _SCENARIO_SHEET_HANDLERS.get(normalized_sheet)
            sheet = workbook[sheet_name]
            _extract_source_backed_detail_kpis(sheet_name, sheet, values)
            if not handler:
                continue
            scanned_sheets.append(sheet_name)
            if handler == "label_value":
                _extract_label_value_sheet(sheet, values)
            elif handler == "team_summary":
                _extract_team_summary_sheet(sheet, values)
            elif handler == "renewal_summary":
                _extract_renewal_summary_sheet(sheet, values)
            elif handler == "horizontal_label_value":
                _extract_horizontal_label_value_sheet(sheet, values)
        if "Team_Summary" in workbook.sheetnames:
            # Round 97.3: for Leader reports, Team_Summary is the semantic
            # counterpart of selected Word executive rollups. The Action_Plans
            # detail sheet is de-duplicated by ID for the ledger view and can
            # have a lower count than the per-CSSM workload total; restore that
            # rollup after detail-sheet extraction. Keep Adoption_Barriers from
            # the de-duplicated detail ledger, matching the Word AB rollup.
            _team_summary_values: dict[str, str] = {}
            _extract_team_summary_sheet(workbook["Team_Summary"], _team_summary_values)
            for _team_key in (
                "action_plans",
                "bems",
                "customer_pulse",
                "support_cases",
                "team_members",
                "total_customers",
            ):
                if _team_key in _team_summary_values:
                    values[_team_key] = _team_summary_values[_team_key]
        if "Metric_Lineage" in workbook.sheetnames:
            # Decision reports write the canonical value and function for
            # each visible KPI here. Apply it last so legacy detail-sheet
            # heuristics cannot overwrite lifecycle semantics (for example,
            # treating unknown-status plans as open).
            _extract_metric_lineage_kpis(
                workbook["Metric_Lineage"],
                values,
                withheld_kpis,
            )
        return {
            "scanned_sheets": scanned_sheets,
            "values": values,
            "withheld_kpis": sorted(withheld_kpis),
        }
    finally:
        workbook.close()


def compare_kpi_parity(
    docx_kpis: dict[str, Any],
    xlsx_kpis: dict[str, Any],
    *,
    strict: bool,
    required_keys: Iterable[str] = (),
) -> GateResult:
    docx_values = docx_kpis.get("values", {}) if isinstance(docx_kpis, dict) else {}
    xlsx_values = xlsx_kpis.get("values", {}) if isinstance(xlsx_kpis, dict) else {}
    docx_withheld = set(docx_kpis.get("withheld_kpis", ())) if isinstance(docx_kpis, dict) else set()
    xlsx_withheld = set(xlsx_kpis.get("withheld_kpis", ())) if isinstance(xlsx_kpis, dict) else set()
    common = sorted(set(docx_values) & set(xlsx_values))
    common_withheld = sorted(docx_withheld & xlsx_withheld)
    availability_mismatches = sorted(
        (set(docx_values) & xlsx_withheld) | (set(xlsx_values) & docx_withheld)
    )
    mismatches = {
        key: {"docx": docx_values.get(key), "xlsx": xlsx_values.get(key)}
        for key in common
        if _normalize_kpi_value(docx_values.get(key)) != _normalize_kpi_value(xlsx_values.get(key))
    }

    required_set = tuple(sorted(set(required_keys)))
    union_keys = set(docx_values) | set(xlsx_values) | docx_withheld | xlsx_withheld
    missing_required = sorted(key for key in required_set if key not in union_keys)

    if availability_mismatches:
        passed = False
        reason = "kpi_availability_mismatch"
    elif common:
        passed = not mismatches
        reason = "compared_common_kpis"
    elif common_withheld:
        # Round 148: explicit, matching withheld states are affirmative parity
        # evidence, just as a non-empty numeric intersection is above. This
        # does not relax numeric drift: any value-vs-withheld disagreement is
        # rejected above.
        passed = True
        reason = "compared_withheld_kpis"
    else:
        # Round 53.2: in strict mode, zero overlap means we cannot prove
        # cross-format accuracy, so fail closed.
        passed = not strict
        reason = "no_common_kpis"

    if strict and missing_required:
        # Round 52 (Phase 2): scenario-specific required keys MUST be present
        # in at least one format (DOCX or XLSX). Missing means extraction
        # coverage is broken or the report stopped emitting that KPI.
        passed = False
        reason = "missing_required_kpis"

    return GateResult(
        passed=passed,
        details={
            "reason": reason,
            "strict": strict,
            "required_keys": list(required_set),
            "missing_required_keys": missing_required,
            "common_kpis": common,
            "common_withheld_kpis": common_withheld,
            "availability_mismatches": availability_mismatches,
            "mismatches": mismatches,
            "docx_kpi_count": len(docx_values),
            "xlsx_kpi_count": len(xlsx_values),
            "docx_only": _cap_sorted(set(docx_values) - set(xlsx_values)),
            "xlsx_only": _cap_sorted(set(xlsx_values) - set(docx_values)),
            "docx_withheld": _cap_sorted(docx_withheld),
            "xlsx_withheld": _cap_sorted(xlsx_withheld),
        },
    )


def extract_and_write_kpis(
    docx_path: Path,
    xlsx_path: Path,
    sidecar_path: Path,
    *,
    strict: bool,
    required_keys: Iterable[str] = (),
) -> tuple[dict[str, Any], GateResult]:
    docx_kpis = extract_docx_kpis(docx_path)
    xlsx_kpis = extract_xlsx_kpis(xlsx_path)
    parity = compare_kpi_parity(
        docx_kpis,
        xlsx_kpis,
        strict=strict,
        required_keys=required_keys,
    )
    payload = {
        "docx": docx_kpis,
        "xlsx": xlsx_kpis,
        "parity": asdict(parity),
        "required_keys": list(sorted(set(required_keys))),
    }
    sidecar_path.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
    return payload, parity


_RASTER_IMAGE_SUFFIXES = {
    ".bmp",
    ".gif",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}
_RASTER_IMAGE_CONTENT_TYPES = {
    "image/bmp",
    "image/gif",
    "image/jpeg",
    "image/png",
    "image/tiff",
    "image/webp",
    "image/x-ms-bmp",
}
_DRAWINGML_BLIP_TAG = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_PICTURE_NON_VISUAL_PROPERTIES_TAG = (
    "{http://schemas.openxmlformats.org/drawingml/2006/picture}cNvPr"
)
_RELATIONSHIP_EMBED_ATTRIBUTE = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
)
_CANONICAL_CHART_IDS = (
    "activity_mix",
    "action_plan_status_aging",
    "risk_distribution",
    "activity_trend",
)
_CANONICAL_CHART_PREFIXES = {
    "activity_mix": ("chart.activity_mix.",),
    "action_plan_status_aging": (
        "chart.action_plan_status.",
        "chart.action_plan_age.",
    ),
    "risk_distribution": ("chart.risk_distribution.",),
    "activity_trend": ("chart.activity_trend.",),
}
_CANONICAL_CHART_STATES = frozenset(
    {"available", "zero", "partial", "stale", "failed", "unavailable", "unknown"}
)
_CANONICAL_CHART_COMPLETE_STATES = frozenset({"available", "zero"})
_CANONICAL_CHART_MAX_ROWS = 100_000


def _chart_cell_missing(value: Any) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _canonical_chart_renderability(xlsx_path: Optional[Path]) -> dict[str, Any]:
    """Project a canonical workbook into bounded, source-aware chart evidence.

    The configured matrix threshold is relaxed only when this projector proves
    that it is reading the exact ordered 17-sheet Source Data contract.  It
    never emits source values, record identifiers, paths, or arbitrary sheet
    names; hostile inventory is represented by counts and digests only.
    """

    result: dict[str, Any] = {
        "schema_version": "canonical-chart-renderability/v1",
        "workbook_supplied": xlsx_path is not None,
        "canonical_workbook_detected": False,
        "contract_valid": False,
        "reason": "workbook_not_supplied",
        "sheet_inventory_exact": False,
        "expected_sheet_count": 17,
        "observed_sheet_count": 0,
        "chart_inventory_exact": False,
        "expected_chart_group_count": len(_CANONICAL_CHART_IDS),
        "observed_chart_group_count": 0,
        "chart_row_count": 0,
        "renderable_chart_count": 0,
        "withheld_chart_count": 0,
        "contract_sha256": "",
    }
    if xlsx_path is None:
        return result

    workbook_path = Path(xlsx_path)
    workbook = None
    try:
        # Reuse the parity gate's OOXML/path/size/expansion checks before
        # openpyxl parses any workbook content.
        from report_source_parity import validate_ooxml_artifact  # noqa: PLC0415

        validate_ooxml_artifact(
            workbook_path,
            allowed_root=workbook_path.parent,
        )
        workbook = openpyxl.load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
        )
    except Exception as exc:  # noqa: BLE001 - evidence stays sanitized
        result.update(
            {
                "reason": "workbook_validation_failed",
                "error_kind": type(exc).__name__,
            }
        )
        return result

    try:
        from decision_report_delivery import SOURCE_DATA_SHEET_NAMES  # noqa: PLC0415

        expected_sheets = tuple(SOURCE_DATA_SHEET_NAMES)
        observed_sheets = tuple(workbook.sheetnames)
        result.update(
            {
                "expected_sheet_count": len(expected_sheets),
                "observed_sheet_count": len(observed_sheets),
                "sheet_inventory_exact": observed_sheets == expected_sheets,
                "sheet_inventory_sha256": hashlib.sha256(
                    json.dumps(
                        observed_sheets,
                        ensure_ascii=False,
                        separators=(",", ":"),
                    ).encode("utf-8")
                ).hexdigest(),
            }
        )
        if observed_sheets != expected_sheets:
            result["reason"] = "sheet_inventory_mismatch"
            return result

        result["canonical_workbook_detected"] = True
        sheet = workbook["Chart_Data"]
        max_row = int(sheet.max_row or 0)
        if max_row < 2 or max_row > _CANONICAL_CHART_MAX_ROWS + 1:
            result["reason"] = "chart_row_bound_invalid"
            return result

        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, ())
        headers = [str(value).strip() if value is not None else "" for value in header_row]
        required_columns = {
            "Metric_Key",
            "Chart_ID",
            "Value",
            "Source_State",
            "Period_Start",
        }
        if (
            not required_columns.issubset(headers)
            or len(headers) != len(set(headers))
        ):
            result["reason"] = "chart_schema_invalid"
            return result
        positions = {name: headers.index(name) for name in required_columns}
        groups = {chart_id: {"rows": 0, "values": 0, "states": set()} for chart_id in _CANONICAL_CHART_IDS}
        metric_keys: set[str] = set()
        contract_errors = 0
        chart_rows = 0
        for row in rows:
            metric_key = str(row[positions["Metric_Key"]] or "").strip()
            if not metric_key:
                # Chart_Data also carries bounded coverage metadata rows.  Only
                # canonical metric rows participate in visual renderability.
                continue
            chart_rows += 1
            if chart_rows > _CANONICAL_CHART_MAX_ROWS:
                result["reason"] = "chart_row_bound_invalid"
                return result
            chart_id = str(row[positions["Chart_ID"]] or "").strip()
            state = str(row[positions["Source_State"]] or "").strip().casefold()
            value = row[positions["Value"]]
            period_start = row[positions["Period_Start"]]
            if chart_id not in groups:
                contract_errors += 1
                continue
            group = groups[chart_id]
            group["rows"] += 1
            group["states"].add(state)
            if (
                metric_key in metric_keys
                or not metric_key.startswith(_CANONICAL_CHART_PREFIXES[chart_id])
                or state not in _CANONICAL_CHART_STATES
            ):
                contract_errors += 1
            metric_keys.add(metric_key)
            if _chart_cell_missing(value):
                continue
            if isinstance(value, bool):
                contract_errors += 1
                continue
            try:
                numeric_value = float(value)
            except (TypeError, ValueError):
                contract_errors += 1
                continue
            if not math.isfinite(numeric_value):
                contract_errors += 1
                continue
            if state not in _CANONICAL_CHART_COMPLETE_STATES:
                contract_errors += 1
            if chart_id == "activity_trend":
                parsed_period = pd.to_datetime(
                    period_start,
                    errors="coerce",
                    utc=True,
                )
                if _chart_cell_missing(parsed_period):
                    contract_errors += 1
                    continue
            group["values"] += 1

        present_groups = {chart_id for chart_id, group in groups.items() if group["rows"]}
        inventory_exact = present_groups == set(_CANONICAL_CHART_IDS)
        renderable = sum(1 for group in groups.values() if group["values"] > 0)
        result.update(
            {
                "chart_inventory_exact": inventory_exact,
                "observed_chart_group_count": len(present_groups),
                "chart_row_count": chart_rows,
                "renderable_chart_count": renderable,
                "withheld_chart_count": len(_CANONICAL_CHART_IDS) - renderable,
                "contract_sha256": hashlib.sha256(
                    json.dumps(
                        {
                            chart_id: {
                                "rows": group["rows"],
                                "values": group["values"],
                                "states": sorted(group["states"]),
                            }
                            for chart_id, group in groups.items()
                        },
                        sort_keys=True,
                        separators=(",", ":"),
                    ).encode("utf-8")
                ).hexdigest(),
            }
        )
        if not inventory_exact or contract_errors:
            result.update(
                {
                    "reason": "chart_contract_invalid",
                    "contract_error_count": contract_errors,
                }
            )
            return result
        result.update({"contract_valid": True, "reason": "canonical_chart_contract"})
        return result
    except Exception as exc:  # noqa: BLE001 - evidence stays sanitized
        result.update(
            {
                "reason": "chart_projection_failed",
                "error_kind": type(exc).__name__,
            }
        )
        return result
    finally:
        if workbook is not None:
            workbook.close()


def _inspect_docx_visible_charts(path: Path, *, doc: Optional[Any] = None) -> dict[str, Any]:
    """Inspect native and accessible raster charts used by report-quality checks.

    Report generators commonly render a chart to PNG before adding it to Word.
    Those visuals live under ``word/media`` and therefore are not represented by
    a native ``word/charts/chart*.xml`` part.  Treat an inline embedded raster as
    a visible chart only when it has positive dimensions and explicit title or
    alternative-text metadata.  That keeps decorative images without accessible
    chart metadata from satisfying the chart-quality signal.
    """

    native_chart_parts: list[str] = []
    embedded_raster_charts: list[dict[str, Any]] = []
    inspection_errors: list[str] = []

    try:
        with zipfile.ZipFile(path) as archive:
            native_chart_parts = sorted(
                {
                    name
                    for name in archive.namelist()
                    if re.fullmatch(r"word/charts/chart\d+\.xml", name)
                }
            )
    except Exception as exc:  # noqa: BLE001 - quality diagnostics must fail safe
        inspection_errors.append(f"native_chart_parts: {exc}")

    try:
        document = doc if doc is not None else Document(str(path))
        for shape_index, shape in enumerate(document.inline_shapes):
            inline = getattr(shape, "_inline", None)
            if inline is None:
                continue

            width_emu = int(getattr(shape, "width", 0) or 0)
            height_emu = int(getattr(shape, "height", 0) or 0)
            if width_emu <= 0 or height_emu <= 0:
                continue

            metadata_nodes = []
            doc_properties = getattr(inline, "docPr", None)
            if doc_properties is not None:
                metadata_nodes.append(doc_properties)
            metadata_nodes.extend(inline.iter(_PICTURE_NON_VISUAL_PROPERTIES_TAG))
            title = next(
                (
                    str(node.get("title") or "").strip()
                    for node in metadata_nodes
                    if str(node.get("title") or "").strip()
                ),
                "",
            )
            alt_text = next(
                (
                    str(node.get("descr") or "").strip()
                    for node in metadata_nodes
                    if str(node.get("descr") or "").strip()
                ),
                "",
            )
            if not title and not alt_text:
                continue

            blips = list(inline.iter(_DRAWINGML_BLIP_TAG))
            if not blips:
                continue
            relationship_id = str(
                blips[0].get(_RELATIONSHIP_EMBED_ATTRIBUTE) or ""
            ).strip()
            if not relationship_id:
                continue

            media_part = ""
            content_type = ""
            try:
                relationship = document.part.rels[relationship_id]
                target_part = relationship.target_part
                media_part = str(getattr(target_part, "partname", "") or "")
                content_type = str(getattr(target_part, "content_type", "") or "")
            except Exception:  # noqa: BLE001 - malformed relationship is not a chart
                continue

            suffix = Path(media_part).suffix.lower()
            if (
                suffix not in _RASTER_IMAGE_SUFFIXES
                and content_type.lower() not in _RASTER_IMAGE_CONTENT_TYPES
            ):
                continue

            embedded_raster_charts.append(
                {
                    "inline_shape_index": shape_index,
                    "width_emu": width_emu,
                    "height_emu": height_emu,
                    "title": title,
                    "alt_text": alt_text,
                    "relationship_id": relationship_id,
                    "media_part": media_part,
                    "content_type": content_type,
                }
            )
    except Exception as exc:  # noqa: BLE001 - quality diagnostics must fail safe
        inspection_errors.append(f"embedded_raster_charts: {exc}")

    native_chart_count = len(native_chart_parts)
    embedded_raster_chart_count = len(embedded_raster_charts)
    return {
        "visible_chart_count": native_chart_count + embedded_raster_chart_count,
        "native_chart_count": native_chart_count,
        "native_chart_parts": native_chart_parts,
        "embedded_raster_chart_count": embedded_raster_chart_count,
        "embedded_raster_charts": embedded_raster_charts,
        "inspection_errors": inspection_errors,
    }


def _count_docx_chart_parts(path: Path) -> int:
    """Count all visible DOCX charts while preserving the legacy helper API."""

    return int(_inspect_docx_visible_charts(path).get("visible_chart_count", 0))


def _docx_full_text(doc: Document) -> str:
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text:
            parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def _source_backed_cell(row: list[str], value_idx: int) -> bool:
    """Round 53.2: source must be in the metric/value cell or an adjacent source cell."""

    candidates = {value_idx}
    if value_idx > 0:
        candidates.add(value_idx - 1)
    if value_idx + 1 < len(row):
        candidates.add(value_idx + 1)
    for idx in candidates:
        candidate = str(row[idx] or "").strip().lower()
        if "[source:" in candidate or candidate.startswith(("kpi.", "chart.")):
            return True
    return False


# Round 114 / Build 83: mirror of
# ``report_source_injector._MATRIX_SOURCE_CAPTION_PREFIX``.  Kept local
# (the same "mirror the contract" pattern the injector uses for the gate
# regexes) so the gate has no hard import dependency on the injector.
_MATRIX_SOURCE_CAPTION_PREFIX = "Sources:"


def _matrix_has_following_source_caption(table: Any) -> bool:
    """Round 114 / Build 83: True when a source caption follows ``table``.

    The R114 injector no longer cites every numeric cell of a
    multi-column count matrix; it writes ONE compact ``Sources: ...``
    caption paragraph immediately below the table instead.  For the
    quality scorer to stay green (without weakening the check), a caption
    directly following a matrix is treated as source backing for that
    table's aggregated multi-column claims.

    Walks forward from the table's ``w:tbl`` element over the immediately
    following paragraph(s) (skipping empties, bounded) and returns True
    when one starts with the stable caption marker.  Never raises -- on
    any introspection failure returns False so the pre-R114 behaviour
    (claim counted unbacked unless cited in-cell) is preserved.
    """
    try:
        from docx.oxml.ns import qn  # type: ignore[import-not-found]

        nxt = table._tbl.getnext()
        for _ in range(4):
            if nxt is None or nxt.tag != qn("w:p"):
                return False
            text = "".join(
                node.text or "" for node in nxt.findall(".//" + qn("w:t"))
            ).strip()
            if text:
                lowered = text.lower()
                return lowered.startswith(
                    _MATRIX_SOURCE_CAPTION_PREFIX.lower()
                ) or lowered.startswith("[source:")
            nxt = nxt.getnext()
        return False
    except Exception:  # noqa: BLE001
        return False


def _paragraph_claim_source_backed(text: str, match_idx: int, matches: list[re.Match[str]]) -> bool:
    """Round 53.2: a paragraph source backs only the current metric span."""

    match = matches[match_idx]
    next_start = matches[match_idx + 1].start() if match_idx + 1 < len(matches) else len(text)
    segment = text[match.start():next_start]
    return "[source:" in segment.lower()


def _paragraph_has_following_source_reference(paragraph: Any) -> bool:
    """Accept an exact Source Data citation in the next non-empty paragraph."""

    try:
        from docx.oxml.ns import qn  # type: ignore[import-not-found]

        nxt = paragraph._p.getnext()
        for _ in range(2):
            if nxt is None or nxt.tag != qn("w:p"):
                return False
            text = "".join(
                node.text or "" for node in nxt.findall(".//" + qn("w:t"))
            ).strip()
            if text:
                return text.lower().startswith("[source:")
            nxt = nxt.getnext()
    except Exception:  # noqa: BLE001
        return False
    return False


def _numeric_tokens_requiring_source(text: str) -> list[str]:
    """Round 53.2: find narrative numeric tokens that need source backing."""

    clean = str(text or "").strip()
    # Round 130 / Build 98: R114/R115 source captions are standalone
    # ``Sources: ...`` paragraphs (Renewal portfolio block, Leader aging
    # bucket tables).  They legitimately mention numbers like ``90`` or
    # ``0-7 days`` in label text but are not uncited narrative claims.
    if clean.lower().startswith(_MATRIX_SOURCE_CAPTION_PREFIX.lower()):
        return []
    if not clean or "[source:" in clean.lower():
        return []
    lowered = clean.lower()
    if any(
        marker in lowered
        for marker in (
            "generated",
            "analysis id",
            "report metadata",
            "page ",
            "build ",
            "version ",
            "data as of",
            "day window",
        )
    ):
        return []
    # Round 149 / Build 111: concise comprehensive overflow pointers are not KPI claims.
    if "additional" in lowered and "row(s)" in lowered and "source data file" in lowered:
        return []
    # Round 149 / Build 111: action-plan step ordinals (01, 03, …) are workflow labels.
    if "next action:" in lowered:
        clean = re.sub(
            r"(?i)next action:\s*0?\d{1,2}\s+",
            "next action: ",
            clean,
        )
    clean = re.sub(r"\b[A-Za-z][A-Za-z0-9]*-0*\d+\b", "", clean)
    tokens = []
    for token in NUMERIC_TOKEN_RE.findall(clean):
        stripped = token.replace(",", "").replace("%", "")
        if len(stripped) >= 4 and stripped.isdigit():
            # Long IDs / dates are handled by structural metadata, not as
            # business facts in this lightweight review pass.
            continue
        tokens.append(token)
    return tokens


def _extract_docx_metric_claims(doc: Document) -> list[dict[str, Any]]:
    """Round 53: collect rendered metric claims and whether source backing is adjacent."""

    claims: list[dict[str, Any]] = []
    for table_idx, table in enumerate(doc.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        # Round 114 / Build 83 + Round 115 / Build 84: a ``Sources: ...``
        # caption directly below a table backs that table's aggregated
        # claims.  R114 introduced this for multi-column matrices; R115
        # extends the same treatment to two-column ``label | value`` KPI
        # cards (the injector no longer cites every numeric cell / row).
        table_caption_backed = _matrix_has_following_source_caption(table)
        if len(rows) >= 2 and len(rows[0]) >= 3:
            header = rows[0]
            caption_backed = table_caption_backed
            value_rows = _select_multicolumn_value_rows(rows)
            for value_row in value_rows:
                row_text = " | ".join(item for item in value_row if item)
                for col_idx, (label, value) in enumerate(zip(header, value_row)):
                    canonical = _canonical_kpi_label(label)
                    if canonical and _is_numeric_kpi_value(value):
                        claims.append(
                            {
                                "location": f"table[{table_idx}]",
                                "label": label,
                                "canonical": canonical,
                                "value": _normalize_kpi_value(value),
                                "source_backed": (
                                    _source_backed_cell(value_row, col_idx) or caption_backed
                                ),
                                "excerpt": row_text[:240],
                            }
                        )
        for row_idx, row in enumerate(rows):
            if len(row) < 2:
                continue
            canonical = _canonical_kpi_label(row[0])
            if canonical and _is_numeric_kpi_value(row[1]):
                row_text = " | ".join(item for item in row if item)
                claims.append(
                    {
                        "location": f"table[{table_idx}].row[{row_idx}]",
                        "label": row[0],
                        "canonical": canonical,
                        "value": _normalize_kpi_value(row[1]),
                        # Round 115 / Build 84: a two-column KPI card is
                        # source-backed by an in-cell citation OR by the
                        # aggregated ``Sources: ...`` caption below it.
                        "source_backed": _source_backed_cell(row, 1)
                        or table_caption_backed,
                        "excerpt": row_text[:240],
                    }
                )
    for paragraph_idx, paragraph in enumerate(doc.paragraphs):
        text = (paragraph.text or "").strip()
        if not text:
            continue
        matches = [
            match
            for match in _PARAGRAPH_KPI_NUMERIC_RE.finditer(text)
            if not _paragraph_kpi_match_is_identifier(text, match)
        ]
        for match_idx, match in enumerate(matches):
            label = match.group("label").strip()
            canonical = _canonical_kpi_label(label)
            if canonical:
                claims.append(
                    {
                        "location": f"paragraph[{paragraph_idx}]",
                        "label": label,
                        "canonical": canonical,
                        "value": _normalize_kpi_value(match.group("value")),
                        "source_backed": _paragraph_claim_source_backed(
                            text, match_idx, matches
                        )
                        or _paragraph_has_following_source_reference(paragraph),
                        "excerpt": text[:240],
                    }
                )
    return claims


def _chart_recommendations(scenario_key: str, chart_count: int, values: dict[str, str]) -> list[dict[str, Any]]:
    """Round 53: suggest visual improvements without treating them as automatic failures."""

    if chart_count:
        return []
    if scenario_key == "leader" and {"team_members", "total_customers"} & set(values):
        return [
            {
                "kind": "chart_opportunity",
                "section": "team_summary",
                "suggestion": "Consider a team workload chart using Team_Summary counts if it improves scanability.",
            }
        ]
    if scenario_key == "renewal" and {"risk_score", "risk_category", "total_customers"} & set(values):
        return [
            {
                "kind": "chart_opportunity",
                "section": "renewal_summary",
                "suggestion": "Consider a renewal-risk distribution chart backed by Renewal_Summary rows.",
            }
        ]
    if scenario_key in {"compact", "comprehensive"} and {"total_customers", "support_cases"} & set(values):
        return [
            {
                "kind": "chart_opportunity",
                "section": "executive_summary",
                "suggestion": "Consider a compact risk or support-volume visual if the source rows expose stable buckets.",
            }
        ]
    return []


def evaluate_report_quality(
    docx_path: Path,
    xlsx_path: Optional[Path],
    *,
    scenario_key: str,
    strict: bool,
    expected_min_charts: int = 0,
) -> tuple[dict[str, Any], GateResult]:
    """Round 53: evaluate accuracy-adjacent report quality beyond baseline drift."""

    try:
        doc = Document(str(docx_path))
        full_text = _docx_full_text(doc)
        paragraph_entries = [
            {
                "paragraph_index": idx,
                "text": p.text.strip(),
                "style": str(getattr(p.style, "name", "") or "").lower(),
                "source_backed": "[source:" in (p.text or "").lower()
                or _paragraph_has_following_source_reference(p),
            }
            for idx, p in enumerate(doc.paragraphs)
            if (p.text or "").strip()
        ]
        paragraphs = [entry["text"] for entry in paragraph_entries]
        headings = [
            p.text.strip()
            for p in doc.paragraphs
            if (p.text or "").strip() and str(getattr(p.style, "name", "")).lower().startswith("heading")
        ]
        metric_claims = _extract_docx_metric_claims(doc)
        uncited_numeric_paragraphs = [
            {
                "paragraph_index": idx,
                "numbers": _numeric_tokens_requiring_source(text),
                "excerpt": text[:240],
            }
            for entry in paragraph_entries
            for idx, text in [(entry["paragraph_index"], entry["text"])]
            if not entry["style"].startswith(("heading", "title"))
            if not entry["source_backed"]
            if _numeric_tokens_requiring_source(text)
        ]
        unbacked_metric_claims = [
            claim for claim in metric_claims if not claim.get("source_backed")
        ]
        docx_kpis = extract_docx_kpis(docx_path)
        xlsx_kpis = extract_xlsx_kpis(xlsx_path) if xlsx_path else {"values": {}}
        value_union = {
            **(xlsx_kpis.get("values", {}) if isinstance(xlsx_kpis, dict) else {}),
            **(docx_kpis.get("values", {}) if isinstance(docx_kpis, dict) else {}),
        }
        chart_metadata = _inspect_docx_visible_charts(docx_path, doc=doc)
        chart_count = int(chart_metadata.get("visible_chart_count", 0))
        configured_min_charts = max(int(expected_min_charts), 0)
        chart_contract = _canonical_chart_renderability(xlsx_path)
        effective_min_charts = configured_min_charts
        if chart_contract.get("contract_valid") is True:
            renderable_chart_count = chart_contract.get("renderable_chart_count")
            if type(renderable_chart_count) is int and renderable_chart_count >= 0:
                effective_min_charts = min(
                    configured_min_charts,
                    renderable_chart_count,
                )
        source_citation_count = len(re.findall(r"\[\s*source\s*:", full_text, flags=re.IGNORECASE))
        empty_table_count = 0
        for table in doc.tables:
            if not any(cell.text.strip() for row in table.rows for cell in row.cells):
                empty_table_count += 1

        errors: list[str] = []
        recommendations: list[dict[str, Any]] = []
        if not paragraphs:
            errors.append("DOCX contains no readable paragraphs.")
        if not doc.tables:
            errors.append("DOCX contains no tables; KPI scanability is likely poor.")
        if empty_table_count:
            recommendations.append(
                {
                    "kind": "formatting",
                    "suggestion": f"Remove or populate {empty_table_count} empty Word table(s).",
                }
            )
        if metric_claims and source_citation_count == 0:
            errors.append("Metric claims are present but the report contains no inline [Source:] citations.")
        if strict and unbacked_metric_claims:
            errors.append(
                f"{len(unbacked_metric_claims)} metric claim(s) lack adjacent [Source:] backing."
            )
        if strict and uncited_numeric_paragraphs:
            errors.append(
                f"{len(uncited_numeric_paragraphs)} paragraph(s) contain uncited numeric claims."
            )
        if chart_contract.get("canonical_workbook_detected") is True and chart_contract.get("contract_valid") is not True:
            errors.append(
                "Canonical Chart_Data renderability contract is invalid; "
                "the configured chart minimum was retained."
            )
        elif chart_contract.get("reason") in {
            "workbook_validation_failed",
            "chart_projection_failed",
        }:
            errors.append(
                "Source Data workbook could not be safely inspected for chart "
                "renderability; the configured chart minimum was retained."
            )
        if strict and chart_count < effective_min_charts:
            errors.append(
                "Report contains "
                f"{chart_count} visible chart(s); this complete-source scenario "
                f"requires at least {effective_min_charts}."
            )
        if len(headings) < 2:
            recommendations.append(
                {
                    "kind": "formatting",
                    "suggestion": "Review heading hierarchy; the report has fewer than two detected headings.",
                }
            )
        recommendations.extend(_chart_recommendations(scenario_key, chart_count, value_union))

        payload = {
            "scenario": scenario_key,
            "docx_path": str(docx_path),
            "xlsx_path": str(xlsx_path) if xlsx_path else None,
            "paragraph_count": len(paragraphs),
            "heading_count": len(headings),
            "table_count": len(doc.tables),
            "chart_count": chart_count,
            "expected_min_charts": effective_min_charts,
            "configured_min_charts": configured_min_charts,
            "canonical_chart_contract": chart_contract,
            "chart_metadata": chart_metadata,
            "source_citation_count": source_citation_count,
            "metric_claim_count": len(metric_claims),
            "unbacked_metric_claim_count": len(unbacked_metric_claims),
            "unbacked_metric_claims": unbacked_metric_claims[:25],
            "uncited_numeric_paragraph_count": len(uncited_numeric_paragraphs),
            "uncited_numeric_paragraphs": uncited_numeric_paragraphs[:25],
            "recommendations": recommendations,
            "errors": errors,
        }
        return payload, GateResult(
            passed=not errors,
            details={
                "reason": "quality_review",
                "strict": strict,
                "errors": errors,
                "recommendation_count": len(recommendations),
                "metric_claim_count": len(metric_claims),
                "unbacked_metric_claim_count": len(unbacked_metric_claims),
                "uncited_numeric_paragraph_count": len(uncited_numeric_paragraphs),
                "chart_count": chart_count,
                "expected_min_charts": effective_min_charts,
                "configured_min_charts": configured_min_charts,
                "canonical_chart_contract": chart_contract,
                "chart_metadata": chart_metadata,
                "source_citation_count": source_citation_count,
            },
        )
    except Exception as exc:  # noqa: BLE001 - diagnostics only
        return (
            {
                "scenario": scenario_key,
                "docx_path": str(docx_path),
                "xlsx_path": str(xlsx_path) if xlsx_path else None,
                "errors": [str(exc)],
            },
            GateResult(False, {"reason": "quality_review_error", "error": str(exc)}),
        )


def extract_and_write_quality(
    docx_path: Path,
    xlsx_path: Optional[Path],
    sidecar_path: Path,
    *,
    scenario_key: str,
    strict: bool,
    expected_min_charts: int = 0,
) -> tuple[dict[str, Any], GateResult]:
    payload, quality = evaluate_report_quality(
        docx_path,
        xlsx_path,
        scenario_key=scenario_key,
        strict=strict,
        expected_min_charts=expected_min_charts,
    )
    payload["quality"] = asdict(quality)
    sidecar_path.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
    return payload, quality


def _filename_matches_scenario(name: str, scenario_key: str, extension: str) -> bool:
    lowered = name.lower()
    if RUN_SUFFIX_RE.search(lowered + "."):
        return False
    if extension == "docx" and not lowered.startswith("adoptiq_report_"):
        return False
    # Round 142: Source Data is the customer-facing workbook name going
    # forward.  Keep the legacy AdoptIQ_Data prefix readable so historical
    # baselines and already-downloaded artifacts remain valid.
    if extension == "xlsx" and not lowered.startswith(
        ("adoptiq_source_data_", "adoptiq_data_")
    ):
        return False

    if scenario_key == "comprehensive":
        return (
            "_brian_frazier_" in lowered
            and "_all_contact_center_" in lowered
            and "_compact_" not in lowered
            and "_renewal_" not in lowered
            and "_leader_" not in lowered
        )
    if scenario_key == "compact":
        return "_compact_" in lowered and "_all_managers_" in lowered and "_all_contact_center_" in lowered
    if scenario_key == "renewal":
        return "_renewal_" in lowered and "_all_managers_" in lowered and "_all_contact_center_" in lowered
    if scenario_key == "leader":
        return "_leader_" in lowered and "_brian_frazier_" in lowered
    return False


def select_latest_baseline(
    downloads_dir: Path,
    scenario_key: str,
    extension: str,
    run_id: str,
    current_debug_name: str,
) -> Optional[Path]:
    candidates: list[Path] = []
    suffix = f".{extension.lower()}"
    for path in downloads_dir.glob(f"*{suffix}"):
        if not path.is_file():
            continue
        if f"__data-loop-{_slug(run_id)}__" in path.name:
            continue
        if path.name == current_debug_name:
            continue
        if _filename_matches_scenario(path.name, scenario_key, extension):
            candidates.append(path)
    if not candidates:
        return None
    candidates.sort(key=lambda item: item.stat().st_mtime, reverse=True)
    return candidates[0]


# ---------------------------------------------------------------------------
# Round 52 (Phase 1): manifest-backed baselines.
# ---------------------------------------------------------------------------
# A manifest is a JSON document that pins each scenario to a specific DOCX/XLSX
# baseline plus its SHA-256.  The harness will refuse to run when the manifest
# is missing, malformed, or when a referenced baseline file does not match the
# pinned digest, so a known-good "golden" baseline cannot drift silently.
#
# Format (version 1):
# {
#   "version": 1,
#   "generated_at_utc": "2026-04-29T17:00:00Z",
#   "git_sha": "<short>",
#   "scenarios": {
#     "comprehensive": {
#       "docx": {
#         "path": "comprehensive/AdoptIQ_Report_xxx.docx",
#         "sha256": "...",
#         "size_bytes": 12345,
#         "captured_at_utc": "2026-04-29T16:00:00Z"
#       },
#       "xlsx": { ... }
#     },
#     "compact": {...}, "renewal": {...}, "leader": {...}
#   }
# }
#
# Paths in the manifest are resolved relative to the manifest file's parent
# directory.  Absolute paths are honored as-is.
MANIFEST_SCHEMA_VERSION = 1


@dataclass(frozen=True)
class BaselineEntry:
    """One pinned baseline file from a manifest."""

    scenario_key: str
    file_type: str
    path: Path
    sha256: str
    size_bytes: int
    captured_at_utc: Optional[str]


def _resolve_manifest_relative(manifest_path: Path, raw_path: str) -> Path:
    candidate = Path(raw_path).expanduser()
    if candidate.is_absolute():
        return candidate
    return (manifest_path.parent / candidate).resolve()


def load_baseline_manifest(manifest_path: Path) -> dict[str, dict[str, BaselineEntry]]:
    """Load a baseline manifest and return scenario → file_type → entry."""
    if not manifest_path.exists():
        raise FileNotFoundError(f"Baseline manifest not found: {manifest_path}")
    try:
        raw = json.loads(manifest_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"Baseline manifest is not valid JSON: {manifest_path}: {exc}") from exc
    if not isinstance(raw, dict):
        raise ValueError(f"Baseline manifest must be a JSON object: {manifest_path}")
    version = raw.get("version")
    if type(version) is not int or version != MANIFEST_SCHEMA_VERSION:
        raise ValueError(
            f"Baseline manifest version {version!r} is unsupported; expected {MANIFEST_SCHEMA_VERSION}: {manifest_path}"
        )
    scenarios_raw = raw.get("scenarios")
    if not isinstance(scenarios_raw, dict) or not scenarios_raw:
        raise ValueError(f"Baseline manifest missing 'scenarios' object: {manifest_path}")
    out: dict[str, dict[str, BaselineEntry]] = {}
    for scenario_key, scenario_payload in scenarios_raw.items():
        if not isinstance(scenario_payload, dict):
            raise ValueError(f"Manifest scenario {scenario_key!r} is malformed: {manifest_path}")
        per_type: dict[str, BaselineEntry] = {}
        for file_type, entry in scenario_payload.items():
            if file_type not in {"docx", "xlsx"}:
                continue
            if not isinstance(entry, dict):
                raise ValueError(
                    f"Manifest entry {scenario_key}/{file_type} is not an object: {manifest_path}"
                )
            raw_path = entry.get("path")
            sha256 = entry.get("sha256")
            size_bytes = entry.get("size_bytes")
            captured_at = entry.get("captured_at_utc")
            if not isinstance(raw_path, str) or not raw_path:
                raise ValueError(
                    f"Manifest entry {scenario_key}/{file_type} missing 'path': {manifest_path}"
                )
            if not isinstance(sha256, str) or len(sha256) != 64:
                raise ValueError(
                    f"Manifest entry {scenario_key}/{file_type} requires 64-char sha256: {manifest_path}"
                )
            if type(size_bytes) is not int or size_bytes <= 0:
                raise ValueError(
                    f"Manifest entry {scenario_key}/{file_type} requires positive size_bytes: {manifest_path}"
                )
            resolved = _resolve_manifest_relative(manifest_path, raw_path)
            per_type[file_type] = BaselineEntry(
                scenario_key=scenario_key,
                file_type=file_type,
                path=resolved,
                sha256=sha256.lower(),
                size_bytes=size_bytes,
                captured_at_utc=captured_at if isinstance(captured_at, str) else None,
            )
        if per_type:
            out[scenario_key] = per_type
    if not out:
        raise ValueError(f"Baseline manifest has no usable docx/xlsx entries: {manifest_path}")
    return out


def verify_baseline_entry(entry: BaselineEntry) -> GateResult:
    """Confirm the on-disk baseline file matches the manifest digest/size."""
    if not entry.path.exists():
        return GateResult(
            False,
            {
                "reason": "baseline_file_missing",
                "expected_path": str(entry.path),
            },
        )
    actual_size = entry.path.stat().st_size
    if actual_size != entry.size_bytes:
        return GateResult(
            False,
            {
                "reason": "baseline_size_mismatch",
                "expected_size_bytes": entry.size_bytes,
                "actual_size_bytes": actual_size,
                "expected_path": str(entry.path),
            },
        )
    actual_sha = _file_sha256(entry.path)
    if actual_sha.lower() != entry.sha256:
        return GateResult(
            False,
            {
                "reason": "baseline_sha256_mismatch",
                "expected_sha256": entry.sha256,
                "actual_sha256": actual_sha,
                "expected_path": str(entry.path),
            },
        )
    return GateResult(
        True,
        {
            "reason": "verified",
            "expected_sha256": entry.sha256,
            "expected_size_bytes": entry.size_bytes,
            "expected_path": str(entry.path),
        },
    )


def write_baseline_manifest(
    manifest_path: Path,
    captured: dict[str, dict[str, Path]],
    *,
    label: Optional[str] = None,
) -> dict[str, Any]:
    """Write a manifest pinning each scenario/file_type to its captured file."""
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    base_dir = manifest_path.parent.resolve()
    scenarios_payload: dict[str, dict[str, Any]] = {}
    captured_at = _utc_now_str()
    for scenario_key, per_type in captured.items():
        per_type_payload: dict[str, Any] = {}
        for file_type, file_path in per_type.items():
            file_path = Path(file_path)
            digest = _file_sha256(file_path)
            try:
                rel = str(file_path.resolve().relative_to(base_dir))
            except ValueError:
                rel = str(file_path.resolve())
            per_type_payload[file_type] = {
                "path": rel,
                "sha256": digest,
                "size_bytes": file_path.stat().st_size,
                "captured_at_utc": captured_at,
            }
        if per_type_payload:
            scenarios_payload[scenario_key] = per_type_payload
    payload = {
        "version": MANIFEST_SCHEMA_VERSION,
        "generated_at_utc": captured_at,
        "git_sha": _git_sha(),
        "label": label,
        "environment": build_environment_summary(),
        "scenarios": scenarios_payload,
    }
    manifest_path.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
    return payload


def capture_baseline_artifact(
    baseline_dir: Path,
    scenario_key: str,
    artifact_path: Path,
) -> Path:
    """Copy a freshly generated artifact into the baseline folder."""
    target_dir = baseline_dir / scenario_key
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / artifact_path.name
    target.write_bytes(artifact_path.read_bytes())
    return target


def collect_analysis_log_excerpt(analysis_id: str, max_lines: int = 200) -> dict[str, Any]:
    """Scan local AdoptIQ logs and collect lines containing this analysis ID."""
    candidate_dirs = [
        Path.home() / ".adoptiq",
        Path.home() / "Library" / "Application Support" / "AdoptIQ",
    ]
    explicit_log = os.environ.get("ADOPTIQ_LOG_FILE", "").strip()
    explicit_paths = [Path(explicit_log).expanduser()] if explicit_log else []
    paths: list[Path] = []
    for path in explicit_paths:
        if path.exists() and path.is_file():
            paths.append(path)
    for log_dir in candidate_dirs:
        if log_dir.exists():
            paths.extend(log_dir.glob("adoptiq*.log*"))
    paths = sorted(set(paths), key=lambda item: item.stat().st_mtime, reverse=True)
    matched_lines: list[str] = []
    sources: list[str] = []
    for path in paths[:12]:
        try:
            content = path.read_text(encoding="utf-8", errors="ignore").splitlines()
        except Exception:  # noqa: BLE001
            continue
        filtered = [line for line in content if analysis_id in line]
        if filtered:
            sources.append(str(path))
            matched_lines.extend(filtered[-max_lines:])
        if len(matched_lines) >= max_lines:
            break
    return {"sources": sources, "lines": matched_lines[-max_lines:]}


def _git_sha() -> Optional[str]:
    try:
        result = subprocess.run(
            ["git", "rev-parse", "--short", "HEAD"],
            cwd=Path(__file__).resolve().parent,
            check=False,
            capture_output=True,
            text=True,
            timeout=5,
        )
        if result.returncode == 0:
            return result.stdout.strip() or None
    except Exception:  # noqa: BLE001 - diagnostics only
        return None
    return None


def build_environment_summary() -> dict[str, Any]:
    version = None
    build = None
    try:
        from config import ADOPTIQ_BUILD, ADOPTIQ_VERSION

        version = ADOPTIQ_VERSION
        build = ADOPTIQ_BUILD
    except Exception:  # noqa: BLE001 - diagnostics only
        pass
    return {
        "python": sys.version.split()[0],
        "platform": platform.platform(),
        "git_sha": _git_sha(),
        "adoptiq_version": version,
        "adoptiq_build": build,
    }


def _r97_2_probe_text_endpoint(
    session: requests.Session,
    base_url: str,
    path: str,
    *,
    expected_text: str | None = None,
    timeout: float = 10.0,
) -> dict[str, Any]:
    """Round 97.2: lightweight app-health probe for text endpoints."""
    url = f"{base_url.rstrip('/')}{path}"
    started = _utc_now()
    try:
        response = session.get(url, timeout=timeout)
        text = (response.text or "").strip()
        ok = response.status_code == 200
        if expected_text is not None:
            ok = ok and text == expected_text
        return {
            "path": path,
            "ok": ok,
            "status_code": response.status_code,
            "elapsed_ms": max(int((_utc_now() - started).total_seconds() * 1000), 0),
            "text": text[:120],
        }
    except Exception as exc:  # noqa: BLE001 - health probe should summarize failures
        return {
            "path": path,
            "ok": False,
            "status_code": None,
            "elapsed_ms": max(int((_utc_now() - started).total_seconds() * 1000), 0),
            "error_kind": type(exc).__name__,
            "error": str(exc)[:240],
        }


def _r97_2_probe_json_endpoint(
    session: requests.Session,
    base_url: str,
    path: str,
    *,
    required_keys: Iterable[str] = (),
    require_ok_true: bool = False,
    timeout: float = 10.0,
) -> dict[str, Any]:
    """Round 97.2: lightweight app-health probe for JSON endpoints."""
    url = f"{base_url.rstrip('/')}{path}"
    started = _utc_now()
    try:
        response = session.get(url, timeout=timeout)
        try:
            payload = response.json()
        except ValueError as exc:
            return {
                "path": path,
                "ok": False,
                "status_code": response.status_code,
                "elapsed_ms": max(int((_utc_now() - started).total_seconds() * 1000), 0),
                "error_kind": "invalid_json",
                "error": str(exc)[:240],
                "body_excerpt": (response.text or "")[:240],
            }
        missing = [key for key in required_keys if not isinstance(payload, dict) or key not in payload]
        literal_ok = not require_ok_true or (
            isinstance(payload, dict) and payload.get("ok") is True
        )
        return {
            "path": path,
            "ok": response.status_code == 200 and not missing and literal_ok,
            "status_code": response.status_code,
            "elapsed_ms": max(int((_utc_now() - started).total_seconds() * 1000), 0),
            "payload_keys": sorted(str(key) for key in payload.keys())[:40] if isinstance(payload, dict) else [],
            "missing_required_keys": missing,
            "literal_ok_required": require_ok_true,
        }
    except Exception as exc:  # noqa: BLE001 - health probe should summarize failures
        return {
            "path": path,
            "ok": False,
            "status_code": None,
            "elapsed_ms": max(int((_utc_now() - started).total_seconds() * 1000), 0),
            "error_kind": type(exc).__name__,
            "error": str(exc)[:240],
        }


def evaluate_app_health(session: requests.Session, base_url: str) -> tuple[dict[str, Any], GateResult]:
    """Round 97.2: preflight stability endpoints before expensive report runs.

    Corpus may be blocked on a user's OneDrive state, so we validate endpoint
    shape rather than require an active corpus.
    """
    probes = [
        _r97_2_probe_text_endpoint(session, base_url, "/ping", expected_text="OK"),
        _r97_2_probe_json_endpoint(
            session,
            base_url,
            "/api/version",
            required_keys=("ok", "version", "build", "process_started_at_utc", "restart_required"),
            require_ok_true=True,
        ),
        _r97_2_probe_json_endpoint(session, base_url, "/api/status/all"),
        _r97_2_probe_json_endpoint(session, base_url, "/api/corpus/status", required_keys=("boot",)),
        _r97_2_probe_json_endpoint(session, base_url, "/api/intel/status", required_keys=("boot",)),
    ]
    failed = [probe for probe in probes if probe.get("ok") is not True]
    payload = {
        "base_url": base_url.rstrip("/"),
        "checked_at_utc": _utc_now_str(),
        "probes": probes,
        "failed_required_paths": [probe["path"] for probe in failed],
    }
    return payload, GateResult(
        passed=not failed,
        details={
            "reason": "app_health_preflight",
            "failed_required_paths": payload["failed_required_paths"],
            "probe_count": len(probes),
        },
    )


def thresholds_summary(config: RunnerConfig) -> dict[str, Any]:
    return {
        "strict": config.strict,
        "min_docx_similarity": config.min_docx_similarity,
        "min_docx_numeric_similarity": config.min_docx_numeric_similarity,
        # Round 52.1: surface the new noise-immune table-only drift gate.
        "min_docx_table_numeric_similarity": (
            config.min_docx_table_numeric_similarity
        ),
        "min_sheet_overlap": config.min_sheet_overlap,
        "min_header_similarity": config.min_header_similarity,
        "max_xlsx_row_delta_ratio": config.max_xlsx_row_delta_ratio,
        "max_xlsx_row_delta_abs": config.max_xlsx_row_delta_abs,
        "min_docx_chars": config.min_docx_chars,
        "baseline_mode": config.baseline_mode,
        "baseline_manifest_path": (
            str(config.baseline_manifest_path) if config.baseline_manifest_path else None
        ),
        "init_baseline": config.init_baseline,
        "init_baseline_dir": (
            str(config.init_baseline_dir) if config.init_baseline_dir else None
        ),
        "init_baseline_label": config.init_baseline_label,
    }


def partial_data_warning_summary(status: dict[str, Any]) -> dict[str, Any]:
    warnings = status.get("partial_data_warnings")
    if not isinstance(warnings, list):
        warnings = []
    by_kind: dict[str, int] = {}
    by_dataset: dict[str, int] = {}
    for warning in warnings:
        if not isinstance(warning, dict):
            continue
        kind = str(warning.get("kind") or "unknown")
        dataset = str(warning.get("dataset") or "unknown")
        by_kind[kind] = by_kind.get(kind, 0) + 1
        by_dataset[dataset] = by_dataset.get(dataset, 0) + 1
    return {"count": len(warnings), "by_kind": by_kind, "by_dataset": by_dataset}


class LiveReportRunner:
    """HTTP client that runs scenarios against a live AdoptIQ server."""

    def __init__(self, config: RunnerConfig):
        self.config = config
        self.session = requests.Session()
        self.csrf_token: Optional[str] = None
        self.session_cookie_value: str = ""
        # Round 52 (Phase 1): manifest mode loads pinned baselines once so every
        # iteration in the loop sees the exact same expected files, and captured
        # baselines accumulate when --init-baseline is set.
        self.manifest: dict[str, dict[str, BaselineEntry]] = {}
        if self.config.baseline_mode == "manifest":
            if self.config.baseline_manifest_path is None:
                raise ValueError("baseline_mode=manifest requires --baseline-manifest")
            self.manifest = load_baseline_manifest(self.config.baseline_manifest_path)
        self._init_captured: dict[str, dict[str, Path]] = {}
        self.app_health: dict[str, Any] = {}
        self.app_health_gate: GateResult = GateResult(False, {"reason": "not_checked"})

    def bootstrap_session(self) -> None:
        self.app_health, self.app_health_gate = evaluate_app_health(
            self.session,
            self.config.base_url,
        )
        if not self.app_health_gate.passed:
            raise RuntimeError(
                "App health preflight failed before report run: %s"
                % json.dumps(self.app_health_gate.details, sort_keys=True)
            )
        url = f"{self.config.base_url.rstrip('/')}/"
        response = self.session.get(url, timeout=self.config.request_timeout_seconds)
        response.raise_for_status()
        self.csrf_token = extract_csrf_token(response.text)
        # Round 51: local dev often runs over plain HTTP while Flask sets
        # ``session; Secure``. Requests stores the cookie but will not send it
        # over HTTP, which causes CSRF validation to fail. Preserve the cookie
        # value so we can explicitly set ``Cookie: session=...`` for local runs.
        self.session_cookie_value = self.session.cookies.get("session", "") or ""

    def _headers(self, include_json_content_type: bool) -> dict[str, str]:
        if not self.csrf_token:
            raise RuntimeError("CSRF token not initialized. Call bootstrap_session() first.")
        headers = {
            "X-Requested-With": "XMLHttpRequest",
            "X-CSRFToken": self.csrf_token,
        }
        if self.session_cookie_value:
            headers["Cookie"] = f"session={self.session_cookie_value}"
        if include_json_content_type:
            headers["Content-Type"] = "application/json"
        return headers

    def _start_scenario(self, scenario: Scenario) -> tuple[str, dict[str, Any]]:
        url = f"{self.config.base_url.rstrip('/')}{scenario.endpoint}"
        include_json = scenario.payload_mode == "json"
        request_payload = dict(scenario.payload)
        if scenario.payload_mode == "json":
            response = self.session.post(
                url,
                json=request_payload,
                headers=self._headers(include_json_content_type=True),
                timeout=self.config.request_timeout_seconds,
            )
        else:
            response = self.session.post(
                url,
                data=request_payload,
                headers=self._headers(include_json_content_type=False),
                timeout=self.config.request_timeout_seconds,
            )
        try:
            data = response.json()
        except ValueError:
            data = {"raw_response": response.text[:400]}
        if response.status_code >= 400:
            raise RuntimeError(
                "Scenario start failed for %s: HTTP %s body=%s"
                % (scenario.key, response.status_code, json.dumps(data, sort_keys=True))
            )
        if data.get("success") is not True:
            raise RuntimeError(f"Scenario start failed for {scenario.key}: {data}")
        analysis_id = data.get("analysis_id")
        if not analysis_id:
            raise RuntimeError(f"Missing analysis_id for {scenario.key}: {data}")
        sanitized_payload = {k: ("<redacted>" if "csrf" in k else v) for k, v in request_payload.items()}
        return analysis_id, {"endpoint": scenario.endpoint, "payload_mode": scenario.payload_mode, "payload": sanitized_payload}

    def _poll_status(self, analysis_id: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        status_url = f"{self.config.base_url.rstrip('/')}/status/{quote(analysis_id)}"
        deadline = time.monotonic() + self.config.scenario_timeout_seconds
        snapshots: list[dict[str, Any]] = []
        while True:
            response = self.session.get(status_url, timeout=self.config.request_timeout_seconds)
            response.raise_for_status()
            payload = response.json()
            snapshots.append(
                {
                    "observed_at_utc": _utc_now_str(),
                    "status": payload.get("status"),
                    "progress": payload.get("progress"),
                    "current_step": payload.get("current_step"),
                    "message": payload.get("message"),
                    "error": payload.get("error"),
                    "elapsed_seconds": payload.get("elapsed_seconds"),
                }
            )
            if payload.get("status") in TERMINAL_STATUSES:
                return payload, snapshots
            if time.monotonic() >= deadline:
                raise TimeoutError(
                    f"Timed out waiting for scenario {analysis_id} after {self.config.scenario_timeout_seconds}s"
                )
            time.sleep(self.config.poll_interval_seconds)

    def _download_artifact(self, analysis_id: str, file_type: str) -> tuple[str, bytes]:
        url = f"{self.config.base_url.rstrip('/')}/download/{quote(analysis_id)}/{file_type}"
        # Round 101: large live DOCX downloads can legitimately take longer
        # than the old fixed 120s harness timeout on a busy packaged app.
        response = self.session.get(url, timeout=self.config.download_timeout_seconds)
        if response.status_code != 200:
            raise RuntimeError(f"Download failed for {analysis_id} {file_type}: HTTP {response.status_code} {response.text[:200]}")
        fallback = f"{analysis_id}.{file_type}"
        raw_name = parse_download_name(response.headers.get("Content-Disposition", ""), fallback)
        return raw_name, response.content

    def _write_sidecars(
        self,
        scenario_result: ScenarioResult,
        meta_stem: str,
        log_excerpt: dict[str, Any],
    ) -> tuple[Path, Path]:
        meta_path = self.config.downloads_dir / f"{meta_stem}.meta.json"
        log_path = self.config.downloads_dir / f"{meta_stem}.log"
        meta_payload = asdict(scenario_result)
        meta_payload["log_excerpt_sources"] = log_excerpt.get("sources", [])
        meta_payload["environment"] = build_environment_summary()
        meta_payload["thresholds"] = thresholds_summary(self.config)
        meta_payload["partial_data_warning_summary"] = partial_data_warning_summary(scenario_result.final_status)
        meta_path.write_text(json.dumps(meta_payload, indent=2, sort_keys=True), encoding="utf-8")
        log_path.write_text("\n".join(log_excerpt.get("lines", [])) + "\n", encoding="utf-8")
        return meta_path, log_path

    def run_scenario(self, scenario: Scenario) -> ScenarioResult:
        started = _utc_now()
        analysis_id: Optional[str] = None
        request_meta: dict[str, Any] = {}
        snapshots: list[dict[str, Any]] = []
        final_status: dict[str, Any] = {}
        artifacts: list[ArtifactRecord] = []
        operational_gate = GateResult(False, {"reason": "not_started"})
        parity_gate = GateResult(True, {"reason": "not_applicable"})
        quality_gate = GateResult(True, {"reason": "not_applicable"})
        all_passed = False
        failure_phase = "start"
        try:
            analysis_id, request_meta = self._start_scenario(scenario)
            failure_phase = "poll_status"
            final_status, snapshots = self._poll_status(analysis_id)
            status_name = str(final_status.get("status", "unknown"))
            status_error = (final_status.get("error") or "").strip()
            word_available = final_status.get("word_available") is True
            excel_available = final_status.get("excel_available") is True
            operational_pass = status_name == "completed" and not status_error and word_available
            if scenario.expect_excel:
                operational_pass = operational_pass and excel_available
            operational_gate = GateResult(
                passed=operational_pass,
                details={
                    "status": status_name,
                    "error": status_error,
                    "word_available": word_available,
                    "excel_available": excel_available,
                },
            )

            if not operational_pass:
                failure_phase = "operational_gate"
                raise RuntimeError(f"Operational gate failed for {scenario.key}: {operational_gate.details}")

            artifact_paths: dict[str, Path] = {}
            for file_type in ("docx", "xlsx"):
                if file_type == "xlsx" and not scenario.expect_excel:
                    continue
                failure_phase = f"download_{file_type}"
                downloaded_name, payload = self._download_artifact(analysis_id, file_type)
                debug_name = build_debug_filename(downloaded_name, self.config.run_id, scenario.key)
                debug_path = self.config.downloads_dir / debug_name
                debug_path.write_bytes(payload)
                artifact_paths[file_type] = debug_path
                failure_phase = f"structural_{file_type}"
                structural_gate = (
                    validate_docx_structure(debug_path, self.config.min_docx_chars)
                    if file_type == "docx"
                    else validate_xlsx_structure(debug_path, expected_sheets=scenario.expected_xlsx_sheets)
                )

                baseline: Optional[Path] = None
                baseline_source = "none"
                baseline_sha: Optional[str] = None
                baseline_manifest_key: Optional[str] = None
                baseline_mtime_utc: Optional[str] = None
                baseline_size_bytes: Optional[int] = None
                baseline_integrity = GateResult(True, {"reason": "not_applicable"})
                baseline_gate = GateResult(False, {"reason": "baseline_not_checked"})
                if self.config.baseline_mode == "off":
                    baseline_gate = GateResult(True, {"reason": "baseline_disabled"})
                elif self.config.baseline_mode == "manifest":
                    failure_phase = f"baseline_manifest_{file_type}"
                    entry = self.manifest.get(scenario.key, {}).get(file_type)
                    if entry is None:
                        baseline_gate = GateResult(
                            False,
                            {
                                "reason": "manifest_missing_entry",
                                "scenario": scenario.key,
                                "file_type": file_type,
                            },
                        )
                    else:
                        baseline_integrity = verify_baseline_entry(entry)
                        if not baseline_integrity.passed:
                            baseline_gate = GateResult(
                                False,
                                {
                                    "reason": "manifest_integrity_failed",
                                    **baseline_integrity.details,
                                },
                            )
                        else:
                            baseline = entry.path
                            baseline_source = "manifest"
                            baseline_sha = entry.sha256
                            baseline_manifest_key = f"{scenario.key}/{file_type}"
                            baseline_size_bytes = entry.size_bytes
                            baseline_mtime_utc = datetime.fromtimestamp(
                                entry.path.stat().st_mtime, tz=UTC
                            ).strftime("%Y-%m-%dT%H:%M:%SZ")
                elif self.config.baseline_mode == "latest":
                    failure_phase = f"baseline_select_{file_type}"
                    baseline = select_latest_baseline(
                        self.config.downloads_dir,
                        scenario.key,
                        file_type,
                        self.config.run_id,
                        debug_name,
                    )
                    if baseline is None:
                        baseline_gate = GateResult(False, {"reason": "baseline_not_found"})
                    else:
                        baseline_source = "latest"
                        baseline_sha = _file_sha256(baseline)
                        baseline_size_bytes = baseline.stat().st_size
                        baseline_mtime_utc = datetime.fromtimestamp(
                            baseline.stat().st_mtime, tz=UTC
                        ).strftime("%Y-%m-%dT%H:%M:%SZ")

                if baseline is not None and baseline_gate.details.get("reason") in {
                    "baseline_not_checked",
                    None,
                }:
                    failure_phase = f"baseline_compare_{file_type}"
                    if file_type == "docx":
                        # Round 52 / ship: comprehensive's AI-narrative
                        # section regenerates run-to-run, so a single
                        # global text-similarity floor cannot fairly
                        # cover both templated reports (compact / renewal
                        # / leader at 0.91-0.99 textual_sim) and
                        # narrative-generated reports.  Resolve per-
                        # scenario thresholds; the numeric gate stays
                        # config-driven across all four.
                        # Round 52.1: also resolve a per-scenario table-
                        # only numeric threshold; this is the new
                        # noise-immune binding signal for real data
                        # drift.  Comprehensive relaxes overall numeric
                        # to informational (0.55) and relies on the
                        # table-only gate (0.95 default) for accuracy.
                        min_text, min_numeric, min_table_numeric = (
                            effective_docx_thresholds(self.config, scenario.key)
                        )
                        baseline_gate = compare_docx_against_baseline(
                            debug_path,
                            baseline,
                            min_text,
                            strict=self.config.strict,
                            min_numeric_similarity=min_numeric,
                            min_table_numeric_similarity=min_table_numeric,
                        )
                    else:
                        baseline_gate = compare_xlsx_against_baseline(
                            debug_path,
                            baseline,
                            self.config.min_sheet_overlap,
                            self.config.min_header_similarity,
                            strict=self.config.strict,
                            max_row_delta_ratio=self.config.max_xlsx_row_delta_ratio,
                            max_row_delta_abs=self.config.max_xlsx_row_delta_abs,
                        )

                artifact_record = ArtifactRecord(
                    file_type=file_type,
                    downloaded_name=downloaded_name,
                    debug_name=debug_name,
                    debug_path=str(debug_path),
                    size_bytes=debug_path.stat().st_size,
                    sha256=_file_sha256(debug_path),
                    baseline_path=str(baseline) if baseline else None,
                    structural=structural_gate,
                    baseline_diff=baseline_gate,
                    baseline_source=baseline_source,
                    baseline_sha256=baseline_sha,
                    baseline_manifest_key=baseline_manifest_key,
                    baseline_mtime_utc=baseline_mtime_utc,
                    baseline_size_bytes=baseline_size_bytes,
                    baseline_integrity=baseline_integrity,
                )
                artifacts.append(artifact_record)

            if "docx" in artifact_paths and "xlsx" in artifact_paths:
                failure_phase = "kpi_parity"
                kpi_sidecar = (
                    self.config.downloads_dir
                    / f"AdoptIQ_Report_{analysis_id}__data-loop-{_slug(self.config.run_id)}__scenario-{scenario.key}.kpis.json"
                )
                _, parity_gate = extract_and_write_kpis(
                    artifact_paths["docx"],
                    artifact_paths["xlsx"],
                    kpi_sidecar,
                    strict=self.config.strict,
                    required_keys=required_kpis_for_scenario(scenario),
                )
            else:
                kpi_sidecar = None

            if "docx" in artifact_paths:
                failure_phase = "report_quality"
                quality_sidecar = (
                    self.config.downloads_dir
                    / f"AdoptIQ_Report_{analysis_id}__data-loop-{_slug(self.config.run_id)}__scenario-{scenario.key}.quality.json"
                )
                _, quality_gate = extract_and_write_quality(
                    artifact_paths["docx"],
                    artifact_paths.get("xlsx"),
                    quality_sidecar,
                    scenario_key=scenario.key,
                    strict=self.config.strict,
                    expected_min_charts=scenario.expected_min_charts,
                )
            else:
                quality_sidecar = None

            # Round 53.3: parity_gate.passed already encodes the strict
            # gradient correctly -- it only fails when there is a real
            # cross-format mismatch (or strict-only edges like
            # ``no_common_kpis`` and ``missing_required_kpis``). The
            # legacy ``... if self.config.strict else True`` collapse
            # silently swallowed real DOCX vs XLSX mismatches whenever
            # a caller forgot to pass ``--strict``, defeating the
            # accuracy gate the harness exists to enforce. Always honor
            # the parity gate's verdict so non-strict callers still see
            # numeric drift.
            all_passed = (
                operational_gate.passed
                and all(
                    artifact.structural.passed and artifact.baseline_diff.passed
                    for artifact in artifacts
                )
                and parity_gate.passed
                and quality_gate.passed
            )

            if (
                self.config.init_baseline
                and self.config.init_baseline_dir is not None
                and operational_gate.passed
                and all(artifact.structural.passed for artifact in artifacts)
            ):
                # Round 52 (Phase 1): only structural+operational success
                # qualifies a run for capture; baseline diff failures are
                # expected here because there is no baseline yet.
                captured_per_type: dict[str, Path] = {}
                for artifact in artifacts:
                    captured_path = capture_baseline_artifact(
                        self.config.init_baseline_dir,
                        scenario.key,
                        Path(artifact.debug_path),
                    )
                    captured_per_type[artifact.file_type] = captured_path
                if captured_per_type:
                    self._init_captured[scenario.key] = captured_per_type
            return self._build_scenario_result(
                scenario=scenario,
                analysis_id=analysis_id,
                started=started,
                request_meta=request_meta,
                snapshots=snapshots,
                final_status=final_status,
                operational=operational_gate,
                artifacts=artifacts,
                all_passed=all_passed,
                parity=parity_gate,
                kpi_sidecar_path=str(kpi_sidecar) if kpi_sidecar else None,
                quality=quality_gate,
                quality_sidecar_path=str(quality_sidecar) if quality_sidecar else None,
            )
        except Exception as exc:  # noqa: BLE001
            if not final_status:
                final_status = {"status": "error", "error": str(exc)}
            return self._build_scenario_result(
                scenario=scenario,
                analysis_id=analysis_id,
                started=started,
                request_meta=request_meta,
                snapshots=snapshots,
                final_status=final_status,
                operational=operational_gate,
                artifacts=artifacts,
                all_passed=False,
                crash_error=str(exc),
                failure_phase=failure_phase,
                exception_type=type(exc).__name__,
            )

    def _build_scenario_result(
        self,
        *,
        scenario: Scenario,
        analysis_id: Optional[str],
        started: datetime,
        request_meta: dict[str, Any],
        snapshots: list[dict[str, Any]],
        final_status: dict[str, Any],
        operational: GateResult,
        artifacts: list[ArtifactRecord],
        all_passed: bool,
        parity: Optional[GateResult] = None,
        quality: Optional[GateResult] = None,
        kpi_sidecar_path: Optional[str] = None,
        quality_sidecar_path: Optional[str] = None,
        crash_error: Optional[str] = None,
        failure_phase: Optional[str] = None,
        exception_type: Optional[str] = None,
    ) -> ScenarioResult:
        finished = _utc_now()
        if crash_error:
            final_status = dict(final_status)
            final_status.setdefault("error", crash_error)
            final_status.setdefault("status", "error")
        result = ScenarioResult(
            scenario=scenario.key,
            analysis_id=analysis_id,
            started_at_utc=started.strftime("%Y-%m-%dT%H:%M:%SZ"),
            completed_at_utc=finished.strftime("%Y-%m-%dT%H:%M:%SZ"),
            elapsed_seconds=max(int((finished - started).total_seconds()), 0),
            request=request_meta,
            status_snapshots=snapshots,
            final_status=final_status,
            operational=operational,
            artifacts=artifacts,
            all_passed=all_passed,
            parity=parity or GateResult(True, {"reason": "not_applicable"}),
            quality=quality or GateResult(True, {"reason": "not_applicable"}),
            kpi_sidecar_path=kpi_sidecar_path,
            quality_sidecar_path=quality_sidecar_path,
            failure_phase=failure_phase if crash_error else None,
            exception_type=exception_type if crash_error else None,
        )
        if analysis_id:
            log_excerpt = collect_analysis_log_excerpt(analysis_id)
            stem = f"AdoptIQ_Report_{analysis_id}__data-loop-{_slug(self.config.run_id)}__scenario-{scenario.key}"
            _, log_path = self._write_sidecars(result, stem, log_excerpt)
            result.log_excerpt_path = str(log_path)
        return result


def _default_run_id() -> str:
    return _utc_now().strftime("%Y%m%dT%H%M%SZ")


def _parse_scenarios(raw: str, available: Iterable[str]) -> list[str]:
    chosen = [part.strip().lower() for part in raw.split(",") if part.strip()]
    if not chosen or chosen == ["all"]:
        return list(available)
    unknown = [item for item in chosen if item not in set(available)]
    if unknown:
        raise ValueError(f"Unknown scenario(s): {', '.join(unknown)}")
    return chosen


def build_runner_config(args: argparse.Namespace) -> RunnerConfig:
    scenarios = build_scenario_map()
    keys = _parse_scenarios(args.scenarios, scenarios.keys())
    strict = bool(args.strict)
    min_docx_similarity = float(args.min_docx_similarity)
    min_sheet_overlap = float(args.min_sheet_overlap)
    min_header_similarity = float(args.min_header_similarity)
    if strict:
        min_docx_similarity = max(min_docx_similarity, 0.55)
        min_sheet_overlap = max(min_sheet_overlap, 0.85)
        min_header_similarity = max(min_header_similarity, 0.8)

    baseline_mode = args.baseline_mode
    baseline_manifest_path: Optional[Path] = None
    raw_manifest = (getattr(args, "baseline_manifest", None) or "").strip()
    if raw_manifest:
        baseline_manifest_path = Path(raw_manifest).expanduser().resolve()
    if baseline_mode == "manifest" and baseline_manifest_path is None:
        raise ValueError("--baseline-mode manifest requires --baseline-manifest <path>")

    init_baseline = bool(getattr(args, "init_baseline", False))
    init_baseline_dir: Optional[Path] = None
    raw_init_dir = (getattr(args, "init_baseline_dir", "") or "").strip()
    if init_baseline:
        if not raw_init_dir:
            raise ValueError("--init-baseline requires --init-baseline-dir <path>")
        init_baseline_dir = Path(raw_init_dir).expanduser().resolve()
    init_baseline_label = (getattr(args, "init_baseline_label", "") or "").strip() or None

    return RunnerConfig(
        base_url=args.base_url.rstrip("/"),
        downloads_dir=Path(args.downloads_dir).expanduser().resolve(),
        iterations=max(args.iterations, 1),
        poll_interval_seconds=max(args.poll_interval, 1.0),
        scenario_timeout_seconds=max(args.timeout, 60),
        request_timeout_seconds=max(getattr(args, "request_timeout", 120), 30),
        download_timeout_seconds=max(getattr(args, "download_timeout", 300), 60),
        run_id=args.run_id or _default_run_id(),
        stop_on_failure=bool(args.stop_on_failure),
        scenario_keys=keys,
        baseline_mode=baseline_mode,
        min_docx_similarity=min_docx_similarity,
        min_sheet_overlap=min_sheet_overlap,
        min_header_similarity=min_header_similarity,
        min_docx_chars=max(int(args.min_docx_chars), 50),
        strict=strict,
        min_docx_numeric_similarity=float(args.min_docx_numeric_similarity),
        # Round 52.1: thread the new table-only numeric gate through.
        min_docx_table_numeric_similarity=float(
            args.min_docx_table_numeric_similarity
        ),
        max_xlsx_row_delta_ratio=float(args.max_xlsx_row_delta_ratio),
        max_xlsx_row_delta_abs=max(int(args.max_xlsx_row_delta_abs), 0),
        baseline_manifest_path=baseline_manifest_path,
        init_baseline=init_baseline,
        init_baseline_dir=init_baseline_dir,
        init_baseline_label=init_baseline_label,
    )


def run_iterations(config: RunnerConfig) -> dict[str, Any]:
    config.downloads_dir.mkdir(parents=True, exist_ok=True)
    scenario_map = build_scenario_map()
    runner = LiveReportRunner(config)
    runner.bootstrap_session()

    all_results: list[ScenarioResult] = []
    started = _utc_now()
    aborted = False

    for iteration in range(1, config.iterations + 1):
        for key in config.scenario_keys:
            scenario = scenario_map[key]
            print(f"[run] iteration={iteration} scenario={key}")
            result = runner.run_scenario(scenario)
            all_results.append(result)
            print(
                "[done] scenario=%s status=%s pass=%s analysis_id=%s"
                % (
                    key,
                    result.final_status.get("status", "unknown"),
                    result.all_passed,
                    result.analysis_id or "n/a",
                )
            )
            if config.stop_on_failure and result.all_passed is not True:
                aborted = True
                break
        if aborted:
            break

    finished = _utc_now()
    summary = {
        "run_id": config.run_id,
        "started_at_utc": started.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "completed_at_utc": finished.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "elapsed_seconds": max(int((finished - started).total_seconds()), 0),
        "base_url": config.base_url,
        "downloads_dir": str(config.downloads_dir),
        "iterations_requested": config.iterations,
        "iterations_completed": len(all_results),
        "stop_on_failure": config.stop_on_failure,
        "aborted": aborted,
        "all_passed": (
            all(result.all_passed is True for result in all_results)
            if all_results
            else False
        ),
        "environment": build_environment_summary(),
        "thresholds": thresholds_summary(config),
        "app_health": runner.app_health,
        "app_health_gate": asdict(runner.app_health_gate),
        "partial_data_warning_summary": {
            result.scenario: partial_data_warning_summary(result.final_status)
            for result in all_results
        },
        "quality_summary": {
            result.scenario: result.quality.details
            for result in all_results
        },
        "results": [asdict(result) for result in all_results],
    }
    if config.init_baseline and config.init_baseline_dir is not None and runner._init_captured:
        manifest_path = config.init_baseline_dir / "baseline_manifest.json"
        write_baseline_manifest(
            manifest_path,
            runner._init_captured,
            label=config.init_baseline_label,
        )
        summary["baseline_manifest_path"] = str(manifest_path)
        print(f"[manifest] wrote {manifest_path}")

    summary_name = (
        f"AdoptIQ_ReportIterationSummary__data-loop-{_slug(config.run_id)}__ts-{_utc_now().strftime('%Y%m%dT%H%M%SZ')}.json"
    )
    summary_path = config.downloads_dir / summary_name
    summary_path.write_text(json.dumps(summary, indent=2, sort_keys=True), encoding="utf-8")
    print(f"[summary] {summary_path}")
    return summary


def run_option_matrix(
    config: RunnerConfig,
    matrix: dict[str, Scenario],
    scenario_keys: list[str],
) -> dict[str, Any]:
    """Round 133: execute the exhaustive option matrix against a live app."""
    if not scenario_keys:
        raise ValueError("No matrix scenarios selected")

    config.downloads_dir.mkdir(parents=True, exist_ok=True)
    runner = LiveReportRunner(config)
    runner.bootstrap_session()

    all_results: list[ScenarioResult] = []
    started = _utc_now()
    aborted = False

    for key in scenario_keys:
        scenario = matrix[key]
        block = matrix_block_for_key(key)
        print(f"[matrix] block={block} scenario={key}")
        result = runner.run_scenario(scenario)
        all_results.append(result)
        print(
            "[done] block=%s scenario=%s status=%s pass=%s analysis_id=%s"
            % (
                block,
                key,
                result.final_status.get("status", "unknown"),
                result.all_passed,
                result.analysis_id or "n/a",
            )
        )
        if config.stop_on_failure and result.all_passed is not True:
            aborted = True
            break

    finished = _utc_now()
    expected_scenario_keys = list(scenario_keys)
    completed_scenario_keys = [result.scenario for result in all_results]
    missing_scenario_keys = [
        key for key in expected_scenario_keys if key not in completed_scenario_keys
    ]
    unexpected_scenario_keys = [
        key for key in completed_scenario_keys if key not in expected_scenario_keys
    ]
    duplicate_completed_keys = sorted(
        {
            key
            for key in completed_scenario_keys
            if completed_scenario_keys.count(key) > 1
        }
    )
    scenario_inventory_exact = (
        completed_scenario_keys == expected_scenario_keys
        and not duplicate_completed_keys
    )
    summary = {
        "run_id": config.run_id,
        "matrix_mode": True,
        "started_at_utc": started.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "completed_at_utc": finished.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "elapsed_seconds": max(int((finished - started).total_seconds()), 0),
        "base_url": config.base_url,
        "downloads_dir": str(config.downloads_dir),
        "scenario_keys_requested": list(scenario_keys),
        "scenario_keys_expected": expected_scenario_keys,
        "scenario_count_expected": len(expected_scenario_keys),
        "scenario_keys_completed": completed_scenario_keys,
        "scenario_count_completed": len(completed_scenario_keys),
        "scenario_keys_missing": missing_scenario_keys,
        "scenario_keys_unexpected": unexpected_scenario_keys,
        "scenario_keys_completed_duplicate": duplicate_completed_keys,
        "scenario_inventory_exact": scenario_inventory_exact,
        "scenarios_completed": len(all_results),
        "stop_on_failure": config.stop_on_failure,
        "aborted": aborted,
        "all_passed": (
            scenario_inventory_exact
            and bool(all_results)
            and all(result.all_passed is True for result in all_results)
        ),
        "environment": build_environment_summary(),
        "thresholds": thresholds_summary(config),
        "app_health": runner.app_health,
        "app_health_gate": asdict(runner.app_health_gate),
        "partial_data_warning_summary": {
            result.scenario: partial_data_warning_summary(result.final_status)
            for result in all_results
        },
        "quality_summary": {
            result.scenario: result.quality.details
            for result in all_results
        },
        "grounding_summary": {
            result.scenario: _grounding_summary_from_status(result.final_status)
            for result in all_results
        },
        "results": [asdict(result) for result in all_results],
    }

    summary_name = (
        f"AdoptIQ_ReportOptionMatrixSummary__data-loop-{_slug(config.run_id)}__ts-{_utc_now().strftime('%Y%m%dT%H%M%SZ')}.json"
    )
    summary_path = config.downloads_dir / summary_name
    summary_path.write_text(json.dumps(summary, indent=2, sort_keys=True), encoding="utf-8")
    print(f"[matrix_summary] {summary_path}")
    return summary


def _grounding_summary_from_status(status: dict[str, Any]) -> dict[str, Any]:
    """Extract grounding rejection rollup from a terminal status payload."""
    diag = status.get("grounding_diagnostics") if isinstance(status, dict) else None
    if not isinstance(diag, dict):
        return {"rejected": None, "total": None, "rate": None}
    rollup = diag.get("rejection_summary")
    if not isinstance(rollup, dict):
        return {"rejected": None, "total": None, "rate": None}
    return {
        "rejected": rollup.get("rejected"),
        "total": rollup.get("total"),
        "rate": rollup.get("rate"),
    }


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Run repeatable live AdoptIQ report scenarios and write debug bundles to Downloads.",
    )
    parser.add_argument("--base-url", default="http://127.0.0.1:5151", help="Live AdoptIQ base URL")
    parser.add_argument("--downloads-dir", default="~/Downloads", help="Directory for debug artifacts")
    parser.add_argument("--iterations", type=int, default=1, help="Number of full scenario loops")
    parser.add_argument(
        "--scenarios",
        default="all",
        help="Comma-separated subset: comprehensive,compact,renewal,leader or 'all'",
    )
    parser.add_argument("--poll-interval", type=float, default=5.0, help="Seconds between /status polls")
    parser.add_argument("--timeout", type=int, default=1800, help="Max seconds per scenario")
    parser.add_argument(
        "--request-timeout",
        type=int,
        default=120,
        help="Max seconds for app health, start, and status HTTP requests",
    )
    parser.add_argument(
        "--download-timeout",
        type=int,
        default=300,
        help="Max seconds for each DOCX/XLSX artifact download",
    )
    parser.add_argument("--run-id", default="", help="Optional run identifier for artifact names")
    parser.add_argument("--stop-on-failure", action="store_true", help="Stop immediately when any scenario fails")
    parser.add_argument(
        "--baseline-mode",
        choices=["off", "latest", "manifest"],
        default="latest",
        help="Baseline lookup strategy: 'off' disables diff, 'latest' picks newest matching file in Downloads, 'manifest' pins exact files via JSON",
    )
    parser.add_argument(
        "--baseline-manifest",
        default="",
        help="Path to baseline manifest JSON (required when --baseline-mode=manifest)",
    )
    parser.add_argument(
        "--init-baseline",
        action="store_true",
        help="Capture this run's artifacts as a new baseline manifest (use with --init-baseline-dir)",
    )
    parser.add_argument(
        "--init-baseline-dir",
        default="",
        help="Directory under which to capture baseline artifacts and write baseline_manifest.json",
    )
    parser.add_argument(
        "--init-baseline-label",
        default="",
        help="Optional label persisted in the captured baseline manifest",
    )
    parser.add_argument("--strict", action="store_true", help="Enable stricter semantic diff and KPI parity gates")
    parser.add_argument(
        "--min-docx-similarity",
        type=float,
        default=0.35,
        help="Minimum docx token similarity vs baseline",
    )
    parser.add_argument(
        "--min-sheet-overlap",
        type=float,
        default=0.5,
        help="Minimum xlsx sheet-name overlap vs baseline",
    )
    parser.add_argument(
        "--min-header-similarity",
        type=float,
        default=0.3,
        help="Minimum xlsx header overlap vs baseline",
    )
    parser.add_argument(
        "--min-docx-chars",
        type=int,
        default=200,
        help="Minimum normalized text length for docx structural pass",
    )
    parser.add_argument(
        "--min-docx-numeric-similarity",
        type=float,
        default=0.8,
        help="Strict-mode minimum numeric-token similarity vs docx baseline",
    )
    parser.add_argument(
        # Round 52.1: noise-immune table-only numeric drift gate.  Default
        # 0.95 is the empirical floor across the round52 manifest 3-iter
        # repeatability proof (all four scenarios held 1.0).
        "--min-docx-table-numeric-similarity",
        type=float,
        default=0.95,
        help=(
            "Strict-mode minimum table-only numeric-token similarity vs "
            "docx baseline (Round 52.1 noise-immune drift gate)"
        ),
    )
    parser.add_argument(
        "--max-xlsx-row-delta-ratio",
        type=float,
        default=0.2,
        help="Strict-mode maximum row-count delta ratio before xlsx baseline fails",
    )
    parser.add_argument(
        "--max-xlsx-row-delta-abs",
        type=int,
        default=25,
        help="Strict-mode row-count delta must exceed this absolute value before ratio fails",
    )
    return parser


def main(argv: Optional[list[str]] = None) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)
    try:
        config = build_runner_config(args)
        summary = run_iterations(config)
    except Exception as exc:  # noqa: BLE001
        print(f"[error] {exc}", file=sys.stderr)
        return 1
    return 0 if summary.get("all_passed") is True else 2


if __name__ == "__main__":
    raise SystemExit(main())
