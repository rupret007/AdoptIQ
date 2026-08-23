"""Independent public-artifact completeness and drill-through audit.

The canonical report contract proves numeric parity.  This companion gate
focuses on the presentation defects that can still make a mathematically
correct report unusable: unexplained ``Unknown`` values, literal renderer
placeholders, export footer rows masquerading as cases, and missing or unsafe
CSConsole record links.  It never infers business values or calls a live
source.
"""

from __future__ import annotations

import re
from typing import Any, Mapping

import pandas as pd

from data_normalization import customer_names_match, normalize_customer_name

from source_record_links import (
    SOURCE_RECORD_URL_COLUMN,
    build_source_record_url,
    is_allowed_source_record_url,
)


_UNDEFINED_RE = re.compile(r"\bundefined\b", re.IGNORECASE)
_GENERIC_UNKNOWN_RE = re.compile(r"\bunknown\b", re.IGNORECASE)
_PLACEHOLDER_TOKENS = frozenset({"nan", "none", "null", "<na>"})
_VALID_SOURCE_STATES = frozenset(
    {
        "available",
        "zero",
        "filtered",
        "partial",
        "stale",
        "failed",
        "unavailable",
    }
)
_IDENTITY_HEADERS = frozenset(
    {
        "account",
        "account / owner",
        "account / next-action owner",
        "customer",
        "customer name",
        "team member",
    }
)
_TAC_SUBSTANTIVE_COLUMNS = (
    "Title",
    "Problem Description",
    "Problem Details",
    "Date/Time Opened",
    "Severity",
    "Highest Priority",
    "Service Tier",
    "Case Owner",
    "Current Contact Email",
    "Tech.",
    "Sub Technology",
)


def _clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _populated(series: pd.Series) -> pd.Series:
    values = series.fillna("").astype(str).str.strip()
    return values.ne("") & ~values.str.casefold().isin(_PLACEHOLDER_TOKENS | {"n/a", "unknown", "undefined"})


def _column_by_token(frame: pd.DataFrame, *tokens: str) -> Any:
    wanted = {re.sub(r"[^a-z0-9]", "", token.casefold()) for token in tokens}
    return next(
        (column for column in frame.columns if re.sub(r"[^a-z0-9]", "", str(column).casefold()) in wanted),
        None,
    )


def _attribution_values(value: Any) -> set[str]:
    token = _clean(value)
    if not token:
        return set()
    return {
        label.strip().casefold()
        for label in re.split(r"\s*(?:;|\|)\s*", token)
        if label.strip() and label.strip().casefold() not in {"unassigned / portfolio", "__legacy_unassigned__"}
    }


def _report_info_map(sheets: Mapping[str, pd.DataFrame]) -> dict[str, Any]:
    frame = sheets.get("Report_Info", pd.DataFrame())
    if not isinstance(frame, pd.DataFrame) or not {"Item", "Value"}.issubset(frame.columns):
        return {}
    return {_clean(row.get("Item")): row.get("Value") for _, row in frame.iterrows() if _clean(row.get("Item"))}


def audit_source_data_frames(
    sheets: Mapping[str, pd.DataFrame],
) -> dict[str, Any]:
    """Audit already-public workbook frames without recalculating metrics."""

    errors: list[str] = []
    undefined_cells: list[str] = []
    placeholder_cells: list[str] = []
    unknown_cells: dict[str, int] = {}
    source_link_rows = 0
    source_link_errors = 0
    attribution_mismatch_rows = 0
    attribution_mismatch_by_sheet: dict[str, dict[str, int]] = {}
    member_summary_mismatch = False
    tac_outside_window_rows = 0
    tac_missing_open_date_rows = 0

    for sheet_name, raw in sheets.items():
        if not isinstance(raw, pd.DataFrame):
            continue
        for column in raw.columns:
            values = raw[column].astype("string").fillna("").str.strip()
            undefined_count = int(values.str.contains(_UNDEFINED_RE, na=False).sum())
            if undefined_count:
                undefined_cells.append(f"{sheet_name}:{column}:{undefined_count}")
            placeholder_count = int(values.str.casefold().isin(_PLACEHOLDER_TOKENS).sum())
            if placeholder_count:
                placeholder_cells.append(f"{sheet_name}:{column}:{placeholder_count}")
            unknown_count = int(values.str.casefold().eq("unknown").sum())
            if unknown_count:
                unknown_cells[f"{sheet_name}:{column}"] = unknown_count

    if undefined_cells:
        errors.append("public workbook contains literal undefined placeholders: " + ", ".join(undefined_cells[:10]))
    if placeholder_cells:
        errors.append("public workbook contains serialized null placeholders: " + ", ".join(placeholder_cells[:10]))

    info = _report_info_map(sheets)
    for item in (
        "Report_Type",
        "Manager",
        "Scope_Type",
        "Scope_Value",
        "Days",
        "Evaluation_As_Of_UTC",
    ):
        token = _clean(info.get(item))
        if not token or token.casefold() in _PLACEHOLDER_TOKENS | {
            "unknown",
            "undefined",
            "n/a",
        }:
            errors.append(f"Report_Info lacks a substantive {item}")
    data_as_of = _clean(info.get("Data_As_Of_UTC"))
    data_as_of_state = _clean(info.get("Data_As_Of_State")).casefold()
    if data_as_of_state == "available" and not data_as_of:
        errors.append("Report_Info claims available freshness without Data_As_Of_UTC")
    if not data_as_of and data_as_of_state not in {
        "partial",
        "stale",
        "failed",
        "unavailable",
        "unknown",
    }:
        errors.append("Report_Info lacks an honest data-as-of state")
    for item, value in info.items():
        if not item.startswith("Source_State:"):
            continue
        state = _clean(value).casefold()
        if state not in _VALID_SOURCE_STATES:
            errors.append(f"Report_Info contains an invalid source state for {item}")

    action_plans = sheets.get("Action_Plans", pd.DataFrame())
    unexplained_status = 0
    unexplained_age = 0
    dishonest_title_fallback = 0
    if isinstance(action_plans, pd.DataFrame) and not action_plans.empty:
        quality = (
            action_plans.get(
                "AdoptIQ_Data_Quality",
                pd.Series("", index=action_plans.index, dtype="object"),
            )
            .fillna("")
            .astype(str)
        )
        buckets = (
            action_plans.get(
                "AdoptIQ_Status_Bucket",
                pd.Series("", index=action_plans.index, dtype="object"),
            )
            .fillna("")
            .astype(str)
        )
        age_bands = (
            action_plans.get(
                "AdoptIQ_Age_Band",
                pd.Series("", index=action_plans.index, dtype="object"),
            )
            .fillna("")
            .astype(str)
        )
        titles = (
            action_plans.get(
                "AdoptIQ_Title",
                pd.Series("", index=action_plans.index, dtype="object"),
            )
            .fillna("")
            .astype(str)
        )
        unexplained_status = int((buckets.eq("Unknown") & ~quality.str.contains("status", case=False)).sum())
        unexplained_age = int((age_bands.eq("Unknown") & ~quality.str.contains("created date", case=False)).sum())
        dishonest_title_fallback = int(
            (titles.eq("Title unavailable") & ~quality.str.contains("Missing title", case=False)).sum()
        )
        if unexplained_status:
            errors.append(f"Action_Plans has {unexplained_status} unexplained Unknown status row(s)")
        if unexplained_age:
            errors.append(f"Action_Plans has {unexplained_age} unexplained Unknown age row(s)")
        if dishonest_title_fallback:
            errors.append("Action_Plans uses Title unavailable without matching data-quality disclosure")

    tac = sheets.get("TAC_Cases", pd.DataFrame())
    tac_non_record_rows = 0
    tac_unexplained_case_type_rows = 0
    if isinstance(tac, pd.DataFrame) and not tac.empty:
        ids = tac.get("Record_ID", pd.Series("", index=tac.index, dtype="object")).fillna("").astype(str).str.strip()
        signal_count = pd.Series(0, index=tac.index, dtype="int64")
        for column in _TAC_SUBSTANTIVE_COLUMNS:
            if column in tac.columns:
                signal_count += _populated(tac[column]).astype("int64")
        tac_non_record_rows = int((ids.eq("") & signal_count.lt(3)).sum())
        if tac_non_record_rows:
            errors.append(f"TAC_Cases contains {tac_non_record_rows} non-record footer/summary row(s)")
        case_type_column = _column_by_token(tac, "case_type_class", "Case Type")
        case_quality_column = _column_by_token(
            tac,
            "case_type_data_quality",
            "Case Type Data Quality",
        )
        if case_type_column is not None:
            case_types = tac[case_type_column].fillna("").astype(str).str.strip().str.casefold()
            unresolved = case_types.isin({"", "unknown", "unclassified"})
            if case_quality_column is None:
                tac_unexplained_case_type_rows = int(unresolved.sum())
            else:
                quality = tac[case_quality_column].fillna("").astype(str).str.strip().str.casefold()
                explained = quality.str.contains(
                    r"not derivable|did not support|case type|source field|evidence",
                    regex=True,
                )
                tac_unexplained_case_type_rows = int((unresolved & ~explained).sum())
            if tac_unexplained_case_type_rows:
                errors.append(
                    "TAC_Cases has "
                    f"{tac_unexplained_case_type_rows} unresolved case-type row(s) without an exact source-data reason"
                )
        opened_column = _column_by_token(
            tac,
            "Date/Time Opened",
            "CREATED_DATE",
            "DATE_OPENED",
            "OPEN_DATE_C",
            "Created Date",
        )
        evaluation_clock = pd.to_datetime(
            info.get("Evaluation_As_Of_UTC"),
            errors="coerce",
            utc=True,
        )
        try:
            window_days = int(float(info.get("Days")))
        except (TypeError, ValueError):
            window_days = 0
        if opened_column is None:
            tac_missing_open_date_rows = len(tac)
        elif not pd.isna(evaluation_clock) and window_days > 0:
            opened = pd.to_datetime(tac[opened_column], errors="coerce", utc=True)
            tac_missing_open_date_rows = int(opened.isna().sum())
            cutoff = evaluation_clock - pd.Timedelta(window_days, unit="D")
            tac_outside_window_rows = int(
                (opened.notna() & ~opened.between(cutoff, evaluation_clock, inclusive="both")).sum()
            )
        if tac_missing_open_date_rows:
            errors.append(
                "TAC_Cases has "
                f"{tac_missing_open_date_rows} row(s) without a verifiable opened date"
            )
        if tac_outside_window_rows:
            errors.append(
                "TAC_Cases has "
                f"{tac_outside_window_rows} row(s) outside the selected report window"
            )

    for sheet_name in (
        "Action_Plans",
        "Adoption_Barriers",
        "Customer_Pulse",
        "Success_Priorities",
    ):
        frame = sheets.get(sheet_name, pd.DataFrame())
        if not isinstance(frame, pd.DataFrame):
            continue
        for _, row in frame.iterrows():
            record_id = _clean(row.get("Record_ID"))
            expected = build_source_record_url(sheet_name, record_id)
            actual = _clean(row.get(SOURCE_RECORD_URL_COLUMN))
            if expected:
                source_link_rows += 1
            if actual != expected or (actual and not is_allowed_source_record_url(actual)):
                source_link_errors += 1
    if source_link_errors:
        errors.append(f"Source Data has {source_link_errors} missing, unsafe, or mismatched CSConsole URL(s)")

    evidence = sheets.get("Evidence_Links", pd.DataFrame())
    evidence_link_errors = 0
    if isinstance(evidence, pd.DataFrame):
        for _, row in evidence.iterrows():
            source_sheet = _clean(row.get("Source_Sheet"))
            record_id = _clean(row.get("Record_ID"))
            expected = build_source_record_url(source_sheet, record_id)
            actual = _clean(row.get(SOURCE_RECORD_URL_COLUMN))
            if actual != expected:
                evidence_link_errors += 1
    if evidence_link_errors:
        errors.append(f"Evidence_Links has {evidence_link_errors} source-record URL mismatch(es)")

    subscriptions = sheets.get("Subscriptions", pd.DataFrame())
    account_attribution: dict[str, set[str]] = {}
    customer_scope_mismatch_rows = 0
    customer_scope_mismatch_by_sheet: dict[str, int] = {}
    if isinstance(subscriptions, pd.DataFrame) and not subscriptions.empty:
        account_column = _column_by_token(
            subscriptions,
            "ACCOUNT_ID_C",
            "Account ID",
        )
        attribution_column = _column_by_token(
            subscriptions,
            "Attributed_Team_Members",
            "Attributed Team Members",
        )
        legacy_type_column = _column_by_token(
            subscriptions,
            "Legacy_Record_Type",
        )
        for _, row in subscriptions.iterrows():
            if (
                legacy_type_column is not None
                and _clean(row.get(legacy_type_column)) == "Family-specific reported fact"
            ):
                continue
            account_id = _clean(row.get(account_column)).casefold() if account_column is not None else ""
            labels = _attribution_values(row.get(attribution_column)) if attribution_column is not None else set()
            if account_id and labels:
                account_attribution.setdefault(account_id, set()).update(labels)

    # A customer-scoped workbook must not contain a source or family-specific
    # fact for another customer/account.  This independently catches scope
    # leaks even when the Word and workbook agree with each other.  Stable
    # subscription account IDs are authoritative; normalized names provide a
    # defensive check for family-fact rows that have no account ID.
    scope_type = _clean(info.get("Scope_Type")).casefold()
    if scope_type == "customer":
        raw_scope_value = _clean(info.get("Scope_Value"))
        scope_value = (
            normalize_customer_name(raw_scope_value).casefold()
            if raw_scope_value
            else ""
        )
        scoped_accounts = set(account_attribution)
        for sheet_name in (
            "Subscriptions",
            "Action_Plans",
            "Adoption_Barriers",
            "Customer_Pulse",
            "TAC_Cases",
            "Success_Priorities",
        ):
            frame = sheets.get(sheet_name, pd.DataFrame())
            if not isinstance(frame, pd.DataFrame) or frame.empty:
                continue
            account_column = _column_by_token(
                frame,
                "ACCOUNT_ID_C",
                "ACCOUNT__C",
                "Account ID",
            )
            customer_column = _column_by_token(
                frame,
                "BU_NAME",
                "Customer Name",
                "customer_name",
                "Customer",
                "RELATED_CUSTOMER__C",
            )
            for _, row in frame.iterrows():
                account_id = (
                    _clean(row.get(account_column)).casefold()
                    if account_column is not None
                    else ""
                )
                raw_customer = (
                    _clean(row.get(customer_column))
                    if customer_column is not None
                    else ""
                )
                customer = (
                    normalize_customer_name(raw_customer).casefold()
                    if raw_customer
                    else ""
                )
                outside = bool(account_id and scoped_accounts and account_id not in scoped_accounts)
                if not outside and customer and scope_value:
                    # Round 168: alias-aware name compare (Round 132 SSoT).
                    # Literal normalize-and-equals treats NYU MEDICAL CENTER
                    # vs NYU LANGONE HEALTH SYSTEMS as two customers.
                    outside = not customer_names_match(raw_customer, raw_scope_value)
                if outside:
                    customer_scope_mismatch_rows += 1
                    customer_scope_mismatch_by_sheet[sheet_name] = (
                        customer_scope_mismatch_by_sheet.get(sheet_name, 0) + 1
                    )
        if customer_scope_mismatch_rows:
            errors.append(
                "Customer-scoped Source Data contains "
                f"{customer_scope_mismatch_rows} out-of-scope row(s)"
            )

    for sheet_name in (
        "Action_Plans",
        "Adoption_Barriers",
        "Customer_Pulse",
        "TAC_Cases",
        "Success_Priorities",
    ):
        frame = sheets.get(sheet_name, pd.DataFrame())
        if not isinstance(frame, pd.DataFrame) or frame.empty:
            continue
        account_column = _column_by_token(
            frame,
            "ACCOUNT_ID_C",
            "ACCOUNT__C",
            "Account ID",
        )
        attribution_column = _column_by_token(
            frame,
            "Attributed_Team_Members",
            "Attributed Team Members",
        )
        if account_column is None or attribution_column is None:
            continue
        for _, row in frame.iterrows():
            account_id = _clean(row.get(account_column)).casefold()
            expected = account_attribution.get(account_id, set())
            if not expected:
                continue
            actual = _attribution_values(row.get(attribution_column))
            # A source may carry one verified direct owner even when the
            # account is shared by multiple scoped members. That is valid.
            # Empty/unassigned attribution and any owner outside the scoped
            # subscription roster are not valid.
            if not actual or not actual.issubset(expected):
                attribution_mismatch_rows += 1
                reason = "missing" if not actual else "out_of_scope"
                sheet_counts = attribution_mismatch_by_sheet.setdefault(
                    sheet_name,
                    {"missing": 0, "out_of_scope": 0},
                )
                sheet_counts[reason] += 1
    if attribution_mismatch_rows:
        diagnostic = ", ".join(
            f"{sheet_name}="
            + "/".join(
                f"{reason}:{count}"
                for reason, count in sorted(counts.items())
                if count
            )
            for sheet_name, counts in sorted(attribution_mismatch_by_sheet.items())
        )
        errors.append(
            "Source Data has "
            f"{attribution_mismatch_rows} account-backed row(s) with missing or "
            "out-of-scope team attribution"
            + (f" ({diagnostic})" if diagnostic else "")
        )

    member_summary = sheets.get("Member_Summary", pd.DataFrame())
    expected_members = {member for members in account_attribution.values() for member in members}
    if scope_type in {"team", "member"} and expected_members:
        member_column = (
            _column_by_token(member_summary, "Team_Member", "Team Member")
            if isinstance(member_summary, pd.DataFrame)
            else None
        )
        actual_members = (
            {_clean(value).casefold() for value in member_summary[member_column] if _clean(value)}
            if member_column is not None
            else set()
        )
        member_summary_mismatch = actual_members != expected_members
        if member_summary_mismatch:
            errors.append("Member_Summary membership does not match the scoped subscription roster")

    return {
        "ok": not errors,
        "errors": errors,
        "undefined_cells": undefined_cells,
        "serialized_placeholder_cells": placeholder_cells,
        "unknown_cells": dict(sorted(unknown_cells.items())),
        "unexplained_action_plan_status_rows": unexplained_status,
        "unexplained_action_plan_age_rows": unexplained_age,
        "dishonest_title_fallback_rows": dishonest_title_fallback,
        "tac_non_record_rows": tac_non_record_rows,
        "tac_unexplained_case_type_rows": tac_unexplained_case_type_rows,
        "tac_outside_window_rows": tac_outside_window_rows,
        "tac_missing_open_date_rows": tac_missing_open_date_rows,
        "source_record_link_rows": source_link_rows,
        "source_record_link_errors": source_link_errors,
        "evidence_link_errors": evidence_link_errors,
        "attribution_mismatch_rows": attribution_mismatch_rows,
        "attribution_mismatch_by_sheet": {
            sheet_name: dict(sorted(counts.items()))
            for sheet_name, counts in sorted(attribution_mismatch_by_sheet.items())
        },
        "member_summary_mismatch": member_summary_mismatch,
        "customer_scope_mismatch_rows": customer_scope_mismatch_rows,
        "customer_scope_mismatch_by_sheet": dict(
            sorted(customer_scope_mismatch_by_sheet.items())
        ),
    }


def audit_written_source_hyperlinks(path: Any) -> dict[str, Any]:
    """Prove every published source URL cell has the matching xlsx relation."""

    from openpyxl import load_workbook  # noqa: PLC0415

    errors: list[str] = []
    url_cells = 0
    clickable_cells = 0
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        for worksheet in workbook.worksheets:
            headers = {cell.value: cell.column for cell in worksheet[1]}
            column = headers.get(SOURCE_RECORD_URL_COLUMN)
            if not column:
                continue
            for row_number in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row_number, column)
                url = _clean(cell.value)
                if not url:
                    continue
                url_cells += 1
                target_url = _clean(getattr(getattr(cell, "hyperlink", None), "target", ""))
                if not is_allowed_source_record_url(url) or target_url != url:
                    errors.append(f"{worksheet.title}!{cell.coordinate}")
                else:
                    clickable_cells += 1
    finally:
        workbook.close()
    return {
        "ok": not errors,
        "errors": errors,
        "url_cells": url_cells,
        "clickable_cells": clickable_cells,
    }


def audit_word_placeholders(document: Any) -> dict[str, Any]:
    """Reject renderer placeholders and unknown customer identities in Word."""

    errors: list[str] = []
    undefined_locations: list[str] = []
    generic_unknown_locations: list[str] = []
    unknown_identity_cells: list[str] = []
    for index, paragraph in enumerate(getattr(document, "paragraphs", ())):
        paragraph_text = str(getattr(paragraph, "text", "") or "")
        if _UNDEFINED_RE.search(paragraph_text):
            undefined_locations.append(f"paragraph:{index + 1}")
        if _GENERIC_UNKNOWN_RE.search(paragraph_text):
            generic_unknown_locations.append(f"paragraph:{index + 1}")
    for table_index, table in enumerate(getattr(document, "tables", ())):
        if not table.rows:
            continue
        headers = [_clean(cell.text).casefold() for cell in table.rows[0].cells]
        for row_index, row in enumerate(table.rows[1:], start=2):
            for column_index, cell in enumerate(row.cells):
                value = _clean(cell.text)
                if _UNDEFINED_RE.search(value):
                    undefined_locations.append(f"table:{table_index + 1}:row:{row_index}:column:{column_index + 1}")
                if _GENERIC_UNKNOWN_RE.search(value):
                    generic_unknown_locations.append(
                        f"table:{table_index + 1}:row:{row_index}:column:{column_index + 1}"
                    )
                header = headers[column_index] if column_index < len(headers) else ""
                if header in _IDENTITY_HEADERS and value.casefold() in {
                    "unknown",
                    "undefined",
                    "n/a",
                    "none",
                    "null",
                }:
                    unknown_identity_cells.append(f"table:{table_index + 1}:row:{row_index}:{header}")
    if undefined_locations:
        errors.append("Word contains literal undefined renderer placeholders")
    if generic_unknown_locations:
        errors.append("Word contains generic Unknown text instead of an exact source-state or data-quality reason")
    if unknown_identity_cells:
        errors.append("Word contains unknown customer/account/team-member identities")
    return {
        "ok": not errors,
        "errors": errors,
        "undefined_locations": undefined_locations,
        "generic_unknown_locations": generic_unknown_locations,
        "unknown_identity_cells": unknown_identity_cells,
    }


__all__ = [
    "audit_source_data_frames",
    "audit_word_placeholders",
    "audit_written_source_hyperlinks",
]
