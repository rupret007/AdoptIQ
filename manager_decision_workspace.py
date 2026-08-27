"""Manager Decision Workspace domain helpers.

Round 146 keeps workspace behavior outside the Flask monolith so report scope
previews, artifact summaries, history filtering, and comparisons share one
deterministic contract.  The module reads only already-generated, formula-free
Source Data workbooks; it never reaches Snowflake or an LLM.
"""

from __future__ import annotations

import copy
import hashlib
import json
import math
import re
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

from offline_validation_receipt import OfflineValidationReceiptStatus
from source_record_links import (
    SOURCE_RECORD_URL_COLUMN,
    build_source_record_url,
    is_allowed_source_record_url,
)


WORKSPACE_SCHEMA = "manager-decision-workspace/v1"
MAX_WORKBOOK_BYTES = 100 * 1024 * 1024
MAX_SHEET_ROWS = 50_000
MAX_PUBLIC_ITEMS = 250
MAX_EVIDENCE_RECORDS = 100

REPORT_TYPES: Mapping[str, Mapping[str, Any]] = {
    "leader": {
        "label": "Leader",
        "purpose": "Manager decisions with Team, Member, and Customer drill-down",
        "sources": ("Subscriptions", "Action Plans", "Adoption Barriers", "Customer Pulse", "TAC Cases"),
        "scopes": ("team", "member", "customer"),
    },
    "comprehensive": {
        "label": "Comprehensive",
        "purpose": "Broad portfolio decisions with complete paired source data",
        "sources": ("Subscriptions", "Action Plans", "Adoption Barriers", "Customer Pulse", "TAC Cases", "External Intelligence"),
        "scopes": ("team", "customer"),
    },
    "compact": {
        "label": "Compact",
        "purpose": "Fast executive briefing focused on material exceptions",
        "sources": ("Subscriptions", "Action Plans", "Adoption Barriers", "Customer Pulse", "TAC Cases"),
        "scopes": ("team", "customer"),
    },
    "renewal_portfolio": {
        "label": "Renewal portfolio",
        "purpose": "Portfolio renewal decisions and risk concentration",
        "sources": ("Subscriptions", "Action Plans", "Adoption Barriers", "Customer Pulse", "TAC Cases", "External Intelligence"),
        "scopes": ("team",),
    },
    "renewal": {
        "label": "Renewal customer",
        "purpose": "Deep renewal-risk explanation for one customer",
        "sources": ("Subscriptions", "Action Plans", "Adoption Barriers", "Customer Pulse", "TAC Cases", "External Intelligence"),
        "scopes": ("customer",),
    },
    "subscription": {
        "label": "Subscription analysis",
        "purpose": "Product adoption, barriers, support, and renewal evidence",
        "sources": ("Subscription", "Action Plans", "Adoption Barriers", "Customer Pulse", "TAC Cases"),
        "scopes": ("subscription",),
    },
}

_TERMINAL_STATUSES = {"completed", "success", "done"}
_COMPLETED_AP_STATUSES = {"completed", "complete", "closed", "done", "resolved", "cancelled", "canceled"}
_COMPARABLE_SOURCE_STATES = {"available", "zero"}
_SAFE_KEY_RE = re.compile(r"[^a-z0-9]+")
_EMAIL_RE = re.compile(r"(?P<email>[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})", re.IGNORECASE)
_EVIDENCE_KEY_RE = re.compile(r"^[A-Za-z0-9._:-]{1,300}$")
_OFFLINE_FIXTURE_PROVENANCE_WARNING = (
    "This report uses controlled local test data; no live source validation was performed."
)

# Only claims frozen by ``decision_report_delivery`` may cross into the web
# workspace.  Keeping this allow-list beside the projection prevents a future
# workbook row (or a tampered one) from becoming customer-facing copy merely
# because its key happens to start with ``insight.``.
_SUPPORTED_DECISION_INSIGHTS: Mapping[str, str] = {
    "insight.support_themes": "Support themes",
    "insight.support_operating_health": "Support operating health",
    "insight.window_momentum": "Momentum within this window",
    "insight.predictive_outlook_30d": "Predictive outlook (next 30 days)",
}


class EvidenceNotFoundError(LookupError):
    """The requested evidence key is not present in the verified workbook."""


class EvidenceIntegrityError(ValueError):
    """An evidence locator no longer matches its immutable workbook row."""


def _verified_source_record_url(
    *,
    source_sheet: object,
    record_id: object,
    evidence_link: Mapping[str, Any],
    source_record: Mapping[str, Any],
) -> str:
    """Return one canonical CSConsole drill-through URL or fail closed.

    The URL is duplicated deliberately in ``Evidence_Links`` and the exact
    source row.  Neither copy is trusted on its own: both must agree with each
    other and with the URL deterministically rebuilt from the verified sheet
    and record identity.  Unsupported source types legitimately return no
    link; a supplied URL for one is rejected.
    """

    linked_url = _text(evidence_link.get(SOURCE_RECORD_URL_COLUMN), 2_000)
    record_url = _text(source_record.get(SOURCE_RECORD_URL_COLUMN), 2_000)
    canonical_url = build_source_record_url(source_sheet, record_id)
    if linked_url != record_url:
        raise EvidenceIntegrityError("Evidence source-record URL does not match its exact source row.")
    if linked_url:
        if not canonical_url or linked_url != canonical_url or not is_allowed_source_record_url(linked_url):
            raise EvidenceIntegrityError("Evidence contains a noncanonical source-record URL.")
        return linked_url
    if canonical_url:
        raise EvidenceIntegrityError("Evidence is missing the canonical CSConsole source-record URL.")
    return ""


_CHART_METADATA: Mapping[str, tuple[str, str]] = {
    "activity_mix": (
        "Activity mix",
        "Complete source-record counts by activity type for the selected scope.",
    ),
    "action_plan_status_aging": (
        "Action Plan status and aging",
        "Mutually exclusive Action Plan lifecycle buckets at the report as-of time.",
    ),
    "risk_distribution": (
        "Customer risk distribution",
        "Customers grouped by the report's canonical risk band.",
    ),
    "activity_trend": (
        "Activity trend",
        "Dated source records grouped into the report's weekly periods.",
    ),
}

_SOURCE_LABELS: Mapping[str, str] = {
    "action_plans": "Action Plans",
    "adoption_barriers": "Adoption Barriers",
    "customer_pulse": "Customer Pulse",
    "success_priorities": "Success Priorities",
    "subscriptions": "Subscriptions",
    "tac_cases": "TAC Cases",
    "csone": "TAC Cases",
    "csone_tac": "TAC Cases",
    "csone_tac_cases": "TAC Cases",
    "external_incidents": "External incidents",
    "external_bugs": "External bugs",
}


def _text(value: object, limit: int = 1_000) -> str:
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    if text.casefold() in {"nan", "none", "nat"}:
        return ""
    return text[:limit]


def _frozen_text(value: object, *, limit: int) -> str:
    """Return an exact bounded workbook string or withhold it unchanged."""

    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        return ""
    if any(ord(character) < 32 and character not in {"\t", "\n", "\r"} for character in value):
        return ""
    return value


def _key(value: object) -> str:
    return _SAFE_KEY_RE.sub("_", _text(value).casefold()).strip("_")


def _number(value: object) -> int | float | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number):
        return None
    rounded = round(number, 6)
    return int(rounded) if rounded.is_integer() else rounded


def _as_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return _text(value).casefold() in {"1", "true", "yes", "y", "overdue"}


def _first(record: Mapping[str, Any], aliases: Iterable[str]) -> Any:
    folded = {_key(name): value for name, value in record.items()}
    for alias in aliases:
        value = folded.get(_key(alias))
        if _text(value):
            return value
    return None


def _info_value(info: Mapping[str, Any], aliases: Iterable[str]) -> str:
    folded = {_key(name): value for name, value in info.items()}
    for alias in aliases:
        value = folded.get(_key(alias))
        if _text(value):
            return _text(value, 4_000)
    return ""


def _public_label(value: object) -> str:
    normalized = _key(value)
    if normalized in _SOURCE_LABELS:
        return _SOURCE_LABELS[normalized]
    words = re.sub(r"[_-]+", " ", _text(value, 120)).strip()
    return words.title() if words else "A report source"


def _public_warning_copy(warning: object) -> str:
    """Convert internal warning shapes into bounded manager-facing copy."""

    source = "A report source"
    tokens: list[str] = []
    if isinstance(warning, Mapping):
        source = _public_label(
            warning.get("dataset")
            or warning.get("source")
            or warning.get("source_sheet")
        )
        tokens = [
            _text(warning.get(name), 1_000).casefold()
            for name in ("kind", "state", "source_state", "error", "effect", "detail")
            if _text(warning.get(name))
        ]
    else:
        raw_warning = _text(warning, 1_000)
        prefix = re.match(r"^([A-Za-z0-9_-]+)(?:\s*\(|\s*\||\s*:)", raw_warning)
        if prefix:
            source = _public_label(prefix.group(1))
        tokens = [raw_warning.casefold()]
    combined = " ".join(tokens)
    if not combined:
        return ""
    if "fixture" in combined or "local acceptance" in combined or "offline test" in combined:
        return "This report uses controlled local test data; no live source validation was performed."
    if "stale" in combined or "older than" in combined:
        return f"{source} may be older than expected; confirm timing before acting on time-sensitive details."
    if any(token in combined for token in ("failed", "unavailable", "fetch_error", "could not")):
        return f"{source} could not be retrieved; related metrics are shown as unavailable, not zero."
    if "scope_excluded" in combined or "filter" in combined or "excluded" in combined:
        return f"{source} was limited to records matching the selected scope; excluded records are not counted."
    if any(token in combined for token in ("partial", "truncat", "limit")):
        return f"{source} is only partially available; related totals may be incomplete."
    if "formula" in combined:
        return "This workbook contains formulas and is not eligible for canonical report comparison."
    return f"{source} has a recorded coverage limitation; review the available Source Data before acting."


def _source_warnings(values: Iterable[object]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        copy_value = _public_warning_copy(value)
        key = copy_value.casefold()
        if copy_value and key not in seen:
            result.append(copy_value)
            seen.add(key)
    return result


def _parse_utc(value: object) -> datetime | None:
    raw = _text(value, 120)
    if not raw:
        return None
    candidate = raw.replace(" UTC", "+00:00")
    if candidate.endswith("Z"):
        candidate = candidate[:-1] + "+00:00"
    try:
        parsed = datetime.fromisoformat(candidate)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _status_is_completed(value: object) -> bool:
    status_key = _key(value)
    completed_keys = {_key(item) for item in _COMPLETED_AP_STATUSES}
    return status_key in completed_keys or any(
        status_key.startswith(prefix)
        for prefix in ("completed_", "closed_", "resolved_", "cancelled_", "canceled_")
    )


def validate_workspace_selection(
    *,
    report_type: object,
    manager: object,
    technology: object,
    days: object,
    scope_type: object = "team",
    scope_value: object = "",
    subscription_id: object = "",
    allowed_managers: Iterable[str] = (),
    allowed_technologies: Iterable[str] = (),
) -> dict[str, Any]:
    """Validate and normalize one pre-generation workspace selection."""

    normalized_report = _text(report_type).casefold()
    if normalized_report not in REPORT_TYPES:
        raise ValueError("Unsupported report type.")
    normalized_scope = _text(scope_type or "team").casefold()
    allowed_scopes = set(REPORT_TYPES[normalized_report]["scopes"])
    if normalized_scope not in allowed_scopes:
        raise ValueError(
            f"{REPORT_TYPES[normalized_report]['label']} does not support the selected scope."
        )
    normalized_manager = _text(manager)
    manager_allowlist = {_text(item) for item in allowed_managers if _text(item)}
    if normalized_report not in {"renewal", "subscription"}:
        if not normalized_manager:
            raise ValueError("Select a manager.")
        if manager_allowlist and normalized_manager not in manager_allowlist:
            raise ValueError("The selected manager is not available.")
    normalized_technology = _text(technology or "All")
    tech_allowlist = {_text(item) for item in allowed_technologies if _text(item)}
    if tech_allowlist and normalized_technology not in tech_allowlist:
        raise ValueError("The selected technology is not available.")
    try:
        normalized_days = int(days)
    except (TypeError, ValueError) as exc:
        raise ValueError("Analysis days must be a whole number.") from exc
    if not 1 <= normalized_days <= 365:
        raise ValueError("Analysis days must be between 1 and 365.")
    normalized_value = _text(scope_value)
    normalized_subscription = _text(subscription_id)
    if normalized_scope in {"member", "customer"} and not normalized_value:
        raise ValueError(f"Select an individual {normalized_scope}.")
    if normalized_scope == "subscription" and not normalized_subscription:
        raise ValueError("Select a subscription.")
    return {
        "report_type": normalized_report,
        "report_label": REPORT_TYPES[normalized_report]["label"],
        "manager": normalized_manager,
        "technology": normalized_technology,
        "days": normalized_days,
        "scope_type": normalized_scope,
        "scope_value": normalized_value,
        "subscription_id": normalized_subscription,
    }


def build_scope_preview(
    selection: Mapping[str, Any],
    *,
    members: Sequence[Mapping[str, Any]] = (),
    customers: Sequence[Mapping[str, Any] | str] = (),
    source_mode: str = "live",
    source_state: str = "unknown",
    warning: str = "",
    data_as_of_utc: str = "",
) -> dict[str, Any]:
    """Build the concise, non-claiming scope preview shown before a run."""

    report_type = _text(selection.get("report_type")).casefold()
    spec = REPORT_TYPES[report_type]
    scope_type = _text(selection.get("scope_type") or "team").casefold()
    scope_value = _text(selection.get("scope_value"))
    manager = _text(selection.get("manager"))
    technology = _text(selection.get("technology") or "All")
    days = int(selection.get("days") or 90)
    member_count = len({
        _text(item.get("email") or item.get("value") or item.get("name")).casefold()
        for item in members
        if isinstance(item, Mapping) and _text(item.get("email") or item.get("value") or item.get("name"))
    })
    customer_values: set[str] = set()
    for item in customers:
        value = item.get("value") or item.get("label") if isinstance(item, Mapping) else item
        if _text(value):
            customer_values.add(_text(value).casefold())
    customer_count: int | None = len(customer_values) if customers else None
    if scope_type == "team":
        scope_label = f"{manager or 'Selected manager'} team"
    elif scope_type == "member":
        scope_label = scope_value or "Selected team member"
    elif scope_type == "customer":
        scope_label = scope_value or "Selected customer"
    else:
        scope_label = _text(selection.get("subscription_id")) or "Selected subscription"
    availability = _text(source_state or "unknown").casefold()
    state_message = {
        "available": "Scope data is available for this preview.",
        "partial": "The preview is partial; the generated report will name missing sources.",
        "stale": "The preview uses stale source context and will disclose that limitation.",
        "failed": "Scope data could not be reached; generation may fail honestly.",
        "unavailable": "Scope data is unavailable; no zero-count claim is being made.",
        "zero": "The selected scope is available and contains zero matching records.",
    }.get(availability, "Source coverage will be confirmed during generation.")
    source_mode_value = _text(source_mode or "live").casefold()
    is_fixture = source_mode_value.startswith("local") or "fixture" in source_mode_value
    limitations = [state_message]
    if warning:
        limitations.append(_text(warning, 500))
    if is_fixture:
        limitations.append("Controlled local fixture only; no live validation was performed.")
    return {
        "schema": WORKSPACE_SCHEMA,
        "report_type": report_type,
        "report_label": spec["label"],
        "purpose": spec["purpose"],
        "manager": manager,
        "technology": technology,
        "days": days,
        "scope_type": scope_type,
        "scope_value": scope_value,
        "scope_label": scope_label,
        "member_count": member_count if member_count else None,
        "customer_count": customer_count,
        "expected_sources": list(spec["sources"]),
        "source_mode": source_mode_value,
        "source_state": availability,
        "data_as_of_utc": _text(data_as_of_utc),
        "limitations": limitations,
        "deliverables": ["Decision-focused Word report", "Paired Source Data workbook"],
        "ask_ai_binding": {
            "manager": manager,
            "technology": technology,
            "days": days,
            "scope_type": scope_type,
            "scope_value": scope_value,
        },
        "live_validation_performed": False if is_fixture else None,
    }


def _sheet_records(worksheet: Any, *, limit: int = MAX_SHEET_ROWS) -> tuple[list[dict[str, Any]], int, int]:
    rows = worksheet.iter_rows(values_only=False)
    try:
        header_cells = next(rows)
    except StopIteration:
        return [], 0, 0
    headers = [_text(cell.value, 200) for cell in header_cells]
    records: list[dict[str, Any]] = []
    formula_count = sum(
        1 for cell in header_cells if getattr(cell, "data_type", "") == "f"
    )
    total_rows = 0
    for cells in rows:
        total_rows += 1
        formula_count += sum(1 for cell in cells if getattr(cell, "data_type", "") == "f")
        if len(records) >= limit:
            continue
        record = {
            header: cell.value
            for header, cell in zip(headers, cells)
            if header
        }
        if any(_text(value) for value in record.values()):
            records.append(record)
    return records, total_rows, formula_count


def _action_plan_projection(
    record: Mapping[str, Any],
    row_number: int,
    *,
    as_of_utc: object = "",
    source_sheet: str = "Action_Plans",
) -> dict[str, Any]:
    record_id = _text(_first(record, ("Record_ID", "ID", "Action Plan ID", "TASK_ID")), 240)
    status = _text(_first(record, ("AdoptIQ_Status_Bucket", "Status", "STATUS_C")), 120)
    due_date = _text(_first(record, ("AdoptIQ_Due_Date", "Due Date", "DUE_DATE_C", "TARGET_COMPLETION_DATE_C")), 80)
    customer = _text(_first(record, ("Customer", "BU_NAME", "Account", "Customer_Name")), 240)
    due_at = _parse_utc(due_date)
    as_of = _parse_utc(as_of_utc)
    completed = _status_is_completed(status)
    derived_overdue = bool(due_at and as_of and due_at < as_of and not completed)
    return {
        "record_id": record_id,
        "record_key": record_id.casefold() if record_id else f"missing-id-row-{row_number}",
        "record_id_quality": "stable" if record_id else "missing",
        "customer": customer,
        "customer_key": customer.casefold(),
        "title": _text(_first(record, ("Title", "SUBJECT_C", "Subject", "NAME")), 360) or "Title unavailable",
        "status": status or "Unknown",
        "status_key": _key(status or "unknown"),
        "owner": _text(_first(record, ("Owner", "OWNER_NAME", "OWNER_NAME_C", "ASSIGNEE_NAME_C", "OWNER_EMAIL_C", "Next Action Owner", "NEXT_ACTION_OWNER_C")), 240) or "Unassigned",
        "due_date": due_date,
        "priority": _text(_first(record, ("Priority", "PRIORITY_C", "AdoptIQ_Priority")), 80),
        "next_action": _text(_first(record, ("Next Action", "NEXT_ACTION_C", "NEXT_STEP_C")), 500),
        "is_overdue": (
            _as_bool(_first(record, ("AdoptIQ_Is_Overdue", "Is Overdue", "Overdue")))
            or _key(status) == "overdue"
            or derived_overdue
        ),
        "source_sheet": source_sheet,
    }


def _derived_projection_state(values: Iterable[object]) -> str:
    states = [
        _text(value, 80).casefold() or "unknown"
        for value in values
    ]
    if not states:
        return "unknown"
    if len(states) == 1 or len(set(states)) == 1:
        return states[0]
    if all(state in {"available", "zero"} for state in states):
        return "zero" if all(state == "zero" for state in states) else "available"
    if all(state in {"failed", "unavailable", "unknown"} for state in states):
        return "unavailable"
    return "partial"


def _risk_projection(
    record: Mapping[str, Any],
    *,
    source_states: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    customer = _text(_first(record, ("Account", "Customer_Name", "Customer", "BU_NAME")), 240)
    normalized_source_states = {
        _key(source): _text(state, 80).casefold() or "unknown"
        for source, state in (source_states or {}).items()
    }

    def field_state(
        state_columns: Sequence[str],
        contributing_sources: Sequence[str],
    ) -> str:
        declared = [
            _text(record.get(column), 80).casefold()
            for column in state_columns
            if _text(record.get(column), 80)
        ]
        fallback = [
            normalized_source_states[_key(source)]
            for source in contributing_sources
            if _key(source) in normalized_source_states
        ]
        if declared:
            return _derived_projection_state((*declared, *fallback))
        if source_states is None:
            # Retired workbook shapes predate field-level coverage metadata.
            # Their projection remains compatibility-only; canonical reports
            # always pass Report_Info source states and therefore fail closed.
            return "available"
        return _derived_projection_state(
            normalized_source_states.get(_key(source), "unknown")
            for source in contributing_sources
        )

    risk_sources = (
        "Subscriptions",
        "Action_Plans",
        "Adoption_Barriers",
        "Customer_Pulse",
        "TAC_Cases",
    )
    field_states = {
        "risk_band": field_state(("Risk_Band_Source_State",), risk_sources),
        "risk_score_0_100": field_state(
            ("Risk_Score_0_100_Source_State",), risk_sources
        ),
        "open_action_plans": field_state(
            ("Open_AP_Source_State",), ("Action_Plans",)
        ),
        "overdue_action_plans": field_state(
            ("Overdue_AP_Source_State",), ("Action_Plans",)
        ),
        "critical_high_barriers": field_state(
            ("Critical_High_Barriers_Source_State",), ("Adoption_Barriers",)
        ),
        "tac_cases": field_state(("TAC_Cases_Source_State",), ("TAC_Cases",)),
    }
    risk_score = _number(
        _first(
            record,
            (
                "Risk_Score_0_100",
                "Risk Score 0 100",
                "Overall_Risk_Score",
                "Risk_Score_0_10",
            ),
        )
    )
    if risk_score is not None and risk_score <= 10 and not _text(
        _first(record, ("Risk_Score_0_100", "Risk Score 0 100"))
    ):
        risk_score = round(float(risk_score) * 10, 1)
    complete = {"available", "zero"}
    return {
        "customer": customer,
        "customer_key": customer.casefold(),
        "risk_band": (
            _text(
                _first(
                    record,
                    ("Risk_Band", "Risk Band", "Risk_Level", "Risk Category"),
                ),
                80,
            ).upper()
            if field_states["risk_band"] in complete
            else ""
        ),
        "risk_score_0_100": (
            risk_score if field_states["risk_score_0_100"] in complete else None
        ),
        "open_action_plans": (
            _number(_first(record, ("Open_AP", "Open Action Plans")))
            if field_states["open_action_plans"] in complete
            else None
        ),
        "overdue_action_plans": (
            _number(_first(record, ("Overdue_AP", "Overdue Action Plans")))
            if field_states["overdue_action_plans"] in complete
            else None
        ),
        "critical_high_barriers": (
            _number(
                _first(
                    record,
                    (
                        "Critical_High_Barriers",
                        "Critical / High Barriers",
                        "Critical/High Barriers",
                    ),
                )
            )
            if field_states["critical_high_barriers"] in complete
            else None
        ),
        "tac_cases": (
            _number(
                _first(
                    record,
                    ("TAC_Cases", "TAC Cases", "Support_Cases", "Support Cases"),
                )
            )
            if field_states["tac_cases"] in complete
            else None
        ),
        "field_states": field_states,
        "source_sheet": "Account_Summary",
    }


def _stable_fingerprint(value: Mapping[str, Any]) -> str:
    raw = json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False, default=str).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _offline_receipt_claim_value(value: Any, *, ancestors: frozenset[int] = frozenset()) -> Any:
    """Freeze the complete JSON-like public value for hashing only."""

    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, Mapping):
        identity = id(value)
        if identity in ancestors:
            raise ValueError("public workspace projection contains a cycle")
        nested_ancestors = ancestors | {identity}
        return {
            str(raw_key): _offline_receipt_claim_value(
                value[raw_key],
                ancestors=nested_ancestors,
            )
            for raw_key in sorted(value, key=lambda item: str(item))
        }
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray)):
        identity = id(value)
        if identity in ancestors:
            raise ValueError("public workspace projection contains a cycle")
        nested_ancestors = ancestors | {identity}
        return [
            _offline_receipt_claim_value(item, ancestors=nested_ancestors)
            for item in value
        ]
    return str(value)


def finalize_public_workspace_snapshot(
    snapshot: Mapping[str, Any],
    *,
    word_download_url: str = "",
    source_data_download_url: str = "",
    ask_ai_report_url: str = "",
) -> dict[str, Any]:
    """Finish the path-free public workspace before receipt hashing."""

    result = dict(snapshot)
    result["downloads"] = {
        "word": _text(word_download_url, 2_000)
        if result.get("word_available")
        else "",
        "source_data": _text(source_data_download_url, 2_000)
        if result.get("excel_available")
        else "",
    }
    # Only manager-facing warning copy crosses the API boundary. Persisted
    # dictionaries can contain internal dataset/error tokens.
    result.pop("partial_data_warnings", None)
    result["ask_ai_binding"] = ask_ai_binding(result)
    result["ask_ai_url"] = (
        _text(ask_ai_report_url, 2_000)
        if result.get("ask_ai_binding_available") is not False
        else ""
    )
    return result


def offline_validation_projection_sha256(snapshot: Mapping[str, Any]) -> str:
    """Hash the complete path-free public workspace before receipt decoration.

    Private values are transient hash input; only the final SHA-256 is stored.
    Receipt/readiness fields are excluded because they are derived from this
    digest and the verified envelope. Internal warning dictionaries are removed
    by the route before this function is called.
    """

    derived_or_internal = {
        "_trusted_offline_validation_receipt",
        "customer_share_readiness",
        "customer_share_validation_receipt",
        "offline_validation_receipt",
        "partial_data_warnings",
    }
    public_snapshot = {
        key: value
        for key, value in snapshot.items()
        if key not in derived_or_internal and not str(key).startswith("_")
    }
    projection = {
        "schema": "adoptiq-offline-web-projection/v2",
        "public_snapshot": _offline_receipt_claim_value(public_snapshot),
    }
    return _stable_fingerprint(projection)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _metric_projection(
    metric_key: str,
    label: str,
    value: object,
    *,
    unit: str,
    source_state: str,
    source_sheet: str,
    provenance: str,
) -> dict[str, Any] | None:
    number = _number(value)
    if number is None:
        return None
    return {
        "metric_key": metric_key,
        "label": label,
        "value": number,
        "display_value": _text(number, 120),
        "unit": unit,
        "source_state": _text(source_state, 80).casefold() or "unknown",
        "source_sheet": source_sheet,
        "provenance": provenance,
    }


def _aggregate_source_state(values: Iterable[object]) -> str:
    states = [_text(value, 80).casefold() for value in values if _text(value)]
    if not states:
        return "unknown"
    order = {
        "available": 0,
        "zero": 1,
        "unknown": 2,
        "stale": 3,
        "partial": 4,
        "unavailable": 5,
        "failed": 6,
    }
    return max(states, key=lambda state: order.get(state, 2))


def _insight_evidence_source_state(
    metric_key: str,
    evidence_links: Sequence[Mapping[str, Any]],
) -> str:
    """Mirror each canonical insight family's source-state rule."""

    states = [
        _text(link.get("Source_State"), 80).casefold() or "unavailable"
        for link in evidence_links
    ]
    if not states:
        return "unavailable"
    if metric_key in {
        "insight.support_themes",
        "insight.support_operating_health",
    }:
        return _aggregate_source_state(states)
    if metric_key == "insight.predictive_outlook_30d" and any(
        _text(link.get("Source_Sheet"), 120) == "TAC_Cases"
        and _text(link.get("Source_State"), 80).casefold()
        in {"failed", "unavailable"}
        for link in evidence_links
    ):
        return "unavailable"
    if any(state not in _COMPARABLE_SOURCE_STATES for state in states):
        return "partial"
    return "zero" if all(state == "zero" for state in states) else "available"


def _insight_evidence_claim_matches(
    evidence_link: Mapping[str, Any],
    claim: str,
) -> bool:
    """Reconcile complete links exactly and incomplete links without invention."""

    link_claim = _frozen_text(evidence_link.get("Metric_Value"), limit=4_000)
    source_state = _text(evidence_link.get("Source_State"), 80).casefold()
    if source_state in _COMPARABLE_SOURCE_STATES:
        return link_claim == claim
    return not link_claim or link_claim == claim


def _chart_projection(
    records: Sequence[Mapping[str, Any]],
    evidence_by_key: Mapping[str, Mapping[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    evidence_by_key = evidence_by_key or {}
    grouped: dict[str, list[Mapping[str, Any]]] = {}
    for record in records:
        chart_id = _text(record.get("Chart_ID"), 120).casefold()
        metric_key = _text(record.get("Metric_Key"), 300)
        if not chart_id or not metric_key or chart_id.endswith("_coverage"):
            continue
        grouped.setdefault(chart_id, []).append(record)
    desired = list(_CHART_METADATA)
    ordered_ids = [item for item in desired if item in grouped]
    ordered_ids.extend(sorted(set(grouped) - set(ordered_ids)))
    charts: list[dict[str, Any]] = []
    for chart_id in ordered_ids:
        rows = grouped[chart_id]
        label, description = _CHART_METADATA.get(
            chart_id,
            (_public_label(chart_id), "Validated values from the paired Source Data workbook."),
        )
        points: list[dict[str, Any]] = []
        units = {_text(row.get("Unit"), 80) for row in rows if _text(row.get("Unit"))}
        for row in rows:
            category = _text(row.get("Category") or row.get("Series"), 240)
            period = _text(row.get("Period_Start"), 80)[:10]
            point_label = (
                " — ".join(part for part in (_text(row.get("Series"), 160), period) if part)
                if chart_id == "activity_trend"
                else category
            )
            value = _number(row.get("Value"))
            if not point_label or value is None:
                continue
            evidence_key = _text(row.get("Metric_Key"), 300)
            evidence_meta = evidence_by_key.get(evidence_key, {})
            points.append({
                "label": point_label,
                "value": value,
                "display_value": _text(value, 120),
                "metric_key": evidence_key,
                "evidence_key": evidence_key if evidence_meta else "",
                "evidence_count": int(evidence_meta.get("total_records") or 0),
                "source_state": _text(row.get("Source_State"), 80).casefold() or "unknown",
            })
        if not points:
            continue
        charts.append({
            "chart_key": chart_id,
            "label": label,
            "description": description,
            "unit": next(iter(units)) if len(units) == 1 else "",
            "source_state": _aggregate_source_state(point["source_state"] for point in points),
            "series": points[:MAX_PUBLIC_ITEMS],
            "provenance": "canonical Chart_Data",
        })
    return charts


def _record_is_unavailable(record: Mapping[str, Any]) -> bool:
    content = " ".join(_text(value, 240).casefold() for value in record.values())
    return any(
        marker in content
        for marker in ("data_unavailable", "data unavailable", "source unavailable", "fetch failed")
    )


def _record_is_placeholder(record: Mapping[str, Any]) -> bool:
    content = " ".join(_text(value, 240).casefold() for value in record.values())
    return any(
        marker in content
        for marker in (
            "no records in the analysis window",
            "no matching records",
            " - zero",
            " - empty",
        )
    )


def _usable_records(records: Sequence[Mapping[str, Any]]) -> list[Mapping[str, Any]]:
    return [
        record
        for record in records
        if not _as_bool(record.get("_adoptiq_provenance_row"))
        and _text(record.get("AdoptIQ_Status")).upper() != "EMPTY"
        and not _record_is_unavailable(record)
        and not _record_is_placeholder(record)
    ]


def _legacy_report_type(info: Mapping[str, Any]) -> str:
    export_type = _key(_info_value(info, ("Report_Type", "Export type", "Export_Type")))
    if "compact" in export_type:
        return "compact"
    if "renewal_portfolio" in export_type:
        return "renewal_portfolio"
    if "renewal" in export_type:
        return "renewal"
    if "subscription" in export_type:
        return "subscription"
    if "comprehensive" in export_type:
        return "comprehensive"
    if "leader" in export_type:
        return "leader"
    return export_type


def _legacy_workbook_projection(
    *,
    info: Mapping[str, Any],
    sheet_names: Sequence[str],
    read: Any,
) -> dict[str, Any]:
    """Project retired workbook shapes without claiming canonical parity."""

    cache: dict[str, list[dict[str, Any]]] = {}

    def records(name: str) -> list[dict[str, Any]]:
        if name not in cache:
            cache[name] = read(name, MAX_SHEET_ROWS)
        return cache[name]

    def first_sheet(names: Sequence[str]) -> tuple[str, list[dict[str, Any]]]:
        for name in names:
            if name in sheet_names:
                return name, records(name)
        return "", []

    report_type = _legacy_report_type(info)
    as_of_utc = _info_value(
        info,
        ("Data_As_Of_UTC", "Report_Generated_At_UTC", "Generated at (UTC)"),
    )
    action_sheet, raw_actions = first_sheet(("Action_Plans", "Customer_Action_Plans"))
    action_records = _usable_records(raw_actions)
    action_plans = [
        _action_plan_projection(
            record,
            index,
            as_of_utc=as_of_utc,
            source_sheet=action_sheet or "Action_Plans",
        )
        for index, record in enumerate(action_records, 2)
    ]
    action_plans.sort(
        key=lambda item: (
            not item["is_overdue"],
            item["due_date"] or "9999-12-31",
            item["customer"].casefold(),
            item["record_key"],
        )
    )

    barrier_sheet, raw_barriers = first_sheet(
        ("All_Adoption_Barriers", "Customer_Adoption_Barriers", "Adoption_Barriers")
    )
    tac_sheet, raw_tac = first_sheet(
        ("All_Support_Cases", "Customer_Support_Cases", "TAC_Cases", "Escalated_Cases")
    )
    pulse_sheet, raw_pulse = first_sheet(("Customer_Customer_Pulse", "Customer_Pulse"))
    priority_sheet, raw_priorities = first_sheet(
        ("Customer_Success_Priorities", "Success_Priorities")
    )
    risk_sheet, raw_risk = first_sheet(("Account_Summary", "Renewal_Summary", "Risk_Summary"))
    barriers = _usable_records(raw_barriers)
    tac_cases = _usable_records(raw_tac)
    pulse = _usable_records(raw_pulse)
    priorities = _usable_records(raw_priorities)
    risk_records = _usable_records(raw_risk)
    accounts = [item for item in (_risk_projection(record) for record in risk_records) if item["customer"]]

    def state(sheet_name: str, raw: Sequence[Mapping[str, Any]], usable: Sequence[Mapping[str, Any]]) -> str:
        if not sheet_name:
            return "unknown"
        if raw and not usable:
            return "unavailable" if any(_record_is_unavailable(item) for item in raw) else "zero"
        return "available" if usable else "zero"

    source_states = {
        "Action_Plans": state(action_sheet, raw_actions, action_records),
        "Adoption_Barriers": state(barrier_sheet, raw_barriers, barriers),
        "Customer_Pulse": state(pulse_sheet, raw_pulse, pulse),
        "Success_Priorities": state(priority_sheet, raw_priorities, priorities),
        "TAC_Cases": state(tac_sheet, raw_tac, tac_cases),
    }

    by_customer_actions: dict[str, list[dict[str, Any]]] = {}
    for item in action_plans:
        if item["customer_key"]:
            by_customer_actions.setdefault(item["customer_key"], []).append(item)
    tac_by_customer: dict[str, int] = {}
    for record in tac_cases:
        customer = _text(_first(record, ("Customer", "Customer Name", "Customer_Name", "BU_NAME")), 240)
        if customer:
            tac_by_customer[customer.casefold()] = tac_by_customer.get(customer.casefold(), 0) + 1
    for account in accounts:
        related_actions = by_customer_actions.get(account["customer_key"], [])
        if account["open_action_plans"] is None:
            account["open_action_plans"] = sum(
                not _is_completed_status(item.get("status")) for item in related_actions
            )
        if account["overdue_action_plans"] is None:
            account["overdue_action_plans"] = sum(bool(item.get("is_overdue")) for item in related_actions)
        if account["tac_cases"] is None:
            account["tac_cases"] = tac_by_customer.get(account["customer_key"], 0)
        account["source_sheet"] = risk_sheet
    accounts.sort(
        key=lambda item: (-(float(item["risk_score_0_100"] or 0)), item["customer"].casefold())
    )

    dashboard_sheet, dashboard_rows = first_sheet(("Executive_Dashboard", "Summary", "Key_Metrics"))
    dashboard_values: dict[str, object] = {}
    for record in dashboard_rows:
        label = _first(record, ("Metric", "Item", "Name"))
        if _text(label):
            dashboard_values[_key(label)] = _first(record, ("Value", "Metric_Value"))

    customers = {account["customer_key"] for account in accounts if account["customer_key"]}
    for collection in (action_records, barriers, tac_cases, pulse, priorities):
        for record in collection:
            customer = _text(
                _first(record, ("Customer", "Customer Name", "Customer_Name", "BU_NAME", "RELATED_CUSTOMER__C")),
                240,
            )
            if customer:
                customers.add(customer.casefold())
    customer_count: object = len(customers) if customers else (
        dashboard_values.get("total_customers_analyzed")
        or dashboard_values.get("customers")
        or dashboard_values.get("customer_count")
    )
    open_actions = [item for item in action_plans if not _is_completed_status(item.get("status"))]
    overdue_actions = [item for item in open_actions if item.get("is_overdue")]
    high_risk = [
        item for item in accounts
        if _key(item.get("risk_band")) in {"critical", "high"}
    ]
    def count_when_known(count: int, source_state: str) -> int | None:
        return count if source_state in {"available", "zero", "partial", "stale"} else None

    metric_inputs = (
        ("kpi.customers", "Customers", customer_count, "records", "available" if customers else "unknown", risk_sheet or dashboard_sheet),
        ("kpi.action_plans_total", "Action Plans", count_when_known(len(action_plans), source_states["Action_Plans"]), "records", source_states["Action_Plans"], action_sheet),
        ("kpi.action_plans_open", "Open Action Plans", count_when_known(len(open_actions), source_states["Action_Plans"]), "records", source_states["Action_Plans"], action_sheet),
        ("kpi.action_plans_overdue", "Overdue Action Plans", count_when_known(len(overdue_actions), source_states["Action_Plans"]), "records", source_states["Action_Plans"], action_sheet),
        ("kpi.adoption_barriers", "Adoption Barriers", count_when_known(len(barriers), source_states["Adoption_Barriers"]), "records", source_states["Adoption_Barriers"], barrier_sheet),
        ("kpi.tac_cases", "TAC Cases", count_when_known(len(tac_cases), source_states["TAC_Cases"]), "records", source_states["TAC_Cases"], tac_sheet),
        ("kpi.high_risk_customers", "High-risk customers", len(high_risk) if accounts else None, "customers", "available" if accounts else "unknown", risk_sheet),
    )
    metrics = [
        metric
        for values in metric_inputs
        if (metric := _metric_projection(
            values[0], values[1], values[2], unit=values[3],
            source_state=values[4], source_sheet=values[5],
            provenance="compatibility projection from named Source Data sheets",
        )) is not None
    ]

    activity_points = [
        {"label": label, "value": count, "display_value": str(count), "source_state": source_states[source]}
        for label, count, source in (
            ("Action Plans", len(action_plans), "Action_Plans"),
            ("Adoption Barriers", len(barriers), "Adoption_Barriers"),
            ("Customer Pulse", len(pulse), "Customer_Pulse"),
            ("TAC Cases", len(tac_cases), "TAC_Cases"),
        )
        if source_states[source] not in {"unknown", "unavailable", "failed"}
    ]
    action_buckets = [
        ("Open", sum(not item.get("is_overdue") for item in open_actions)),
        ("Overdue", len(overdue_actions)),
        ("Completed", sum(_is_completed_status(item.get("status")) for item in action_plans)),
    ]
    risk_bands: dict[str, int] = {}
    for account in accounts:
        label = _text(account.get("risk_band"), 80).upper() or "UNKNOWN"
        risk_bands[label] = risk_bands.get(label, 0) + 1
    charts: list[dict[str, Any]] = []
    if activity_points:
        charts.append({
            "chart_key": "activity_mix",
            "label": _CHART_METADATA["activity_mix"][0],
            "description": _CHART_METADATA["activity_mix"][1],
            "unit": "records",
            "source_state": _aggregate_source_state(point["source_state"] for point in activity_points),
            "series": activity_points,
            "provenance": "compatibility projection from named Source Data sheets",
        })
    if action_sheet and source_states["Action_Plans"] not in {"unavailable", "failed"}:
        charts.append({
            "chart_key": "action_plan_status_aging",
            "label": _CHART_METADATA["action_plan_status_aging"][0],
            "description": _CHART_METADATA["action_plan_status_aging"][1],
            "unit": "records",
            "source_state": source_states["Action_Plans"],
            "series": [
                {"label": label, "value": value, "display_value": str(value), "source_state": source_states["Action_Plans"]}
                for label, value in action_buckets
            ],
            "provenance": "compatibility projection from Action Plan rows",
        })
    if risk_bands:
        charts.append({
            "chart_key": "risk_distribution",
            "label": _CHART_METADATA["risk_distribution"][0],
            "description": _CHART_METADATA["risk_distribution"][1],
            "unit": "customers",
            "source_state": "available",
            "series": [
                {"label": label.title(), "value": value, "display_value": str(value), "source_state": "available"}
                for label, value in sorted(risk_bands.items())
            ],
            "provenance": f"compatibility projection from {risk_sheet}",
        })

    customer_name = _info_value(info, ("Customer Name", "Customer_Name"))
    subscription_id = _info_value(info, ("Subscription ID", "Subscription_Id"))
    scope_type = "subscription" if subscription_id else (
        "customer"
        if report_type == "renewal" and customer_name and "portfolio" not in customer_name.casefold()
        else "team"
    )
    scope_value = subscription_id if subscription_id else (customer_name if scope_type == "customer" else "")
    days = _number(_info_value(info, ("Days", "Analysis Period (Days)", "Analysis_Period")))
    return {
        "report_type": report_type,
        "manager": _info_value(info, ("Manager",)),
        "technology": _info_value(info, ("Technology",)),
        "scope_type": scope_type,
        "scope_value": scope_value,
        "days": days,
        "data_as_of_utc": as_of_utc,
        "metrics": sorted(metrics, key=lambda item: item["metric_key"]),
        "decision_insights": [],
        "decision_insight_integrity": "legacy_unavailable",
        "source_states": dict(sorted(source_states.items())),
        "action_plans": action_plans,
        "accounts": accounts,
        "charts": charts,
        "evidence_available": False,
        "evidence_notice": (
            "This older compatibility workbook does not contain exact Evidence_Links. "
            "Generate a new report to enable source-record drill-down."
        ),
        "evidence_manifest": [],
        "evidence_contract": "unavailable",
        "canonical_snapshot": False,
        "snapshot_contract": "legacy-compatibility-projection/v1",
        "compatibility_source_sheets": sorted(cache),
    }


def _evidence_manifest_projection(
    records: Sequence[Mapping[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    """Project row-level evidence links without exposing source-row contents."""

    grouped: dict[str, dict[str, Any]] = {}
    locators: dict[str, set[tuple[str, int]]] = {}
    for record in records:
        evidence_key = _text(record.get("Evidence_Key"), 300)
        if not evidence_key:
            continue
        item = grouped.setdefault(
            evidence_key,
            {
                "evidence_key": evidence_key,
                "evidence_type": _text(record.get("Evidence_Type"), 80),
                "label": _text(record.get("Display_Label"), 300) or evidence_key,
                "source_state": _text(record.get("Source_State"), 80).casefold() or "unknown",
                "metric_value": _number(record.get("Metric_Value")),
                "unit": _text(record.get("Unit"), 80),
                "sources": [],
                "total_records": 0,
                "has_derivation_state": False,
            },
        )
        source = _text(record.get("Source_Sheet"), 120)
        if source and source not in item["sources"]:
            item["sources"].append(source)
        item["source_state"] = _aggregate_source_state(
            (item.get("source_state"), record.get("Source_State"))
        )
        row_number = _number(record.get("Source_Row_Number"))
        if row_number is None:
            item["has_derivation_state"] = True
            continue
        locator = (source, int(row_number))
        if locator not in locators.setdefault(evidence_key, set()):
            locators[evidence_key].add(locator)
            item["total_records"] += 1
    manifest = []
    for evidence_key in sorted(grouped):
        item = grouped[evidence_key]
        item["sources"] = sorted(item["sources"])
        manifest.append(item)
    return manifest[:1_000], {item["evidence_key"]: item for item in manifest}


def _canonical_workbook_projection(
    *,
    info: Mapping[str, Any],
    read: Any,
) -> dict[str, Any]:
    evidence_records = read("Evidence_Links", MAX_SHEET_ROWS)
    evidence_manifest, evidence_by_key = _evidence_manifest_projection(evidence_records)
    evidence_records_by_key: dict[str, list[Mapping[str, Any]]] = {}
    for evidence_record in evidence_records:
        evidence_key = _text(evidence_record.get("Evidence_Key"), 300)
        if evidence_key:
            evidence_records_by_key.setdefault(evidence_key, []).append(evidence_record)
    lineage_records = read("Metric_Lineage", 10_000)
    metrics: list[dict[str, Any]] = []
    insight_candidates: dict[str, dict[str, Any]] = {}
    seen_insight_keys: set[str] = set()
    duplicate_insight_keys: set[str] = set()
    invalid_insight_keys: set[str] = set()
    for record in lineage_records:
        metric_key = _text(record.get("Metric_Key"), 300)
        if metric_key in _SUPPORTED_DECISION_INSIGHTS:
            if metric_key in seen_insight_keys:
                duplicate_insight_keys.add(metric_key)
                continue
            seen_insight_keys.add(metric_key)
            evidence_meta = evidence_by_key.get(metric_key, {})
            evidence_links = evidence_records_by_key.get(metric_key, [])
            source_sheets = [
                sheet.strip()
                for sheet in _text(record.get("Source_Sheet"), 1_000).split(";")
                if sheet.strip()
            ]
            claim = _frozen_text(record.get("Metric_Value"), limit=4_000)
            if not claim or not source_sheets:
                invalid_insight_keys.add(metric_key)
                continue
            raw_caveat = record.get("Caveat")
            caveat = (
                ""
                if raw_caveat is None or not str(raw_caveat).strip()
                else _frozen_text(raw_caveat, limit=1_000)
            )
            lineage_state = (
                _text(record.get("Source_State"), 80).casefold() or "unknown"
            )
            evidence_sources = {
                _text(link.get("Source_Sheet"), 120)
                for link in evidence_links
                if _text(link.get("Source_Sheet"), 120)
            }
            evidence_state = _insight_evidence_source_state(
                metric_key,
                evidence_links,
            )
            evidence_binding_valid = (
                bool(evidence_meta)
                and bool(evidence_links)
                and all(_key(link.get("Evidence_Type")) == "insight" for link in evidence_links)
                and all(
                    _insight_evidence_claim_matches(link, claim)
                    for link in evidence_links
                )
                and evidence_sources == set(source_sheets)
                and evidence_state == lineage_state
                and (lineage_state != "available" or int(evidence_meta.get("total_records") or 0) > 0)
                and (raw_caveat is None or not str(raw_caveat).strip() or bool(caveat))
            )
            if not evidence_binding_valid:
                invalid_insight_keys.add(metric_key)
                continue
            insight_candidates[metric_key] = {
                "insight_key": metric_key,
                "evidence_key": metric_key,
                "evidence_count": int(evidence_meta.get("total_records") or 0),
                "label": (
                    _text(record.get("Display_Label"), 240)
                    or _SUPPORTED_DECISION_INSIGHTS[metric_key]
                ),
                "claim": claim,
                "caveat": caveat,
                "source_state": lineage_state,
                "source_sheets": source_sheets,
                "provenance": "canonical Metric_Lineage",
            }
            continue
        if not metric_key.startswith("kpi."):
            continue
        evidence_meta = evidence_by_key.get(metric_key, {})
        metrics.append({
            "metric_key": metric_key,
            "evidence_key": metric_key if evidence_meta else "",
            "evidence_count": int(evidence_meta.get("total_records") or 0),
            "label": _text(record.get("Display_Label"), 240) or metric_key,
            "value": _number(record.get("Metric_Value")),
            "display_value": _text(record.get("Metric_Value"), 120),
            "unit": _text(record.get("Unit"), 80),
            "source_state": _text(record.get("Source_State"), 80).casefold() or "unknown",
            "source_sheet": _text(record.get("Source_Sheet"), 300),
            "provenance": "canonical Metric_Lineage",
        })
    metrics.sort(key=lambda item: item["metric_key"])
    insights = [
        insight_candidates[key]
        for key in _SUPPORTED_DECISION_INSIGHTS
        if key in insight_candidates and key not in duplicate_insight_keys
    ]

    source_states = {
        key.split(":", 1)[1]: value.casefold()
        for key, value in info.items()
        if key.startswith("Source_State:")
    }
    if not source_states:
        for metric in metrics:
            for source_sheet in metric["source_sheet"].split(";"):
                source_sheet = source_sheet.strip()
                if source_sheet and source_sheet not in source_states:
                    source_states[source_sheet] = metric["source_state"]

    as_of_utc = _info_value(info, ("Data_As_Of_UTC",))
    data_as_of_state = (
        _info_value(info, ("Data_As_Of_State",)).casefold()
        or ("available" if as_of_utc else "unavailable")
    )
    evaluation_as_of_utc = _info_value(info, ("Evaluation_As_Of_UTC",))
    retrieval_attempted_at_utc = _info_value(
        info, ("Retrieval_Attempted_At_UTC",)
    )
    if data_as_of_state not in {"available", "zero"}:
        source_states["Data_Freshness"] = data_as_of_state
    action_records = read("Action_Plans", MAX_SHEET_ROWS)
    action_plans = [
        _action_plan_projection(record, index, as_of_utc=as_of_utc)
        for index, record in enumerate(action_records, 2)
        if not _as_bool(record.get("_adoptiq_provenance_row"))
        and _text(record.get("AdoptIQ_Status")).upper() != "EMPTY"
    ]
    action_plans.sort(
        key=lambda item: (
            not item["is_overdue"],
            item["due_date"] or "9999-12-31",
            item["customer"].casefold(),
            item["record_key"],
        )
    )
    action_evidence_by_id = {
        _text(record.get("Record_ID"), 240).casefold(): _text(record.get("Evidence_Key"), 300)
        for record in evidence_records
        if _key(record.get("Evidence_Type")) == "action_plan"
        and _text(record.get("Record_ID"))
        and _text(record.get("Evidence_Key"))
    }
    for item in action_plans:
        item["evidence_key"] = action_evidence_by_id.get(
            _text(item.get("record_id"), 240).casefold(), ""
        )

    account_records = read("Account_Summary", MAX_SHEET_ROWS)
    accounts = [
        _risk_projection(record, source_states=source_states)
        for record in account_records
    ]
    accounts = [item for item in accounts if item["customer"]]
    accounts.sort(
        key=lambda item: (
            item["risk_score_0_100"] is None,
            -(float(item["risk_score_0_100"] or 0)),
            item["customer"].casefold(),
        )
    )
    account_evidence_by_label = {
        _text(record.get("Display_Label"), 240).casefold(): _text(record.get("Evidence_Key"), 300)
        for record in evidence_records
        if _key(record.get("Evidence_Type")) == "account"
        and _text(record.get("Display_Label"))
        and _text(record.get("Evidence_Key"))
    }
    for item in accounts:
        item["evidence_key"] = account_evidence_by_label.get(
            _text(item.get("customer"), 240).casefold(), ""
        )
    scope_type = _info_value(info, ("Scope_Type",)) or "team"
    scope_value = _info_value(info, ("Scope_Value",))
    email_match = _EMAIL_RE.search(scope_value) if _key(scope_type) == "member" else None
    return {
        "report_type": _info_value(info, ("Report_Type",)).casefold(),
        "manager": _info_value(info, ("Manager",)),
        "technology": _info_value(info, ("Technology",)),
        "scope_type": scope_type.casefold(),
        "scope_value": scope_value,
        "scope_member": email_match.group("email").casefold() if email_match else "",
        "days": _number(_info_value(info, ("Days",))),
        "data_as_of_utc": as_of_utc,
        "data_as_of_state": data_as_of_state,
        "retrieval_attempted_at_utc": retrieval_attempted_at_utc,
        "evaluation_as_of_utc": evaluation_as_of_utc,
        "metrics": metrics,
        "decision_insights": insights,
        "decision_insight_integrity": (
            "duplicate_keys"
            if duplicate_insight_keys
            else "evidence_mismatch"
            if invalid_insight_keys
            else "verified_shape"
        ),
        "source_states": dict(sorted(source_states.items())),
        "action_plans": action_plans,
        "accounts": accounts,
        "charts": _chart_projection(read("Chart_Data", MAX_SHEET_ROWS), evidence_by_key),
        "evidence_available": bool(evidence_records and evidence_manifest),
        "evidence_notice": (
            "Canonical evidence links are present. Each drill-down is reverified against the paired Source Data workbook before records are displayed."
            if evidence_records and evidence_manifest
            else "This canonical workbook predates record-level evidence links; generate a new report to enable drill-down."
        ),
        "evidence_manifest": evidence_manifest,
        "evidence_contract": "canonical-evidence-links/v1" if evidence_records else "unavailable",
        "canonical_snapshot": True,
        "snapshot_contract": "canonical-source-data/v1",
    }


def _load_workbook_snapshot(path_text: str, mtime_ns: int, size: int) -> dict[str, Any]:
    del mtime_ns
    path = Path(path_text)
    if size > MAX_WORKBOOK_BYTES:
        raise ValueError("Source Data workbook exceeds the workspace size limit.")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet_names = list(workbook.sheetnames)
        shapes: dict[str, dict[str, int]] = {}
        formulas = 0

        record_cache: dict[str, list[dict[str, Any]]] = {}

        def read(name: str, limit: int = MAX_SHEET_ROWS) -> list[dict[str, Any]]:
            nonlocal formulas
            if name not in workbook.sheetnames:
                return []
            if name not in record_cache:
                records, total_rows, formula_count = _sheet_records(
                    workbook[name], limit=limit
                )
                formulas += formula_count
                shapes[name] = {
                    "rows": total_rows,
                    "columns": int(getattr(workbook[name], "max_column", 0) or 0),
                }
                record_cache[name] = records
            return record_cache[name][:limit]

        info_records = read("Report_Info", 2_000)
        info: dict[str, str] = {}
        for record in info_records:
            item = _text(_first(record, ("Item", "Field", "Key")), 300)
            if item:
                info[item] = _text(_first(record, ("Value", "Metric_Value")), 4_000)

        canonical = all(
            name in sheet_names
            for name in ("Report_Info", "Metric_Lineage", "Chart_Data", "Action_Plans", "Account_Summary")
        ) and bool(info.get("Fact_Contract_SHA256"))
        public_core = (
            _canonical_workbook_projection(info=info, read=read)
            if canonical
            else _legacy_workbook_projection(
                info=info,
                sheet_names=sheet_names,
                read=read,
            )
        )
        # Inspect formula cells in every remaining sheet. Formula-free output
        # is a canonical Source Data requirement and must not depend on which
        # cards happened to be rendered in the workspace.
        for sheet_name in sheet_names:
            if sheet_name not in record_cache:
                read(sheet_name, 0)

        warning_inputs: list[object] = []
        for record in info_records:
            item = _text(_first(record, ("Item", "Field", "Key")), 300)
            item_key = _key(item)
            if item_key == "partial_data_warning" or re.fullmatch(
                r"partial_data_warning_\d+", item_key
            ):
                warning_value = _first(record, ("Value", "Metric_Value"))
                warning_detail = _first(record, ("Detail", "Description"))
                warning_inputs.append(
                    {"source": warning_value, "detail": warning_detail}
                    if _text(warning_detail)
                    else warning_value
                )
        if formulas:
            warning_inputs.append("formula")
        warnings = _source_warnings(warning_inputs)
        if not canonical:
            compatibility_copy = (
                "This older workbook is shown through a compatibility projection; "
                "generate a new report before using canonical comparison."
            )
            if compatibility_copy not in warnings:
                warnings.append(compatibility_copy)
        fingerprint = info.get("Fact_Contract_SHA256") or _stable_fingerprint(public_core)
        action_plans = list(public_core.get("action_plans") or [])
        return {
            "schema": WORKSPACE_SCHEMA,
            **public_core,
            "fact_fingerprint": fingerprint,
            "workbook_sha256": _file_sha256(path),
            "sheet_names": sheet_names,
            "sheet_shapes": shapes,
            "formula_cells": formulas,
            "warnings": warnings,
            "source_warnings": warnings,
            "action_plan_count": len(action_plans),
            "missing_action_plan_id_count": sum(1 for item in action_plans if item["record_id_quality"] == "missing"),
        }
    finally:
        workbook.close()


@lru_cache(maxsize=64)
def _cached_snapshot(path_text: str, mtime_ns: int, size: int) -> dict[str, Any]:
    return _load_workbook_snapshot(path_text, mtime_ns, size)


def load_workbook_snapshot(path: str | Path) -> dict[str, Any]:
    """Read a bounded Source Data snapshot, cached by path/mtime/size."""

    resolved = Path(path).expanduser().resolve()
    if resolved.suffix.casefold() != ".xlsx" or not resolved.is_file():
        raise ValueError("A readable Source Data .xlsx file is required.")
    stat = resolved.stat()
    return copy.deepcopy(_cached_snapshot(str(resolved), stat.st_mtime_ns, stat.st_size))


@lru_cache(maxsize=64)
def _cached_evidence_workbook_integrity(
    path_text: str,
    mtime_ns: int,
    size: int,
    expected_fingerprint: str,
) -> dict[str, Any]:
    """Recompute the serialized sheet digests behind Evidence_Links."""

    del mtime_ns, size
    import pandas as pd  # noqa: PLC0415
    from decision_report_delivery import (  # noqa: PLC0415
        _SOURCE_CONTRACT_SHEETS,
        _frame_content_digest,
        validate_evidence_links,
    )

    workbook = load_workbook(path_text, read_only=True, data_only=False)
    try:
        required = {"Report_Info", *_SOURCE_CONTRACT_SHEETS}
        missing = sorted(required - set(workbook.sheetnames))
        if missing:
            return {
                "ok": False,
                "errors": ["Canonical evidence workbook is missing: " + ", ".join(missing)],
            }
        frames: dict[str, Any] = {}
        formula_cells = 0
        for sheet_name in required:
            worksheet = workbook[sheet_name]
            iterator = worksheet.iter_rows()
            try:
                header_cells = next(iterator)
            except StopIteration:
                frames[sheet_name] = pd.DataFrame()
                continue
            headers = [_text(cell.value, 300) for cell in header_cells]
            rows = []
            for row in iterator:
                formula_cells += sum(
                    1 for cell in row if getattr(cell, "data_type", "") == "f"
                )
                rows.append([cell.value for cell in row])
            frames[sheet_name] = pd.DataFrame(rows, columns=headers)
        errors: list[str] = []
        if formula_cells:
            errors.append(
                f"Canonical evidence workbook contains {formula_cells} formula cell(s)."
            )
        report_info = frames["Report_Info"]
        info = {
            _text(row.get("Item"), 300): _text(row.get("Value"), 4_000)
            for _, row in report_info.iterrows()
            if _text(row.get("Item"), 300)
        }
        workbook_fingerprint = _text(info.get("Fact_Contract_SHA256"), 128)
        if expected_fingerprint and workbook_fingerprint != expected_fingerprint:
            errors.append("Canonical evidence workbook fingerprint does not match the report.")
        for sheet_name in _SOURCE_CONTRACT_SHEETS:
            expected_digest = _text(info.get(f"Sheet_SHA256:{sheet_name}"), 128)
            actual_digest = _frame_content_digest(
                frames[sheet_name],
                sheet_name=sheet_name,
                already_exported=True,
            )
            if not expected_digest or actual_digest != expected_digest:
                errors.append(f"Canonical evidence sheet digest failed for {sheet_name}.")
        evidence_result = validate_evidence_links(frames, already_exported=True)
        errors.extend(evidence_result.get("errors") or [])
        return {
            "ok": not errors,
            "errors": errors,
            "fact_fingerprint": workbook_fingerprint,
            "formula_cells": formula_cells,
            "evidence": evidence_result,
        }
    finally:
        workbook.close()


def verify_canonical_evidence_workbook(
    path: str | Path,
    *,
    expected_fingerprint: object = "",
) -> dict[str, Any]:
    """Fail closed unless every serialized canonical sheet digest reconciles."""

    resolved = Path(path).expanduser().resolve()
    if resolved.suffix.casefold() != ".xlsx" or not resolved.is_file():
        raise ValueError("A readable Source Data .xlsx file is required.")
    stat = resolved.stat()
    result = copy.deepcopy(
        _cached_evidence_workbook_integrity(
            str(resolved),
            stat.st_mtime_ns,
            stat.st_size,
            _text(expected_fingerprint, 128),
        )
    )
    if not result.get("ok"):
        raise EvidenceIntegrityError(
            "; ".join(result.get("errors") or ["Canonical evidence integrity failed."])
        )
    return result


def load_workbook_evidence(
    path: str | Path,
    evidence_key: object,
    *,
    expected_fingerprint: object = "",
    limit: int = 25,
) -> dict[str, Any]:
    """Resolve a canonical Evidence_Link to bounded, exact workbook rows."""

    key = _text(evidence_key, 300)
    if not _EVIDENCE_KEY_RE.fullmatch(key):
        raise ValueError("A valid evidence key is required.")
    try:
        bounded_limit = max(1, min(int(limit), MAX_EVIDENCE_RECORDS))
    except (TypeError, ValueError) as exc:
        raise ValueError("Evidence limit must be a whole number.") from exc
    resolved = Path(path).expanduser().resolve()
    if resolved.suffix.casefold() != ".xlsx" or not resolved.is_file():
        raise ValueError("A readable Source Data .xlsx file is required.")
    if resolved.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ValueError("Source Data workbook exceeds the evidence size limit.")

    verify_canonical_evidence_workbook(
        resolved,
        expected_fingerprint=expected_fingerprint,
    )

    from decision_report_delivery import _evidence_row_fingerprint  # noqa: PLC0415

    workbook = load_workbook(resolved, read_only=True, data_only=False)
    try:
        required = {"Report_Info", "Evidence_Links"}
        if not required.issubset(workbook.sheetnames):
            raise EvidenceIntegrityError(
                "This report does not contain the canonical Evidence_Links contract."
            )
        info_records, _, info_formulas = _sheet_records(workbook["Report_Info"], limit=2_000)
        links, _, link_formulas = _sheet_records(workbook["Evidence_Links"], limit=MAX_SHEET_ROWS)
        if info_formulas or link_formulas:
            raise EvidenceIntegrityError("Evidence metadata contains formulas and cannot be verified.")
        info = {
            _text(record.get("Item"), 300): _text(record.get("Value"), 4_000)
            for record in info_records
            if _text(record.get("Item"))
        }
        workbook_fingerprint = _text(info.get("Fact_Contract_SHA256"), 128)
        expected = _text(expected_fingerprint, 128)
        if expected and workbook_fingerprint != expected:
            raise EvidenceIntegrityError(
                "The report fingerprint does not match its paired Source Data workbook."
            )
        selected = [record for record in links if _text(record.get("Evidence_Key"), 300) == key]
        if not selected:
            raise EvidenceNotFoundError("Evidence was not found for this report item.")

        source_states = [
            _text(record.get("Source_State"), 80).casefold() or "unknown"
            for record in selected
        ]
        label = _text(selected[0].get("Display_Label"), 300) or key
        limitations: list[str] = []
        for state in sorted(set(source_states)):
            if state not in {"available", "zero"}:
                limitations.append(
                    f"Source state is {state}; the returned records may not represent complete current coverage."
                )
        if any(_number(record.get("Source_Row_Number")) is None for record in selected):
            roles = sorted({
                _text(record.get("Evidence_Role"), 120).replace("_", " ")
                for record in selected
                if _number(record.get("Source_Row_Number")) is None
            })
            limitations.append(
                "This item includes an explicit derivation state without a source row"
                + (f" ({', '.join(role for role in roles if role)})." if roles else ".")
            )

        locators: list[tuple[Mapping[str, Any], str, int]] = []
        seen_locators: set[tuple[str, int]] = set()
        for link in selected:
            row_number = _number(link.get("Source_Row_Number"))
            if row_number is None:
                continue
            source_sheet = _text(link.get("Source_Sheet"), 120)
            row_number_int = int(row_number)
            locator = (source_sheet, row_number_int)
            if locator in seen_locators:
                raise EvidenceIntegrityError("Evidence metadata contains a duplicate source-row locator.")
            seen_locators.add(locator)
            if source_sheet not in workbook.sheetnames or row_number_int < 2:
                raise EvidenceIntegrityError("Evidence references a missing source row.")
            locators.append((link, source_sheet, row_number_int))

        total_records = len(locators)
        if total_records > bounded_limit:
            limitations.append(
                f"Showing the first {bounded_limit} of {total_records} exact supporting records."
            )
        records: list[dict[str, Any]] = []
        for link, source_sheet, row_number in locators[:bounded_limit]:
            worksheet = workbook[source_sheet]
            header_cells = next(worksheet.iter_rows(min_row=1, max_row=1), ())
            headers = [_text(cell.value, 200) for cell in header_cells]
            row_cells = next(
                worksheet.iter_rows(min_row=row_number, max_row=row_number),
                (),
            )
            if not row_cells or any(getattr(cell, "data_type", "") == "f" for cell in row_cells):
                raise EvidenceIntegrityError("Evidence references a missing or formula-backed source row.")
            record = {
                header: cell.value
                for header, cell in zip(headers, row_cells)
                if header
            }
            expected_hash = _text(link.get("Source_Row_SHA256"), 128)
            if not expected_hash or _evidence_row_fingerprint(record) != expected_hash:
                raise EvidenceIntegrityError(
                    f"Evidence row integrity failed for {source_sheet}!{row_number}."
                )
            expected_record_id = _text(link.get("Record_ID"), 240)
            actual_record_id = _text(
                _first(record, ("Record_ID", "Metric_Key")),
                240,
            )
            if expected_record_id and actual_record_id != expected_record_id:
                raise EvidenceIntegrityError(
                    f"Evidence record identity failed for {source_sheet}!{row_number}."
                )
            source_record_url = _verified_source_record_url(
                source_sheet=source_sheet,
                record_id=actual_record_id or expected_record_id,
                evidence_link=link,
                source_record=record,
            )

            customer = _text(
                _first(
                    record,
                    (
                        "Customer", "Customer Name", "Customer_Name", "BU_NAME",
                        "Account", "RELATED_CUSTOMER__C",
                    ),
                ),
                240,
            )
            title = _text(
                _first(
                    record,
                    (
                        "AdoptIQ_Title", "Title", "Subject", "SUBJECT_C", "NAME",
                        "Problem Description", "Description",
                    ),
                ),
                360,
            )
            status = _text(
                _first(record, ("AdoptIQ_Status_Bucket", "Status", "STATUS_C", "Case Status")),
                120,
            )
            date_value = _text(
                _first(
                    record,
                    (
                        "AdoptIQ_Due_Date", "Due Date", "DUE_DATE_C", "Date/Time Opened",
                        "Open Date", "OPEN_DATE_C", "Pulse Date", "CREATED_DATE_C",
                    ),
                ),
                120,
            )
            owner = _text(
                _first(
                    record,
                    (
                        "Owner", "Next Action Owner", "NEXT_ACTION_OWNER_C", "OWNER_NAME_C",
                        "CSSM", "Team_Member",
                    ),
                ),
                240,
            )
            summary_parts = []
            for field in (
                "Next Action", "NEXT_ACTION_C", "Description", "Problem Description",
                "Component_Details_JSON", "Risk_Band", "Risk_Score_0_100",
            ):
                value = _text(record.get(field), 500)
                if value:
                    summary_parts.append(f"{_public_label(field)}: {value}")
                if len(summary_parts) >= 3:
                    break
            # A report metric and a source-row locator are independent parts
            # of the Evidence_Links contract.  Preserve the one exact public
            # source field that can prove a non-count scalar so report-bound
            # Ask AI can reconcile the two instead of trusting the group
            # metric alone.  Keep this deliberately allowlisted: count
            # summaries still require their row-level count contract and do
            # not become assertable merely because a summary cell repeats the
            # number.
            metric_value_evidence: dict[str, Any] | None = None
            metric_suffix = key.rsplit(".", 1)[-1].casefold()
            scalar_fields = {
                "risk_score": ("Risk_Score_0_100", "Risk Score 0 100"),
            }.get(metric_suffix, ())
            scalar_matches = [
                (field, _number(record.get(field)))
                for field in scalar_fields
                if field in record and _number(record.get(field)) is not None
            ]
            if len(scalar_matches) == 1:
                metric_field, exact_metric_value = scalar_matches[0]
                metric_value_evidence = {
                    "field": metric_field,
                    "value": exact_metric_value,
                    "source_sheet": source_sheet,
                    "source_row_number": row_number,
                }
            records.append(
                {
                    "source_sheet": source_sheet,
                    "source_row_number": row_number,
                    "record_id": actual_record_id or expected_record_id,
                    "record_id_quality": _text(link.get("Record_ID_Data_Quality"), 240),
                    "source_record_url": source_record_url,
                    "customer": customer,
                    "title": title,
                    "status": status,
                    "date": date_value,
                    "owner": owner,
                    "summary": " | ".join(summary_parts),
                    "metric_value_evidence": metric_value_evidence,
                }
            )

        scope_type = _text(info.get("Scope_Type") or "team", 80).casefold()
        scope_value = _text(info.get("Scope_Value"), 300)
        manager = _text(info.get("Manager"), 240)
        scope_label = scope_value or (f"{manager} team" if manager else scope_type.title())
        return {
            "evidence_key": key,
            "label": label,
            "evidence_type": _text(selected[0].get("Evidence_Type"), 80),
            "metric_value": _number(selected[0].get("Metric_Value")),
            "unit": _text(selected[0].get("Unit"), 80),
            "evidence_roles": sorted({
                _text(record.get("Evidence_Role"), 120)
                for record in selected
                if _text(record.get("Evidence_Role"), 120)
            }),
            "source_state": _aggregate_source_state(source_states),
            "total_records": total_records,
            "records": records,
            "limitations": limitations,
            "scope_label": scope_label,
            "scope_type": scope_type,
            "scope_value": scope_value,
            "data_as_of_utc": _text(info.get("Data_As_Of_UTC"), 120),
            "data_as_of_state": (
                _text(info.get("Data_As_Of_State"), 80).casefold()
                or ("available" if _text(info.get("Data_As_Of_UTC"), 120) else "unavailable")
            ),
            "retrieval_attempted_at_utc": _text(
                info.get("Retrieval_Attempted_At_UTC"), 120
            ),
            "fact_fingerprint": workbook_fingerprint,
            "truncated": total_records > bounded_limit,
        }
    finally:
        workbook.close()


_EVIDENCE_QUESTION_STOP_WORDS = {
    "about", "after", "again", "also", "and", "are", "can", "could",
    "data", "does", "for", "from", "give", "have", "how", "into", "its",
    "list", "more", "report", "show", "source", "tell", "that", "the",
    "their", "these", "this", "those", "what", "when", "where", "which",
    "with", "would",
}


def _evidence_question_terms(value: object) -> set[str]:
    """Return conservative terms used only to rank already-authorized evidence."""

    return {
        token
        for token in re.findall(r"[a-z0-9]{2,}", _text(value, 4_000).casefold())
        if token not in _EVIDENCE_QUESTION_STOP_WORDS
    }


def select_report_bound_evidence(
    path: str | Path,
    snapshot: Mapping[str, Any],
    question: object,
    *,
    preferred_evidence_key: object = "",
    max_keys: int = 12,
    max_records: int = 60,
) -> dict[str, Any]:
    """Resolve a bounded, question-ranked set of exact report evidence rows.

    Selection is performed exclusively over the server-owned manifest from the
    already-verified workbook snapshot. Every returned record is then resolved
    again through ``load_workbook_evidence`` so its row identity and SHA-256 are
    checked immediately before it becomes report-bound Ask AI context.
    """

    if snapshot.get("canonical_snapshot") is not True:
        raise EvidenceIntegrityError("Report-bound evidence requires a canonical snapshot.")
    if snapshot.get("evidence_contract") != "canonical-evidence-links/v1":
        raise EvidenceIntegrityError("Report-bound evidence requires canonical Evidence_Links.")
    manifest = [
        item for item in (snapshot.get("evidence_manifest") or [])
        if isinstance(item, Mapping) and _text(item.get("evidence_key"), 300)
    ]
    if not manifest:
        raise EvidenceIntegrityError("The report evidence manifest is empty.")
    try:
        key_limit = max(1, min(int(max_keys), 24))
        record_limit = max(1, min(int(max_records), 100))
    except (TypeError, ValueError) as exc:
        raise ValueError("Evidence selection limits must be whole numbers.") from exc

    preferred = _text(preferred_evidence_key, 300)
    manifest_keys = {_text(item.get("evidence_key"), 300) for item in manifest}
    if preferred and preferred not in manifest_keys:
        raise EvidenceNotFoundError("The requested evidence key is not part of this report.")
    question_text = _text(question, 4_000).casefold()
    terms = _evidence_question_terms(question_text)
    wants_records = any(
        phrase in question_text
        for phrase in ("actual record", "records", "details", "deep dive", "examples", "which ")
    )
    intent_types: set[str] = set()
    wants_next_action = any(
        term in question_text
        for term in (
            "next action", "do next", "should ", "recommend", "priority",
            "prioritize", "owner", "action", "plan", "due", "overdue",
        )
    )
    if wants_next_action:
        intent_types.update({"action_plan", "action", "recommendation"})
    if any(term in question_text for term in ("customer", "account", "risk", "attention")):
        intent_types.update({"account", "recommendation", "summary"})
    if any(term in question_text for term in ("chart", "trend", "change", "compare")):
        intent_types.update({"chart", "chart_point", "trend"})
    if any(
        term in question_text
        for term in ("metric", "count", "how many", "total", "rate", "score", "value")
    ):
        intent_types.update({"metric", "kpi", "summary", "reported_fact"})

    status_key_hints: tuple[str, ...] = ()
    if "overdue" in question_text:
        status_key_hints = ("action_plans_overdue",)
    elif "due soon" in question_text:
        status_key_hints = ("action_plans_due_soon",)
    elif any(term in question_text for term in ("blocked", "on hold")):
        status_key_hints = ("action_plans_blocked",)
    elif any(term in question_text for term in ("completed", "complete")):
        status_key_hints = ("action_plans_completed",)
    elif "open" in question_text and "action" in question_text:
        status_key_hints = ("action_plans_open",)
    wants_risk_score = "risk score" in question_text or (
        "risk" in question_text and "score" in question_text
    )

    def rank(item: Mapping[str, Any]) -> tuple[int, int, str]:
        evidence_key = _text(item.get("evidence_key"), 300)
        evidence_type = _key(item.get("evidence_type"))
        searchable = " ".join((
            evidence_key,
            _text(item.get("label"), 300),
            evidence_type,
            " ".join(_text(source, 120) for source in item.get("sources") or []),
        ))
        overlap = len(terms & _evidence_question_terms(searchable))
        score = overlap * 20
        if preferred and evidence_key == preferred:
            score += 10_000
        if evidence_type in intent_types:
            score += 25
        if wants_risk_score and evidence_key.endswith(".risk_score"):
            score += 120
        if status_key_hints and any(hint in evidence_key for hint in status_key_hints):
            score += 120
        if wants_next_action and evidence_type == "recommendation":
            score += 55
        elif wants_next_action and evidence_type == "action_plan":
            score += 40
        elif wants_next_action and evidence_key in {
            "kpi.action_plans_overdue", "kpi.action_plans_due_soon",
            "kpi.action_plans_open",
        }:
            score += 30
        if wants_records and int(item.get("total_records") or 0) > 0:
            score += 8
        if evidence_key.startswith("kpi."):
            score += 2
        return (-score, -int(item.get("total_records") or 0), evidence_key)

    ranked = sorted(manifest, key=rank)
    # With no meaningful lexical match, retain the small canonical decision KPI
    # surface instead of arbitrarily selecting a high-volume raw table.
    if terms and all(rank(item)[0] == 0 for item in ranked):
        ranked = sorted(
            manifest,
            key=lambda item: (
                not _text(item.get("evidence_key"), 300).startswith("kpi."),
                _text(item.get("evidence_key"), 300),
            ),
        )
    selected = ranked[:key_limit]
    groups: list[dict[str, Any]] = []
    remaining = record_limit
    total_available_records = 0
    for selection_rank, item in enumerate(selected, 1):
        key = _text(item.get("evidence_key"), 300)
        per_key_limit = max(1, min(remaining or 1, 15))
        group = load_workbook_evidence(
            path,
            key,
            expected_fingerprint=snapshot.get("fact_fingerprint"),
            limit=per_key_limit,
        )
        group["question_relevance"] = max(0, -rank(item)[0])
        group["selection_rank"] = selection_rank
        total_available_records += int(group.get("total_records") or 0)
        if remaining <= 0:
            group["records"] = []
            group["truncated"] = bool(group.get("total_records"))
            group.setdefault("limitations", []).append(
                "The report-bound evidence record limit was reached before this group."
            )
        else:
            remaining -= len(group.get("records") or [])
        groups.append(group)
    returned_records = sum(len(group.get("records") or []) for group in groups)
    return {
        "schema": "report-bound-evidence/v1",
        "evidence_contract": "canonical-evidence-links/v1",
        "fact_fingerprint": _text(snapshot.get("fact_fingerprint"), 128),
        "data_as_of_utc": _text(snapshot.get("data_as_of_utc"), 120),
        "selected_key_count": len(groups),
        "manifest_key_count": len(manifest),
        "returned_record_count": returned_records,
        "selected_record_count": total_available_records,
        "truncated": len(manifest) > len(groups) or total_available_records > returned_records,
        "groups": groups,
    }


def assess_customer_share_readiness(
    snapshot: Mapping[str, Any],
    offline_receipt_status: OfflineValidationReceiptStatus | None = None,
) -> dict[str, Any]:
    """Return a fail-closed, non-actionable customer-share status.

    This function does not publish or export anything. A signed, persisted
    *offline fixture* receipt can make fixture-validation evidence visible,
    but it is a distinct contract from live validation or owner approval and
    can never turn a release/customer-share stamp true.
    """

    reasons: list[str] = []
    receipt_status = (
        offline_receipt_status
        if isinstance(offline_receipt_status, OfflineValidationReceiptStatus)
        else OfflineValidationReceiptStatus(state="missing")
    )
    receipt_summary = receipt_status.public_summary()

    def block(condition: bool, reason: str) -> None:
        if condition and reason not in reasons:
            reasons.append(reason)

    status = _text(snapshot.get("status"), 80).casefold()
    scope_type = _text(snapshot.get("scope_type"), 80).casefold()
    fingerprint = _text(snapshot.get("fact_fingerprint"), 128).casefold()
    insights = [
        item
        for item in (snapshot.get("decision_insights") or [])
        if isinstance(item, Mapping)
    ]
    raw_source_states = snapshot.get("source_states")
    raw_source_states = raw_source_states if isinstance(raw_source_states, Mapping) else {}
    source_states = {
        _text(source, 240): _text(state, 80).casefold() or "unknown"
        for source, state in raw_source_states.items()
    }
    degraded_sources = sorted(
        source for source, state in source_states.items()
        if state not in _COMPARABLE_SOURCE_STATES
    )

    block(not _status_is_completed(status), "The report is not complete.")
    block(
        scope_type not in {"customer", "subscription"},
        "Customer sharing requires an exact customer or subscription scope.",
    )
    block(
        snapshot.get("canonical_snapshot") is not True,
        "A current canonical Source Data workbook is required.",
    )
    block(
        re.fullmatch(r"[0-9a-f]{64}", fingerprint) is None,
        "The report fact fingerprint is missing or invalid.",
    )
    block(
        snapshot.get("persisted_workbook_hash_verified") is not True,
        "The persisted workbook hash has not been verified against the current bytes.",
    )
    block(
        snapshot.get("evidence_integrity_verified") is not True,
        "The canonical evidence and sheet digests have not been fully verified.",
    )
    block(
        snapshot.get("evidence_available") is not True,
        "Record-level evidence is unavailable.",
    )
    block(
        _number(snapshot.get("formula_cells")) != 0,
        "The workbook contains formulas and cannot be treated as frozen evidence.",
    )
    block(
        _text(snapshot.get("data_as_of_state"), 80).casefold()
        not in _COMPARABLE_SOURCE_STATES,
        "The report data-as-of state is not complete enough for customer use.",
    )
    block(not source_states, "The canonical required-source inventory is missing.")
    block(bool(degraded_sources), "One or more report sources are incomplete or unavailable.")
    block(not insights, "No supported evidence-backed executive insights are available.")
    block(
        any(
            _text(item.get("insight_key"), 300) not in _SUPPORTED_DECISION_INSIGHTS
            or _text(item.get("evidence_key"), 300)
            != _text(item.get("insight_key"), 300)
            or not _frozen_text(item.get("claim"), limit=4_000)
            or not isinstance(item.get("source_sheets"), Sequence)
            or isinstance(item.get("source_sheets"), (str, bytes))
            or not item.get("source_sheets")
            or _text(item.get("source_state"), 80).casefold()
            not in _COMPARABLE_SOURCE_STATES
            or (
                _text(item.get("source_state"), 80).casefold() == "available"
                and int(_number(item.get("evidence_count")) or 0) <= 0
            )
            or any(
                source_states.get(_text(sheet, 240), "unknown")
                not in _COMPARABLE_SOURCE_STATES
                for sheet in (item.get("source_sheets") or [])
            )
            for item in insights
        ),
        "At least one displayed insight lacks complete canonical evidence.",
    )
    block(
        snapshot.get("decision_insight_integrity") != "verified_shape",
        "The executive-insight projection did not pass its uniqueness check.",
    )
    block(bool(snapshot.get("source_warnings")), "The report still contains source warnings.")
    preview_evidence_ready = not reasons
    if receipt_status.verified:
        block(
            True,
            "A trusted offline-fixture validation receipt is persisted and verified; "
            "fixture validation cannot establish live accuracy or customer-share permission.",
        )
    else:
        block(
            True,
            "Customer-share release receipts are not enabled in this internal preview; "
            "a missing or invalid offline receipt also cannot establish live accuracy.",
        )
    return {
        "state": "internal_preview",
        "ready": False,
        "customer_shareable": False,
        "preview_evidence_ready": preview_evidence_ready,
        "offline_validation_receipt_state": receipt_summary["state"],
        "fixture_validation_performed": receipt_summary[
            "fixture_validation_performed"
        ],
        "fixture_validation_passed": receipt_summary["fixture_validation_passed"],
        "live_validation_performed": False,
        "production_accuracy_claimed": False,
        "release_ready": False,
        "manual_source_reconciliation_complete": False,
        "owner_customer_share_approved": False,
        "reasons": reasons,
    }


def apply_offline_validation_receipt_status(
    snapshot: Mapping[str, Any],
    receipt_status: OfflineValidationReceiptStatus | None,
) -> dict[str, Any]:
    """Attach only a redacted receipt summary and recompute closed readiness."""

    result = dict(snapshot)
    status = (
        receipt_status
        if isinstance(receipt_status, OfflineValidationReceiptStatus)
        else OfflineValidationReceiptStatus(state="missing")
    )
    if status.verified:
        # Fixture validation is intentionally distinct from customer-share
        # readiness. The shipped guarded fixture declares partial source states
        # because it is not live; those exact states are signed and projection-
        # bound. We still recompute the artifact/evidence gates from current
        # bytes and reject failed, unavailable, stale, unknown, or unbounded
        # fixture state. Customer-share readiness below remains hard false.
        fixture_gate = assess_offline_fixture_validation_readiness(result)
        if fixture_gate.get("ready") is not True:
            status = OfflineValidationReceiptStatus(state="artifact_mismatch")
    result["offline_validation_receipt"] = status.public_summary()
    result["customer_share_readiness"] = assess_customer_share_readiness(
        result,
        status,
    )
    for prohibited_flag in (
        "live_validation_attempted",
        "live_validation_performed",
        "live_validation_passed",
        "production_accuracy_claimed",
        "manual_source_reconciliation_complete",
        "release_ready",
        "customer_shareable",
        "owner_customer_share_approved",
        "ready_for_live_cisco",
    ):
        result[prohibited_flag] = False
    # Neither a status/client dictionary nor the signed envelope is public API.
    result.pop("customer_share_validation_receipt", None)
    result.pop("_trusted_offline_validation_receipt", None)
    return result


def assess_offline_fixture_validation_readiness(
    snapshot: Mapping[str, Any],
) -> dict[str, Any]:
    """Recompute current-byte gates appropriate to guarded fixture evidence."""

    reasons: list[str] = []

    def block(condition: bool, reason: str) -> None:
        if condition and reason not in reasons:
            reasons.append(reason)

    status = _text(snapshot.get("status"), 80).casefold()
    scope_type = _text(snapshot.get("scope_type"), 80).casefold()
    fingerprint = _text(snapshot.get("fact_fingerprint"), 128).casefold()
    source_states = {
        _text(source, 240): _text(state, 80).casefold() or "unknown"
        for source, state in (snapshot.get("source_states") or {}).items()
    } if isinstance(snapshot.get("source_states"), Mapping) else {}
    disallowed_source_states = sorted(
        source
        for source, state in source_states.items()
        if state not in {"available", "zero", "partial"}
    )
    disallowed_warnings = [
        _text(warning, 1_000)
        for warning in (snapshot.get("source_warnings") or [])
        if _text(warning, 1_000) != _OFFLINE_FIXTURE_PROVENANCE_WARNING
    ]

    block(not _status_is_completed(status), "The fixture report is not complete.")
    block(
        scope_type not in {"customer", "subscription"},
        "The fixture receipt requires an exact customer or subscription scope.",
    )
    block(
        snapshot.get("canonical_snapshot") is not True,
        "A canonical fixture workbook is required.",
    )
    block(
        re.fullmatch(r"[0-9a-f]{64}", fingerprint) is None,
        "The fixture fact fingerprint is missing or invalid.",
    )
    block(
        snapshot.get("persisted_workbook_hash_verified") is not True,
        "The current fixture workbook bytes do not match report history.",
    )
    block(
        snapshot.get("evidence_integrity_verified") is not True,
        "Canonical fixture evidence digests did not reconcile.",
    )
    block(
        snapshot.get("evidence_available") is not True,
        "Canonical fixture evidence links are unavailable.",
    )
    block(
        _number(snapshot.get("formula_cells")) != 0,
        "The fixture workbook is not formula-free.",
    )
    block(
        _text(snapshot.get("data_as_of_state"), 80).casefold()
        not in {"available", "zero"},
        "The fixture data-as-of state is not deterministic.",
    )
    block(not source_states, "The fixture source-state inventory is missing.")
    block(
        bool(disallowed_source_states),
        "The fixture contains failed, unavailable, stale, or unknown sources.",
    )
    block(bool(disallowed_warnings), "The fixture contains a non-provenance warning.")
    block(
        snapshot.get("decision_insight_integrity") != "verified_shape",
        "The fixture insight projection shape did not reconcile.",
    )
    return {"ready": not reasons, "reasons": reasons}


def snapshot_from_status(status: Mapping[str, Any], workbook_snapshot: Mapping[str, Any] | None = None) -> dict[str, Any]:
    """Merge public run status with its canonical workbook snapshot."""

    snapshot = dict(workbook_snapshot or {})
    workbook_scope_value = _text(snapshot.get("scope_value"), 300)
    # Status is the server-canonical request envelope.  Workbook Scope_Value is
    # intentionally presentation-friendly for some report writers (for
    # example, a Leader member's display name), so it must not replace the
    # authorization identifier used by report-bound Ask AI.  Subscription
    # reports historically omitted Scope_Type and rendered a customer label in
    # Scope_Value; the persisted subscription_id is the only canonical value.
    report_type = _text(
        status.get("report_type") or snapshot.get("report_type"), 80
    ).casefold()
    scope_type = _text(
        status.get("scope_type")
        or (
            "subscription"
            if status.get("subscription_id")
            else snapshot.get("scope_type")
            or (
                "subscription"
                if report_type == "subscription"
                else "customer" if status.get("customer_name") else "team"
            )
        ),
        80,
    ).casefold()
    if scope_type == "subscription":
        scope_value = _text(
            status.get("subscription_id")
            or status.get("scope_value")
            or snapshot.get("scope_value"),
            300,
        )
    elif scope_type == "team":
        # A team scope has no authorization selector value.  Canonical Word/
        # workbook metadata may carry a presentation label such as
        # ``Manager One team`` in Scope_Value, but returning that label as the
        # immutable Ask AI selector is inconsistent with the scoped pipeline,
        # which correctly normalizes team scope to an empty value.  Preserve
        # the workbook value below as ``scope_label`` only.
        scope_value = ""
    else:
        scope_value = _text(
            status.get("scope_value")
            or snapshot.get("scope_value")
            or status.get("customer_name"),
            300,
        )
    scope_member = _text(
        status.get("scope_member")
        or snapshot.get("scope_member")
        or status.get("member_email"),
        320,
    )
    binding_available = True
    binding_notice = ""
    status_has_canonical_scope = bool(status.get("scope_type")) and (
        scope_type == "team"
        or bool(status.get("scope_value") or status.get("subscription_id"))
    )
    if report_type == "leader" and scope_type == "member" and not status_has_canonical_scope:
        email_match = _EMAIL_RE.search(scope_value)
        if email_match:
            scope_value = email_match.group("email").casefold()
            scope_member = scope_value
        else:
            binding_available = False
    elif report_type == "leader" and scope_type == "customer" and not status_has_canonical_scope:
        # Historical Leader customer workbooks may store a presentation label
        # such as "Customer (Member Name)". That composite label is not an
        # authorization identifier, so old rows fail closed instead of
        # widening or querying an unverified customer scope.
        if re.search(r"\s+\([^()]+\)\s*$", scope_value):
            binding_available = False
    if not binding_available:
        binding_notice = (
            "Ask AI is unavailable for this older scoped report because its "
            "canonical authorization scope was not retained. Generate a new report to ask scoped questions."
        )
    persisted_hash = _text(status.get("excel_hash"), 128).casefold()
    workbook_hash = _text(snapshot.get("workbook_sha256"), 128).casefold()
    persisted_workbook_hash_verified = (
        re.fullmatch(r"[0-9a-f]{64}", persisted_hash) is not None
        and re.fullmatch(r"[0-9a-f]{64}", workbook_hash) is not None
        and persisted_hash == workbook_hash
    )
    snapshot.update({
        "schema": WORKSPACE_SCHEMA,
        "analysis_id": _text(status.get("analysis_id"), 500),
        "status": _text(status.get("status") or "unknown", 80).casefold(),
        "report_type": report_type,
        "manager": _text(snapshot.get("manager") or status.get("manager"), 240),
        "technology": _text(snapshot.get("technology") or status.get("technology") or status.get("tech"), 240),
        "scope_type": scope_type,
        "scope_value": scope_value,
        "scope_member": scope_member,
        "days": _number(snapshot.get("days") if snapshot.get("days") is not None else status.get("days")),
        "data_as_of_utc": _text(snapshot.get("data_as_of_utc") or status.get("data_retrieved_at") or status.get("data_as_of_utc"), 120),
        "data_as_of_state": _text(
            snapshot.get("data_as_of_state")
            or status.get("data_as_of_state")
            or ("available" if snapshot.get("data_as_of_utc") or status.get("data_retrieved_at") else "unavailable"),
            80,
        ).casefold(),
        "retrieval_attempted_at_utc": _text(
            snapshot.get("retrieval_attempted_at_utc")
            or status.get("retrieval_attempted_at_utc"),
            120,
        ),
        "evaluation_as_of_utc": _text(
            snapshot.get("evaluation_as_of_utc")
            or status.get("evaluation_as_of_utc"),
            120,
        ),
        "started_at": _text(status.get("start_time") or status.get("started_at"), 120),
        "completed_at": _text(status.get("completion_time") or status.get("end_time"), 120),
        "word_available": bool(status.get("word_available") or status.get("word_report") or status.get("report_path")),
        "excel_available": bool(status.get("excel_available") or status.get("excel_report")),
        "ask_ai_binding_available": binding_available,
        "ask_ai_binding_notice": binding_notice,
        "persisted_workbook_hash_verified": persisted_workbook_hash_verified,
    })
    snapshot["scope_label"] = _text(status.get("scope_display"), 300) or workbook_scope_value or snapshot["scope_value"] or (
        f"{snapshot['manager']} team" if snapshot["manager"] else "Portfolio"
    )
    if snapshot["data_as_of_state"] not in {"available", "zero"}:
        snapshot.setdefault("source_states", {})["Data_Freshness"] = snapshot[
            "data_as_of_state"
        ]
    snapshot["source_limitations"] = [
        {"source": _public_label(source), "state": state}
        for source, state in sorted((snapshot.get("source_states") or {}).items())
        if state not in {"available", "zero"}
    ]
    snapshot["evidence_available"] = bool(snapshot.get("evidence_available"))
    snapshot["evidence_notice"] = _text(
        snapshot.get("evidence_notice")
        or (
            "Record-level evidence is unavailable because the paired Source Data workbook could not be verified."
            if not workbook_snapshot
            else "Generate a current canonical report to enable source-record drill-down."
        ),
        500,
    )
    source_warnings: list[str] = []
    seen_warnings: set[str] = set()
    for warning in (
        [_text(item, 1_000) for item in snapshot.get("source_warnings") or []]
        + _source_warnings(status.get("partial_data_warnings") or [])
    ):
        warning_key = warning.casefold()
        if warning and warning_key not in seen_warnings:
            source_warnings.append(warning)
            seen_warnings.add(warning_key)
    snapshot["source_warnings"] = source_warnings
    if binding_notice:
        snapshot["source_warnings"].append(binding_notice)
    snapshot["decision_metrics"] = [
        item for item in snapshot.get("metrics", [])
        if item.get("metric_key") in {
            "kpi.customers", "kpi.team_members", "kpi.action_plans_total", "kpi.action_plans_open",
            "kpi.action_plans_overdue", "kpi.action_plans_due_soon",
            "kpi.adoption_barriers", "kpi.tac_cases", "kpi.bems",
            "kpi.high_risk_customers",
        }
    ]
    projected_insights = [
        dict(item)
        for item in snapshot.get("decision_insights", [])
        if isinstance(item, Mapping)
        and _text(item.get("insight_key"), 300) in _SUPPORTED_DECISION_INSIGHTS
    ]
    insight_integrity = snapshot.get("decision_insight_integrity")
    if insight_integrity not in {"verified_shape", "legacy_unavailable"}:
        snapshot["decision_insights"] = []
        snapshot["source_warnings"].append(
            "Evidence-backed insights are withheld because their canonical lineage and evidence binding did not reconcile."
        )
    elif projected_insights and snapshot.get("persisted_workbook_hash_verified") is not True:
        snapshot["decision_insights"] = []
        snapshot["decision_insight_integrity"] = "workbook_hash_unverified"
        snapshot["source_warnings"].append(
            "Evidence-backed insights are withheld because the current workbook bytes are not bound to the persisted report hash."
        )
    else:
        snapshot["decision_insights"] = projected_insights
    snapshot["top_action_plans"] = list(snapshot.get("action_plans") or [])[:20]
    snapshot["top_accounts"] = list(snapshot.get("accounts") or [])[:12]
    return apply_offline_validation_receipt_status(snapshot, None)


def _ap_index(snapshot: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    return {
        _text(item.get("record_id")).casefold(): item
        for item in snapshot.get("action_plans", [])
        if isinstance(item, Mapping) and _text(item.get("record_id"))
    }


def _is_completed_status(value: object) -> bool:
    return _status_is_completed(value)


def _comparison_source_state(snapshot: Mapping[str, Any], item: Mapping[str, Any]) -> str:
    """Return the least-comparable state declared for one metric's lineage."""

    declared_state = _text(item.get("source_state"), 80).casefold() or "unknown"
    source_states = {
        _key(source): _text(state, 80).casefold() or "unknown"
        for source, state in (snapshot.get("source_states") or {}).items()
    }
    lineage_states = [
        source_states[_key(source)]
        for source in _text(item.get("source_sheet"), 300).split(";")
        if _key(source) in source_states
    ]
    if declared_state != "unknown":
        lineage_states.append(declared_state)
    return _aggregate_source_state(lineage_states) if lineage_states else declared_state


def _action_plan_source_state(snapshot: Mapping[str, Any]) -> str:
    """Resolve Action Plan coverage without inferring availability from row presence."""

    states = [
        _text(state, 80).casefold() or "unknown"
        for source, state in (snapshot.get("source_states") or {}).items()
        if _key(source) in {"action_plans", "customer_action_plans"}
    ]
    return _aggregate_source_state(states) if states else "unknown"


def compare_snapshots(before: Mapping[str, Any], after: Mapping[str, Any]) -> dict[str, Any]:
    """Compare two canonical report snapshots with source-change separation."""

    comparison_fields = ("manager", "technology", "scope_type", "scope_value")
    incompatible_fields = [
        field
        for field in comparison_fields
        if _text(before.get(field)).casefold() != _text(after.get(field)).casefold()
    ]
    if _number(before.get("days")) != _number(after.get("days")):
        incompatible_fields.append("days")
    before_report_type = _text(before.get("report_type"), 80).casefold()
    after_report_type = _text(after.get("report_type"), 80).casefold()
    if before_report_type and after_report_type and before_report_type != after_report_type:
        incompatible_fields.append("report_type")
    same_scope = not incompatible_fields
    if not same_scope:
        return {
            "schema": WORKSPACE_SCHEMA,
            "same_scope": False,
            "comparison_state": "incompatible_scope",
            "incompatible_fields": incompatible_fields,
            "before_fingerprint": _text(before.get("fact_fingerprint"), 128),
            "after_fingerprint": _text(after.get("fact_fingerprint"), 128),
            "metric_changes": [],
            "action_plan_changes": [],
            "source_driven_metric_changes": [],
            "source_driven_action_plan_changes": [],
            "source_state_changes": [],
            "business_change_count": 0,
            "source_driven_change_count": 0,
            "source_change_count": 0,
            "caveats": [
                "The selected reports do not have the same report family, manager, "
                "technology, scope, and window. No business deltas were calculated."
            ],
        }

    before_metrics = {item["metric_key"]: item for item in before.get("metrics", []) if isinstance(item, Mapping) and item.get("metric_key")}
    after_metrics = {item["metric_key"]: item for item in after.get("metrics", []) if isinstance(item, Mapping) and item.get("metric_key")}
    metric_changes: list[dict[str, Any]] = []
    source_driven_metric_changes: list[dict[str, Any]] = []
    for metric_key in sorted(set(before_metrics) | set(after_metrics)):
        old = before_metrics.get(metric_key, {})
        new = after_metrics.get(metric_key, {})
        old_value = _number(old.get("value"))
        new_value = _number(new.get("value"))
        if old_value == new_value:
            continue
        before_source_state = _comparison_source_state(before, old)
        after_source_state = _comparison_source_state(after, new)
        old_source_sheet = _text(old.get("source_sheet"), 300)
        new_source_sheet = _text(new.get("source_sheet"), 300)
        change = {
            "metric_key": metric_key,
            "label": _text(new.get("label") or old.get("label") or metric_key, 240),
            "before": old_value,
            "after": new_value,
            "delta": (new_value - old_value) if old_value is not None and new_value is not None else None,
            "source_sheet": new_source_sheet or old_source_sheet,
            "before_source_state": before_source_state,
            "after_source_state": after_source_state,
        }
        comparable = (
            bool(old)
            and bool(new)
            and before_source_state in _COMPARABLE_SOURCE_STATES
            and after_source_state in _COMPARABLE_SOURCE_STATES
            and _key(old_source_sheet) == _key(new_source_sheet)
        )
        if comparable:
            metric_changes.append(change)
        else:
            change["classification"] = "source_availability"
            source_driven_metric_changes.append(change)

    before_states = dict(before.get("source_states") or {})
    after_states = dict(after.get("source_states") or {})
    source_changes = [
        {"source": source, "before": before_states.get(source, "unknown"), "after": after_states.get(source, "unknown")}
        for source in sorted(set(before_states) | set(after_states))
        if before_states.get(source, "unknown") != after_states.get(source, "unknown")
    ]

    old_aps = _ap_index(before)
    new_aps = _ap_index(after)
    detected_action_changes: list[dict[str, Any]] = []
    for record_key in sorted(set(old_aps) | set(new_aps)):
        old = old_aps.get(record_key)
        new = new_aps.get(record_key)
        if old is None and new is not None:
            detected_action_changes.append({"record_id": new["record_id"], "change": "new", "customer": new.get("customer"), "title": new.get("title")})
            continue
        if new is None and old is not None:
            detected_action_changes.append({"record_id": old["record_id"], "change": "absent", "customer": old.get("customer"), "title": old.get("title")})
            continue
        assert old is not None and new is not None
        if not _is_completed_status(old.get("status")) and _is_completed_status(new.get("status")):
            detected_action_changes.append({"record_id": new["record_id"], "change": "completed", "before": old.get("status"), "after": new.get("status")})
        elif _is_completed_status(old.get("status")) and not _is_completed_status(new.get("status")):
            detected_action_changes.append({"record_id": new["record_id"], "change": "reopened", "before": old.get("status"), "after": new.get("status")})
        if not old.get("is_overdue") and new.get("is_overdue"):
            detected_action_changes.append({"record_id": new["record_id"], "change": "became_overdue", "before": old.get("due_date"), "after": new.get("due_date")})
        for field, change_name in (("owner", "owner_changed"), ("due_date", "due_date_changed"), ("status", "status_changed")):
            if _text(old.get(field)) != _text(new.get(field)):
                detected_action_changes.append({"record_id": new["record_id"], "change": change_name, "before": old.get(field), "after": new.get(field)})

    before_action_state = _action_plan_source_state(before)
    after_action_state = _action_plan_source_state(after)
    action_sources_comparable = (
        before_action_state in _COMPARABLE_SOURCE_STATES
        and after_action_state in _COMPARABLE_SOURCE_STATES
    )
    if action_sources_comparable:
        action_changes = detected_action_changes
        source_driven_action_changes: list[dict[str, Any]] = []
    else:
        action_changes = []
        source_driven_action_changes = [
            {
                **change,
                "classification": "source_availability",
                "before_source_state": before_action_state,
                "after_source_state": after_action_state,
            }
            for change in detected_action_changes
        ]

    caveats = []
    if source_changes:
        caveats.append("Source coverage changed; non-comparable KPI and Action Plan differences are excluded from business movement.")
    if source_driven_metric_changes or source_driven_action_changes:
        caveats.append(
            f"Excluded {len(source_driven_metric_changes)} KPI and "
            f"{len(source_driven_action_changes)} Action Plan difference(s) because source coverage was not comparable."
        )
    if before.get("missing_action_plan_id_count") or after.get("missing_action_plan_id_count"):
        caveats.append("Action Plans without stable IDs are retained in each artifact but cannot be compared reliably across runs.")
    return {
        "schema": WORKSPACE_SCHEMA,
        "same_scope": same_scope,
        "comparison_state": "comparable",
        "incompatible_fields": [],
        "before_fingerprint": _text(before.get("fact_fingerprint"), 128),
        "after_fingerprint": _text(after.get("fact_fingerprint"), 128),
        "metric_changes": metric_changes[:MAX_PUBLIC_ITEMS],
        "action_plan_changes": action_changes[:MAX_PUBLIC_ITEMS],
        "source_driven_metric_changes": source_driven_metric_changes[:MAX_PUBLIC_ITEMS],
        "source_driven_action_plan_changes": source_driven_action_changes[:MAX_PUBLIC_ITEMS],
        "source_state_changes": source_changes,
        "business_change_count": len(metric_changes) + len(action_changes),
        "source_driven_change_count": len(source_driven_metric_changes) + len(source_driven_action_changes),
        "source_change_count": len(source_changes),
        "caveats": caveats,
    }


def normalize_history_record(record: Mapping[str, Any]) -> dict[str, Any]:
    report_type = _text(record.get("report_type"), 80).casefold()
    return {
        "analysis_id": _text(record.get("request_id") or record.get("analysis_id"), 500),
        "report_type": report_type,
        "report_label": REPORT_TYPES.get(report_type, {}).get("label", report_type.replace("_", " ").title() or "Report"),
        "manager": _text(record.get("manager"), 240),
        "technology": _text(record.get("technology") or record.get("tech"), 240),
        "customer": _text(record.get("customer_name") or record.get("scope_value"), 300),
        "scope_type": _text(record.get("scope_type") or ("customer" if record.get("customer_name") else "team"), 80).casefold(),
        "days": _number(record.get("days")),
        "status": _text(record.get("status"), 80).casefold(),
        "started_at": _text(record.get("start_time") or record.get("started_at"), 120),
        "completed_at": _text(record.get("end_time") or record.get("completion_time") or record.get("created_at"), 120),
        "word_available": bool(record.get("word_path") or record.get("word_report") or record.get("report_path")),
        "excel_available": bool(record.get("excel_path") or record.get("excel_report")),
        "is_completed": _text(record.get("status")).casefold() in _TERMINAL_STATUSES,
    }


def filter_history(
    records: Iterable[Mapping[str, Any]],
    *,
    manager: str = "",
    report_type: str = "",
    scope_type: str = "",
    customer: str = "",
    technology: str = "",
    status: str = "",
) -> list[dict[str, Any]]:
    """Normalize and filter history without exposing filesystem paths."""

    wanted = {
        "manager": _text(manager).casefold(),
        "report_type": _text(report_type).casefold(),
        "scope_type": _text(scope_type).casefold(),
        "customer": _text(customer).casefold(),
        "technology": _text(technology).casefold(),
        "status": _text(status).casefold(),
    }
    normalized = [normalize_history_record(record) for record in records if not record.get("_placeholder")]
    result = []
    for record in normalized:
        if wanted["manager"] and record["manager"].casefold() != wanted["manager"]:
            continue
        if wanted["report_type"] and record["report_type"] != wanted["report_type"]:
            continue
        if wanted["scope_type"] and record["scope_type"] != wanted["scope_type"]:
            continue
        if wanted["customer"] and wanted["customer"] not in record["customer"].casefold():
            continue
        if wanted["technology"] and record["technology"].casefold() != wanted["technology"]:
            continue
        if wanted["status"] and record["status"] != wanted["status"]:
            continue
        result.append(record)
    result.sort(key=lambda item: (item["completed_at"], item["analysis_id"]), reverse=True)
    return result


def ask_ai_binding(snapshot: Mapping[str, Any]) -> dict[str, Any]:
    """Return the immutable public context that binds Ask AI to one report."""

    binding_available = snapshot.get("ask_ai_binding_available") is not False
    return {
        "analysis_id": _text(snapshot.get("analysis_id"), 500),
        "report_type": _text(snapshot.get("report_type"), 80),
        "manager": _text(snapshot.get("manager"), 240),
        "technology": _text(snapshot.get("technology"), 240),
        "days": _number(snapshot.get("days")),
        "scope_type": _text(snapshot.get("scope_type") or "team", 80),
        "scope_value": _text(snapshot.get("scope_value"), 300) if binding_available else "",
        "scope_member": _text(snapshot.get("scope_member"), 320) if binding_available else "",
        "data_as_of_utc": _text(snapshot.get("data_as_of_utc"), 120),
        "data_as_of_state": _text(
            snapshot.get("data_as_of_state") or "unknown", 80
        ).casefold(),
        "retrieval_attempted_at_utc": _text(
            snapshot.get("retrieval_attempted_at_utc"), 120
        ),
        "fact_fingerprint": _text(snapshot.get("fact_fingerprint"), 128),
        "source_states": dict(sorted((snapshot.get("source_states") or {}).items())),
        "decision_metrics": list(snapshot.get("decision_metrics") or []),
        "evidence_available": bool(snapshot.get("evidence_available")),
        "evidence_contract": _text(snapshot.get("evidence_contract"), 120),
        "evidence_manifest": list(snapshot.get("evidence_manifest") or []),
        "binding_available": binding_available,
        "binding_notice": _text(snapshot.get("ask_ai_binding_notice"), 500),
        "live_validation_performed": None,
    }
